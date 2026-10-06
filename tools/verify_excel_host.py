"""Create a synthetic Office.js host fixture and verify the saved Excel result.

This tool does not automate Excel. Host interaction and save/reopen must happen
in real Excel; --verify checks the resulting file and bridge operation record.
The developer reference follows Microsoft's Office-Addin-Scripts sideload template.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.workbook.defined_name import DefinedName


def create_fixture(path: Path, manifest: Path) -> None:
    if path.exists():
        raise FileExistsError(path)
    ns = "{http://schemas.microsoft.com/office/appforoffice/1.1}"
    app = ET.fromstring(manifest.read_bytes())
    addin_id, version = app.findtext(ns + "Id"), app.findtext(ns + "Version")
    if not addin_id or not version:
        raise ValueError("Manifest needs Id and Version")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "经营分析"
    for row in (("management_account", "amount"), ("revenue", 120), ("cost", 70), ("expense", 20)):
        sheet.append(row)
    for cell, value in {"D1": "profit", "D2": 0, "F1": "经理备注", "F2": "保留人工判断",
                        "G1": "区域外公式", "G2": "=B2-B3-B4", "H1": "结果引用", "H2": "=D2*2"}.items():
        sheet[cell] = value
    workbook.defined_names.add(DefinedName("FoundryResult", attr_text="'经营分析'!$D$1:$D$2"))
    chart = BarChart()
    chart.title = "原工作簿图表"
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=4))
    sheet.add_chart(chart, "A7")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    with zipfile.ZipFile(path) as source:
        parts = {name: source.read(name) for name in source.namelist()}
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    ct_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    rels = ET.fromstring(parts["_rels/.rels"])
    ET.SubElement(rels, f"{{{rel_ns}}}Relationship", {
        "Id": "rIdFoundryTaskpane", "Target": "/xl/webextensions/taskpanes.xml",
        "Type": "http://schemas.microsoft.com/office/2011/relationships/webextensiontaskpanes"})
    parts["_rels/.rels"] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
    types = ET.fromstring(parts["[Content_Types].xml"])
    for name, kind in (("taskpanes", "webextensiontaskpanes"), ("webextension", "webextension")):
        ET.SubElement(types, f"{{{ct_ns}}}Override", {
            "PartName": f"/xl/webextensions/{name}.xml", "ContentType": f"application/vnd.ms-office.{kind}+xml"})
    parts["[Content_Types].xml"] = ET.tostring(types, encoding="utf-8", xml_declaration=True)
    parts["xl/webextensions/webextension.xml"] = (
        '<we:webextension xmlns:we="http://schemas.microsoft.com/office/webextensions/webextension/2010/11" '
        f'id="{{{addin_id}}}"><we:reference id="{addin_id}" version="{version}" '
        'store="developer" storeType="Registry"/><we:alternateReferences/><we:properties/>'
        '<we:bindings/><we:snapshot/></we:webextension>').encode()
    parts["xl/webextensions/taskpanes.xml"] = (
        '<wetp:taskpanes xmlns:wetp="http://schemas.microsoft.com/office/webextensions/taskpanes/2010/11" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<wetp:taskpane dockstate="right" visibility="1" width="350" row="1">'
        '<wetp:webextensionref r:id="rId1"/></wetp:taskpane></wetp:taskpanes>').encode()
    parts["xl/webextensions/_rels/taskpanes.xml.rels"] = (
        f'<Relationships xmlns="{rel_ns}"><Relationship Id="rId1" '
        'Type="http://schemas.microsoft.com/office/2011/relationships/webextension" '
        'Target="/xl/webextensions/webextension.xml"/></Relationships>').encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for name, data in parts.items():
            out.writestr(name, data)


def verify(path: Path, bridge_state: Path) -> dict:
    values = load_workbook(path, data_only=True)["经营分析"]
    formulas = load_workbook(path, data_only=False)["经营分析"]
    # Each session keeps its host events in a file of its own beside the state; the state holds those logged before.
    # Sessions that only logged host events leave no state behind; with neither, the bridge was never used.
    logs = sorted((bridge_state.parent / "telemetry").glob("*.json"))
    state = json.loads(bridge_state.read_text()) if bridge_state.exists() or not logs else {}
    events = list(state.get("telemetry", []))
    for log in logs:
        events += json.loads(log.read_text()).get("events", [])
    counts = {name: sum(event.get("event") == name for event in events) for name in (
        "input_imported", "metric_run", "output_previewed", "output_applied", "full_recalc", "workbook_saved")}
    checks = {"modified_input": values["B2"].value == 130,
              "refreshed_profit": values["D2"].value == 40,
              "excel_recalculation": values["G2"].value == 40 and values["H2"].value == 80,
              "outside_formulas_preserved": formulas["G2"].value == "=B2-B3-B4" and formulas["H2"].value == "=D2*2",
              "manager_note_preserved": values["F2"].value == "保留人工判断",
              "chart_preserved": len(formulas._charts) == 1,
              "two_real_imports": counts["input_imported"] >= 2,
              "two_real_metrics": counts["metric_run"] >= 2,
              "host_recalc_and_save": counts["full_recalc"] >= 1 and counts["workbook_saved"] >= 1,
              "conflict_keep_observed": any(event.get("event") == "output_applied" and event.get("details", {}).get("kept", 0) >= 1 for event in events)}
    return {"schema_version": 1, "status": "passed" if all(checks.values()) else "failed", "checks": checks,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "host_event_counts": counts,
            "hosts": [book.get("host") for book in state.get("workbooks", {}).values()],
            "boundary": "Requires independent native Excel open/save/reopen observation; file checks alone are insufficient."}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--fixture", type=Path)
    parser.add_argument("--manifest", type=Path, default=Path("excel-addin/manifest.xml"))
    parser.add_argument("--verify", type=Path)
    parser.add_argument("--bridge-state", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    if args.fixture:
        create_fixture(args.fixture, args.manifest)
        print(f"Created {args.fixture}")
        return 0
    if not args.verify or not args.bridge_state or not args.output:
        parser.error("Use --fixture, or --verify with --bridge-state and --output")
    result = verify(args.verify, args.bridge_state)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(result, ensure_ascii=False))
    return 0 if result["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
