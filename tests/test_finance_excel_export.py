from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from foundry.analysis import AnalysisStore
from foundry.errors import FoundryError
from foundry.task import TaskStore
from foundry.workbook_bindings import create_binding_contract, preview_output_patch, apply_output_patch


def test_report_normal_amount_numeric_and_large_amount_exact_text(tmp_path):
    service = AnalysisStore(TaskStore(tmp_path / "var").create(title="金额导出"))
    path = tmp_path / "data.json"
    path.write_text('[{"amount":"120.35","large":"123456789012345.67","code":"00123","text":"=1+1"}]')
    data = service.add_dataset(path, name="金额")
    revision = service.save_analysis({"question": "金额输出", "bindings": {"a": data["id"]}, "sql": "SELECT CAST(amount AS DECIMAL(38,2)) AS amount, CAST(large AS DECIMAL(38,2)) AS large, code, text FROM a"})
    run = service.run_analysis(revision["id"])
    draft = service.save_draft({"title": "结果", "run_ids": [run["id"]]})
    exported = service.export_draft(draft["id"], "xlsx")
    book = load_workbook(exported["path"])
    assert book['Result 1']['A2'].value == 120.35
    assert book['Result 1']['A2'].data_type == 'n'
    assert book['Result 1']['B2'].value == '123456789012345.67'
    assert book['Result 1']['C2'].value == '00123'
    assert book['Result 1']['D2'].value == '=1+1'
    assert book['Result 1']['D2'].data_type == 's'
    book.close()


def workbook(tmp_path):
    source = tmp_path / "input.xlsx"
    book = Workbook()
    book.active.title = 'Model'
    book.active.append(['amount', 'code'])
    book.active.append([10, '001'])
    book.save(source)
    book.close()
    contract = create_binding_contract(source, {"kind": "range", "sheet": "Model", "range": "A1:B2"}, [{"name": "amount", "policy": "decimal"}, {"name": "code", "policy": "identifier"}], role="output", writable=True)
    return source, contract


def test_stale_preview_refuses_write(tmp_path):
    source, contract = workbook(tmp_path)
    patch = preview_output_patch(source, contract, [{"cell": "A2", "field": "amount", "value": "15"}])
    book = load_workbook(source)
    book['Model']['A2'] = 22
    book.save(source)
    book.close()
    with pytest.raises(FoundryError, match="changed.*preview"):
        apply_output_patch(source, tmp_path / 'out.xlsx', patch, {'A2': 'apply'})


def test_independent_output_does_not_execute_formula_text(tmp_path):
    source, contract = workbook(tmp_path)
    patch = preview_output_patch(source, contract, [{"cell": "B2", "field": "code", "value": "=HYPERLINK(\"https://example.org\",\"x\")"}])
    patch['unsupported_features'] = ['drawings']
    result = apply_output_patch(source, tmp_path / 'out.xlsx', patch)
    book = load_workbook(Path(result['path']))
    assert book['Results']['C2'].data_type == 's'
    book.close()
