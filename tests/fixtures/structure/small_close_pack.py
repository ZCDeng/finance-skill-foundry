"""The small 2026-06 practice workbook the app shipped before the 华宸 example pack.

Three sheets, one planted tie-out break (the bad-debt column of the aging table is 50 万 over its
合计). Kept here because it is a quick, readable workbook for tests that need one break and nothing else.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

SAMPLE_DATASET = "示例集团_2026年06月_月结套表"


def build_sample_workbook(path: Path) -> Path:
    book = Workbook()
    cover = book.active
    cover.title = "00·封面"
    cover.append(["示例集团（虚构）"])
    cover.append(["示例集团  |  2026年06月  |  月结练习套表"])
    cover.append([])
    cover.append(["说明", "全部数字为合成数据，不对应任何真实企业。应收账龄表的坏账准备分项故意与合计相差 50 万元，供练习勾稽控制与差异处置。"])

    regions = book.create_sheet("01·区域收入")
    regions.append(["区域收入"])
    regions.append(["示例集团  |  2026年06月  |  单位：万元  |  币种：人民币"])
    regions.append([])
    regions.append(["区域", "本月收入", "本月成本", "毛利率"])
    for name, revenue, cost in (("华东", 1200, 950), ("华南", 860, 700), ("华北", 640, 530)):
        regions.append([name, revenue, cost, round(1 - cost / revenue, 4)])
    regions.append(["合计", 2700, 2180, round(1 - 2180 / 2700, 4)])

    aging = book.create_sheet("02·应收账龄")
    aging.append(["应收账款账龄与坏账准备"])
    aging.append(["示例集团  |  2026年06月  |  单位：万元  |  币种：人民币"])
    aging.append([])
    aging.append(["账龄", "账面余额", "坏账准备", "计提比例"])
    for bucket, balance, provision in (("1年以内", 3000, 150), ("1至2年", 1200, 180), ("2至3年", 500, 175), ("3年以上", 300, 300)):
        aging.append([bucket, balance, provision, round(provision / balance, 4)])
    # The deliberate break: the details provide 805, the total says 755.
    aging.append(["合计", 5000, 755, round(755 / 5000, 4)])
    book.save(path)
    return path
