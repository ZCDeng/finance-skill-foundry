"""U1: the structure proposal engine reads a real-shaped workbook."""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

from foundry.analysis_data import MAX_COLUMNS
from foundry.structure_proposal import (
    has_undecided_direction,
    normalize_header,
    propose_workbook_structure,
)

sys.path.insert(0, str(Path(__file__).parent / "fixtures" / "structure"))
import build_huachen_like as fixture  # noqa: E402


def _sheet(proposal: dict, prefix: str) -> dict:
    return next(item for item in proposal["sheets"] if item["sheet"].startswith(prefix))


def _pack() -> dict:
    return propose_workbook_structure(fixture.monthly_pack())


def _bytes(book: Workbook) -> bytes:
    import io

    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def test_ae1_ledger_header_row_pairs_sections_and_direction() -> None:
    ledger = _sheet(_pack(), "02")
    assert ledger["decision"] == "table"
    assert ledger["header_row"] == 4
    assert ledger["title_rows"] == [1, 2]

    pairs = {pair["label"]: pair for pair in ledger["amount_pairs"]}
    period_pair = pairs["本期发生额"]
    assert period_pair["primary"] is True
    columns = {column["letter"]: column for column in ledger["columns"]}
    assert columns[period_pair["debit"]]["display"] == "本期发生额借方"
    assert columns[period_pair["credit"]]["display"] == "本期发生额贷方"

    # Direction follows the category column: 收入 sits on credit, 成本费用 on debit.
    assert period_pair["direction"] == {"kind": "by_category", "column": "C"}
    assert has_undecided_direction(ledger) is False

    sections = [item["text"] for item in ledger["excluded_rows"] if item["reason"] == "section"]
    assert sections == ["资产类科目", "负债类科目", "权益类科目", "收入类科目", "成本费用类科目"]
    subtotals = [item["text"] for item in ledger["excluded_rows"] if item["reason"] == "subtotal"]
    assert "合　计" in subtotals
    assert any("平衡校验" in text for text in subtotals)
    assert sum(1 for item in ledger["excluded_rows"] if item["reason"] == "note") == 3


def test_ae3_journal_fill_down_and_subtotal_exclusion() -> None:
    journal = _sheet(_pack(), "03")
    letters = {column["role"]: column["letter"] for column in journal["columns"]}
    assert sorted(journal["fill_down_columns"]) == sorted(
        [letters["doc_id"], letters["date"], letters["summary"]]
    )

    excluded_rows = {item["row"] for item in journal["excluded_rows"]}
    assert all(item["reason"] == "subtotal" for item in journal["excluded_rows"])
    assert len(excluded_rows) == len(fixture.VOUCHERS) + 1  # one 小计 each, plus 合计

    # Every detail line survives. The real pack has 50 vouchers and 183 lines;
    # this fixture carries fewer, so the invariant is asserted, not the count.
    data_rows = _data_row_count(journal, fixture.monthly_pack(), "03·GL月结分录清单")
    assert data_rows == fixture.journal_line_count()


def _data_row_count(proposal: dict, data: bytes, sheet_name: str) -> int:
    from foundry.workbook_bindings import open_workbook_bytes

    book = open_workbook_bytes(data, data_only=True)
    try:
        sheet = book[sheet_name]
        skip = {item["row"] for item in proposal["excluded_rows"]}
        return sum(
            1
            for number, cells in enumerate(sheet.iter_rows(), start=1)
            if number > proposal["header_row"]
            and number not in skip
            and any(cell.value is not None for cell in cells)
        )
    finally:
        book.close()


def test_ae4_presentation_sheets_are_excluded_with_a_reason() -> None:
    decisions = {item["sheet"]: item for item in _pack()["sheets"]}
    for name in ("00·封面与编制说明", "01·核心指标看板", "13·行业对标与假设"):
        assert decisions[name]["decision"] == "excluded"
        assert decisions[name]["reason"] == "presentation_sheet"
    for name in ("02·GL科目余额表", "03·GL月结分录清单", "10·区域经营损益表", "11·城市公司明细表"):
        assert decisions[name]["decision"] == "table"


def test_ae5_scope_hints_come_from_the_title_and_never_invent_currency() -> None:
    ledger = _sheet(_pack(), "02")
    assert ledger["scope_hints"] == {"period": "2026-08", "currency": "", "unit": "元"}


def test_ae5_a_period_span_is_not_reported_as_one_period() -> None:
    region = _sheet(_pack(), "10")
    # 「2026年1-8月」 is a span; declaring 2026-01 would invent scope.
    assert region["scope_hints"]["period"] == ""
    assert region["scope_hints"]["unit"] == "万元"


def test_ae9_debit_credit_without_a_category_column_stays_undecided() -> None:
    sheet = propose_workbook_structure(fixture.no_category_ledger())["sheets"][0]
    assert sheet["decision"] == "table"
    pair = sheet["amount_pairs"][0]
    assert pair["direction"] == {"kind": "undecided"}
    assert has_undecided_direction(sheet) is True
    sides = [column["direction"] for column in sheet["columns"] if column["role"] in {"debit", "credit"}]
    assert sides == ["undecided", "undecided"]


def test_journal_without_a_category_column_reads_the_populated_side() -> None:
    journal = _sheet(_pack(), "03")
    assert journal["amount_pairs"][0]["direction"] == {"kind": "by_populated_side"}
    assert has_undecided_direction(journal) is False


def test_line_broken_header_still_pairs_and_keeps_its_column_letter() -> None:
    ledger = _sheet(_pack(), "02")
    columns = {column["letter"]: column for column in ledger["columns"]}
    assert columns["G"]["header"] == "本期发生额\n借方"
    assert columns["G"]["display"] == "本期发生额借方"
    assert columns["G"]["role"] == "debit"
    assert normalize_header("本期发生额\n借方") == "本期发生额借方"
    assert normalize_header("合　计") == "合计"


def test_header_below_row_four_is_still_located() -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "余额表"
    for row in range(1, 6):
        sheet.cell(row=row, column=1, value=f"说明第 {row} 行")
    for index, text in enumerate(["科目代码", "科目名称", "金额"], start=1):
        sheet.cell(row=6, column=index, value=text)
    for offset, values in enumerate([("6602", "管理费用", 20), ("6601", "销售费用", 10)]):
        for index, value in enumerate(values, start=1):
            sheet.cell(row=7 + offset, column=index, value=value)
    proposal = propose_workbook_structure(_bytes(book))["sheets"][0]
    assert proposal["decision"] == "table"
    assert proposal["header_row"] == 6
    assert proposal["title_rows"] == [1, 2, 3, 4, 5]


def test_empty_and_title_only_sheets_are_excluded_for_lack_of_a_header() -> None:
    book = Workbook()
    empty = book.active
    empty.title = "空表"
    titled = book.create_sheet("仅标题")
    titled["A1"] = "科目余额表"
    titled["A2"] = "华宸物业"
    proposals = {item["sheet"]: item for item in propose_workbook_structure(_bytes(book))["sheets"]}
    assert proposals["空表"]["reason"] == "empty_sheet"
    assert proposals["仅标题"]["reason"] == "no_header"
    assert all(item["decision"] == "excluded" for item in proposals.values())


def test_an_oversized_sheet_is_excluded_without_costing_the_workbook() -> None:
    book = Workbook()
    wide = book.active
    wide.title = "过宽"
    for column in range(1, MAX_COLUMNS + 3):
        wide.cell(row=1, column=column, value=f"列{column}")
        wide.cell(row=2, column=column, value=column)
    fixture.ledger_sheet(book)
    proposals = {item["sheet"]: item for item in propose_workbook_structure(_bytes(book))["sheets"]}
    assert proposals["过宽"]["decision"] == "excluded"
    assert proposals["过宽"]["reason"] == "too_large"
    assert proposals["02·GL科目余额表"]["decision"] == "table"


def test_duplicate_header_text_keeps_both_columns_addressable() -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "明细"
    sheet["A1"] = "科目余额表"
    for index, text in enumerate(["科目名称", "金额", "金额"], start=1):
        sheet.cell(row=2, column=index, value=text)
    for offset, values in enumerate([("管理费用", 20, 21), ("销售费用", 10, 11)]):
        for index, value in enumerate(values, start=1):
            sheet.cell(row=3 + offset, column=index, value=value)
    proposal = propose_workbook_structure(_bytes(book))["sheets"][0]
    letters = [(column["letter"], column["role"]) for column in proposal["columns"]]
    assert letters == [("A", "account_name"), ("B", "amount"), ("C", "unknown")]


def test_the_engine_reads_bytes_and_touches_no_task_directory(tmp_path: Path) -> None:
    before = set(tmp_path.iterdir())
    propose_workbook_structure(fixture.monthly_pack())
    assert set(tmp_path.iterdir()) == before


def test_a_single_sheet_export_is_proposed_too() -> None:
    proposal = propose_workbook_structure(fixture.ledger_only())
    assert len(proposal["sheets"]) == 1
    assert proposal["sheets"][0]["decision"] == "table"
    assert proposal["sheets"][0]["header_row"] == 4


def test_a_corrupt_workbook_is_refused_with_a_foundry_error() -> None:
    from foundry.errors import FoundryError

    with pytest.raises(FoundryError):
        propose_workbook_structure(b"not a workbook at all")


def test_the_proposal_shows_the_rows_around_the_header_it_picked() -> None:
    ledger = _sheet(_pack(), "02")
    preview = ledger["preview"]
    header = next(row for row in preview if row["is_header"])
    assert header["row"] == ledger["header_row"]
    assert [cell["text"] for cell in header["cells"]][:2] == ["科目代码", "科目名称"]
    # The title above and the first data rows below, so a wrong guess is visible.
    assert {row["row"] for row in preview} >= {1, 2, 5, 6}
    assert any("科目余额表" in cell["text"] for row in preview for cell in row["cells"])


def test_preview_text_is_clipped_and_single_line() -> None:
    ledger = _sheet(_pack(), "02")
    for row in ledger["preview"]:
        for cell in row["cells"]:
            assert "\n" not in cell["text"]
            assert len(cell["text"]) <= 40


def test_a_compound_takes_its_role_from_its_head_noun() -> None:
    """期间费用 is a kind of 费用, not an accounting period.

    Role aliases used to match anywhere in a header, so a column of money
    labelled 期间费用 was read as the accounting period and the amounts landed
    in a column called `period` — data, not a visible guess. In a Chinese
    compound the head noun sits last, so that is where an alias counts.
    """
    from foundry.structure_proposal import _role_for

    assert _role_for("期间费用") == "unknown"
    assert _role_for("期间费用预算") == "unknown"
    assert _role_for("期间") == "period"
    # Heads the vocabulary does know, reached through a modifier.
    assert _role_for("供应商类别") == "category"
    assert _role_for("城市公司") == "organization"
    assert _role_for("流失金额") == "amount"
    # A trailing parenthetical qualifies the column, it does not rename it.
    assert _role_for("年化合同金额（万元）") == "amount"
    # 借贷 markers are read before the vocabulary and are unaffected.
    assert _role_for("借方金额") == "debit"
    assert _role_for("本期发生额贷方") == "credit"


def test_a_table_nothing_can_be_read_from_is_named_rather_than_offered() -> None:
    """It is still a table; there is just nothing here to derive.

    A 区域经营损益表 is plainly tabular — one row per region, numbers across —
    so calling it "not a table" would be false. But none of 营业收入, 营业成本,
    期间费用 means anything to a ledger vocabulary, so confirming it could only
    fail. AE4 keeps the decision; `blocked_reason` carries the refusal.
    """
    region = _sheet(_pack(), "10")
    assert region["decision"] == "table"
    assert region["blocked_reason"] == "unreadable_columns"
    # The headers it did read come back, because "I understood none of these"
    # is only actionable if she can see which ones.
    assert [column["display"] for column in region["columns"]][:2] == ["区域", "营业收入"]
    assert all(column["role"] == "unknown" for column in region["columns"])


def test_a_readable_table_carries_no_blocked_reason() -> None:
    ledger = _sheet(_pack(), "02")
    assert ledger["decision"] == "table"
    assert ledger["blocked_reason"] is None


def test_every_sheet_answers_why_it_cannot_be_confirmed() -> None:
    for sheet in _pack()["sheets"]:
        if sheet["decision"] == "excluded":
            assert sheet["blocked_reason"] == sheet["reason"]


def test_every_unblocked_table_can_actually_be_derived() -> None:
    """The gate and the snapshot builder must not drift apart.

    `_carries_data` decides what the card offers and `_derived_headers` decides
    what can be built. They live in different modules and answer the same
    question, so this holds them together on a real-shaped pack.
    """
    from foundry.structure_snapshot import derive_table, undeclarable_reason

    data = fixture.monthly_pack()
    for sheet in propose_workbook_structure(data)["sheets"]:
        if sheet["decision"] != "table":
            continue
        if sheet["blocked_reason"]:
            assert undeclarable_reason(sheet) is not None
            continue
        assert undeclarable_reason(sheet) is None
        derive_table(data, sheet)
