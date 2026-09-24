from __future__ import annotations

from decimal import Decimal

import pytest

from foundry.analysis import AnalysisStore, _new
from foundry.errors import FoundryError
from foundry.finance_explanations import validate_explanation
from foundry.task import TaskStore


def _run(
    analysis: AnalysisStore,
    *,
    value: str,
    suffix: str,
    source_hash: str = "source-hash-a",
    seq: int | None = None,
    row_extra: dict | None = None,
) -> dict:
    record = _new(
        analysis.task,
        id=suffix * 32,
        analysis_id="a" * 32,
        engine="finance",
        execution_status="succeeded",
        reproducibility_status="unverified",
        bindings={"budget": "dataset-a"},
        input_hashes={"budget": "data-hash-a"},
        source_hashes={"budget": source_hash},
        finance_scope={
            "budget": {
                "period": "2026-01..2026-08",
                "scenario": "BUD",
                "organization": "集团",
                "unit": "万元",
                "currency": "CNY",
            }
        },
        **({"seq": seq} if seq is not None else {}),
        result={
            "columns": [{"name": "metric"}, {"name": "actual"}],
            "rows": [{
                "metric_key": "revenue",
                "metric": "收入",
                "actual": value,
                "unit": "万元",
                "currency": "CNY",
                **(row_extra or {}),
            }],
            "row_count": 1,
            "truncated": False,
        },
    )
    return analysis._write(analysis._object_path("runs", record["id"]), record)


def _huachen_value_run(
    analysis: AnalysisStore,
    *,
    actual_value: str,
    suffix: str,
    source_hash: str = "source-hash-a",
) -> dict:
    record = _new(
        analysis.task,
        id=suffix * 32,
        analysis_id="a" * 32,
        engine="python",
        execution_status="succeeded",
        reproducibility_status="unverified",
        bindings={"budget": "dataset-a", "close": "dataset-b"},
        input_hashes={"budget": "data-hash-a", "close": "data-hash-b"},
        source_hashes={"budget": source_hash, "close": "source-hash-b"},
        result={
            "columns": [
                {"name": "metric_key", "type": "VARCHAR"},
                {"name": "period", "type": "VARCHAR"},
                {"name": "organization", "type": "VARCHAR"},
                {"name": "scenario", "type": "VARCHAR"},
                {"name": "unit", "type": "VARCHAR"},
                {"name": "value", "type": "DECIMAL(38,18)"},
            ],
            "rows": [
                {
                    "metric_key": "revenue",
                    "scenario": "actual_ytd",
                    "period": "2026-01:2026-08",
                    "organization": "集团合并",
                    "value": actual_value,
                    "unit": "万元",
                },
                {
                    "metric_key": "revenue",
                    "scenario": "budget_ytd",
                    "period": "2026-01:2026-08",
                    "organization": "集团合并",
                    "value": "344487.0552",
                    "unit": "万元",
                },
            ],
            "row_count": 2,
            "truncated": False,
        },
    )
    return analysis._write(analysis._object_path("runs", record["id"]), record)


def test_revision_compare_flags_unquoted_same_basis_numeric_regression(tmp_path):
    task = TaskStore(tmp_path / "var").create(title="revision compare")
    analysis = AnalysisStore(task)
    old_run = _run(analysis, value="325644", suffix="1")
    new_run = _run(analysis, value="344487.0552", suffix="2")
    parent = analysis.save_draft({
        "title": "旧稿",
        "body": "收入完成情况。",
        "run_ids": [old_run["id"]],
    })

    revised = analysis.save_draft({
        "parent_id": parent["id"],
        "body": "只修正其他收益预算。",
        "run_ids": [new_run["id"]],
    })

    issues = revised["revision_comparison"]["issues"]
    assert revised["revision_comparison"]["summary"]["status"] == "needs_attention"
    assert issues == [{
        "code": "revision_numeric_regression",
        "business_key": issues[0]["business_key"],
        "column": "actual",
        "old_value": "325644",
        "new_value": "344487.0552",
        "old_run_id": old_run["id"],
        "new_run_id": new_run["id"],
        "severity": "error",
        "message": "同一来源和口径下的关键数值发生变化，需要解释或修正。",
    }]


def test_revision_compare_flags_huachen_value_column_regression(tmp_path):
    task = TaskStore(tmp_path / "var").create(title="revision compare")
    analysis = AnalysisStore(task)
    old_run = _huachen_value_run(analysis, actual_value="325644.0000", suffix="1")
    new_run = _huachen_value_run(analysis, actual_value="344487.0552", suffix="2")
    parent = analysis.save_draft({"title": "旧稿", "body": "收入完成情况。", "run_ids": [old_run["id"]]})

    revised = analysis.save_draft({
        "parent_id": parent["id"],
        "body": "只修正其他收益预算。",
        "run_ids": [new_run["id"]],
    })

    issue = next(item for item in revised["revision_comparison"]["issues"] if item["code"] == "revision_numeric_regression")
    assert issue["column"] == "value"
    assert issue["old_value"] == "325644.0000"
    assert issue["new_value"] == "344487.0552"


def test_revision_compare_does_not_treat_arbitrary_value_column_as_financial(tmp_path):
    task = TaskStore(tmp_path / "var").create(title="revision compare")
    analysis = AnalysisStore(task)
    old_run = _new(
        analysis.task,
        id="c" * 32,
        analysis_id="a" * 32,
        engine="finance",
        execution_status="succeeded",
        reproducibility_status="unverified",
        result={"rows": [{"label": "备注", "value": "1"}]},
    )
    new_run = _new(
        analysis.task,
        id="d" * 32,
        analysis_id="a" * 32,
        engine="finance",
        execution_status="succeeded",
        reproducibility_status="unverified",
        result={"rows": [{"label": "备注", "value": "2"}]},
    )
    old_run = analysis._write(analysis._object_path("runs", old_run["id"]), old_run)
    new_run = analysis._write(analysis._object_path("runs", new_run["id"]), new_run)
    parent = analysis.save_draft({"title": "旧稿", "body": "备注。", "run_ids": [old_run["id"]]})

    revised = analysis.save_draft({"parent_id": parent["id"], "body": "备注。", "run_ids": [new_run["id"]]})

    assert revised["revision_comparison"]["summary"]["status"] == "no_comparable_runs"


def test_revision_compare_discloses_basis_change_without_numeric_regression(tmp_path):
    task = TaskStore(tmp_path / "var").create(title="revision compare")
    analysis = AnalysisStore(task)
    old_run = _run(analysis, value="325644", suffix="1")
    new_run = _run(analysis, value="344487.0552", suffix="2", source_hash="source-hash-b")
    parent = analysis.save_draft({"body": "旧稿", "run_ids": [old_run["id"]]})

    revised = analysis.save_draft({"parent_id": parent["id"], "body": "新来源", "run_ids": [new_run["id"]]})

    assert revised["revision_comparison"]["issues"][0]["code"] == "revision_basis_changed"


def test_revision_basis_change_includes_changed_values_and_hash_aliases(tmp_path):
    task = TaskStore(tmp_path / "var").create(title="revision compare")
    analysis = AnalysisStore(task)
    old_run = _run(analysis, value="325644", suffix="1", row_extra={"source": "close workbook"})
    new_run = _run(
        analysis,
        value="344487.0552",
        suffix="2",
        source_hash="budget-hash-b",
        row_extra={"source": "close workbook"},
    )
    parent = analysis.save_draft({"body": "旧稿", "run_ids": [old_run["id"]]})

    revised = analysis.save_draft({"parent_id": parent["id"], "body": "新来源", "run_ids": [new_run["id"]]})

    issue = revised["revision_comparison"]["issues"][0]
    assert issue["code"] == "revision_basis_changed"
    assert issue["severity"] == "warning"
    assert issue["changed_columns"] == [{
        "column": "actual",
        "old_value": "325644",
        "new_value": "344487.0552",
    }]
    assert issue["changed_hash_aliases"] == [{
        "kind": "source_hashes",
        "alias": "budget",
        "old_value": "source-hash-a",
        "new_value": "budget-hash-b",
    }]
    assert not any(item["code"] == "revision_numeric_regression" for item in revised["revision_comparison"]["issues"])


def test_revision_compare_uses_legacy_period_org_and_latest_run_for_regression(tmp_path):
    task = TaskStore(tmp_path / "var").create(title="revision compare")
    analysis = AnalysisStore(task)
    row_scope = {"period": "2026-01..2026-08", "organization": "集团", "scenario": "actual_vs_budget"}
    old_run = _run(analysis, value="325644", suffix="3", seq=1, row_extra=row_scope)
    copied_old_run = _run(analysis, value="325644", suffix="4", seq=2, row_extra=row_scope)
    new_run = _run(analysis, value="344487.0552", suffix="5", seq=3, row_extra=row_scope)
    # Legacy script rows often carry scope directly on the row rather than under
    # finance_scope. These fields must be part of the business identity so YTD,
    # month, and regional rows do not collapse into one ambiguous metric bucket.
    parent = analysis.save_draft({"title": "旧稿", "body": "收入。", "run_ids": [old_run["id"]]})

    revised = analysis.save_draft({
        "parent_id": parent["id"],
        "body": "只修其他损益，收入不应改变。",
        "run_ids": [copied_old_run["id"], new_run["id"]],
    })

    issue = next(item for item in revised["revision_comparison"]["issues"] if item["code"] == "revision_numeric_regression")
    assert issue["old_value"] == "325644"
    assert issue["new_value"] == "344487.0552"
    assert issue["new_run_id"] == new_run["id"]
    assert any(item["code"] == "revision_numeric_regression" for item in revised["quality"]["issues"])


def test_revision_compare_zero_business_intersection_is_not_pass(tmp_path):
    task = TaskStore(tmp_path / "var").create(title="revision compare")
    analysis = AnalysisStore(task)
    old_run = _run(analysis, value="325644", suffix="a")
    new_run = _run(analysis, value="42", suffix="b", row_extra={"metric_key": "profit", "metric": "利润"})
    parent = analysis.save_draft({"title": "旧稿", "body": "收入。", "run_ids": [old_run["id"]]})

    revised = analysis.save_draft({"parent_id": parent["id"], "body": "利润。", "run_ids": [new_run["id"]]})

    assert revised["revision_comparison"]["summary"]["status"] == "no_comparable_runs"
    assert revised["revision_comparison"]["summary"]["compared_business_keys"] == 0


def test_evidence_recipe_mismatch_enters_draft_quality_and_validate_explanation(tmp_path):
    from openpyxl import Workbook

    task = TaskStore(tmp_path / "var").create(title="evidence recipes")
    analysis = AnalysisStore(task)
    source = tmp_path / "budget.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "预算"
    sheet["B2"] = 100
    sheet["C2"] = 23
    book.save(source)
    book.close()
    dataset = analysis.add_dataset(source, name="budget", metadata={"mode": "workbook"})
    run = _run(analysis, value="120", suffix="6")
    recipe = {
        "id": "budget-sum",
        "label": "预算合计",
        "kind": "sum",
        "dataset_id": dataset["id"],
        "source_hash": dataset["source_hash"],
        "sheet": "预算",
        "range": "B2:C2",
        "unit": "万元",
        "claim": {"run_id": run["id"], "row": 0, "column": "actual"},
        "business_key": {"metric": "budget", "period": "2026-08"},
    }

    check = validate_explanation(analysis, {"body": "预算为120万元。", "run_ids": [run["id"]], "evidence_recipes": [recipe]})
    issue = next(item for item in check["issues"] if item["code"] == "evidence_recipe_mismatch")
    assert issue["actual"] == "123"
    assert issue["expected"] == "120"
    assert issue["location"]["sheet"] == "预算"

    draft = analysis.save_draft({"title": "说明", "body": "预算为120万元。", "run_ids": [run["id"]], "evidence_recipes": [recipe]})
    assert draft["evidence_report"]["summary"]["status"] == "needs_attention"
    assert any(item["code"] == "evidence_recipe_mismatch" for item in draft["quality"]["issues"])


def test_evidence_recipe_zero_and_invalid_ranges_are_not_silent_passes(tmp_path):
    from foundry.financial_evidence import evaluate_evidence_recipes
    from foundry.errors import FoundryError

    task = TaskStore(tmp_path / "var").create(title="evidence recipes")
    analysis = AnalysisStore(task)
    assert evaluate_evidence_recipes(analysis, [])["summary"]["status"] == "not_configured"

    bad = {
        "id": "full-column",
        "label": "full-column",
        "kind": "sum",
        "dataset_id": "0" * 32,
        "source_hash": "a" * 64,
        "sheet": "预算",
        "range": "A:A",
        "claim": {"run_id": "1" * 32, "row": 0, "column": "actual"},
    }
    try:
        evaluate_evidence_recipes(analysis, [bad])
    except FoundryError as exc:
        assert "bounded A1" in str(exc)
    else:
        raise AssertionError("full-column recipe should be rejected")

    duplicate = [{**bad, "range": "A1:A1", "id": "dup"}, {**bad, "range": "B1:B1", "id": "dup"}]
    try:
        evaluate_evidence_recipes(analysis, duplicate)
    except FoundryError as exc:
        assert "unique" in str(exc)
    else:
        raise AssertionError("duplicate recipe ids should be rejected")


def _factor_workbook_dataset(tmp_path, analysis: AnalysisStore, *, denominator: int = 100):
    from openpyxl import Workbook

    source = tmp_path / "factor.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "测算"
    sheet["B2"] = 30
    sheet["C2"] = denominator
    sheet["D2"] = 60
    book.save(source)
    book.close()
    return analysis.add_dataset(source, name="factor", metadata={"mode": "workbook"})


def _factor_run(analysis: AnalysisStore, *, value: str, suffix: str, metric_key: str = "gross_margin", period: str = "2026-08") -> dict:
    return _run(
        analysis,
        value=value,
        suffix=suffix,
        row_extra={"metric_key": metric_key, "metric": metric_key, "period": period},
    )


def _factor_recipe_pair(dataset: dict, factor_run: dict, result_run: dict, *, metric_key: str = "gross_margin", period: str = "2026-08") -> list[dict]:
    business_key = {"metric_key": metric_key, "period": period, "organization": "集团"}
    factor = {
        "id": "gross-margin-factor",
        "label": "毛利率",
        "kind": "ratio",
        "left": {"dataset_id": dataset["id"], "source_hash": dataset["source_hash"], "sheet": "测算", "range": "B2"},
        "right": {"dataset_id": dataset["id"], "source_hash": dataset["source_hash"], "sheet": "测算", "range": "C2"},
        "claim": {"run_id": factor_run["id"], "row": 0, "column": "actual"},
        "business_key": business_key,
    }
    result = {
        "id": "gross-margin-budget",
        "label": "毛利预算",
        "kind": "direct",
        "dataset_id": dataset["id"],
        "source_hash": dataset["source_hash"],
        "sheet": "测算",
        "range": "D2",
        "factor_recipe_id": "gross-margin-factor",
        "claim": {"run_id": result_run["id"], "row": 0, "column": "actual"},
        "business_key": business_key,
    }
    return [factor, result]


def test_evidence_recipe_factor_recipe_computes_archived_ratio_for_direct_result(tmp_path):
    task = TaskStore(tmp_path / "var").create(title="factor recipe")
    analysis = AnalysisStore(task)
    dataset = _factor_workbook_dataset(tmp_path, analysis)
    factor_run = _factor_run(analysis, value="0.3", suffix="1")
    result_run = _factor_run(analysis, value="18", suffix="2")

    draft = analysis.save_draft({
        "title": "正确",
        "body": "毛利预算18。",
        "run_ids": [factor_run["id"], result_run["id"]],
        "evidence_recipes": _factor_recipe_pair(dataset, factor_run, result_run),
    })

    assert draft["evidence_report"]["summary"]["status"] == "pass"
    result = next(item for item in draft["evidence_report"]["results"] if item["recipe_id"] == "gross-margin-budget")
    assert result["actual"] == "18.0"
    assert result["expected"] == "18"
    assert result["factor_recipe_id"] == "gross-margin-factor"
    assert result["factor_value"] == "0.3"


def test_evidence_recipe_factor_recipe_rejects_wrong_metric_identity(tmp_path):
    from foundry.financial_evidence import evaluate_evidence_recipes

    task = TaskStore(tmp_path / "var").create(title="factor recipe")
    analysis = AnalysisStore(task)
    dataset = _factor_workbook_dataset(tmp_path, analysis)
    factor_run = _factor_run(analysis, value="0.3", suffix="3", metric_key="operating_margin")
    result_run = _factor_run(analysis, value="18", suffix="4", metric_key="gross_margin")
    recipes = _factor_recipe_pair(dataset, factor_run, result_run)
    recipes[0]["business_key"] = {"metric_key": "operating_margin", "period": "2026-08", "organization": "集团"}

    with pytest.raises(FoundryError, match="metric_key and period"):
        evaluate_evidence_recipes(analysis, recipes)

    with pytest.raises(FoundryError, match="explicit scale"):
        evaluate_evidence_recipes(analysis, [{**recipes[1], "id": "scaled-dependent", "scale": "0.3"}])

    scaled_factor = _factor_recipe_pair(dataset, _factor_run(analysis, value="30", suffix="5"), result_run)
    scaled_factor[0]["scale"] = "100"
    with pytest.raises(FoundryError, match="factor ratio recipe scale"):
        evaluate_evidence_recipes(analysis, scaled_factor)


def test_evidence_recipe_factor_recipe_rejects_mislabeled_claim_rows(tmp_path):
    task = TaskStore(tmp_path / "var").create(title="factor recipe")
    analysis = AnalysisStore(task)
    dataset = _factor_workbook_dataset(tmp_path, analysis)
    factor_run = _factor_run(analysis, value="0.3", suffix="b", metric_key="operating_margin")
    result_run = _factor_run(analysis, value="18", suffix="c", metric_key="gross_margin")
    factor_mislabeled = analysis.save_draft({
        "title": "错误",
        "body": "毛利预算18。",
        "run_ids": [factor_run["id"], result_run["id"]],
        "evidence_recipes": _factor_recipe_pair(dataset, factor_run, result_run, metric_key="gross_margin"),
    })

    assert factor_mislabeled["evidence_report"]["summary"]["status"] == "needs_attention"
    assert any(
        item["recipe_id"] == "gross-margin-budget" and "business_key metric_key and period" in item["message"]
        for item in factor_mislabeled["evidence_report"]["issues"]
    )

    dependent_analysis = AnalysisStore(TaskStore(tmp_path / "dependent-var").create(title="dependent identity"))
    dependent_dataset = _factor_workbook_dataset(tmp_path, dependent_analysis)
    dependent_factor_run = _factor_run(dependent_analysis, value="0.3", suffix="d", metric_key="gross_margin")
    dependent_result_run = _factor_run(dependent_analysis, value="18", suffix="e", metric_key="operating_margin")
    dependent_mislabeled = dependent_analysis.save_draft({
        "title": "错误",
        "body": "毛利预算18。",
        "run_ids": [dependent_factor_run["id"], dependent_result_run["id"]],
        "evidence_recipes": _factor_recipe_pair(dependent_dataset, dependent_factor_run, dependent_result_run, metric_key="gross_margin"),
    })

    assert dependent_mislabeled["evidence_report"]["summary"]["status"] == "needs_attention"
    assert any(
        item["recipe_id"] == "gross-margin-budget" and "business_key metric_key and period" in item["message"]
        for item in dependent_mislabeled["evidence_report"]["issues"]
    )


def test_evidence_recipe_factor_recipe_requires_factor_claim_to_match_recomputation(tmp_path):
    task = TaskStore(tmp_path / "var").create(title="factor recipe")
    analysis = AnalysisStore(task)
    dataset = _factor_workbook_dataset(tmp_path, analysis)
    factor_run = _factor_run(analysis, value="0.2", suffix="5")
    result_run = _factor_run(analysis, value="18", suffix="6")

    draft = analysis.save_draft({
        "title": "错误",
        "body": "毛利预算18。",
        "run_ids": [factor_run["id"], result_run["id"]],
        "evidence_recipes": _factor_recipe_pair(dataset, factor_run, result_run),
    })

    assert draft["evidence_report"]["summary"]["status"] == "needs_attention"
    factor_issue = next(item for item in draft["evidence_report"]["issues"] if item["recipe_id"] == "gross-margin-factor")
    dependent_issue = next(item for item in draft["evidence_report"]["issues"] if item["recipe_id"] == "gross-margin-budget")
    assert factor_issue["code"] == "evidence_recipe_mismatch"
    assert factor_issue["actual"] == "0.3"
    assert "测算!B2 ÷ 测算!C2" in factor_issue["message"]
    assert dependent_issue["code"] == "evidence_recipe_unreadable"
    assert "系数规则未通过归档来源的精确复算" in dependent_issue["message"]


def test_evidence_recipe_factor_recipe_preserves_bad_hash_and_zero_denominator_failures(tmp_path):
    task = TaskStore(tmp_path / "var").create(title="factor recipe")
    analysis = AnalysisStore(task)
    dataset = _factor_workbook_dataset(tmp_path, analysis)
    factor_run = _factor_run(analysis, value="0.3", suffix="7")
    result_run = _factor_run(analysis, value="18", suffix="8")
    bad_hash = _factor_recipe_pair(dataset, factor_run, result_run)
    bad_hash[0]["left"]["source_hash"] = "0" * 64

    bad_hash_draft = analysis.save_draft({
        "title": "坏hash",
        "body": "毛利预算18。",
        "run_ids": [factor_run["id"], result_run["id"]],
        "evidence_recipes": bad_hash,
    })

    assert bad_hash_draft["evidence_report"]["summary"]["status"] == "needs_attention"
    assert any("source_hash" in item["message"] for item in bad_hash_draft["evidence_report"]["issues"])

    zero_analysis = AnalysisStore(TaskStore(tmp_path / "zero-var").create(title="zero denominator"))
    zero_dataset = _factor_workbook_dataset(tmp_path, zero_analysis, denominator=0)
    zero_factor_run = _factor_run(zero_analysis, value="0.3", suffix="9")
    zero_result_run = _factor_run(zero_analysis, value="18", suffix="a")
    zero_draft = zero_analysis.save_draft({
        "title": "零分母",
        "body": "毛利预算18。",
        "run_ids": [zero_factor_run["id"], zero_result_run["id"]],
        "evidence_recipes": _factor_recipe_pair(zero_dataset, zero_factor_run, zero_result_run),
    })

    assert zero_draft["evidence_report"]["summary"]["status"] == "needs_attention"
    assert any("denominator is zero" in item["message"] for item in zero_draft["evidence_report"]["issues"])

def test_consistency_rules_read_values_from_authorized_sources(tmp_path):
    from openpyxl import Workbook

    task = TaskStore(tmp_path / "var").create(title="consistency")
    analysis = AnalysisStore(task)
    source = tmp_path / "variance.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "差异"
    sheet["B2"] = 8
    sheet["C2"] = 12
    book.save(source)
    book.close()
    numbers = analysis.add_dataset(source, name="variance", metadata={"mode": "workbook"})
    text_path = tmp_path / "assertion.md"
    text_path.write_text("所有关键差异均超过10万元。\n", encoding="utf-8")
    text_source = analysis.add_dataset(text_path, name="assertion")
    run = _run(analysis, value="20", suffix="7")

    draft = analysis.save_draft({
        "title": "一致性",
        "body": "所有关键差异均超过10万元。",
        "run_ids": [run["id"]],
        "consistency_sources": {
            "a": {"kind": "number", "dataset_id": numbers["id"], "source_hash": numbers["source_hash"], "sheet": "差异", "range": "B2"},
            "b": {"kind": "number", "dataset_id": numbers["id"], "source_hash": numbers["source_hash"], "sheet": "差异", "range": "C2"},
            "assertion": {"kind": "text", "dataset_id": text_source["id"], "source_hash": text_source["source_hash"], "paragraph": 1},
        },
        "consistency_rules": [{
            "id": "all-material",
            "kind": "all_absolute_exceed",
            "threshold": "10",
            "asserted_result": True,
            "assertion_source_id": "assertion",
            "values": [{"id": "a", "source_id": "a"}, {"id": "b", "source_id": "b"}],
        }],
    })

    assert draft["consistency_report"]["summary"]["status"] == "needs_attention"
    issue = next(item for item in draft["quality"]["issues"] if item["code"] == "financial_consistency_mismatch")
    assert issue["computed_result"] is False
    assert issue["original_text"] == "所有关键差异均超过10万元。"
    assert draft["quality"]["status_dimensions"]["source_validation"]["status"] == "needs_attention"


def test_consistency_rules_reject_caller_supplied_values(tmp_path):
    task = TaskStore(tmp_path / "var").create(title="consistency")
    analysis = AnalysisStore(task)
    run = _run(analysis, value="20", suffix="8")
    try:
        analysis.save_draft({
            "title": "一致性",
            "body": "差异超过10。",
            "run_ids": [run["id"]],
            "consistency_rules": [{
                "id": "fake",
                "kind": "all_absolute_exceed",
                "threshold": "10",
                "asserted_result": True,
                "assertion_source_id": "assertion",
                "values": [{"id": "fake", "source_id": "missing", "value": "999"}],
            }],
        })
    except FoundryError as exc:
        assert "authorized source" in str(exc)
    else:
        raise AssertionError("caller supplied consistency values should be rejected")


def test_quality_status_dimensions_keep_missing_source_validation_distinct(tmp_path):
    task = TaskStore(tmp_path / "var").create(title="dimensions")
    analysis = AnalysisStore(task)
    run = _run(analysis, value="20", suffix="9")
    draft = analysis.save_draft({"title": "说明", "body": "说明。", "run_ids": [run["id"]]})
    dimensions = draft["quality"]["status_dimensions"]
    assert set(dimensions) == {"execution", "numeric_citation", "source_validation", "professional_review"}
    assert dimensions["execution"]["status"] == "pass"
    assert dimensions["source_validation"]["status"] == "not_configured"
    assert dimensions["professional_review"]["status"] == "unreviewed"

def test_huachen_direct_source_recipe_accepts_3185_and_rejects_proportional_3063(tmp_path):
    from openpyxl import Workbook

    task = TaskStore(tmp_path / "var").create(title="huachen source recipe")
    analysis = AnalysisStore(task)
    source = tmp_path / "budget.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "预算"
    sheet["B36"] = 3185
    book.save(source)
    book.close()
    dataset = analysis.add_dataset(source, name="budget", metadata={"mode": "workbook"})
    correct = _run(analysis, value="3185", suffix="b")
    wrong = _run(analysis, value="3063.0494", suffix="c")

    base_recipe = {
        "id": "other-gains-budget-direct",
        "label": "其他收益预算直接取数",
        "kind": "direct",
        "dataset_id": dataset["id"],
        "source_hash": dataset["source_hash"],
        "sheet": "预算",
        "range": "B36",
        "claim": {"run_id": correct["id"], "row": 0, "column": "actual"},
        "business_key": {"metric": "other_gains_budget", "period": "2026-01..2026-08", "organization": "集团"},
    }
    ok = analysis.save_draft({"title": "正确", "body": "其他收益预算3185。", "run_ids": [correct["id"]], "evidence_recipes": [base_recipe]})
    assert ok["evidence_report"]["summary"]["status"] == "pass"

    bad_recipe = {**base_recipe, "claim": {"run_id": wrong["id"], "row": 0, "column": "actual"}}
    bad = analysis.save_draft({"title": "错误", "body": "其他收益预算3063.0494。", "run_ids": [wrong["id"]], "evidence_recipes": [bad_recipe]})
    issue = next(item for item in bad["quality"]["issues"] if item["code"] == "evidence_recipe_mismatch")
    assert issue["actual"] == "3185"
    assert issue["expected"] == "3063.0494"
    assert issue["location"]["range"] == "B36"

def test_exact_unit_power_mismatch_suggests_scale_and_scaled_retry_passes(tmp_path):
    """Huachen rerun: 04·合并利润表!F5 is 元 (5000000000), the run cites 万元 (500000)."""
    from openpyxl import Workbook

    task = TaskStore(tmp_path / "var").create(title="unit scale hint")
    analysis = AnalysisStore(task)
    source = tmp_path / "close.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "04·合并利润表"
    sheet["F5"] = 5000000000
    book.save(source)
    book.close()
    dataset = analysis.add_dataset(source, name="close", metadata={"mode": "workbook"})
    run = _run(analysis, value="500000.0000", suffix="d")
    recipe = {
        "id": "le-revenue",
        "label": "全年预算列（LE 收入）",
        "kind": "direct",
        "dataset_id": dataset["id"],
        "source_hash": dataset["source_hash"],
        "sheet": "04·合并利润表",
        "range": "F5",
        "claim": {"run_id": run["id"], "row": 0, "column": "actual"},
    }

    missing = analysis.save_draft({"title": "缺换算", "body": "LE 收入500000.0000。", "run_ids": [run["id"]], "evidence_recipes": [recipe]})
    issue = next(item for item in missing["quality"]["issues"] if item["code"] == "evidence_recipe_mismatch")
    assert issue["suggested_scale"] == "0.0001"
    assert "相差 10000 倍" in issue["message"]
    assert "scale = 0.0001" in issue["message"]

    fixed = analysis.save_draft({"title": "已换算", "body": "LE 收入500000.0000。", "run_ids": [run["id"]], "evidence_recipes": [{**recipe, "scale": issue["suggested_scale"]}]})
    assert fixed["evidence_report"]["summary"]["status"] == "pass"


def test_non_unit_mismatch_has_no_scale_hint(tmp_path):
    from foundry.financial_evidence import _unit_scale_hint

    assert _unit_scale_hint(Decimal("3185"), Decimal("3063.0494")) is None
    assert _unit_scale_hint(Decimal("100"), Decimal("1000")) is None
    assert _unit_scale_hint(Decimal("64441.376900000001"), Decimal("64441.3769")) is None
    assert _unit_scale_hint(Decimal("0"), Decimal("500000")) is None


def _income_statement_dataset(tmp_path, analysis: AnalysisStore, *, value=5000000000) -> dict:
    from openpyxl import Workbook

    source = tmp_path / "income.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "04·合并利润表"
    sheet["F5"] = value
    book.save(source)
    book.close()
    return analysis.add_dataset(source, name="income", metadata={"mode": "workbook"})


def _le_revenue_recipe(dataset: dict, run: dict, **extra) -> dict:
    return {
        "id": "le-revenue",
        "label": "全年预算列（LE 收入）",
        "kind": "direct",
        "dataset_id": dataset["id"],
        "source_hash": dataset["source_hash"],
        "sheet": "04·合并利润表",
        "range": "F5",
        "claim": {"run_id": run["id"], "row": 0, "column": "actual"},
        **extra,
    }


def _mismatch(draft: dict, recipe_id: str) -> dict:
    issue = next(item for item in draft["evidence_report"]["issues"] if item["recipe_id"] == recipe_id)
    assert issue["code"] == "evidence_recipe_mismatch"
    return issue


def test_unit_scale_hint_names_the_total_scale_when_the_recipe_already_has_one(tmp_path):
    """F5 is 元 and the recipe already converts to 万元 (0.0001), but the run cites 亿元."""
    analysis = AnalysisStore(TaskStore(tmp_path / "var").create(title="existing scale"))
    dataset = _income_statement_dataset(tmp_path, analysis)
    yi_run = _run(analysis, value="50", suffix="1")
    recipe = _le_revenue_recipe(dataset, yi_run, scale="0.0001")

    wrong = analysis.save_draft({"title": "亿元", "body": "LE 收入50。", "run_ids": [yi_run["id"]], "evidence_recipes": [recipe]})
    issue = _mismatch(wrong, "le-revenue")
    assert issue["suggested_scale"] == "0.00000001"
    assert "相差 10000 倍" in issue["message"]
    assert "疑似单位换算有误" in issue["message"]
    assert "scale = 0.00000001" in issue["message"]
    fixed = analysis.save_draft({"title": "亿元", "body": "LE 收入50。", "run_ids": [yi_run["id"]], "evidence_recipes": [{**recipe, "scale": issue["suggested_scale"]}]})
    assert fixed["evidence_report"]["summary"]["status"] == "pass"

    # Reverse direction: the run cites 元, so the existing 万元 conversion has to go.
    yuan_run = _run(analysis, value="5000000000", suffix="2")
    reverse = _le_revenue_recipe(dataset, yuan_run, scale="0.0001")
    wrong = analysis.save_draft({"title": "元", "body": "LE 收入5000000000。", "run_ids": [yuan_run["id"]], "evidence_recipes": [reverse]})
    issue = _mismatch(wrong, "le-revenue")
    assert issue["suggested_scale"] == "1"
    assert "scale = 1。" in issue["message"]
    fixed = analysis.save_draft({"title": "元", "body": "LE 收入5000000000。", "run_ids": [yuan_run["id"]], "evidence_recipes": [{**reverse, "scale": issue["suggested_scale"]}]})
    assert fixed["evidence_report"]["summary"]["status"] == "pass"


def test_factor_pair_mismatch_keeps_plain_message_without_scale_hint(tmp_path):
    """Neither the factor-bound recipe nor its factor ratio may carry a scale, so neither is told to set one."""
    analysis = AnalysisStore(TaskStore(tmp_path / "var").create(title="factor scale hint"))
    dataset = _factor_workbook_dataset(tmp_path, analysis)
    factor_run = _factor_run(analysis, value="0.3", suffix="1")
    wan_run = _factor_run(analysis, value="0.0018", suffix="2")

    dependent = analysis.save_draft({
        "title": "毛利预算",
        "body": "毛利预算0.0018。",
        "run_ids": [factor_run["id"], wan_run["id"]],
        "evidence_recipes": _factor_recipe_pair(dataset, factor_run, wan_run),
    })
    issue = _mismatch(dependent, "gross-margin-budget")
    assert "suggested_scale" not in issue
    assert issue["message"] == "来源区域 测算!D2 的独立计算为 18.0，运行引用值为 0.0018。"

    percent_run = _factor_run(analysis, value="3000", suffix="3")
    result_run = _factor_run(analysis, value="18", suffix="4")
    factor = analysis.save_draft({
        "title": "毛利率",
        "body": "毛利率3000，毛利预算18。",
        "run_ids": [percent_run["id"], result_run["id"]],
        "evidence_recipes": _factor_recipe_pair(dataset, percent_run, result_run),
    })
    issue = _mismatch(factor, "gross-margin-factor")
    assert "suggested_scale" not in issue
    assert issue["message"] == "来源区域 测算!B2 ÷ 测算!C2 的独立计算为 0.3，运行引用值为 3000。"

    # The same ratio outside a factor pair may take a scale, so it keeps the hint and the retry passes.
    standalone = {**_factor_recipe_pair(dataset, percent_run, result_run)[0], "id": "standalone-margin"}
    alone = analysis.save_draft({"title": "毛利率", "body": "毛利率3000。", "run_ids": [percent_run["id"]], "evidence_recipes": [standalone]})
    issue = _mismatch(alone, "standalone-margin")
    assert issue["suggested_scale"] == "10000"
    fixed = analysis.save_draft({"title": "毛利率", "body": "毛利率3000。", "run_ids": [percent_run["id"]], "evidence_recipes": [{**standalone, "scale": issue["suggested_scale"]}]})
    assert fixed["evidence_report"]["summary"]["status"] == "pass"


def test_unit_scale_hint_requires_identical_digits_and_sign(tmp_path):
    from foundry.financial_evidence import _unit_scale_hint

    # Within 1e-28 of a unit power: a 28-digit division would round these onto 10^k.
    assert _unit_scale_hint(Decimal("5000000000"), Decimal("499999.99999999999999999999999999999")) is None
    assert _unit_scale_hint(Decimal("5000000000"), Decimal("500000.00000000000000000000001")) is None
    assert _unit_scale_hint(Decimal("3"), Decimal("0.0029999999999999999999999999999")) is None
    # Float tails near a unit power; the last two have trailing-digit gaps of exactly 10^3 and 10^4.
    assert _unit_scale_hint(Decimal("3185"), Decimal("31850000.5")) is None
    assert _unit_scale_hint(Decimal("64441.376900000001"), Decimal("644413769")) is None
    assert _unit_scale_hint(Decimal("64441.3769"), Decimal("644413769.0000001")) is None
    assert _unit_scale_hint(Decimal("3185"), Decimal("31860000")) is None
    # Opposite signs, and a power of ten that is not a unit step.
    assert _unit_scale_hint(Decimal("-5000000000"), Decimal("500000")) is None
    assert _unit_scale_hint(Decimal("1"), Decimal("100000")) is None
    # Exact unit powers in both directions, including negative values and trailing zeros.
    assert _unit_scale_hint(Decimal("5000000000"), Decimal("500000.0000")) == -4
    assert _unit_scale_hint(Decimal("500000"), Decimal("5000000000")) == 4
    assert _unit_scale_hint(Decimal("-3185"), Decimal("-3.185")) == -3
    assert _unit_scale_hint(Decimal("-2.5"), Decimal("-2500000.00")) == 6
    assert _unit_scale_hint(Decimal("64441.3769"), Decimal("644413769")) == 4
    assert _unit_scale_hint(Decimal("1.5"), Decimal("150000000")) == 8


def test_unit_scale_hint_never_crashes_on_extreme_exponents(tmp_path):
    from foundry.financial_evidence import _unit_scale_hint, evaluate_evidence_recipes

    assert _unit_scale_hint(Decimal("1E-300"), Decimal("1E+999999")) is None
    assert _unit_scale_hint(Decimal("1E-1000000"), Decimal("1000")) is None
    assert _unit_scale_hint(Decimal("1E+999995"), Decimal("1E+1000003")) == 8
    assert _unit_scale_hint(Decimal("Infinity"), Decimal("5")) is None

    analysis = AnalysisStore(TaskStore(tmp_path / "var").create(title="extreme scale"))
    dataset = _income_statement_dataset(tmp_path, analysis, value=1)
    run = _run(analysis, value="1000", suffix="1")
    draft = analysis.save_draft({"title": "极端", "body": "收入1000。", "run_ids": [run["id"]], "evidence_recipes": [_le_revenue_recipe(dataset, run, scale="1E-1000000")]})
    assert draft["evidence_report"]["summary"]["status"] == "needs_attention"
    assert "suggested_scale" not in _mismatch(draft, "le-revenue")

    # An exact unit gap on top of a huge existing scale: the total is built without context arithmetic.
    huge_run = _run(analysis, value="1E+1000007", suffix="2")
    report = evaluate_evidence_recipes(analysis, [_le_revenue_recipe(dataset, huge_run, scale="1E+999999")])
    assert report["issues"][0]["suggested_scale"] == "1" + "0" * 1000007


def test_consistency_text_source_can_read_archived_workbook_single_cell(tmp_path):
    from openpyxl import Workbook

    task = TaskStore(tmp_path / "var").create(title="workbook text source")
    analysis = AnalysisStore(task)
    source = tmp_path / "budget.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "预算"
    sheet["A32"] = "所有关键差异均超过10万元。"
    sheet["B2"] = 8
    sheet["C2"] = 12
    book.save(source)
    book.close()
    dataset = analysis.add_dataset(source, name="budget", metadata={"mode": "workbook"})
    run = _run(analysis, value="20", suffix="d")

    spec = {
        "title": "一致性",
        "body": "所有关键差异均超过10万元。",
        "run_ids": [run["id"]],
        "consistency_sources": {
            "a": {"kind": "number", "dataset_id": dataset["id"], "source_hash": dataset["source_hash"], "sheet": "预算", "range": "B2"},
            "b": {"kind": "number", "dataset_id": dataset["id"], "source_hash": dataset["source_hash"], "sheet": "预算", "range": "C2"},
            "assertion": {"kind": "text", "dataset_id": dataset["id"], "source_hash": dataset["source_hash"], "sheet": "预算", "range": "A32"},
        },
        "consistency_rules": [{
            "id": "all-material",
            "kind": "all_absolute_exceed",
            "threshold": "10",
            "asserted_result": True,
            "assertion_source_id": "assertion",
            "values": [{"id": "a", "source_id": "a"}, {"id": "b", "source_id": "b"}],
        }],
    }
    check = validate_explanation(analysis, {key: value for key, value in spec.items() if key != "title"})
    issue = next(item for item in check["issues"] if item["code"] == "financial_consistency_mismatch")
    assert issue["original_text"] == "所有关键差异均超过10万元。"
    assert issue["assertion_source"]["reference"]["sheet"] == "预算"
    assert issue["assertion_source"]["reference"]["range"] == "A32"
    assert check["status_dimensions"]["source_validation"]["status"] == "needs_attention"

    draft = analysis.save_draft(spec)
    assert draft["consistency_report"]["issues"][0]["original_text"] == "所有关键差异均超过10万元。"


def test_consistency_assertion_source_must_be_archived_text_not_numeric_label(tmp_path):
    from openpyxl import Workbook
    from foundry.financial_consistency import evaluate_consistency_rules
    from foundry.financial_evidence import prepare_consistency_inputs

    task = TaskStore(tmp_path / "var").create(title="consistency sources")
    analysis = AnalysisStore(task)
    source = tmp_path / "budget.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "预算"
    sheet["B2"] = 0.2
    book.save(source)
    book.close()
    dataset = analysis.add_dataset(source, name="budget", metadata={"mode": "workbook"})
    sources = {
        "numeric-cell": {
            "kind": "number",
            "dataset_id": dataset["id"],
            "source_hash": dataset["source_hash"],
            "sheet": "预算",
            "range": "B2",
            "label": "caller label is not original text",
        }
    }
    rules = [{
        "id": "fake-assertion",
        "kind": "all_absolute_exceed",
        "threshold": "0.10",
        "asserted_result": True,
        "assertion_source_id": "numeric-cell",
        "values": [{"id": "v", "source_id": "numeric-cell"}],
    }]

    resolved_rules, source_auth = prepare_consistency_inputs(analysis, rules, sources)

    assert source_auth["numeric-cell"]["kind"] == "number"
    assert "text" not in source_auth["numeric-cell"]
    with pytest.raises(FoundryError, match="authorized text source"):
        evaluate_consistency_rules(resolved_rules, source_auth=source_auth)


def test_quality_dimensions_do_not_pass_without_numeric_work(tmp_path):
    task = TaskStore(tmp_path / "var").create(title="dimensions")
    analysis = AnalysisStore(task)
    check = validate_explanation(analysis, {"body": "没有数字。"})
    assert check["status_dimensions"]["numeric_citation"]["status"] == "not_configured"
    assert check["status_dimensions"]["source_validation"]["status"] == "not_configured"
    draft = analysis.save_draft({"title": "空", "body": "没有数字。"})
    assert draft["quality"]["status_dimensions"]["numeric_citation"]["status"] == "not_configured"


def test_revision_duplicate_business_key_marks_summary_needs_attention(tmp_path):
    task = TaskStore(tmp_path / "var").create(title="revision duplicate")
    analysis = AnalysisStore(task)
    row_scope = {"period": "2026-01..2026-08", "organization": "集团", "scenario": "actual"}
    old_run = _run(analysis, value="325644", suffix="e", seq=1, row_extra=row_scope)
    duplicate_run = _new(
        analysis.task,
        id="f" * 32,
        analysis_id="a" * 32,
        engine="finance",
        execution_status="succeeded",
        reproducibility_status="unverified",
        bindings={"budget": "dataset-a"},
        input_hashes={"budget": "data-hash-a"},
        source_hashes={"budget": "source-hash-a"},
        finance_scope={"budget": {"period": "2026-01..2026-08", "scenario": "BUD", "organization": "集团", "unit": "万元", "currency": "CNY"}},
        result={
            "columns": [{"name": "metric"}, {"name": "actual"}],
            "rows": [
                {"metric_key": "revenue", "metric": "收入", "actual": "325644", "unit": "万元", "currency": "CNY", **row_scope},
                {"metric_key": "revenue", "metric": "收入", "actual": "325644", "unit": "万元", "currency": "CNY", **row_scope},
            ],
            "row_count": 2,
            "truncated": False,
        },
    )
    duplicate_run = analysis._write(analysis._object_path("runs", duplicate_run["id"]), duplicate_run)
    parent = analysis.save_draft({"title": "旧稿", "body": "收入。", "run_ids": [old_run["id"]]})
    revised = analysis.save_draft({"parent_id": parent["id"], "body": "收入。", "run_ids": [duplicate_run["id"]]})
    assert revised["revision_comparison"]["summary"]["status"] == "needs_attention"
    assert any(item["code"] == "revision_comparison_ambiguous" for item in revised["quality"]["issues"])


def test_recipe_and_consistency_count_caps_are_enforced(tmp_path):
    from foundry.financial_evidence import evaluate_evidence_recipes, prepare_consistency_inputs

    task = TaskStore(tmp_path / "var").create(title="caps")
    analysis = AnalysisStore(task)
    recipe = {"kind": "sum", "dataset_id": "0" * 32, "source_hash": "a" * 64, "sheet": "预算", "range": "A1", "claim": {"run_id": "1" * 32, "row": 0, "column": "actual"}}
    try:
        evaluate_evidence_recipes(analysis, [{**recipe, "id": f"r{i}"} for i in range(257)])
    except FoundryError as exc:
        assert "count limit" in str(exc)
    else:
        raise AssertionError("recipe count cap should be enforced")
    try:
        prepare_consistency_inputs(analysis, [{"id": f"r{i}", "kind": "all_absolute_exceed", "threshold": "1", "asserted_result": True, "assertion_source_id": "s", "values": []} for i in range(257)], {})
    except FoundryError as exc:
        assert "rule count limit" in str(exc)
    else:
        raise AssertionError("rule count cap should be enforced")

def test_source_count_and_excel_bounds_are_enforced(tmp_path):
    from foundry.financial_evidence import evaluate_evidence_recipes, prepare_consistency_inputs

    task = TaskStore(tmp_path / "var").create(title="bounds")
    analysis = AnalysisStore(task)
    try:
        prepare_consistency_inputs(analysis, [], {f"s{i}": {"kind": "text", "dataset_id": "0" * 32, "source_hash": "a" * 64, "paragraph": 1} for i in range(257)})
    except FoundryError as exc:
        assert "source count limit" in str(exc)
    else:
        raise AssertionError("consistency source count cap should be enforced")
    try:
        evaluate_evidence_recipes(analysis, [{"id": "bad-bounds", "kind": "sum", "dataset_id": "0" * 32, "source_hash": "a" * 64, "sheet": "预算", "range": "XFE1:XFE1", "claim": {"run_id": "1" * 32, "row": 0, "column": "actual"}}])
    except FoundryError as exc:
        assert "Excel worksheet bounds" in str(exc)
    else:
        raise AssertionError("Excel bounds should be enforced")

def test_analysis_and_script_persist_result_role_scope_to_runs(tmp_path):
    import csv

    task = TaskStore(tmp_path / "var").create(title="metadata")
    analysis = AnalysisStore(task)
    csv_path = tmp_path / "input.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["value"])
        writer.writerow(["1"])
    dataset = analysis.add_dataset(csv_path, name="input")
    saved = analysis.save_analysis({
        "question": "count rows",
        "sql": "select count(*) as actual from source",
        "bindings": {"source": dataset["id"]},
        "result_role": "probe",
        "scope": {"period": "2026-08", "organization": "集团"},
    })
    assert saved["result_role"] == "probe"
    assert saved["scope"]["period"] == "2026-08"
    run = analysis.run_analysis(saved["id"])
    assert run["result_role"] == "probe"
    assert run["scope"]["organization"] == "集团"

    script_run = analysis.run_script({
        "language": "python",
        "code": "print('ok')",
        "result_role": "activity",
        "scope": {"period": "2026-09"},
    })
    assert script_run["result_role"] == "activity"
    assert script_run["scope"] == {"period": "2026-09"}
    default_script_run = analysis.run_script({
        "language": "python",
        "code": "print('probe')",
    })
    assert default_script_run["result_role"] == "activity"


def test_projection_for_draft_resolves_recipe_and_run_datasets(tmp_path):
    from openpyxl import Workbook

    task = TaskStore(tmp_path / "var").create(title="projection sources")
    analysis = AnalysisStore(task)
    source = tmp_path / "budget.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "预算"
    sheet["B36"] = 3185
    book.save(source)
    book.close()
    dataset = analysis.add_dataset(source, name="华宸预算", metadata={"mode": "workbook"})
    run = _new(
        analysis.task,
        id="a" * 32,
        analysis_id="a" * 32,
        engine="finance",
        execution_status="succeeded",
        reproducibility_status="unverified",
        bindings={"budget": dataset["id"]},
        input_hashes={"budget": dataset["data_hash"]},
        source_hashes={"budget": dataset["source_hash"]},
        result={"columns": [{"name": "actual"}], "rows": [{"metric_key": "other_gains", "actual": "3185"}], "row_count": 1, "truncated": False},
    )
    run = analysis._write(analysis._object_path("runs", run["id"]), run)
    draft = analysis.save_draft({
        "title": "报告",
        "body": "其他收益预算3185。",
        "run_ids": [run["id"]],
        "evidence_recipes": [{"id": "direct", "kind": "direct", "dataset_id": dataset["id"], "source_hash": dataset["source_hash"], "sheet": "预算", "range": "B36", "claim": {"run_id": run["id"], "row": 0, "column": "actual"}}],
    })
    projection = analysis.projection_for_draft(draft["id"], "audit")
    assert any("华宸预算" in item and dataset["id"] in item for item in projection["sources"])

def test_runtime_cannot_promote_professional_review_and_edits_reset_review(tmp_path):
    from foundry.finance_service import FinanceStore

    task = TaskStore(tmp_path / "var").create(title="review authority")
    analysis = AnalysisStore(task)
    runtime = FinanceStore(task, session_id="runtime")
    try:
        runtime._save_explanation_draft({"body": "待复核", "review_status": "reviewed"})
    except FoundryError as exc:
        assert "professional review" in str(exc)
    else:
        raise AssertionError("runtime save should not promote review status")

    reviewed = analysis.save_draft({"title": "已复核", "body": "数字正确。", "review_status": "reviewed"})
    unchanged = validate_explanation(analysis, {"draft_id": reviewed["id"]})
    assert unchanged["status_dimensions"]["professional_review"]["status"] == "pass"
    spoofed = validate_explanation(analysis, {"draft_id": reviewed["id"], "body": "数字改了。", "review_status": "accepted"})
    assert spoofed["status_dimensions"]["professional_review"]["status"] == "unreviewed"
    child = analysis.save_draft({"parent_id": reviewed["id"], "body": "数字改了。", "review_status": "reviewed"})
    assert child["review_status"] == "unreviewed"


def test_revision_that_swaps_runs_does_not_inherit_facts_from_dropped_runs(tmp_path):
    """Huachen rerun: a child saved with only a new run_ids list kept 296 parent fact_ids -> 296 unknown_fact."""
    from foundry.finance_explanations import prepare_explanation

    task = TaskStore(tmp_path / "var").create(title="fact inheritance")
    analysis = AnalysisStore(task)
    old_run = _run(analysis, value="325644", suffix="1")
    new_run = _run(analysis, value="325644", suffix="2")
    old_fact_ids = [fact["id"] for fact in prepare_explanation(analysis, [old_run])["facts"]]
    assert old_fact_ids
    parent = analysis.save_draft({"title": "旧稿", "body": "收入完成情况。", "run_ids": [old_run["id"]], "fact_ids": old_fact_ids})

    swapped = analysis.save_draft({"parent_id": parent["id"], "body": "收入完成情况。", "run_ids": [new_run["id"]]})
    same_runs = analysis.save_draft({"parent_id": parent["id"], "body": "收入完成情况。"})

    assert swapped["fact_ids"] == []
    assert not [issue for issue in swapped["quality"]["issues"] if issue["code"] == "unknown_fact"]
    assert same_runs["fact_ids"] == old_fact_ids


def test_service_save_path_does_not_inherit_facts_from_dropped_runs(tmp_path):
    """Review of PR #16: the MCP/CLI save_explanation_draft path copied parent fact_ids itself, skipping the filter."""
    from foundry.finance_explanations import prepare_explanation
    from foundry.finance_service import FinanceStore

    task = TaskStore(tmp_path / "var").create(title="service fact inheritance")
    analysis = AnalysisStore(task)
    service = FinanceStore(task)

    def finance_run(suffix):
        record = _new(
            analysis.task, id=suffix * 32, analysis_id="a" * 32, engine="finance",
            finance_context_id="c" * 32, account_map_id=None, assumptions={},
            execution_status="succeeded", reproducibility_status="unverified",
            bindings={"budget": "d" * 32}, input_hashes={"budget": "data-hash-a"},
            source_hashes={"budget": "source-hash-a"},
            finance_scope={"budget": {"period": "2026-01..2026-08", "scenario": "BUD", "organization": "集团",
                                      "unit": "万元", "currency": "CNY"}},
            result={"columns": [{"name": "metric"}, {"name": "actual"}],
                    "rows": [{"metric_key": "revenue", "metric": "收入", "actual": "325644", "unit": "万元",
                              "currency": "CNY"}],
                    "row_count": 1, "truncated": False},
        )
        return analysis._write(analysis._object_path("runs", record["id"]), record)

    old_run, new_run = finance_run("1"), finance_run("2")
    old_ids = [fact["id"] for fact in prepare_explanation(analysis, [old_run])["facts"]]
    parent = service.save_explanation_draft({"title": "旧稿", "body": "收入完成情况。",
                                             "run_ids": [old_run["id"]], "fact_ids": old_ids})

    swapped = service.save_explanation_draft({"parent_id": parent["id"], "body": "收入完成情况。",
                                              "run_ids": [new_run["id"]], "claims": []})
    same_runs = service.save_explanation_draft({"parent_id": parent["id"], "body": "收入完成情况。"})

    assert not set(old_ids) & set(swapped["fact_ids"])
    assert not [issue for issue in swapped["quality"]["issues"] if issue["code"] == "unknown_fact"]
    assert set(old_ids) <= set(same_runs["fact_ids"])


def test_swapped_runs_with_malformed_inherited_fact_id_report_instead_of_crashing(tmp_path):
    """Review of PR #16: a stored non-string fact id made the new filter raise a bare TypeError."""
    task = TaskStore(tmp_path / "var").create(title="malformed fact id")
    analysis = AnalysisStore(task)
    old_run = _run(analysis, value="325644", suffix="1")
    new_run = _run(analysis, value="325644", suffix="2")
    parent = analysis.save_draft({"title": "旧稿", "body": "收入。", "run_ids": [old_run["id"]], "fact_ids": [{}]})

    child = analysis.save_draft({"parent_id": parent["id"], "body": "收入。", "run_ids": [new_run["id"]]})

    assert [issue["code"] for issue in child["quality"]["issues"]].count("unknown_fact") == 1
