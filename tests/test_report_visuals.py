from __future__ import annotations

import copy
import json
import shutil
import subprocess
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from foundry.analysis import AnalysisStore
from foundry.chart_deliverables import normalize_chart
from foundry.cli import main
from foundry.errors import FoundryError
from foundry.finance_service import FinanceStore
from foundry.report_documents import validate_report_document
from foundry.report_visuals import add_visual_worksheets, validate_visual, visual_height, visual_markdown
from foundry.report_visuals import ESTIMATE_NOTICE, mark_visual_estimates
from foundry.task import TaskStore


def document(visual):
    return {"schema_version": 1, "template_version": "annual-v1", "metadata": {},
            "content": [{"type": "reportVisual", "attrs": {"visual": visual}}]}


def quantitative(**changes):
    return {"version": 1, "kind": "bar", "title": "地区金额比较", "takeaway": "保留负值与缺失。",
            "chart_id": "a" * 32, "category": "region", "value_columns": ["amount", "budget"], "unit": "万元",
            "series": {"columns": [{"name": "region"}, {"name": "amount"}, {"name": "budget"}],
                       "rows": [{"region": "华东", "amount": "-2.1200000000000000001", "budget": "0"},
                                {"region": "西南", "amount": None, "budget": "5.00"}]},
            "run_id": "b" * 32, "snapshot_hash": "c" * 64, "source_label": "已核算地区金额", **changes}


def flow():
    return {"version": 1, "kind": "flow", "title": "业务跟进路径", "takeaway": "先核实，再行动。",
            "steps": [{"label": "访谈", "detail": "请负责人说明。", "connector": "确认后"},
                      {"label": "行动", "detail": "登记后续动作。", "connector": ""}]}


@pytest.fixture
def bound(tmp_path):
    task = TaskStore(tmp_path / "var").create(title="图表文稿")
    task.record_session("empty", role="analysis")
    store = AnalysisStore(task)
    source = tmp_path / "amounts.csv"
    source.write_text("region,amount,budget\n华东,-2.12,0\n西南,,5.00\n", encoding="utf8")
    dataset = store.add_dataset(source, name="地区金额")
    revision = store.save_analysis({"question": "比较地区金额", "sql": "select region, amount, budget from actual order by region",
                                    "bindings": {"actual": dataset["id"]}})
    run = store.run_analysis(revision["id"])
    chart = FinanceStore(task).chart_deliverable_save({"run_id": run["id"], "label": "地区金额原计算"})
    visual = quantitative(chart_id=chart["id"])
    return store, dataset, run, chart, visual


def test_visual_save_overwrites_client_proof_and_binds_claims(bound):
    store, dataset, run, chart, visual = bound
    visual.update(series="spoofed", run_id="foreign", snapshot_hash="fake", source_label="自行改名")
    draft = store.save_draft({"title": "本月分析", "document": document(visual)})
    saved = draft["document"]["content"][0]["attrs"]["visual"]
    assert saved["series"] == run["result"] or saved["series"] == {key: run["result"][key] for key in ("columns", "rows")}
    assert saved["run_id"] == run["id"]
    assert saved["snapshot_hash"] == chart["snapshot_hash"]
    assert saved["source_label"] == chart["label"]
    assert draft["run_ids"] == [run["id"]]
    assert len(draft["claims"]) == 3
    assert all(claim["verification_status"] == "matches_run" for claim in draft["claims"])
    assert dataset["id"] in json.dumps(draft["claims"])
    assert "缺失" in draft["body"]
    changed = copy.deepcopy(draft["document"])
    changed["content"][0]["attrs"]["visual"]["title"] = "地区差异"
    revision = store.save_draft({"parent_id": draft["id"], "document": changed})
    assert revision["content_hash"] != draft["content_hash"]
    assert store.get_draft(draft["id"])["document"] == draft["document"]


def test_visual_exports_use_frozen_data_even_after_original_snapshot_is_missing(bound):
    store, dataset, run, _, visual = bound
    draft = store.save_draft({"title": "本月分析", "document": document(visual)})
    projection = store.report_projection(draft["id"])
    shutil.rmtree(store.root / "datasets" / dataset["id"])
    assert store.report_projection(draft["id"])["document"]["content"][0] == projection["document"]["content"][0]
    exported = store.export_draft(draft["id"], "md")
    markdown = Path(exported["path"]).read_text()
    assert "-2.12" in markdown
    saved = draft["document"]["content"][0]["attrs"]["visual"]
    assert f"Figure: {saved['title']}; chart_id: {saved['chart_id']}; run_id: {run['id']}; snapshot_hash: {saved['snapshot_hash']}" in markdown
    assert saved["snapshot_hash"] not in markdown.split("\n---\n", 1)[0]


def test_visual_cannot_reference_other_task_or_outside_session(bound, tmp_path):
    store, _, _, _, visual = bound
    other = AnalysisStore(TaskStore(tmp_path / "other").create(title="其他任务"))
    with pytest.raises(FoundryError, match="unavailable"):
        other.save_draft({"title": "越界", "document": document(visual)})
    restricted = FinanceStore(store.task, snapshot_ids=[])
    with pytest.raises(FoundryError, match="outside this session"):
        restricted.save_explanation_draft({"title": "越界", "document": document(visual)})
    parent = store.save_draft({"title": "正常保存", "document": document(visual)})
    with pytest.raises(FoundryError, match="outside this session"):
        restricted.save_explanation_draft({"parent_id": parent["id"], "title": "越界继承"})


def test_analysis_cli_visual_only_input_cannot_widen_session(bound, capsys):
    store, dataset, _, _, visual = bound
    store.task.record_session("empty", role="analysis")
    store.task.update_session_context("empty", snapshot_ids=[])
    args = ["--var", str(store.task.store.var_root), "--session", "empty", "analysis-draft", "--spec",
            json.dumps({"title": "仅图表", "document": document(visual)})]
    assert main(args) == 3
    assert "outside this session" in capsys.readouterr().err
    store.task.update_session_context("empty", snapshot_ids=[dataset["id"]])
    assert main(args) == 0
    draft = json.loads(capsys.readouterr().out)
    assert draft["run_ids"]


def test_finance_service_visual_only_input_keeps_its_run_and_exact_claims(bound):
    store, dataset, run, _, visual = bound
    draft = FinanceStore(store.task, snapshot_ids=[dataset["id"]]).save_explanation_draft(
        {"title": "图表说明", "document": document(visual)})
    assert draft["run_ids"] == [run["id"]]
    assert len(draft["claims"]) == 3
    assert draft["document"]["content"][0]["type"] == "reportVisual"


def test_freezing_preserves_large_integer_and_float_display_through_json(bound, monkeypatch):
    store, dataset, _, _, _ = bound
    revision = store.save_analysis({"question": "保留数值原精度",
                                    "sql": "select 9007199254740993::BIGINT as region, 9007199254740995::BIGINT as amount, 1.0::DOUBLE as budget, 7::INTEGER as extra from actual limit 1",
                                    "bindings": {"actual": dataset["id"]}})
    run = store.run_analysis(revision["id"])
    chart = FinanceStore(store.task).chart_deliverable_save({"run_id": run["id"], "label": "原精度"})
    # SQL currently serializes numerics as text; other successful run producers may
    # supply JSON numbers. Exercise that allowed contract before the Electron boundary.
    native_run = copy.deepcopy(run)
    native_run["result"]["rows"][0].update(region=9007199254740993, amount=9007199254740995, budget=1.0, extra=7)
    original_get_run = store.get_run
    monkeypatch.setattr(store, "get_run", lambda run_id: copy.deepcopy(native_run) if run_id == run["id"] else original_get_run(run_id))
    draft = store.save_draft({"document": document(quantitative(chart_id=chart["id"]))})
    row = draft["document"]["content"][0]["attrs"]["visual"]["series"]["rows"][0]
    assert row["region"] == "9007199254740993"
    assert row["amount"] == "9007199254740995"
    assert row["budget"] == "1.0"
    assert row["extra"] == 7
    assert original_get_run(run["id"])["result"] == run["result"]
    assert "9007199254740993 | 9007199254740995 | 1.0" in draft["body"]
    assert all(claim["verification_status"] == "matches_run" for claim in draft["claims"])


def test_chart_open_adds_lossless_report_projection_without_changing_archive(bound, monkeypatch):
    store, _, _, chart, _ = bound
    service = FinanceStore(store.task)
    original_read = service._read
    numeric_chart = copy.deepcopy(chart)
    numeric_chart["series"]["rows"][0].update(region=9007199254740993, amount=9007199254740995, budget=1.0)
    monkeypatch.setattr(service, "_read", lambda kind, identity: copy.deepcopy(numeric_chart) if kind == "charts" else original_read(kind, identity))
    opened = service.chart_deliverable_open({"id": chart["id"]})
    assert opened["series"] == numeric_chart["series"]
    assert opened["content_hash"] == chart["content_hash"]
    assert opened["report_series"]["rows"][0] == {"region": "9007199254740993", "amount": "9007199254740995", "budget": "1.0"}
    assert original_read("charts", chart["id"])["series"] == chart["series"]


def test_xlsx_contains_native_bar_line_and_exact_values_flow_relations(bound):
    store, _, run, _, visual = bound
    doc = document(visual)
    doc["content"].extend([document({**visual, "kind": "line"})["content"][0], document(flow())["content"][0]])
    draft = store.save_draft({"title": "三张图", "document": doc})
    exported = store.export_draft(draft["id"], "xlsx")
    book = load_workbook(exported["path"])
    assert len(book["图表 1"]._charts) == len(book["图表 2"]._charts) == 1
    assert book["图表 1"]._charts[0].type == "bar"
    assert book["图表 1"].cell(8, 2).value == str(run["result"]["rows"][0]["amount"])
    assert book["图表 1"].cell(8, 2).data_type == "s"
    assert book["图表 3"].cell(6, 4).value == "确认后"
    with ZipFile(exported["path"]) as archive:
        charts = [name for name in archive.namelist() if name.startswith("xl/charts/chart")]
        assert len(charts) == 2
        assert b'<dispBlanksAs val="gap"' in archive.read(charts[1])


def test_visual_xlsx_writes_untrusted_labels_as_literal_text(tmp_path):
    visual = quantitative(title='=HYPERLINK("https://example.com")', takeaway="=1+1", source_label="=1+1")
    visual["series"]["rows"][0]["region"] = "=1+1"
    visual["category"] = "=region"
    visual["series"]["columns"][0]["name"] = "=region"
    for row in visual["series"]["rows"]:
        row["=region"] = row.pop("region")
    steps = flow()
    steps["steps"][0].update(label="=1+1", detail="=1+1", connector="=1+1")
    doc = document(visual)
    doc["content"].extend(document(steps)["content"])
    book = Workbook()
    add_visual_worksheets(book, doc)
    path = tmp_path / "literal.xlsx"
    book.save(path)
    reloaded = load_workbook(path)
    for sheet in reloaded:
        assert all(cell.data_type != "f" for row in sheet for cell in row)
    assert reloaded["图表 1"].cell(8, 1).value == "=1+1"
    assert reloaded["图表 1"].cell(14, 1).value == "=1+1"
    with ZipFile(path) as archive:
        assert all(b"<f>" not in archive.read(name) for name in archive.namelist() if name.startswith("xl/worksheets/"))


@pytest.mark.parametrize("values,minimum,maximum", [(["2", "3"], 0, None), (["-2", "-3"], None, 0), (["-2", "3"], None, None)])
def test_native_line_chart_includes_zero_and_has_deliberate_colors(values, minimum, maximum):
    visual = quantitative(kind="line", value_columns=["amount"])
    for row, value in zip(visual["series"]["rows"], values):
        row["amount"] = value
    book = Workbook()
    add_visual_worksheets(book, document(visual))
    chart = book["图表 1"]._charts[0]
    assert chart.y_axis.scaling.min == minimum
    assert chart.y_axis.scaling.max == maximum
    assert chart.series[0].graphicalProperties.line.solidFill.srgbClr == "8A6729"


def test_markdown_and_validation_keep_decimal_negative_zero_null_and_flow_text():
    visual = quantitative()
    checked = validate_report_document(document(visual))
    assert checked["content"][0]["attrs"]["visual"]["series"] == visual["series"]
    markdown = visual_markdown(visual)
    assert "-2.1200000000000000001" in markdown and "| 缺失 | 5.00 |" in markdown
    assert "0" in markdown
    assert "作者分析，不是计算证据" in visual_markdown(flow())
    unresolved = {key: value for key, value in visual.items() if key not in {"series", "run_id", "snapshot_hash", "source_label"}}
    assert validate_report_document(document(unresolved))["content"][0]["type"] == "reportVisual"


@pytest.mark.parametrize("bad", [True, "NaN", "Infinity", "1e1000", "1e-1000", "", " 3", {}, [1]])
def test_visual_rejects_invalid_numeric_values(bad):
    visual = quantitative()
    visual["series"]["rows"][0]["amount"] = bad
    with pytest.raises(FoundryError):
        validate_visual(visual)


@pytest.mark.parametrize("mutation", ["dense", "duplicate", "blank", "null_series", "absent", "bool_category", "duplicate_column"])
def test_visual_rejects_silent_loss_and_ambiguous_data(mutation):
    visual = quantitative()
    if mutation == "dense":
        visual["series"]["rows"] = [{"region": str(i), "amount": "1", "budget": "2"} for i in range(13)]
    elif mutation == "duplicate":
        visual["series"]["rows"][1]["region"] = " 华东 "
    elif mutation == "blank":
        visual["series"]["rows"][0]["region"] = " "
    elif mutation == "null_series":
        visual["series"]["rows"][0]["amount"] = None
    elif mutation == "absent":
        del visual["series"]["rows"][0]["amount"]
    elif mutation == "bool_category":
        visual["series"]["rows"][0]["region"] = True
    else:
        visual["value_columns"] = ["amount", "amount"]
    with pytest.raises(FoundryError):
        validate_visual(visual)


def test_flow_rejects_trailing_relation_and_visual_children():
    visual = flow()
    visual["steps"][-1]["connector"] = "未指向步骤"
    with pytest.raises(FoundryError, match="last connector"):
        validate_visual(visual)
    doc = document(flow())
    doc["content"][0]["content"] = [{"type": "text", "text": "hidden"}]
    with pytest.raises(FoundryError, match="child nodes"):
        validate_report_document(doc)


def test_chart_archive_rejects_unsuccessful_run():
    with pytest.raises(FoundryError, match="successful run"):
        normalize_chart({}, run={"execution_status": "failed"})


def test_visual_limit_and_escaped_markup():
    doc = document(flow())
    doc["content"] *= 21
    with pytest.raises(FoundryError, match="at most 20"):
        validate_report_document(doc)
    visual = quantitative(title="<script>alert(1)</script>")
    visual["series"]["rows"][0]["region"] = "<img src=x onerror=alert(1)>"
    markdown = visual_markdown(visual)
    assert "<script>" not in markdown and "<img" not in markdown
    assert "&lt;script&gt;" in markdown


def test_estimate_notice_is_selected_cell_based_and_preserved_in_projection():
    from foundry.report_projection import build_report_projection

    visual = quantitative()
    doc = document(copy.deepcopy(visual))
    key = (visual["run_id"], 0, "amount")
    assert mark_visual_estimates(doc, [{"cell": (visual["run_id"], 0, "unselected")}]) == 0
    assert mark_visual_estimates(doc, [{"cell": key}]) == 1
    assert ESTIMATE_NOTICE in visual_markdown(doc["content"][0]["attrs"]["visual"])
    draft = {"id": "d" * 32, "content_hash": "e" * 64, "document": document(visual),
             "claims": [{"basis": "calculation", "run_id": key[0], "row": 0, "column": key[2], "value": "-2.1200000000000000001"}],
             "explanation_check": {"requirements": {"items": [{"label": "地区金额", "estimated_grains": [
                 {"locations": [{"run_id": key[0], "row": 0, "column": key[2]}]}]}]}}}
    projection = build_report_projection(draft)
    frozen = next(node for node in projection["document"]["content"] if node["type"] == "reportVisual")["attrs"]["visual"]
    assert frozen["estimate_notice"] == ESTIMATE_NOTICE
    assert ESTIMATE_NOTICE in projection["body"]
    assert "estimate_notice" not in draft["document"]["content"][0]["attrs"]["visual"]


def test_forged_estimate_notice_does_not_survive_save(bound):
    store, _, _, _, visual = bound
    visual["estimate_notice"] = ESTIMATE_NOTICE
    draft = store.save_draft({"title": "图", "document": document(visual)})
    assert "estimate_notice" not in draft["document"]["content"][0]["attrs"]["visual"]


def test_python_and_typescript_markdown_are_identical():
    root = Path(__file__).resolve().parents[1]
    compiler = root / "app/node_modules/typescript"
    if not shutil.which("node") or not compiler.exists():
        pytest.skip("the existing frontend TypeScript runtime is not installed")
    visuals = [quantitative(), flow(), quantitative(kind="line", estimate_notice=ESTIMATE_NOTICE)]
    escaped = quantitative(title="利润 [核对] <b> & (备注)", takeaway="保留 [^n]，*重要*与`原值`。")
    escaped["series"]["rows"][0]["region"] = "![图](https://example.com/x) | A\\B\nC"
    escaped["series"]["rows"][0]["amount"] = "1.234567890123456789e-3"
    visuals.append(escaped)
    script = r"""
const fs = require('node:fs');
const ts = require(process.argv[1]);
const source = fs.readFileSync(process.argv[2], 'utf8');
const code = ts.transpileModule(source, { compilerOptions: { module: ts.ModuleKind.CommonJS, target: ts.ScriptTarget.ES2022 } }).outputText;
const output = {};
new Function('exports', code)(output);
const cases = JSON.parse(fs.readFileSync(0, 'utf8'));
process.stdout.write(JSON.stringify(cases.map(output.reportVisualMarkdown)));
"""
    result = subprocess.run(["node", "-e", script, str(compiler), str(root / "app/electron/reports/visuals.ts")],
                            input=json.dumps(visuals, ensure_ascii=False), text=True, capture_output=True, timeout=20)
    assert result.returncode == 0, result.stderr
    assert json.loads(result.stdout) == [visual_markdown(visual) for visual in visuals]


@pytest.mark.parametrize("case", json.loads((Path(__file__).parent / "fixtures/report_visual_contract.json").read_text()), ids=lambda case: case["name"])
def test_shared_visual_density_contract(case):
    assert visual_height(case["visual"]) == case["height"]
    if case["accepted"]:
        assert validate_visual(case["visual"])
    else:
        with pytest.raises(FoundryError):
            validate_visual(case["visual"])
