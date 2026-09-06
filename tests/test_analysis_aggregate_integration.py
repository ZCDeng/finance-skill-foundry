import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from foundry.analysis import AnalysisStore
from foundry.errors import FoundryError
from foundry.finance_service import FinanceStore
from foundry.task import TaskStore


def setup_analysis(tmp_path):
    task = TaskStore(tmp_path / 'var').create(title='来源追踪')
    store = AnalysisStore(task)
    path = tmp_path / 'source.json'
    path.write_text(json.dumps([
        {'region': '华东', 'amount': '100', 'currency': 'CNY', 'unit': '元'},
        {'region': '华南', 'amount': '100', 'currency': 'CNY', 'unit': '元'},
    ], ensure_ascii=False))
    data = store.add_dataset(path, name='actual')
    plan = {
        'dataset_alias': 'actual', 'group_by': ['region'],
        'filters': [{'field': 'region', 'op': 'eq', 'value': '华东'}],
        'measures': [{'name': 'revenue', 'field': 'amount', 'aggregation': 'sum'}],
        'dedupe': {'policy': 'none'},
    }
    return store, data, plan


def test_structured_run_rerun_and_real_source_drilldown(tmp_path):
    store, data, plan = setup_analysis(tmp_path)
    definition = store.save_analysis({
        'question': '华东收入', 'engine': 'structured_aggregate',
        'bindings': {'actual': data['id']}, 'aggregate_plan': plan,
    })
    run = store.run_analysis(definition['id'])
    assert run['execution_status'] == 'succeeded'
    assert run['result']['rows'][0]['revenue'] == '100'
    assert run['engine'] == 'structured_aggregate'
    assert run['input_hashes']['actual'] == data['data_hash']
    assert store.run_analysis(definition['id'])['reproducibility_status'] == 'reproducible'
    finance = FinanceStore(store.task)
    trace = finance.drilldown_fact({'run_id': run['id'], 'row': 0, 'column': 'revenue'})
    assert trace['provenance_status'] == 'verified'
    assert [r['data']['region'] for r in trace['rows']] == ['华东']
    excluded = finance.drilldown_fact({
        'run_id': run['id'], 'row': 0, 'section': 'excluded_detail_rows',
    })
    assert [r['data']['region'] for r in excluded['rows']] == ['华南']
    draft = store.save_draft({'body': '华东收入计算与来源。', 'run_ids': [run['id']]})
    markdown = Path(store.export_draft(draft['id'], 'md')['path']).read_text()
    assert '来源核验：verified' in markdown
    assert 'aggregate' in markdown and data['source_hash'] in markdown
    workbook = load_workbook(store.export_draft(draft['id'], 'xlsx')['path'])
    assert {'Result 1', 'Evidence 1', 'Sources 1', 'Excluded 1'} <= set(workbook.sheetnames)
    assert workbook['Sources 1'].max_row == 2
    assert workbook['Excluded 1'].max_row == 2


def test_aggregate_definition_cannot_smuggle_other_bindings_or_sql(tmp_path):
    store, data, plan = setup_analysis(tmp_path)
    spec = {'question': '收入', 'engine': 'structured_aggregate',
            'bindings': {'actual': data['id']}, 'aggregate_plan': plan}
    with pytest.raises(FoundryError):
        store.save_analysis({**spec, 'sql': 'SELECT 100 AS revenue'})
    with pytest.raises(FoundryError):
        store.save_analysis({**spec, 'bindings': {'other': data['id']}})
    with pytest.raises(FoundryError):
        store.save_analysis({**spec, 'engine': 'made_up_engine'})
    with pytest.raises(FoundryError):
        store.save_analysis({**spec, 'aggregate_plan': {**plan, 'dataset_alias': []}})


def test_free_sql_same_amount_does_not_gain_verified_sources(tmp_path):
    store, data, _ = setup_analysis(tmp_path)
    definition = store.save_analysis({'question': '临时测算',
        'bindings': {'actual': data['id']}, 'sql': 'SELECT 100 AS revenue'})
    run = store.run_analysis(definition['id'])
    assert run['execution_status'] == 'succeeded'
    trace = FinanceStore(store.task).drilldown_fact({'run_id': run['id'], 'row': 0, 'column': 'revenue'})
    assert trace['provenance_status'] != 'verified'


def test_explanation_facts_keep_each_group_currency_and_unit(tmp_path):
    store, _, plan = setup_analysis(tmp_path)
    path = tmp_path / 'multi_currency.json'
    path.write_text(json.dumps([
        {'region': '华东', 'amount': '100', 'currency': 'CNY', 'unit': '元'},
        {'region': '海外', 'amount': '100', 'currency': 'USD', 'unit': '美元'},
    ], ensure_ascii=False))
    data = store.add_dataset(path, name='mixed')
    definition = store.save_analysis({
        'question': '各地收入', 'engine': 'structured_aggregate',
        'bindings': {'actual': data['id']}, 'aggregate_plan': {**plan, 'filters': []},
    })
    run = store.run_analysis(definition['id'])
    prepared = FinanceStore(store.task).prepare_explanation({'run_ids': [run['id']]})
    facts = {fact['scope']['region']: fact for fact in prepared['facts'] if fact['column'] == 'revenue'}
    assert facts['华东']['scope']['currency'] == 'CNY'
    assert facts['海外']['scope']['currency'] == 'USD'
    assert facts['海外']['unit'] == '美元'
    assert facts['海外']['source_trace']['provenance_status'] == 'verified'
