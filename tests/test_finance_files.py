from __future__ import annotations

import json
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo

from foundry.analysis import AnalysisStore
from foundry.errors import FoundryError
from foundry.finance_files import FinanceFiles
from foundry.task import TaskStore, sha256_bytes


def _source_book(path: Path, *, result: str = "old", amount: str = "10") -> Path:
    wb = Workbook()
    data = wb.active
    data.title = "Data"
    data.append(["account_id", "amount", "note"])
    data.append(["001234567890123456", Decimal(amount), "ok"])
    table = Table(displayName="InputTable", ref="A1:C2")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    data.add_table(table)
    out = wb.create_sheet("Output")
    out.append(["result", "comment"])
    out.append([result, ""])
    wb.defined_names.add(DefinedName("ResultCell", attr_text="'Output'!$A$2"))
    wb.save(path)
    wb.close()
    return path


@pytest.fixture
def files(tmp_path: Path):
    task = TaskStore(tmp_path / "var").create(title="finance files")
    analysis = AnalysisStore(task)
    dataset = analysis.add_dataset(
        _source_book(tmp_path / "source.xlsx"),
        name="source",
        metadata={"mode": "table"},
    )
    service = FinanceFiles(task)
    return service, analysis, dataset, tmp_path


def test_inspect_bind_and_read_use_authorized_snapshot_source(files) -> None:
    service, _, dataset, _ = files

    inspection = service.inspect({"dataset_id": dataset["id"]})
    binding = service.bind(
        {
            "dataset_id": dataset["id"],
            "source": {"kind": "table", "name": "InputTable"},
            "fields": [
                {"name": "account_id", "policy": "identifier", "unit": "account"},
                {"name": "amount", "policy": "amount", "unit": "CNY"},
                {"name": "note", "policy": "text"},
            ],
        }
    )
    read = service.read({"binding_id": binding["id"]})

    assert inspection["source_hash"] == dataset["source_hash"]
    assert "InputTable" in inspection["tables"]
    assert read["rows"] == [
        {"account_id": "001234567890123456", "amount": "10", "note": "ok"}
    ]
    assert read["diagnostics"] == []


def test_preview_and_apply_persist_immutable_patch_output_and_original(files) -> None:
    service, analysis, dataset, _ = files
    run = analysis.run_analysis(
        analysis.save_analysis(
            {
                "question": "sum",
                "sql": "SELECT SUM(amount) AS result FROM actual",
                "bindings": {"actual": dataset["id"]},
            }
        )["id"]
    )
    binding = service.bind(
        {
            "dataset_id": dataset["id"],
            "source": {"kind": "named_range", "name": "ResultCell"},
            "fields": [{"name": "result", "policy": "amount", "unit": "CNY"}],
            "role": "output",
            "writable": True,
            "header": False,
        }
    )
    patch = service.preview(
        {
            "binding_id": binding["id"],
            "run_id": run["id"],
            "writes": [
                {
                    "row": 1,
                    "field": "result",
                    "value": "10",
                    "result_row": 0,
                    "result_column": "result",
                }
            ],
        }
    )

    output = service.apply({"patch_id": patch["id"]})
    original = load_workbook(analysis.root / "datasets" / dataset["id"] / "source.xlsx")
    result_book = load_workbook(output["path"], data_only=True)
    try:
        assert original["Output"]["A2"].value == "old"
        assert result_book["Output"]["A2"].value == 10
        assert isinstance(result_book["Output"]["A2"].value, int)
    finally:
        original.close()
        result_book.close()
    assert output["artifact_path"].endswith(".xlsx")
    assert output["sha256"] == sha256_bytes(Path(output["path"]).read_bytes())
    assert output["verification"][0]["status"] == "matches_run"
    stored = service._read("workbook_outputs", output["id"])
    assert stored["patch_id"] == patch["id"]
    assert stored["sha256"] == output["sha256"]


def test_manager_only_write_never_gets_false_verified_status(files) -> None:
    service, _, dataset, _ = files
    binding = service.bind(
        {
            "dataset_id": dataset["id"],
            "source": {"kind": "named_range", "name": "ResultCell"},
            "fields": [{"name": "result", "policy": "amount"}],
            "role": "output",
            "writable": True,
            "header": False,
        }
    )

    patch = service.preview(
        {
            "binding_id": binding["id"],
            "writes": [{"row": 1, "field": "result", "value": "15"}],
        }
    )

    assert patch["verification"] == [
        {"write_index": 0, "status": "manager_write_unverified", "reason": "no run_id supplied"}
    ]


def test_run_verified_write_must_match_actual_result(files) -> None:
    service, analysis, dataset, _ = files
    revision = analysis.save_analysis(
        {
            "question": "sum",
            "sql": "SELECT SUM(amount) AS result FROM actual",
            "bindings": {"actual": dataset["id"]},
        }
    )
    run = analysis.run_analysis(revision["id"])
    binding = service.bind(
        {
            "dataset_id": dataset["id"],
            "source": {"kind": "named_range", "name": "ResultCell"},
            "fields": [{"name": "result", "policy": "amount"}],
            "role": "output",
            "writable": True,
            "header": False,
        }
    )

    with pytest.raises(FoundryError, match="differs"):
        service.preview(
            {
                "binding_id": binding["id"],
                "run_id": run["id"],
                "writes": [
                    {
                        "row": 1,
                        "field": "result",
                        "value": "999",
                        "result_row": 0,
                        "result_column": "result",
                    }
                ],
            }
        )


def test_session_ceiling_and_cross_task_snapshots_are_refused(files, tmp_path: Path) -> None:
    service, _, dataset, _ = files
    restricted = FinanceFiles(service.task, snapshot_ids=[])
    other_task = TaskStore(tmp_path / "other-var").create(title="other")
    other_analysis = AnalysisStore(other_task)
    other_dataset = other_analysis.add_dataset(
        _source_book(tmp_path / "other.xlsx"),
        name="other",
        metadata={"mode": "table"},
    )

    with pytest.raises(FoundryError, match="outside this session context"):
        restricted.inspect({"dataset_id": dataset["id"]})
    with pytest.raises(FoundryError, match="integrity|unavailable"):
        service.inspect({"dataset_id": other_dataset["id"]})


def test_saved_patch_tamper_is_rejected_by_analysis_integrity(files) -> None:
    service, _, dataset, _ = files
    binding = service.bind(
        {
            "dataset_id": dataset["id"],
            "source": {"kind": "named_range", "name": "ResultCell"},
            "fields": [{"name": "result", "policy": "amount"}],
            "role": "output",
            "writable": True,
            "header": False,
        }
    )
    patch = service.preview(
        {
            "binding_id": binding["id"],
            "writes": [{"row": 1, "field": "result", "value": "10"}],
        }
    )
    patch_path = service._path("workbook_patches", patch["id"])
    raw = json.loads(patch_path.read_text())
    raw["patch"]["cells"][0]["new_value"] = 999
    patch_path.write_text(json.dumps(raw))

    with pytest.raises(FoundryError, match="integrity"):
        service.apply({"patch_id": patch["id"]})


def test_newer_snapshot_uses_old_baseline_for_conflict_choices(files, tmp_path: Path) -> None:
    service, analysis, dataset, _ = files
    binding = service.bind(
        {
            "dataset_id": dataset["id"],
            "source": {"kind": "named_range", "name": "ResultCell"},
            "fields": [{"name": "result", "policy": "amount"}],
            "role": "output",
            "writable": True,
            "header": False,
        }
    )
    newer = analysis.add_dataset(
        _source_book(tmp_path / "newer.xlsx", result="manager edit"),
        name="newer",
        metadata={"mode": "table"},
    )

    patch = service.preview(
        {
            "binding_id": binding["id"],
            "dataset_id": newer["id"],
            "writes": [{"row": 1, "field": "result", "value": "20"}],
        }
    )

    assert patch["patch"]["conflicts"][0]["old_value"] == "manager edit"
    kept = service.apply({"patch_id": patch["id"], "resolutions": {"A2": "keep"}})
    applied = service.apply({"patch_id": patch["id"], "resolutions": {"A2": "apply"}})
    keep_book = load_workbook(kept["path"])
    apply_book = load_workbook(applied["path"], data_only=True)
    try:
        assert keep_book["Output"]["A2"].value == "manager edit"
        assert apply_book["Output"]["A2"].value == 20
    finally:
        keep_book.close()
        apply_book.close()


def test_text_output_formula_literal_and_long_decimal_diagnostic(files) -> None:
    service, _, dataset, _ = files
    binding = service.bind(
        {
            "dataset_id": dataset["id"],
            "source": {"kind": "range", "sheet": "Output", "range": "A1:B2"},
            "fields": [
                {"name": "result", "policy": "amount"},
                {"name": "comment", "policy": "text"},
            ],
            "role": "output",
            "writable": True,
        }
    )

    patch = service.preview(
        {
            "binding_id": binding["id"],
            "writes": [
                {"row": 1, "field": "result", "value": "123456789012345.67"},
                {"row": 1, "field": "comment", "value": "=literal"},
            ],
        }
    )
    output = service.apply({"patch_id": patch["id"]})
    book = load_workbook(output["path"], data_only=False)
    try:
        assert book["Output"]["A2"].value == "123456789012345.67"
        assert book["Output"]["B2"].value == "=literal"
        assert book["Output"]["B2"].data_type == "s"
    finally:
        book.close()
    assert [diag["code"] for diag in patch["patch"]["diagnostics"]] == [
        "decimal_written_as_text",
        "formula_like_text",
    ]


def test_an_inspection_reply_carries_no_host_path(tmp_path: Path) -> None:
    """The reply went out with /Users/... in it."""
    import json as _json

    from foundry.analysis import AnalysisStore
    from foundry.finance_files import FinanceFiles
    from foundry.task import TaskStore
    from openpyxl import Workbook

    task = TaskStore(tmp_path / "var").create(title="检视")
    analysis = AnalysisStore(task)
    source = tmp_path / "book.xlsx"
    book = Workbook()
    book.active["A1"] = "科目代码"
    book.active["A2"] = "6001"
    book.save(source)
    book.close()
    dataset = analysis.add_dataset(source, name="册", metadata={"mode": "workbook"})

    reply = FinanceFiles(task).inspect({"dataset_id": dataset["id"]})
    assert "path" not in reply
    blob = _json.dumps(reply, ensure_ascii=False)
    assert str(tmp_path) not in blob and "/Users/" not in blob
    # What a caller actually needs is still there.
    assert reply["dataset_id"] == dataset["id"] and reply["sheet_names"]
