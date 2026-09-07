from __future__ import annotations

import hashlib
import io
import json
import os
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from foundry.analysis import AnalysisStore
from foundry.analysis_data import MAX_SOURCE_BYTES, parse_source
from foundry.errors import FoundryError
from foundry.task import TaskStore


@pytest.fixture
def work(tmp_path: Path) -> AnalysisStore:
    task = TaskStore(tmp_path / "var").create(title="财务套表导入")
    return AnalysisStore(task)


def _multi_sheet_cover(path: Path) -> Path:
    wb = Workbook()
    cover = wb.active
    cover.title = "00·封面"
    cover.merge_cells("A1:D1")
    cover["A1"] = "华宸物业月结财务套表"
    cover["A3"] = "编制基础"
    cover["B3"] = "=1+1"
    data = wb.create_sheet("04·合并利润表")
    data.append(["项目", "本月", "累计"])
    data.append(["营业收入", "1280.00", "32560.00"])
    wb.save(path)
    return path


def _single_sheet_table(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Actual"
    ws.append(["account_code", "amount"])
    ws.append(["1001", "1280.00"])
    wb.save(path)
    return path


def _multi_sheet_with_flat_first_sheet(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Input"
    ws.append(["account_code", "amount"])
    ws.append(["1001", "1280.00"])
    wb.create_sheet("Output")
    wb.save(path)
    return path


def test_multisheet_cover_xlsx_registers_as_complete_workbook_snapshot(
    work: AnalysisStore, tmp_path: Path
) -> None:
    source = _multi_sheet_cover(tmp_path / "month.xlsx")
    dataset = work.add_dataset(source, name="月结套表")
    read = work.read_dataset(dataset["id"])

    assert dataset["kind"] == "workbook"
    assert dataset["schema"] == []
    assert dataset["row_count"] == 0
    assert dataset["transformation"]["mode"] == "workbook_snapshot"
    assert dataset["transformation"]["sheet_count"] == 2
    assert read["workbook"]["sheets"] == [
        {"name": "00·封面", "rows": 3, "columns": 4},
        {"name": "04·合并利润表", "rows": 2, "columns": 3},
    ]
    archived = work.root / "datasets" / dataset["id"] / "source.xlsx"
    assert archived.read_bytes() == source.read_bytes()
    assert dataset["source_hash"] == hashlib.sha256(source.read_bytes()).hexdigest()


def test_single_sheet_xlsx_keeps_flat_table_import_behavior(
    work: AnalysisStore, tmp_path: Path
) -> None:
    dataset = work.add_dataset(_single_sheet_table(tmp_path / "actual.xlsx"), name="actual")
    read = work.read_dataset(dataset["id"])

    assert dataset["kind"] == "table"
    assert dataset["row_count"] == 1
    assert [column["name"] for column in dataset["schema"]] == ["account_code", "amount"]
    assert read["rows"] == [{"account_code": "1001", "amount": "1280.00"}]


def test_multisheet_xlsx_with_valid_first_sheet_still_keeps_complete_workbook(
    work: AnalysisStore, tmp_path: Path
) -> None:
    dataset = work.add_dataset(
        _multi_sheet_with_flat_first_sheet(tmp_path / "multi.xlsx"), name="多表输入"
    )

    assert dataset["kind"] == "workbook"
    assert dataset["row_count"] == 0
    assert dataset["transformation"]["sheet_count"] == 2


def test_multisheet_xlsx_can_be_forced_to_strict_table_mode(
    work: AnalysisStore, tmp_path: Path
) -> None:
    dataset = work.add_dataset(
        _multi_sheet_with_flat_first_sheet(tmp_path / "multi.xlsx"),
        name="Input",
        metadata={"mode": "table"},
    )

    assert dataset["kind"] == "table"
    assert work.read_dataset(dataset["id"])["rows"] == [
        {"account_code": "1001", "amount": "1280.00"}
    ]


def test_explicit_sheet_in_multisheet_xlsx_keeps_strict_table_import(
    work: AnalysisStore, tmp_path: Path
) -> None:
    source = _multi_sheet_cover(tmp_path / "month.xlsx")
    dataset = work.add_dataset(
        source, name="利润表", metadata={"sheet": "04·合并利润表"}
    )

    assert dataset["kind"] == "table"
    assert dataset["row_count"] == 1
    assert dataset["transformation"]["sheet"] == "04·合并利润表"
    assert work.read_dataset(dataset["id"])["rows"][0]["本月"] == "1280.00"


def test_workbook_snapshot_source_file_is_available_to_scripts(
    work: AnalysisStore, tmp_path: Path
) -> None:
    source = _multi_sheet_cover(tmp_path / "month.xlsx")
    dataset = work.add_dataset(source, name="月结套表")
    expected_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    code = f"""
import hashlib, json, os
from pathlib import Path
from openpyxl import load_workbook
inputs = Path(os.environ["FOUNDRY_INPUTS"])
meta = json.loads((inputs / "book.json").read_text())
raw = (inputs / "book.source.xlsx").read_bytes()
book = load_workbook(inputs / "book.source.xlsx", data_only=False)
assert meta["kind"] == "workbook"
assert hashlib.sha256(raw).hexdigest() == {expected_hash!r}
assert book["04·合并利润表"]["B2"].value == "1280.00"
assert book["00·封面"]["B3"].value == "=1+1"
Path("ok.json").write_text(json.dumps({{"sheets": book.sheetnames}}, ensure_ascii=False))
"""
    run = work.run_script(
        {
            "code": code,
            "language": "python",
            "bindings": {"book": dataset["id"]},
            "timeout_seconds": 30,
        }
    )
    if run.get("runner_status") == "unavailable":
        pytest.skip("OS-isolated runner unavailable on this platform")
    assert run["execution_status"] == "succeeded", run


def test_sql_rejects_workbook_snapshots_with_actionable_message(
    work: AnalysisStore, tmp_path: Path
) -> None:
    dataset = work.add_dataset(_multi_sheet_cover(tmp_path / "month.xlsx"), name="月结套表")
    analysis = work.save_analysis(
        {
            "question": "查营业收入",
            "bindings": {"book": dataset["id"]},
            "sql": "SELECT 1 AS value",
        }
    )
    run = work.run_analysis(analysis["id"])

    assert run["execution_status"] == "failed"
    assert "SQL analysis requires table snapshots" in run["error"]
    assert "workbook_inspect/workbook_read" in run["error"]


def test_workbook_snapshot_keeps_zip_expansion_guard() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/workbook.xml", b"0" * (MAX_SOURCE_BYTES * 4 + 1))

    with pytest.raises(FoundryError, match="expands beyond"):
        parse_source(buffer.getvalue(), ".xlsx", {"mode": "workbook"})


def test_user_property_workbooks_register_readonly_in_tmp(tmp_path: Path) -> None:
    fixture = os.environ.get("FOUNDRY_PROPERTY_WORKBOOK_FIXTURE")
    if not fixture:
        pytest.skip("set FOUNDRY_PROPERTY_WORKBOOK_FIXTURE to run local property workbook fixture")
    folder = Path(fixture)
    files = sorted(folder.glob("*.xlsx"))
    if len(files) != 2:
        pytest.skip("user property workbook fixture is not present on this machine")
    work = AnalysisStore(TaskStore(tmp_path / "var").create(title="真实财务套表"))

    imported = [work.add_dataset(path, name=path.stem) for path in files]

    assert [row["kind"] for row in imported] == ["workbook", "workbook"]
    assert [row["transformation"]["sheet_count"] for row in imported] == [14, 16]
    for row, source in zip(imported, files):
        read = work.read_dataset(row["id"])
        assert read["workbook"]["sheets"][0]["name"].startswith("00")
        assert (work.root / "datasets" / row["id"] / "source.xlsx").read_bytes() == source.read_bytes()
