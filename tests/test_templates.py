"""The 补充数据模板 ship as the generator writes them and read as tables a manager confirms like any file."""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from foundry.analysis import AnalysisStore
from foundry.cli import main
from foundry.data_templates import export_templates, template_names, templates_dir
from foundry.errors import FoundryError
from foundry.finance_service import FinanceStore
from foundry.task import TaskStore

GENERATOR = Path(__file__).parent / "fixtures" / "templates" / "build_templates.py"
TABLES = {
    "补充数据模板_在管面积台账.xlsx": "面积台账",
    "补充数据模板_物业服务合同台账.xlsx": "合同台账",
    "补充数据模板_人数.xlsx": "人数",
}


def _cells(path: Path) -> dict:
    book = load_workbook(path)
    return {(sheet.title, cell.coordinate): cell.value for sheet in book.worksheets for row in sheet.iter_rows() for cell in row
            if cell.value is not None}


def test_the_bundled_templates_are_what_the_generator_writes(tmp_path):
    subprocess.run([sys.executable, "-I", str(GENERATOR), str(tmp_path)], check=True, capture_output=True)
    assert template_names() == sorted(TABLES) == sorted(path.name for path in tmp_path.iterdir())
    for name in TABLES:
        assert _cells(tmp_path / name) == _cells(templates_dir() / name), f"{name}: regenerate foundry/resources/data_templates"


@pytest.mark.parametrize("name", sorted(TABLES))
def test_a_template_reads_as_one_table_with_its_cover_left_out(tmp_path, name):
    task = TaskStore(tmp_path / "var").create(title="补充数据")
    dataset = AnalysisStore(task).add_dataset(templates_dir() / name, name=name)
    store = FinanceStore(task)
    store.structure_propose(dataset["id"], source=store.actor)
    sheets = {sheet["sheet"]: sheet for sheet in store.structure_proposal_list({"dataset_id": dataset["id"]})["proposals"][0]["sheets"]}
    assert sheets["00·封面与填写说明"]["decision"] == "excluded"
    table = sheets[TABLES[name]]
    assert table["decision"] == "table" and not table.get("blocked_reason")
    assert table["columns"][0]["role"] == "organization"
    # Only the money ledger says what it is in; no column is read as an amount it is not.
    assert [column["role"] for column in table["columns"]].count("amount") == 0


def test_export_copies_the_templates_and_never_overwrites(tmp_path, capsys):
    target = tmp_path / "模板"
    assert main(["--var", str(tmp_path / "var"), "templates-export", "--dest", str(target)]) == 0
    assert json.loads(capsys.readouterr().out)["files"] == sorted(TABLES)
    with pytest.raises(FoundryError, match="new or empty folder"):
        export_templates(target)
