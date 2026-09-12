from __future__ import annotations

import io
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from foundry.analysis import AnalysisStore
from foundry.errors import FoundryError
from foundry.finance_service import FinanceStore
from foundry.report_documents import report_document_markdown, validate_report_document
from foundry.report_templates import apply_excel_named_ranges, bundled_templates, excel_brief_template_bytes
from foundry.task import TaskStore

SCOPE = {
    "ledger": "service-co",
    "organization": "华东",
    "period": "2026-08",
    "currency": "CNY",
    "unit": "万元",
    "scenario": "actual",
}


def _finance(tmp_path: Path):
    task = TaskStore(tmp_path / "var").create(title="公司模板")
    return FinanceStore(task), task


def test_cf16_management_brief_keeps_annual_v1_and_hashes_template_id(tmp_path: Path) -> None:
    service, _task = _finance(tmp_path)
    listed = {row["template_id"]: row for row in service.report_template_list({})["templates"]}
    assert listed["management-brief-v1"] == {"template_id": "management-brief-v1", "label": "经营简报", "kind": "word"}
    assert set(listed["management-brief-v1"]) == {"template_id", "label", "kind"}
    draft = service.report_template_apply({"template_id": "management-brief-v1", "scope": SCOPE})["draft"]
    document = draft["document"]
    assert document["template_version"] == "annual-v1"
    assert document["template_id"] == "management-brief-v1"
    assert "经营简报" in draft["body"]
    assert "账套：service-co" in draft["body"]
    assert "组织：华东" in draft["body"]
    assert "期间：2026-08" in draft["body"]
    assert "单位：万元" in draft["body"]
    assert "template_id" not in draft["body"]
    missing = validate_report_document(
        {
            "schema_version": 1,
            "template_version": "annual-v1",
            "template_id": "management-brief-v1",
            "metadata": {},
            "content": [{"type": "paragraph", "content": [{"type": "text", "text": "空页眉"}]}],
        }
    )
    markdown = report_document_markdown(missing)
    assert "账套：未声明" in markdown
    assert "组织：未声明" in markdown
    clone = dict(document)
    clone["template_id"] = "variance-note-v1"
    changed = _finance(tmp_path)[1]
    other = AnalysisStore(TaskStore(tmp_path / "var").create(title="哈希对照"))
    first = other.save_draft({"title": "经营简报", "document": document})
    second = other.save_draft({"title": "经营简报", "document": clone})
    assert first["content_hash"] != second["content_hash"]
    assert "template_id" in first["document"]


G1_EXCEL_VALUES = {"Profit": "30", "Volume": "20", "Price": "-10", "Mix": "-5"}


def test_cf17_excel_named_ranges_write_g1_amounts_as_summable_numbers() -> None:
    data = apply_excel_named_ranges(excel_brief_template_bytes(), G1_EXCEL_VALUES)
    workbook = load_workbook(io.BytesIO(data))
    sheet = workbook["Brief"]
    assert sheet["B2"].value == Decimal("30")
    assert sheet["B2"].data_type == "n"
    assert sheet["B3"].value == Decimal("20")
    assert sheet["B4"].value == Decimal("-10")
    assert sheet["B5"].value == Decimal("-5")
    assert sheet["B6"].value == "=SUM(B2:B5)"
    workbook.close()


def test_cf18_apply_writes_new_draft_and_rejects_unknown_ids(tmp_path: Path) -> None:
    service, task = _finance(tmp_path)
    analysis = AnalysisStore(task)
    original = analysis.save_draft(
        {
            "title": "旧稿",
            "document": {
                "schema_version": 1,
                "template_version": "annual-v1",
                "metadata": {"period": "2026-01"},
                "content": [{"type": "paragraph", "content": [{"type": "text", "text": "旧正文留着"}]}],
            },
        }
    )
    before = dict(original)
    applied = service.report_template_apply({"template_id": "management-brief-v1", "scope": SCOPE})["draft"]
    assert applied["id"] != original["id"]
    reread = analysis.get_draft(original["id"])
    assert reread["body"] == before["body"]
    assert reread["content_hash"] == before["content_hash"]
    assert "旧正文留着" in reread["body"]
    with pytest.raises(FoundryError, match="unknown report template"):
        service.report_template_apply({"template_id": "not-a-company-file"})


def test_cf18f_foreign_task_cannot_list_or_apply_published_snapshot(tmp_path: Path) -> None:
    service, task = _finance(tmp_path)
    source = tmp_path / "company.docx"
    source.write_bytes(_minimal_docx())
    published = service.report_template_publish({"path": str(source), "label": "本公司简报"})
    assert published["kind"] == "word"
    listed = service.report_template_list({})["templates"]
    assert any(row["template_id"] == published["template_id"] for row in listed)
    for row in listed:
        blob = " ".join(row.values())
        assert "http" not in blob
        assert "/" not in blob
        assert "path" not in row
        assert "bytes" not in row
    other = FinanceStore(TaskStore(tmp_path / "var").create(title="外任务"))
    foreign_ids = {row["template_id"] for row in other.report_template_list({})["templates"]}
    assert published["template_id"] not in foreign_ids
    with pytest.raises(FoundryError, match="unknown report template"):
        other.report_template_apply({"template_id": published["template_id"]})
    applied = service.report_template_apply({"template_id": published["template_id"], "scope": SCOPE})
    assert applied["draft"]["document"]["template_id"] == published["template_id"]


def test_cf34_publish_rejects_pptx(tmp_path: Path) -> None:
    service, _task = _finance(tmp_path)
    pptx = tmp_path / "deck.pptx"
    pptx.write_bytes(b"PK\x03\x04ppt/presentation.xml fake")
    with pytest.raises(FoundryError, match="pptx"):
        service.report_template_publish({"path": str(pptx)})


def test_bundled_templates_are_metadata_only() -> None:
    for row in bundled_templates():
        assert set(row) == {"template_id", "label", "kind"}


def _minimal_docx() -> bytes:
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types></Types>")
        archive.writestr("word/document.xml", "<w:document></w:document>")
    return buffer.getvalue()


def test_published_word_apply_cites_snapshot_bytes(tmp_path: Path) -> None:
    service, _task = _finance(tmp_path)
    source = tmp_path / "company.docx"
    source.write_bytes(_minimal_docx())
    published = service.report_template_publish({"path": str(source), "label": "本公司简报"})
    applied = service.report_template_apply({"template_id": published["template_id"], "scope": SCOPE})
    document = applied["draft"]["document"]
    sources = document.get("sources") or []
    assert sources
    assert sources[0]["format"] == "docx"
    assert "已套用已发布公司模板" in applied["draft"]["body"]
    assert "账套：service-co" in applied["draft"]["body"]


def _g1_run(service: FinanceStore, tmp_path: Path):
    import csv

    fixtures = Path(__file__).parent / "fixtures" / "finance_baseline" / "scenarios"
    with (fixtures / "service_ledger.csv").open(encoding="utf-8-sig", newline="") as handle:
        rows = [
            row
            for row in csv.DictReader(handle)
            if row["period"] == "2026-08" and row["organization"] == "华东" and row["scenario"] == "actual"
        ]
    path = tmp_path / "g1.csv"
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
    dataset = service.analysis.add_dataset(path, name="g1")
    context = service.save_context({"bindings": {"actual": dataset["id"]}, "scopes": {"actual": SCOPE}})
    return service.run_metric({"context_id": context["id"], "metric": {"kind": "management_pnl"}})


def _source_workbook_cells(service: FinanceStore, draft: dict) -> dict[str, object]:
    source_id = draft["document"]["sources"][0]["id"]
    record = service.analysis._read(service.analysis.root / "report_sources" / (source_id + ".json"))
    data = (service.analysis.root / "report_sources" / record["id"] / record["source_file"]).read_bytes()
    workbook = load_workbook(io.BytesIO(data))
    sheet = workbook.active
    cells = {name: sheet[name].value for name in ("B2", "B3", "B4", "B5", "B6")}
    workbook.close()
    return cells


def test_bundled_excel_apply_without_run_writes_no_fixture_numbers(tmp_path: Path) -> None:
    service, _task = _finance(tmp_path)
    applied = service.report_template_apply({"template_id": "management-brief-xlsx-v1", "scope": SCOPE})
    cells = _source_workbook_cells(service, applied["draft"])
    assert cells["B2"] in (0, 0.0, None)
    assert cells["B3"] in (0, 0.0, None)
    assert applied["filled_named_ranges"] == []


def test_published_excel_apply_fills_named_ranges_from_authorized_run(tmp_path: Path) -> None:
    service, task = _finance(tmp_path)
    run = _g1_run(service, tmp_path)
    published_path = tmp_path / "company.xlsx"
    published_path.write_bytes(excel_brief_template_bytes())
    datasets_before = len(service.analysis.state()["datasets"])
    published = service.report_template_publish({"path": str(published_path), "label": "本公司Excel"})
    applied = service.report_template_apply(
        {"template_id": published["template_id"], "scope": SCOPE, "run_id": run["id"]}
    )
    assert applied["filled_named_ranges"] == ["Profit"]
    assert set(applied["unfilled_named_ranges"]) == {"Volume", "Price", "Mix", "Total"}
    cells = _source_workbook_cells(service, applied["draft"])
    assert Decimal(str(cells["B2"])) == Decimal("30")
    assert cells["B6"] == "=SUM(B2:B5)"
    assert len(service.analysis.state()["datasets"]) == datasets_before
    with pytest.raises(FoundryError, match="run_id"):
        service.report_template_apply(
            {"template_id": published["template_id"], "scope": SCOPE, "values": {"Profit": "999"}}
        )
    narrow = FinanceStore(task, snapshot_ids=set())
    with pytest.raises(FoundryError):
        narrow.report_template_apply({"template_id": published["template_id"], "run_id": run["id"]})


def test_publish_rejects_oversized_template(tmp_path: Path, monkeypatch) -> None:
    from foundry import report_templates

    service, _task = _finance(tmp_path)
    source = tmp_path / "big.xlsx"
    source.write_bytes(excel_brief_template_bytes())
    monkeypatch.setattr(report_templates, "TEMPLATE_MAX_BYTES", 16)
    with pytest.raises(FoundryError, match="limit"):
        service.report_template_publish({"path": str(source), "label": "太大"})


def test_excel_apply_unwraps_report_source_and_keeps_published_amounts(tmp_path: Path) -> None:
    service, _task = _finance(tmp_path)
    bundled = service.report_template_apply({"template_id": "management-brief-xlsx-v1", "scope": SCOPE})
    assert bundled["draft"]["document"]["sources"]
    published_path = tmp_path / "company.xlsx"
    published_path.write_bytes(excel_brief_template_bytes())
    published = service.report_template_publish({"path": str(published_path), "label": "本公司Excel"})
    applied = service.report_template_apply({"template_id": published["template_id"], "scope": SCOPE})
    source_id = applied["draft"]["document"]["sources"][0]["id"]
    from openpyxl import load_workbook
    record = service.analysis._read(service.analysis.root / "report_sources" / (source_id + ".json"))
    data = (service.analysis.root / "report_sources" / record["id"] / record["source_file"]).read_bytes()
    workbook = load_workbook(io.BytesIO(data))
    sheet = workbook.active
    assert sheet["B2"].value in (0, 0.0, None)
    workbook.close()


def test_an_applied_template_names_its_source_after_the_template(tmp_path: Path) -> None:
    """The draft said tmppcj2w5jg.docx where 经营简报 belonged.

    The temporary file is how the bytes reach the archive, not what the manager
    applied.
    """
    from foundry.analysis import AnalysisStore
    from foundry.finance_service import FinanceStore
    from foundry.task import TaskStore

    task = TaskStore(tmp_path / "var").create(title="模板")
    service = FinanceStore(task, actor="manager")
    del AnalysisStore
    excel = next(
        item
        for item in service.report_template_list({})["templates"]
        if item["kind"] == "excel"
    )
    applied = service.report_template_apply({"template_id": excel["template_id"]})
    names = [
        source.get("name")
        for source in applied["draft"]["document"].get("sources") or []
    ]
    assert names == [f"{excel['label']}.xlsx"]
    assert not any(str(name).startswith("tmp") for name in names)
