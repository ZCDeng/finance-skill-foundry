"""滚动到下月: last month's workbook, this month's figures, written into a copy."""

from __future__ import annotations

import json
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Protection
from openpyxl.worksheet.formula import ArrayFormula

from foundry.analysis import AnalysisStore
from foundry.cli import main
from foundry.errors import FoundryError
from foundry.finance_files import FinanceFiles
from foundry.task import TaskStore, sha256_bytes


def _august(path: Path, *, chart: bool = False) -> Path:
    """Last month's pack: inputs, formulas, a merge, an array formula, an error, a protected sheet."""
    book = Workbook()
    report = book.active
    report.title = "利润表"
    report["A1"] = "华东公司利润表 2026年08月"
    report["A4"], report["B4"] = "项目", "本月"
    report["A5"], report["B5"] = "营业收入", 1000
    report["B5"].number_format = "#,##0.00;[Red]-#,##0.00"
    report["A6"], report["B6"] = "营业成本", 600
    report["A7"], report["B7"] = "毛利", "=B5-B6"
    report["C5"] = ArrayFormula("C5:C6", "=B5:B6*2")
    report.merge_cells("D5:E6")
    report["D5"] = "说明"
    report["F5"] = "#N/A"
    report["G5"] = "上月备注"
    report["H5"] = 12.5
    locked = book.create_sheet("锁定表")
    locked["A1"] = 1
    locked["B1"] = 2
    locked["B1"].protection = Protection(locked=False)
    locked.protection.sheet = True
    if chart:
        figure = BarChart()
        figure.add_data(Reference(report, min_col=2, min_row=5, max_row=6))
        report.add_chart(figure, "J2")
    book.save(path)
    book.close()
    return path


def _september(path: Path) -> Path:
    """This month's balances: exact amounts, a text cell, a formula with no cached result, a date."""
    from datetime import date

    book = Workbook()
    sheet = book.active
    sheet.title = "科目余额"
    sheet["A1"], sheet["B1"] = "主营业务收入", Decimal("1234.56")
    sheet["A2"], sheet["B2"] = "其他业务收入", Decimal("0.1")
    sheet["A3"], sheet["B3"] = "主营业务成本", Decimal("0.2")
    sheet["A4"], sheet["B4"] = "备注", "未结账"
    sheet["B5"] = "=B1*2"
    sheet["B6"] = date(2026, 9, 30)
    sheet["B7"] = "-抵消后"
    book.save(path)
    book.close()
    return path


@pytest.fixture
def pack(tmp_path: Path):
    task = TaskStore(tmp_path / "var").create(title="九月月结", period="2026-09")
    analysis = AnalysisStore(task)
    august = analysis.add_dataset(_august(tmp_path / "八月套表.xlsx"), name="八月套表", metadata={"mode": "workbook"})
    balances = analysis.add_dataset(_september(tmp_path / "九月余额.xlsx"), name="九月余额", metadata={"mode": "workbook"})
    return task, analysis, august, balances, tmp_path


def _cell(balances: dict, cell_range: str, sheet: str = "科目余额") -> dict:
    return {
        "kind": "dataset",
        "dataset_id": balances["id"],
        "source_hash": balances["source_hash"],
        "sheet": sheet,
        "range": cell_range,
    }


def _map(sheet: str, cell: str, source: dict) -> dict:
    return {"target": {"sheet": sheet, "cell": cell}, "source": source}


def _preview(files: FinanceFiles, august: dict, mappings: list[dict], period: str = "2026-09") -> dict:
    return files.rollforward_preview(
        {"dataset_id": august["id"], "source_hash": august["source_hash"], "period": period, "mappings": mappings}
    )


def _by_target(plan: dict) -> dict[str, dict]:
    return {row["target"]: row for row in plan["rows"]}


def _archived(analysis: AnalysisStore, dataset: dict) -> Path:
    return analysis.root / "datasets" / dataset["id"] / "source.xlsx"


def test_every_unsafe_target_is_listed_with_its_reason_and_never_written(pack) -> None:
    task, analysis, august, balances, _ = pack
    files = FinanceFiles(task, actor="manager")
    plan = _preview(files, august, [
        _map("利润表", "B7", _cell(balances, "B1")),
        _map("利润表", "C5", _cell(balances, "B1")),
        _map("利润表", "C6", _cell(balances, "B1")),
        _map("利润表", "E6", _cell(balances, "B1")),
        _map("没有这张表", "B5", _cell(balances, "B1")),
        _map("利润表", "B5:B6", _cell(balances, "B1")),
        _map("", "B5", _cell(balances, "B1")),
        _map("锁定表", "A1", _cell(balances, "B1")),
        _map("锁定表", "B1", _cell(balances, "B1")),
        _map("利润表", "F5", _cell(balances, "B1")),
        _map("利润表", "G5", _cell(balances, "B4:B4")),
        _map("利润表", "H5", _cell(balances, "B1")),
        _map("利润表", "$h$5", _cell(balances, "B2")),
    ])
    rows = _by_target(plan)
    reasons = {target: row.get("reason_code") for target, row in rows.items()}
    assert reasons["利润表!B7"] == "formula_cell"
    assert reasons["利润表!C5"] == "formula_cell"  # the array formula's anchor is a formula too
    assert reasons["利润表!C6"] == "array_formula_member"
    assert reasons["利润表!E6"] == "merged_non_anchor"
    assert reasons["没有这张表!B5"] == "missing_sheet"
    assert reasons["利润表!B5:B6"] == "bad_address"
    assert reasons["!B5"] == "bad_address"
    assert reasons["锁定表!A1"] == "protected_cell"
    assert rows["锁定表!B1"]["status"] == "write"  # unlocked on a protected sheet
    assert reasons["利润表!F5"] == "error_cell"
    assert reasons["利润表!G5"] is None and rows["利润表!G5"]["new_value"] == "未结账"
    # The same cell mapped twice: neither row wins.
    assert reasons["利润表!H5"] == "duplicate_target"
    assert [row["status"] for row in plan["rows"] if row["cell"] == "H5"] == ["blocked", "blocked"]
    for row in plan["rows"]:
        if row["status"] == "blocked":
            assert row["reason"], row
    # A blocked target still shows what its source holds, so it can be typed by hand.
    assert rows["利润表!B7"]["new_value"] == "1234.56"
    assert rows["利润表!B7"]["old_value"] == "=B5-B6"
    assert plan["counts"] == {"write": 2, "blocked": 11, "unchanged": 0}

    run = files.rollforward_apply({"plan_id": plan["plan_id"]})
    copy = load_workbook(_archived(analysis, {"id": run["output_dataset_id"]}))
    try:
        report = copy["利润表"]
        assert report["B7"].value == "=B5-B6"
        assert isinstance(report["C5"].value, ArrayFormula)
        assert report["C6"].value is None
        assert report["F5"].value == "#N/A"
        assert report["H5"].value == 12.5
        assert report["G5"].value == "未结账"
        assert copy["锁定表"]["A1"].value == 1
        assert copy["锁定表"]["B1"].value == 1234.56
    finally:
        copy.close()


def test_a_written_cell_keeps_its_number_format_and_exact_sum(pack) -> None:
    task, analysis, august, balances, _ = pack
    files = FinanceFiles(task, actor="manager")
    plan = _preview(files, august, [
        _map("利润表", "B5", _cell(balances, "B1:B2")),
        _map("利润表", "B6", _cell(balances, "B2:B3")),
    ])
    rows = _by_target(plan)
    assert rows["利润表!B5"]["new_value"] == "1234.66"
    # 0.1 + 0.2 is 0.3 here, not 0.30000000000000004.
    assert rows["利润表!B6"]["new_value"] == "0.3"
    assert rows["利润表!B5"]["number_format"] == "#,##0.00;[Red]-#,##0.00"
    run = files.rollforward_apply({"plan_id": plan["plan_id"]})
    copy = load_workbook(_archived(analysis, {"id": run["output_dataset_id"]}))
    try:
        cell = copy["利润表"]["B5"]
        assert cell.value == 1234.66 and cell.data_type == "n"
        assert cell.number_format == "#,##0.00;[Red]-#,##0.00"
        assert copy["利润表"]["B6"].value == 0.3
        assert copy.calculation.fullCalcOnLoad is True
    finally:
        copy.close()
    assert run["formula_cache_cleared"] is True


def test_a_long_decimal_is_written_as_its_exact_digits_and_formula_like_text_stays_text(pack) -> None:
    task, analysis, august, balances, _ = pack
    files = FinanceFiles(task, actor="manager")
    plan = _preview(files, august, [
        _map("利润表", "B5", {"kind": "value", "value": "123456789012345678.12"}),
        _map("利润表", "G5", {"kind": "value", "value": "=HYPERLINK(\"http://x\")", "note": "备注"}),
        _map("利润表", "I5", _cell(balances, "B7")),
        _map("利润表", "B6", {"kind": "value", "value": "-25.5"}),
    ])
    rows = _by_target(plan)
    assert rows["利润表!B5"]["cell_type"] == "text"
    assert rows["利润表!B5"]["notes"] == ["decimal_written_as_text"]
    assert rows["利润表!G5"]["notes"] == ["formula_like_text"]
    assert rows["利润表!G5"]["source"] == {"kind": "value", "manager_value": "=HYPERLINK(\"http://x\")", "note": "备注"}
    assert rows["利润表!B6"]["cell_type"] == "number"
    run = files.rollforward_apply({"plan_id": plan["plan_id"]})
    copy = load_workbook(_archived(analysis, {"id": run["output_dataset_id"]}))
    try:
        report = copy["利润表"]
        assert report["B5"].value == "123456789012345678.12" and report["B5"].data_type == "s"
        assert report["B5"].number_format == "#,##0.00;[Red]-#,##0.00"
        assert report["G5"].value == "=HYPERLINK(\"http://x\")" and report["G5"].data_type == "s"
        assert report["I5"].value == "-抵消后" and report["I5"].data_type == "s"
        assert report["B6"].value == -25.5
    finally:
        copy.close()
    # The change log pastes into Excel without turning text into a formula.
    tsv = run["change_log"]["tsv"].splitlines()
    assert tsv[0] == "单元格\t原值\t新值\t来源\t状态/原因"
    assert any(line.startswith("利润表!G5\t上月备注\t'=HYPERLINK") for line in tsv)
    assert any(line.startswith("利润表!B6\t600\t-25.5\t") for line in tsv)
    # A figure Excel would round pastes as text, so the record shows what was written.
    assert any(line.startswith("利润表!B5\t1000\t'123456789012345678.12\t") for line in tsv)
    assert "| 利润表!B5 | 1000 | 123456789012345678.12 | 手填 | 已写入 |" in run["change_log"]["markdown"]


def test_the_original_snapshot_is_untouched_and_the_copy_is_a_new_snapshot(pack) -> None:
    task, analysis, august, balances, _ = pack
    before = sha256_bytes(_archived(analysis, august).read_bytes())
    files = FinanceFiles(task, actor="manager")
    plan = _preview(files, august, [
        _map("利润表", "B5", _cell(balances, "B1")),
        _map("利润表", "A1", {"kind": "value", "value": "华东公司利润表 2026年09月"}),
    ])
    run = files.rollforward_apply({"plan_id": plan["plan_id"]})
    assert sha256_bytes(_archived(analysis, august).read_bytes()) == before == august["source_hash"]
    output = analysis.read_dataset(run["output_dataset_id"])
    assert output["name"] == "八月套表（2026-09 滚动）"
    assert output["kind"] == "workbook"
    assert output["source_name"] == "八月套表（2026-09 滚动）.xlsx"
    assert output["metadata"] == {
        "mode": "workbook",
        "period": "2026-09",
        "rollforward": {"plan_id": plan["plan_id"], "source_dataset_id": august["id"], "source_hash": august["source_hash"]},
    }
    assert output["source_hash"] == run["output_sha256"]
    assert run["output_dataset_id"] in {row["id"] for row in analysis.state()["datasets"]}
    # No structure proposal is filed for a copy whose formulas have no cached results yet.
    assert not (analysis.root / "finance" / "structure_proposals").exists()
    # Nothing is left staged, and nothing in the record names a path.
    assert not (analysis.root / ".rollforward-staging").exists()
    stored = files._read("rollforward_runs", run["id"])
    assert stored["applied_by"] == "manager"
    assert stored["plan_id"] == plan["plan_id"]
    assert [row["target"] for row in stored["written"]] == ["利润表!B5", "利润表!A1"]
    assert str(task.store.var_root) not in json.dumps(stored, ensure_ascii=False)
    with pytest.raises(FoundryError, match="already applied"):
        files.rollforward_apply({"plan_id": plan["plan_id"]})


def test_an_unchanged_cell_is_not_written(pack) -> None:
    task, _, august, _, _ = pack
    plan = _preview(FinanceFiles(task), august, [
        _map("利润表", "B6", {"kind": "value", "value": "600"}),
        _map("利润表", "G5", {"kind": "value", "value": "上月备注"}),
    ])
    assert [row["status"] for row in plan["rows"]] == ["unchanged", "unchanged"]
    with pytest.raises(FoundryError, match="no cells to write"):
        FinanceFiles(task, actor="manager").rollforward_apply({"plan_id": plan["plan_id"]})


def test_sources_that_cannot_give_one_figure_are_blocked_with_a_reason(pack) -> None:
    task, analysis, august, balances, tmp_path = pack
    files = FinanceFiles(task)
    stale = {**_cell(balances, "B1"), "source_hash": "0" * 64}
    plan = _preview(files, august, [
        _map("利润表", "B5", _cell(balances, "B1:B4")),
        _map("利润表", "B6", _cell(balances, "B5")),
        _map("利润表", "G5", _cell(balances, "B6")),
        _map("利润表", "H5", _cell(balances, "B9")),
        _map("锁定表", "B1", stale),
        _map("利润表", "A10", {**_cell(balances, "B1"), "dataset_id": ""}),
        _map("利润表", "A11", _cell(balances, "B1", sheet="不存在")),
        _map("利润表", "A12", _cell(balances, "B:B")),
        _map("利润表", "A13", {"kind": "value", "value": "  "}),
        # The formula without a cached result blocks only the rows that read it.
        _map("利润表", "A14", _cell(balances, "B1")),
    ])
    rows = _by_target(plan)
    reasons = {target: (row.get("reason_code"), row.get("reason", "")) for target, row in rows.items()}
    assert reasons["利润表!B5"][0] == "source_unreadable" and "B4" in reasons["利润表!B5"][1] and "文字" in reasons["利润表!B5"][1]
    assert "没有保存计算结果" in reasons["利润表!B6"][1]
    assert "日期" in reasons["利润表!G5"][1]
    assert "空的" in reasons["利润表!H5"][1]
    assert "版本不一致" in reasons["锁定表!B1"][1]
    assert "还没有选择来源资料" in reasons["利润表!A10"][1]
    assert "没有工作表" in reasons["利润表!A11"][1]
    assert "工作表!B5" in reasons["利润表!A12"][1]
    assert "手填值是空的" in reasons["利润表!A13"][1]
    assert rows["利润表!A14"]["status"] == "write" and rows["利润表!A14"]["new_value"] == "1234.56"
    assert rows["利润表!B5"]["source"]["name"] == "九月余额"

    # A later upload under the same name replaces the snapshot a mapping points at.
    analysis.add_dataset(_september(tmp_path / "九月余额-更正.xlsx"), name="九月余额", metadata={"mode": "workbook"}, allow_duplicate=True)
    replaced = _preview(files, august, [_map("利润表", "B5", _cell(balances, "B1"))])
    assert "取代" in replaced["rows"][0]["reason"]


def test_apply_refuses_when_a_source_or_the_workbook_changed_since_preview(pack) -> None:
    task, analysis, august, balances, tmp_path = pack
    files = FinanceFiles(task, actor="manager")
    plan = _preview(files, august, [_map("利润表", "B5", _cell(balances, "B1"))])
    analysis.add_dataset(_september(tmp_path / "九月余额2.xlsx"), name="九月余额", metadata={"mode": "workbook"}, allow_duplicate=True)
    with pytest.raises(FoundryError, match="source changed since preview"):
        files.rollforward_apply({"plan_id": plan["plan_id"]})

    other = _preview(files, august, [_map("利润表", "B5", {"kind": "value", "value": "1"})])
    archived = _archived(analysis, august)
    archived.write_bytes(archived.read_bytes() + b"\0")
    with pytest.raises(FoundryError, match="integrity check failed"):
        files.rollforward_apply({"plan_id": other["plan_id"]})
    assert len(analysis.state()["datasets"]) == 3  # nothing was filed


def test_the_preview_refuses_a_workbook_version_it_was_not_given(pack) -> None:
    task, _, august, balances, _ = pack
    files = FinanceFiles(task)
    with pytest.raises(FoundryError, match="source_hash does not match"):
        files.rollforward_preview({"dataset_id": august["id"], "source_hash": "f" * 64, "period": "2026-09",
                                   "mappings": [_map("利润表", "B5", _cell(balances, "B1"))]})
    for period in ("2026-9", "2026-13", "", None, "本月"):
        with pytest.raises(FoundryError, match="period must be YYYY-MM"):
            _preview(files, august, [_map("利润表", "B5", _cell(balances, "B1"))], period=period)
    with pytest.raises(FoundryError, match="mappings must be a list"):
        _preview(files, august, [])
    with pytest.raises(FoundryError, match="mappings must be a list"):
        _preview(files, august, [_map("利润表", "B5", {"kind": "value", "value": "1"})] * 2001)
    with pytest.raises(FoundryError, match="source kind"):
        _preview(files, august, [_map("利润表", "B5", {"kind": "path", "path": "/etc/passwd"})])


def test_a_workbook_with_parts_a_copy_would_lose_gets_only_the_list(tmp_path: Path) -> None:
    task = TaskStore(tmp_path / "var").create(title="九月")
    analysis = AnalysisStore(task)
    august = analysis.add_dataset(_august(tmp_path / "带图表.xlsx", chart=True), name="带图表", metadata={"mode": "workbook"})
    files = FinanceFiles(task, actor="manager")
    plan = _preview(files, august, [_map("利润表", "B5", {"kind": "value", "value": "1"})])
    assert plan["workbook_supported"] is False
    assert "charts" in plan["unsupported_features"]
    assert plan["rows"][0]["status"] == "write"
    # No copy will be made, so nothing is "about to be written": it is hers to fill in.
    assert "需手工填写" in plan["change_log"]["tsv"] and "将写入" not in plan["change_log"]["tsv"]
    assert "不生成副本" in plan["change_log"]["markdown"]
    assert "需手工填写 1 格" in plan["change_log"]["markdown"]
    with pytest.raises(FoundryError, match="no copy was made"):
        files.rollforward_apply({"plan_id": plan["plan_id"]})
    assert [row["id"] for row in analysis.state()["datasets"]] == [august["id"]]


def test_only_the_manager_applies_or_saves_a_template(pack) -> None:
    task, analysis, august, balances, _ = pack
    assistant = FinanceFiles(task)
    plan = _preview(assistant, august, [_map("利润表", "B5", _cell(balances, "B1"))])
    with pytest.raises(FoundryError, match="only the manager can apply a workbook roll-forward"):
        assistant.rollforward_apply({"plan_id": plan["plan_id"]})
    with pytest.raises(FoundryError, match="only the manager can save a roll-forward template"):
        assistant.rollforward_template_save({"plan_id": plan["plan_id"], "name": "月度"})
    assert len(analysis.state()["datasets"]) == 2
    with pytest.raises(FoundryError, match="actor must be manager or assistant"):
        FinanceFiles(task, actor="admin")


def test_a_run_cell_is_read_by_its_column_type_inside_the_session_ceiling(tmp_path: Path) -> None:
    task = TaskStore(tmp_path / "var").create(title="九月")
    analysis = AnalysisStore(task)
    august = analysis.add_dataset(_august(tmp_path / "八月.xlsx"), name="八月", metadata={"mode": "workbook"})
    ledger = tmp_path / "ledger.csv"
    ledger.write_text("org,amount\n华东,100.10\n华东,0.15\n华南,7\n", encoding="utf8")
    table = analysis.add_dataset(ledger, name="明细")
    run = analysis.run_analysis(analysis.save_analysis({
        "question": "合计",
        "bindings": {"actual": table["id"]},
        "sql": "SELECT org, SUM(amount) AS amount FROM actual GROUP BY org ORDER BY org",
    })["id"])
    files = FinanceFiles(task, actor="manager")
    plan = _preview(files, august, [
        _map("利润表", "B5", {"kind": "run", "run_id": run["id"], "row": 0, "column": "amount"}),
        _map("利润表", "G5", {"kind": "run", "run_id": run["id"], "row": 1, "column": "org"}),
        _map("利润表", "H5", {"kind": "run", "run_id": run["id"], "row": 5, "column": "amount"}),
        _map("利润表", "A10", {"kind": "run", "run_id": run["id"], "row": 0, "column": "missing"}),
        _map("利润表", "A11", {"kind": "run", "run_id": "", "row": 0, "column": "amount"}),
    ])
    rows = _by_target(plan)
    assert rows["利润表!B5"]["new_value"] == "100.25" and rows["利润表!B5"]["cell_type"] == "number"
    assert rows["利润表!B5"]["source"]["run_id"] == run["id"]
    assert rows["利润表!G5"]["new_value"] == "华南" and rows["利润表!G5"]["cell_type"] == "text"
    assert "第 6 行" in rows["利润表!H5"]["reason"]
    assert "missing" in rows["利润表!A10"]["reason"]
    assert "还没有选择计算结果" in rows["利润表!A11"]["reason"]
    with pytest.raises(FoundryError, match="only the manager"):
        FinanceFiles(task, snapshot_ids={august["id"]}).rollforward_apply({"plan_id": plan["plan_id"]})
    # A session that holds the workbook but not the run's snapshot cannot read the run.
    narrowed = FinanceFiles(task, snapshot_ids={august["id"]})
    with pytest.raises(FoundryError, match="run bindings are outside this session context"):
        _preview(narrowed, august, [_map("利润表", "B5", {"kind": "run", "run_id": run["id"], "row": 0, "column": "amount"})])
    with pytest.raises(FoundryError, match="outside this session context"):
        _preview(FinanceFiles(task, snapshot_ids={table["id"]}), august, [_map("利润表", "B5", {"kind": "value", "value": "1"})])


def test_a_manager_session_sees_the_copy_it_made(pack) -> None:
    task, analysis, august, balances, _ = pack
    task.record_session("manager-session", role="analysis")
    task.update_session_context("manager-session", snapshot_ids=[august["id"], balances["id"]])
    files = FinanceFiles(task, snapshot_ids={august["id"], balances["id"]}, session_id="manager-session", actor="manager")
    plan = _preview(files, august, [_map("利润表", "B5", _cell(balances, "B1"))])
    run = files.rollforward_apply({"plan_id": plan["plan_id"]})
    session = next(row for row in task.sessions() if row["session_id"] == "manager-session")
    assert run["output_dataset_id"] in session["snapshot_ids"]


def test_the_title_still_naming_last_month_is_pointed_out(pack) -> None:
    task, _, august, balances, _ = pack
    files = FinanceFiles(task)
    plan = _preview(files, august, [_map("利润表", "B5", _cell(balances, "B1"))])
    assert plan["title_hints"] == [{
        "sheet": "利润表", "cell": "A1", "target": "利润表!A1", "text": "华东公司利润表 2026年08月",
        "period": "2026-08", "mapped": False,
    }]
    mapped = _preview(files, august, [_map("利润表", "A1", {"kind": "value", "value": "华东公司利润表 2026年09月"})])
    assert mapped["title_hints"][0]["mapped"] is True
    # A title that already names the new month is not a hazard.
    assert _preview(files, august, [_map("利润表", "B5", _cell(balances, "B1"))], period="2026-08")["title_hints"] == []


def test_a_template_keeps_the_mapping_without_this_months_snapshots(pack, tmp_path: Path) -> None:
    task, analysis, august, balances, _ = pack
    files = FinanceFiles(task, actor="manager")
    plan = _preview(files, august, [
        _map("利润表", "B5", _cell(balances, "B1:B2")),
        _map("利润表", "B6", {"kind": "run", "run_id": "a" * 32, "row": 0, "column": "amount"}),
        _map("利润表", "A1", {"kind": "value", "value": "华东公司利润表 2026年09月", "note": "标题"}),
        _map("利润表", "H5", {"kind": "value", "value": "98765.43", "note": "调整"}),
        _map("利润表", "G5", {"kind": "value", "value": "1,234.50"}),
    ])
    saved = files.rollforward_template_save({"plan_id": plan["plan_id"], "name": "华东月度套表"})
    assert saved["sheet_pattern"] == "利润表|锁定表"
    assert saved["mappings"] == [
        {"target": {"sheet": "利润表", "cell": "B5"}, "source": {"kind": "dataset", "sheet": "科目余额", "range": "B1:B2", "hint": "九月余额"}},
        {"target": {"sheet": "利润表", "cell": "B6"}, "source": {"kind": "run", "row": 0, "column": "amount"}},
        {"target": {"sheet": "利润表", "cell": "A1"}, "source": {"kind": "value", "note": "标题", "last_value": "华东公司利润表 2026年09月"}},
        # This month's typed amounts are not carried to other months' tasks.
        {"target": {"sheet": "利润表", "cell": "H5"}, "source": {"kind": "value", "note": "调整"}},
        {"target": {"sheet": "利润表", "cell": "G5"}, "source": {"kind": "value", "note": ""}},
    ]
    text = json.dumps(saved, ensure_ascii=False)
    for month_specific in (balances["id"], balances["source_hash"], "a" * 32, august["id"], "98765.43", "1,234.50"):
        assert month_specific not in text
    with pytest.raises(FoundryError, match="template name"):
        files.rollforward_template_save({"plan_id": plan["plan_id"], "name": "  "})

    # October's task sees September's template for the same pack; another pack does not.
    october = TaskStore(tmp_path / "var").create(title="十月月结", period="2026-10")
    october_store = AnalysisStore(october)
    rolled = october_store.add_dataset(_august(tmp_path / "九月套表.xlsx"), name="九月套表", metadata={"mode": "workbook"})
    other = october_store.add_dataset(_september(tmp_path / "余额.xlsx"), name="余额", metadata={"mode": "workbook"})
    later = FinanceFiles(october, actor="manager")
    listed = later.rollforward_templates({"dataset_id": rolled["id"]})
    assert [row["name"] for row in listed["templates"]] == ["华东月度套表"]
    assert listed["templates"][0]["mapping_count"] == 5
    assert later.rollforward_templates({"dataset_id": other["id"]})["templates"] == []
    assert len(later.rollforward_templates({})["templates"]) == 1
    registry = Path(task.store.var_root) / "rollforward_templates"
    record = next(path for path in registry.glob("*.json") if path.stem != "sequence")
    record.write_text(json.dumps({**json.loads(record.read_text()), "name": "改过的名字"}))
    with pytest.raises(FoundryError, match="integrity check failed"):
        later.rollforward_templates({})


def test_the_cli_previews_as_anyone_and_applies_only_as_the_manager(pack, capsys) -> None:
    task, analysis, august, balances, _ = pack
    root = task.store.var_root
    spec = {"dataset_id": august["id"], "source_hash": august["source_hash"], "period": "2026-09",
            "mappings": [_map("利润表", "B5", _cell(balances, "B1"))]}

    def call(*args: str) -> tuple[int, str, str]:
        code = main(["--var", str(root), "--task", task.id, *args])
        captured = capsys.readouterr()
        return code, captured.out, captured.err

    code, out, _ = call("workbook-rollforward-preview", "--spec", json.dumps(spec))
    assert code == 0
    plan = json.loads(out)
    assert plan["counts"] == {"write": 1, "blocked": 0, "unchanged": 0}
    code, _, err = call("workbook-rollforward-apply", "--spec", json.dumps({"plan_id": plan["plan_id"]}))
    assert code == 3 and "only the manager" in err
    code, out, _ = call("--actor", "manager", "workbook-rollforward-apply", "--spec", json.dumps({"plan_id": plan["plan_id"]}))
    assert code == 0
    applied = json.loads(out)
    assert applied["run_id"] == plan["plan_id"]
    assert applied["change_log"]["tsv"].startswith("单元格\t原值\t新值")
    assert str(root) not in out
    code, out, _ = call("--actor", "manager", "workbook-rollforward-template-save", "--spec", json.dumps({"plan_id": plan["plan_id"], "name": "月度"}))
    assert code == 0 and json.loads(out)["name"] == "月度"
    code, out, _ = call("workbook-rollforward-templates", "--spec", json.dumps({"dataset_id": august["id"]}))
    assert code == 0 and [row["name"] for row in json.loads(out)["templates"]] == ["月度"]


def test_the_assistant_is_not_offered_the_roll_forward() -> None:
    profile = json.loads(Path("dsh-plugin/package.json").read_text())["dsh"]["foundry"]["profiles"]["analysis"]
    assert not [tool for tool in [*profile["tools"], *profile["childTools"]] if "rollforward" in tool]
    assert "rollforward" not in Path("dsh-plugin/finance-tools.js").read_text()
    from foundry.finance_mcp import TOOL_METHODS

    assert not [name for name in TOOL_METHODS if "rollforward" in name]


def _kinds_book(path: Path) -> Path:
    """Cells whose kind a write could change: amounts with a format, a date, a blank amount cell."""
    from datetime import date

    book = Workbook()
    sheet = book.active
    sheet.title = "利润表"
    sheet["B5"], sheet["B6"], sheet["B7"] = 1000, 600, "=SUM(B5:B6)"
    sheet["B5"].number_format = "#,##0.00"
    sheet["C5"] = date(2026, 8, 31)
    sheet["C5"].number_format = "yyyy-mm-dd"
    sheet["D5"].number_format = "#,##0.00"  # a blank input cell set up for an amount
    sheet["E5"] = "备注"
    book.save(path)
    book.close()
    return path


def test_text_never_turns_a_number_cell_into_text_that_a_total_would_skip(tmp_path: Path) -> None:
    task = TaskStore(tmp_path / "var").create(title="九月")
    analysis = AnalysisStore(task)
    pack = analysis.add_dataset(_kinds_book(tmp_path / "套表.xlsx"), name="套表", metadata={"mode": "workbook"})
    export = Workbook()
    export.active.title = "TB"
    export.active["A1"] = "1,234.50"  # an ERP export that stores the amount as text
    export.save(tmp_path / "erp.xlsx")
    erp = analysis.add_dataset(tmp_path / "erp.xlsx", name="ERP", metadata={"mode": "workbook"})
    files = FinanceFiles(task, actor="manager")
    plan = _preview(files, pack, [
        _map("利润表", "B5", {"kind": "value", "value": "1,200"}),
        _map("利润表", "B6", _cell(erp, "A1", sheet="TB")),
        _map("利润表", "D5", {"kind": "value", "value": "不适用"}),
        _map("利润表", "C5", {"kind": "value", "value": "1500"}),
        _map("利润表", "C6", {"kind": "value", "value": "2026-09-30"}),
        _map("利润表", "E5", {"kind": "value", "value": "已结账"}),
        _map("利润表", "F5", {"kind": "value", "value": "(25.00)"}),
        _map("利润表", "D6", {"kind": "value", "value": "123456789012345678.12"}),
    ])
    rows = _by_target(plan)
    assert rows["利润表!B5"]["reason_code"] == "text_into_number_cell"
    assert "合计" in rows["利润表!B5"]["reason"] and "不带千分位" in rows["利润表!B5"]["reason"]
    assert rows["利润表!B6"]["reason_code"] == "text_into_number_cell"
    assert "按文字存的" in rows["利润表!B6"]["reason"]
    assert rows["利润表!D5"]["reason_code"] == "text_into_number_cell"
    assert "手工改" in rows["利润表!D5"]["reason"]
    assert rows["利润表!C5"]["reason_code"] == "number_into_date_cell"
    assert rows["利润表!C5"]["old_value"] == "2026-08-31T00:00:00"
    # A blank General cell, a text cell and a figure too long for Excel are still written.
    assert rows["利润表!C6"]["status"] == "write"
    assert rows["利润表!E5"]["status"] == "write"
    assert rows["利润表!F5"]["status"] == "write" and rows["利润表!F5"]["cell_type"] == "text"
    assert rows["利润表!D6"]["status"] == "write" and rows["利润表!D6"]["notes"] == ["decimal_written_as_text"]
    run = files.rollforward_apply({"plan_id": plan["plan_id"]})
    copy = load_workbook(_archived(analysis, {"id": run["output_dataset_id"]}))
    try:
        sheet = copy["利润表"]
        assert (sheet["B5"].value, sheet["B5"].data_type) == (1000, "n")
        assert (sheet["B6"].value, sheet["B6"].data_type) == (600, "n")
        assert sheet["C5"].is_date
    finally:
        copy.close()
    assert "未写入：目标原来是数字" in run["change_log"]["tsv"]


def _with_sparkline(path: Path) -> Path:
    """The x14 worksheet extension Excel writes for a sparkline; openpyxl reads past it and drops it."""
    import zipfile

    extension = (
        b'<extLst><ext uri="{05C60535-1F16-4fd2-B633-F4F36F0B64E0}" '
        b'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
        b'<x14:sparklineGroups xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main"><x14:sparklineGroup>'
        b'<x14:colorSeries rgb="FF376092"/><x14:sparklines><x14:sparkline><xm:f>\xe5\x88\xa9\xe6\xb6\xa6\xe8\xa1\xa8!B5:B6</xm:f>'
        b'<xm:sqref>I5</xm:sqref></x14:sparkline></x14:sparklines></x14:sparklineGroup></x14:sparklineGroups></ext></extLst>'
    )
    parts = {}
    with zipfile.ZipFile(path) as archive:
        for item in archive.infolist():
            parts[item.filename] = archive.read(item.filename)
    parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(b"</worksheet>", extension + b"</worksheet>")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return path


def test_a_workbook_with_sparklines_or_x14_extensions_gets_only_the_list(tmp_path: Path) -> None:
    task = TaskStore(tmp_path / "var").create(title="九月")
    analysis = AnalysisStore(task)
    august = analysis.add_dataset(_with_sparkline(_august(tmp_path / "迷你图.xlsx")), name="迷你图", metadata={"mode": "workbook"})
    files = FinanceFiles(task, actor="manager")
    plan = _preview(files, august, [_map("利润表", "B5", {"kind": "value", "value": "1"})])
    assert plan["workbook_supported"] is False
    assert plan["unsupported_features"] == ["extensions"]
    with pytest.raises(FoundryError, match="no copy was made"):
        files.rollforward_apply({"plan_id": plan["plan_id"]})
    assert [row["id"] for row in analysis.state()["datasets"]] == [august["id"]]


def test_a_chartsheet_is_neither_a_target_nor_a_source(tmp_path: Path) -> None:
    task = TaskStore(tmp_path / "var").create(title="九月")
    analysis = AnalysisStore(task)
    book = Workbook()
    sheet = book.active
    sheet.title = "数据"
    sheet["A1"], sheet["A2"] = 1, 2
    chartsheet = book.create_chartsheet("图表")
    figure = BarChart()
    figure.add_data(Reference(sheet, min_col=1, min_row=1, max_row=2))
    chartsheet.add_chart(figure)
    book.save(tmp_path / "图表.xlsx")
    charted = analysis.add_dataset(tmp_path / "图表.xlsx", name="图表", metadata={"mode": "workbook"})
    plan = _preview(FinanceFiles(task), charted, [
        _map("图表", "A1", {"kind": "value", "value": "1"}),
        _map("数据", "A1", _cell(charted, "A1", sheet="图表")),
    ])
    rows = _by_target(plan)
    assert rows["图表!A1"]["reason_code"] == "chart_sheet"
    assert rows["数据!A1"]["reason_code"] == "source_unreadable" and "图表工作表" in rows["数据!A1"]["reason"]


def test_a_second_apply_cannot_slip_in_before_the_first_is_recorded(pack) -> None:
    task, analysis, august, balances, _ = pack
    files = FinanceFiles(task, actor="manager")
    plan = _preview(files, august, [_map("利润表", "B5", _cell(balances, "B1"))])
    widen, refused = files._widen_session, []

    def meanwhile(snapshot_ids: list[str]) -> None:
        # The copy is registered and its record not yet filed: another window applies the same plan.
        try:
            FinanceFiles(task, actor="manager").rollforward_apply({"plan_id": plan["plan_id"]})
        except FoundryError as exc:
            refused.append(str(exc))
        widen(snapshot_ids)

    files._widen_session = meanwhile
    run = files.rollforward_apply({"plan_id": plan["plan_id"]})
    assert refused == ["rollforward plan is already being applied"]
    rolled = [row for row in analysis.state()["datasets"] if (row.get("metadata") or {}).get("rollforward")]
    assert [row["id"] for row in rolled] == [run["output_dataset_id"]]
    assert not (analysis.root / ".rollforward-staging").exists()


def test_apply_refuses_when_the_source_of_a_blocked_row_changed(pack) -> None:
    task, analysis, august, balances, tmp_path = pack
    files = FinanceFiles(task, actor="manager")
    plan = _preview(files, august, [
        _map("利润表", "B7", _cell(balances, "B1")),  # a formula: listed with the figure to type by hand
        _map("利润表", "B5", {"kind": "value", "value": "1"}),
    ])
    assert [row["status"] for row in plan["rows"]] == ["blocked", "write"]
    analysis.add_dataset(_september(tmp_path / "九月余额-更正.xlsx"), name="九月余额", metadata={"mode": "workbook"}, allow_duplicate=True)
    with pytest.raises(FoundryError, match="source changed since preview"):
        files.rollforward_apply({"plan_id": plan["plan_id"]})


def _monthly(path: Path, month: int) -> Path:
    book = Workbook()
    report = book.active
    report.title = f"{month}月利润表"
    report["B5"] = 1000
    balance = book.create_sheet(f"{month}月资产负债表")
    balance["C9"] = 5
    book.save(path)
    book.close()
    return path


def test_a_template_fills_in_this_months_sheet_names(tmp_path: Path) -> None:
    store = TaskStore(tmp_path / "var")
    august_task = store.create(title="八月月结", period="2026-08")
    august_store = AnalysisStore(august_task)
    august = august_store.add_dataset(_monthly(tmp_path / "7月套表.xlsx", 7), name="七月套表", metadata={"mode": "workbook"})
    balances = august_store.add_dataset(_monthly(tmp_path / "8月余额.xlsx", 8), name="八月余额", metadata={"mode": "workbook"})
    files = FinanceFiles(august_task, actor="manager")
    plan = _preview(files, august, [
        _map("7月利润表", "B5", _cell(balances, "C9", sheet="8月资产负债表")),
        _map("7月资产负债表", "C9", {"kind": "value", "value": "6"}),
    ], period="2026-08")
    assert plan["counts"]["write"] == 2
    files.rollforward_template_save({"plan_id": plan["plan_id"], "name": "月度套表"})

    september_task = store.create(title="九月月结", period="2026-09")
    september_store = AnalysisStore(september_task)
    rolled = september_store.add_dataset(_monthly(tmp_path / "8月套表.xlsx", 8), name="八月套表", metadata={"mode": "workbook"})
    later = FinanceFiles(september_task, actor="manager")
    template = later.rollforward_templates({"dataset_id": rolled["id"]})["templates"][0]
    assert [mapping["target"] for mapping in template["mappings"]] == [
        {"sheet": "8月利润表", "cell": "B5"}, {"sheet": "8月资产负债表", "cell": "C9"},
    ]
    # Listed without a workbook, the template shows the names it was saved with.
    assert later.rollforward_templates({})["templates"][0]["mappings"][0]["target"]["sheet"] == "7月利润表"
    # A source sheet still named for last month says which of this month's sheets it means.
    september = september_store.add_dataset(_monthly(tmp_path / "9月余额.xlsx", 9), name="九月余额", metadata={"mode": "workbook"})
    source = {**template["mappings"][0]["source"], "dataset_id": september["id"], "source_hash": september["source_hash"]}
    source.pop("hint")
    rerun = _preview(later, rolled, [
        {"target": template["mappings"][0]["target"], "source": source},
        {"target": template["mappings"][1]["target"], "source": {"kind": "value", "value": "7"}},
    ])
    rows = rerun["rows"]
    assert rows[0]["reason_code"] == "source_unreadable"
    assert "「9月资产负债表」" in rows[0]["reason"] and "9月资产负债表!C9" in rows[0]["reason"]
    assert rows[1]["status"] == "write"


def test_the_manager_exports_the_copy_to_open_it_in_excel(pack, capsys) -> None:
    task, analysis, august, balances, tmp_path = pack
    root = task.store.var_root
    files = FinanceFiles(task, actor="manager")
    plan = _preview(files, august, [_map("利润表", "B5", _cell(balances, "B1"))])
    run = files.rollforward_apply({"plan_id": plan["plan_id"]})
    archived = _archived(analysis, {"id": run["output_dataset_id"]})
    outside = tmp_path / "导出"
    outside.mkdir()
    destination = outside / "八月套表（2026-09 滚动）.xlsx"

    def call(*args: str) -> tuple[int, str, str]:
        code = main(["--var", str(root), "--task", task.id, *args])
        captured = capsys.readouterr()
        return code, captured.out, captured.err

    spec = json.dumps({"run_id": run["run_id"], "destination": str(destination)})
    code, _, err = call("workbook-rollforward-export", "--spec", spec)
    assert code == 3 and "only the manager" in err and not destination.exists()
    code, out, _ = call("--actor", "manager", "workbook-rollforward-export", "--spec", spec)
    assert code == 0
    answer = json.loads(out)
    assert answer["filename"] == destination.name and answer["sha256"] == run["output_sha256"]
    assert str(outside) not in out
    assert destination.read_bytes() == archived.read_bytes()
    assert [path.name for path in outside.iterdir()] == [destination.name]
    # Saving it in Excel changes the export, never the filed copy.
    destination.write_bytes(b"changed in Excel")
    assert sha256_bytes(archived.read_bytes()) == run["output_sha256"]

    for bad, message in (
        ("relative.xlsx", "absolute .xlsx path"),
        (str(outside / "copy.csv"), "absolute .xlsx path"),
        (str(Path(root) / "copy.xlsx"), "inside the store"),
        (str(outside / "missing" / "copy.xlsx"), "folder does not exist"),
    ):
        with pytest.raises(FoundryError, match=message):
            files.rollforward_export({"run_id": run["run_id"], "destination": bad})
    task.record_session("manager-session", role="analysis")
    with pytest.raises(FoundryError, match="requires the task editor"):
        FinanceFiles(task, session_id="manager-session", actor="manager").rollforward_export(
            {"run_id": run["run_id"], "destination": str(outside / "again.xlsx")}
        )
    (outside / "已存在.xlsx").mkdir()
    with pytest.raises(FoundryError, match="destination is a folder"):
        files.rollforward_export({"run_id": run["run_id"], "destination": str(outside / "已存在.xlsx")})
    archived.write_bytes(archived.read_bytes() + b"\0")
    with pytest.raises(FoundryError, match="integrity check failed"):
        files.rollforward_export({"run_id": run["run_id"], "destination": str(outside / "again.xlsx")})
    assert not (outside / "again.xlsx").exists()
    assert sorted(path.name for path in outside.iterdir()) == sorted([destination.name, "已存在.xlsx"])


def test_an_export_that_cannot_be_saved_says_so_without_the_path(pack, monkeypatch) -> None:
    import os

    task, _, august, balances, tmp_path = pack
    files = FinanceFiles(task, actor="manager")
    run = files.rollforward_apply({"plan_id": _preview(files, august, [_map("利润表", "B5", _cell(balances, "B1"))])["plan_id"]})
    outside = tmp_path / "导出"
    outside.mkdir()

    def locked(*_args, **_kwargs):
        raise PermissionError(13, "file is open", str(outside / "副本.xlsx"))

    monkeypatch.setattr(os, "replace", locked)
    with pytest.raises(FoundryError) as refused:
        files.rollforward_export({"run_id": run["run_id"], "destination": str(outside / "副本.xlsx")})
    assert str(refused.value) == "rollforward export could not save the file"
    assert list(outside.iterdir()) == []
