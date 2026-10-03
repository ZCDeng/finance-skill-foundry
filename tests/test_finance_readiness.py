from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook

from foundry.analysis import AnalysisStore
from foundry.cli import main
from foundry.finance_service import FinanceStore
from foundry.task import TaskStore

FIXTURES = Path(__file__).parent / "fixtures" / "finance_baseline" / "scenarios"


def _csv_rows(name: str) -> list[dict]:
    with (FIXTURES / name).open(encoding="utf-8-sig", newline="") as source:
        return list(csv.DictReader(source))


def _write_csv(path: Path, rows: list[dict], fieldnames: list[str] | None = None) -> Path:
    if not rows:
        raise AssertionError("fixture rows required")
    names = fieldnames or list(rows[0].keys())
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=names, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    return path


def _task(tmp_path: Path):
    task = TaskStore(tmp_path / "var").create(title="就绪")
    return task, AnalysisStore(task), FinanceStore(task)


def test_declared_hints_preserve_annual_budget_and_forecast_context() -> None:
    from foundry.finance_readiness import _declared_hints

    dataset = {
        "transformation": {
            "structure": {
                "declaration": {
                    "scope_hints": {
                        "period": "",
                        "currency": "",
                        "unit": "万元",
                        "granularity": "annual",
                        "scenario": "budget",
                        "forecast": "LE",
                    }
                }
            }
        }
    }
    assert _declared_hints(dataset) == {
        "period": "",
        "currency": "",
        "unit": "万元",
        "granularity": "annual",
        "as_of_date": "",
        "scenario": "budget",
        "forecast": "LE",
    }


def test_declared_hints_preserve_as_of_date_without_guessing_monthly_period() -> None:
    from foundry.finance_readiness import _declared_hints

    dataset = {
        "transformation": {
            "structure": {
                "declaration": {
                    "scope_hints": {
                        "period": "",
                        "unit": "万元",
                        "granularity": "as_of",
                        "as_of_date": "2026-08-31",
                    }
                }
            }
        }
    }
    hints = _declared_hints(dataset)
    assert hints["period"] == ""
    assert hints["granularity"] == "as_of"
    assert hints["as_of_date"] == "2026-08-31"


def test_cf01_g1_files_can_run_management_pnl_without_writing_context(tmp_path):
    task, analysis, store = _task(tmp_path)
    actual = _write_csv(tmp_path / "actual.csv", _csv_rows("service_ledger.csv"))
    budget = _write_csv(tmp_path / "budget.csv", _csv_rows("service_budget.csv"))
    analysis.add_dataset(actual, name="g1-actual")
    analysis.add_dataset(budget, name="g1-budget")
    before = analysis.state()
    report = store.finance_readiness({})
    after = analysis.state()
    assert set(report["can_run"]) >= {"management_pnl", "control_totals"}
    assert report["proposed_scope"]["period"]
    assert report["proposed_scope"]["currency"]
    assert report["erp_connected"] is False
    assert len(after["datasets"]) == len(before["datasets"])
    assert store._list("contexts") == []


def test_cf01e_no_authorized_dataset_blocks(tmp_path):
    _task_obj, _analysis, store = _task(tmp_path)
    report = store.finance_readiness({})
    assert report["can_run"] == []
    assert report["erp_connected"] is False
    blocked = report["blocked"]
    assert any(
        row.get("capability") == "management_pnl"
        and "dataset" in row.get("missing_fields", [])
        and row.get("reason") == "no_dataset"
        for row in blocked
    )


def test_cf02_missing_period_is_blocked_not_filled(tmp_path):
    task, analysis, store = _task(tmp_path)
    rows = []
    for row in _csv_rows("service_ledger.csv"):
        dropped = dict(row)
        dropped.pop("period", None)
        rows.append(dropped)
    path = _write_csv(tmp_path / "no-period.csv", rows, fieldnames=[k for k in rows[0] if k != "period"])
    analysis.add_dataset(path, name="no-period")
    report = store.finance_readiness({})
    assert "management_pnl" not in report["can_run"]
    assert any(
        "period" in row.get("missing_fields", []) for row in report["blocked"]
    )
    text = str(report)
    assert "本月" not in text
    assert report["proposed_scope"].get("period") in (None, "", "未声明")


def test_cf02m_unmapped_accounts_block_management_pnl(tmp_path):
    task, analysis, store = _task(tmp_path)
    rows = []
    for row in _csv_rows("service_ledger.csv"):
        dropped = dict(row)
        dropped.pop("management_account", None)
        rows.append(dropped)
    names = [k for k in rows[0] if k != "management_account"]
    analysis.add_dataset(_write_csv(tmp_path / "unmap.csv", rows, fieldnames=names), name="unmap")
    report = store.finance_readiness({})
    assert "management_pnl" not in report["can_run"]
    assert any(row.get("reason") == "mapping_incomplete" for row in report["blocked"])


def test_cf03_workbook_does_not_union_sheets(tmp_path):
    book = Workbook()
    mapped = book.active
    mapped.title = "损益"
    mapped.append(["period", "currency", "unit", "management_account", "amount"])
    mapped.append(["2026-08", "CNY", "万元", "revenue", 120])
    mapped.append(["2026-08", "CNY", "万元", "cost", 70])
    mapped.append(["2026-08", "CNY", "万元", "expense", 20])
    other = book.create_sheet("说明")
    other.append(["备注"])
    other.append(["不可映射说明页"])
    path = tmp_path / "multi.xlsx"
    book.save(path)
    task, analysis, store = _task(tmp_path)
    analysis.add_dataset(path, name="multi-sheet")
    report = store.finance_readiness({})
    names = {row.get("sheet") or row.get("name") for row in report.get("sheet_candidates") or []}
    assert "损益" in names
    assert "说明" in names
    assert report.get("merged_sheets") is not True


def test_cf03h_kingdee_header_identity_does_not_can_run(tmp_path):
    path = tmp_path / "kingdee.csv"
    path.write_text(
        "科目名称,借方,贷方\n主营业务收入,0,120\n管理费用,20,0\n",
        encoding="utf-8",
    )
    task, analysis, store = _task(tmp_path)
    analysis.add_dataset(path, name="kingdee-header")
    report = store.finance_readiness({})
    assert "management_pnl" not in report["can_run"]
    assert any(row.get("reason") == "identity_failed" for row in report["blocked"])


def test_cf04_file_product_is_not_erp_connected(tmp_path, capsys):
    task, analysis, store = _task(tmp_path)
    analysis.add_dataset(
        _write_csv(tmp_path / "actual.csv", _csv_rows("service_ledger.csv")),
        name="g1",
    )
    report = store.finance_readiness({})
    assert report["erp_connected"] is False
    blob = str(report)
    assert "金蝶" not in blob
    assert "用友" not in blob
    root = tmp_path / "var"
    assert main(["--var", str(root), "--task", task.id, "finance-readiness"]) == 0
    out = capsys.readouterr().out
    assert '"erp_connected":false' in out.replace(" ", "")
    setup = Path(__file__).resolve().parents[1] / "method" / "setup" / "SKILL.md"
    setup_text = setup.read_text(encoding="utf-8")
    assert "finance_readiness" in setup_text
    assert "不得拒绝 run_metric" in setup_text or "do not treat this pack as a gate" in setup_text.lower()
    assert "已连接金蝶" not in setup_text
    assert "已连接用友" not in setup_text


def test_cf05_identity_policy_follows_document_line_keys(tmp_path):
    task, analysis, store = _task(tmp_path)
    analysis.add_dataset(
        _write_csv(tmp_path / "actual.csv", _csv_rows("service_ledger.csv")),
        name="g1",
    )
    report = store.finance_readiness({})
    assert report["file_identity_policy"] == "business_key"

    rows = []
    for row in _csv_rows("service_ledger.csv"):
        dropped = dict(row)
        dropped.pop("doc_id", None)
        dropped.pop("line_id", None)
        rows.append(dropped)
    names = [key for key in rows[0] if key not in {"doc_id", "line_id"}]
    other = TaskStore(tmp_path / "var-balance").create(title="余额")
    analysis = AnalysisStore(other)
    store = FinanceStore(other)
    analysis.add_dataset(_write_csv(tmp_path / "balance.csv", rows, fieldnames=names), name="balance")
    balance_report = store.finance_readiness({})
    assert balance_report["file_identity_policy"] == "snapshot_only"


def test_cf03c_renminbi_title_can_run(tmp_path):
    path = tmp_path / "kingdee.csv"
    path.write_text(
        "科目名称,管理科目,金额,期间\n金蝶 币种：人民币 2026年08月,revenue,120,2026-08\n金蝶,cost,70,2026-08\n金蝶,expense,20,2026-08\n",
        encoding="utf-8",
    )
    task, analysis, store = _task(tmp_path)
    analysis.add_dataset(path, name="kingdee-rmb")
    report = store.finance_readiness({})
    assert report["proposed_scope"]["currency"] == "人民币"
    assert "management_pnl" in report["can_run"]
    assert not any(row.get("reason") == "identity_failed" for row in report["blocked"])


def test_cf03_split_sheets_do_not_can_run(tmp_path):
    book = Workbook()
    income = book.active
    income.title = "收入"
    income.append(["period", "currency", "unit", "management_account", "amount"])
    income.append(["2026-08", "CNY", "万元", "revenue", 120])
    income.append(["2026-08", "CNY", "万元", "cost", 70])
    expense = book.create_sheet("费用")
    expense.append(["period", "currency", "unit", "management_account", "amount"])
    expense.append(["2026-08", "CNY", "万元", "expense", 20])
    path = tmp_path / "split.xlsx"
    book.save(path)
    _task_obj, analysis, store = _task(tmp_path)
    analysis.add_dataset(path, name="split-sheet")
    report = store.finance_readiness({})
    assert "management_pnl" not in report["can_run"]


def test_blank_first_xlsx_row_still_reads_headers(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.append([None, None, None, None, None])
    sheet.append(["period", "currency", "unit", "management_account", "amount"])
    sheet.append(["2026-08", "CNY", "万元", "revenue", 120])
    sheet.append(["2026-08", "CNY", "万元", "cost", 70])
    sheet.append(["2026-08", "CNY", "万元", "expense", 20])
    path = tmp_path / "blank-header.xlsx"
    book.save(path)
    _task_obj, analysis, store = _task(tmp_path)
    analysis.add_dataset(path, name="blank-first")
    report = store.finance_readiness({})
    assert "management_pnl" in report["can_run"]


def test_cf02c_missing_currency_is_blocked_not_filled_with_cny(tmp_path):
    import json

    task, analysis, store = _task(tmp_path)
    rows = []
    for row in _csv_rows("service_ledger.csv"):
        dropped = dict(row)
        dropped.pop("currency", None)
        rows.append(dropped)
    names = [key for key in rows[0] if key != "currency"]
    analysis.add_dataset(_write_csv(tmp_path / "no-currency.csv", rows, fieldnames=names), name="no-currency")
    report = store.finance_readiness({})
    assert "management_pnl" not in report["can_run"]
    assert any(
        row["capability"] == "management_pnl" and row["missing_fields"] == ["currency"] and row["reason"] == "missing_scope"
        for row in report["blocked"]
    )
    assert report["proposed_scope"]["currency"] == ""
    assert "CNY" not in json.dumps(report, ensure_ascii=False)


def test_empty_workbook_shells_cannot_run_management_pnl(tmp_path):
    book = Workbook()
    first = book.active
    first.title = "封面"
    second = book.create_sheet("空表")
    second.append(["period", "currency", "unit", "management_account", "amount"])
    path = tmp_path / "empty.xlsx"
    book.save(path)
    _task_obj, analysis, store = _task(tmp_path)
    analysis.add_dataset(path, name="empty-shell")
    report = store.finance_readiness({})
    assert report["can_run"] == []
    assert any(row["capability"] == "management_pnl" for row in report["blocked"])
    assert all(candidate["mappable"] is False for candidate in report["sheet_candidates"])


def test_business_values_are_reported_verbatim(tmp_path):
    import json

    rows = []
    for row in _csv_rows("service_ledger.csv"):
        changed = dict(row)
        changed["period"] = "2026.09"
        changed["organization"] = "用友网络科技股份有限公司"
        changed["ledger"] = "KIS账套001"
        rows.append(changed)
    _task_obj, analysis, store = _task(tmp_path)
    analysis.add_dataset(_write_csv(tmp_path / "dotted.csv", rows), name="dotted")
    report = store.finance_readiness({})
    assert report["proposed_scope"]["period"] == "2026.09"
    assert report["proposed_scope"]["organization"] == "用友网络科技股份有限公司"
    assert report["proposed_scope"]["ledger"] == "KIS账套001"
    assert "[redacted]" not in json.dumps(report, ensure_ascii=False)
    assert "management_pnl" in report["can_run"]


def test_parse_failure_diagnostic_carries_no_path_or_exception_text(tmp_path):
    import json

    book = Workbook()
    sheet = book.active
    sheet.title = "损益"
    sheet.append(["period", "currency", "unit", "management_account", "amount"])
    sheet.append(["2026-08", "CNY", "万元", "revenue", 120])
    path = tmp_path / "broken.xlsx"
    book.save(path)
    _task_obj, analysis, store = _task(tmp_path)
    dataset = analysis.add_dataset(path, name="broken")
    stored = analysis.root / "datasets" / dataset["id"] / "source.xlsx"
    stored.write_bytes(b"not a workbook at all")
    report = store.finance_readiness({})
    blob = json.dumps(report, ensure_ascii=False)
    assert "management_pnl" not in report["can_run"]
    assert any(diag["code"] in {"parse_failed", "dataset_unreadable"} for diag in report["diagnostics"])
    for marker in (str(tmp_path), "/Users", "Traceback", "file://", "http", "Error"):
        assert marker not in blob


def test_kingdee_header_with_covered_mapping_still_fails_identity(tmp_path):
    path = tmp_path / "kingdee-mapped.csv"
    path.write_text(
        "科目名称,借方,贷方,管理科目\n主营业务收入,0,120,revenue\n主营业务成本,70,0,cost\n管理费用,20,0,expense\n",
        encoding="utf-8",
    )
    _task_obj, analysis, store = _task(tmp_path)
    analysis.add_dataset(path, name="kingdee-mapped")
    report = store.finance_readiness({})
    assert "management_pnl" not in report["can_run"]
    blocked = next(row for row in report["blocked"] if row["capability"] == "management_pnl")
    assert blocked["reason"] == "identity_failed"
    assert set(blocked["missing_fields"]) == {"period", "currency"}
    assert "management_account" not in blocked["missing_fields"]


def _scoped_rows(period: str, unit: str) -> list[dict]:
    return [
        {"account_code": "6001", "amount": "400", "period": period,
         "currency": "CNY", "unit": unit, "management_account": "revenue"},
        {"account_code": "6401", "amount": "300", "period": period,
         "currency": "CNY", "unit": unit, "management_account": "cost"},
        {"account_code": "6602", "amount": "50", "period": period,
         "currency": "CNY", "unit": unit, "management_account": "expense"},
    ]


def test_two_files_closing_different_months_do_not_share_a_period(tmp_path):
    """Joining them produced 2026-08,2026-01, which is not a period at all."""
    _task_obj, analysis, store = _task(tmp_path)
    analysis.add_dataset(
        _write_csv(tmp_path / "aug.csv", _scoped_rows("2026-08", "元")), name="八月"
    )
    analysis.add_dataset(
        _write_csv(tmp_path / "jan.csv", _scoped_rows("2026-01", "万元")), name="一月"
    )
    report = store.finance_readiness({})
    assert report["proposed_scope"]["period"] == ""
    assert report["proposed_scope"]["unit"] == ""
    # Each file still answers for itself.
    periods = {row["name"]: row["scope"]["period"] for row in report["datasets"]}
    assert periods == {"八月": "2026-08", "一月": "2026-01"}
    units = {row["name"]: row["scope"]["unit"] for row in report["datasets"]}
    assert units == {"八月": "元", "一月": "万元"}


def test_one_file_missing_its_currency_does_not_block_the_file_that_has_one(tmp_path):
    """A weak file used to kill the whole task's management_pnl."""
    _task_obj, analysis, store = _task(tmp_path)
    analysis.add_dataset(
        _write_csv(tmp_path / "good.csv", _scoped_rows("2026-08", "元")), name="有币种"
    )
    bare = [{k: v for k, v in row.items() if k != "currency"}
            for row in _scoped_rows("2026-08", "元")]
    analysis.add_dataset(_write_csv(tmp_path / "bare.csv", bare), name="缺币种")

    report = store.finance_readiness({})
    assert "management_pnl" in report["can_run"]
    weak = next(row for row in report["datasets"] if row["name"] == "缺币种")
    assert weak["missing_fields"] == ["currency"]
    assert "management_pnl" not in weak["can_run"]
    strong = next(row for row in report["datasets"] if row["name"] == "有币种")
    assert "management_pnl" in strong["can_run"]
    # Nothing is reported as blocked once some file can answer for it.
    assert not [row for row in report["blocked"]
                if row["capability"] == "management_pnl"]


def test_a_column_with_many_values_is_a_dimension_not_the_files_scope(tmp_path):
    """组织 came back as 上海城市公司 because the first row happened to say so."""
    _task_obj, analysis, store = _task(tmp_path)
    rows = _scoped_rows("2026-08", "元")
    rows[0]["organization"] = "上海城市公司"
    rows[1]["organization"] = "宁波城市公司"
    analysis.add_dataset(_write_csv(tmp_path / "multi.csv", rows), name="多组织")
    report = store.finance_readiness({})
    assert report["proposed_scope"]["organization"] == ""
    assert report["datasets"][0]["scope"]["organization"] == ""


def test_a_file_that_names_one_organization_still_reports_it(tmp_path):
    _task_obj, analysis, store = _task(tmp_path)
    rows = [{**row, "organization": "集团合并"} for row in _scoped_rows("2026-08", "元")]
    analysis.add_dataset(_write_csv(tmp_path / "one.csv", rows), name="单组织")
    report = store.finance_readiness({})
    assert report["proposed_scope"]["organization"] == "集团合并"


def _report_pack(path: Path, subtitles: dict[str, str]) -> Path:
    """A month-end pack laid out the way real ones are: a title, a subtitle naming the period, then a table."""
    book = Workbook()
    book.remove(book.active)
    for sheet, subtitle in subtitles.items():
        page = book.create_sheet(sheet)
        page.append([sheet])
        page.append([f"华宸物业  |  {subtitle}  |  单位：万元"])
        page.append([])
        page.append(["项目", "金额"])
        page.append(["营业收入", 100])
    book.save(path)
    return path


def test_a_period_span_in_a_title_declares_no_month() -> None:
    """「2026年1-8月」 used to read as 2026-01; it is a span, and a span is not a period."""
    from foundry.finance_readiness import _period_match

    for title in (
        "2026年1-8月",
        "预算执行分析（1-8 月） 华宸物业 | 2026年1-8月",
        "2026年1月至8月",
        "2026年1－8月",  # full-width dash, as Excel files carry it
        "2026年1〜8月",
        "2026年1−8月",
        "2026-1-8月",
        "2026.1-8月",
        "2025年12月-2026年3月",
        "2026-08-01至2026-08-31",
        "2026年1月1日至2026年8月31日",
        "2026年8-12月预测",
    ):
        assert _period_match(title) == ("", ""), title


def test_period_title_still_reads_single_months_and_dates() -> None:
    from foundry.finance_readiness import _period_match

    assert _period_match("2026年08月 月结") == ("2026-08", "2026年08月")
    assert _period_match("2026年01月 月结")[0] == "2026-01"
    assert _period_match("2026年8月31日")[0] == "2026-08"
    assert _period_match("台账基准日 2026-08-31")[0] == "2026-08"
    assert _period_match("2026.1")[0] == "2026-01"
    # The first declaration still wins: a comparison span later in the title is not the period.
    assert _period_match("2026年08月 对比 2025年1-12月")[0] == "2026-08"
    # Days inside one month are not a span of months.
    assert _period_match("2026年8月1日-31日")[0] == "2026-08"
    assert _period_match("2026年08月第1-2周")[0] == "2026-08"


def test_a_dash_cell_between_two_dates_does_not_make_a_span(tmp_path):
    """A 「-」 placeholder cell sat between two date cells and read as 2026-08-15 至 2026-09-14."""
    from foundry.finance_readiness import _title_period

    assert _title_period(["开票日期", "金额", "到期日", "2026-08-15", "-", "2026-09-14"]) == ("2026-08", "2026-08")
    assert _title_period(["2026年1-8月", "2026年08月"]) == ("", "")


def test_period_title_reads_no_period_where_none_is_declared() -> None:
    from foundry.finance_readiness import _period_match

    assert _period_match("2026 年度") == ("", "")
    assert _period_match("BUD-2026 → LE（8月底滚动预测）") == ("", "")
    assert _period_match("2026年13月") == ("", "")
    assert _period_match("共 2026年08月 - 2 张")[0] == "2026-08"


def test_month_and_year_to_date_sheets_of_one_pack_agree_on_the_month(tmp_path):
    """The 2026-08 pack proposed no period: its 1-8 月 sheets were read as 2026-01.

    They now declare nothing, and the two sheets that name the month decide.
    """
    path = _report_pack(tmp_path / "华宸物业_2026年08月_月结.xlsx", {
        "00·封面": "2026年08月",
        "05·资产负债表": "2026年8月31日",
        "08·收缴率": "2026年1-8月",
    })
    _task_obj, analysis, store = _task(tmp_path)
    analysis.add_dataset(path, name=path.stem)
    report = store.finance_readiness({})
    assert report["proposed_scope"]["period"] == "2026-08"
    (row,) = report["datasets"]
    assert row["scope"]["period"] == "2026-08"
    assert row["scope_sources"]["period"] == {
        "origin": "title", "sheet": "00·封面", "text": "2026年08月", "sheet_count": 2,
    }
    assert row["name_scope"] == {"period": "2026-08", "text": "2026年08月"}


def test_a_pack_whose_sheets_name_different_months_still_proposes_none(tmp_path):
    path = _report_pack(tmp_path / "mixed.xlsx", {
        "本月": "2026年08月",
        "上月": "2026年07月",
    })
    _task_obj, analysis, store = _task(tmp_path)
    analysis.add_dataset(path, name="mixed")
    report = store.finance_readiness({})
    assert report["proposed_scope"]["period"] == ""
    (row,) = report["datasets"]
    assert row["scope"]["period"] == ""
    assert row["scope_sources"] == {}
    assert row["name_scope"] == {}


def test_a_file_name_period_is_offered_but_never_fills_the_scope(tmp_path):
    path = _report_pack(tmp_path / "华宸_2026年08月.xlsx", {"说明": "年度汇总"})
    _task_obj, analysis, store = _task(tmp_path)
    analysis.add_dataset(path, name=path.stem)
    report = store.finance_readiness({})
    (row,) = report["datasets"]
    assert row["scope"]["period"] == ""
    assert "period" in row["missing_fields"]
    assert row["name_scope"]["period"] == "2026-08"
