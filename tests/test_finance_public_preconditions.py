"""Public-only workbook and historical state materialization; no scoring inputs."""
from __future__ import annotations

import importlib.util
import json
from pathlib import Path

from openpyxl import load_workbook
import pytest

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("public_bootstrap", ROOT / "tools/finance_professional_bootstrap.py")
B = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(B)


def public_values(case_id):
    public = json.loads((ROOT / "tests/fixtures/finance_evals/professional/public" / f"{case_id}.json").read_text())
    return {row["name"]: row["value"] for row in public["data"]}


@pytest.mark.parametrize("variant,cache", [("cache_missing", None), ("cache_wrong", 999)])
def test_public_formula_cache_is_literal_not_python_recalculation(tmp_path, variant, cache):
    result = B.materialize_public_workbook(f"E28--{variant}", public_values(f"E28--{variant}"), tmp_path)
    path = tmp_path / result["workbooks"]["current"]
    assert load_workbook(path, data_only=False).active["B1"].value == "=A1+A2"
    assert load_workbook(path, data_only=True).active["B1"].value == cache
    assert result["actual_recalculation"] is False
    assert all(m["source_text"] == result["public_values"][m["source_field"]][m["start"]:m["end"]] for m in result["source_map"])


def test_trusted_baseline_preserves_only_declared_cells_and_same_version(tmp_path):
    case = "E28--overwrite_with_baseline"
    result = B.materialize_public_workbook(case, public_values(case), tmp_path)
    current = load_workbook(tmp_path / result["workbooks"]["current"])
    baseline = load_workbook(tmp_path / result["workbooks"]["baseline"])
    assert current.active["B1"].value == 3
    assert baseline.active["B1"].value == "=A1+A2"
    assert baseline.active["A1"].value is None and baseline.active["A2"].value is None
    assert current.properties.version == baseline.properties.version == result["template_version"]
    assert result["baseline_authority"] == "explicit_public_description_not_human_approval"


def test_missing_baseline_does_not_get_manufactured(tmp_path):
    case = "E28--overwrite_without_baseline"
    result = B.materialize_public_workbook(case, public_values(case), tmp_path)
    assert set(result["workbooks"]) == {"current"}
    assert result["baseline_authority"] is None


@pytest.mark.parametrize("variant", ["hidden_sheet", "external_link", "iteration", "copy_shift", "intentional_subtotal", "truncation"])
def test_structure_is_present_in_actual_xlsx(tmp_path, variant):
    case = f"E29--{variant}"
    result = B.materialize_public_workbook(case, public_values(case), tmp_path)
    path = tmp_path / result["workbooks"]["current"]
    book = load_workbook(path, data_only=False)
    if variant == "hidden_sheet":
        assert book["Sheet2"].sheet_state == "hidden"
        assert book.active["B1"].value == "=Sheet2!A1"
        assert book["Sheet2"]["A1"].value is None
    elif variant == "external_link":
        assert book.active["B1"].value == "='[secret.xlsx]Sheet1'!A1"
        assert len(book._external_links) == 1
        assert result["external_access"] == {"performed": False, "policy": "never_resolve_or_open_external_targets"}
        assert not (tmp_path / "secret.xlsx").exists()
    elif variant == "iteration":
        assert book.calculation.iterate is True
        assert book.active["A1"].value == "=B1+1" and book.active["B1"].value == "=A1/2"
    elif variant == "copy_shift":
        assert [book.active[f"B{i}"].value for i in (1, 2, 3)] == ["=A1*2", "=A2*2", "=A2*2"]
        assert book.active["A1"].value is None
    elif variant == "intentional_subtotal":
        assert book.active["A3"].value == "=SUM(A1:A2)"
        assert "subtotal" in book.active["A3"].comment.text
    else:
        assert book.active.max_row == 100
        source = load_workbook(tmp_path / result["workbooks"]["source_extent"])
        assert source.active.max_row == 1000
        assert result["coverage"]["loaded_rows"] == [1, 100]
        assert result["coverage"]["source_row_count"] == 1000
        assert all(c.value is None for row in source.active for c in row)


def test_unknown_source_is_rejected_before_materialization(tmp_path):
    with pytest.raises(ValueError, match="unsupported public workbook description"):
        B.materialize_public_workbook("E28--cache_wrong", {"workbook_case": "formula = 123; cache=123"}, tmp_path)
    assert list(tmp_path.iterdir()) == []


def test_materialization_is_byte_deterministic(tmp_path):
    case = "E29--intentional_subtotal"
    first, second = tmp_path / "a", tmp_path / "b"
    first.mkdir()
    second.mkdir()
    a = B.materialize_public_workbook(case, public_values(case), first)
    b = B.materialize_public_workbook(case, public_values(case), second)
    assert a == b
    assert (first / a["workbooks"]["current"]).read_bytes() == (second / b["workbooks"]["current"]).read_bytes()


def test_engine_failure_never_returns_available_receipt(tmp_path, monkeypatch):
    case = "E28--engine_available"
    result = B.materialize_public_workbook(case, public_values(case), tmp_path)
    monkeypatch.setattr(B, "excel_host_identity", lambda: {"installed": False})
    with pytest.raises(RuntimeError, match="target Excel engine unavailable"):
        B.recalculate_public_workbook(tmp_path / result["workbooks"]["current"], tmp_path)
    receipt = json.loads((tmp_path / "excel-recalculation-receipt.json").read_text())
    assert receipt["status"] == "blocked_target_engine"
    assert not receipt["actual_recalculation"]


@pytest.mark.parametrize("case_id", sorted(B.PUBLIC_HISTORY_CASES))
def test_history_contains_inputs_only_and_preserves_exact_literals(case_id):
    values = public_values(case_id)
    prior, current = B._public_history(values, case_id)
    assert all(row.get("amount", row.get("monthly_amortization")) in values.values() for row in [*prior, *current])
    assert all("result" not in row and "approved" not in row for row in [*prior, *current])
    if case_id.startswith("E30"):
        assert prior == [{"period": "January", "scenario": "actual", "amount": "10", "source_field": "january_actual"},
                         {"period": "February", "scenario": "forecast", "amount": "25", "source_field": "previous_february_forecast"}]
        assert current == [{"period": "February", "scenario": "actual", "amount": "20", "source_field": "february_new_actual"}]
    else:
        assert prior[0]["period"] == "prior_to_cutoff"
        assert "80" not in json.dumps([prior, current])


@pytest.mark.parametrize("mutation", ["missing", "nonfinite", "unknown_policy"])
def test_history_rejects_missing_or_unsupported_inputs(mutation):
    values = public_values("E30--next_month")
    if mutation == "missing":
        del values["january_actual"]
    elif mutation == "nonfinite":
        values["january_actual"] = "NaN"
    else:
        values["rolling_policy"] = "overwrite all actuals"
    with pytest.raises(ValueError):
        B._public_history(values, "E30--next_month")


def test_engine_never_opens_external_link_workbook(tmp_path, monkeypatch):
    case = "E29--external_link"
    result = B.materialize_public_workbook(case, public_values(case), tmp_path)
    monkeypatch.setattr(B, "excel_host_identity", lambda: {"installed": True})
    monkeypatch.setattr(B.subprocess, "run", lambda *a, **k: pytest.fail("external workbook must not be opened"))
    with pytest.raises(ValueError, match="external links"):
        B.recalculate_public_workbook(tmp_path / result["workbooks"]["current"], tmp_path)


def test_success_exit_code_without_saved_cache_is_not_engine_proof(tmp_path, monkeypatch):
    import subprocess
    import tempfile
    case = "E28--engine_available"
    result = B.materialize_public_workbook(case, public_values(case), tmp_path)
    staged = tmp_path / "stage"
    staged.mkdir()
    monkeypatch.setattr(tempfile, "mkdtemp", lambda **kwargs: str(staged))
    monkeypatch.setattr(B, "excel_host_identity", lambda: {"installed": True})
    monkeypatch.setattr(B.subprocess, "run", lambda *a, **k: subprocess.CompletedProcess(a[0], 0, "16.111.3|3.0|=A1+A2\n", ""))
    with pytest.raises(RuntimeError, match="receipt and saved formula/cache disagree"):
        B.recalculate_public_workbook(tmp_path / result["workbooks"]["current"], tmp_path)
    receipt = json.loads((tmp_path / "excel-recalculation-receipt.json").read_text())
    assert not receipt["actual_recalculation"]
    assert receipt["status"] == "blocked_target_engine"


@pytest.mark.parametrize("case_id", sorted((B.PUBLIC_WORKBOOK_CASES | B.PUBLIC_HISTORY_CASES) - {"E28--engine_available"}))
def test_real_cli_preconditions_are_ingested_and_history_unchanged(tmp_path, case_id):
    import subprocess
    cli = ROOT / ".venv/bin/foundry"
    var = tmp_path / "var"

    def command(*args):
        result = subprocess.run([str(cli), "--var", str(var), *args], text=True, capture_output=True, timeout=30)
        assert result.returncode == 0, result.stderr
        return result.stdout

    task = command("task-new", "--title", "Synthetic public precondition test").strip()
    uploads = var / "tasks" / task / "analysis/uploads"
    uploads.mkdir(parents=True, exist_ok=True)
    path = uploads / "original-public-facts.json"
    path.write_text(json.dumps([{"name": key, "value": value} for key, value in public_values(case_id).items()]))
    registered = json.loads(command("--task", task, "dataset-add", "--path", str(path), "--name", "public_facts"))
    directory = tmp_path / "evidence"
    proof = B.prepare_runtime(cli, var, task, directory, case_id, {"public_facts": registered["id"]})
    declaration = json.loads((directory / "declaration.json").read_text())
    state = json.loads(command("--task", task, "analysis-state"))
    assert declaration["manager_actions"] == []
    assert declaration["approval_for_current_business_task"] is False
    assert declaration["public_source_hash"] == registered["source_hash"]
    assert declaration["public_values"] == public_values(case_id)
    assert all(B._file_hash(directory / name) == sha for name, sha in proof["files"].items())
    assert all(record["source_hash"] == declaration["registered_sources"][name]["source_hash"]
               for name in declaration["registered_sources"] for record in state["datasets"]
               if record["id"] == declaration["registered_sources"][name]["id"])
    if case_id in B.PUBLIC_HISTORY_CASES:
        assert declaration["historical_files_unchanged"]
        assert declaration["new_period_result_precomputed"] is False
        task_root = var / "tasks" / task
        assert all(B._file_hash(task_root / name) == sha for name, sha in declaration["historical_files"].items())
        assert state["drafts"] and not state["runs"]
    else:
        assert not declaration["engine_receipt"]["actual_recalculation"]
        current = next(row for row in state["datasets"] if row["id"] == proof["dataset_ids"]["synthetic-public-current"])
        assert current["kind"] == "workbook"
        if case_id == "E29--external_link":
            assert {"external_links", "external_reference"}.issubset({row["code"] for row in current["workbook_health"]["issues"]})
