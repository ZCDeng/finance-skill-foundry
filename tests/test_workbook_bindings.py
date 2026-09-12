from __future__ import annotations

import json
import zipfile
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.styles import PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from foundry.errors import FoundryError
from foundry.workbook_bindings import (
    apply_output_patch,
    create_binding_contract,
    inspect_workbook,
    preview_output_patch,
    read_input,
)


def _save_template(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Model"
    ws.append(["account_id", "amount", "note"])
    ws.append(["001234567890123456", Decimal("282670.7525"), "ok"])
    ws.append(["0002", Decimal("10"), "=literal"])
    ws["C3"].data_type = "s"
    table = Table(displayName="InputTable", ref="A1:C3")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws["E1"] = "Total"
    ws["E2"] = "=SUM(B2:B3)"
    ws["G1"] = "Result"
    ws["H1"] = "Comment"
    ws["G2"] = "old"
    ws["G2"].fill = PatternFill("solid", fgColor="FFEEAA")
    wb.defined_names.add(DefinedName("ResultCell", attr_text="'Model'!$G$2"))
    wb.defined_names.add(DefinedName("InputBlock", attr_text="'Model'!$A$1:$C$3"))
    wb.save(path)
    wb.close()
    return path


def test_named_table_input_preserves_identifier_and_decimal_policy(tmp_path: Path) -> None:
    source = _save_template(tmp_path / "model.xlsx")
    contract = create_binding_contract(
        source,
        {"kind": "table", "name": "InputTable"},
        [
            {"name": "account_id", "policy": "identifier", "unit": "account"},
            {"name": "amount", "policy": "amount", "unit": "CNY"},
            {"name": "note", "policy": "text"},
        ],
    )

    rows = read_input(source, contract)

    assert rows["rows"] == [
        {"account_id": "001234567890123456", "amount": "282670.7525", "note": "ok"},
        {"account_id": "0002", "amount": "10", "note": "=literal"},
    ]
    assert rows["diagnostics"] == []
    assert contract["resolved"] == {"sheet": "Model", "range": "A1:C3", "rows": 3, "columns": 3}


def test_copy_patch_writes_summable_amounts_and_preserves_formula_style_and_names(tmp_path: Path) -> None:
    source = _save_template(tmp_path / "model.xlsx")
    contract = create_binding_contract(
        source,
        {"kind": "range", "sheet": "Model", "range": "G1:H3"},
        [
            {"name": "result", "policy": "amount", "unit": "CNY"},
            {"name": "comment", "policy": "text"},
        ],
        role="output",
        writable=True,
    )
    patch = preview_output_patch(
        source,
        contract,
        [
            {"row": 1, "field": "result", "value": "1234.5"},
            {"row": 1, "field": "comment", "value": "=not a formula"},
        ],
    )

    result = apply_output_patch(source, tmp_path / "patched.xlsx", patch)
    book = load_workbook(result["path"], data_only=False)
    try:
        ws = book["Model"]
        assert ws["G2"].value == pytest.approx(1234.5)
        assert ws["H2"].value == "=not a formula"
        assert ws["H2"].data_type == "s"
        assert ws["E2"].value == "=SUM(B2:B3)"
        assert ws["G2"].fill.fgColor.rgb == "00FFEEAA"
        assert book.defined_names.get("ResultCell") is not None
        assert list(book.defined_names["ResultCell"].destinations) == [("Model", "$G$2")]
    finally:
        book.close()
    data_book = load_workbook(result["path"], data_only=True)
    try:
        assert isinstance(data_book["Model"]["G2"].value, float)
    finally:
        data_book.close()


def test_long_decimal_output_is_text_with_diagnostic(tmp_path: Path) -> None:
    source = _save_template(tmp_path / "model.xlsx")
    contract = create_binding_contract(
        source,
        {"kind": "named_range", "name": "ResultCell"},
        [{"name": "result", "policy": "amount", "unit": "CNY"}],
        role="output",
        writable=True,
        header=False,
    )

    patch = preview_output_patch(source, contract, [{"row": 1, "field": "result", "value": "123456789012345.67"}])
    result = apply_output_patch(source, tmp_path / "long.xlsx", patch)
    book = load_workbook(result["path"])
    try:
        assert book["Model"]["G2"].value == "123456789012345.67"
        assert isinstance(book["Model"]["G2"].value, str)
    finally:
        book.close()
    assert patch["diagnostics"][0]["code"] == "decimal_written_as_text"


def test_concurrent_cell_conflict_requires_and_honors_resolution(tmp_path: Path) -> None:
    source = _save_template(tmp_path / "model.xlsx")
    contract = create_binding_contract(
        source,
        {"kind": "named_range", "name": "ResultCell"},
        [{"name": "result", "policy": "amount"}],
        role="output",
        writable=True,
        header=False,
    )
    edited = load_workbook(source)
    edited["Model"]["G2"] = "manager edit"
    edited.save(source)
    edited.close()

    patch = preview_output_patch(source, contract, [{"row": 1, "field": "result", "value": "20"}])

    assert patch["conflicts"][0]["cell"] == "G2"
    with pytest.raises(FoundryError, match="conflict at G2"):
        apply_output_patch(source, tmp_path / "blocked.xlsx", patch)
    kept = apply_output_patch(source, tmp_path / "kept.xlsx", patch, {"G2": "keep"})
    applied = apply_output_patch(source, tmp_path / "applied.xlsx", patch, {"G2": "apply"})
    keep_book = load_workbook(kept["path"])
    apply_book = load_workbook(applied["path"])
    try:
        assert keep_book["Model"]["G2"].value == "manager edit"
        assert apply_book["Model"]["G2"].value == 20
    finally:
        keep_book.close()
        apply_book.close()


def test_apply_rechecks_cell_hash_after_preview(tmp_path: Path) -> None:
    source = _save_template(tmp_path / "model.xlsx")
    contract = create_binding_contract(
        source,
        {"kind": "named_range", "name": "ResultCell"},
        [{"name": "result", "policy": "amount"}],
        role="output",
        writable=True,
        header=False,
    )
    patch = preview_output_patch(source, contract, [{"row": 1, "field": "result", "value": "20"}])
    edited = load_workbook(source)
    edited["Model"]["G2"] = "late manager edit"
    edited.save(source)
    edited.close()

    with pytest.raises(FoundryError, match="changed since preview"):
        apply_output_patch(source, tmp_path / "blocked-late.xlsx", patch)
    patch = preview_output_patch(source, contract, [{"row": 1, "field": "result", "value": "20"}])
    result = apply_output_patch(source, tmp_path / "late-applied.xlsx", patch, {"G2": "apply"})
    book = load_workbook(result["path"])
    try:
        assert book["Model"]["G2"].value == 20
    finally:
        book.close()


def test_complex_workbook_uses_independent_result_workbook_fallback(tmp_path: Path) -> None:
    source = _save_template(tmp_path / "with-drawing.xlsx")
    with zipfile.ZipFile(source, "a") as archive:
        archive.writestr("xl/drawings/drawing1.xml", "<drawing />")
    contract = create_binding_contract(
        source,
        {"kind": "named_range", "name": "ResultCell"},
        [{"name": "result", "policy": "amount"}],
        role="output",
        writable=True,
        header=False,
    )

    inspection = inspect_workbook(source)
    patch = preview_output_patch(source, contract, [{"row": 1, "field": "result", "value": "30"}])
    result = apply_output_patch(source, tmp_path / "fallback.xlsx", patch)

    assert inspection["copy_rewrite_supported"] is False
    assert patch["requires_independent_result_workbook"] is True
    assert result["mode"] == "independent_result_workbook"
    original = load_workbook(source)
    fallback = load_workbook(result["path"])
    try:
        assert original["Model"]["G2"].value == "old"
        assert fallback["Results"]["C2"].value == 30
    finally:
        original.close()
        fallback.close()


def test_writes_outside_declared_area_are_refused(tmp_path: Path) -> None:
    source = _save_template(tmp_path / "model.xlsx")
    contract = create_binding_contract(
        source,
        {"kind": "range", "sheet": "Model", "range": "G2:G2"},
        [{"name": "result", "policy": "amount"}],
        role="output",
        writable=True,
        header=False,
    )

    with pytest.raises(FoundryError, match="outside declared writable area"):
        preview_output_patch(source, contract, [{"cell": "H2", "field": "result", "value": "1"}])


def test_a_header_that_wraps_in_the_cell_binds_by_what_she_reads(tmp_path: Path) -> None:
    """「本期发生额\\n借方」 is one header shown on two lines.

    Matching the raw cell text made the manager type the line break to bind a
    column she can see perfectly well.
    """
    from foundry.workbook_bindings import create_binding_contract, read_input

    source = tmp_path / "ledger.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "余额表"
    sheet["A1"] = "科目代码"
    sheet["B1"] = "本期发生额\n借方"
    sheet["A2"] = "6001"
    sheet["B2"] = 100
    book.save(source)
    book.close()

    spec = {"kind": "range", "sheet": "余额表", "range": "A1:B2"}
    for spelling in ("本期发生额借方", "本期发生额\n借方", "本期发生额 借方"):
        contract = create_binding_contract(
            source, spec, [{"name": spelling, "type": "decimal"}], header=True
        )
        assert len(read_input(source, contract)["rows"]) == 1

    absent = create_binding_contract(
        source, spec, [{"name": "期末余额借方", "type": "decimal"}], header=True
    )
    with pytest.raises(FoundryError, match="binding header missing fields"):
        read_input(source, absent)


def test_two_headers_that_read_alike_are_refused_rather_than_guessed(tmp_path: Path) -> None:
    from foundry.workbook_bindings import create_binding_contract, read_input

    source = tmp_path / "twin.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "余额表"
    sheet["A1"] = "本期发生额\n借方"
    sheet["B1"] = "本期发生额借方"
    sheet["A2"] = 1
    sheet["B2"] = 2
    book.save(source)
    book.close()

    contract = create_binding_contract(
        source,
        {"kind": "range", "sheet": "余额表", "range": "A1:B2"},
        [{"name": "本期发生额借方", "type": "decimal"}],
        header=True,
    )
    with pytest.raises(FoundryError, match="ambiguous"):
        read_input(source, contract)


def test_an_inspection_does_not_carry_the_hosts_path(tmp_path: Path) -> None:
    from foundry.workbook_bindings import inspect_workbook

    source = tmp_path / "book.xlsx"
    book = Workbook()
    book.active["A1"] = 1
    book.save(source)
    book.close()
    inspection = inspect_workbook(source)
    assert "path" not in inspection
    assert str(tmp_path) not in json.dumps(inspection, ensure_ascii=False)
