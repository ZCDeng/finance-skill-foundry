"""Exports always label estimates 估算 (issue #24; phase 1 decision 2, spec item 13).

A cell is an estimate when it is a cited cell of a grain the server's requirement check for the draft found
`estimated`; the assistant's text and any report a caller supplies decide nothing. Every export format is built
from the report projection, so the projection writes （估算） beside the numbers bound to those cells in prose,
tables and captions: written as the cell's claim value, its presentation or its {{fact:…}} template renders it,
never by a value that is merely close, and never for 0. The xlsx run sheets mark the estimated cells themselves.
Where the text writes an estimate the export cannot label (a rounded copy bound to nothing, a code block), the
draft gets a quality issue and the export says so instead of shipping it unlabelled.
"""
import hashlib
import io
import json
import re
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from foundry.analysis import AnalysisStore
from foundry.finance_requirements import save_requirements
from foundry.finance_service import FinanceStore
from foundry.report_documents import report_document_from_markdown_body
from foundry.report_projection import build_report_projection
from foundry.task import TaskStore
from test_finance_requirements import _claim, _recipe, _script_run, _workbook, spec

LABEL = '（估算）'
# 华东 1-8 月收入预算 = 华东全年预算 05!C5 188503.3 × 集团 1-8 月预算进度 (01!B12:C12 300 ÷ 01!N12 480).
ESTIMATE = '117814.5625'
VALUES = ['325644', '300', '109742.028', ESTIMATE, '0.625']
PRESENTED = {'decimal_places': 2, 'scale': 'identity', 'rounding': 'half_up'}
TABLE = '\n'.join([
    '| 区域 | 1-8 月实际 | 1-8 月预算 |',
    '| --- | ---: | ---: |',
    '| 集团合并 | 325,644 | 300 |',
    '| 华东区域 | 109,742.028 | 117,814.5625 |',
])
BODY = '## 区域收入\n\n华东区域 1-8 月收入实际 109,742.028 万元，同期预算 117,814.56 万元。\n\n' + TABLE


@pytest.fixture
def analysis(tmp_path):
    return AnalysisStore(TaskStore(tmp_path / 'var').create(title='估算标注'))


def _case(analysis, tmp_path, *, declare=True, east=(188503.3, ESTIMATE, 109742.028, VALUES[2])):
    """Huachen revenue_1_8 in miniature, the run typed like a real script run (named, typed columns).

    `east` is 华东's full-year budget cell, its 1-8 estimate, its 1-8 actual cell and that actual as run output.
    """
    full_year, estimate, actual_cell, actual = east
    values = [VALUES[0], VALUES[1], actual, estimate, VALUES[4]]
    budget = _workbook(analysis, tmp_path, 'budget', {
        '01·年度预算总表': {'B12': 100, 'C12': 200, 'N12': 480},
        '05·区域预算分解': {'C5': full_year},
    })
    close = _workbook(analysis, tmp_path, 'close', {'10·区域经营损益表': {'D5': actual_cell, 'D11': 325644}})
    declared = [{'organizations': ['华东区域'], 'scenarios': ['budget_1_8'], 'basis': 'estimate',
                 'method': '区域全年收入预算 × 集团1-8月收入预算进度'}]
    requirement = save_requirements(analysis, spec(
        question='1-8 月收入实际与同期预算',
        scope={'ledger': '预算工作簿与月结套表', 'currency': 'CNY', 'unit': '万元'},
        bindings={'budget': budget['id'], 'close': close['id']},
        items=[{'id': 'revenue_1_8', 'metric': 'revenue', 'organizations': ['集团合并', '华东区域'],
                'periods': ['2026-01:2026-08'], 'scenarios': ['actual_1_8', 'budget_1_8'], 'requires_drilldown': True,
                **({'grain_basis': declared} if declare else {})}],
    ))
    base = {'metric_key': 'revenue', 'period': '2026-01:2026-08', 'unit': '万元', 'currency': 'CNY'}
    rows = [
        {**base, 'organization': '集团合并', 'scenario': 'actual_1_8', 'value': values[0]},
        {**base, 'organization': '集团合并', 'scenario': 'budget_1_8', 'value': values[1]},
        {**base, 'organization': '华东区域', 'scenario': 'actual_1_8', 'value': values[2]},
        {**base, 'organization': '华东区域', 'scenario': 'budget_1_8', 'value': values[3]},
        {**base, 'organization': '集团合并', 'scenario': 'budget_phasing_ratio', 'value': values[4], 'unit': '-'},
    ]
    names = list(dict.fromkeys(key for row in rows for key in row))
    run = _script_run(analysis, {'budget': budget, 'close': close}, rows, suffix='estimate-labels',
                      requirements_id=requirement['id'],
                      requirements_semantic_fingerprint=requirement['semantic_fingerprint'],
                      result={'columns': [{'name': name, 'type': 'DECIMAL(38,10)' if name == 'value' else 'VARCHAR'}
                                          for name in names], 'rows': rows})
    key = {'metric_key': 'revenue', 'period': '2026-01:2026-08'}

    def cell(reference):
        return {'dataset_id': budget['id'], 'source_hash': budget['source_hash'], 'sheet': '01·年度预算总表',
                'range': reference}

    recipes = [
        _recipe(close, '10·区域经营损益表', 'D11', run, 0, recipe_id='group_actual'),
        _recipe(budget, '01·年度预算总表', 'B12:C12', run, 1, kind='sum', recipe_id='group_budget'),
        _recipe(close, '10·区域经营损益表', 'D5', run, 2, recipe_id='east_actual'),
        {'id': 'phasing', 'kind': 'ratio', 'business_key': key, 'claim': {'run_id': run['id'], 'row': 4, 'column': 'value'},
         'left': cell('B12:C12'), 'right': cell('N12')},
        {**_recipe(budget, '05·区域预算分解', 'C5', run, 3, recipe_id='east_budget_estimate'),
         'factor_recipe_id': 'phasing', 'business_key': key},
    ]
    service = FinanceStore(analysis.task)
    facts = service.prepare_explanation({'run_ids': [run['id']]})['facts']
    estimate_fact = next(fact for fact in facts if fact['row'] == 3 and fact['column'] == 'value')
    presented = {'basis': 'calculation', 'fact_id': estimate_fact['id'], 'run_id': run['id'], 'row': 3,
                 'column': 'value', 'value': estimate, 'presentation': PRESENTED}
    claims = [_claim(run, row=index, column='value', value=value) for index, value in enumerate(values)]
    return SimpleNamespace(service=service, run=run, requirement=requirement, recipes=recipes, claims=claims,
                           presented=presented)


def _save(case, body, *, presented=True, **extra):
    saved = case.service.save_explanation_draft({
        'body': body, 'run_ids': [case.run['id']], 'evidence_recipes': case.recipes,
        'claims': [*case.claims, *([case.presented] if presented else [])], **extra})
    return {**case.service.analysis.get_draft(saved['id']), 'completion': saved['completion']}


def _outside_code(text):
    return re.sub(r'```.*?```', '', text, flags=re.S)


def _unlabelled_estimates(text):
    """Places outside code blocks where the estimate (exact or shown to 2 places) has no 估算 beside it."""
    text = _outside_code(text)
    misses = []
    for match in re.finditer(r'117,?814\.56(?:25)?', text):
        after = re.sub(r'^ ?万元', '', text[match.end():match.end() + 8])
        if not after.startswith(LABEL) and '估算' not in text[max(0, match.start() - 4):match.start()]:
            misses.append(match.group(0))
    return misses


def _item(draft):
    return draft['quality']['requirements']['items'][0]


# --- which numbers are labelled -------------------------------------------------------------------------


def test_every_export_mode_labels_the_estimate_in_prose_and_tables_and_changes_no_number(analysis, tmp_path):
    case = _case(analysis, tmp_path)
    draft = _save(case, BODY)
    assert _item(draft)['status'] == 'estimated'

    management = build_report_projection(draft, mode='management')
    audit = build_report_projection(draft, mode='audit')

    for projection in (management, audit):
        body = projection['body']
        assert '同期预算 117,814.56 万元（估算）。' in body
        assert '| 华东区域 | 109,742.028 | 117,814.5625（估算） |' in body
        assert _unlabelled_estimates(body) == []
        # Numbers from answered grains stay as they were.
        assert '109,742.028（估算）' not in body and '109,742.028 万元（估算）' not in body
        assert '325,644（估算）' not in body and '| 300（估算） |' not in body
        assert LABEL + LABEL not in body
        # Only the label is added: every number, rounding and unit is what the unlabelled export shows.
        unlabelled = build_report_projection({**draft, 'quality': {}}, mode=projection['content_mode'])
        assert LABEL not in _outside_code(unlabelled['body'])
        assert _outside_code(body).replace(LABEL, '') == _outside_code(unlabelled['body'])
    # The audit JSON is data: it stays machine-readable and keeps the cited values as they are.
    code = next(node for node in audit['document']['content'] if node['type'] == 'codeBlock')
    recorded = json.loads(code['content'][0]['text'])
    assert [claim['value'] for claim in recorded['claims'] if claim.get('row') == 3] == [ESTIMATE, ESTIMATE]
    assert not [issue for issue in draft['quality']['issues'] if issue['code'] == 'estimate_label_unplaced']


def test_the_trusted_appendix_value_of_an_estimate_is_labelled_in_the_audit_export(analysis, tmp_path):
    case = _case(analysis, tmp_path)
    draft = _save(case, BODY)
    appendix = '## 计算依据\n\n- 华东区域 1-8 月预算：117,814.5625 万元（已引用）\n- 华东区域 1-8 月实际：109,742.028 万元（已引用）'
    stored = {**draft, 'evidence_appendix': appendix,
              'evidence_appendix_meta': {'body_sha256': hashlib.sha256(appendix.encode('utf-8')).hexdigest()}}

    audit = build_report_projection(stored, mode='audit')['body']
    management = build_report_projection(stored, mode='management')['body']

    assert '- 华东区域 1-8 月预算：117,814.5625 万元（估算）（已引用）' in audit
    assert '- 华东区域 1-8 月实际：109,742.028 万元（已引用）' in audit
    assert '1-8 月预算：117,814.5625' not in management


def test_text_that_already_says_估算_beside_the_value_is_not_labelled_again(analysis, tmp_path):
    case = _case(analysis, tmp_path)
    body = '\n\n'.join([
        '华东区域同期预算（估算）117,814.56 万元。',
        '华东区域同期预算 117,814.56 万元（估算），待经理确认。',
        '按系数估算为 117,814.5625 万元。',
        TABLE.replace('117,814.5625 |', '117,814.5625（估算） |'),
    ])
    projection = build_report_projection(_save(case, body), mode='management')['body']

    assert LABEL + LABEL not in projection
    assert '同期预算（估算）117,814.56 万元。' in projection
    assert '117,814.56 万元（估算），待经理确认。' in projection
    assert '估算为 117,814.5625 万元。' in projection
    assert '| 华东区域 | 109,742.028 | 117,814.5625（估算） |' in projection
    assert projection.count('估算') == body.count('估算')


def test_limitations_that_cite_the_estimate_are_labelled_in_the_projection_and_the_markdown_export(analysis, tmp_path):
    case = _case(analysis, tmp_path)
    draft = _save(case, BODY, limitations=['华东区域 1-8 月预算 117,814.5625 万元按集团进度推算，经理确认前不作批复预算。'])

    projection = build_report_projection(draft, mode='management')
    exported = analysis.export_draft(draft['id'], 'md', mode='management')
    text = open(exported['path'], encoding='utf-8').read()

    assert projection['limitations'] == ['华东区域 1-8 月预算 117,814.5625 万元（估算）按集团进度推算，经理确认前不作批复预算。']
    assert '- 华东区域 1-8 月预算 117,814.5625 万元（估算）按集团进度推算' in projection['body']
    assert _unlabelled_estimates(text) == []
    assert text.count('117,814.5625 万元（估算）按集团进度推算') == 2


SAME = (188503.3, ESTIMATE, 117814.5625, ESTIMATE)  # 华东's counted actual equals its estimate


def test_an_estimate_equal_to_a_counted_number_is_labelled_only_where_its_own_form_tells_them_apart(analysis, tmp_path):
    """An estimate and a counted actual with the same value: the estimate shown as 117,814.56 (its presentation) is
    labelled; 117,814.5625, which the counted actual also states, never is."""
    case = _case(analysis, tmp_path, east=SAME)
    table = '| 区域 | 1-8 月实际 | 1-8 月预算 |\n| --- | ---: | ---: |\n| 华东区域 | 117,814.5625 | 117,814.56 |'
    draft = _save(case, '华东区域 1-8 月收入实际 117,814.5625 万元，同期预算 117,814.56 万元。\n\n' + table)
    assert _item(draft)['status'] == 'estimated'

    body = build_report_projection(draft, mode='management')['body']

    assert '实际 117,814.5625 万元，同期预算 117,814.56 万元（估算）。' in body
    assert '| 华东区域 | 117,814.5625 | 117,814.56（估算） |' in body
    assert [issue for issue in draft['quality']['issues'] if issue['code'] == 'estimate_label_unplaced'] == []


def test_an_estimate_the_text_cannot_tell_from_a_counted_number_is_not_labelled_and_is_a_quality_issue(analysis, tmp_path):
    case = _case(analysis, tmp_path, east=SAME)
    table = '| 区域 | 1-8 月实际 | 1-8 月预算 |\n| --- | ---: | ---: |\n| 华东区域 | 117,814.5625 | 117,814.5625 |'
    draft = _save(case, table, presented=False)

    body = build_report_projection(draft, mode='management')['body']
    issues = [issue for issue in draft['quality']['issues'] if issue['code'] == 'estimate_label_unplaced']

    assert '| 华东区域 | 117,814.5625 | 117,814.5625 |' in body
    assert len(issues) == 1 and issues[0]['same_as_counted_numbers'] == ['117,814.5625']
    assert '## 估算标注待核对' in body and issues[0]['message'] in body
    # The assistant can fix the text, so the draft is not handed to the manager as only an estimate.
    assert draft['completion']['status'] != 'awaiting_manager'
    assert draft['completion']['repair_allowed'] is True


# --- the server decides, never the text or a caller's report ---------------------------------------------


def test_an_estimate_on_a_grain_the_check_does_not_find_estimated_is_not_labelled(analysis, tmp_path):
    """Undeclared, the grain is pending (estimate_cited_for_source_grain); writing 估算 in prose makes nothing an estimate."""
    case = _case(analysis, tmp_path, declare=False)
    draft = _save(case, BODY + '\n\n以上区域预算为估算口径。')
    assert _item(draft)['status'] == 'pending'

    projection = build_report_projection(draft, mode='management')

    assert LABEL not in projection['body']
    assert '同期预算 117,814.56 万元。' in projection['body']


def test_a_caller_supplied_quality_report_labels_nothing(analysis, tmp_path):
    case = _case(analysis, tmp_path)
    forged = {'requirements': {'items': [{'id': 'revenue_1_8', 'estimated_grains': [
        {'organization': '华东区域', 'scenario': 'actual_1_8',
         'locations': [{'run_id': case.run['id'], 'row': 2, 'column': 'value'}]}]}]}}
    draft = analysis.save_draft({'title': '区域收入', 'body': BODY, 'run_ids': [case.run['id']],
                                 'claims': [*case.claims, case.presented], 'evidence_recipes': case.recipes,
                                 'requirements_id': case.requirement['id'], 'quality': forged,
                                 'explanation_check': forged})

    body = build_report_projection(draft, mode='management')['body']

    assert '109,742.028 万元，' in body and '109,742.028（估算）' not in body
    assert '117,814.56 万元（估算）' in body


def test_a_draft_without_estimated_grains_exports_exactly_as_before(analysis, tmp_path):
    case = _case(analysis, tmp_path, declare=False)
    draft = _save(case, BODY)

    management = build_report_projection(draft, mode='management')
    assert management['body'] == build_report_projection({**draft, 'quality': {}}, mode='management')['body']
    assert LABEL not in build_report_projection(draft, mode='audit')['body']


# --- every format ----------------------------------------------------------------------------------------


@pytest.mark.parametrize('mode', ['management', 'audit'])
def test_markdown_and_xlsx_exports_carry_the_label(analysis, tmp_path, mode):
    case = _case(analysis, tmp_path)
    draft = _save(case, BODY)

    markdown = open(analysis.export_draft(draft['id'], 'md', mode=mode)['path'], encoding='utf-8').read()
    assert '同期预算 117,814.56 万元（估算）。' in markdown
    assert '| 华东区域 | 109,742.028 | 117,814.5625（估算） |' in markdown
    assert _unlabelled_estimates(markdown) == []

    path = analysis.export_draft(draft['id'], 'xlsx', mode=mode)['path']
    book = load_workbook(path)
    assert '同期预算 117,814.56 万元（估算）。' in book['Report']['A3'].value
    table_texts = [row[4] for row in book['财务表格'].iter_rows(min_row=2, values_only=True)]
    assert '117,814.5625（估算）' in table_texts and '109,742.028' in table_texts
    assert not [text for text in table_texts if text and '109,742.028' in text and LABEL in text]
    prose = [row[3] for row in book['分段正文'].iter_rows(min_row=2, values_only=True) if row[3]]
    assert any('117,814.56 万元（估算）' in text for text in prose)
    # The run sheet keeps every value a number; only the estimated cell's display carries the label.
    result = book['Result 1']
    header = [cell.value for cell in result[1]]
    value_column = header.index('value') + 1
    marked = {}
    for row in result.iter_rows(min_row=2):
        cell = row[value_column - 1]
        marked[(row[header.index('organization')].value, row[header.index('scenario')].value)] = cell
        assert cell.data_type == 'n'
    assert marked[('华东区域', 'budget_1_8')].value == pytest.approx(117814.5625)
    assert LABEL in marked[('华东区域', 'budget_1_8')].number_format
    assert [key for key, cell in marked.items() if LABEL in cell.number_format] == [('华东区域', 'budget_1_8')]
    with open(path, 'rb') as handle:
        assert '估算'.encode() in _zip_member(handle.read(), 'xl/styles.xml')


def _zip_member(data, name):
    from zipfile import ZipFile

    with ZipFile(io.BytesIO(data)) as archive:
        return archive.read(name)


# --- fail loud -------------------------------------------------------------------------------------------


def test_an_estimate_written_in_prose_without_a_claim_is_a_quality_issue_and_the_export_says_so(analysis, tmp_path):
    case = _case(analysis, tmp_path)
    draft = _save(case, '华东区域 1-8 月收入实际 109,742.028 万元，同期预算约 117,815 万元。', presented=False)

    issues = [issue for issue in draft['quality']['issues'] if issue['code'] == 'estimate_label_unplaced']
    assert len(issues) == 1
    assert (issues[0]['run_id'], issues[0]['row'], issues[0]['column']) == (case.run['id'], 3, 'value')
    assert issues[0]['requirement_item_id'] == 'revenue_1_8'
    assert issues[0]['uncited_numbers'] == ['117,815']
    assert '估算' in issues[0]['message'] and '117,815' in issues[0]['message']

    for mode in ('management', 'audit'):
        body = build_report_projection(draft, mode=mode)['body']
        assert '## 估算标注待核对' in body
        assert issues[0]['message'] in body
    revalidated = case.service.validate_explanation({'draft_id': draft['id']})
    assert [issue['code'] for issue in revalidated['issues']].count('estimate_label_unplaced') == 1


def test_no_label_issue_when_the_text_binds_every_number_or_already_says_估算(analysis, tmp_path):
    case = _case(analysis, tmp_path)
    silent = _save(case, '华东区域 1-8 月收入实际 109,742.028 万元。')
    written = _save(case, '华东区域 1-8 月收入实际 109,742.028 万元，同期预算估算约 117,815 万元。', presented=False)

    for draft in (silent, written):
        assert [issue for issue in draft['quality']['issues'] if issue['code'] == 'estimate_label_unplaced'] == []
        assert '估算标注待核对' not in build_report_projection(draft, mode='management')['body']


# --- PR #29 review: forms of text the label and the check must reach ---------------------------------------


def _unplaced(draft):
    return [issue for issue in draft['quality']['issues'] if issue['code'] == 'estimate_label_unplaced']


def _document_text(document):
    def walk(node):
        if node.get('type') == 'text':
            yield node['text']
        for child in node.get('content') or []:
            yield from walk(child)

    return ''.join(walk({'content': document['content']}))


def test_an_estimate_in_slash_shorthand_is_labelled(analysis, tmp_path):
    """'实际/预算 a/b' is common shorthand; the number check skips numbers beside '/', the label must not."""
    case = _case(analysis, tmp_path)
    draft = _save(case, '华东区域 1-8 月收入实际/同期预算 109,742.028/117,814.56 万元。\n\n'
                        '华东区域预算/实际：117,814.56/109,742.028 万元。')

    body = build_report_projection(draft, mode='management')['body']

    assert '实际/同期预算 109,742.028/117,814.56 万元（估算）。' in body
    assert '预算/实际：117,814.56（估算）/109,742.028 万元。' in body
    assert '109,742.028（估算）' not in body and '109,742.028 万元（估算）' not in body
    assert _unplaced(draft) == []


def test_an_unbound_estimate_in_slash_shorthand_is_a_quality_issue(analysis, tmp_path):
    case = _case(analysis, tmp_path)
    draft = _save(case, '华东区域 1-8 月收入实际/同期预算 109,742.028/117,815 万元。', presented=False)

    issues = _unplaced(draft)
    assert len(issues) == 1 and issues[0]['uncited_numbers'] == ['117,815']
    assert issues[0]['message'] in build_report_projection(draft, mode='management')['body']


def test_a_negated_估算_beside_the_estimate_is_not_its_label(analysis, tmp_path):
    case = _case(analysis, tmp_path)
    draft = _save(case, '\n\n'.join([
        '华东区域同期预算为已批复数、非估算：117,814.56 万元。',
        '华东区域同期预算 117,814.56 万元估算口径不适用。',
        '华东区域同期预算 117,814.56 万元（非估算）。',
        TABLE,
    ]))

    body = build_report_projection(draft, mode='management')['body']

    assert '非估算：117,814.56 万元（估算）。' in body
    assert '117,814.56 万元（估算）估算口径不适用。' in body
    assert '117,814.56 万元（估算）（非估算）。' in body
    assert _unlabelled_estimates(body) == []


def test_wording_that_already_marks_the_estimate_gets_no_second_label(analysis, tmp_path):
    case = _case(analysis, tmp_path)
    lines = ['华东区域同期预算（估算值）117,814.56 万元。',
             '华东区域同期预算 117,814.56 万元，系估算。',
             '华东区域同期预算 117,814.56 万元（按集团进度估算）。']
    draft = _save(case, '\n\n'.join(lines) + '\n\n' + TABLE)

    body = build_report_projection(draft, mode='management')['body']

    for line in lines:
        assert line in body
    assert '| 华东区域 | 109,742.028 | 117,814.5625（估算） |' in body


def test_a_counted_number_stated_without_its_own_claim_is_not_labelled_as_the_estimate_it_equals(analysis, tmp_path):
    """The counted actual equals the estimate. Written without its own claim, the actual still is a counted cell
    of the run, so the text cannot mean the estimate by it."""
    case = _case(analysis, tmp_path, east=SAME)
    claims = [claim for claim in case.claims if claim['row'] != 2]
    table = '| 区域 | 1-8 月实际 | 1-8 月预算 |\n| --- | ---: | ---: |\n| 华东区域 | 117,814.5625 | 117,814.56 |'
    saved = case.service.save_explanation_draft({'body': table, 'run_ids': [case.run['id']],
                                                 'evidence_recipes': case.recipes, 'claims': [*claims, case.presented]})
    draft = analysis.get_draft(saved['id'])
    assert _item(draft).get('estimated_grains')

    body = build_report_projection(draft, mode='management')['body']

    assert '| 华东区域 | 117,814.5625 | 117,814.56（估算） |' in body
    assert _unplaced(draft) == []


def test_an_estimate_a_limitation_states_in_a_form_that_cannot_be_labelled_is_a_quality_issue(analysis, tmp_path):
    case = _case(analysis, tmp_path)
    draft = _save(case, '华东区域 1-8 月收入实际 109,742.028 万元。',
                  limitations=['华东区域 1-8 月预算约 11.78 亿元，按集团进度推算，经理确认前不作批复预算。'])

    issues = _unplaced(draft)
    assert len(issues) == 1 and issues[0]['uncited_numbers'] == ['11.78']
    body = build_report_projection(draft, mode='management')['body']
    assert '## 估算标注待核对' in body and issues[0]['message'] in body
    revalidated = case.service.validate_explanation({'draft_id': draft['id']})
    assert [issue['code'] for issue in revalidated['issues']].count('estimate_label_unplaced') == 1


def test_the_title_labels_the_estimate_it_states_in_every_export_and_is_checked(analysis, tmp_path):
    case = _case(analysis, tmp_path)
    title = '华东区域 1-8 月同期预算 117,814.56 万元'
    draft = _save(case, '华东区域 1-8 月收入实际 109,742.028 万元。', title=title)

    projection = build_report_projection(draft, mode='management')

    assert projection['title'] == title + LABEL
    assert '# ' + title + LABEL + '\n' in projection['body']
    markdown = open(analysis.export_draft(draft['id'], 'md')['path'], encoding='utf-8').read()
    assert '# ' + title + LABEL + '\n' in markdown
    assert load_workbook(analysis.export_draft(draft['id'], 'xlsx')['path'])['Report']['A1'].value == title + LABEL
    assert _unplaced(draft) == []

    rounded = _save(case, '华东区域 1-8 月收入实际 109,742.028 万元。', presented=False,
                    title='华东区域 1-8 月同期预算约 11.78 亿元')
    assert [issue['uncited_numbers'] for issue in _unplaced(rounded)] == [['11.78']]


def test_the_label_goes_after_the_whole_unit(analysis, tmp_path):
    case = _case(analysis, tmp_path)
    lines = {
        '甲：同期预算 117,814.56 万。': '甲：同期预算 117,814.56 万（估算）。',
        '乙：同期预算 117,814.56万。': '乙：同期预算 117,814.56万（估算）。',
        '丙：人均预算 117,814.56 万元/人。': '丙：人均预算 117,814.56 万元/人（估算）。',
        '丁：单价预算 117,814.56 元/吨。': '丁：单价预算 117,814.56 元/吨（估算）。',
        '戊：同期预算 117,814.56 万元人民币。': '戊：同期预算 117,814.56 万元人民币（估算）。',
        '己：单价预算 117,814.56 元/平方米。': '己：单价预算 117,814.56 元/平方米（估算）。',
    }
    draft = _save(case, '\n\n'.join(lines) + '\n\n' + TABLE)

    body = build_report_projection(draft, mode='management')['body']

    for labelled in lines.values():
        assert labelled in body


@pytest.mark.parametrize('row', [[3], {'i': 3}])
def test_validate_reports_a_malformed_claim_location_on_an_estimate_draft_instead_of_crashing(analysis, tmp_path, row):
    case = _case(analysis, tmp_path)
    draft = _save(case, BODY)
    bad = {'basis': 'calculation', 'run_id': case.run['id'], 'row': row, 'column': 'value', 'value': ESTIMATE}

    result = case.service.validate_explanation({'draft_id': draft['id'], 'claims': [*case.claims, bad]})

    assert 'invalid_claim_location' in [issue['code'] for issue in result['issues']]


def test_labels_do_not_push_a_draft_whose_export_fits_past_the_document_text_limit(analysis, tmp_path, monkeypatch):
    import foundry.report_documents as documents
    from foundry.errors import FoundryError

    case = _case(analysis, tmp_path)
    draft = _save(case, BODY + '\n\n' + '\n\n'.join(['华东区域同期预算 117,814.56 万元。'] * 100))
    stored = len(_document_text(report_document_from_markdown_body(draft['body'])))
    sizes = {}
    for mode in ('management', 'audit'):
        text = _document_text(analysis.report_projection(draft['id'], mode=mode)['document'])
        sizes[mode] = (len(text) - len(LABEL) * text.count(LABEL), len(text))
    limit = max(stored, sizes['management'][0])
    # The draft and its management export without labels fit; the labels take that export past the limit.
    assert sizes['management'][1] > limit

    monkeypatch.setattr(documents, 'MAX_DOCUMENT_TEXT', limit)
    exported = {fmt: analysis.export_draft(draft['id'], fmt, mode='management')['path'] for fmt in ('md', 'xlsx')}
    assert open(exported['md'], encoding='utf-8').read().count('117,814.56 万元（估算）') == 101
    assert load_workbook(exported['xlsx'])['Report']['A3'].value.count('117,814.56 万元（估算）') == 101
    # Only the labels get the room: the audit export's own text is over the limit and still fails.
    assert sizes['audit'][0] > limit
    with pytest.raises(FoundryError, match='exceeds|report text'):
        analysis.report_projection(draft['id'], mode='audit')


def test_labelling_a_paragraph_of_many_inline_nodes_costs_about_what_the_projection_costs(analysis, tmp_path):
    import time

    case = _case(analysis, tmp_path)
    draft = _save(case, BODY)
    content = []
    for _ in range(9000):
        content += [{'type': 'text', 'text': '甲', 'marks': [{'type': 'bold'}]}, {'type': 'text', 'text': '乙'}]
    content.append({'type': 'text', 'text': '同期预算 117,814.56 万元。'})
    stored = {**draft, 'document': {'schema_version': 1, 'template_version': 'annual-v1', 'metadata': {},
                                    'content': [{'type': 'paragraph', 'content': content}]}}

    def timed(record):
        best = float('inf')
        for _ in range(2):
            start = time.perf_counter()
            projection = build_report_projection(record, mode='management')
            best = min(best, time.perf_counter() - start)
        return best, projection

    plain, _ = timed({**stored, 'quality': {}})
    labelled, projection = timed(stored)

    assert '同期预算 117,814.56 万元（估算）。' in projection['body']
    assert labelled < 3 * plain + 0.2, (labelled, plain)


# --- PR #29 review round 2: a number is labelled only where it is bound to the estimate ---------------------
# A number is the estimate where it is written as the estimate's own claim, presentation or fact template renders
# it. A number bound to nothing that equals an estimate rounded to the digits it shows is reported instead of
# labelled. Zeros, counts, percentages and unrelated numbers are neither labelled nor reported.


def test_a_fact_template_of_the_estimate_is_labelled(analysis, tmp_path):
    case = _case(analysis, tmp_path)
    facts = case.service.prepare_explanation({'run_ids': [case.run['id']]})['facts']
    estimate = next(fact for fact in facts if fact['row'] == 3 and fact['column'] == 'value')
    actual = next(fact for fact in facts if fact['row'] == 2 and fact['column'] == 'value')
    draft = _save(case, f'华东区域实际 {{{{fact:{actual["id"]}}}}}，同期预算 {{{{fact:{estimate["id"]}}}}}。', presented=False)

    body = build_report_projection(draft, mode='management')['body']

    assert f'同期预算 {estimate["display_value"]}（估算）。' in body
    assert f'实际 {actual["display_value"]}，' in body
    assert _unplaced(draft) == []


def test_the_exact_value_padded_with_zeros_is_the_estimate_and_a_rounded_copy_is_reported(analysis, tmp_path):
    case = _case(analysis, tmp_path)
    draft = _save(case, '华东区域同期预算 117,814.5625000 万元，实际 109,742.028 万元；预算约 117,814.6 万元。',
                  presented=False)

    body = build_report_projection(draft, mode='management')['body']

    assert '同期预算 117,814.5625000 万元（估算），实际 109,742.028 万元；预算约 117,814.6 万元。' in body
    assert [issue['uncited_numbers'] for issue in _unplaced(draft)] == [['117,814.6']]


def test_zeros_are_never_labelled_or_reported_even_beside_a_zero_estimate(analysis, tmp_path):
    """Huachen: 集团本部及共享中心 has a 0 full-year budget, so its 1-8 estimate is 0. A 0 is never bound to an
    estimate: no bare 0, 0 项, 0 个 or 0% becomes an estimate, and the 0 estimate itself stays as written."""
    case = _case(analysis, tmp_path, east=(0, '0', 109742.028, '109742.028'))
    body = ('华东区域 1-8 月收入实际 109,742.028 万元，同期预算 0.00 万元。\n\n'
            '本期共 0 项调整事项，共 0 个区域超出全年预算，环比变动 0%，差异 0 万元。')
    draft = _save(case, body)
    assert _item(draft)['status'] == 'estimated'

    for mode in ('management', 'audit'):
        projection = build_report_projection(draft, mode=mode)['body']
        assert LABEL not in _outside_code(projection)
        assert '本期共 0 项调整事项，共 0 个区域超出全年预算，环比变动 0%，差异 0 万元。' in projection
        assert '同期预算 0.00 万元。' in projection
    assert _unplaced(draft) == []


def test_a_count_or_a_percentage_equal_to_an_estimate_is_not_labelled(analysis, tmp_path):
    case = _case(analysis, tmp_path, east=(9.6, '6', 109742.028, '109742.028'))
    draft = _save(case, '华东区域 1-8 月收入实际 109,742.028 万元，同期预算 6 万元；共 6 个区域超预算，环比增长 6%。',
                  presented=False, title='6 个区域 1-8 月收入', limitations=['其余 6 家子公司未纳入。'])
    assert _item(draft)['status'] == 'estimated'

    body = build_report_projection(draft, mode='management')['body']

    assert '同期预算 6 万元（估算）；共 6 个区域超预算，环比增长 6%。' in body
    assert '# 6 个区域 1-8 月收入\n' in body and '- 其余 6 家子公司未纳入。' in body
    assert _outside_code(body).count(LABEL) == 1
    assert _unplaced(draft) == []


def test_numbers_unrelated_to_an_estimate_the_text_does_not_state_raise_nothing(analysis, tmp_path):
    """The estimate is cited by claim but not written, so there is nothing to label. A count, a percentage, a
    slash or any other number in the body, limitations or title is not the estimate: the draft waits for the
    manager exactly as on main and the export gains no 估算标注待核对 section."""
    case = _case(analysis, tmp_path)
    base = '华东区域 1-8 月收入实际 109,742.028 万元，集团合并实际 325,644 万元，同期预算 300 万元，预算进度 0.625。'
    reference = _save(case, base)
    assert reference['completion']['status'] == 'awaiting_manager'
    variants = [
        {'limitations': ['共 2 个组织，华东同期预算为推算值。']},
        {'limitations': ['差异超过 5% 的区域需经理复核，其余区域差异合计 1,234.5 万元。']},
        {'title': '2 个组织 1-8 月收入'},
        {'body': base + '（1/2 级）'},
    ]
    for extra in variants:
        name = json.dumps(extra, ensure_ascii=False)
        draft = _save(case, extra.pop('body', base), **extra)
        assert _unplaced(draft) == [], name
        assert draft['completion']['status'] == reference['completion']['status'], name
        assert '估算标注待核对' not in build_report_projection(draft, mode='management')['body'], name


def test_a_note_that_calls_the_value_approved_is_no_estimate_label(analysis, tmp_path):
    case = _case(analysis, tmp_path)
    draft = _save(case, '华东区域 1-8 月收入实际 109,742.028 万元，同期预算 117,814.56 万元（批复数，已取代估算）。\n\n'
                        '华东区域取代估算：117,814.56 万元。\n\n' + TABLE)

    body = build_report_projection(draft, mode='management')['body']

    assert '同期预算 117,814.56 万元（估算）（批复数，已取代估算）。' in body
    assert '取代估算：117,814.56 万元（估算）。' in body
    assert _unlabelled_estimates(body) == []


def test_估算_before_the_noun_of_the_estimate_is_its_label(analysis, tmp_path):
    case = _case(analysis, tmp_path)
    lines = ['华东区域估算预算 117,814.56 万元。', '华东区域同期估算预算为 117,814.56 万元。']
    draft = _save(case, '\n\n'.join(lines) + '\n\n' + TABLE)

    body = build_report_projection(draft, mode='management')['body']

    for line in lines:
        assert line in body
    assert LABEL + LABEL not in body and '万元（估算）' not in body


def test_an_estimate_right_after_a_period_or_beside_a_minus_sign_is_labelled_or_reported(analysis, tmp_path):
    case = _case(analysis, tmp_path)
    draft = _save(case, '华东区域实际 109,742.028 万元；华东区域1-8月117,814.56万元为同期预算。\n\n'
                        '华东区域预算缺口 117,814.56-109,742.028 万元。')

    body = build_report_projection(draft, mode='management')['body']

    assert '华东区域1-8月117,814.56万元（估算）为同期预算。' in body
    assert '预算缺口 117,814.56（估算）-109,742.028 万元。' in body
    assert _unplaced(draft) == []

    unbound = _save(case, '华东区域实际 109,742.028 万元；华东区域1-8月117,815万元为同期预算。', presented=False)
    assert [issue['uncited_numbers'] for issue in _unplaced(unbound)] == [['117,815']]


def test_the_label_keeps_a_non_money_unit_whole(analysis, tmp_path):
    case = _case(analysis, tmp_path)
    lines = {
        '甲：同期计划 117,814.56 万吨。': '甲：同期计划 117,814.56 万吨（估算）。',
        '乙：同期计划 117,814.56 万人。': '乙：同期计划 117,814.56 万人（估算）。',
        '丙：同期计划 117,814.56 万平方米。': '丙：同期计划 117,814.56 万平方米（估算）。',
        '丁：同期计划 117,814.56 亿股。': '丁：同期计划 117,814.56 亿股（估算）。',
        '戊：同期计划 117,814.56 万吨/年。': '戊：同期计划 117,814.56 万吨/年（估算）。',
        '己：同期计划 117,814.56 吨。': '己：同期计划 117,814.56 吨（估算）。',
    }
    draft = _save(case, '\n\n'.join(lines) + '\n\n' + TABLE)

    body = build_report_projection(draft, mode='management')['body']

    for labelled in lines.values():
        assert labelled in body


def test_an_estimate_in_a_code_block_is_reported_and_one_in_a_table_caption_is_labelled(analysis, tmp_path):
    case = _case(analysis, tmp_path)
    coded = _save(case, '华东区域 1-8 月收入实际 109,742.028 万元，同期预算 117,814.56 万元。\n\n'
                        '```text\n华东区域  109,742.028    117,814.56\n```')
    issues = _unplaced(coded)
    assert len(issues) == 1 and issues[0]['in_code_numbers'] == ['117,814.56']
    body = build_report_projection(coded, mode='management')['body']
    assert '同期预算 117,814.56 万元（估算）。' in body
    assert '华东区域  109,742.028    117,814.56\n```' in body and issues[0]['message'] in body

    def with_caption(caption):
        document = report_document_from_markdown_body(BODY)
        table = next(node for node in document['content'] if node['type'] == 'table')
        table['attrs'] = {**table.get('attrs', {}), 'caption': caption}
        saved = case.service.save_explanation_draft({'document': document, 'run_ids': [case.run['id']],
                                                     'evidence_recipes': case.recipes,
                                                     'claims': [*case.claims, case.presented]})
        draft = analysis.get_draft(saved['id'])
        projection = build_report_projection(draft, mode='management')
        return draft, next(node for node in projection['document']['content'] if node['type'] == 'table')['attrs']['caption']

    captioned, caption = with_caption('华东区域 1-8 月同期预算 117,814.56 万元')
    assert _unplaced(captioned) == []
    assert caption == '华东区域 1-8 月同期预算 117,814.56 万元（估算）'
    markdown = open(analysis.export_draft(captioned['id'], 'md')['path'], encoding='utf-8').read()
    assert '表题：华东区域 1-8 月同期预算 117,814.56 万元（估算）' in markdown

    # A caption with no room for the label under the 300-character caption limit stays as written and is reported.
    full = '华东区域 1-8 月同期预算 117,814.56 万元'
    full += '。' * (300 - len(full))
    crowded, caption = with_caption(full)
    assert caption == full
    assert [issue['in_caption_numbers'] for issue in _unplaced(crowded)] == [['117,814.56']]


def test_labelled_limitations_and_the_label_section_do_not_push_an_export_that_fits_past_the_text_limit(
        analysis, tmp_path, monkeypatch):
    """Main appends the limitations unlabelled and has no 估算标注待核对 section: an export that fits there fits here."""
    import foundry.report_documents as documents

    case = _case(analysis, tmp_path)
    draft = _save(case, BODY + '\n\n华东区域同期预算约 117,815 万元。',
                  limitations=['华东区域同期预算 117,814.56 万元，按集团进度推算。'] * 60)
    assert len(_unplaced(draft)) == 1
    stored = len(_document_text(report_document_from_markdown_body(draft['body'])))
    with monkeypatch.context() as unlabelled:
        # The export as main builds it: the same draft with no estimates to label.
        read = analysis.get_draft
        unlabelled.setattr(analysis, 'get_draft', lambda draft_id: {**read(draft_id), 'quality': {}})
        plain = len(_document_text(analysis.report_projection(draft['id'], mode='management')['document']))
    labelled = _document_text(analysis.report_projection(draft['id'], mode='management')['document'])
    assert labelled.count(LABEL) >= 61 and '估算标注待核对' in labelled and len(labelled) > max(stored, plain)
    monkeypatch.setattr(documents, 'MAX_DOCUMENT_TEXT', max(stored, plain))

    markdown = open(analysis.export_draft(draft['id'], 'md', mode='management')['path'], encoding='utf-8').read()
    assert markdown.count('117,814.56 万元（估算），按集团进度推算。') >= 60
    book = load_workbook(analysis.export_draft(draft['id'], 'xlsx', mode='management')['path'])
    assert '同期预算 117,814.56 万元（估算）。' in book['Report']['A3'].value
