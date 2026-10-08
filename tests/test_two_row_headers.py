"""A two-row header (「本月」 over 「实际」「预算」) reads as one row of names, never as colliding columns."""

from __future__ import annotations

import io

from openpyxl import Workbook

from foundry.structure_proposal import propose_workbook_structure
from foundry.structure_snapshot import derive_table


def _book(build) -> bytes:
    book = Workbook()
    build(book.active)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _sheet(data: bytes) -> dict:
    return propose_workbook_structure(data)["sheets"][0]


def _period_columns(sheet, *, merge: bool) -> None:
    sheet.title = "区域对比"
    sheet["A1"] = "华宸物业 区域收入对比"
    sheet.merge_cells("A1:F1")
    sheet["A2"], sheet["E2"] = "期间：2026年8月", "单位：万元"
    sheet["A3"], sheet["B3"], sheet["C3"], sheet["E3"] = "组织", "业态", "本月", "本年累计"
    if merge:
        for cells in ("A3:A4", "B3:B4", "C3:D3", "E3:F3"):
            sheet.merge_cells(cells)
    for letter, label in zip("CDEF", ("实际", "预算", "实际", "预算")):
        sheet[f"{letter}4"] = label
    for row in range(5, 8):
        sheet.append([f"城市{row}", "住宅", row * 10, row * 11, row * 100, row * 110])


def _names(sheet: dict) -> list[str]:
    return [column["display"] for column in sheet["columns"]]


def test_a_merged_band_names_the_columns_under_it():
    sheet = _sheet(_book(lambda ws: _period_columns(ws, merge=True)))
    assert _names(sheet) == ["组织", "业态", "本月实际", "本月预算", "本年累计实际", "本年累计预算"]
    assert [column["role"] for column in sheet["columns"][:2]] == ["organization", "product"]
    assert (sheet["header_row"], sheet["header_band_row"]) == (4, 3)
    # The band row is part of the header, not a title line; the scope line above still is.
    assert sheet["title_rows"] == [1, 2]
    assert sheet["scope_hints"]["period"] == "2026-08"


def test_a_band_with_nothing_merged_is_read_from_the_repeated_labels_under_it():
    assert _names(_sheet(_book(lambda ws: _period_columns(ws, merge=False)))) == [
        "组织", "业态", "本月实际", "本月预算", "本年累计实际", "本年累计预算"]


def test_the_confirmed_table_carries_each_band_column_separately():
    data = _book(lambda ws: _period_columns(ws, merge=True))
    table = derive_table(data, _sheet(data))
    first = table["rows"][0]
    values = {name: first[name] for name in first if "本" in name}
    assert sorted(int(value) for value in values.values()) == [50, 55, 500, 550]
    assert len(table["rows"]) == 3


def test_a_trial_balance_band_over_debit_and_credit_reads_as_its_pairs():
    def build(ws):
        ws["A1"], ws["B1"] = "科目代码", "科目名称"
        ws.merge_cells("A1:A2")
        ws.merge_cells("B1:B2")
        for first, label in (("C", "期初余额"), ("E", "本期发生额"), ("G", "期末余额")):
            ws[f"{first}1"] = label
            ws.merge_cells(f"{first}1:{chr(ord(first) + 1)}1")
        for letter, label in zip("CDEFGH", ("借方", "贷方") * 3):
            ws[f"{letter}2"] = label
        ws.append(["6001", "主营业务收入", 0, 0, 0, 120, 0, 120])
        ws.append(["6401", "主营业务成本", 0, 0, 90, 0, 90, 0])

    sheet = _sheet(_book(build))
    assert _names(sheet)[2:] == ["期初余额借方", "期初余额贷方", "本期发生额借方", "本期发生额贷方", "期末余额借方", "期末余额贷方"]
    assert next(pair for pair in sheet["amount_pairs"] if pair["primary"])["label"] == "本期发生额"


def test_a_title_merged_across_the_table_is_not_a_band():
    def build(ws):
        ws["A1"] = "城市公司明细"
        ws.merge_cells("A1:C1")
        ws.append(["组织", "实际", "实际"])
        ws.append(["城市1", 1, 2])
        ws.append(["城市2", 3, 4])

    sheet = _sheet(_book(build))
    assert sheet["header_band_row"] is None
    assert _names(sheet) == ["组织", "实际", "实际"]


def test_a_one_row_header_is_left_as_it_is():
    def build(ws):
        ws.append(["组织", "业态", "金额"])
        ws.append(["城市1", "住宅", 1])
        ws.append(["城市2", "商办", 2])

    sheet = _sheet(_book(build))
    assert (sheet["header_row"], sheet["header_band_row"]) == (1, None)
    assert _names(sheet) == ["组织", "业态", "金额"]


def test_a_table_above_a_two_row_header_ends_before_the_band():
    def build(ws):
        ws.append(["组织", "金额"])
        ws.append(["城市1", 1])
        ws.append(["城市2", 2])
        ws.append(["城市3", 3])
        ws.append(["城市4", 4])
        ws.append(["组织", "本月", None, "本年", None])
        ws.append([None, "实际", "预算", "实际", "预算"])
        for cells in ("A6:A7", "B6:C6", "D6:E6"):
            ws.merge_cells(cells)
        ws.append(["城市1", 1, 2, 3, 4])

    data = _book(build)
    sheet = _sheet(data)
    above, below = sheet["table_blocks"]
    assert (above["start_row"], above["end_row"]) == (1, 5)
    assert (below["header_row"], below["start_row"]) == (7, 6)
    assert [column["display"] for column in below["columns"]] == ["组织", "本月实际", "本月预算", "本年实际", "本年预算"]
    first = {**sheet, **above, "header_row": above["header_row"]}
    assert len(derive_table(data, first)["rows"]) == 4
