import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName

from foundry.analysis import AnalysisStore
from foundry.cli import main
from foundry import finance_cli
from foundry.task import TaskStore


def test_completion_recovery_cli_is_task_editor_only(tmp_path, capsys):
    from foundry.finance_completion_operations import CompletionOperationStore

    root = tmp_path / "var"
    task = TaskStore(root).create(title="恢复正文")
    task.record_session("assistant", role="analysis")
    operations = CompletionOperationStore(task, "assistant")
    operation = operations.begin("human-message")
    operations.reserve(operation["id"], call_id="interrupted-save", payload={
        "title": "未保存分析", "body": "待恢复的正文", "run_ids": [],
    })
    operations.stop(operation["id"], "cancelled")
    assert main(["--var", str(root), "--task", task.id, "finance-completion-recoveries"]) == 0
    result = json.loads(capsys.readouterr().out)
    assert result["recoveries"][0]["body"] == "待恢复的正文"
    assert "session_id" not in result["recoveries"][0]
    assert main(["--var", str(root), "--session", "assistant", "finance-completion-recoveries"]) == 3
    assert "task editor" in capsys.readouterr().err
    other = TaskStore(root).create(title="另一任务")
    assert main(["--var", str(root), "--task", other.id, "finance-completion-recoveries"]) == 0
    assert json.loads(capsys.readouterr().out) == {"recoveries": []}


def test_source_refresh_cli_dedupes_and_preserves_old_run(tmp_path, capsys):
    root = tmp_path / "var"
    task = TaskStore(root).create(title="月度刷新")
    task.record_session("manager", role="analysis")
    args = ["--var", str(root), "--session", "manager"]
    path = task.root / "analysis" / "uploads" / "actual.csv"
    path.parent.mkdir(parents=True)
    path.write_text("management_account,amount\nrevenue,120\ncost,70\nexpense,20\n")
    spec = {"path": str(path), "scope": {
        "source_id": "monthly-file", "tenant_id": "T1", "ledger_id": "L1",
        "organization_id": "O1", "period": "2026-09", "currency": "CNY", "cutoff_at": "2026-09-06",
    }}

    def call(command, spec=None):
        assert main([*args, command, *(["--spec", json.dumps(spec)] if spec is not None else [])]) == 0
        return json.loads(capsys.readouterr().out)

    batch = call("finance-source-refresh", spec)
    context = call("finance-context-save", {"bindings": {"actual": batch["dataset_id"]}, "scopes": {
        "actual": {"ledger": "L1", "organization": "O1", "period": "2026-09", "currency": "CNY", "unit": "元", "scenario": "actual"},
    }})
    run = call("run-metric", {"context_id": context["id"]})
    assert run["result"]["rows"][0]["profit"] == "30"
    repeat = call("finance-source-refresh", spec)
    assert repeat["effective_dataset_id"] == batch["dataset_id"]
    assert repeat["duplicate_of"] == batch["id"]
    path.write_text("management_account,amount\nrevenue,130\ncost,70\nexpense,20\n")
    newer = call("finance-source-refresh", {**spec, "scope": {**spec["scope"], "cutoff_at": "2026-09-07"}})
    assert newer["dataset_id"] != batch["dataset_id"]
    assert len(call("finance-source-list")["batches"]) == 3
    old_again = AnalysisStore(task).run_analysis(run["analysis_id"])
    assert old_again["output_hash"] == run["output_hash"]
    assert len(task.sessions()[0]["snapshot_ids"]) == 2


def test_source_refresh_cannot_escape_uploads_or_child_ceiling(tmp_path, capsys):
    root = tmp_path / "var"
    task = TaskStore(root).create(title="隔离")
    task.record_session("manager", role="analysis")
    task.record_session("child", role="analysis", parent_session_id="manager")
    path = tmp_path / "outside.csv"
    path.write_text("amount\n1\n")
    spec = {"path": str(path), "scope": {
        "source_id": "file", "tenant_id": "T1", "ledger_id": "L1", "organization_id": "O1",
        "period": "2026-09", "currency": "CNY", "cutoff_at": "2026-09-06",
    }}
    for session in ("manager", "child", "unknown"):
        assert main(["--var", str(root), "--session", session, "finance-source-refresh", "--spec", json.dumps(spec)]) == 3
        capsys.readouterr()
    assert AnalysisStore(task).state()["datasets"] == []


def test_workbook_cli_writes_numeric_new_artifact_with_real_run(tmp_path, capsys):
    root = tmp_path / "var"
    task = TaskStore(root).create(title="Excel")
    source = tmp_path / "source.xlsx"
    book = Workbook()
    book.active.title = "Input"
    book.active.append(["amount"])
    book.active.append([120.35])
    output = book.create_sheet("Output")
    output["A1"] = "金额"
    output["A2"] = 0
    book.defined_names.add(DefinedName("Result", attr_text="'Output'!$A$2"))
    book.save(source)
    store = AnalysisStore(task)
    dataset = store.add_dataset(source, name="原工作簿", metadata={"mode": "table"})
    run = store.run_analysis(store.save_analysis({"question": "合计", "bindings": {"input": dataset["id"]}, "sql": "SELECT SUM(amount) AS amount FROM input"})["id"])
    args = ["--var", str(root), "--task", task.id]

    def call(command, spec):
        assert main([*args, command, "--spec", json.dumps(spec)]) == 0
        return json.loads(capsys.readouterr().out)

    assert "Result" in call("workbook-inspect", {"dataset_id": dataset["id"]})["named_ranges"]
    binding = call("workbook-bind", {"dataset_id": dataset["id"], "source": {"kind": "named_range", "name": "Result"},
        "role": "output", "writable": True, "header": False, "fields": [{"name": "amount", "policy": "amount"}]})
    patch = call("workbook-preview", {"binding_id": binding["id"], "run_id": run["id"],
        "writes": [{"cell": "A2", "field": "amount", "value": "120.35", "result_row": 0, "result_column": "amount"}]})
    result = call("workbook-apply", {"patch_id": patch["id"]})
    assert Path(result["path"]).is_file()
    result_book = load_workbook(result["path"], data_only=True)
    original = load_workbook(source)
    assert result_book["Output"]["A2"].data_type == "n"
    assert result_book["Output"]["A2"].value == 120.35
    assert original["Output"]["A2"].value == 0
    result_book.close()
    original.close()


def test_explanation_quality_cli_commands_route_to_finance_service(tmp_path, capsys, monkeypatch):
    root = tmp_path / "var"
    task = TaskStore(root).create(title="解释质量")
    task.record_session("manager", role="analysis")
    calls = []

    class FakeFinanceStore:
        def __init__(self, task, **kwargs):
            self.task = task
            self.kwargs = kwargs

        def prepare_explanation(self, spec):
            calls.append(("prepare_explanation", spec, self.kwargs))
            return {"ok": "prepare"}

        def validate_explanation(self, spec):
            calls.append(("validate_explanation", spec, self.kwargs))
            return {"ok": "validate"}

        def finance_background_save(self, spec):
            calls.append(("finance_background_save", spec, self.kwargs))
            return {"ok": "save"}

        def finance_background_list(self, spec):
            calls.append(("finance_background_list", spec, self.kwargs))
            return {"ok": "list"}

        def finance_adjustment_save(self, spec):
            calls.append(("finance_adjustment_save", spec, self.kwargs))
            return {"ok": "adj-save"}

        def finance_adjustment_list(self, spec):
            calls.append(("finance_adjustment_list", spec, self.kwargs))
            return {"ok": "adj-list"}

        def finance_adjustment_confirm(self, spec):
            calls.append(("finance_adjustment_confirm", spec, self.kwargs))
            return {"ok": "adj-confirm"}

    monkeypatch.setattr(finance_cli, "FinanceStore", FakeFinanceStore)
    args = ["--var", str(root), "--session", "manager"]
    for command, payload, expected in [
        ("prepare-explanation", {"run_ids": ["r1"]}, "prepare"),
        ("validate-explanation", {"body": "收入 10", "run_ids": ["r1"]}, "validate"),
        ("finance-background-save", {"text": "客户流失", "basis": "professional_judgment"}, "save"),
        ("finance-background-list", {"as_of": "2026-09"}, "list"),
        ("finance-adjustment-save", {"context_id": "ctx-a", "label": "总部费用调整", "amount": "5", "profit_impact": "5"}, "adj-save"),
        ("finance-adjustment-list", {"context_id": "ctx-a"}, "adj-list"),
        ("finance-adjustment-confirm", {"id": "adj-a"}, "adj-confirm"),
    ]:
        assert main([*args, command, "--spec", json.dumps(payload)]) == 0
        assert json.loads(capsys.readouterr().out) == {"ok": expected}
    assert [call[0] for call in calls] == [
        "prepare_explanation",
        "validate_explanation",
        "finance_background_save",
        "finance_background_list",
        "finance_adjustment_save",
        "finance_adjustment_list",
        "finance_adjustment_confirm",
    ]
    assert all(call[2]["session_id"] == "manager" for call in calls)


def test_completion_runtime_cli_scopes_identity_and_preserves_manual_save(tmp_path, capsys):
    task = TaskStore(tmp_path / "var").create(title="自动收尾")
    task.record_session("runtime-one", role="analysis")
    task.record_session("runtime-other", role="analysis")
    args = ["--var", str(tmp_path / "var"), "--session", "runtime-one"]
    assert main([*args, "finance-completion-begin", "--spec", json.dumps({
        "request_key": "user-event-1", "policy": {"max_repairs": 0, "timeout_seconds": 300},
    })]) == 0
    operation = json.loads(capsys.readouterr().out)
    assert operation["policy"]["max_repairs"] == 0
    spec = json.dumps({"title": "待核对", "body": "本次判断仍需资料支持。"})
    bound = [*args, "save-explanation-draft", "--completion-operation", operation["id"],
             "--completion-call-id", "actual-runtime-call", "--spec", spec]
    assert main(bound) == 0
    first = json.loads(capsys.readouterr().out)
    assert first["completion_operation"]["status"] == "needs_review"
    assert main(bound) == 0
    again = json.loads(capsys.readouterr().out)
    assert again["id"] == first["id"]
    assert main(["--var", str(tmp_path / "var"), "--session", "runtime-other",
                 "finance-completion-status", "--spec", json.dumps({"operation_id": operation["id"]})]) == 3
    capsys.readouterr()
    assert main([*args, "save-explanation-draft", "--spec", spec]) == 0
    manual = json.loads(capsys.readouterr().out)
    assert manual["id"] != first["id"]
    assert manual["completion_operation_id"] is None
