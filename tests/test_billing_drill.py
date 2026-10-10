"""往下拆: a collection rate taken apart 城市公司 → 项目 → 房产 → 账单, each level adding up to the one above.

Worked by hand from the rows below (as of 2026-08-31, a calendar year):
本年应收 = 100 + 100 + 200 + 300 + 40 = 740 (B8 void); 当期已收 = 100 + 60 + 150 = 310 (R7 lands after the date);
海口 (P1, P2) 400 / 310, 三亚 (P3) 300 / 0, and 40 / 0 on P9, which the property master does not have.
年初欠费 = (50 − 20) + 150 = 180, 本年清欠 = 30 + 50 = 80; P4 is the developer's (产权归属 地产).
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from foundry.billing_detail import BillingDetailStore
from foundry.billing_drill import Drill, export
from foundry.cli import main
from foundry.errors import FoundryError
from foundry.task import TaskStore

RECEIVABLES = [
    "应收单号,房产编号,项目,收费项目,账期,应收净额,单据状态",
    "B1,P1,椰林湾,物业服务费,2026-01,100,正常",
    "B2,P1,椰林湾,物业服务费,2026-08,100,正常",
    "B3,P2,观澜府,物业服务费,2026-03,200,正常",
    "B4,P3,海棠湾,物业服务费,2026-05,300,正常",
    "B5,P1,椰林湾,物业服务费,2025-12,50,正常",
    "B6,P4,海棠湾,物业服务费,2025-06,150,正常",
    "B7,P9,观澜府,车位管理费,2026-06,40,正常",
    "B8,P2,观澜府,物业服务费,2026-04,999,作废",
]
RECEIPTS = [
    "收款单号,收款日期,应收单号,收款金额",
    "R1,2026-01-15,B1,100",
    "R2,2026-08-20,B2,60",
    "R3,2026-04-01,B3,150",
    "R4,2025-12-20,B5,20",
    "R5,2026-02-01,B5,30",
    "R6,2026-05-01,B6,50",
    "R7,2026-09-05,B2,40",
]
PROPERTIES = ["房产编号,城市公司,产权归属", "P1,海口城市公司,业主", "P2,海口城市公司,业主", "P3,三亚城市公司,业主", "P4,三亚城市公司,地产"]


@pytest.fixture
def drilled(tmp_path):
    task = TaskStore(tmp_path / "var").create(title="往下拆")
    files = tmp_path / "files"
    files.mkdir()
    store = BillingDetailStore(task, actor="manager")
    ids = {}
    for kind, name, lines in (("receivable", "应收.csv", RECEIVABLES), ("receipt", "实收.csv", RECEIPTS), ("property", "房产.csv", PROPERTIES)):
        path = files / name
        path.write_text("\n".join(lines) + "\n", encoding="utf-8")
        ids[kind] = store.add(kind, [path], name=name[:-4])["id"]
        store.confirm(ids[kind])
    run = store.collection_run({**ids, "as_of": "2026-08-31"})
    return store, ids, run


def _rows(level):
    return [(row["label"], row["base"], row["done"], row["outstanding"]) for row in level["rows"]]


def test_each_level_adds_up_to_the_row_it_was_opened_from(drilled):
    store, _ids, run = drilled
    drill = Drill(store, run["id"])
    top = drill.level("collection_rate_current")
    assert [level["key"] for level in top["levels"]] == ["city", "project", "property_id", "bill"]
    assert top["run_check"] == {"base": "740.00", "done": "310.00", "ties": True}
    assert top["level"] == {"key": "city", "label": "城市公司"} and top["ties"]
    # The most still owed first; the bills on a property the master does not have stay in, as （空）.
    assert _rows(top) == [("三亚城市公司", "300.00", "0.00", "300.00"), ("海口城市公司", "400.00", "310.00", "90.00"),
                          ("（空）", "40.00", "0.00", "40.00")]
    assert top["rows"][2]["value"] is None and top["total"]["rate"] == "0.4189"
    city = drill.level("collection_rate_current", ["海口城市公司"])
    assert [(step["label"], step["base"], step["done"]) for step in city["path"]] == [("海口城市公司", "400.00", "310.00")]
    assert _rows(city) == [("观澜府", "200.00", "150.00", "50.00"), ("椰林湾", "200.00", "160.00", "40.00")]
    assert city["ties"] and city["gap"] == {"base": "0.00", "done": "0.00"}
    assert [row["properties"] for row in city["rows"]] == [1, 1]
    empty = drill.level("collection_rate_current", [None])
    assert _rows(empty) == [("观澜府", "40.00", "0.00", "40.00")]


def test_a_property_lists_its_bills_with_the_lines_they_were_read_from(drilled):
    store, _ids, run = drilled
    bills = Drill(store, run["id"]).level("collection_rate_current", ["海口城市公司", "椰林湾", "P1"])
    assert bills["level"]["key"] == "bill" and bills["ties"]
    assert [(row["label"], row["period"], row["outstanding"]) for row in bills["rows"]] == [("B2", "2026-08", "40.00"), ("B1", "2026-01", "0.00")]
    owed = bills["rows"][0]
    assert owed["lines"] == ["应收.csv 第 3 行"] and owed["fee_item"] == "物业服务费"
    # R7 came in after the statistics date: it is not one of this bill's receipts here.
    assert owed["receipts"] == ["实收.csv 第 3 行"] and owed["receipt_count"] == 1


def test_every_rate_comes_apart_by_its_own_base(drilled):
    store, ids, run = drilled
    drill = Drill(store, run["id"])
    arrears = drill.level("arrears_recovery_rate")
    assert arrears["run_check"] == {"base": "180.00", "done": "80.00", "ties": True}
    assert _rows(arrears) == [("三亚城市公司", "150.00", "50.00", "100.00"), ("海口城市公司", "30.00", "30.00", "0.00")]
    total = drill.level("collection_rate_total", ["海口城市公司", "椰林湾", "P1"])
    assert [(row["label"], row["place"]) for row in total["rows"]] == [("B2", "current"), ("B5", "arrears"), ("B1", "current")]
    assert drill.level("collection_rate_total")["run_check"] == {"base": "920.00", "done": "390.00", "ties": True}
    assert drill.level("collection_rate_month")["run_check"] == {"base": "100.00", "done": "60.00", "ties": True}
    every_receipt = store.collection_run({**ids, "as_of": "2026-08-31", "total_method": "all_receipts"})
    with pytest.raises(FoundryError, match="拆不到账单"):
        Drill(store, every_receipt["id"]).level("collection_rate_total")
    # Leaving the developer's property out leaves it out of every level.
    without = store.collection_run({**ids, "as_of": "2026-08-31", "exclude_related": True})
    assert _rows(Drill(store, without["id"]).level("arrears_recovery_rate")) == [("海口城市公司", "30.00", "30.00", "0.00")]


def test_it_reads_the_run_as_filed(drilled, monkeypatch):
    store, ids, run = drilled
    # A mapping confirmed after the run does not change how the run comes apart.
    store.confirm(ids["receivable"], {"bill_id": "应收单号", "property_id": "房产编号", "period": "账期", "amount_net": "应收净额"})
    drill = Drill(store, run["id"])
    assert [level["key"] for level in drill.level("collection_rate_current")["levels"]] == ["city", "project", "property_id", "bill"]
    # The bills are worked out once and kept beside the run.
    cached = list((store.root / "runs").glob(f"{run['id']}.drill-v*.duckdb"))
    assert len(cached) == 1
    monkeypatch.setattr("foundry.billing_drill._prepare", lambda *args: (_ for _ in ()).throw(AssertionError("worked out again")))
    assert Drill(store, run["id"]).level("collection_rate_current", ["三亚城市公司"])["ties"]
    with pytest.raises(FoundryError, match="has no 项目 不存在"):
        drill.level("collection_rate_current", ["海口城市公司", "不存在"])
    with pytest.raises(FoundryError, match="path names up to 3 values"):
        drill.level("collection_rate_current", ["a", "b", "c", "d"])
    with pytest.raises(FoundryError, match="figure is one of"):
        drill.level("waiver_rate")
    filed = store.root / "runs" / f"{run['id']}.json"
    record = json.loads(filed.read_text(encoding="utf-8"))
    filed.write_text(json.dumps({**record, "engine_version": 0}), encoding="utf-8")
    with pytest.raises(FoundryError, match="run it again"):
        Drill(store, run["id"])


def test_without_a_property_master_or_projects_it_starts_at_the_property(tmp_path):
    task = TaskStore(tmp_path / "var").create(title="往下拆")
    store = BillingDetailStore(task, actor="manager")
    receivable, receipt = tmp_path / "应收.csv", tmp_path / "实收.csv"
    receivable.write_text("应收单号,房产编号,账期,应收净额\nB1,P1,2026-01,100\nB2,P2,2026-02,50\n", encoding="utf-8")
    receipt.write_text("收款日期,应收单号,收款金额\n2026-01-10,B1,100\n", encoding="utf-8")
    ids = {"receivable": store.add("receivable", [receivable], name="应收")["id"], "receipt": store.add("receipt", [receipt], name="实收")["id"]}
    for detail in ids.values():
        store.confirm(detail)
    run = store.collection_run({**ids, "as_of": "2026-08-31"})
    level = Drill(store, run["id"]).level("collection_rate_current")
    assert [step["key"] for step in level["levels"]] == ["property_id", "bill"]
    assert _rows(level) == [("P2", "50.00", "0.00", "50.00"), ("P1", "100.00", "100.00", "0.00")]


def test_a_level_exports_every_row_to_excel(drilled, tmp_path):
    store, _ids, run = drilled
    drill = Drill(store, run["id"])
    target = tmp_path / "out" / "海口.xlsx"
    assert export(drill, "collection_rate_current", ["海口城市公司"], str(target)) == {
        "path": str(target), "rows": 2, "level": {"key": "project", "label": "项目"}}
    sheet = load_workbook(target).active
    rows = [row for row in sheet.iter_rows(values_only=True)]
    assert rows[2] == ("位置：全部 › 海口城市公司",) + (None,) * (len(rows[2]) - 1)
    header = rows.index(("项目", "户数", "本年应收", "当期已收", "未收", "当期收缴率"))
    assert rows[header + 1] == ("观澜府", 1, 200.0, 150.0, 50.0, 0.75)
    bills = tmp_path / "bills.xlsx"
    export(drill, "collection_rate_current", ["海口城市公司", "椰林湾", "P1"], str(bills))
    found = [row for row in load_workbook(bills).active.iter_rows(values_only=True) if row and row[0] == "B2"]
    assert found == [("B2", "物业服务费", "2026-08", 100.0, 60.0, 40.0, 0.6, "应收.csv 第 3 行", "实收.csv 第 3 行")]
    with pytest.raises(FoundryError, match="absolute path ending in .xlsx"):
        export(drill, "collection_rate_current", [], "out.xlsx")


def test_the_cli_takes_a_run_apart(drilled, capsys):
    store, _ids, run = drilled
    var = Path(store.task.root).parents[1]
    spec = json.dumps({"run": run["id"], "figure": "collection_rate_current", "path": ["三亚城市公司"]}, ensure_ascii=False)
    assert main(["--var", str(var), "--task", store.task.id, "collection-drill", "--spec", spec]) == 0
    level = json.loads(capsys.readouterr().out)
    assert [row["label"] for row in level["rows"]] == ["海棠湾"] and level["ties"]


def test_a_level_reads_out_only_the_rows_it_shows_and_a_bill_counts_every_line_it_came_from(tmp_path):
    # No property master and no project column: the first level is every property. B1 was exported over seven lines.
    task = TaskStore(tmp_path / "var").create(title="往下拆")
    files = tmp_path / "files"
    files.mkdir()
    receivables = ["应收单号,房产编号,收费项目,账期,应收净额,单据状态", *["B1,P1,物业服务费,2026-01,10,正常"] * 7,
                   "B2,P2,物业服务费,2026-02,50,正常", "B3,P3,物业服务费,2026-03,40,正常",
                   "B4,P4,物业服务费,2026-04,30,正常", "B5,P5,物业服务费,2026-05,20,正常"]
    receipts = ["收款单号,收款日期,应收单号,收款金额", "R1,2026-03-01,B2,50"]
    store = BillingDetailStore(task, actor="manager")
    ids = {}
    for kind, name, lines in (("receivable", "应收.csv", receivables), ("receipt", "实收.csv", receipts)):
        path = files / name
        path.write_text("\n".join(lines) + "\n", encoding="utf-8")
        ids[kind] = store.add(kind, [path], name=name[:-4])["id"]
        store.confirm(ids[kind])
    run = store.collection_run({**ids, "as_of": "2026-08-31"})
    drill = Drill(store, run["id"])
    every = drill.level("collection_rate_current", [], limit=None)
    shown = drill.level("collection_rate_current", [], limit=2)
    assert [row["label"] for row in every["rows"]] == ["P1", "P3", "P4", "P5", "P2"]
    # The two with most still owed, while the count and the totals are still the whole level's.
    assert shown["rows"] == every["rows"][:2]
    assert (shown["row_count"], shown["total"], shown["ties"]) == (5, every["total"], True)
    assert shown["total"]["base"] == "210.00" and shown["total"]["done"] == "50.00"
    [bill] = drill.level("collection_rate_current", ["P1"])["rows"]
    assert (bill["outstanding"], bill["line_count"], len(bill["lines"])) == ("70.00", 7, 5)
    # The workbook names every line.
    target = tmp_path / "P1.xlsx"
    export(drill, "collection_rate_current", ["P1"], str(target))
    [found] = [row for row in load_workbook(target).active.iter_rows(values_only=True) if row and row[0] == "B1"]
    assert found[7] == "；".join(f"应收.csv 第 {line} 行" for line in range(2, 9))
