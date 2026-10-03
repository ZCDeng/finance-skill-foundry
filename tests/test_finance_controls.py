from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from foundry.analysis import AnalysisStore
from foundry.errors import FoundryError
from foundry.finance_service import FinanceStore
from foundry.source_refresh import SourceRefreshService, SourceScope
from foundry.task import TaskStore


def _task(tmp_path: Path):
    task = TaskStore(tmp_path / "var").create(title="勾稽")
    return task, AnalysisStore(task), FinanceStore(task), FinanceStore(task, actor="manager")


def _ar_pack(path: Path, *, provisions: list[int], total: int) -> Path:
    """The F-01 shape: bill-year detail rows, a 合计 row, a ratio column and a rate column."""
    book = Workbook()
    sheet = book.active
    sheet.title = "08·AR"
    sheet.append(["欠费结构"])
    sheet.append(["华宸物业  |  2026年1-8月  |  单位：万元"])
    sheet.append([])
    sheet.append(["账单年度", "欠费余额", "占比", "已计提坏账", "计提比例", "预计可收回"])
    balances = [70500, 35700, 19800, 11700, 6750, 5550]
    for year, balance, provision in zip(range(2026, 2020, -1), balances, provisions):
        sheet.append([f"{year}年账单", balance, balance / 150000, provision, provision / balance, balance - provision])
    sheet.append(["合　计", 150000, 1, total, total / 150000, 150000 - total])
    sheet.append([])
    sheet.append(["区域", "本年账单", "收缴率"])
    sheet.append(["华东", 100, 0.9])
    sheet.append(["华南", 50, 0.8])
    sheet.append(["合计 / 加权", 150, 0.8667])
    book.save(path)
    return path


def _ledger(path: Path, collections: list[float], *, sheet_name="11·合同台账") -> Path:
    """The F-02 shape: a contract ledger with more rows than an evidence recipe may read at once."""
    book = Workbook()
    sheet = book.active
    sheet.title = sheet_name
    sheet.append(["物业服务合同台账"])
    sheet.append(["华宸物业  |  台账基准日 2026-08-31  |  金额单位：万元"])
    sheet.append([])
    sheet.append(["合同编号", "本年实收（万元）"])
    for index, amount in enumerate(collections, start=1):
        sheet.append([f"WY-{index:04d}", amount])
    sheet.append([f"合　计（{len(collections)} 份）", sum(collections)])
    book.save(path)
    return path


def _side(dataset: dict, sheet: str, cells: str, **extra) -> dict:
    return {"dataset_id": dataset["id"], "source_hash": dataset["source_hash"], "sheet": sheet, "range": cells, **extra}


def test_suggestions_find_a_detail_column_that_does_not_add_up_to_its_total(tmp_path):
    _task_obj, analysis, assistant, _manager = _task(tmp_path)
    analysis.add_dataset(_ar_pack(tmp_path / "ar.xlsx", provisions=[3525, 5355, 6930, 7020, 5400, 5550], total=22125), name="月结")
    suggestions = assistant.finance_controls_propose({})["suggestions"]
    by_column = {row["control"]["left"]["range"][0] + ":" + row["control"]["right"]["range"]: row for row in suggestions}
    # 欠费余额 and 预计可收回 / 已计提坏账 are sums; 占比 and 计提比例 are not offered.
    assert by_column["B:B11"]["preview"]["status"] == "pass"
    assert by_column["D:D11"]["preview"]["status"] == "break"
    assert by_column["D:D11"]["preview"]["difference"] == "11655"
    assert by_column["F:F11"]["preview"]["difference"] == "-11655"
    assert not any(key.startswith(("C:", "E:")) for key in by_column)
    # Under a weighted total, a column that is not an amount (收缴率) is never offered; 本年账单 is.
    assert "B:B16" in by_column and "C:C16" not in by_column
    assert by_column["D:D11"]["control"]["left"]["range"] == "D5:D10"


def test_a_pack_whose_details_add_up_raises_no_break(tmp_path):
    _task_obj, analysis, assistant, _manager = _task(tmp_path)
    analysis.add_dataset(_ar_pack(tmp_path / "ar.xlsx", provisions=[2069, 3144, 4069, 4122, 3171, 5550], total=22125), name="月结")
    previews = [row["preview"]["status"] for row in assistant.finance_controls_propose({})["suggestions"]]
    assert previews and set(previews) == {"pass"}


def test_a_declared_tie_out_across_two_workbooks_finds_the_gap_over_six_hundred_rows(tmp_path):
    _task_obj, analysis, assistant, manager = _task(tmp_path)
    ledger = analysis.add_dataset(_ledger(tmp_path / "ledger.xlsx", [409.5] * 600), name="台账")
    close = analysis.add_dataset(_ar_pack(tmp_path / "ar.xlsx", provisions=[2069, 3144, 4069, 4122, 3171, 5550], total=22125), name="月结")
    saved = assistant.finance_controls_save({"controls": [{
        "kind": "tie_out", "label": "台账实收 = 月结区域账单", "unit": "万元",
        "left": _side(ledger, "11·合同台账", "B5:B604"), "right": _side(close, "08·AR", "B11"),
    }]})["controls"][0]
    assert saved["state"] == "proposed" and saved["origin"] == "declared"
    run = assistant.finance_controls_run({"ids": [saved["id"]]})
    (result,) = run["results"]
    assert result["status"] == "break"
    assert result["left"]["value"] == "245700" and result["left"]["cells"] == 600
    assert result["difference"] == "95700"
    # Not confirmed, so it does not count toward the close yet.
    assert run["summary"]["confirmed"] == 0 and run["summary"]["break"] == 0
    manager.finance_controls_confirm({"id": saved["id"]})
    assert assistant.finance_controls_run({})["summary"]["break"] == 1


def test_a_unit_conversion_is_allowed_and_a_coefficient_is_not(tmp_path):
    _task_obj, analysis, assistant, _manager = _task(tmp_path)
    close = analysis.add_dataset(_ar_pack(tmp_path / "ar.xlsx", provisions=[2069, 3144, 4069, 4122, 3171, 5550], total=22125), name="月结")
    control = {"label": "元换万元", "left": _side(close, "08·AR", "B11", scale="10000"), "right": _side(close, "08·AR", "B11")}
    saved = assistant.finance_controls_save({"control": control})["controls"][0]
    assert assistant.finance_controls_run({"ids": [saved["id"]]})["results"][0]["difference"] == "1499850000"
    with pytest.raises(FoundryError, match="unit conversion"):
        assistant.finance_controls_save({"control": {**control, "left": _side(close, "08·AR", "B11", scale="1.048")}})


def test_only_the_manager_confirms_or_retires_and_a_repeat_save_returns_the_live_control(tmp_path):
    _task_obj, analysis, assistant, manager = _task(tmp_path)
    close = analysis.add_dataset(_ar_pack(tmp_path / "ar.xlsx", provisions=[3525, 5355, 6930, 7020, 5400, 5550], total=22125), name="月结")
    control = {"kind": "subtotal", "label": "坏账分项", "left": _side(close, "08·AR", "D5:D10"), "right": _side(close, "08·AR", "D11")}
    first = assistant.finance_controls_save({"control": control})["controls"][0]
    assert assistant.finance_controls_save({"control": control})["controls"][0]["id"] == first["id"]
    renamed = assistant.finance_controls_save({"control": {**control, "label": "坏账分项合计"}})["controls"][0]
    assert renamed["parent_id"] == first["id"]
    first = renamed
    with pytest.raises(FoundryError, match="only the manager"):
        assistant.finance_controls_confirm({"id": first["id"]})
    with pytest.raises(FoundryError, match="only the manager"):
        assistant.finance_controls_save({"control": control, "confirm": True})
    confirmed = manager.finance_controls_confirm({"id": first["id"]})
    assert confirmed["state"] == "confirmed" and confirmed["parent_id"] == first["id"]
    with pytest.raises(FoundryError, match="only the manager"):
        assistant.finance_controls_retire({"id": confirmed["id"], "reason": "不需要"})
    with pytest.raises(FoundryError, match="requires a reason"):
        manager.finance_controls_retire({"id": confirmed["id"]})
    manager.finance_controls_retire({"id": confirmed["id"], "reason": "改用跨表控制"})
    assert assistant.finance_controls_list({})["controls"] == []
    assert assistant.finance_controls_list({"include_retired": True})["controls"][0]["state"] == "retired"


def test_a_subtotal_must_stay_on_one_sheet(tmp_path):
    _task_obj, analysis, assistant, _manager = _task(tmp_path)
    close = analysis.add_dataset(_ar_pack(tmp_path / "ar.xlsx", provisions=[2069, 3144, 4069, 4122, 3171, 5550], total=22125), name="月结")
    ledger = analysis.add_dataset(_ledger(tmp_path / "ledger.xlsx", [1.0, 2.0]), name="台账")
    with pytest.raises(FoundryError, match="one sheet"):
        assistant.finance_controls_save({"control": {"kind": "subtotal", "label": "x", "left": _side(close, "08·AR", "B5:B10"), "right": _side(ledger, "11·合同台账", "B7")}})


def test_a_control_bound_to_a_snapshot_the_task_no_longer_holds_is_stale_not_green(tmp_path):
    _task_obj, analysis, assistant, _manager = _task(tmp_path)
    close = analysis.add_dataset(_ar_pack(tmp_path / "ar.xlsx", provisions=[2069, 3144, 4069, 4122, 3171, 5550], total=22125), name="月结")
    control = {"label": "余额", "left": _side(close, "08·AR", "B5:B10"), "right": {**_side(close, "08·AR", "B11"), "source_hash": "0" * 64}}
    saved = assistant.finance_controls_save({"control": control})["controls"][0]
    (result,) = assistant.finance_controls_run({"ids": [saved["id"]]})["results"]
    assert result["status"] == "stale"


def _ledger_export(path: Path, amount: int) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = "明细"
    sheet.append(["document_id", "line_id", "amount"])
    sheet.append(["JV-1", "1", amount])
    sheet.append(["JV-1", "2", 20])
    book.save(path)
    return path


def test_a_refreshed_source_makes_the_old_result_stale(tmp_path):
    task, analysis, assistant, _manager = _task(tmp_path)
    scope = SourceScope(source_id="erp", tenant_id="t", ledger_id="l", organization_id="o", period="2026-08",
                        currency="CNY", cutoff_at="2026-09-05T00:00:00Z")
    refresh = SourceRefreshService(task)
    first = refresh.refresh_file(_ledger_export(tmp_path / "ledger.xlsx", 1), scope=scope, name="ledger")
    dataset = analysis.read_dataset(first["dataset_id"])
    control = {"label": "两行合计", "left": _side(dataset, "明细", "C2:C3"), "right": _side(dataset, "明细", "C2:C3")}
    saved = assistant.finance_controls_save({"control": control})["controls"][0]
    # A refreshed ledger is filed as rows, not a workbook: said, never summed as zero.
    before = assistant.finance_controls_run({})["results"][0]
    assert before["status"] == "error" and "Excel" in before["reason"]
    refresh.refresh_file(_ledger_export(tmp_path / "ledger.xlsx", 5), scope=scope, name="ledger")
    listed = assistant.finance_controls_list({})["controls"][0]
    assert listed["id"] == saved["id"] and listed["result"]["status"] == "stale"
    assert assistant.finance_controls_run({})["results"][0]["status"] == "stale"


def test_a_session_narrowed_to_one_snapshot_cannot_declare_a_control_on_another(tmp_path):
    task, analysis, _assistant, _manager = _task(tmp_path)
    close = analysis.add_dataset(_ar_pack(tmp_path / "ar.xlsx", provisions=[2069, 3144, 4069, 4122, 3171, 5550], total=22125), name="月结")
    ledger = analysis.add_dataset(_ledger(tmp_path / "ledger.xlsx", [1.0, 2.0]), name="台账")
    narrowed = FinanceStore(task, snapshot_ids=[close["id"]])
    with pytest.raises(FoundryError, match="outside this session"):
        narrowed.finance_controls_save({"control": {"label": "x", "left": _side(close, "08·AR", "B11"), "right": _side(ledger, "11·合同台账", "B7")}})
    assert all(row["control"]["left"]["dataset_id"] == close["id"] for row in narrowed.finance_controls_propose({})["suggestions"])


def test_a_large_gap_over_many_whole_numbers_is_a_break_and_never_rounding(tmp_path):
    """Review: the rounding allowance grew with the row count and swallowed a 300 万 gap over 600 integer rows."""
    _task_obj, analysis, assistant, manager = _task(tmp_path)
    ledger = analysis.add_dataset(_ledger(tmp_path / "ledger.xlsx", [10000.0] * 600), name="台账")
    close = analysis.add_dataset(_ar_pack(tmp_path / "ar.xlsx", provisions=[2069, 3144, 4069, 4122, 3171, 5550], total=22125), name="月结")
    # 600 × 10,000 元 = 600 万元 against 150,000 万元 at B11: an enormous break, scaled 元 → 万元.
    saved = manager.finance_controls_save({"confirm": True, "control": {
        "label": "元换万元", "left": _side(ledger, "11·合同台账", "B5:B604", scale="0.0001"), "right": _side(close, "08·AR", "B11")}})["controls"][0]
    run = assistant.finance_controls_run({"ids": [saved["id"]]})
    (result,) = run["results"]
    assert result["status"] == "break" and result["within_rounding"] is False
    assert run["summary"]["break"] == 1


def test_a_gap_display_rounding_could_explain_is_still_counted_as_a_break(tmp_path):
    _task_obj, analysis, assistant, manager = _task(tmp_path)
    ledger = analysis.add_dataset(_ledger(tmp_path / "ledger.xlsx", [0.3333, 0.3333, 0.3333]), name="台账")
    path = tmp_path / "total.xlsx"
    book = Workbook(); book.active.title = "S"; book.active.append(["合计"]); book.active.append([1.0]); book.save(path)
    total = analysis.add_dataset(path, name="合计")
    saved = manager.finance_controls_save({"confirm": True, "control": {
        "label": "三分之一", "left": _side(ledger, "11·合同台账", "B5:B7"), "right": _side(total, "S", "A2")}})["controls"][0]
    run = assistant.finance_controls_run({"ids": [saved["id"]]})
    assert run["results"][0]["within_rounding"] is True
    assert run["summary"]["break"] == 1 and run["summary"]["break_within_rounding"] == 1


def test_a_number_typed_as_text_inside_a_range_is_an_error_not_a_skipped_blank(tmp_path):
    _task_obj, analysis, assistant, _manager = _task(tmp_path)
    path = tmp_path / "text.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "S"
    for row in (["项目", "金额"], ["甲", 100], ["乙", "1,200"], ["丙", 50], ["合计", 150]):
        sheet.append(row)
    book.save(path)
    dataset = analysis.add_dataset(path, name="文字数字")
    saved = assistant.finance_controls_save({"control": {"kind": "subtotal", "label": "x", "left": _side(dataset, "S", "B2:B4"), "right": _side(dataset, "S", "B5")}})["controls"][0]
    (result,) = assistant.finance_controls_run({"ids": [saved["id"]]})["results"]
    assert result["status"] == "error" and "文字" in result["reason"]
    assert result["left"]["text_cells"] == 1


def test_uploading_a_corrected_pack_again_makes_controls_on_the_old_one_stale(tmp_path):
    _task_obj, analysis, assistant, manager = _task(tmp_path)
    old = analysis.add_dataset(_ar_pack(tmp_path / "v1.xlsx", provisions=[2069, 3144, 4069, 4122, 3171, 5550], total=22125), name="月结")
    suggestions = [row["control"] for row in assistant.finance_controls_propose({})["suggestions"]]
    manager.finance_controls_save({"confirm": True, "origin": "suggested", "controls": suggestions})
    assert assistant.finance_controls_run({})["summary"]["stale"] == 0
    new = analysis.add_dataset(_ar_pack(tmp_path / "v2.xlsx", provisions=[3525, 5355, 6930, 7020, 5400, 5550], total=22125), name="月结")
    summary = assistant.finance_controls_run({})["summary"]
    assert summary["pass"] == 0 and summary["stale"] == len(suggestions)
    assert {row["control"]["left"]["dataset_id"] for row in assistant.finance_controls_propose({})["suggestions"]} == {new["id"]}
    assert old["id"] != new["id"]


def test_a_retired_check_saved_again_revives_its_own_lineage_and_never_runs_twice(tmp_path):
    _task_obj, analysis, assistant, manager = _task(tmp_path)
    close = analysis.add_dataset(_ar_pack(tmp_path / "ar.xlsx", provisions=[3525, 5355, 6930, 7020, 5400, 5550], total=22125), name="月结")
    control = {"kind": "subtotal", "label": "坏账分项", "left": _side(close, "08·AR", "D5:D10"), "right": _side(close, "08·AR", "D11")}
    confirmed = manager.finance_controls_save({"control": control, "confirm": True})["controls"][0]
    retired = manager.finance_controls_retire({"id": confirmed["id"], "reason": "暂不检查"})
    revived = assistant.finance_controls_save({"control": {**control, "tolerance": "1000000"}})["controls"][0]
    assert revived["parent_id"] == retired["id"] and revived["state"] == "proposed"
    manager.finance_controls_confirm({"id": revived["id"]})
    assert len(assistant.finance_controls_list({})["controls"]) == 1


def test_the_assistant_cannot_quietly_change_a_confirmed_control(tmp_path):
    _task_obj, analysis, assistant, manager = _task(tmp_path)
    close = analysis.add_dataset(_ar_pack(tmp_path / "ar.xlsx", provisions=[3525, 5355, 6930, 7020, 5400, 5550], total=22125), name="月结")
    control = {"kind": "subtotal", "label": "坏账分项", "left": _side(close, "08·AR", "D5:D10"), "right": _side(close, "08·AR", "D11")}
    manager.finance_controls_save({"control": control, "confirm": True})
    with pytest.raises(FoundryError, match="only be changed by the manager"):
        assistant.finance_controls_save({"control": {**control, "tolerance": "1000000"}})
    widened = manager.finance_controls_save({"control": {**control, "tolerance": "10"}})["controls"][0]
    assert widened["state"] == "confirmed" and widened["control"]["tolerance"] == "10"


def test_breakdown_rows_and_provision_notes_do_not_become_totals(tmp_path):
    _task_obj, analysis, assistant, _manager = _task(tmp_path)
    path = tmp_path / "breakdown.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "S"
    for row in (["项目", "金额"], ["营业收入", 100], ["其中：物业费", 80], ["其中：增值服务", 20], ["营业成本", 60], ["合计", 160],
                [], ["项目", "金额"], ["甲", 1], ["乙", 2], ["按组合计提", 3], ["丙", 4], ["合计", 10]):
        sheet.append(row)
    book.save(path)
    analysis.add_dataset(path, name="分解")
    suggestions = assistant.finance_controls_propose({})["suggestions"]
    ranges = {row["control"]["left"]["range"] for row in suggestions}
    assert "B2:B5" not in ranges  # 其中 rows would count 营业收入 twice
    assert ranges == {"B9:B12"}
    assert all(row["preview"]["status"] == "pass" for row in suggestions)


def test_a_drafted_two_place_comparison_previews_its_gap_without_saving(tmp_path):
    task, analysis, assistant, _manager = _task(tmp_path)
    ledger = analysis.add_dataset(_ledger(tmp_path / "ledger.xlsx", [409.5] * 600), name="台账")
    close = analysis.add_dataset(_ar_pack(tmp_path / "ar.xlsx", provisions=[2069, 3144, 4069, 4122, 3171, 5550], total=22125), name="月结")
    sources = {row["name"]: row for row in assistant.finance_controls_sources({})["sources"]}
    assert sources["台账"]["sheets"] == ["11·合同台账"] and sources["月结"]["sheets"] == ["08·AR"]
    assert sources["台账"]["source_hash"] == ledger["source_hash"]
    drafted = {"kind": "tie_out", "label": "台账实收 对 月结区域账单",
               "left": _side(ledger, "11·合同台账", "B605"), "right": _side(close, "08·AR", "B11")}
    preview = assistant.finance_controls_preview({"control": drafted})
    assert preview["preview"]["status"] == "break" and preview["preview"]["difference"] == "95700"
    assert assistant.finance_controls_list({})["controls"] == []
    narrowed = FinanceStore(task, snapshot_ids=[close["id"]])
    with pytest.raises(FoundryError, match="outside this session"):
        narrowed.finance_controls_preview({"control": drafted})
    assert [row["name"] for row in narrowed.finance_controls_sources({})["sources"]] == ["月结"]
