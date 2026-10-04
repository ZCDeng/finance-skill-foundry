"""工作簿公式体检: hints read from an uploaded workbook, never a refusal."""

from __future__ import annotations

import io
import json
import re
import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from foundry import analysis_data, workbook_health
from foundry.analysis import AnalysisStore, _hash
from foundry.cli import main
from foundry.errors import FoundryError
from foundry.finance_files import FinanceFiles
from foundry.structure_proposal import propose_workbook_structure
from foundry.task import TaskStore
from foundry.workbook_health import check_workbook

sys.path.insert(0, str(Path(__file__).parent / "fixtures" / "structure"))
import build_huachen_like as fixture  # noqa: E402

SHEET1 = "xl/worksheets/sheet1.xml"


def _bytes(book: Workbook) -> bytes:
    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def _book(formulas: dict[str, object], *, title: str = "汇总") -> Workbook:
    book = Workbook()
    sheet = book.active
    sheet.title = title
    for row in range(1, 11):
        sheet.cell(row, 1, row)
    for cell, value in formulas.items():
        sheet[cell] = value
    return book


def _flagged(result: dict, code: str) -> list[tuple[str, str | None]]:
    issue = _issue(result, code)
    return [(sample["cell"], sample.get("detail")) for sample in issue["samples"]] if issue else []


def _rewrite(data: bytes, edits: dict[str, object] | None = None, extra: dict[str, bytes] | None = None) -> bytes:
    """Rewrite package members the way a non-openpyxl writer (Excel) would have written them."""
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as rewritten:
        for item in source.infolist():
            content = source.read(item.filename)
            edit = (edits or {}).get(item.filename)
            if edit is not None:
                content = edit(content.decode()).encode()
            rewritten.writestr(item, content)
        for name, content in (extra or {}).items():
            rewritten.writestr(name, content)
    return out.getvalue()


def _cells(replacements: dict[str, str]):
    """Replace whole <c> elements of a sheet by address."""

    def edit(xml: str) -> str:
        for address, cell in replacements.items():
            xml, count = re.subn(rf'<c r="{address}"(?:[^>]*/>|[^>]*>.*?</c>)', cell, xml, flags=re.S)
            assert count == 1, address
        return xml

    return edit


def _issue(result: dict, code: str) -> dict | None:
    return next((issue for issue in result["issues"] if issue["code"] == code), None)


def _codes(result: dict) -> list[str]:
    return [issue["code"] for issue in result["issues"]]


def _only_plain_values(value) -> bool:
    if isinstance(value, dict):
        return all(isinstance(key, str) and _only_plain_values(item) for key, item in value.items())
    if isinstance(value, list):
        return all(_only_plain_values(item) for item in value)
    return isinstance(value, (str, int, bool))


def test_clean_workbook_has_no_issues_and_the_documented_shape() -> None:
    book = _book({})
    book.active["B1"] = 10
    result = check_workbook(_bytes(book))
    assert result == {
        "schema_version": 1,
        "status": "checked",
        "checked_formula_cells": 0,
        "truncated": False,
        "sheet_count": 1,
        "issue_count": 0,
        "issues": [],
    }


def test_hardcoded_numbers_are_flagged_and_structural_arguments_are_not() -> None:
    flagged = {"B2": "=A5*1.13", "B3": "=SUM(A2:A9)+2793.28", "B4": "=A5*13%", "B5": "=DATE(2026,9,30)"}
    quiet = {
        "C2": "=ROUND(A5,2)", "C3": "=INDEX(A1:A9,3)", "C4": "=VLOOKUP(A1,A1:A9,3,0)", "C5": "=MID(A1,3,4)",
        "C6": "=A5/10000", "C7": "=A5/100000000", "C8": "=A5*12/365", "C9": "=-1*A5", "C10": "=SUM({1,2,3;4,5,6})",
        "D2": '=IF(A1="1.13",1,0)', "D3": "=OFFSET(A1,3,4,5,6)", "D4": "=CHOOSE(3,A1,A2,A3)", "D5": "=LEFT(A1,7)",
        "D6": "=MATCH(A1,A1:A9,-1)", "D7": "=ROUNDDOWN((A5+A6)/1E+4,4)", "D8": "=A5*(2)", "D9": "=TEXT(A1,7)",
    }
    result = check_workbook(_bytes(_book({**flagged, **quiet})))
    issue = _issue(result, "hardcoded_number")
    assert issue is not None and issue["severity"] == "warning"
    assert issue["count"] == 4
    assert [(sample["cell"], sample["formula"], sample["detail"]) for sample in issue["samples"]] == [
        ("B2", "=A5*1.13", "1.13"),
        ("B3", "=SUM(A2:A9)+2793.28", "2793.28"),
        ("B4", "=A5*13%", "13%"),
        ("B5", "=DATE(2026,9,30)", "2026, 9, 30"),
    ]
    assert issue["sheets"] == ["汇总"]


def test_error_literals_in_formulas_but_not_quoted_text() -> None:
    data = _bytes(_book({"B2": "=Sheet1!#REF!+1", "B3": "=#REF!", "B4": '=IF(A1="#REF!",1,0)', "B5": '="#N/A"'}))
    issue = _issue(check_workbook(data), "error_in_formula")
    assert issue is not None and issue["count"] == 2
    assert [sample["cell"] for sample in issue["samples"]] == ["B2", "B3"]
    assert issue["samples"][0]["detail"] == "#REF!"


def test_a_deliberate_error_literal_is_not_a_broken_link() -> None:
    # Only #REF! is what Excel leaves behind; #N/A here leaves a chart gap on purpose.
    deliberate = {f"B{row}": f'=IF(A{row}="",#N/A,A{row})' for row in range(1, 6)}
    deliberate["C1"] = "=IFERROR(A1/A2,#DIV/0!)"
    deliberate["C2"] = "=IF(A2>0,#REF!,A2)"
    result = check_workbook(_bytes(_book(deliberate)))
    assert _flagged(result, "error_in_formula") == [("C2", "#REF!")]


def test_a_percentage_is_judged_by_its_value() -> None:
    rates = {"B1": "=A1*10%", "B2": "=A1*1%", "B3": "=A1*2%", "B4": "=A1*12 %", "B5": "=A1*13%",
             "C1": "=A1*100%", "C2": "=A1*1000%", "C3": "=A1*0%", "C4": "=A1*12", "C5": "=A1*(1+10)"}
    result = check_workbook(_bytes(_book(rates)))
    assert _issue(result, "hardcoded_number")["count"] == 5
    assert _flagged(result, "hardcoded_number") == [("B1", "10%"), ("B2", "1%"), ("B3", "2%"), ("B4", "12%"), ("B5", "13%")]


def test_volatile_functions_ignore_the_xlfn_prefix() -> None:
    data = _bytes(_book({"B2": "=OFFSET(A1,0,0)+TODAY()", "B3": '=INDIRECT("A1")', "B4": "=_xlfn.RANDBETWEEN(1,6)",
                         "B5": "=SUM(A1:A9)"}))
    issue = _issue(check_workbook(data), "volatile_function")
    assert issue is not None and issue["severity"] == "info" and issue["count"] == 3
    assert [sample["detail"] for sample in issue["samples"]] == ["OFFSET, TODAY", "INDIRECT", "RANDBETWEEN"]


def test_external_references_keep_the_book_name_but_never_its_directory() -> None:
    data = _bytes(_book({
        "B2": "=[1]Sheet1!A1",
        "B3": "='C:\\Users\\deng\\Desktop\\[预算.xlsx]Sheet 1'!A1*2",
        "B4": "=[2]!Rate",
        "B5": "=Sheet1!A1+Table1[Amount]",
    }))
    result = check_workbook(data)
    issue = _issue(result, "external_reference")
    assert issue is not None and issue["count"] == 3
    assert [sample["cell"] for sample in issue["samples"]] == ["B2", "B3", "B4"]
    assert issue["samples"][1]["formula"] == "='[预算.xlsx]Sheet 1'!A1*2"
    assert "Users" not in json.dumps(result, ensure_ascii=False)


def test_sample_formulas_never_carry_a_path() -> None:
    data = _bytes(_book({
        "B1": '=HYPERLINK("/Volumes/财务共享/2026/应付明细.xlsx","底稿")&TODAY()',
        "B2": '=HYPERLINK("D:/财务/2026/预算.xlsx","预算")*1.13',
        "B3": '=HYPERLINK("~/Documents/预算.xlsx","预算")*1.13',
        "B4": '=HYPERLINK("\\\\fs01\\财务\\预算.xlsx","预算")*1.13',
        "B5": '=HYPERLINK("财务共享/2026/凭证-0930.pdf","凭证")*1.13',
        "B6": "='D:/财务/2026/[预算.xlsx]Sheet1'!A1*1.13",
        # Any text with a slash is held back, even a harmless date format.
        "B7": '=TEXT(TODAY(),"yyyy/mm")&"/"&A8',
    }))
    result = check_workbook(data)
    shown = {sample["cell"]: sample.get("formula") for issue in result["issues"] for sample in issue["samples"]}
    assert shown["B1"] == '=HYPERLINK("…","底稿")&TODAY()'
    assert shown["B2"] == '=HYPERLINK("…","预算")*1.13'
    assert shown["B5"] == '=HYPERLINK("…","凭证")*1.13'
    assert shown["B6"] == "='[预算.xlsx]Sheet1'!A1*1.13"
    assert shown["B7"] == '=TEXT(TODAY(),"…")&"…"&A8'
    text = json.dumps(result, ensure_ascii=False)
    for fragment in ("Volumes", "D:/", "~/", "fs01", "财务共享", "Documents"):
        assert fragment not in text
    for pattern in ("/Volumes/x", "D:/x", "d:\\x", "~/x", "/private/var/x", "/mnt/c/x", "smb://fs01/x"):
        assert workbook_health._PATH_LIKE.search(pattern), pattern


def test_external_link_parts_report_count_and_file_names_only() -> None:
    link_rels = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLinkPath"'
        ' Target="file:///C:\\Users\\deng\\Desktop\\2026%20预算.xlsx" TargetMode="External"/></Relationships>'
    ).encode()
    link = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<externalLink xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<externalBook xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rId1"/>'
        "</externalLink>"
    ).encode()

    def with_link(xml: str) -> str:
        return xml.replace(
            "</Relationships>",
            '<Relationship Id="rIdLink" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLink"'
            ' Target="externalLinks/externalLink1.xml"/></Relationships>',
        )

    data = _rewrite(
        _bytes(_book({})),
        {"xl/_rels/workbook.xml.rels": with_link},
        {"xl/externalLinks/externalLink1.xml": link, "xl/externalLinks/_rels/externalLink1.xml.rels": link_rels},
    )
    result = check_workbook(data)
    issue = _issue(result, "external_links")
    assert issue == {"code": "external_links", "severity": "warning", "count": 1, "sheets": [], "samples": [],
                     "targets": ["2026 预算.xlsx"]}
    text = json.dumps(result, ensure_ascii=False)
    assert "Users" not in text and "C:" not in text and "file:" not in text
    assert _issue(check_workbook(_bytes(_book({}))), "external_links") is None


def test_iterative_calculation_and_hidden_sheets() -> None:
    book = _book({})
    book.calculation.iterate = True
    book.create_sheet("参数").sheet_state = "hidden"
    book.create_sheet("底稿").sheet_state = "veryHidden"
    book.create_sheet("说明")
    result = check_workbook(_bytes(book))
    assert _issue(result, "iterative_calculation") == {
        "code": "iterative_calculation", "severity": "warning", "count": 1, "sheets": [], "samples": [],
    }
    hidden = _issue(result, "hidden_sheets")
    assert hidden["severity"] == "info" and hidden["count"] == 2
    assert hidden["sheets"] == ["参数", "底稿"]
    assert hidden["samples"] == [{"sheet": "参数", "detail": "hidden"}, {"sheet": "底稿", "detail": "veryHidden"}]
    assert result["sheet_count"] == 4
    plain = check_workbook(_bytes(_book({})))
    assert _issue(plain, "iterative_calculation") is None and _issue(plain, "hidden_sheets") is None


def test_cached_values_follow_the_catalog_rule() -> None:
    data = _rewrite(_bytes(_book({"B2": "=A2*2", "B3": "=A3*2", "B4": "=A4*2", "B5": "=A5*2", "B6": "=A6*2", "B7": "x"})), {
        SHEET1: _cells({
            "B2": '<c r="B2"><f>A2*2</f><v>4</v></c>',
            "B3": '<c r="B3" t="str"><f>IF(A3=1,"",A3)</f><v></v></c>',
            "B4": '<c r="B4" t="e"><f>A4/0</f><v>#DIV/0!</v></c>',
            "B7": '<c r="B7" t="e"><v>#N/A</v></c>',
        }),
    })
    result = check_workbook(data)
    uncached = _issue(result, "formula_without_cache")
    assert uncached["severity"] == "info"
    assert [sample["cell"] for sample in uncached["samples"]] == ["B5", "B6"]
    errors = _issue(result, "cached_error")
    assert errors["count"] == 2
    assert errors["samples"] == [
        {"sheet": "汇总", "cell": "B4", "range": "A1:B9", "formula": "=A4/0", "detail": "#DIV/0!"},
        {"sheet": "汇总", "cell": "B7", "range": "A3:B10", "detail": "#N/A"},
    ]


def test_inconsistent_row_formula_inside_a_copied_run() -> None:
    row = {f"{column}10": f"=SUM({column}2:{column}9)" for column in "BCDEFG"}
    row["E10"] = "=SUM(E2:E8)"
    # A ratio at the end of the run is where it belongs, not a broken copy.
    row["H10"] = "=G10/B10"
    issue = _issue(check_workbook(_bytes(_book(row))), "inconsistent_row_formula")
    assert issue["count"] == 1
    assert issue["samples"] == [{"sheet": "汇总", "cell": "E10", "range": "B6:H10", "formula": "=SUM(E2:E8)", "detail": "D10"}]


def test_inconsistent_column_formula_keeps_absolute_parts_and_spares_the_total() -> None:
    column = {f"D{row}": f"=B{row}*C{row}*$A$1" for row in range(2, 10)}
    column["D6"] = "=B6*C5*$A$1"
    column["D10"] = "=SUM(D2:D9)"
    result = check_workbook(_bytes(_book(column)))
    issue = _issue(result, "inconsistent_column_formula")
    assert [(sample["cell"], sample["detail"]) for sample in issue["samples"]] == [("D6", "D5")]
    assert _issue(result, "inconsistent_row_formula") is None


def test_the_cell_right_after_a_copied_run_is_checked_too() -> None:
    from openpyxl.utils import get_column_letter as letter

    # September was added to the right of a copied run and still reads August.
    rolled = {f"{letter(column)}10": f"={letter(column)}8-{letter(column)}9" for column in range(2, 13)}
    rolled["M10"] = "=L8-L9"
    rolled["N10"] = "=SUM(B10:M10)"
    assert _flagged(check_workbook(_bytes(_book(rolled))), "inconsistent_row_formula") == [("M10", "L10")]
    # Once September reads its own month, the year total after it reads the run and is not a broken copy.
    rolled["M10"] = "=M8-M9"
    assert _issue(check_workbook(_bytes(_book(rolled))), "inconsistent_row_formula") is None
    # A row appended at the bottom of a copied column that still reads the row above.
    appended = {f"D{row}": f"=B{row}*C{row}" for row in range(2, 10)}
    appended["D10"] = "=B9*C9"
    assert _flagged(check_workbook(_bytes(_book(appended))), "inconsistent_column_formula") == [("D10", "D9")]
    # An opening balance before the run is where it belongs.
    opening = {"B10": "=期初!B10", **{f"{letter(column)}10": f"={letter(column - 1)}10+{letter(column)}9" for column in range(3, 9)}}
    assert _issue(check_workbook(_bytes(_book(opening))), "inconsistent_row_formula") is None


def test_subtotals_of_the_run_are_spared_but_not_in_a_running_balance() -> None:
    from openpyxl.utils import get_column_letter as letter

    # 1-3月, 一季度, 4-6月, 二季度 ...: each quarter adds up the months before it on its own row.
    quarters = {}
    for row in range(2, 8):
        for quarter in (5, 9, 13, 17):
            for column in range(quarter - 3, quarter):
                quarters[f"{letter(column)}{row}"] = f"=明细!{letter(column)}{row}"
            quarters[f"{letter(quarter)}{row}"] = f"=SUM({letter(quarter - 3)}{row}:{letter(quarter - 1)}{row})"
    quarters["G3"] = "=明细!F3"
    result = check_workbook(_bytes(_book(quarters)))
    assert _flagged(result, "inconsistent_row_formula") == [("G3", "F3")]
    # A running balance reads the row above, so a balance that skips a row is still a broken copy.
    balance = {f"D{row}": f"=D{row - 1}+B{row}-C{row}" for row in range(3, 11)}
    balance["D7"] = "=D5+B7-C7"
    assert _flagged(check_workbook(_bytes(_book(balance))), "inconsistent_column_formula") == [("D7", "D6")]


def test_whole_column_and_row_references_are_compared_relative_to_their_cell() -> None:
    from openpyxl.utils import get_column_letter as letter

    # SUMIF over whole columns, copied across the months; E6 still reads D:D.
    lookups = {f"{letter(column)}{row}": f"=SUMIF(明细!$A:$A,$A{row},明细!{letter(column)}:{letter(column)})"
               for row in range(5, 9) for column in range(3, 9)}
    lookups["E6"] = "=SUMIF(明细!$A:$A,$A6,明细!D:D)"
    result = check_workbook(_bytes(_book(lookups)))
    assert _flagged(result, "inconsistent_row_formula") == [("E6", "D6")]
    assert _flagged(result, "inconsistent_column_formula") == [("E6", "E5")]
    totals = {f"{letter(column)}10": f"=SUM({letter(column)}:{letter(column)})" for column in range(2, 8)}
    totals["D10"] = "=SUM(C:C)"
    assert _flagged(check_workbook(_bytes(_book(totals))), "inconsistent_row_formula") == [("D10", "C10")]
    rows = {f"F{row}": f"=SUM({row - 1}:{row - 1})" for row in range(2, 10)}
    rows["F6"] = "=SUM(4:4)"
    assert _flagged(check_workbook(_bytes(_book(rows))), "inconsistent_column_formula") == [("F6", "F5")]

    # A correct copy stays quiet, whatever another sheet happens to hold first.
    book = _book({}, title="明细")
    book.active["C10"] = "=SUM(C:C)"
    summary = book.create_sheet("汇总")
    for column in range(3, 9):
        summary[f"{letter(column)}10"] = f"=SUM({letter(column - 1)}:{letter(column - 1)})"
    assert _issue(check_workbook(_bytes(book)), "inconsistent_row_formula") is None


def test_a_shared_shape_always_means_a_shared_signature() -> None:
    cases = [
        ("SUM(C:C)", 10, 4), ("SUM(C:C)", 10, 5), ("SUM($C:$C)", 10, 6), ("SUM($C:$C)", 10, 7), ("$A:B", 1, 3),
        ("SUM(4:4)", 5, 2), ("SUM(4:4)", 6, 2), ("SUM($4:$4)", 7, 2), ("1:$3", 2, 2), ("2:$3", 3, 2),
        ("SUMIF(明细!$A:$A,$A5,明细!C:C)", 5, 3), ("SUMIF(明细!$A:$A,$A6,明细!C:C)", 6, 3),
        ("[1]Sheet1!A1", 1, 2), ("[1]Sheet1!A2", 2, 2), ("[1]Sheet1!A2", 3, 2), ("[2]!Rate*B5", 5, 3),
        ("'C:\\d\\[b.xlsx]S 1'!B5*2", 5, 3), ("'C:\\d\\[b.xlsx]S 1'!B6*2", 6, 3),
        ("B5*C5", 5, 4), ("B6*C6", 6, 4), ("A$1:$B2", 3, 3), ("A$1:$B3", 4, 3), ("a1+b2", 3, 3), ("A2+B3", 4, 3),
        ("Table1[Amount]*B5", 5, 3), ("Table1[Amount]*B6", 6, 3), ("'Q1 2026'!B5+$A$1", 5, 3), ("'Q1 2026'!B6+$A$1", 6, 3),
        ("SUM(A1:INDEX(B:B,5))", 5, 3), ("SUM(A2:INDEX(B:B,5))", 6, 3), ("Sheet1!#REF!+A5", 5, 2), ("Sheet1!#REF!+A6", 6, 2),
    ]
    cache = workbook_health._FormulaCache()
    for text, row, column in cases:
        formula, signature, (origin_row, origin_column) = cache.at(text, row, column)
        rendered = {number: key for key, number in cache.signatures.items()}[signature]
        template = workbook_health._analyse(text).template
        expected = tuple(piece if isinstance(piece, str) else workbook_health._render_reference(piece, row, column)
                         for piece in template)
        assert rendered == expected, (text, row, column)
        # The shared analysis, moved from the cell it was read at, is this cell's own reading.
        shift = {"R": row - origin_row, "C": column - origin_column}
        moved = tuple(
            piece if isinstance(piece, str) else tuple(
                part if isinstance(part, str) else (part[0], part[1], part[2] if part[1] else part[2] + shift[part[0]])
                for part in piece
            )
            for piece in formula.template
        )
        assert moved == template, (text, row, column)


def test_no_pattern_without_a_dominant_signature() -> None:
    # Two of three agree: below the three-cell floor. Copies with a moving absolute anchor differ by design.
    data = _bytes(_book({"B10": "=SUM(B2:B9)", "C10": "=SUM(C2:C9)", "D10": "=SUM(D2:D8)",
                         "B12": "=B2*$A$1", "C12": "=C2*$A$2", "D12": "=D2*$A$3", "E12": "=E2*$A$4"}))
    result = check_workbook(data)
    assert _issue(result, "inconsistent_row_formula") is None


def test_copies_share_a_shape_and_text_never_does() -> None:
    shape = workbook_health._shape
    assert shape("B5*C5", 5, 4) == shape("B6*C6", 6, 4) == "R[0]C[-2]*R[0]C[-1]"
    assert shape("Sheet2!B5+$A$1", 5, 3) == shape("Sheet2!C5+$A$1", 5, 4)
    # Quoted text, quoted sheet names, function names and exponents are not references.
    assert shape('IF(A1="B5",LOG10(A1),1E5)', 9, 9) == 'IF(R[-8]C[-8]="B5",LOG10(R[-8]C[-8]),1E5)'
    assert shape("'Q1 2026'!B5", 1, 1) == "'Q1 2026'!R[4]C[1]"
    # The same shape at another cell is the same analysis; a different literal is a different shape.
    assert shape("B5*1.13", 5, 3) != shape("B6*1.15", 6, 3)
    data = _bytes(_book({f"C{row}": f"=Sheet1!A{row}*2" for row in range(2, 10)} | {"C6": "=Sheet1!A6*3"}))
    result = check_workbook(data)
    assert [sample["cell"] for sample in _issue(result, "inconsistent_column_formula")["samples"]] == ["C6"]
    assert _issue(result, "hardcoded_number")["samples"] == [
        {"sheet": "汇总", "cell": "C6", "range": "A2:C10", "formula": "=Sheet1!A6*3", "detail": "3"},
    ]


def test_shared_formulas_count_every_copy_and_show_the_copy_text() -> None:
    data = _rewrite(_bytes(_book({f"C{row}": "x" for row in range(2, 10)})), {
        SHEET1: _cells({
            "C2": '<c r="C2"><f t="shared" ref="C2:C9" si="0">A2*1.13</f><v>1.13</v></c>',
            **{f"C{row}": f'<c r="C{row}"><f t="shared" si="0"/><v>{row}</v></c>' for row in range(3, 10) if row != 6},
            "C6": '<c r="C6"><f>A6*1.15</f><v>6.9</v></c>',
        }),
    })
    result = check_workbook(data)
    assert result["checked_formula_cells"] == 8
    numbers = _issue(result, "hardcoded_number")
    assert numbers["count"] == 8
    assert numbers["samples"][1]["formula"] == "=A3*1.13"
    assert _issue(result, "formula_without_cache") is None
    broken = _issue(result, "inconsistent_column_formula")
    assert [(sample["cell"], sample["formula"]) for sample in broken["samples"]] == [("C6", "=A6*1.15")]


def test_an_oversized_sheet_is_skipped_and_reported_not_fatal() -> None:
    book = _book({"B2": "=A2*1.13"})
    wide = book.create_sheet("宽表")
    wide.cell(1, analysis_data.MAX_COLUMNS + 1, 1)
    wide["A2"] = "=A1*1.5"
    result = check_workbook(_bytes(book))
    assert result["status"] == "checked"
    too_large = _issue(result, "sheet_too_large")
    assert too_large == {"code": "sheet_too_large", "severity": "info", "count": 1, "sheets": ["宽表"],
                         "samples": [{"sheet": "宽表"}]}
    assert _issue(result, "hardcoded_number")["sheets"] == ["汇总"]
    assert result["checked_formula_cells"] == 1


def test_scan_stops_at_the_formula_cap_and_says_so(monkeypatch) -> None:
    monkeypatch.setattr(workbook_health, "MAX_FORMULA_CELLS", 5)
    book = _book({f"B{row}": f"=A{row}*2" for row in range(1, 9)})
    book.create_sheet("后面").append(["=1*7"])
    result = check_workbook(_bytes(book))
    assert result["truncated"] is True
    assert result["checked_formula_cells"] == 5
    assert _issue(result, "formula_without_cache")["count"] == 5
    monkeypatch.setattr(workbook_health, "MAX_FORMULA_CELLS", 9)
    assert check_workbook(_bytes(book))["truncated"] is False


def test_distinct_formulas_stop_at_a_work_budget_and_copies_do_not_count(monkeypatch) -> None:
    # openpyxl stores "=1*2" as the text 1*2.
    monkeypatch.setattr(workbook_health, "MAX_ANALYSED_CHARACTERS", 3 * (len("1*2") + workbook_health.ANALYSIS_OVERHEAD))
    # Hand-typed amounts: every formula is read afresh.
    typed = check_workbook(_bytes(_book({f"B{row}": f"={row}*2" for row in range(1, 9)})))
    assert typed["truncated"] is True
    assert typed["checked_formula_cells"] == 3
    # One formula copied down: read once, however many copies.
    copied = check_workbook(_bytes(_book({f"B{row}": f"=A{row}*2" for row in range(1, 11)})))
    assert copied["truncated"] is False
    assert copied["checked_formula_cells"] == 10


def test_unreadable_input_is_unavailable_and_never_raises(monkeypatch) -> None:
    assert check_workbook(b"not a workbook") == {"schema_version": 1, "status": "unavailable", "reason": "unreadable_workbook"}
    empty = io.BytesIO()
    with zipfile.ZipFile(empty, "w") as archive:
        archive.writestr("readme.txt", "hello")
    assert check_workbook(empty.getvalue())["reason"] == "unreadable_workbook"
    data = _bytes(_book({"B2": "=A2*2"}))
    assert check_workbook(_rewrite(data, {SHEET1: lambda xml: xml[: len(xml) // 2]}))["status"] == "unavailable"
    monkeypatch.setattr(analysis_data, "MAX_SOURCE_BYTES", 16)
    assert check_workbook(data)["reason"] == "workbook_too_large"
    monkeypatch.undo()

    def broken(*_args, **_kwargs):
        raise RuntimeError("unexpected")

    monkeypatch.setattr(workbook_health, "_scan_sheet", broken)
    assert check_workbook(data) == {"schema_version": 1, "status": "unavailable", "reason": "check_failed"}


def test_result_is_deterministic_and_holds_no_floats() -> None:
    book = _book({"B2": "=A2*1.13", "B3": "=OFFSET(A1,0,0)", "B4": "=[1]S!A1", "B5": "=#REF!"})
    book.create_sheet("隐藏").sheet_state = "hidden"
    data = _bytes(book)
    first = check_workbook(data)
    assert first == check_workbook(data)
    assert _only_plain_values(first)
    assert first["issue_count"] == len(first["issues"])
    order = [code for code, _ in workbook_health.CODES]
    assert _codes(first) == sorted(_codes(first), key=order.index)
    for issue in first["issues"]:
        assert len(issue["samples"]) <= workbook_health.MAX_SAMPLES
        for sample in issue["samples"]:
            assert len(sample.get("formula", "")) <= workbook_health.MAX_FORMULA_TEXT


# The import path.


@pytest.fixture
def work(tmp_path: Path) -> AnalysisStore:
    return AnalysisStore(TaskStore(tmp_path / "var").create(title="公式体检"))


def _multi_sheet(path: Path) -> Path:
    book = Workbook()
    cover = book.active
    cover.title = "00·封面"
    cover["B3"] = "=1+1"
    cover["B4"] = "=B3*1.13"
    data = book.create_sheet("04·利润表")
    data.append(["项目", "本月"])
    data.append(["营业收入", "1280.00"])
    book.save(path)
    return path


def test_add_dataset_stores_health_beside_metadata(work: AnalysisStore, tmp_path: Path) -> None:
    dataset = work.add_dataset(_multi_sheet(tmp_path / "月结.xlsx"), name="月结套表", metadata={"period": "2026-08"})
    health = dataset["workbook_health"]
    assert health["status"] == "checked"
    assert _codes(health) == ["hardcoded_number", "formula_without_cache"]
    assert "workbook_health" not in dataset["metadata"]
    meta = json.loads((work.root / "datasets" / dataset["id"] / "meta.json").read_text(encoding="utf-8"))
    assert meta["workbook_health"] == health
    assert work.read_dataset(dataset["id"])["workbook_health"] == health
    assert work.state()["datasets"][0]["workbook_health"] == health

    again = work.add_dataset(_multi_sheet(tmp_path / "月结.xlsx"), name="月结套表")
    assert again["duplicate_of"] == dataset["id"]
    assert again["workbook_health"] == health

    csv = tmp_path / "rows.csv"
    csv.write_text("amount\n1\n", encoding="utf-8")
    assert "workbook_health" not in work.add_dataset(csv, name="rows")


def test_a_failed_check_still_imports_the_workbook(work: AnalysisStore, tmp_path: Path, monkeypatch) -> None:
    def broken(*_args, **_kwargs):
        raise RuntimeError("unexpected")

    monkeypatch.setattr(workbook_health, "_check", broken)
    dataset = work.add_dataset(_multi_sheet(tmp_path / "月结.xlsx"), name="月结套表")
    assert dataset["workbook_health"] == {"schema_version": 1, "status": "unavailable", "reason": "check_failed"}
    assert work.read_dataset(dataset["id"])["kind"] == "workbook"


def test_single_sheet_uncached_formula_refusal_is_unchanged(work: AnalysisStore, tmp_path: Path) -> None:
    book = Workbook()
    book.active.append(["amount"])
    book.active.append(["=1+2"])
    path = tmp_path / "formula.xlsx"
    book.save(path)
    with pytest.raises(FoundryError, match="formula has no cached value at Sheet!A2"):
        work.add_dataset(path, name="bad")
    assert work.state()["datasets"] == []


def test_a_derived_table_does_not_rerun_the_check(work: AnalysisStore, tmp_path: Path) -> None:
    data = fixture.monthly_pack()
    path = tmp_path / "pack.xlsx"
    path.write_bytes(data)
    shell = work.add_dataset(path, name="月结套表", metadata={"mode": "workbook"})
    assert shell["workbook_health"]["status"] == "checked"
    declaration = next(sheet for sheet in propose_workbook_structure(data)["sheets"] if sheet["sheet"].startswith("02"))
    snapshot = work.add_table_snapshot(shell["id"], declaration, name="02")
    assert "workbook_health" not in snapshot


# On demand, for snapshots filed before the check existed.


def _as_older_snapshot(work: AnalysisStore, dataset_id: str) -> None:
    path = work.root / "datasets" / dataset_id / "meta.json"
    record = json.loads(path.read_text(encoding="utf-8"))
    record.pop("workbook_health")
    record.pop("_integrity_hash")
    record["_integrity_hash"] = _hash(record)
    path.write_text(json.dumps(record, ensure_ascii=False), encoding="utf-8")


def test_on_demand_check_reads_the_verified_archive_and_writes_nothing(work: AnalysisStore, tmp_path: Path) -> None:
    dataset = work.add_dataset(_multi_sheet(tmp_path / "月结.xlsx"), name="月结套表")
    stored = dataset["workbook_health"]
    _as_older_snapshot(work, dataset["id"])
    assert "workbook_health" not in work.read_dataset(dataset["id"])
    before = sorted(str(path) for path in work.root.rglob("*"))
    result = FinanceFiles(work.task).health({"dataset_id": dataset["id"]})
    assert result == {"dataset_id": dataset["id"], "source_hash": dataset["source_hash"], **stored}
    assert sorted(str(path) for path in work.root.rglob("*")) == before
    assert "workbook_health" not in work.read_dataset(dataset["id"])

    with pytest.raises(FoundryError, match="outside this session context"):
        FinanceFiles(work.task, snapshot_ids=[]).health({"dataset_id": dataset["id"]})
    assert FinanceFiles(work.task, snapshot_ids=[dataset["id"]]).health({"dataset_id": dataset["id"]})["status"] == "checked"

    csv = tmp_path / "rows.csv"
    csv.write_text("amount\n1\n", encoding="utf-8")
    rows = work.add_dataset(csv, name="rows")
    with pytest.raises(FoundryError, match="require an XLSX snapshot"):
        FinanceFiles(work.task).health({"dataset_id": rows["id"]})

    archived = work.root / "datasets" / dataset["id"] / "source.xlsx"
    archived.write_bytes(archived.read_bytes() + b"tampered")
    with pytest.raises(FoundryError, match="integrity check failed"):
        FinanceFiles(work.task).health({"dataset_id": dataset["id"]})


def test_cli_round_trip_and_session_ceiling(tmp_path: Path, capsys) -> None:
    root = tmp_path / "var"
    task = TaskStore(root).create(title="公式体检 CLI")
    # Registered before the upload, so the file is outside its context until it is added to it.
    task.record_session("limited", role="analysis")
    source = _multi_sheet(tmp_path / "月结.xlsx")
    operator = ["--var", str(root), "--task", task.id, "--actor", "manager"]
    assert main([*operator, "dataset-add", "--path", str(source), "--name", "月结套表"]) == 0
    added = json.loads(capsys.readouterr().out)
    assert added["workbook_health"]["status"] == "checked"
    spec = json.dumps({"dataset_id": added["id"]})
    assert main([*operator, "workbook-health", "--spec", spec]) == 0
    checked = json.loads(capsys.readouterr().out)
    assert {key: value for key, value in checked.items() if key not in ("dataset_id", "source_hash")} == added["workbook_health"]

    session = ["--var", str(root), "--session", "limited"]
    assert main([*session, "workbook-health", "--spec", spec]) == 3
    assert "outside this session context" in capsys.readouterr().err
    task.update_session_context("limited", snapshot_ids=[added["id"]])
    assert main([*session, "workbook-health", "--spec", spec]) == 0
    assert json.loads(capsys.readouterr().out)["dataset_id"] == added["id"]
    assert main([*operator, "workbook-health", "--spec", json.dumps({"dataset_id": "nope"})]) == 3
    assert "invalid analysis object id" in capsys.readouterr().err


def test_a_cached_error_sample_opens_in_the_range_preview(tmp_path: Path) -> None:
    task = TaskStore(tmp_path / "var").create(title="错误单元格")
    data = _rewrite(_bytes(_book({"B4": "=A4/0", "B5": 5})), {
        SHEET1: _cells({"B4": '<c r="B4" t="e"><f>A4/0</f><v>#DIV/0!</v></c>'}),
    })
    path = tmp_path / "errors.xlsx"
    path.write_bytes(data)
    dataset = AnalysisStore(task).add_dataset(path, name="errors", metadata={"mode": "workbook"})
    sample = _issue(dataset["workbook_health"], "cached_error")["samples"][0]
    preview = FinanceFiles(task).range_preview({"dataset_id": dataset["id"], "sheet": sample["sheet"], "range": sample["range"]})
    cells = {cell["address"]: cell for cell in preview["cells"]}
    assert cells["B4"]["formula"] == "=A4/0"
    assert cells["B5"]["exact"] == "5"
