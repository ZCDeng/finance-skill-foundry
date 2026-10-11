from __future__ import annotations

import io
import json
import struct
import zipfile
import zlib
from pathlib import Path
from xml.etree import ElementTree

import pytest
from openpyxl import load_workbook

from foundry import dashboard_snapshot
from foundry.analysis import AnalysisStore
from foundry.dashboard_snapshot import (
    _check_figure,
    _reads,
    normalize,
    ranked,
    save,
    snapshot_path,
    snapshots,
)
from foundry.errors import FoundryError
from foundry.finance_service import FinanceStore
from foundry.task import TaskStore

KPI = [
    {"指标": "营业收入", "本年累计": "1000"},
    {"指标": "营业成本", "本年累计": "813"},
    {"指标": "期间费用", "本年累计": "70"},
    {"指标": "本年账单收缴率", "本年累计": "87.32%"},
]
CITY = [
    {"城市": "海口", "收入": "400", "成本": "300", "费用": "30", "收缴率": "92%"},
    {"城市": "三亚", "收入": "300", "成本": "270", "费用": "25", "收缴率": "80%"},
    {"城市": "儋州", "收入": "200", "成本": "160", "费用": "20", "收缴率": "86%"},
]


def _cell(name: str) -> dict:
    return {"alias": "kpi", "match": {"指标": name}, "column": "本年累计"}


def _figure(run: dict, key: str, state: str) -> dict:
    return {"kind": "figure", "title": key, "state": state, "run_id": run["id"],
            "select": {"match": {"entity": "集团", "metric_key": key, "scenario": "本期"}, "column": "value"}}


@pytest.fixture
def board(tmp_path: Path):
    """A manager's board on one collection-health run: two accepted figures and a ranking, one figure waiting, a 情景."""
    task = TaskStore(tmp_path / "var").create(title="华宸物业 8 月")
    manager = FinanceStore(task, actor="manager")
    bindings = {}
    for alias, rows in {"kpi": KPI, "city": CITY}.items():
        path = tmp_path / f"{alias}.json"
        path.write_text(json.dumps(rows, ensure_ascii=False), encoding="utf-8")
        bindings[alias] = AnalysisStore(task).add_dataset(path, name=alias)["id"]
    context = manager.save_context({"bindings": bindings, "scopes": {alias: {"currency": "CNY", "unit": "万元"} for alias in bindings}})
    scenario = manager.save_scenario({
        "context_id": context["id"], "kind": "collection_health", "label": "回款健康",
        "inputs": {"entities": [{"name": "集团", "values": {
            "revenue": _cell("营业收入"), "cost": _cell("营业成本"), "period_expense": _cell("期间费用"),
            "collection_rate_current": _cell("本年账单收缴率")}}],
            "entity_tables": [{"alias": "city", "name_column": "城市", "columns": {
                "revenue": "收入", "cost": "成本", "period_expense": "费用", "collection_rate_current": "收缴率"}}]},
    })
    run = manager.run_scenario({"scenario_id": scenario["id"]})
    cards = [
        _figure(run, "collection_gap", "accepted"),
        _figure(run, "break_even_collection", "accepted"),
        _figure(run, "gross_margin", "pinned"),
        {"kind": "rank", "title": "距生死线（从低到高）", "state": "accepted", "run_id": run["id"],
         "select": {"match": {"metric_key": "collection_gap", "scenario": "本期", "entity_origin": "table"},
                    "label": "entity", "column": "value", "order": "asc", "limit": 10, "reference": {"value": "0", "label": "生死线"}}},
        {"kind": "scenario", "title": "回款情景", "state": "pinned", "run_id": run["id"],
         "select": {"match": {"entity": "集团", "scenario": "本期"}}},
    ]
    saved = manager.dashboard_save({"parent_id": None, "cards": cards})
    return task, manager, run, saved


def _model(manager: FinanceStore, board: dict, run: dict, *, page: int | None = None) -> dict:
    """What the desktop sends: the accepted cards as the board shows them, the first ``page`` of them on the page."""
    taken = [card for card in board["cards"] if card["state"] == "accepted" and card["kind"] != "scenario"]
    on_page, beyond = (taken, []) if page is None else (taken[:page], taken[page:])
    rows = run["result"]["rows"]
    figures, charts, sources = [], [], []
    for number, card in enumerate(on_page, start=1):
        sources.append({"n": number, "card_id": card["id"], "title": card["title"], "run": "A", "lines": ["公式：示例"]})
        if card["kind"] == "figure":
            match = card["select"]["match"]
            row = next(row for row in rows if all(str(row.get(key)) == value for key, value in match.items()))
            figures.append({"card_id": card["id"], "title": card["title"], "subject": "集团", "part": "本期",
                            "value": {"text": row["display_value"], "negative": row["display_value"].startswith("-")},
                            "level": row.get("level"), "level_label": row.get("level_label") or "", "beside": [], "note": "",
                            "source": number, "cell": {"column": "value", "value": str(row["value"])}})
        else:
            charts.append({"card_id": card["id"], "kind": "rank", "title": card["title"], "unit": "个百分点", "note": "", "source": number,
                           "items": [{"label": row["entity"], "value": value, "shown": {"text": row["display_value"], "negative": value < 0}, "note": ""}
                                     for row, value in ranked(card, run)],
                           "reference": {"value": 0.0, "label": "生死线"}})
    pinned = [card for card in board["cards"] if card["state"] == "pinned" and card["kind"] != "scenario"]
    scenario = [card for card in board["cards"] if card["kind"] == "scenario"]
    named = lambda cards: [{"card_id": card["id"], "title": card["title"]} for card in cards]  # noqa: E731
    return {"version": 1, "board_id": board["id"], "task": "华宸物业 8 月", "period": "2026年8月", "figures": figures, "charts": charts,
            "sources": sources, "runs": [{"key": "A", "run_id": run["id"], "label": "回款健康", "computed": "10-10 14:20 计算", "files": ["kpi"]}],
            "beyond": named(beyond), "left_out": {"pinned": named(pinned), "scenario": named(scenario)}}


def _png(width: int = 4, height: int = 3) -> bytes:
    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))
    raw = b"".join(b"\x00" + b"\xff\xff\xff" * width for _ in range(height))
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


def test_a_workbook_snapshot_holds_every_accepted_card_as_its_run_holds_it(board):
    task, manager, run, saved = board
    filed = save(task, {"format": "xlsx", "model": _model(manager, saved, run)}, actor="manager")
    record = filed["snapshot"]
    assert [card["title"] for card in record["cards"]] == ["collection_gap", "break_even_collection", "距生死线（从低到高）"]
    assert record["left_out"] == {"pinned": ["gross_margin"], "scenario": ["回款情景"]} and record["beyond"] == []
    assert record["filename"] == "经营看板_华宸物业8月_2026年8月.xlsx"

    book = load_workbook(io.BytesIO(Path(filed["path"]).read_bytes()))
    sheet = [list(row) for row in book["看板快照"].iter_rows(values_only=True)]
    gap = next(row for row in run["result"]["rows"] if row["entity"] == "集团" and row["metric_key"] == "collection_gap" and row["scenario"] == "本期")
    shown = next(row for row in sheet if row[0] == "collection_gap")
    assert shown[6] == float(gap["value"]) and shown[3] == gap["display_value"]
    assert any(row[0] and str(row[0]).startswith("没有放进来：待认可 1 张（gross_margin）；情景 1 张") for row in sheet)
    places = [row[1] for row in sheet if isinstance(row[0], int)]
    assert places == [str(row["entity"]) for row, _ in ranked(saved["cards"][3], run)]
    sources = [list(row) for row in book["来源"].iter_rows(values_only=True)]
    assert sources[1][4] == run["id"]

    assert snapshot_path(task, record["id"]) == Path(filed["path"])
    assert [row["id"] for row in snapshots(task)] == [record["id"]]


def test_a_workbook_snapshot_says_who_exported_it_when_a_name_was_given(board):
    task, manager, run, saved = board
    signed = save(task, {"format": "xlsx", "model": {**_model(manager, saved, run), "signed_by": "财务部 张三"}}, actor="manager")
    sheet = [row[0] for row in load_workbook(io.BytesIO(Path(signed["path"]).read_bytes()))["看板快照"].iter_rows(values_only=True)]
    assert "导出人：财务部 张三" in sheet
    unsigned = normalize(_model(manager, saved, run))
    assert "signed_by" not in unsigned
    with pytest.raises(FoundryError, match="signed by must be text of at most 80"):
        normalize({**_model(manager, saved, run), "signed_by": "张" * 81})


def test_a_page_names_the_cards_left_off_it_and_the_files_come_from_the_desktop(board, tmp_path):
    task, manager, run, saved = board
    model = _model(manager, saved, run, page=2)
    pdf = tmp_path / "page.pdf"
    pdf.write_bytes(b"%PDF-1.7\n%%EOF\n")
    filed = save(task, {"format": "pdf", "model": model, "file": str(pdf)}, actor="manager")
    assert Path(filed["path"]).read_bytes() == pdf.read_bytes() and filed["snapshot"]["beyond"] == ["距生死线（从低到高）"]
    with pytest.raises(FoundryError, match="accepted cards and no others"):
        save(task, {"format": "xlsx", "model": model}, actor="manager")
    with pytest.raises(FoundryError, match="not a PNG"):
        save(task, {"format": "pptx", "model": model, "file": str(pdf)}, actor="manager")

    png = tmp_path / "page.png"
    png.write_bytes(_png(1123, 794))
    deck = Path(save(task, {"format": "pptx", "model": _model(manager, saved, run), "file": str(png)}, actor="manager")["path"])
    with zipfile.ZipFile(deck) as archive:
        for name in archive.namelist():
            if name.endswith((".xml", ".rels")):
                ElementTree.fromstring(archive.read(name))
        assert archive.read("ppt/media/image1.png") == png.read_bytes()
        slide = archive.read("ppt/slides/slide1.xml").decode("utf-8")
        presentation = archive.read("ppt/presentation.xml").decode("utf-8")
    assert '<p:sldSz cx="10692000" cy="7560000"/>' in presentation
    gap = next(item for item in _model(manager, saved, run)["figures"] if item["title"] == "collection_gap")
    assert f"collection_gap：{gap['value']['text']}" in slide and "海口" in slide


def test_a_snapshot_that_is_not_what_the_board_shows_is_refused(board):
    task, manager, run, saved = board
    honest = _model(manager, saved, run)

    def refused(change, message, *, actor="manager"):
        model = json.loads(json.dumps(honest))
        change(model)
        with pytest.raises(FoundryError, match=message):
            save(task, {"format": "xlsx", "model": model}, actor=actor)

    refused(lambda model: model["figures"][0]["cell"].update(value="1.5"), "is not the figure its run holds")
    # The run's number sent beside other words: what the page would print is held to the run as well.
    refused(lambda model: model["figures"][0]["value"].update(text="-0.89 个百分点"), "does not print the figure its run holds")
    refused(lambda model: model["figures"][0].update(level="green", level_label="绿"), "does not print the figure its run holds")
    refused(lambda model: model["figures"][0]["beside"].append({"label": "上年同期", "text": "+1.20 个百分点"}), "does not print the figure")
    refused(lambda model: model["charts"][0]["items"][0]["shown"].update(text="(3.00)"), "does not print the figures its run holds")
    refused(lambda model: model["charts"][0]["items"].reverse(), "is not the ranking its run holds")
    refused(lambda model: model["charts"][0].update(reference=None), "own reference line")
    refused(lambda model: model["left_out"].update(pinned=[]), "cards it leaves out")
    refused(lambda model: model["sources"][0].update(run="B"), "does not name the run")
    refused(lambda model: model.update(board_id="0" * 32), "changed since")
    refused(lambda model: None, "only the manager", actor="assistant")
    # A card waiting for the manager is not on a snapshot, however it is sent.
    pinned = next(card for card in saved["cards"] if card["state"] == "pinned" and card["kind"] == "figure")
    refused(lambda model: model["figures"].append({**model["figures"][0], "card_id": pinned["id"]}), "accepted cards and no others")

    later = manager.dashboard_save({"parent_id": saved["id"], "cards": [{**card, "state": "pinned"} if card["kind"] != "scenario" else card
                                                                       for card in saved["cards"]]})
    with pytest.raises(FoundryError, match="no card on the dashboard is accepted"):
        save(task, {"format": "xlsx", "model": {**honest, "board_id": later["id"]}}, actor="manager")


def test_a_page_prints_each_figure_as_the_board_writes_it(board):
    """The board's own forms pass: a bar's figure without its unit, a negative in accounting parentheses."""
    task, manager, run, saved = board
    model = _model(manager, saved, run)
    for item in model["charts"][0]["items"]:
        figure = item["shown"]["text"].split(" ")[0].lstrip("+")
        item["shown"]["text"] = f"({figure[1:]})" if figure.startswith("-") else figure
    assert save(task, {"format": "xlsx", "model": model}, actor="manager")["snapshot"]["cards"]


def test_a_figure_is_read_at_the_places_the_page_shows():
    assert _reads("0.8732", "87.32%") and _reads("87.32", "87.32%") and _reads("-0.0098", "-0.98 个百分点")
    assert _reads("8.33333333333333333333333333333333333", "8.33%") and _reads(1234.5, "1,234.50")
    assert _reads("-18843.06", "(18,843.06)") and _reads("325644", "325,644.00 万元（估算）")
    assert not _reads("1234.5", "1,234.60") and not _reads("0.8732", "87.42%") and not _reads("18843.06", "(18,843.06)")
    # Text that writes no figure states no amount; a cell that is no number is shown as it is.
    assert _reads("2026-08", "2026 年 8 月") and _reads("87.32%", "87.32%") and not _reads("待定", "1.00")


def test_a_figure_prints_the_same_measure_s_other_lines_beside_it_as_the_run_wrote_them():
    line = lambda scenario, value, shown: {"entity": "集团", "metric_key": "collection_rate_current", "scenario": scenario,
                                           "value": value, "display_value": shown, "level": "yellow", "level_label": "黄"}
    rows = [line("本期", "0.8732", "87.32%"), line("上年同期", "0.8811", "88.11%"), line("变动", "-0.79", "-0.79 个百分点")]
    run = {"engine": "finance", "result": {"scenario_kind": "collection_health", "rows": rows}}
    card = {"title": "收缴率", "select": {"match": {"entity": "集团", "metric_key": "collection_rate_current", "scenario": "本期"}, "column": "value"}}
    shown = {"value": {"text": "87.32%"}, "level": "yellow", "level_label": "黄", "cell": {"column": "value", "value": "0.8732"},
             "beside": [{"label": "上年同期", "text": "88.11%"}, {"label": "变动", "text": "-0.79 个百分点"}]}
    _check_figure(card, run, shown)
    shown["beside"][1]["text"] = "+0.79 个百分点"
    with pytest.raises(FoundryError, match="does not print"):
        _check_figure(card, run, shown)


def test_a_board_saved_while_the_snapshot_is_written_is_not_filed(board, monkeypatch):
    task, manager, run, saved = board
    workbook = dashboard_snapshot.workbook

    def saved_meanwhile(*args, **kwargs):
        manager.dashboard_save({"parent_id": saved["id"], "cards": saved["cards"][:-1]})
        return workbook(*args, **kwargs)

    monkeypatch.setattr(dashboard_snapshot, "workbook", saved_meanwhile)
    with pytest.raises(FoundryError, match="changed since"):
        save(task, {"format": "xlsx", "model": _model(manager, saved, run)}, actor="manager")
    assert snapshots(task) == []


def test_the_library_lists_a_snapshot_and_refuses_one_whose_bytes_changed(board):
    from foundry.overview import library

    task, manager, run, saved = board
    filed = save(task, {"format": "xlsx", "model": _model(manager, saved, run)}, actor="manager")
    rows = [row for row in library(TaskStore(task.store.var_root)) if row["type"] == "dashboard"]
    assert len(rows) == 1 and rows[0]["key"] == f"dashboard:{filed['snapshot']['id']}"
    assert rows[0]["snapshot"]["cards"] == ["collection_gap", "break_even_collection", "距生死线（从低到高）"]
    assert rows[0]["snapshot"]["format"] == "xlsx" and rows[0]["snapshot"]["left_out"]["scenario"] == ["回款情景"]
    Path(filed["path"]).write_bytes(b"changed")
    with pytest.raises(FoundryError, match="integrity"):
        snapshot_path(task, filed["snapshot"]["id"])


def test_the_desktop_files_a_snapshot_through_the_cli_and_finds_it_again(board, capsys):
    """Main's call: the CLI reads the model's decimals as Decimal, and no session can export."""
    from foundry.cli import main

    task, manager, run, saved = board
    spec = json.dumps({"format": "xlsx", "model": _model(manager, saved, run)}, ensure_ascii=False)
    var = str(task.store.var_root)
    assert main(["--var", var, "--task", task.id, "--actor", "manager", "dashboard-snapshot", "--spec", spec]) == 0
    filed = json.loads(capsys.readouterr().out)
    assert main(["--var", var, "--task", task.id, "dashboard-snapshot-path", "--snapshot", filed["snapshot"]["id"]]) == 0
    assert json.loads(capsys.readouterr().out)["path"] == filed["path"]
    assert main(["--var", var, "--task", task.id, "--actor", "assistant", "dashboard-snapshot", "--spec", spec]) == 3
    assert "only the manager" in capsys.readouterr().err
