from __future__ import annotations

import json
from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from foundry.analysis import AnalysisStore
from foundry.errors import FoundryError
from foundry.task import TaskStore


@pytest.fixture
def work(tmp_path):
    task = TaskStore(tmp_path / "var").create(title="利润分析")
    return AnalysisStore(task)


def add(work, tmp_path, rows, name="actual", **metadata):
    path = tmp_path / (name + ".json")
    path.write_text(json.dumps(rows, ensure_ascii=False), encoding="utf-8")
    return work.add_dataset(path, name=name, metadata=metadata)


def spec(dataset, sql, **kw):
    return {
        "question": "分析利润",
        "bindings": {"actual": dataset["id"]},
        "sql": sql,
        **kw,
    }


def test_b01_budget_bridge_b02_period_comparison(work, tmp_path):
    actual = add(
        work,
        tmp_path,
        [
            {
                "period": "2026-08",
                "region": "华东",
                "project": "A",
                "revenue": 120,
                "cost": 70,
                "expense": 20,
            },
            {
                "period": "2026-07",
                "region": "华东",
                "project": "A",
                "revenue": 100,
                "cost": 65,
                "expense": 15,
            },
            {
                "period": "2026-08",
                "region": "西南",
                "project": "B",
                "revenue": 50,
                "cost": 30,
                "expense": 10,
            },
        ],
        grain=["period", "region", "project"],
        currency="CNY",
        unit="万元",
    )
    budget = add(
        work,
        tmp_path,
        [
            {
                "period": "2026-08",
                "region": "华东",
                "project": "A",
                "revenue": 100,
                "cost": 60,
                "expense": 15,
            }
        ],
        name="budget",
        grain=["period", "region", "project"],
        currency="CNY",
        unit="万元",
    )
    analysis = work.save_analysis(
        {
            "question": "华东利润为何超预算",
            "bindings": {"actual": actual["id"], "budget": budget["id"]},
            "relationships": [
                {
                    "left": "actual",
                    "right": "budget",
                    "keys": ["period", "region", "project"],
                    "cardinality": "one-to-one",
                }
            ],
            "sql": """SELECT a.revenue-a.cost-a.expense AS actual_profit,b.revenue-b.cost-b.expense AS budget_profit,
(a.revenue-a.cost-a.expense)-(b.revenue-b.cost-b.expense) AS variance,
((a.revenue-a.cost-a.expense)-(b.revenue-b.cost-b.expense))/(b.revenue-b.cost-b.expense)*100 AS variance_pct,
a.revenue-b.revenue AS revenue_effect,b.cost-a.cost AS cost_effect,b.expense-a.expense AS expense_effect
FROM actual a JOIN budget b USING(period,region,project) WHERE a.region='华东' AND a.period='2026-08' """,
        }
    )
    run = work.run_analysis(analysis["id"])
    assert run["execution_status"] == "succeeded", run
    row = run["result"]["rows"][0]
    assert {k: Decimal(v) for k, v in row.items()} == {
        "actual_profit": 30,
        "budget_profit": 25,
        "variance": 5,
        "variance_pct": 20,
        "revenue_effect": 20,
        "cost_effect": -10,
        "expense_effect": -5,
    }
    assert run["relationship_checks"][0]["left_unmatched_rows"] == 2
    following = work.save_analysis(
        spec(
            actual,
            "SELECT SUM(CASE WHEN period='2026-08' THEN revenue-cost-expense ELSE 0 END)-SUM(CASE WHEN period='2026-07' THEN revenue-cost-expense ELSE 0 END) AS change FROM actual WHERE region='华东'",
            parent_id=analysis["id"],
        )
    )
    assert (
        Decimal(work.run_analysis(following["id"])["result"]["rows"][0]["change"]) == 10
    )


def test_north_share_and_deterministic_reexecution(work, tmp_path):
    dataset = add(
        work,
        tmp_path,
        [{"region": "North", "revenue": 70}, {"region": "South", "revenue": 30}],
    )
    a = work.save_analysis(
        spec(
            dataset,
            "SELECT SUM(CASE WHEN region=? THEN revenue ELSE 0 END)/NULLIF(SUM(revenue),0)*100 AS share FROM actual",
            parameters=["North"],
        )
    )
    first, second = work.run_analysis(a["id"]), work.run_analysis(a["id"])
    assert Decimal(first["result"]["rows"][0]["share"]) == 70
    assert first["output_hash"] == second["output_hash"]
    assert second["reproducibility_status"] == "reproducible"
    assert first["review_status"] == "unreviewed"


def test_b03_missing_budget_is_visible_without_blocking_actual(work, tmp_path):
    dataset = add(work, tmp_path, [{"profit": 30}])
    a = work.save_analysis(
        spec(
            dataset,
            "SELECT SUM(profit) AS profit FROM actual",
            missing_inputs=["budget"],
        )
    )
    run = work.run_analysis(a["id"])
    assert run["execution_status"] == "succeeded"
    assert run["missing_inputs"] == ["budget"]
    assert Decimal(run["result"]["rows"][0]["profit"]) == 30


def test_b04_b05_immutable_assumptions_snapshot_and_runs(work, tmp_path):
    dataset = add(work, tmp_path, [{"profit": 30}])
    old = work.save_analysis(spec(dataset, "SELECT profit FROM actual"))
    prior = work.run_analysis(old["id"])
    revised = work.save_analysis(
        spec(
            dataset,
            "SELECT profit+CAST(? AS DECIMAL(38,2)) AS profit FROM actual",
            parameters=["5"],
            assumptions=[
                {"name": "exclude_hq", "value": "5", "basis": "user_assumption"}
            ],
            parent_id=old["id"],
        )
    )
    assert (
        Decimal(work.run_analysis(revised["id"])["result"]["rows"][0]["profit"]) == 35
    )
    newdata = add(work, tmp_path, [{"profit": 40}])
    new = work.save_analysis(
        spec(newdata, "SELECT profit FROM actual", parent_id=revised["id"])
    )
    assert Decimal(work.run_analysis(new["id"])["result"]["rows"][0]["profit"]) == 40
    assert work.get_run(prior["id"])["input_hashes"]["actual"] == dataset["data_hash"]
    assert Decimal(work.get_run(prior["id"])["result"]["rows"][0]["profit"]) == 30


def test_b06_b07_draft_archive_revision_override_and_export(work, tmp_path):
    dataset = add(work, tmp_path, [{"profit": 30}])
    a = work.save_analysis(spec(dataset, "SELECT profit FROM actual"))
    run = work.run_analysis(a["id"])
    original = work.save_draft(
        {
            "title": "经营简报",
            "body": "利润 30。",
            "run_ids": [run["id"]],
            "claims": [
                {
                    "basis": "calculation",
                    "run_id": run["id"],
                    "row": 0,
                    "column": "profit",
                    "value": "30",
                }
            ],
            "publication_status": "archived",
            "review_status": "accepted",
        }
    )
    exported = work.export_draft(original["id"], "md")
    oldbytes = Path(exported["path"]).read_bytes()
    revised = work.save_draft(
        {
            "parent_id": original["id"],
            "body": "利润 35，包含经理调整。",
            "claims": [
                {
                    "basis": "user_assumption",
                    "value": "35",
                    "original_value": "30",
                    "reason": "管理调整",
                }
            ],
        }
    )
    assert revised["run_ids"] == [run["id"]]
    assert revised["reproducibility_status"] == "unverified"
    assert revised["publication_status"] == "draft"
    assert revised["review_status"] == "unreviewed"
    assert Path(exported["path"]).read_bytes() == oldbytes
    assert work.get_draft(original["id"])["publication_status"] == "archived"
    workbook = work.export_draft(revised["id"], "xlsx")
    assert load_workbook(workbook["path"]).active["A1"].value == "经营简报"


def test_forged_claim_and_history_reference_refused(work, tmp_path):
    dataset = add(work, tmp_path, [{"profit": 30}])
    a = work.save_analysis(spec(dataset, "SELECT profit FROM actual"))
    run = work.run_analysis(a["id"])
    with pytest.raises(FoundryError, match="claim"):
        work.save_draft(
            {
                "body": "利润35",
                "run_ids": [run["id"]],
                "claims": [
                    {
                        "basis": "calculation",
                        "run_id": run["id"],
                        "row": 0,
                        "column": "profit",
                        "value": "35",
                    }
                ],
            }
        )
    with pytest.raises(FoundryError):
        work.save_draft({"body": "x", "run_ids": ["0" * 32]})


def test_b08_method_has_no_data_or_original_parameters(work, tmp_path):
    dataset = add(work, tmp_path, [{"region": "SecretClient", "profit": 30}])
    a = work.save_analysis(
        spec(
            dataset,
            "SELECT SUM(profit) AS profit FROM actual WHERE region=?",
            parameters=["SecretClient"],
            assumptions=[{"sensitive": "SecretClient"}],
        )
    )
    method = work.save_method(a["id"], "利润方法")
    assert dataset["id"] not in json.dumps(method)
    assert "SecretClient" not in json.dumps(method)
    nextdata = add(work, tmp_path, [{"region": "NextClient", "profit": 50}])
    applied = work.apply_method(
        method["id"], {"actual": nextdata["id"]}, parameters=["NextClient"]
    )
    assert (
        Decimal(work.run_analysis(applied["id"])["result"]["rows"][0]["profit"]) == 50
    )


def test_b09_background_and_unsupported_format(work, tmp_path):
    note = tmp_path / "note.md"
    note.write_text("本月经营情况\n\n新店开业可能推高费用。", encoding="utf8")
    dataset = work.add_dataset(note, name="经营说明")
    read = work.read_dataset(dataset["id"])
    assert read["kind"] == "background"
    assert read["paragraphs"][1]["text"] == "新店开业可能推高费用。"
    assert read["paragraphs"][1]["source_hash"] == dataset["source_hash"]
    pdf = tmp_path / "x.pdf"
    pdf.write_bytes(b"fake")
    with pytest.raises(FoundryError, match="unsupported"):
        work.add_dataset(pdf, name="pdf")


@pytest.mark.parametrize(
    "sql",
    [
        "SELECT * FROM read_text('/etc/passwd')",
        "SELECT * FROM read_csv('/etc/passwd')",
        "ATTACH '/tmp/evil.db' AS evil",
        "COPY (SELECT 1) TO '/tmp/evil.csv'",
        "INSTALL httpfs",
        "LOAD httpfs",
        "SELECT 1;CREATE TABLE evil(a INTEGER)",
        "SET enable_external_access=true",
        "SELECT * FROM query('SELECT * FROM actual')",
        "SELECT * FROM information_schema.tables",
        "SELECT * FROM duckdb_settings()",
        "SELECT getenv('HOME')",
    ],
)
def test_sql_boundary(work, tmp_path, sql):
    dataset = add(work, tmp_path, [{"profit": 30}])
    a = work.save_analysis(spec(dataset, sql))
    run = work.run_analysis(a["id"])
    assert run["execution_status"] == "failed"
    assert run["reproducibility_status"] == "unverified"
    assert run["output_hash"] is None


def test_exact_json_csv_xlsx_and_export(work, tmp_path):
    path = tmp_path / "amounts.json"
    path.write_text('[{"amount":900719925474099.99}]')
    dataset = work.add_dataset(path, name="large")
    assert work.read_dataset(dataset["id"])["rows"][0]["amount"] == "900719925474099.99"
    a = work.save_analysis(spec(dataset, "SELECT SUM(amount) AS amount FROM actual"))
    run = work.run_analysis(a["id"])
    assert run["result"]["rows"][0]["amount"] == "900719925474099.99"
    d = work.save_draft({"body": "精确金额", "run_ids": [run["id"]]})
    ex = work.export_draft(d["id"], "xlsx")
    book = load_workbook(ex["path"])
    assert book["Result 1"]["A2"].value == "900719925474099.99"
    wb = Workbook()
    wb.active.append(["amount"])
    wb.active.append(["1234567890.1234567"])
    xp = tmp_path / "small.xlsx"
    wb.save(xp)
    xdata = work.add_dataset(xp, name="small")
    assert work.read_dataset(xdata["id"])["rows"][0]["amount"] == "1234567890.1234567"


def test_xlsx_uncached_formula_refused(work, tmp_path):
    wb = Workbook()
    wb.active.append(["amount"])
    wb.active.append(["=1+2"])
    p = tmp_path / "formula.xlsx"
    wb.save(p)
    with pytest.raises(FoundryError, match="formula"):
        work.add_dataset(p, name="bad")


def test_duplicate_relationship_rejected_before_sum(work, tmp_path):
    actual = add(work, tmp_path, [{"id": 1, "revenue": 70}])
    mapping = add(
        work, tmp_path, [{"id": 1, "name": "A"}, {"id": 1, "name": "B"}], name="map"
    )
    a = work.save_analysis(
        {
            "question": "收入",
            "sql": "SELECT SUM(a.revenue) AS revenue FROM actual a JOIN mapping m USING(id)",
            "bindings": {"actual": actual["id"], "mapping": mapping["id"]},
            "relationships": [
                {
                    "left": "actual",
                    "right": "mapping",
                    "keys": ["id"],
                    "cardinality": "many-to-one",
                }
            ],
        }
    )
    run = work.run_analysis(a["id"])
    assert run["execution_status"] == "failed"
    assert "duplicate" in run["error"].lower()


def test_malformed_relationships_refused_when_saved(work, tmp_path):
    actual = add(work, tmp_path, [{"id": 1, "revenue": 70}])
    budget = add(work, tmp_path, [{"id": 1, "revenue": 60}], name="budget")
    base = {
        "question": "收入差异",
        "sql": "SELECT a.revenue - b.revenue AS variance FROM actual a JOIN budget b USING(id)",
        "bindings": {"actual": actual["id"], "budget": budget["id"]},
    }

    with pytest.raises(FoundryError, match="relationship #1 must be an object"):
        work.save_analysis({**base, "relationships": ["actual to budget on id"]})
    with pytest.raises(FoundryError, match="left and right"):
        work.save_analysis(
            {
                **base,
                "relationships": [
                    {
                        "source": "actual",
                        "target": "budget",
                        "on": ["id"],
                        "cardinality": "1:1",
                    }
                ],
            }
        )
    with pytest.raises(FoundryError, match="cardinality"):
        work.save_analysis(
            {
                **base,
                "relationships": [
                    {
                        "left": "actual",
                        "right": "budget",
                        "left_keys": ["id"],
                        "right_keys": ["id"],
                        "cardinality": "1:1",
                    }
                ],
            }
        )


def test_relationship_schema_uses_left_and_right_keys(work, tmp_path):
    actual = add(work, tmp_path, [{"id": 1, "revenue": 70}])
    budget = add(work, tmp_path, [{"id": 1, "target": 60}], name="budget")
    analysis = work.save_analysis(
        {
            "question": "收入差异",
            "sql": "SELECT a.revenue - b.target AS variance FROM actual a JOIN budget b ON a.id = b.id",
            "bindings": {"actual": actual["id"], "budget": budget["id"]},
            "relationships": [
                {
                    "left": "actual",
                    "right": "budget",
                    "left_keys": ["id"],
                    "right_keys": ["id"],
                    "cardinality": "one-to-one",
                }
            ],
        }
    )
    run = work.run_analysis(analysis["id"])
    assert run["execution_status"] == "succeeded"
    assert run["relationship_checks"][0]["left_keys"] == ["id"]
    assert run["relationship_checks"][0]["right_keys"] == ["id"]


def test_cross_task_path_and_tampered_snapshot_refused(work, tmp_path):
    other = AnalysisStore(work.task.store.create(title="Other"))
    data = add(other, tmp_path, [{"profit": 30}])
    with pytest.raises(FoundryError):
        work.read_dataset(data["id"])
    with pytest.raises(FoundryError):
        work.read_dataset("../../x")
    own = add(work, tmp_path, [{"profit": 30}])
    p = work.root / "datasets" / own["id"] / "data.json"
    p.write_text('{"rows":[]}')
    with pytest.raises(FoundryError, match="integrity"):
        work.read_dataset(own["id"])


def test_concurrent_revision_ids_unique(work, tmp_path):
    data = add(work, tmp_path, [{"profit": 30}])
    with ThreadPoolExecutor(max_workers=4) as pool:
        records = list(
            pool.map(
                lambda _: work.save_analysis(spec(data, "SELECT profit FROM actual")),
                range(12),
            )
        )
    assert len({r["id"] for r in records}) == 12
    assert len(work.state()["analyses"]) == 12


def test_result_limit_timeout_and_nonfinite(work, tmp_path):
    data = add(work, tmp_path, [{"profit": 30}])
    limited = work.save_analysis(spec(data, "SELECT * FROM range(100)", max_rows=5))
    r = work.run_analysis(limited["id"])
    assert r["result"]["row_count"] == 5
    assert r["result"]["truncated"]
    huge = work.save_analysis(
        spec(
            data,
            "SELECT SUM(a.range*b.range) FROM range(1000000000) a CROSS JOIN range(1000000000) b",
            timeout_seconds=0.02,
        )
    )
    r = work.run_analysis(huge["id"])
    assert r["execution_status"] == "failed"
    assert r["error_code"] == "timeout"
    nonfinite = work.save_analysis(spec(data, "SELECT 1.0/0.0 AS ratio"))
    r = work.run_analysis(nonfinite["id"])
    assert r["execution_status"] == "failed"
    assert "finite" in r["error"]


def test_relationship_compares_decimal_keys_with_sql_semantics(work, tmp_path):
    a = tmp_path / "a.json"
    a.write_text('[{"id":1,"revenue":70}]')
    actual = work.add_dataset(a, name="actual")
    b = tmp_path / "b.json"
    b.write_text('[{"id":1.0,"label":"A"},{"id":1.00,"label":"B"}]')
    mapping = work.add_dataset(b, name="mapping")
    revision = work.save_analysis(
        {
            "question": "收入",
            "sql": "SELECT SUM(a.revenue) AS revenue FROM actual a JOIN mapping m USING(id)",
            "bindings": {"actual": actual["id"], "mapping": mapping["id"]},
            "relationships": [
                {
                    "left": "actual",
                    "right": "mapping",
                    "keys": ["id"],
                    "cardinality": "many-to-one",
                }
            ],
        }
    )
    run = work.run_analysis(revision["id"])
    assert run["execution_status"] == "failed"
    assert "duplicate" in run["error"]


@pytest.mark.parametrize(
    "condition",
    ["ON TRUE", "ON a.id=m.id OR TRUE", "ON a.id=m.id OR a.id=1", "ON a.id+0=m.id"],
)
def test_relationship_must_match_actual_sql_join(work, tmp_path, condition):
    actual = add(work, tmp_path, [{"id": 1, "revenue": 70}])
    mapping = add(
        work, tmp_path, [{"id": 1, "name": "A"}, {"id": 2, "name": "B"}], name="map"
    )
    revision = work.save_analysis(
        {
            "question": "收入",
            "sql": f"SELECT SUM(a.revenue) AS revenue FROM actual a JOIN mapping m {condition}",
            "bindings": {"actual": actual["id"], "mapping": mapping["id"]},
            "relationships": [
                {
                    "left": "actual",
                    "right": "mapping",
                    "keys": ["id"],
                    "cardinality": "many-to-one",
                }
            ],
        }
    )
    run = work.run_analysis(revision["id"])
    assert run["execution_status"] == "failed"
    assert "JOIN" in run["error"]


def test_xlsx_sparse_cells_and_dates(work, tmp_path):
    from datetime import date

    wb = Workbook()
    ws = wb.active
    ws.append(["project", "amount", "invoice_date"])
    ws.append([None, 100, date(2026, 8, 1)])
    path = tmp_path / "sparse.xlsx"
    wb.save(path)
    dataset = work.add_dataset(path, name="sparse")
    row = work.read_dataset(dataset["id"])["rows"][0]
    assert row["project"] is None
    assert row["invoice_date"].startswith("2026-08-01")


def test_script_run_has_frozen_inputs_and_immutable_output(work, tmp_path):
    dataset = add(work, tmp_path, [{"profit": "123456789012345.67"}])
    code = "import os,json\nfrom pathlib import Path\nfrom decimal import Decimal\ndata=json.loads((Path(os.environ['FOUNDRY_INPUTS'])/'actual.json').read_text())\nvalue=sum(Decimal(r['profit']) for r in data['rows'])\nPath('profit.txt').write_text(str(value))\nprint(value)"
    run = work.run_script(
        {"code": code, "language": "python", "bindings": {"actual": dataset["id"]}},
        session_id="test-session",
        context_revision=7,
    )
    if run.get("runner_status") == "unavailable":
        pytest.skip("OS-isolated runner unavailable on this platform")
    assert run["execution_status"] == "succeeded", run
    assert run["stdout"].strip() == "123456789012345.67"
    assert run["context_revision"] == 7
    assert run["input_hashes"]["actual"] == dataset["data_hash"]
    assert run["outputs"][0]["path"] == "profit.txt"
    assert (
        work.root / run["outputs"][0]["artifact_path"]
    ).read_text() == "123456789012345.67"


def test_method_removes_comments_and_literal_business_values(work, tmp_path):
    data = add(work, tmp_path, [{"revenue": 70, "region": "PrivateClient"}])
    revision = work.save_analysis(
        spec(
            data,
            "SELECT revenue+500000 AS total FROM actual WHERE region='PrivateClient' -- PrivateClient bank account 123456",
        )
    )
    method = work.save_method(revision["id"], "Reusable")
    encoded = json.dumps(method)
    assert (
        "PrivateClient" not in encoded
        and "500000" not in encoded
        and "123456" not in encoded
    )
    applied = work.apply_method(
        method["id"], {"actual": data["id"]}, parameters=[100, "PrivateClient"]
    )
    result = work.run_analysis(applied["id"])
    assert result["execution_status"] == "succeeded", result
    assert Decimal(result["result"]["rows"][0]["total"]) == 170


def test_reexport_detects_corrupted_export(work, tmp_path):
    draft = work.save_draft({"body": "original", "publication_status": "archived"})
    export = work.export_draft(draft["id"], "md")
    Path(export["path"]).write_text("tampered")
    with pytest.raises(FoundryError, match="integrity"):
        work.export_draft(draft["id"], "md")


def test_derived_join_cannot_hide_duplicate_mapping(work, tmp_path):
    actual = add(work, tmp_path, [{"id": 1, "revenue": 70}])
    mapping = add(work, tmp_path, [{"id": 1, "name": "A"}], name="map")
    revision = work.save_analysis(
        {
            "question": "收入",
            "sql": "WITH doubled AS (SELECT * FROM mapping UNION ALL SELECT * FROM mapping) SELECT SUM(a.revenue) AS revenue FROM actual a JOIN doubled d USING(id)",
            "bindings": {"actual": actual["id"], "mapping": mapping["id"]},
            "relationships": [
                {
                    "left": "actual",
                    "right": "mapping",
                    "keys": ["id"],
                    "cardinality": "many-to-one",
                }
            ],
        }
    )
    assert work.run_analysis(revision["id"])["execution_status"] == "failed"


def test_xlsx_physical_row_locations_and_empty_typed_column(work, tmp_path):
    wb = Workbook()
    wb.active.append(["amount"])
    wb.active.append([None])
    wb.active.append([100])
    p = tmp_path / "rows.xlsx"
    wb.save(p)
    data = work.add_dataset(p, name="rows")
    assert work.read_dataset(data["id"])["transformation"]["source_row_numbers"] == [3]
    p = tmp_path / "empty.json"
    p.write_text('[{"amount":null}]')
    data = work.add_dataset(
        p, name="empty", metadata={"column_types": {"amount": "DECIMAL(38,2)"}}
    )
    assert data["schema"][0]["type"] == "DECIMAL(38,2)"
    a = work.save_analysis(spec(data, "SELECT SUM(amount) AS amount FROM actual"))
    r = work.run_analysis(a["id"])
    assert r["execution_status"] == "succeeded"
    assert r["result"]["rows"][0]["amount"] is None


def test_realpath_aliases_do_not_loop_in_path_ownership_check(work):
    expected = (work.root / "future-output").resolve()
    assert work._inside(expected) == expected


def test_script_export_includes_code_output_and_hashes(work):
    run = work.run_script(
        {
            "code": "from pathlib import Path\nPath('value.txt').write_text('0.3')\nprint('computed 0.3')",
            "bindings": {},
        }
    )
    if run.get("runner_status") == "unavailable":
        pytest.skip("OS-isolated runner unavailable")
    assert run["execution_status"] == "succeeded"
    draft = work.save_draft({"body": "脚本测算", "run_ids": [run["id"]]})
    export = work.export_draft(draft["id"], "md")
    text = Path(export["path"]).read_text()
    assert (
        "Path('value.txt')" in text
        and "computed 0.3" in text
        and run["outputs"][0]["sha256"] in text
    )
    workbook = load_workbook(work.export_draft(draft["id"], "xlsx")["path"])
    evidence = dict(workbook["Evidence 1"].values)
    assert "computed 0.3" in evidence["stdout"] and "value.txt" in evidence["outputs"]
    (work.root / run["outputs"][0]["artifact_path"]).write_text("tampered")
    with pytest.raises(FoundryError, match="integrity"):
        work.get_run(run["id"])


def test_xlsx_long_body_keeps_tail(work):
    body = "a" * 40000 + "MUST_KEEP_TAIL"
    draft = work.save_draft({"body": body})
    workbook = load_workbook(work.export_draft(draft["id"], "xlsx")["path"])
    content = "".join(str(row[0].value or "") for row in workbook["Report"])
    assert body in content


def test_script_revision_can_rerun_via_unified_analysis_run(work):
    first = work.run_script({"code": "print('result')", "bindings": {}})
    if first.get("runner_status") == "unavailable":
        pytest.skip("OS-isolated runner unavailable")
    second = work.run_analysis(first["analysis_id"])
    assert second["analysis_id"] == first["analysis_id"]
    assert second["stdout"] == first["stdout"]
    assert second["execution_status"] == "succeeded"
    assert second["id"] != first["id"]


def test_version_order_is_monotonic_not_uuid_order(work, monkeypatch):
    import foundry.analysis as module

    original = module._new
    ids = iter(["f" * 32, "0" * 32])

    def new(task, **fields):
        record = original(task, **fields)
        record["id"] = next(ids)
        record["created_at"] = "2026-09-05T00:00:00Z"
        return record

    monkeypatch.setattr(module, "_new", new)
    older = work.save_draft({"body": "older"})
    newer = work.save_draft({"body": "newer", "parent_id": older["id"]})
    assert newer["seq"] > older["seq"]
    assert work.state()["drafts"][-1]["id"] == newer["id"]


def test_saved_methods_reuse_across_tasks_without_source_data(work, tmp_path):
    original = add(work, tmp_path, [{"region": "SecretClient", "profit": 30}])
    analysis = work.save_analysis(
        spec(
            original,
            "SELECT SUM(profit) AS profit FROM actual WHERE region=?",
            parameters=["SecretClient"],
        )
    )
    source_run = work.run_analysis(analysis["id"])
    method = work.save_method(analysis["id"], "利润分析")
    other = AnalysisStore(work.task.store.create(title="下一次分析"))
    data = add(other, tmp_path, [{"region": "NewClient", "profit": 90}])
    methods = other.state()["methods"]
    assert method["id"] in [m["id"] for m in methods]
    assert "SecretClient" not in json.dumps(methods) and original[
        "id"
    ] not in json.dumps(methods)
    applied = other.apply_method(
        method["id"], {"actual": data["id"]}, parameters=["NewClient"]
    )
    result = other.run_analysis(applied["id"])
    assert result["task_id"] == other.task.id
    assert Decimal(result["result"]["rows"][0]["profit"]) == 90
    assert methods[0]["source_task_id"] == work.task.id
    with pytest.raises(FoundryError):
        other.read_dataset(original["id"])
    with pytest.raises(FoundryError):
        other.get_run(source_run["id"])
    source_path = work.root / "methods" / (method["id"] + ".json")
    altered = json.loads(source_path.read_text())
    altered["sql"] = "SELECT 999"
    source_path.write_text(json.dumps(altered))
    with pytest.raises(FoundryError, match="integrity"):
        other.apply_method(
            method["id"], {"actual": data["id"]}, parameters=["NewClient"]
        )


def test_malformed_workbook_returns_domain_refusal(work, tmp_path):
    path = tmp_path / "broken.xlsx"
    path.write_text("not a zip")
    with pytest.raises(FoundryError, match="cannot parse"):
        work.add_dataset(path, name="broken")
