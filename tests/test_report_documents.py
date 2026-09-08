from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from foundry.analysis import AnalysisStore
from foundry.cli import main
from foundry.errors import FoundryError
from foundry.finance_service import FinanceStore
from foundry.report_documents import report_document_from_markdown_body, report_document_markdown, validate_report_document
from foundry.task import TaskStore


def _store(tmp_path: Path) -> AnalysisStore:
    task = TaskStore(tmp_path / "var").create(title="专业文稿")
    return AnalysisStore(task)


def _document(source: dict | None = None) -> dict:
    attrs = {"caption": "利润预算差异", "period": "2026-08", "currency": "CNY", "unit": "万元"}
    if source:
        attrs.update({"source_id": source["id"], "sheet": "月结", "range": "A1:B2"})
    return {
        "schema_version": 1,
        "template_version": "annual-v1",
        "metadata": {
            "company": "华宸物业",
            "period": "2026-08",
            "currency": "CNY",
            "unit": "万元",
        },
        **({"sources": [source]} if source else {}),
        "content": [
            {
                "type": "heading",
                "attrs": {"level": 1},
                "content": [{"type": "text", "text": "管理摘要"}],
            },
            {
                "type": "paragraph",
                "content": [{"type": "text", "text": "利润超预算 5.00 万元。"}],
            },
            {
                "type": "table",
                "attrs": attrs,
                "content": [
                    {
                        "type": "tableRow",
                        "content": [
                            {"type": "tableHeader", "content": [{"type": "text", "text": "项目"}]},
                            {"type": "tableHeader", "content": [{"type": "text", "text": "金额"}]},
                        ],
                    },
                    {
                        "type": "tableRow",
                        "content": [
                            {"type": "tableCell", "content": [{"type": "text", "text": "利润"}]},
                            {
                                "type": "tableCell",
                                "attrs": {"textAlign": "right"},
                                "content": [{"type": "text", "text": "5.00"}],
                            },
                        ],
                    },
                ],
            },
        ],
    }


def test_structured_report_document_projects_body_and_hashes_document(tmp_path: Path) -> None:
    store = _store(tmp_path)
    draft = store.save_draft({"title": "客户传来的标题", "body": "客户端旧正文", "document": _document()})
    assert "客户端旧正文" not in draft["body"]
    assert "企业：华宸物业；期间：2026-08；币种：CNY；单位：万元" in draft["body"]
    assert "利润超预算 5.00 万元。" in draft["body"]
    assert "表题：利润预算差异；期间：2026-08；币种：CNY；单位：万元" in draft["body"]
    assert "| 项目 | 金额 |" in draft["body"]
    assert draft["document"]["template_version"] == "annual-v1"
    same_text = store.save_draft({"title": "客户传来的标题", "body": draft["body"]})
    assert same_text["content_hash"] != draft["content_hash"]
    assert "document" not in same_text


def test_python_projection_matches_renderer_markdown_rules_for_marks_notes_and_tables() -> None:
    doc = {
        "schema_version": 1,
        "template_version": "annual-v1",
        "metadata": {},
        "content": [
            {
                "type": "paragraph",
                "content": [
                    {"type": "text", "text": "粗体", "marks": [{"type": "bold"}]},
                    {"type": "text", "text": "斜体", "marks": [{"type": "italic"}]},
                    {"type": "text", "text": "代码", "marks": [{"type": "code"}]},
                    {"type": "text", "text": "链接", "marks": [{"type": "link", "attrs": {"href": "https://example.com/report"}}]},
                ],
            },
            {
                "type": "reportNote",
                "attrs": {"id": "n1"},
                "content": [{"type": "paragraph", "content": [{"type": "text", "text": "附注正文"}]}],
            },
            {"type": "reportPageBreak"},
            {
                "type": "table",
                "content": [
                    {"type": "tableRow", "content": [{"type": "tableHeader", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "项目"}]}]}, {"type": "tableHeader", "content": [{"type": "text", "text": "金额"}]}]},
                    {"type": "tableRow", "content": [{"type": "tableCell", "attrs": {"colspan": 2}, "content": [{"type": "text", "text": "合并行"}]}]},
                ],
            },
        ],
    }
    assert report_document_markdown(doc) == (
        "**粗体**斜体`代码`[链接](https://example.com/report)\n\n"
        "[^n1]: 附注正文\n\n"
        '<div class="report-page-break"></div>\n\n'
        "| 项目 | 金额 |\n"
        "| --- | --- |\n"
        "| 合并行 |  |\n"
    )



def test_python_table_projection_escapes_pipes_backslashes_and_line_breaks_deterministically() -> None:
    def header(text: str) -> dict:
        return {"type": "tableHeader", "content": [{"type": "paragraph", "content": [{"type": "text", "text": text}]}]}

    def cell(text: str) -> dict:
        return {"type": "tableCell", "content": [{"type": "paragraph", "content": [{"type": "text", "text": text}]}]}

    doc = {
        "schema_version": 1,
        "template_version": "annual-v1",
        "metadata": {},
        "content": [
            {
                "type": "table",
                "content": [
                    {"type": "tableRow", "content": [header("项目"), header("说明"), header("金额")]},
                    {
                        "type": "tableRow",
                        "content": [
                            cell("收入|成本"),
                            cell("字面\\|组合\\路径\n第二行"),
                            cell("1234567890.123400"),
                        ],
                    },
                ],
            }
        ],
    }

    body = report_document_markdown(doc)
    assert body == (
        "| 项目 | 说明 | 金额 |\n"
        "| --- | --- | --- |\n"
        "| 收入\\|成本 | 字面\\\\\\|组合\\\\路径<br>第二行 | 1234567890.123400 |\n"
    )
    parsed = report_document_from_markdown_body(body)
    row = parsed["content"][0]["content"][1]["content"]
    assert len(row) == 3
    assert row[0]["content"][0]["content"][0]["text"] == "收入|成本"
    assert row[1]["content"][0]["content"][0]["text"] == "字面\\|组合\\路径<br>第二行"
    assert row[2]["content"][0]["content"][0]["text"] == "1234567890.123400"


def test_tiptap_optional_null_attrs_are_normalized_without_weakening_required_attrs() -> None:
    doc = {
        "schema_version": 1,
        "template_version": "annual-v1",
        "metadata": {},
        "content": [
            {
                "type": "paragraph",
                "attrs": {"source_id": None, "page": None, "sheet": None, "range": None, "textAlign": None},
                "content": [{"type": "text", "text": "普通编辑段落"}],
            },
            {
                "type": "table",
                "attrs": {"caption": None, "period": None, "currency": None, "unit": None, "source_id": None},
                "content": [
                    {
                        "type": "tableRow",
                        "content": [
                            {
                                "type": "tableHeader",
                                "attrs": {"textAlign": None, "colwidth": None, "source_id": None},
                                "content": [{"type": "paragraph", "content": [{"type": "text", "text": "项目"}]}],
                            }
                        ],
                    }
                ],
            },
        ],
    }

    clean = validate_report_document(doc)
    assert "attrs" not in clean["content"][0]
    assert "attrs" not in clean["content"][1]
    assert "attrs" not in clean["content"][1]["content"][0]["content"][0]
    assert report_document_markdown(clean) == "普通编辑段落\n\n| 项目 |\n| --- |\n"
    bad_image = {**doc, "content": [{"type": "image", "attrs": {"src": None}}]}
    with pytest.raises(FoundryError, match="image src"):
        validate_report_document(bad_image)

def test_body_only_revision_drops_inherited_report_document(tmp_path: Path) -> None:
    store = _store(tmp_path)
    rich = store.save_draft({"title": "经营分析", "document": _document()})
    inherited = store.save_draft({"parent_id": rich["id"], "review_status": "reviewed"})
    assert inherited["document"] == rich["document"]
    body_only = store.save_draft({"parent_id": rich["id"], "body": "经理仅保留正文修订。"})
    assert body_only["body"] == "经理仅保留正文修订。"
    assert "document" not in body_only


def test_report_source_xlsx_read_preserves_source_and_exact_visible_numbers(tmp_path: Path) -> None:
    store = _store(tmp_path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "月结"
    sheet.append(["项目", "金额", "公式"])
    sheet.append(["利润", "1234567890.1234", "=1+2"])
    sheet.merge_cells("A4:B4")
    sheet["A4"] = "附注：不信任导入审核状态"
    source_path = tmp_path / "monthly.xlsx"
    workbook.save(source_path)

    source = store.add_report_source(source_path)
    result = store.read_report_xlsx(source["source"]["id"], sheet="月结", range_ref="A1:C4")
    assert result["source"] == source["source"]
    assert result["sheets"] == ["月结"]
    assert result["document"]["sources"][0]["id"] == source["source"]["id"]
    table = next(node for node in result["document"]["content"] if node["type"] == "table")
    assert table["content"][0]["content"][0]["content"][0]["type"] == "paragraph"
    note = next(node for node in result["document"]["content"] if node["type"] == "reportNote")
    assert note["content"][0]["type"] == "paragraph"
    text = store.save_draft({"title": result["title"], "document": result["document"]})["body"]
    assert "1234567890.1234" in text
    assert "来源：monthly.xlsx；工作表：月结；区域：A1:C4。" in text
    spoofed = result["document"]
    spoofed["sources"][0] = {**spoofed["sources"][0], "name": "伪造.xlsx", "sha256": "b" * 64}
    saved = store.save_draft({"title": "回填来源", "document": spoofed})
    assert saved["document"]["sources"][0] == source["source"]


def test_report_document_refuses_unregistered_source_and_remote_image(tmp_path: Path) -> None:
    store = _store(tmp_path)
    source = {"id": "0" * 32, "name": "outside.xlsx", "sha256": "a" * 64, "format": "xlsx"}
    with pytest.raises(FoundryError, match="outside this task|not task-owned"):
        store.save_draft({"document": _document(source)})
    bad = _document()
    bad["content"].append({"type": "image", "attrs": {"src": "https://example.com/a.png"}})
    with pytest.raises(FoundryError, match="data URL|PNG/JPEG"):
        store.save_draft({"document": bad})


def test_xlsx_draft_export_includes_structured_chinese_report_sheets(tmp_path: Path) -> None:
    store = _store(tmp_path)
    draft = store.save_draft({"title": "经营分析", "document": _document()})
    export = store.export_draft(draft["id"], "xlsx")
    workbook = load_workbook(export["path"], data_only=True)
    try:
        assert workbook.sheetnames[:6] == ["Report", "报告概览", "分段正文", "财务表格", "附注", "来源"]
        assert workbook["报告概览"]["A1"].value == "字段"
        assert workbook["报告概览"]["B8"].value == "华宸物业"
        assert workbook["分段正文"]["D2"].value == "管理摘要"
        assert workbook["财务表格"]["I2"].value == "利润预算差异"
        assert workbook["财务表格"]["K2"].value == "CNY"
        assert workbook["Report"]["A1"].value == "经营分析"
    finally:
        workbook.close()


def test_report_export_register_is_immutable_and_readable(tmp_path: Path) -> None:
    store = _store(tmp_path)
    draft = store.save_draft({"title": "经营分析", "document": _document()})
    html = tmp_path / "report.html"
    html.write_text("<!doctype html><meta charset='utf-8'><h1>经营分析</h1>", encoding="utf-8")
    receipt = store.register_report_export(draft["id"], "html", html)
    assert Path(receipt["path"]).read_text(encoding="utf-8").startswith("<!doctype html>")
    assert receipt["content_hash"] == draft["content_hash"]
    html.write_text("<!doctype html><h1>changed</h1>", encoding="utf-8")
    replay = store.register_report_export(draft["id"], "html", html)
    assert replay["sha256"] == receipt["sha256"]
    assert Path(replay["path"]).read_text(encoding="utf-8").startswith("<!doctype html><meta")


def test_report_cli_spec_file_source_and_draft_read(tmp_path: Path, capsys) -> None:
    root = tmp_path / "var"
    task = TaskStore(root).create(title="CLI")
    args = ["--var", str(root), "--task", task.id]
    spec_file = tmp_path / "draft.json"
    spec_file.write_text(json.dumps({"title": "大文稿", "document": _document()}, ensure_ascii=False), encoding="utf-8")
    assert main([*args, "analysis-draft", "--spec-file", str(spec_file)]) == 0
    draft = json.loads(capsys.readouterr().out)
    assert draft["title"] == "大文稿"
    assert main([*args, "report-draft-read", "--draft", draft["id"]]) == 0
    assert json.loads(capsys.readouterr().out)["id"] == draft["id"]

    html = tmp_path / "report.html"
    html.write_text("<h1>报告</h1>", encoding="utf-8")
    assert main([*args, "report-source-add", "--path", str(html)]) == 0
    source = json.loads(capsys.readouterr().out)
    assert source["source"]["format"] == "html"
    assert main([*args, "report-source-read", "--source", source["source"]["id"]]) == 0
    reread = json.loads(capsys.readouterr().out)
    assert reread["source"] == source["source"]
    assert Path(reread["path"]).is_file()


def test_finance_save_explanation_draft_spec_file_uses_report_projection(tmp_path: Path, capsys) -> None:
    root = tmp_path / "var"
    task = TaskStore(root).create(title="Finance CLI")
    args = ["--var", str(root), "--task", task.id]
    spec_file = tmp_path / "finance-draft.json"
    spec_file.write_text(
        json.dumps({"title": "结构化财务稿", "body": " stale ", "document": _document()}, ensure_ascii=False),
        encoding="utf-8",
    )
    assert main([*args, "save-explanation-draft", "--spec-file", str(spec_file)]) == 0
    draft = json.loads(capsys.readouterr().out)
    assert " stale " not in draft["body"]
    assert "利润超预算 5.00 万元。" in draft["body"]
    assert draft["document"]["template_version"] == "annual-v1"



def test_finance_rich_fact_template_preserves_structured_attrs_sources_and_images(tmp_path: Path) -> None:
    task = TaskStore(tmp_path / "var").create(title="财务结构化 token 保存")
    analysis = AnalysisStore(task)
    bindings = {}
    for scenario, amounts in (("actual", [120, 70, 20]), ("budget", [100, 60, 15])):
        path = tmp_path / f"{scenario}.json"
        path.write_text(
            json.dumps(
                [
                    {
                        "account_code": code,
                        "management_account": account,
                        "amount": amount,
                    }
                    for code, account, amount in zip(
                        ("001", "002", "003"),
                        ("revenue", "cost", "expense"),
                        amounts,
                    )
                ]
            ),
            encoding="utf-8",
        )
        bindings[scenario] = analysis.add_dataset(path, name=scenario)["id"]
    source_file = tmp_path / "evidence.html"
    source_file.write_text("<h1>本地证据</h1>", encoding="utf-8")
    report_source = analysis.add_report_source(source_file)["source"]
    finance = FinanceStore(task)
    context = finance.save_context(
        {
            "bindings": bindings,
            "scopes": {
                alias: {
                    "ledger": "L1",
                    "organization": "O1",
                    "period": "2026-09",
                    "currency": "CNY",
                    "unit": "万元",
                    "scenario": alias,
                }
                for alias in bindings
            },
        }
    )
    run = finance.explain_variance({"context_id": context["id"]})
    fact = next(
        item
        for item in finance.prepare_explanation({"run_ids": [run["id"]]})["facts"]
        if item["row"] == 3 and item["column"] == "contribution"
    )
    document = _document(report_source)
    document["content"][1]["content"][0]["text"] = "利润超预算 {{fact:%s}}。" % fact["id"]
    table = next(node for node in document["content"] if node["type"] == "table")
    table["content"][1]["content"][0]["attrs"] = {"colspan": 2, "source_id": report_source["id"], "sheet": "月结", "range": "A2:B2"}
    document["content"].append(
        {
            "type": "image",
            "attrs": {
                "source_id": report_source["id"],
                "src": "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAFgwJ/lOjG0wAAAABJRU5ErkJggg==",
                "alt": "利润趋势图",
                "title": "本地图片证据",
            },
        }
    )

    draft = finance.save_explanation_draft({"title": "结构化 fact 文稿", "document": document, "run_ids": [run["id"]]})
    saved_table = next(node for node in draft["document"]["content"] if node["type"] == "table")
    saved_image = next(node for node in draft["document"]["content"] if node["type"] == "image")

    assert "{{fact:" not in draft["body"]
    assert "利润超预算 5 万元。" in draft["body"]
    assert "{{fact:" in draft["body_template"]
    assert saved_table["attrs"] == {
        "caption": "利润预算差异",
        "period": "2026-08",
        "currency": "CNY",
        "unit": "万元",
        "source_id": report_source["id"],
        "sheet": "月结",
        "range": "A1:B2",
    }
    assert saved_table["content"][1]["content"][0]["attrs"] == {
        "source_id": report_source["id"],
        "sheet": "月结",
        "range": "A2:B2",
        "colspan": 2,
    }
    assert saved_image["attrs"]["source_id"] == report_source["id"]
    assert saved_image["attrs"]["alt"] == "利润趋势图"
    assert report_document_markdown(draft["document"]) == draft["body"]
    body_only = analysis.save_draft({"title": "结构化 fact 文稿", "body": draft["body"]})
    assert body_only["content_hash"] != draft["content_hash"]

def test_finance_rich_document_save_appends_evidence_into_authoritative_document(tmp_path: Path) -> None:
    task = TaskStore(tmp_path / "var").create(title="财务结构化保存")
    analysis = AnalysisStore(task)
    bindings = {}
    for scenario, amounts in (("actual", [120, 70, 20]), ("budget", [100, 60, 15])):
        path = tmp_path / f"{scenario}.json"
        path.write_text(
            json.dumps(
                [
                    {
                        "account_code": code,
                        "management_account": account,
                        "amount": amount,
                    }
                    for code, account, amount in zip(
                        ("001", "002", "003"),
                        ("revenue", "cost", "expense"),
                        amounts,
                    )
                ]
            ),
            encoding="utf-8",
        )
        bindings[scenario] = analysis.add_dataset(path, name=scenario)["id"]
    finance = FinanceStore(task)
    context = finance.save_context(
        {
            "bindings": bindings,
            "scopes": {
                alias: {
                    "ledger": "L1",
                    "organization": "O1",
                    "period": "2026-09",
                    "currency": "CNY",
                    "unit": "万元",
                    "scenario": alias,
                }
                for alias in bindings
            },
        }
    )
    run = finance.explain_variance({"context_id": context["id"]})
    draft = finance.save_explanation_draft(
        {
            "title": "结构化财务稿",
            "document": _document(),
            "run_ids": [run["id"]],
        }
    )
    assert "## 计算依据" in draft["body"]
    assert report_document_markdown(draft["document"]) == draft["body"]
    assert any(node["type"] == "heading" and node["content"][0]["text"] == "计算依据" for node in draft["document"]["content"])
