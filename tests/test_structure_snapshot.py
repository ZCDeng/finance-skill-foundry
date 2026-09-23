"""U2: a confirmed declaration becomes a calculable table that still points home."""

from __future__ import annotations

import sys
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

from foundry.analysis import AnalysisStore
from foundry.errors import FoundryError
from foundry.finance_model import normalize_finance_dataset
from foundry.finance_service import FinanceStore
from foundry.structure_proposal import propose_workbook_structure
from foundry.structure_snapshot import derive_table
from foundry.task import TaskStore

sys.path.insert(0, str(Path(__file__).parent / "fixtures" / "structure"))
import build_huachen_like as fixture  # noqa: E402

SCOPE = {
    "ledger": "华宸物业",
    "organization": "集团合并",
    "period": "2026-08",
    "currency": "CNY",
    "unit": "元",
    "scenario": "actual",
}


def _bytes(book: Workbook) -> bytes:
    import io

    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def _ledger_declaration(data: bytes) -> dict:
    return next(
        sheet
        for sheet in propose_workbook_structure(data)["sheets"]
        if sheet["sheet"].startswith("02")
    )


def _task(tmp_path: Path):
    task = TaskStore(tmp_path / "var").create(title="结构快照")
    return task, AnalysisStore(task), FinanceStore(task)


def _shell(analysis: AnalysisStore, tmp_path: Path, data: bytes, name: str = "月结套表") -> dict:
    path = tmp_path / f"{name}.xlsx"
    path.write_bytes(data)
    return analysis.add_dataset(path, name=name, metadata={"mode": "workbook"})


def _account_map(service: FinanceStore) -> dict:
    rows = []
    for code, *_rest in fixture.LEDGER_ROWS:
        account = fixture.management_account(code)
        if account:
            rows.append({"account_code": code, "management_account": account})
    return service.save_account_map({"name": "华宸损益映射", "rows": rows})


def test_derived_table_drops_non_data_rows_and_folds_the_pair() -> None:
    data = fixture.monthly_pack()
    payload = derive_table(data, _ledger_declaration(data))
    names = [column["name"] for column in payload["schema"]]
    assert names == [
        "account_code",
        "account_name",
        "category",
        "account_direction",
        "amount",
        "dc_flag",
    ]
    assert payload["row_count"] == len(fixture.LEDGER_ROWS)
    codes = [row["account_code"] for row in payload["rows"]]
    assert "资产类科目" not in codes and "合　计" not in codes

    by_code = {row["account_code"]: row for row in payload["rows"]}
    # 收入 takes the credit side, 成本费用 the debit side — the category decides.
    assert by_code["6001.01"]["amount"] == "257295000"
    assert by_code["6001.01"]["dc_flag"] == "贷"
    assert by_code["6602"]["amount"] == "30750000"
    assert by_code["6602"]["dc_flag"] == "借"


def test_ae2_a_located_figure_points_at_the_original_debit_or_credit_cell(
    tmp_path: Path,
) -> None:
    data = fixture.monthly_pack()
    task, analysis, service = _task(tmp_path)
    shell = _shell(analysis, tmp_path, data)
    snapshot = analysis.add_table_snapshot(
        shell["id"], _ledger_declaration(data), name="02·GL科目余额表"
    )
    loaded = analysis.read_dataset(snapshot["id"])
    normalized = normalize_finance_dataset(loaded, SCOPE, _account_map(service)["rows"])

    rows = {row["account_code"]: row for row in normalized["rows"]}
    expense = rows["6602"]["source_locator"]
    assert expense["sheet"] == "02·GL科目余额表"
    # The cell the manager would land on in the file she uploaded, not a row
    # number in the derived copy.
    assert expense["excel_range"] == f"02·GL科目余额表!G{_original_row(data, '6602')}"
    revenue = rows["6001.01"]["source_locator"]
    # Same snapshot, other side of the pair.
    assert revenue["excel_range"] == f"02·GL科目余额表!H{_original_row(data, '6001.01')}"
    assert "/" not in str(expense) and "file://" not in str(expense)


def _original_row(data: bytes, code: str) -> int:
    from foundry.workbook_bindings import open_workbook_bytes

    book = open_workbook_bytes(data, data_only=True)
    try:
        sheet = book["02·GL科目余额表"]
        for number, cells in enumerate(sheet.iter_rows(), start=1):
            if cells and str(cells[0].value or "") == code:
                return number
    finally:
        book.close()
    raise AssertionError(f"{code} not in the fixture ledger")


def test_snapshot_keeps_the_original_bytes_and_passes_its_integrity_check(
    tmp_path: Path,
) -> None:
    data = fixture.monthly_pack()
    _task_obj, analysis, _service = _task(tmp_path)
    shell = _shell(analysis, tmp_path, data)
    snapshot = analysis.add_table_snapshot(
        shell["id"], _ledger_declaration(data), name="02"
    )
    assert snapshot["source_hash"] == shell["source_hash"]
    assert snapshot["kind"] == "table"
    loaded = analysis.read_dataset(snapshot["id"])
    assert loaded["transformation"]["structure"]["derived_from"] == shell["id"]


def test_management_pnl_matches_the_consolidated_income_statement(tmp_path: Path) -> None:
    data = fixture.monthly_pack()
    task, analysis, service = _task(tmp_path)
    shell = _shell(analysis, tmp_path, data)
    snapshot = analysis.add_table_snapshot(
        shell["id"], _ledger_declaration(data), name="02"
    )
    account_map = _account_map(service)
    context = service.save_context(
        {
            "name": "华宸 2026-08 管理损益",
            "bindings": {"actual": snapshot["id"]},
            "account_map_id": account_map["id"],
            "scopes": {"actual": SCOPE},
        }
    )
    run = service.run_metric(
        {"context_id": context["id"], "metric": {"kind": "management_pnl"}}
    )
    row = run["result"]["rows"][0]
    expected = fixture.expected_pnl()
    assert Decimal(row["revenue"]) == Decimal(expected["revenue"])
    assert Decimal(row["cost"]) == Decimal(expected["cost"])
    assert Decimal(row["expense"]) == Decimal(expected["expense"])


def test_fill_down_values_are_recorded_and_excluded_rows_are_gone() -> None:
    data = fixture.monthly_pack()
    journal = next(
        sheet
        for sheet in propose_workbook_structure(data)["sheets"]
        if sheet["sheet"].startswith("03")
    )
    payload = derive_table(data, journal)
    assert payload["row_count"] == fixture.journal_line_count()
    assert all(row["doc_id"] for row in payload["rows"])
    provenance = payload["transformation"]["structure"]["fill_provenance"]
    assert provenance["doc_id"]
    assert all(row["summary"] for row in payload["rows"])
    # Every 小计 and the 合计 are gone.
    assert all("小" not in str(row["summary"]) for row in payload["rows"])


def test_reconfirming_writes_a_new_snapshot_and_leaves_old_runs_alone(
    tmp_path: Path,
) -> None:
    data = fixture.monthly_pack()
    task, analysis, service = _task(tmp_path)
    shell = _shell(analysis, tmp_path, data)
    declaration = _ledger_declaration(data)
    first = analysis.add_table_snapshot(shell["id"], declaration, name="02")
    account_map = _account_map(service)
    context = service.save_context(
        {
            "bindings": {"actual": first["id"]},
            "account_map_id": account_map["id"],
            "scopes": {"actual": SCOPE},
        }
    )
    run = service.run_metric(
        {"context_id": context["id"], "metric": {"kind": "management_pnl"}}
    )
    second = analysis.add_table_snapshot(shell["id"], declaration, name="02 修订")
    assert second["id"] != first["id"]
    assert service.analysis.get_run(run["id"])["bindings"] == {"actual": first["id"]}


def test_readiness_treats_the_snapshot_as_an_ordinary_table(tmp_path: Path) -> None:
    from foundry.finance_readiness import _is_workbook_shell

    data = fixture.monthly_pack()
    _task_obj, analysis, _service = _task(tmp_path)
    shell = _shell(analysis, tmp_path, data)
    snapshot = analysis.add_table_snapshot(
        shell["id"], _ledger_declaration(data), name="02"
    )
    loaded = analysis.read_dataset(snapshot["id"])
    assert _is_workbook_shell(loaded) is False
    assert _is_workbook_shell(analysis.read_dataset(shell["id"])) is True


def test_an_undecided_debit_credit_side_refuses_to_become_a_table() -> None:
    data = fixture.no_category_ledger()
    declaration = propose_workbook_structure(data)["sheets"][0]
    with pytest.raises(FoundryError, match="undecided"):
        derive_table(data, declaration)


def test_a_manager_chosen_fixed_side_derives_the_table() -> None:
    data = fixture.no_category_ledger()
    declaration = propose_workbook_structure(data)["sheets"][0]
    declaration["amount_pairs"][0]["direction"] = {"kind": "fixed", "side": "credit"}
    payload = derive_table(data, declaration)
    assert [row["dc_flag"] for row in payload["rows"]] == ["贷", "贷"]


def test_undeclarable_reason_names_what_cannot_be_built() -> None:
    from foundry.structure_snapshot import undeclarable_reason

    unreadable = {
        "sheet": "10·区域经营损益表",
        "header_row": 4,
        "columns": [
            {"letter": "A", "display": "区域", "role": "unknown"},
            {"letter": "B", "display": "营业收入", "role": "unknown"},
        ],
        "amount_pairs": [],
    }
    assert undeclarable_reason(unreadable) is None
    assert undeclarable_reason({"columns": []}) == "no column was declared"
    # One named role is enough; a 借贷 pair on its own is too.
    assert undeclarable_reason({**unreadable, "columns": [
        {"letter": "A", "display": "科目代码", "role": "account_code"},
    ]}) is None


def test_unknown_numeric_columns_survive_with_source_letters() -> None:
    data = fixture.monthly_pack()
    declaration = next(
        sheet
        for sheet in propose_workbook_structure(data)["sheets"]
        if sheet["sheet"].startswith("10")
    )
    payload = derive_table(data, declaration)

    names = [column["name"] for column in payload["schema"]]
    assert names == ["col_A_区域", "col_B_营业收入", "col_C_营业成本", "col_D_期间费用", "col_E_营业利润"]
    assert payload["rows"][0]["col_A_区域"] == "华东区域"
    assert payload["rows"][0]["col_B_营业收入"] == "109742.028"
    structure = payload["transformation"]["structure"]
    assert payload["transformation"]["header_columns"]["col_B_营业收入"] == "B"
    assert structure["column_roles"]["col_B_营业收入"] == "unknown"
    assert structure["numeric_coverage"]["numeric_columns_included"] == 4


def test_explicitly_excluded_unknown_column_records_reason() -> None:
    data = fixture.monthly_pack()
    declaration = next(
        sheet
        for sheet in propose_workbook_structure(data)["sheets"]
        if sheet["sheet"].startswith("10")
    )
    declaration["columns"] = [
        {
            **column,
            "include": False,
            "exclusion_reason": "not needed for this calculation",
        }
        if column["letter"] == "D"
        else column
        for column in declaration["columns"]
    ]

    payload = derive_table(data, declaration)
    names = [column["name"] for column in payload["schema"]]
    assert "col_D_期间费用" not in names
    excluded = payload["transformation"]["structure"]["excluded_columns"]
    assert excluded == [
        {
            "letter": "D",
            "header": "期间费用",
            "stable_name": "col_D_期间费用",
            "reason": "not needed for this calculation",
        }
    ]


def test_selected_block_bounds_keep_independent_tables_from_bleeding() -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "预算表"
    sheet["A1"] = "年度预算"
    for index, text in enumerate(["项目", "预算", "实际"], start=1):
        sheet.cell(row=2, column=index, value=text)
    sheet.append(["住宅", 10, 9])
    sheet.append(["商业", 20, 19])
    for index, text in enumerate(["城市公司", "合同金额"], start=1):
        sheet.cell(row=6, column=index, value=text)
    sheet.cell(row=7, column=1, value="宁波")
    sheet.cell(row=7, column=2, value=100)
    sheet.cell(row=8, column=1, value="常州")
    sheet.cell(row=8, column=2, value=200)
    data = _bytes(book)
    proposal = propose_workbook_structure(data)["sheets"][0]
    second = proposal["table_blocks"][1]
    declaration = {
        **proposal,
        "selected_block_id": second["id"],
        "header_row": second["header_row"],
        "start_row": second["start_row"],
        "end_row": second["end_row"],
        "start_column": second["start_column"],
        "end_column": second["end_column"],
        "columns": second["columns"],
        "amount_pairs": second["amount_pairs"],
        "excluded_rows": second["excluded_rows"],
        "fill_down_columns": second["fill_down_columns"],
    }

    payload = derive_table(data, declaration)
    assert payload["row_count"] == 2
    assert payload["transformation"]["bounds"]["start_row"] == 6
    assert payload["transformation"]["bounds"]["end_row"] == 8
    assert [row["organization"] for row in payload["rows"]] == ["宁波", "常州"]
    assert "住宅" not in str(payload["rows"])
