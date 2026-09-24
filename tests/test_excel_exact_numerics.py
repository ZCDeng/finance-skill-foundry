from __future__ import annotations

import io
import re
import sys
import zipfile
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from openpyxl import Workbook, load_workbook

from foundry import analysis_data
from foundry.analysis import AnalysisStore, _new
from foundry.analysis_data import excel_number, parse_source, read_exact_numerics
from foundry.finance_files import FinanceFiles
from foundry.financial_evidence import evaluate_evidence_recipes
from foundry.source_refresh import SourceRefreshService, SourceScope
from foundry.task import TaskStore

sys.path.insert(0, str(Path(__file__).parent / "fixtures" / "structure"))
import build_huachen_like as fixture  # noqa: E402

# Every XLSX import carried this label until the reading rule was versioned:
# first with the literal digits, then for a while with the shortest form.
OLD_LABEL = "xlsx-original-xml-decimal"
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
REL_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


def _rewrite_literals(data: bytes, sheet_name: str, literals: dict[str, str]) -> bytes:
    """Write each literal into the sheet XML verbatim, as a non-openpyxl writer would."""
    source = zipfile.ZipFile(io.BytesIO(data))
    sheets = ET.fromstring(source.read("xl/workbook.xml")).find("s:sheets", NS)
    rid = next(sheet.attrib[REL_ID] for sheet in sheets if sheet.attrib["name"] == sheet_name)
    rels = ET.fromstring(source.read("xl/_rels/workbook.xml.rels"))
    target = next(rel.attrib["Target"] for rel in rels if rel.attrib["Id"] == rid)
    member = target.lstrip("/") if target.startswith("/") else "xl/" + target
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as rewritten:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == member:
                xml = content.decode()
                for coordinate, literal in literals.items():
                    xml, count = re.subn(
                        rf'(<c r="{coordinate}"[^>]*>\s*<v>)[^<]*(</v>)',
                        rf"\g<1>{literal}\g<2>",
                        xml,
                    )
                    assert count == 1, coordinate
                content = xml.encode()
            rewritten.writestr(item, content)
    return out.getvalue()


def _xlsx_with_literals(
    literals: dict[str, str], texts: dict[str, str] | None = None, *, title: str | None = None
) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "预算"
    for coordinate, text in (texts or {}).items():
        sheet[coordinate] = text
    for coordinate in literals:
        sheet[coordinate] = 1
    if title:
        book.properties.title = title
    buffer = io.BytesIO()
    book.save(buffer)
    return _rewrite_literals(buffer.getvalue(), "预算", literals)


def test_seventeen_digit_literal_reads_as_shortest_form_of_the_same_double():
    assert excel_number("64441.376900000001") == Decimal("64441.3769")
    assert excel_number("527999.9999000001") == Decimal("527999.9999")
    assert excel_number("6.4441376900000001E4") == Decimal("64441.3769")
    assert float("64441.376900000001") == float("64441.3769")


def test_round_trip_literals_keep_their_digits():
    for literal in ("3185", "6984.7414", "0.1", "-12012.0347", "1E-7"):
        value = excel_number(literal)
        assert value == Decimal(literal)
        assert str(value) == str(Decimal(literal))


def test_distinct_doubles_stay_distinct():
    assert excel_number("64441.3769") != excel_number("64441.37690000002")


def test_integer_literals_are_kept_as_written():
    # Writers emit int64 IDs verbatim. Above 2**53 two of them share a double,
    # so reading them as that double would merge two different documents.
    assert excel_number("1790234567890123457") == Decimal("1790234567890123457")
    assert excel_number("1790234567890123457") != excel_number("1790234567890123458")
    assert str(excel_number("9007199254740993")) == "9007199254740993"


def test_workbook_reader_uses_the_stored_double_not_its_serialisation_tail():
    data = _xlsx_with_literals({"B2": "64441.376900000001", "C2": "3185", "D2": "527999.9999000001"})

    values = read_exact_numerics(data, "预算")

    assert values["B2"] == Decimal("64441.3769")
    assert values["C2"] == Decimal("3185")
    assert values["D2"] == Decimal("527999.9999")


def test_table_import_keeps_integer_ids_apart_and_records_the_rule():
    data = _xlsx_with_literals(
        {
            "A2": "1790234567890123457",
            "A3": "1790234567890123458",
            "A4": "9007199254740993",
            "B2": "64441.376900000001",
            "B3": "10",
            "B4": "527999.9999000001",
        },
        {"A1": "doc_id", "B1": "amount"},
    )

    parsed = parse_source(data, ".xlsx", {"sheet": "预算"})

    assert [row["doc_id"] for row in parsed["rows"]] == [
        "1790234567890123457",
        "1790234567890123458",
        "9007199254740993",
    ]
    assert [row["amount"] for row in parsed["rows"]] == ["64441.3769", "10", "527999.9999"]
    assert {column["name"]: column["type"] for column in parsed["schema"]}["doc_id"] == "DECIMAL(38,0)"
    assert parsed["transformation"]["numeric_parser"] == "xlsx-shortest-double"


def _store(tmp_path: Path, title: str = "exact numerics") -> AnalysisStore:
    return AnalysisStore(TaskStore(tmp_path / "var").create(title=title))


def _budget(tmp_path: Path, literals: dict[str, str], texts: dict[str, str] | None = None) -> Path:
    path = tmp_path / "预算.xlsx"
    path.write_bytes(_xlsx_with_literals(literals, texts))
    return path


def _file_under_old_label(analysis: AnalysisStore, path: Path, monkeypatch, *, verbatim: bool) -> dict:
    """File a table the way an earlier reader did, under the old label.

    ``verbatim`` is the reader before the shortest-double rule (the literal
    digits); otherwise the reader in between (the shortest form).
    """
    real = analysis_data.parse_source

    def parse(data, suffix, metadata):
        parsed = real(data, suffix, metadata)
        parsed["transformation"]["numeric_parser"] = OLD_LABEL
        return parsed

    with monkeypatch.context() as patch:
        if verbatim:
            patch.setattr(analysis_data, "excel_number", Decimal)
        patch.setattr("foundry.analysis.parse_source", parse)
        return analysis.add_dataset(path, name="预算", metadata={"sheet": "预算"})


def _sql_run(analysis: AnalysisStore, dataset: dict) -> dict:
    definition = analysis.save_analysis(
        {"question": "预算值", "sql": "SELECT value FROM budget", "bindings": {"budget": dataset["id"]}}
    )
    return analysis.run_analysis(definition["id"])


def _claim_run(analysis: AnalysisStore, value: str) -> dict:
    record = _new(
        analysis.task,
        analysis_id="a" * 32,
        engine="python",
        execution_status="succeeded",
        reproducibility_status="unverified",
        bindings={},
        input_hashes={},
        source_hashes={},
        result={"columns": [{"name": "value"}], "rows": [{"value": value}], "row_count": 1, "truncated": False},
    )
    return analysis._write(analysis._object_path("runs", record["id"]), record)


def _recipe_outcome(
    analysis: AnalysisStore, dataset: dict, run: dict, cells: str = "A2", sheet: str = "预算"
) -> tuple[str, str | None]:
    recipe = {
        "id": "source-check",
        "kind": "sum" if ":" in cells else "direct",
        "dataset_id": dataset["id"],
        "source_hash": dataset["source_hash"],
        "sheet": sheet,
        "range": cells,
        "claim": {"run_id": run["id"], "row": 0, "column": "value"},
    }
    result = evaluate_evidence_recipes(analysis, [recipe])["results"][0]
    return result["status"], result.get("actual")


def test_snapshot_filed_with_verbatim_literals_still_passes_its_recipe(tmp_path, monkeypatch):
    analysis = _store(tmp_path)
    path = _budget(tmp_path, {"A2": "64441.376900000001"}, {"A1": "value"})
    dataset = _file_under_old_label(analysis, path, monkeypatch, verbatim=True)
    assert analysis.read_dataset(dataset["id"])["rows"] == [{"value": "64441.376900000001"}]
    run = _sql_run(analysis, dataset)
    assert run["result"]["rows"] == [{"value": "64441.376900000001"}]

    assert _recipe_outcome(analysis, dataset, run) == ("pass", "64441.376900000001")
    # Read the way it was filed, not loosened: on this snapshot the shortest
    # spelling of the same double is a different decimal.
    assert _recipe_outcome(analysis, dataset, _claim_run(analysis, "64441.3769"))[0] == "mismatch"


def test_snapshot_filed_with_the_shortest_form_under_the_old_label_still_passes(tmp_path, monkeypatch):
    analysis = _store(tmp_path)
    path = _budget(tmp_path, {"A2": "64441.376900000001"}, {"A1": "value"})
    dataset = _file_under_old_label(analysis, path, monkeypatch, verbatim=False)
    assert analysis.read_dataset(dataset["id"])["rows"] == [{"value": "64441.3769"}]

    assert _recipe_outcome(analysis, dataset, _sql_run(analysis, dataset)) == ("pass", "64441.3769")


def test_new_snapshot_reads_the_shortest_double_and_distinct_doubles_stay_distinct(tmp_path):
    analysis = _store(tmp_path)
    path = _budget(tmp_path, {"A2": "64441.376900000001"}, {"A1": "value"})
    dataset = analysis.add_dataset(path, name="预算", metadata={"sheet": "预算"})
    assert dataset["transformation"]["numeric_parser"] == "xlsx-shortest-double"

    assert _recipe_outcome(analysis, dataset, _sql_run(analysis, dataset)) == ("pass", "64441.3769")
    assert _recipe_outcome(analysis, dataset, _claim_run(analysis, "64441.37690000002"))[0] == "mismatch"


def test_workbook_snapshot_sums_the_stored_doubles(tmp_path):
    # A workbook snapshot files no cell values, so nothing it stored can
    # disagree with the current rule. This is the Huachen B24:I24 shape.
    analysis = _store(tmp_path)
    path = _budget(tmp_path, {"B2": "8200.931200000001", "C2": "56240.4457"})
    dataset = analysis.add_dataset(path, name="预算", metadata={"mode": "workbook"})

    outcome = _recipe_outcome(analysis, dataset, _claim_run(analysis, "64441.3769"), "B2:C2")
    assert outcome == ("pass", "64441.3769")


def _source_row(data: bytes, sheet_name: str, code: str) -> int:
    book = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    try:
        for number, cells in enumerate(book[sheet_name].iter_rows(), start=1):
            if cells and str(cells[0].value or "") == code:
                return number
    finally:
        book.close()
    raise AssertionError(f"{code} not in {sheet_name}")


def test_structure_derived_snapshot_filed_with_verbatim_literals_still_passes(tmp_path, monkeypatch):
    import foundry.structure_snapshot as structure_snapshot
    from foundry.structure_proposal import propose_workbook_structure

    sheet = "02·GL科目余额表"
    pack = fixture.monthly_pack()
    row = _source_row(pack, sheet, "6001.01")
    # The revenue line sits on the credit side of the 借方/贷方 pair.
    data = _rewrite_literals(pack, sheet, {f"H{row}": "257295000.00000001"})
    path = tmp_path / "月结套表.xlsx"
    path.write_bytes(data)
    analysis = _store(tmp_path)
    shell = analysis.add_dataset(path, name="月结套表", metadata={"mode": "workbook"})
    declaration = next(item for item in propose_workbook_structure(data)["sheets"] if item["sheet"] == sheet)
    real = structure_snapshot.derive_table

    def derive(data, declaration):
        payload = real(data, declaration)
        payload["transformation"]["numeric_parser"] = OLD_LABEL
        return payload

    with monkeypatch.context() as patch:
        patch.setattr(analysis_data, "excel_number", Decimal)
        patch.setattr(structure_snapshot, "derive_table", derive)
        derived = analysis.add_table_snapshot(shell["id"], declaration, name="02")
    rows = analysis.read_dataset(derived["id"])["rows"]
    assert next(item for item in rows if item["account_code"] == "6001.01")["amount"] == "257295000.00000001"

    run = _claim_run(analysis, "257295000.00000001")
    assert _recipe_outcome(analysis, derived, run, f"H{row}", sheet) == ("pass", "257295000.00000001")


def test_preview_of_a_snapshot_shows_the_digits_it_filed(tmp_path, monkeypatch):
    for verbatim, exact in ((True, "64441.376900000001"), (False, "64441.3769")):
        analysis = _store(tmp_path / str(verbatim))
        path = _budget(tmp_path, {"A2": "64441.376900000001"}, {"A1": "value"})
        dataset = _file_under_old_label(analysis, path, monkeypatch, verbatim=verbatim)

        preview = FinanceFiles(analysis.task).range_preview(
            {"dataset_id": dataset["id"], "sheet": "预算", "range": "A2"}
        )

        assert preview["cells"][0]["exact"] == exact


def test_preview_without_formula_cache_reads_the_stored_double_and_keeps_integers(tmp_path):
    analysis = _store(tmp_path)
    # C2 has no cached value, so the preview takes its formula-cache fallback.
    path = _budget(tmp_path, {"A2": "8.199999999999999", "B2": "9007199254740993"}, {"C2": "=A2*2"})
    dataset = analysis.add_dataset(path, name="预算", metadata={"mode": "workbook"})

    preview = FinanceFiles(analysis.task).range_preview(
        {"dataset_id": dataset["id"], "sheet": "预算", "range": "A2:B2"}
    )

    assert {cell["address"]: cell["exact"] for cell in preview["cells"]} == {
        "A2": "8.2",
        "B2": "9007199254740993",
    }


def test_report_xlsx_reads_the_stored_double_and_keeps_integers(tmp_path):
    analysis = _store(tmp_path)
    path = _budget(tmp_path, {"A2": "8.199999999999999", "B2": "9007199254740993"}, {"A1": "金额", "B1": "编号"})
    source = analysis.add_report_source(path)

    result = analysis.read_report_xlsx(source["source"]["id"], sheet="预算", range_ref="A1:B2")

    table = next(node for node in result["document"]["content"] if node["type"] == "table")
    texts = [
        [cell["content"][0]["content"][0]["text"] for cell in table_row["content"]]
        for table_row in table["content"]
    ]
    assert texts[1] == ["8.2", "9007199254740993"]


def _refresh_scope() -> SourceScope:
    return SourceScope(
        source_id="erp",
        tenant_id="tenant-a",
        ledger_id="ledger-a",
        organization_id="org-east",
        period="2026-08",
        currency="CNY",
        cutoff_at="2026-09-05T00:00:00Z",
        mode="full_period",
    )


def _ledger_export(path: Path, amount: str, *, title: str) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "预算"
    sheet.append(["document_id", "line_id", "amount"])
    sheet.append(["JV-1", "1", 1])
    sheet.append(["JV-1", "2", 20])
    book.properties.title = title
    buffer = io.BytesIO()
    book.save(buffer)
    path.write_bytes(_rewrite_literals(buffer.getvalue(), "预算", {"C2": amount}))


@pytest.mark.parametrize("verbatim", [True, False])
def test_refresh_reports_an_unchanged_re_export_as_unchanged(tmp_path, monkeypatch, verbatim):
    task = TaskStore(tmp_path / "var").create(title="refresh")
    refresh = SourceRefreshService(task)
    path = tmp_path / "ledger.xlsx"
    _ledger_export(path, "64441.376900000001", title="v1")
    with monkeypatch.context() as patch:
        if verbatim:
            patch.setattr(analysis_data, "excel_number", Decimal)
        first = refresh.refresh_file(path, scope=_refresh_scope(), name="ledger")
    filed = AnalysisStore(task).read_dataset(first["dataset_id"])["rows"]
    assert filed[0]["amount"] == ("64441.376900000001" if verbatim else "64441.3769")

    # Re-exported with the same figures: new bytes, the same stored doubles.
    _ledger_export(path, "64441.376900000001", title="v2")
    second = refresh.refresh_file(path, scope=_refresh_scope(), name="ledger")
    assert second["change_counts"] == {"inserted": 0, "updated": 0, "deleted": 0, "voided": 0, "unchanged": 2}
    assert AnalysisStore(task).read_dataset(second["dataset_id"])["rows"][0]["amount"] == "64441.3769"

    # A different double is still an update.
    _ledger_export(path, "64441.37690000002", title="v3")
    third = refresh.refresh_file(path, scope=_refresh_scope(), name="ledger")
    assert third["change_counts"]["updated"] == 1
    assert third["change_counts"]["unchanged"] == 1
