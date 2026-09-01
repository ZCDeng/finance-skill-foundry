from __future__ import annotations

import io
import json
import time
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from foundry.cli import main
from foundry.resources import fixture_root
from foundry.sheet import put_number, survives_a_cell, workbook_bytes
from foundry.tags import EXTERNAL, NOT_IN_WINDOW, OUTSIDE, RECOMPUTED
from foundry.task import TaskStore
from foundry.window import WindowStore


def _out(capsys: pytest.CaptureFixture[str]) -> str:
    return capsys.readouterr().out.strip()


def _closed(var: Path, capsys: pytest.CaptureFixture[str]) -> tuple[str, str]:
    """A window taken all the way through a close, on the kit's own fixtures."""
    fixtures = fixture_root(None)
    assert main(["--var", str(var), "task-new", "--title", "2026-08 月结", "--period", "2026-08"]) == 0
    task_id = _out(capsys)
    base = ["--var", str(var), "--task", task_id]
    assert main([*base, "create"]) == 0
    assert (
        main(
            [
                *base,
                "import",
                "--coa",
                str(fixtures / "coa.xlsx"),
                "--units",
                str(fixtures / "units.xlsx"),
                "--map",
                str(fixtures / "map.xlsx"),
                "--named",
                str(fixtures / "named.json"),
                "--drop-dir",
                str(fixtures),
            ]
        )
        == 0
    )
    assert main([*base, "lock"]) == 0
    assert main([*base, "record-questions", "--questions", str(fixtures / "questions.json")]) == 0
    assert main([*base, "emit-brief"]) == 0
    assert (
        main(
            [
                *base,
                "record-reference",
                "--references",
                json.dumps(
                    [
                        {
                            "id": "ref-1",
                            "text": "行业收缴率中位数 92.4%",
                            "source": "中指研究院",
                            "at": "2026-08-30",
                        }
                    ],
                    ensure_ascii=False,
                ),
            ]
        )
        == 0
    )
    capsys.readouterr()
    return task_id, WindowStore(var).current().id


def _book(var: Path, sha: str) -> Workbook:
    task = TaskStore(var).current()
    assert task is not None
    return load_workbook(task.artifact_path(sha, ext=".xlsx"))


def _rows(sheet) -> list[list[str]]:
    return [
        ["" if cell is None else str(cell) for cell in row]
        for row in sheet.iter_rows(min_row=2, values_only=True)
    ]


# -- the number rule --------------------------------------------------------


def test_a_figure_the_format_cannot_hold_goes_in_as_text() -> None:
    """xlsx numbers are binary64, and the catalog is not.

    `catalog.py` refuses to IMPORT a value beyond 2^53 because a spreadsheet
    would already have rounded it. The way out has the same problem and a
    different answer: the figure is real and the operator needs to see it, so
    it goes in exact, as text, rather than as a number that is nearly right.
    """
    book = Workbook()
    sheet = book.active
    put_number(sheet["A1"], 9007199254740993)
    put_number(sheet["A2"], "282670.7525")
    put_number(sheet["A3"], 5000)

    buffer = io.BytesIO()
    book.save(buffer)
    read = load_workbook(io.BytesIO(buffer.getvalue())).active

    assert read["A1"].value == "9007199254740993"
    assert isinstance(read["A1"].value, str)
    # These two survive, so they stay arithmetic.
    assert read["A2"].value == pytest.approx(282670.7525)
    assert read["A3"].value == 5000


def test_the_round_trip_test_is_the_actual_question() -> None:
    assert survives_a_cell(Decimal("282670.7525"))
    assert survives_a_cell(Decimal(5000))
    assert not survives_a_cell(Decimal("9007199254740993"))


# -- the workbook -----------------------------------------------------------


def test_the_workbook_carries_the_lock_it_was_written_from(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    """Without the hashes it is a spreadsheet, not a working paper."""
    var = tmp_path / "var"
    _closed(var, capsys)
    window = WindowStore(var).current()
    book = load_workbook(io.BytesIO(workbook_bytes(window)))

    cover = "\n".join(
        str(cell) for row in book["本期"].iter_rows(values_only=True) for cell in row if cell
    )
    assert window.id in cover
    for digest in window.hashes()["named"].values():
        assert digest in cover
    # And every word a reader will meet, explained on the page they land on.
    for tag in (RECOMPUTED, EXTERNAL, NOT_IN_WINDOW, OUTSIDE):
        assert tag in cover


def test_every_line_carries_which_schedule_it_belongs_to(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    """The three columns that make it pivotable rather than a dump.

    A schedule heading above a block of rows is a markdown idea. In a sheet the
    item, its tag and its origin have to be ON the row, or filtering the 底稿
    for 我重算 hands back lines with nothing saying what they belong to.
    """
    var = tmp_path / "var"
    _closed(var, capsys)
    window = WindowStore(var).current()
    book = load_workbook(io.BytesIO(workbook_bytes(window)))
    rows = _rows(book["底稿"])

    assert rows, "the workpaper had schedules and the sheet has no rows"
    for row in rows:
        assert row[0], f"a line with no 条目: {row}"
        assert row[2], f"a line with no 标签: {row}"
        assert row[3], f"a line with no 来源: {row}"

    # An account code is an identifier that happens to be digits.
    codes = [cell.value for cell in book["底稿"]["F"][1:] if cell.value is not None]
    assert codes and all(isinstance(code, str) for code in codes)


def test_a_refusal_and_a_reference_each_keep_a_row(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    """Neither has arithmetic, and a sheet that dropped them would hold fewer
    schedules than the workpaper it renders."""
    var = tmp_path / "var"
    _closed(var, capsys)
    window = WindowStore(var).current()
    book = load_workbook(io.BytesIO(workbook_bytes(window)))
    rows = _rows(book["底稿"])

    refusal = next(row for row in rows if row[2] == NOT_IN_WINDOW)
    assert "拒绝作答" in refusal[10]
    assert refusal[8] == "", "a refusal must not carry a figure"

    reference = next(row for row in rows if row[2] == EXTERNAL)
    assert "中指研究院" in reference[10]
    assert reference[8] == "", "a cited line has no arithmetic to show"

    # And its own sheet, because a reader looking for what came from outside
    # should not have to filter a 400-row schedule to find four of them.
    assert [row[0] for row in _rows(book[EXTERNAL])] == ["ref-1"]


def test_the_sheet_and_the_workpaper_say_the_same_thing_about_a_follow_up(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    """One rule, two renderings.

    A follow-up's item is stored with whatever the recompute called it and gets
    `稿外追问` put in front at render time. Both documents call `display_tag`,
    so the same schedule cannot read as two different claims depending on which
    file was opened.
    """
    var = tmp_path / "var"
    task_id, _ = _closed(var, capsys)
    assert (
        main(
            [
                "--var",
                str(var),
                "--task",
                task_id,
                "--session",
                "sess-leader",
                "query-locked",
                "--text",
                "按核算单元拆一下",
                "--tables",
                "gl",
                "--split-by",
                "unit_id",
            ]
        )
        == 0
    )
    capsys.readouterr()

    window = WindowStore(var).current()
    book = load_workbook(io.BytesIO(workbook_bytes(window)))
    row = next(row for row in _rows(book["底稿"]) if row[0] == "followup-1")
    paper = window.workpaper_path.read_text()

    assert row[2] == f"{OUTSIDE} / {RECOMPUTED}"
    assert f"标签: {row[2]}" in paper
    assert f"来源: {row[3]}" in paper


# -- the deliverable --------------------------------------------------------


def test_the_export_files_a_version_and_the_bytes_are_a_workbook(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    var = tmp_path / "var"
    task_id, window_id = _closed(var, capsys)
    assert main(["--var", str(var), "--task", task_id, "export", "--format", "xlsx"]) == 0
    row = json.loads(_out(capsys))

    assert row["kind"] == "export"
    assert row["ext"] == ".xlsx"
    assert row["cause"] == "export:xlsx"
    assert row["window_id"] == window_id
    assert _book(var, row["sha256"]).sheetnames[:3] == ["本期", "结论", "底稿"]


def test_a_workbook_is_the_same_bytes_whenever_it_is_rendered(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    """The clock must not reach the file.

    openpyxl stamps `docProps/core.xml` from the clock, and a deliverable
    version is keyed on its bytes -- so an unrepaired workbook files a new
    version of an unchanged window every time the second ticks over. It cost a
    day to find, because two renders in the same instant match and the suite
    only failed when the pair straddled a boundary.

    Renders across a real tick rather than back to back, which is the only
    version of this test that would have caught it.
    """
    var = tmp_path / "var"
    _closed(var, capsys)
    window = WindowStore(var).current()

    first = workbook_bytes(window)
    time.sleep(1.1)
    second = workbook_bytes(window)
    assert first == second


def test_re_exporting_unchanged_files_no_second_version(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    """The same rule the markdown deliverables already have."""
    var = tmp_path / "var"
    task_id, _ = _closed(var, capsys)
    base = ["--var", str(var), "--task", task_id, "export", "--format", "xlsx"]
    assert main(base) == 0
    first = json.loads(_out(capsys))
    time.sleep(1.1)
    assert main(base) == 0
    second = json.loads(_out(capsys))

    assert first["sha256"] == second["sha256"]
    assert first["seq"] == second["seq"]


def test_a_workbook_is_not_offered_as_text(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    """`deliverable-text` decodes; a workbook does not decode.

    Refusing here rather than deep in a codec is what lets the deliverables
    screen ask for a path instead, which is what it wants anyway.
    """
    from foundry.errors import FoundryError

    var = tmp_path / "var"
    task_id, _ = _closed(var, capsys)
    assert main(["--var", str(var), "--task", task_id, "export", "--format", "xlsx"]) == 0
    sha = json.loads(_out(capsys))["sha256"]

    task = TaskStore(var).get(task_id)
    with pytest.raises(FoundryError, match="not text"):
        task.artifact_text(sha, ext=".xlsx")
    assert task.artifact_path(sha, ext=".xlsx").is_file()


def test_the_markdown_export_still_files_alongside_the_workbook(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    """Two exports of one window are two versions, not one overwriting the other.

    They dedupe against their own extension: a workbook and a brief are both
    `kind: export` on the same window, and comparing a workbook's digest to the
    last markdown's would file every alternating export as new.
    """
    var = tmp_path / "var"
    task_id, _ = _closed(var, capsys)
    base = ["--var", str(var), "--task", task_id, "export"]
    assert main([*base, "--format", "xlsx"]) == 0
    capsys.readouterr()
    assert main([*base, "--kind", "brief"]) == 0
    capsys.readouterr()
    assert main([*base, "--format", "xlsx"]) == 0
    again = json.loads(_out(capsys))

    exports = [row for row in TaskStore(var).get(task_id).deliverables() if row["kind"] == "export"]
    assert [row.get("ext", ".md") for row in exports] == [".xlsx", ".md"]
    # The third call matched the first workbook and filed nothing new.
    assert again["seq"] == exports[0]["seq"]
