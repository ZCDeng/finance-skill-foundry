"""Check the numbers an imported report states against the task's own workbooks and calculations.

The method follows a deck/report QC practice: every number in a deliverable should tie to an
identified source. Each number is read with its unit and place in the document, then looked
for among the task's current, authorised workbook cells, table columns and calculation facts.
Deterministic findings (一致, 四舍五入后一致) are kept apart from the ones a person has to judge
(疑似单位换算, 不一致, 需人工核对). Nothing here certifies a report for circulation, edits the
original, or turns the report's text into source truth: a matching number only says where the
same value was found.
"""

from __future__ import annotations

import bisect
import io
import re
from datetime import date, datetime, time
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation, localcontext
from typing import Any

from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, get_column_letter, range_boundaries

from foundry.errors import FoundryError
from foundry.finance_explanations import METRIC_LABELS, _body_numbers, _is_ratio_column
from foundry.financial_evidence import UNIT_SCALE_EXPONENTS
from foundry.report_documents import _markdown_numeric_like, validate_report_document

MAX_NUMBERS = 3000
MAX_CANDIDATES_PER_ITEM = 3
MAX_HITS_PER_QUERY = 60
MAX_SHEET_ROWS = 5000
MAX_SHEET_COLUMNS = 60
HEADER_SCAN_ROWS = 30
CONTEXT_PRECISION = 80

# Unit → power of ten of 元 (or of the ratio, for %).
MONEY_EXPONENTS = {"元": 0, "千元": 3, "万元": 4, "百万元": 6, "亿元": 8, "万": 4, "亿": 8}
PERCENT = "%"
STATUSES = ("match", "match_rounded", "match_scaled", "mismatch", "needs_review")
SCALE_SHIFTS = (0, *sorted({sign * exponent for exponent in UNIT_SCALE_EXPONENTS for sign in (1, -1)}))

# Full-width digits, ％ and the two minus signs read as their ASCII forms. One code point each,
# so offsets into the normalised text are offsets into the original.
_FULL_WIDTH = str.maketrans({**{chr(0xFF10 + digit): str(digit) for digit in range(10)},
                             "％": "%", "−": "-", "－": "-", "．": "."})
_SPACES = "    　"
_TOKEN_UNIT = re.compile(r"[    　]*(亿元|百万元|千元|万元|元|%|万|亿)")
# 「120 万左右」「1.2 亿的收入」 are still amounts; 「3 万人」「2 万件」 count something else.
_SCALE_FOLLOWERS = "左以多余的和或至到及与"
# A superset of structure_proposal.UNIT_IN_TITLE: 千元 and 百万元 are units too.
_UNIT_LINE = re.compile(r"(?:金额)?单位[:：\s]*(亿元|百万元|千元|万元|元)")
_UNIT_IN_LABEL = re.compile(
    r"(?:[（(\[【_／/]|单位[:：]?)\s*(亿元|百万元|千元|万元|元|%)\s*(?:[)）\]】]|$)|(亿元|百万元|千元|万元)$"
)
# 「120 人」「3 家」「15 天」 count something; such a number is never compared with an amount.
_COUNT_UNIT = re.compile(r"[    　]*(平方米|平米|小时|人次|个月|㎡|人|名|位|户|家|个|项|笔|张|台|套|间|次|天|件|辆|吨|倍|年)")
_CLAUSE = re.compile(r"[，；;。\n|]|,(?!\d{3})")
_REFERENCE_BEFORE = ("第", "表", "图", "注", "页", "附件", "章", "节")
_REFERENCE_AFTER = re.compile(r"[  　]*(?:页|季度|季|周|章|节)")
_CODE_HEADER = re.compile(r"编码|代码|编号|序号|科目号|代号|单号|凭证号|行次|^no\.?$|^id$|(?:^|_)(?:code|id)$", re.I)
# A rate's row or column: a 单位：万元 line above it does not make it an amount.
_RATE_WORDS = ("率", "比例", "占比", "比重", "%", "％")
_LEADING_ENUMERATION = re.compile(r"^(?:[一二三四五六七八九十]+[、.．]|\d+[、.．)]|[（(][一二三四五六七八九十\d]+[)）]|其中[:：]?|减[:：]|加[:：])+")
_PARENTHETICAL = re.compile(r"[（(][^（()）]*[)）]")
CHANGE_WORDS = ("同比", "环比", "增长", "增减", "下降", "变动", "增加", "减少", "变化", "增幅", "降幅")
PERIOD_WORDS = ("本月", "当月", "单月", "本期", "上月", "累计", "年初至今", "本年", "全年", "上年", "去年", "同期",
                "预算", "实际", "预测", "差异", "季度")
SCENARIO_WORDS = {
    "actual": "实际", "budget": "预算", "forecast": "预测", "variance": "差异", "prior": "上期",
    "actual_ytd": "累计实际", "budget_ytd": "累计预算", "actual_month": "单月实际", "budget_month": "单月预算",
    "actual_py": "上年同期", "baseline": "对照", "contribution": "差异", "profit_impact": "利润影响",
    "ratio": "变动率", "amount": "金额", "balance": "余额",
}
OCR_HEADING = "OCR 补充草稿"
CLOSING_NOTE = "复核只比对数字，不判断表述是否恰当，也不代表报告可以对外发送。"


# --------------------------------------------------------------------------- numbers in the document

def _plain(nodes: list[dict[str, Any]]) -> str:
    parts = []
    for node in nodes or []:
        if node.get("type") == "text":
            parts.append(str(node.get("text", "")))
        elif node.get("type") == "hardBreak":
            parts.append("\n")
        elif node.get("type") == "paragraph":
            if parts:
                parts.append("\n")
            parts.append(_plain(node.get("content", [])))
    return "".join(parts)


def _int(value: Any, default: int = 1) -> int:
    return value if isinstance(value, int) and not isinstance(value, bool) and value > 0 else default


def _unit_in_label(text: str) -> str | None:
    match = _UNIT_IN_LABEL.search(text or "")
    return (match.group(1) or match.group(2)) if match else None


def _norm_label(text: str) -> str:
    text = re.sub(r"[\s　]+", "", str(text or ""))
    text = _PARENTHETICAL.sub("", text)
    text = _LEADING_ENUMERATION.sub("", text)
    return text.strip("：:·-—")


def _decimal_text(text: str) -> Decimal | None:
    try:
        value = Decimal(text.replace(",", ""))
    except InvalidOperation:
        return None
    return value if value.is_finite() else None


class _Walker:
    """Walks a validated document in order, keeping the heading path and the 单位 line in force."""

    def __init__(self, document: dict[str, Any]):
        self.document = document
        self.formats = {source["id"]: source.get("format") for source in document.get("sources", [])}
        self.metadata_unit = str(document.get("metadata", {}).get("unit") or "").strip() or None
        if self.metadata_unit not in MONEY_EXPONENTS and self.metadata_unit != PERCENT:
            self.metadata_unit = None
        self.headings: list[tuple[int, str]] = []
        self.unit_lines: list[tuple[int, str]] = []
        self.ocr_depth: int | None = None
        self.tables = 0
        self.numbers: list[dict[str, Any]] = []

    def section(self) -> str:
        return " / ".join(text for _, text in self.headings if text)

    def unit_line(self) -> str | None:
        return self.unit_lines[-1][1] if self.unit_lines else None

    def heading(self, level: int, text: str) -> None:
        while self.headings and self.headings[-1][0] >= level:
            self.headings.pop()
        if self.ocr_depth is not None and len(self.headings) < self.ocr_depth:
            self.ocr_depth = None
        self.headings.append((level, text.strip()))
        depth = len(self.headings)
        self.unit_lines = [(at, unit) for at, unit in self.unit_lines if at < depth]
        if text.strip().startswith(OCR_HEADING):
            self.ocr_depth = depth

    def note_unit_line(self, text: str) -> None:
        match = _UNIT_LINE.search(text)
        if match:
            self.unit_lines.append((len(self.headings), match.group(1)))

    def walk(self, nodes: list[dict[str, Any]], inherited: dict[str, Any]) -> None:
        for node in nodes:
            kind = node.get("type")
            attrs = {**inherited, **{key: value for key, value in (node.get("attrs") or {}).items()
                                     if key in {"source_id", "page", "sheet", "range"}}}
            if kind == "heading":
                text = _plain(node.get("content", []))
                self.heading(_int(node.get("attrs", {}).get("level")), text)
                self.scan(text, attrs, "heading")
            elif kind in {"paragraph", "blockquote"}:
                text = _plain(node.get("content", []))
                self.scan(text, attrs, "paragraph")
                self.note_unit_line(text)
            elif kind in {"bulletList", "orderedList"}:
                for item in node.get("content", []):
                    for child in item.get("content", []):
                        if child.get("type") in {"bulletList", "orderedList"}:
                            self.walk([child], attrs)
                        else:
                            self.scan(_plain(child.get("content", [])), {**attrs, **(child.get("attrs") or {})}, "list")
            elif kind == "reportNote":
                self.scan(_plain(node.get("content", [])), attrs, "note")
            elif kind == "table":
                self.table(node, attrs)
            elif kind == "doc":
                self.walk(node.get("content", []), attrs)

    def table(self, node: dict[str, Any], attrs: dict[str, Any]) -> None:
        self.tables += 1
        table_attrs = node.get("attrs") or {}
        caption = str(table_attrs.get("caption") or "")
        self.note_unit_line(caption)
        table_unit = str(table_attrs.get("unit") or "").strip() or _unit_in_label(caption)
        rows = node.get("content", [])
        placed: list[dict[str, Any]] = []
        grid: dict[tuple[int, int], int] = {}
        for row_index, row in enumerate(rows):
            column = 0
            for cell in row.get("content", []):
                while (row_index, column) in grid:
                    column += 1
                cell_attrs = cell.get("attrs") or {}
                colspan = min(_int(cell_attrs.get("colspan")), 100)
                rowspan = min(_int(cell_attrs.get("rowspan")), max(1, len(rows) - row_index))
                index = len(placed)
                placed.append({"row": row_index, "column": column, "cell": cell, "text": _plain(cell.get("content", [])).strip(),
                               "header": cell.get("type") == "tableHeader"})
                for down in range(rowspan):
                    for across in range(colspan):
                        grid[(row_index + down, column + across)] = index
                column += colspan
        header_rows = {item["row"] for item in placed if item["header"]}
        header_rows = {row for row in header_rows if all(item["header"] for item in placed if item["row"] == row)}
        if not header_rows and len(rows) > 1:
            header_rows = {0}
        # A deck or a pasted table often marks only its first row as a header; the word-only rows
        # right under it (「实际」「预算」) still head the columns.
        for row_index in range(max(header_rows, default=-1) + 1, len(rows) - 1):
            texts = [item["text"] for item in placed if item["row"] == row_index and item["text"]]
            if not header_rows or not texts or any(_markdown_numeric_like(text) for text in texts):
                break
            header_rows.add(row_index)
        origin = None
        if table_attrs.get("sheet") and table_attrs.get("range"):
            try:
                min_col, min_row, _max_col, _max_row = range_boundaries(str(table_attrs["range"]))
                origin = (min_col, min_row)
            except (ValueError, TypeError):
                origin = None
        for item in placed:
            row, column = item["row"], item["column"]
            row_label = ""
            for left in range(column):
                other = placed[grid[(row, left)]] if (row, left) in grid else None
                if other and other["text"] and not _markdown_numeric_like(other["text"]):
                    row_label = other["text"]
                    break
            headers = []
            for above in sorted(header_rows):
                if above >= row or (above, column) not in grid:
                    continue
                text = placed[grid[(above, column)]]["text"]
                if text and text not in headers:
                    headers.append(text)
            column_label = " ".join(headers)
            # 序号、科目编码、行次 columns hold identifiers, not amounts.
            if _CODE_HEADER.search(column_label.strip()):
                continue
            location = {**attrs, **{key: value for key, value in (item["cell"].get("attrs") or {}).items()
                                    if key in {"source_id", "page", "sheet", "range"}}}
            if origin is not None:
                location["sheet"] = str(table_attrs["sheet"])
                location["range"] = f"{get_column_letter(origin[0] + column)}{origin[1] + row}"
            fallback = _unit_in_label(row_label) or _unit_in_label(column_label)
            rate = any(word in row_label + column_label for word in _RATE_WORDS)
            if fallback is None and not rate:
                fallback = table_unit
            self.scan(item["text"], location, "table_cell", table={
                "table_index": self.tables, "row": row + 1, "column": column + 1,
                "row_label": row_label, "column_label": column_label, "caption": caption,
                "header": item["header"] or row in header_rows, "unit": fallback, "rate": rate,
            })

    def scan(self, text: str, attrs: dict[str, Any], kind: str, table: dict[str, Any] | None = None) -> None:
        if not text.strip():
            return
        normalized = text.translate(_FULL_WIDTH)
        for found in _body_numbers(normalized):
            number = _read_number(normalized, text, found, kind, table)
            if number is None:
                continue
            number["location"] = self.location(attrs, kind, table)
            number["fallback_unit"] = (table or {}).get("unit")
            rate = bool((table or {}).get("rate"))
            number["unit_line"] = None if rate else self.unit_line()
            number["metadata_unit"] = None if rate else self.metadata_unit
            number["context"] = _context(normalized, number["start"], number["end"], table, self.section())
            number["ocr"] = self.ocr_depth is not None or (
                self.formats.get(str(attrs.get("source_id") or "")) == "pdf" and not attrs.get("page"))
            self.numbers.append(number)

    def location(self, attrs: dict[str, Any], kind: str, table: dict[str, Any] | None) -> dict[str, Any]:
        source_id = attrs.get("source_id")
        location: dict[str, Any] = {
            "source_id": source_id,
            "source_format": self.formats.get(str(source_id or "")),
            "page": attrs.get("page"),
            "section": self.section(),
            "kind": kind,
            "sheet": attrs.get("sheet"),
            "range": attrs.get("range"),
        }
        if table:
            location.update(table_index=table["table_index"], row=table["row"], column=table["column"])
        return {key: value for key, value in location.items() if value not in (None, "")}


def _read_number(normalized: str, original: str, found: dict[str, Any], kind: str,
                 table: dict[str, Any] | None) -> dict[str, Any] | None:
    token, start, end = found["text"], found["start"], found["end"]
    digits = token.lstrip("+-").rstrip("%")
    whole_integer = re.fullmatch(r"\d+", digits) is not None
    before = normalized[:start].rstrip(_SPACES)
    after = normalized[end:]
    value = _decimal_text(token.rstrip("%"))
    if value is None:
        return None
    opening = re.search(r"[(（][ \u00a0\u3000¥￥$]*$", normalized[:start])
    inner_close = re.match(r"[ \u00a0\u3000]*[)）]", after) if opening else None
    unit_from = end + inner_close.end() if inner_close else end
    unit, unit_end = None, unit_from
    if token.endswith("%"):
        unit = PERCENT
    elif match := _TOKEN_UNIT.match(normalized[unit_from:]):
        unit = match.group(1)
        unit_end = unit_from + match.end()
        follower = normalized[unit_end:unit_end + 1]
        if unit in {"万", "亿"} and follower and "\u4e00" <= follower <= "\u9fff" and follower not in _SCALE_FOLLOWERS:
            unit, unit_end = unit + follower, unit_end + 1
        elif unit.endswith("元") and follower in {"/", "每"}:
            unit, unit_end = unit + normalized[unit_end:unit_end + 2], unit_end + 2
    elif whole_integer and 1 <= int(digits) <= 12 and re.match(r"[ \u00a0\u3000]*月", after):
        return None
    elif match := _COUNT_UNIT.match(normalized[unit_from:]):
        unit = match.group(1)
        unit_end = unit_from + match.end()
    amount_unit = unit in MONEY_EXPONENTS or unit == PERCENT
    # 「2026 年度」「FY2026」 name a year; so does a lone 2026 in a heading or header row, unless an
    # amount unit follows it (「8 月营收 2030 万元」 on a slide title is an amount).
    if whole_integer and len(digits) == 4 and 1990 <= int(digits) <= 2100 and (
        before.endswith(("FY", "财年", "年度")) or re.match(r"[ \u00a0\u3000]*(?:年|财年|年度|FY)", after)
        or (not amount_unit and (kind == "heading" or (table and table["header"])))
    ):
        return None
    # 第 3 页、表 2、图 1、注 4：the number names a place, not an amount (「利润表 500 万元」 still is one).
    if unit is None and whole_integer and len(digits) <= 3 and (
        before.endswith(_REFERENCE_BEFORE) or _REFERENCE_AFTER.match(after)
    ):
        return None
    span_start, span_end = start, unit_end
    outer_close = re.match(r"[ \u00a0\u3000]*[)）]", normalized[unit_end:]) if opening and not inner_close else None
    if opening and (inner_close or outer_close):
        whole_cell = kind == "table_cell" and not re.sub(r"[\s(（)）¥￥$,.\d%万亿千百元]", "", normalized)
        # 「（120.5 万元）」「（5.25%）」 in a sentence state an amount in passing: the unit sits inside.
        unit_inside = outer_close is not None or token.endswith("%")
        # （1,234.56）, (1,234.56), （790）万元 and a whole table cell (790 万元) are negatives; an explicit
        # plus sign never is.
        negative = not token.startswith("+") and (
            whole_cell or (not unit_inside and (amount_unit or "," in digits or "." in digits)))
        if negative:
            value = -abs(value)
        elif not unit_inside and not token.startswith("+"):
            return None  # 「（3）收入」「（1）（2）」: an enumeration
        if unit_inside and not negative:
            span_end = end if inner_close else unit_end
        else:
            span_start = opening.start()
            span_end = unit_end if inner_close else unit_end + outer_close.end()
    if token.startswith("-") and value > 0:
        value = -value
    return {
        "start": start,
        "end": end,
        "text": original[span_start:span_end].strip(),
        "raw": value,
        "places": max(0, -value.as_tuple().exponent),
        "suffix_unit": unit,
        "grouped": "," in token,
        "excerpt": _excerpt(original, span_start, span_end),
    }


def _excerpt(text: str, start: int, end: int, width: int = 80) -> str:
    text_length = len(text)
    room = max(0, width - (end - start))
    left = max(0, start - room // 2)
    right = min(text_length, end + (room - (start - left)))
    left = max(0, min(left, right - width))
    piece = text[left:right].replace("\n", " ").strip()
    return ("…" if left > 0 else "") + piece + ("…" if right < text_length else "")


def _context(text: str, start: int, end: int, table: dict[str, Any] | None, section: str) -> dict[str, str]:
    prefix = _CLAUSE.split(text[:start])[-1][-60:]
    suffix = _CLAUSE.split(text[end:], maxsplit=1)[0][:16]
    row_label = (table or {}).get("row_label", "")
    column_label = (table or {}).get("column_label", "")
    caption = (table or {}).get("caption", "")
    return {
        "label_text": f"{row_label} {prefix} {suffix}".strip(),
        "row_label": row_label,
        "period_text": f"{column_label} {prefix} {suffix} {caption} {section}",
    }


def extract_report_numbers(document: dict[str, Any]) -> list[dict[str, Any]]:
    """Every number the document states, with its unit, place and nearby words, in reading order."""
    clean = validate_report_document(document)
    walker = _Walker(clean)
    walker.walk(clean["content"], {})
    for number in walker.numbers:
        suffix = number.pop("suffix_unit")
        if suffix is not None:
            number["unit"], number["unit_source"] = suffix, "suffix"
        elif number.get("fallback_unit"):
            number["unit"], number["unit_source"] = number["fallback_unit"], "table"
        elif number.get("unit_line"):
            number["unit"], number["unit_source"] = number["unit_line"], "unit_line"
        elif number.get("metadata_unit"):
            number["unit"], number["unit_source"] = number["metadata_unit"], "metadata"
        else:
            number["unit"], number["unit_source"] = None, None
        for key in ("fallback_unit", "unit_line", "metadata_unit"):
            number.pop(key, None)
    return walker.numbers


# --------------------------------------------------------------------------- candidates from the task

def _scale(value: Decimal, exponent: int) -> Decimal:
    """value × 10^exponent, exactly (the digits are kept; only the exponent moves)."""
    sign, digits, current = value.as_tuple()
    return Decimal((sign, digits, current + exponent))


def _exponent(unit: str | None) -> int | None:
    if unit in MONEY_EXPONENTS:
        return MONEY_EXPONENTS[unit]
    if unit == PERCENT:
        return -2
    return None


def _is_rate(label: str, header: str) -> bool:
    return any(word in label + header for word in _RATE_WORDS)


def _candidate(kind: str, value: Decimal, unit: str | None, *, order: tuple, label: str = "", header: str = "",
               **fields: Any) -> dict[str, Any]:
    exponent = _exponent(unit)
    return {
        "kind": kind,
        "raw": value,
        "unit": unit,
        "exponent": exponent,
        # A unitless figure on a rate's row or column (毛利率 0.352) is a ratio a report shows as 35.2%.
        "rate": unit is None and _is_rate(label, header),
        # A count or a day figure has a unit of its own: it is compared only with a number that names none.
        "other_unit": unit not in (None, "") and exponent is None,
        "base": _scale(value, exponent) if exponent is not None else None,
        "label": label,
        "header": header,
        "norm_label": _norm_label(label),
        "order": order,
        **{key: value for key, value in fields.items() if value not in (None, "")},
    }


def fact_candidates(facts: list[dict[str, Any]], bindings: dict[str, dict[str, str]]) -> list[dict[str, Any]]:
    found = []
    for index, fact in enumerate(facts):
        value = _decimal_text(str(fact.get("value", "")).rstrip("%"))
        if value is None:
            continue
        unit = fact.get("unit") if isinstance(fact.get("unit"), str) else None
        column = str(fact.get("column") or "")
        # As the fact's own display shows it: a ratio column holds 0.352 for 35.2%, while
        # a _pct column or a row that declares % holds the percentage itself.
        if unit == PERCENT and _is_ratio_column(column) and not column.endswith("_pct"):
            value = _scale(value, 2)
        parts = str(fact.get("label") or "").split(" / ")
        metric = (fact.get("semantics") or {}).get("metric") or parts[0]
        label = METRIC_LABELS.get(str(metric), str(metric))
        header = " ".join(SCENARIO_WORDS.get(part, part) for part in parts[1:] if part)
        locator = (fact.get("source_trace") or {}).get("source_locator") or {}
        sheet, cell = None, None
        excel_range = str(locator.get("excel_range") or "")
        if "!" in excel_range:
            sheet, cell = excel_range.rsplit("!", 1)
            sheet = sheet.strip("'")
        run_bindings = bindings.get(str(fact.get("run_id")), {})
        found.append(_candidate(
            "fact", value, unit, order=(2, str(fact.get("run_id")), int(fact.get("row") or 0), column, index),
            label=label, header=header, fact_id=fact.get("id"), run_id=fact.get("run_id"),
            dataset_id=locator.get("dataset_id") or next(iter(run_bindings.values()), None),
            source_hash=locator.get("source_hash"), sheet=locator.get("sheet") or sheet, range=cell,
            display=fact.get("display_value"),
        ))
    return found


def table_candidates(dataset: dict[str, Any], scope_unit: str | None, scenario: str | None = None) -> list[dict[str, Any]]:
    """DECIMAL columns of a table snapshot, with the sheet cell where the snapshot names one."""
    schema = dataset.get("schema") or []
    decimals = [column["name"] for column in schema if str(column.get("type", "")).startswith("DECIMAL")
                and not _CODE_HEADER.search(str(column["name"]))]
    texts = [column["name"] for column in schema if column.get("type") == "VARCHAR"]
    transformation = dataset.get("transformation") or {}
    header_columns = transformation.get("header_columns") or {}
    row_numbers = transformation.get("source_row_numbers") or []
    sheet = transformation.get("sheet") if transformation.get("format") == "xlsx" else None
    found = []
    for index, row in enumerate(dataset.get("rows") or []):
        label = next((str(row[name]) for name in texts if row.get(name) not in (None, "")), "")
        label = METRIC_LABELS.get(label, label)
        for name in decimals:
            value = _decimal_text(str(row.get(name))) if row.get(name) not in (None, "") else None
            if value is None:
                continue
            cell = None
            if sheet and name in header_columns and index < len(row_numbers):
                cell = f"{header_columns[name]}{row_numbers[index]}"
            unit = _unit_in_label(name)
            if unit is None and not _is_rate(label, name):
                unit = scope_unit
            # The context's scenario (实际/预算) tells two same-named lines of two snapshots apart.
            header = " ".join(part for part in (SCENARIO_WORDS.get(scenario or "", ""), SCENARIO_WORDS.get(name, name)) if part)
            found.append(_candidate(
                "table", value, unit, order=(0, dataset["id"], sheet or "", index, name),
                label=label, header=header, dataset_id=dataset["id"], source_hash=dataset.get("source_hash"),
                sheet=sheet if cell else None, range=cell, row=None if cell else index + 1,
                column=None if cell else name,
            ))
    return found


def workbook_candidates(dataset: dict[str, Any], data: bytes, notes: list[str]) -> list[dict[str, Any]]:
    """Every numeric cell of an archived workbook, labelled by its row's first text and the text above it."""
    from openpyxl import load_workbook

    from foundry.analysis_data import XLSX_SHORTEST_DOUBLE, archived_numeric_parser, read_exact_numerics
    from foundry.finance_controls import _leading_text

    name = dataset.get("name") or dataset["id"]
    found: list[dict[str, Any]] = []
    try:
        parser = archived_numeric_parser(dataset, data)
    except FoundryError:
        parser = XLSX_SHORTEST_DOUBLE
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        for sheet in book.sheetnames:
            page = book[sheet]
            if (page.max_row or 0) > MAX_SHEET_ROWS or (page.max_column or 0) > MAX_SHEET_COLUMNS:
                notes.append(f"{name} 的工作表「{sheet}」超过 {MAX_SHEET_ROWS} 行或 {MAX_SHEET_COLUMNS} 列，未纳入比对。")
                continue
            try:
                exact = read_exact_numerics(data, sheet, parser)
            except FoundryError:
                notes.append(f"{name} 的工作表「{sheet}」有公式没有缓存值或有错误值，未纳入比对；在 Excel 中重新计算并保存后再导入。")
                continue
            rows, formats = [], []
            for row in page.iter_rows():
                rows.append([cell.value for cell in row])
                formats.append([getattr(cell, "number_format", None) for cell in row])
            sheet_unit = None
            for row in rows[:10]:
                for value in row:
                    if isinstance(value, str) and (match := _UNIT_LINE.search(value)):
                        sheet_unit = match.group(1)
                        break
                if sheet_unit:
                    break
            for coordinate, value in exact.items():
                letters, row_number = coordinate_from_string(coordinate)
                column = column_index_from_string(letters)
                row_values = rows[row_number - 1] if row_number - 1 < len(rows) else []
                shown = row_values[column - 1] if column - 1 < len(row_values) else None
                if isinstance(shown, (datetime, date, time)):
                    continue
                label = _leading_text(row_values)
                header = ""
                for above in range(row_number - 2, max(-1, row_number - 2 - HEADER_SCAN_ROWS), -1):
                    cells = rows[above] if above < len(rows) else []
                    text = cells[column - 1] if column - 1 < len(cells) else None
                    if isinstance(text, str) and text.strip():
                        header = re.sub(r"[\s　]+", "", text)
                        break
                if _CODE_HEADER.search(header):
                    continue
                shown_as = formats[row_number - 1] if row_number - 1 < len(formats) else []
                if _percent_format(shown_as[column - 1] if column - 1 < len(shown_as) else None):
                    # Excel shows 0.352 formatted as 0.0% as 35.2%: that is the figure a report quotes.
                    value, unit = _scale(value, 2), PERCENT
                else:
                    unit = _unit_in_label(label) or _unit_in_label(header)
                    if unit is None and not _is_rate(label, header):
                        unit = sheet_unit
                found.append(_candidate(
                    "workbook", value, unit, order=(1, dataset["id"], sheet, row_number, column),
                    label=label, header=header, dataset_id=dataset["id"], source_hash=dataset.get("source_hash"),
                    sheet=sheet, range=coordinate,
                ))
    finally:
        book.close()
    return found


def _percent_format(number_format: Any) -> bool:
    """Whether an Excel number format shows the value ×100 with a % sign (a quoted or escaped % does not)."""
    if not isinstance(number_format, str):
        return False
    return "%" in re.sub(r'"[^"]*"|\\.', "", number_format)


def gather_candidates(store: Any) -> tuple[list[dict[str, Any]], list[str]]:
    """The values a report may tie to: current, authorised snapshots and the runs that read only those."""
    from foundry.finance_controls import snapshot_bytes, superseded_snapshots
    from foundry.finance_explanations import prepare_explanation

    analysis = store.analysis
    state = analysis.state()
    datasets = state["datasets"]
    replaced = superseded_snapshots(store.task, datasets)
    # A table read from a replaced workbook is named after its sheet, so one the new version renamed is
    # not replaced by name; it goes with its workbook (as the month-close workbook list leaves it out).
    replaced |= {row["id"] for row in datasets if (row.get("metadata") or {}).get("derived_from") in replaced}
    published = {row["id"] for row in datasets}
    current = [row for row in datasets if row["id"] not in replaced
               and (store.allowed is None or row["id"] in store.allowed)]
    notes: list[str] = []
    if replaced:
        notes.append(f"{len(replaced)} 份已被新版本取代的资料没有参与比对。")
    scope_units: dict[str, set[str]] = {}
    scope_scenarios: dict[str, set[str]] = {}
    for context in store._list("contexts"):
        for alias, dataset_id in (context.get("bindings") or {}).items():
            scope = (context.get("scopes") or {}).get(alias) or {}
            if isinstance(scope.get("unit"), str) and scope["unit"] in MONEY_EXPONENTS:
                scope_units.setdefault(dataset_id, set()).add(scope["unit"])
            if isinstance(scope.get("scenario"), str) and scope["scenario"] in SCENARIO_WORDS:
                scope_scenarios.setdefault(dataset_id, set()).add(scope["scenario"])

    found: list[dict[str, Any]] = []
    seen_cells: set[tuple[str, str, str]] = set()
    for dataset in current:
        if dataset.get("kind") != "table":
            continue
        try:
            record = analysis.read_dataset(dataset["id"])
        except FoundryError:
            notes.append(f"{dataset.get('name') or dataset['id']} 读取失败，未纳入比对。")
            continue
        units = scope_units.get(dataset["id"], set())
        scenarios = scope_scenarios.get(dataset["id"], set())
        for candidate in table_candidates(record, next(iter(units)) if len(units) == 1 else None,
                                          next(iter(scenarios)) if len(scenarios) == 1 else None):
            if candidate.get("sheet") and candidate.get("range"):
                seen_cells.add((str(candidate.get("source_hash")), candidate["sheet"], candidate["range"]))
            found.append(candidate)
    workbook_hashes: set[str] = set()
    for dataset in current:
        if not str(dataset.get("source_file") or "").endswith(".xlsx") or dataset["source_hash"] in workbook_hashes:
            continue
        workbook_hashes.add(dataset["source_hash"])
        try:
            record = analysis.read_dataset(dataset["id"])
            data = snapshot_bytes(analysis, record)
        except FoundryError:
            notes.append(f"{dataset.get('name') or dataset['id']} 读取失败，未纳入比对。")
            continue
        for candidate in workbook_candidates(record, data, notes):
            key = (str(candidate.get("source_hash")), candidate["sheet"], candidate["range"])
            if key not in seen_cells:
                seen_cells.add(key)
                found.append(candidate)

    latest: dict[str, dict[str, Any]] = {}
    for run in state["runs"]:
        bindings = run.get("bindings") or {}
        if run.get("execution_status") != "succeeded" or not bindings:
            continue
        ids = set(bindings.values())
        if not ids <= published or ids & replaced or (store.allowed is not None and not ids <= store.allowed):
            continue
        key = str(run.get("analysis_id") or run["id"])
        if key not in latest or (latest[key].get("seq", 0), latest[key]["created_at"]) < (run.get("seq", 0), run["created_at"]):
            latest[key] = run
    failed = 0
    for run in sorted(latest.values(), key=lambda row: (row.get("seq", 0), row["id"])):
        try:
            facts = prepare_explanation(analysis, [run], backgrounds=[], materiality={})["facts"]
        except FoundryError:
            failed += 1
            continue
        found.extend(fact_candidates(facts, {run["id"]: run.get("bindings") or {}}))
    if failed:
        notes.append(f"{failed} 次计算的结果读取失败，未纳入比对。")
    found.sort(key=lambda candidate: candidate["order"])
    return found, notes


# --------------------------------------------------------------------------- classification

class _Index:
    def __init__(self, candidates: list[dict[str, Any]]):
        self.candidates = candidates
        self.sorted: dict[str, tuple[list[Decimal], list[dict[str, Any]]]] = {}
        groups = {
            "base": [(c["base"], c) for c in candidates if c["base"] is not None],
            "raw_unknown": [(c["raw"], c) for c in candidates if c["exponent"] is None and not c["other_unit"]],
            "raw_all": [(c["raw"], c) for c in candidates],
        }
        for name, pairs in groups.items():
            pairs.sort(key=lambda pair: pair[0])
            self.sorted[name] = ([pair[0] for pair in pairs], [pair[1] for pair in pairs])
        self.by_label: dict[str, list[dict[str, Any]]] = {}
        self.label_bigrams: dict[str, set[str]] = {}
        for candidate in candidates:
            label = candidate["norm_label"]
            if len(label) < 2:
                continue
            self.by_label.setdefault(label, []).append(candidate)
            for offset in range(len(label) - 1):
                self.label_bigrams.setdefault(label[offset:offset + 2], set()).add(label)

    def between(self, group: str, low: Decimal, high: Decimal) -> list[dict[str, Any]]:
        keys, values = self.sorted[group]
        start = bisect.bisect_left(keys, low)
        end = bisect.bisect_right(keys, high)
        return values[start:min(end, start + MAX_HITS_PER_QUERY)]

    def labels_in(self, context: dict[str, str]) -> set[str]:
        """Candidate labels the number's own words name: inside its clause or row label, or containing its row label."""
        text = re.sub(r"[\s　]+", "", context["label_text"])
        hits: set[str] = set()
        for offset in range(len(text) - 1):
            for label in self.label_bigrams.get(text[offset:offset + 2], ()):
                if text.startswith(label, offset):
                    hits.add(label)
        row_label = _norm_label(context["row_label"])
        # A row named exactly like a source line is that line: 「收入」 is not 「其他收入」.
        exact_row = len(row_label) >= 2 and row_label in self.by_label
        if len(row_label) >= 2 and not exact_row:
            for label in self.label_bigrams.get(row_label[:2], ()):
                if row_label in label:
                    hits.add(label)
        # 「营业收入」 names a line; 「收入」 inside it is not a second, competing name.
        return {label for label in hits
                if (exact_row and label == row_label) or not any(label != other and label in other for other in hits)}


def _words(text: str, vocabulary) -> set[str]:
    return {word for word in vocabulary if word in text}


def _score(number: dict[str, Any], candidate: dict[str, Any], labels: set[str]) -> tuple[int, bool]:
    """How well a candidate's words fit the number's, and whether they name the same line and period."""
    context = number["context"]
    theirs_text = candidate["label"] + candidate["header"]
    score, label_hit = 0, bool(candidate["norm_label"] and candidate["norm_label"] in labels)
    if label_hit:
        score += 3
    metrics = _words(context["label_text"], METRIC_LABELS.values()) & _words(theirs_text, METRIC_LABELS.values())
    if metrics:
        score += 1
    mine = _words(context["period_text"], PERIOD_WORDS)
    theirs = _words(theirs_text, PERIOD_WORDS)
    shared = mine & theirs
    period_conflict = bool(mine and theirs and not shared)
    score += min(2, len(shared)) - (1 if period_conflict else 0)
    # A change (同比、增减) is not the level it changes, and the other way round.
    change_conflict = bool(_words(context["period_text"], CHANGE_WORDS)) != bool(_words(theirs_text, CHANGE_WORDS))
    if change_conflict:
        score -= 1
    # A shared metric word (收入) stands in for a line name only when the number's words name no line.
    strong = (label_hit or bool(metrics and shared and not labels)) and not period_conflict and not change_conflict
    return score, strong


def _comparable(number: dict[str, Any], candidate: dict[str, Any]) -> bool:
    """Whether the two could be the same kind of figure: an amount with an amount, a rate with a rate."""
    mine = number["unit"]
    if mine is None or candidate["unit"] is None:
        return not (mine is not None and candidate["other_unit"])
    return (mine == PERCENT) == (candidate["unit"] == PERCENT) and not candidate["other_unit"]


def _views(number: dict[str, Any]) -> list[tuple[str, int, str]]:
    """Ways to compare: (candidate group, power of ten that turns its value into the number's unit, assumption)."""
    exponent = _exponent(number["unit"])
    if exponent is None:
        return [("raw_all", 0, "no_unit")]
    views = [("base", exponent, ""), ("raw_unknown", 0, "source_no_unit")]
    if number["unit"] == PERCENT:
        views.append(("raw_unknown", -2, "source_ratio"))
    return views


def _rounds_to(shown: Decimal, raw: Decimal) -> bool:
    """Whether the source value, rounded half up (away from zero, as Excel shows it) to the number's places, is the number."""
    exponent = raw.as_tuple().exponent
    if not isinstance(exponent, int) or shown.as_tuple().exponent >= exponent:
        return False
    try:
        return shown.quantize(Decimal((0, (1,), exponent)), rounding=ROUND_HALF_UP) == raw
    except InvalidOperation:
        return False


def _in_group(candidate: dict[str, Any], group: str) -> bool:
    if group == "base":
        return candidate["base"] is not None
    if group == "raw_unknown":
        return candidate["exponent"] is None and not candidate["other_unit"]
    return True


def _hits(number: dict[str, Any], index: _Index, named: list[dict[str, Any]] = ()) -> list[dict[str, Any]]:
    """Candidates whose value is the number's: from the value index, and among `named` (the lines its words
    name), which are tested one by one so a crowd of equal values cannot push them out of the capped hits."""
    raw, places = number["raw"], number["places"]
    half = Decimal((0, (5,), -places - 1))
    hits: dict[int, dict[str, Any]] = {}

    def test(candidate: dict[str, Any], group: str, exponent: int, shift: int, assumption: str) -> None:
        shown = _scale(candidate["base"] if group == "base" else candidate["raw"], shift - exponent)
        if shown == raw:
            how = "exact"
        elif _rounds_to(shown, raw):
            how = "rounded"
        else:
            return
        kind = "scaled" if shift else how
        previous = hits.get(id(candidate))
        rank = ("exact", "rounded", "scaled").index(kind)
        if previous is None or rank < previous["rank"]:
            hits[id(candidate)] = {"candidate": candidate, "how": kind, "rounded": how == "rounded",
                                   "shift": shift, "assumption": assumption, "rank": rank}

    with localcontext() as context:
        context.prec = CONTEXT_PRECISION
        for group, exponent, assumption in _views(number):
            for shift in SCALE_SHIFTS:
                low = _scale(raw - half, exponent - shift)
                high = _scale(raw + half, exponent - shift)
                if low > high:
                    low, high = high, low
                for candidate in index.between(group, low, high):
                    test(candidate, group, exponent, shift, assumption)
                for candidate in named:
                    if _in_group(candidate, group):
                        test(candidate, group, exponent, shift, assumption)
    return list(hits.values())


def _in_number_unit(number: dict[str, Any], candidate: dict[str, Any]) -> tuple[Decimal, str | None]:
    """The candidate's value in the unit the number is written in, and how it got there.

    "same" or "unit" (converted between known units), "ratio" (a fraction on a rate's row, ×100 for a
    percentage), or None when the units cannot be tied together: then the value is the source's own.
    """
    mine, theirs = _exponent(number["unit"]), candidate["exponent"]
    if mine is not None and theirs is not None:
        with localcontext() as context:
            context.prec = CONTEXT_PRECISION
            return _scale(candidate["base"], -mine), "same" if mine == theirs else "unit"
    # 0.352 on a 毛利率 row is 35.2%; a 35.2 there may already be the percentage, so it stays unknown.
    if number["unit"] == PERCENT and candidate["unit"] is None and candidate.get("rate") and abs(candidate["raw"]) <= 1:
        return _scale(candidate["raw"], 2), "ratio"
    return candidate["raw"], None


def _public_candidate(candidate: dict[str, Any]) -> dict[str, Any]:
    keys = ("kind", "dataset_id", "source_hash", "sheet", "range", "row", "column", "fact_id", "run_id", "display")
    public = {key: candidate[key] for key in keys if candidate.get(key) not in (None, "")}
    label = " · ".join(part for part in (candidate["label"], candidate["header"]) if part)
    if label:
        public["label"] = label
    public["value"] = format(candidate["raw"], "f")
    public["unit"] = candidate["unit"]
    return public


def _place(location: dict[str, Any]) -> str:
    parts = []
    if location.get("page"):
        parts.append(f"第 {location['page']} 页")
    if location.get("kind") == "table_cell":
        parts.append(f"表 {location.get('table_index')} 第 {location.get('row')} 行第 {location.get('column')} 列")
    elif location.get("kind") == "note":
        parts.append("附注")
    elif location.get("section") and not location.get("page"):
        parts.append(f"「{location['section'].split(' / ')[-1]}」一节")
    return " ".join(parts) or "正文"


def _source_words(candidate: dict[str, Any]) -> str:
    where = f"{candidate['sheet']}!{candidate['range']}" if candidate.get("sheet") and candidate.get("range") else ""
    kind = {"workbook": "套表", "table": "资料表", "fact": "计算结果"}[candidate["kind"]]
    label = " · ".join(part for part in (candidate["label"], candidate["header"]) if part)
    return " ".join(part for part in (kind, where, f"「{label}」" if label else "") if part)


def _value_words(value: Decimal, unit: str | None) -> str:
    text = format(value, "f")
    sign = "-" if text.startswith("-") else ""
    whole, dot, fraction = text.lstrip("-").partition(".")
    grouped = f"{int(whole):,}" if whole.isdigit() else whole
    return f"{sign}{grouped}{dot}{fraction}" + (f" {unit}" if unit and unit != PERCENT else unit or "")


def _gap_words(difference: Decimal, unit: str | None) -> str:
    """「文稿比来源多 20 万元」; for a percentage the gap is in 个百分点, not a relative change."""
    more = difference > 0
    if unit == PERCENT:
        return f"文稿比来源{'高' if more else '低'} {_value_words(abs(difference), None)} 个百分点"
    return f"文稿比来源{'多' if more else '少'} {_value_words(abs(difference), unit)}"


def _compared_words(outcome: dict[str, Any], unit: str | None) -> str:
    source = outcome["counterpart"]
    shown = _value_words(outcome["compared"], unit)
    if outcome["basis"] == "ratio":
        return f"{_value_words(source['raw'], None)}（比例值，即 {shown}）"
    if outcome["basis"] == "unit":
        return f"{_value_words(source['raw'], source['unit'])}（折成文稿单位是 {shown}）"
    return shown


ASSUMPTION_NOTES = {
    "source_no_unit": "来源没有写明单位，按数字相同判断。",
    "source_ratio": "来源以比例值保存，已按百分比折算。",
    "no_unit": "文稿没有写明单位，按数字相同判断。",
}


def _classify(number: dict[str, Any], index: _Index) -> dict[str, Any]:
    result: dict[str, Any] = {"status": "needs_review", "note": "", "candidates": []}
    if number["unit"] is not None and _exponent(number["unit"]) is None:
        result["note"] = f"单位是「{number['unit']}」，不是金额或比例，没有比对。"
        return result
    labels = index.labels_in(number["context"])
    labelled = [candidate for label in sorted(labels) for candidate in index.by_label.get(label, [])
                if _comparable(number, candidate)]
    strong_labelled = []
    for candidate in labelled:
        score, strong = _score(number, candidate, labels)
        if strong:
            strong_labelled.append((score, candidate))
    hits = [hit for hit in _hits(number, index, [candidate for _, candidate in strong_labelled])
            if _comparable(number, hit["candidate"])]
    for hit in hits:
        hit["score"], hit["strong"] = _score(number, hit["candidate"], labels)
    weak = len(number["raw"].as_tuple().digits) <= 2

    def ranked(rows):
        return sorted(rows, key=lambda hit: (-hit["score"], hit["rank"], hit["candidate"]["order"]))

    for how, status in (("exact", "match"), ("rounded", "match_rounded"), ("scaled", "match_scaled")):
        group = ranked([hit for hit in hits if hit["how"] == how])
        if not group:
            continue
        best = group[0]
        result["candidates"] = [hit["candidate"] for hit in group[:MAX_CANDIDATES_PER_ITEM]]
        if strong_labelled and not any(hit["strong"] for hit in group):
            result["note"] = "这个数字能在别处找到，但与同名项目的数值不同，需要人工确认对应哪一项。"
            result["candidates"] = [*[candidate for _, candidate in sorted(strong_labelled, key=lambda row: (-row[0], row[1]["order"]))][:1],
                                    *result["candidates"]][:MAX_CANDIDATES_PER_ITEM]
            return result
        if best["score"] <= 0 and weak:
            result["note"] = "数字位数太少，只凭数值不能确认对应来源；请对照上下文核对。"
            return result
        if how == "scaled" and number["unit"] is None:
            result["note"] = f"文稿没有写明单位；数值与来源相差 {10 ** abs(best['shift'])} 倍，无法判断是否单位换算。"
            return result
        values = {format(hit["candidate"]["raw"], "f") for hit in group if hit["score"] == best["score"]}
        if how != "exact" and len(values) > 1 and best["score"] <= 0:
            result["note"] = "有多个来源的数值都接近，无法确定对应哪一项。"
            return result
        result["status"] = status
        notes = []
        if how == "rounded":
            notes.append(f"按文稿显示的 {number['places']} 位小数四舍五入后一致。")
        if how == "scaled":
            notes.append(f"与来源相差 {10 ** abs(best['shift'])} 倍，疑似单位换算（如元与万元），请核对单位。")
            if best["rounded"]:
                notes.append("换算后还需四舍五入才一致。")
        if best["assumption"]:
            notes.append(ASSUMPTION_NOTES[best["assumption"]])
        result["note"] = "".join(notes)
        return result
    if strong_labelled:
        top = max(score for score, _ in strong_labelled)
        tops = [candidate for score, candidate in sorted(strong_labelled, key=lambda row: (-row[0], row[1]["order"])) if score == top]
        converted = {format(_in_number_unit(number, candidate)[0], "f") for candidate in tops}
        result["candidates"] = tops[:MAX_CANDIDATES_PER_ITEM]
        if len(converted) > 1:
            result["note"] = "有多个同名项目，数值各不相同，需要人工确认对应哪一项。"
            return result
        counterpart = tops[0]
        value, basis = _in_number_unit(number, counterpart)
        mine = _words(number["context"]["period_text"], PERIOD_WORDS)
        theirs = _words(counterpart["label"] + counterpart["header"], PERIOD_WORDS)
        # 不一致 is a finding only when the unit and the period tie the two together; otherwise the
        # difference could be a unit or a scope, and guessing either is not this check's to do.
        doubts = []
        if basis is None:
            doubts.append("来源没有标明是比例值还是百分数" if number["unit"] == PERCENT and counterpart["unit"] is None
                          else "有一方没有写明单位，无法换算到同一单位")
        if not mine & theirs:
            doubts.append("无法确认是同一期间（如本月、累计、预算）")
        if doubts:
            result["note"] = f"来源里有同名项目，数值不同；但{'，也'.join(doubts)}，请人工核对。"
            return result
        with localcontext() as context:
            context.prec = CONTEXT_PRECISION
            difference = number["raw"] - value
        result.update(status="mismatch", counterpart=counterpart, compared=value, difference=difference, basis=basis)
        result["note"] = ASSUMPTION_NOTES["source_ratio"] if basis == "ratio" else ""
        return result
    result["note"] = "没有在当前资料和计算结果中找到这个数字。"
    if labelled:
        result["candidates"] = labelled[:MAX_CANDIDATES_PER_ITEM]
    return result


def review_numbers(numbers: list[dict[str, Any]], candidates: list[dict[str, Any]], notes: list[str]) -> dict[str, Any]:
    index = _Index(candidates)
    truncated = len(numbers) > MAX_NUMBERS
    items: list[dict[str, Any]] = []
    issues: list[dict[str, Any]] = []
    for order, number in enumerate(numbers[:MAX_NUMBERS], start=1):
        outcome = _classify(number, index)
        note = outcome["note"]
        if number["ocr"]:
            if outcome["status"] != "needs_review":
                note = (note + " " if note else "") + f"（按数值可对应：{_source_words(outcome['candidates'][0])}）"
            outcome["status"] = "needs_review"
            note = "OCR 识别的文字没有页码，需对照原 PDF 核对。" + (f" {note}" if note else "")
        amount_like = number["unit"] in MONEY_EXPONENTS or number["unit"] == PERCENT or number["grouped"] or number["places"] > 0
        item: dict[str, Any] = {
            "id": f"n{order}",
            "text": number["text"],
            "value": format(number["raw"], "f"),
            "unit": number["unit"],
            "unit_source": number["unit_source"],
            "location": number["location"],
            "excerpt": number["excerpt"],
            "status": outcome["status"],
            "candidates": [_public_candidate(candidate) for candidate in outcome["candidates"]],
            "note": note,
            "amount_like": amount_like,
        }
        if outcome["status"] == "mismatch":
            item["compared_value"] = format(outcome["compared"], "f")
            item["compared_basis"] = outcome["basis"]
            item["difference"] = format(outcome["difference"], "f")
        items.append(item)
        issue = _issue(item, number, outcome)
        if issue:
            issues.append(issue)
    order_of = {item["id"]: position for position, item in enumerate(items)}
    # Numbers that do not look like amounts and found nothing still count; they are listed after the rest.
    items.sort(key=lambda item: (item["status"] == "needs_review" and not item["amount_like"], order_of[item["id"]]))
    counts = {status: sum(1 for item in items if item["status"] == status) for status in STATUSES}
    attention = counts["mismatch"] + counts["match_scaled"] + counts["needs_review"]
    summary = {
        # Numbers past the cap were never checked: a cut-off review cannot pass.
        "status": "pass" if items and not attention and not truncated else "needs_attention",
        "total": len(items),
        **counts,
        "needs_review_not_amount": sum(1 for item in items if item["status"] == "needs_review" and not item["amount_like"]),
    }
    review_notes = list(dict.fromkeys(notes))
    if truncated:
        summary["truncated"] = True
        review_notes.append(f"文稿里的数字超过 {MAX_NUMBERS} 个，只复核了前 {MAX_NUMBERS} 个。")
    if not items:
        review_notes.append("文稿里没有找到可比对的数字。")
    if not candidates:
        review_notes.append("当前任务没有可比对的套表、资料表或计算结果。")
    return {"schema_version": 1, "summary": summary, "items": items, "issues": issues,
            "notes": review_notes, "closing_note": CLOSING_NOTE}


def _issue(item: dict[str, Any], number: dict[str, Any], outcome: dict[str, Any]) -> dict[str, Any] | None:
    status = item["status"]
    if status in {"match", "match_rounded"}:
        return None
    place = _place(item["location"])
    best = outcome.get("counterpart") or (outcome["candidates"][0] if outcome["candidates"] else None)
    if status == "mismatch":
        code, severity = "report_review_mismatch", "error"
        message = (f"{place}写 {item['text']}，{_source_words(best)} 为 {_compared_words(outcome, number['unit'])}，"
                   f"{_gap_words(outcome['difference'], number['unit'])}。")
    elif status == "match_scaled":
        code, severity = "report_review_scaled", "warning"
        message = f"{place}写 {item['text']}，{_source_words(best)} 为 {_value_words(best['raw'], best['unit'])}。{item['note']}"
    else:
        code, severity = "report_review_unmatched", "warning"
        message = f"{place}写 {item['text']}：{item['note']}"
    issue: dict[str, Any] = {
        "code": code,
        "severity": severity,
        "message": message,
        "business_key": item["id"],
        "label": place,
        "excerpt": item["excerpt"],
        "original_text": item["text"],
    }
    if best and best.get("dataset_id") and best.get("sheet") and best.get("range"):
        issue["location"] = {key: best[key] for key in ("dataset_id", "sheet", "range", "source_hash") if best.get(key)}
    return issue


def review_report(document: dict[str, Any], candidates: list[dict[str, Any]], notes: list[str]) -> dict[str, Any]:
    return review_numbers(extract_report_numbers(document), candidates, notes)
