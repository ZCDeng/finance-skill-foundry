"""Roll last month's workbook forward into a copy for the new month.

The manager says where each of this month's figures goes (「来源到单元格」): one
cell or a summed range of a task snapshot, one cell of a calculation result, or
a value she types. Every target is read before anything is written, and only
plain input cells of a copy are written. A formula, the hidden part of a merged
cell, a member of an array formula, a locked cell on a protected sheet and an
error value are listed instead, with the figure they should have held, so she
updates them in Excel herself. The archived snapshot is never opened for writing.

openpyxl keeps formulas when it saves but drops their cached results, and it
cannot calculate. The copy asks Excel to recalculate when it is opened; until it
has been opened and saved there once, nothing here can read a formula result in
it. A workbook with parts openpyxl would drop (charts, pictures, external links,
macros) gets no copy at all, only the list.
"""

from __future__ import annotations

import io
import re
import uuid
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation, localcontext
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, MergedCell
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from foundry.analysis_data import _open_xlsx_archive, read_exact_numerics, xlsx_number_reader
from foundry.errors import FoundryError
from foundry.sheet import survives_a_cell
from foundry.workbook_bindings import FORMULA_LIKE, _open

MAX_MAPPINGS = 2000
# The same bound a tie-out control reads; a sum wider than this is not one figure.
MAX_SUM_CELLS = 20_000
MAX_TITLE_HINTS = 50
PERIOD = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
A1 = re.compile(r"^\$?([A-Za-z]{1,3})\$?([1-9][0-9]{0,6})$")
A1_RANGE = re.compile(r"^\$?[A-Za-z]{1,3}\$?[1-9][0-9]{0,6}(?::\$?[A-Za-z]{1,3}\$?[1-9][0-9]{0,6})?$")
# What a typed value must look like to be written as a number. Anything else,
# 「1,200」 and 「0123」 included, is text exactly as typed.
PLAIN_DECIMAL = re.compile(r"^-?(?:0|[1-9]\d*)(?:\.\d+)?$")
NUMERIC_RESULT_TYPE = re.compile(
    r"^(DECIMAL|NUMERIC|HUGEINT|UHUGEINT|BIGINT|UBIGINT|INTEGER|UINTEGER|SMALLINT|USMALLINT|TINYINT|UTINYINT|DOUBLE|FLOAT|REAL|INT\d*)\b",
    re.I,
)
ID = re.compile(r"^[0-9a-f]{32}$")
SOURCE_KINDS = frozenset({"dataset", "run", "value"})
WRITE, BLOCKED, UNCHANGED = "write", "blocked", "unchanged"
EXCEL_MAX_ROW, EXCEL_MAX_COLUMN = 1_048_576, 16_384

REASONS = {
    "bad_address": "目标要写成「工作表!单元格」，例如 利润表!B5，而且只能是一个单元格",
    "missing_sheet": "工作簿里没有这张工作表",
    "duplicate_target": "同一个单元格映射了不止一次，这几行都不写",
    "merged_non_anchor": "目标在合并单元格里，但不是左上角那一格；只能写左上角",
    "array_formula_member": "目标在数组公式的区域里，不能单独改",
    "formula_cell": "目标是公式单元格，公式不会被覆盖",
    "protected_cell": "工作表受保护，这一格是锁定的",
    "error_cell": "目标单元格是错误值（例如 #N/A），请先在 Excel 里处理",
    "source_unreadable": "来源读不出",
}


class SourceUnreadable(Exception):
    """A mapping whose source cannot give one figure; the row is listed, not written."""


@dataclass(frozen=True)
class SourceValue:
    number: Decimal | None = None
    text: str | None = None


# ---------------------------------------------------------------- the request


def period_of(value: Any) -> str:
    """The month the copy is for, as the manager typed it. Never the clock."""
    if not isinstance(value, str) or not PERIOD.fullmatch(value.strip()):
        raise FoundryError("rollforward period must be YYYY-MM")
    return value.strip()


def _text(value: Any, label: str, maximum: int) -> str:
    if value is None:
        return ""
    if not isinstance(value, str) or len(value) > maximum:
        raise FoundryError(f"rollforward {label} must be text (at most {maximum} characters)")
    return value


def normalize_mappings(raw: Any) -> list[dict[str, Any]]:
    """The mapping rows as they will be filed. Shape errors refuse; content problems become blocked rows."""
    if not isinstance(raw, list) or not raw or len(raw) > MAX_MAPPINGS:
        raise FoundryError(f"rollforward mappings must be a list of 1 to {MAX_MAPPINGS} rows")
    rows = []
    for index, item in enumerate(raw, start=1):
        if not isinstance(item, dict) or not isinstance(item.get("target"), dict) or not isinstance(item.get("source"), dict):
            raise FoundryError(f"rollforward mapping {index} needs a target object and a source object")
        target = {
            "sheet": _text(item["target"].get("sheet"), "target sheet", 200).strip(),
            "cell": _text(item["target"].get("cell"), "target cell", 40).strip(),
        }
        rows.append({"target": target, "source": _source(item["source"], index)})
    return rows


def _source(source: dict[str, Any], index: int) -> dict[str, Any]:
    kind = source.get("kind")
    if kind not in SOURCE_KINDS:
        raise FoundryError(f"rollforward mapping {index} source kind must be dataset, run or value")
    if kind == "dataset":
        return {
            "kind": "dataset",
            "dataset_id": _text(source.get("dataset_id"), "source dataset_id", 64).strip(),
            "source_hash": _text(source.get("source_hash"), "source source_hash", 64).strip(),
            "sheet": _text(source.get("sheet"), "source sheet", 200).strip(),
            "range": _text(source.get("range"), "source range", 40).strip(),
        }
    if kind == "run":
        row = source.get("row")
        if isinstance(row, bool) or not isinstance(row, int) or row < 0:
            raise FoundryError(f"rollforward mapping {index} run row must be a zero-based integer")
        return {
            "kind": "run",
            "run_id": _text(source.get("run_id"), "source run_id", 64).strip(),
            "row": row,
            "column": _text(source.get("column"), "source column", 200),
        }
    # Excel holds at most 32,767 characters in a cell.
    return {
        "kind": "value",
        "manager_value": _text(source.get("value", source.get("manager_value")), "manager value", 32767),
        "note": _text(source.get("note"), "manager note", 500).strip(),
    }


def cell_address(text: str) -> str | None:
    """One A1 cell, upper-cased and without $, or None when it is anything else."""
    match = A1.fullmatch(text or "")
    if not match:
        return None
    column, row = match.group(1).upper(), int(match.group(2))
    if column_index_from_string(column) > EXCEL_MAX_COLUMN or row > EXCEL_MAX_ROW:
        return None
    return f"{column}{row}"


def range_bounds(text: str) -> tuple[int, int, int, int] | None:
    """(min_col, min_row, max_col, max_row) of a bounded A1 rectangle, or None."""
    if not A1_RANGE.fullmatch(text or ""):
        return None
    try:
        min_col, min_row, max_col, max_row = range_boundaries(text.replace("$", "").upper())
    except ValueError:
        return None
    if None in (min_col, min_row, max_col, max_row) or max_col > EXCEL_MAX_COLUMN or max_row > EXCEL_MAX_ROW:
        return None
    return min(min_col, max_col), min(min_row, max_row), max(min_col, max_col), max(min_row, max_row)


# ------------------------------------------------------------- reading cells


@dataclass
class SheetCells:
    """What a sheet's XML says about its cells, read once per sheet."""

    numbers: dict[str, Decimal]
    uncached: set[str]
    errors: set[str]


def sheet_cells(data: bytes, sheet_name: str, numeric_parser: str) -> SheetCells:
    """Exact numbers of one sheet, and which cells cannot give one.

    ``read_exact_numerics`` refuses the whole sheet over one formula without a
    cached result or one error value. A roll-forward maps single cells, so only
    the rows that point at such a cell are blocked; the rest of the sheet reads.
    """
    try:
        return SheetCells(read_exact_numerics(data, sheet_name, numeric_parser), set(), set())
    except FoundryError as exc:
        if not str(exc).startswith(("formula has no cached value", "worksheet error")):
            raise
    number = xlsx_number_reader(numeric_parser)
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    relationship = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    with _open_xlsx_archive(data) as archive:
        sheets = ET.fromstring(archive.read("xl/workbook.xml")).find("s:sheets", ns)
        rid = next((item.attrib[relationship] for item in sheets if item.attrib.get("name") == sheet_name), None)
        if rid is None:
            raise FoundryError(f"no worksheet: {sheet_name}")
        rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        target = next(item.attrib["Target"] for item in rels if item.attrib["Id"] == rid)
        member = target.lstrip("/") if target.startswith("/") else "xl/" + target
        found = SheetCells({}, set(), set())
        for _, cell in ET.iterparse(io.BytesIO(archive.read(member)), events=("end",)):
            if cell.tag != "{" + ns["s"] + "}c":
                continue
            coordinate = cell.attrib.get("r", "")
            value, formula = cell.find("s:v", ns), cell.find("s:f", ns)
            if cell.attrib.get("t") == "e":
                found.errors.add(coordinate)
            elif formula is not None and (value is None or value.text is None):
                found.uncached.add(coordinate)
            elif cell.attrib.get("t", "n") == "n" and value is not None and value.text:
                found.numbers[coordinate] = number(value.text)
            cell.clear()
        return found


def _is_date(value: Any) -> bool:
    return isinstance(value, (datetime, date, time, timedelta))


def read_workbook_source(value_book: Any, cells: SheetCells, sheet: str, bounds: tuple[int, int, int, int]) -> SourceValue:
    """One cell as it stands, or the exact sum of a range's numbers."""
    worksheet = value_book[sheet]
    min_col, min_row, max_col, max_row = bounds
    if (min_col, min_row) == (max_col, max_row):
        coordinate = f"{get_column_letter(min_col)}{min_row}"
        if coordinate in cells.errors:
            raise SourceUnreadable(f"来源单元格 {coordinate} 是错误值（例如 #N/A）")
        if coordinate in cells.uncached:
            raise SourceUnreadable(f"来源单元格 {coordinate} 是公式，但文件里没有保存计算结果；请在 Excel 里打开、保存后重新添加这份资料")
        value = worksheet.cell(row=min_row, column=min_col).value
        if _is_date(value):
            raise SourceUnreadable(f"来源单元格 {coordinate} 是日期或时间；这里只搬数字和文字，请改用手填值")
        if isinstance(value, bool):
            raise SourceUnreadable(f"来源单元格 {coordinate} 是 TRUE/FALSE；这里只搬数字和文字")
        if coordinate in cells.numbers:
            return SourceValue(number=cells.numbers[coordinate])
        if isinstance(value, str) and value.strip():
            return SourceValue(text=value)
        raise SourceUnreadable(f"来源单元格 {coordinate} 是空的；如果本月就是空白或 0，请改用手填值")
    if (max_col - min_col + 1) * (max_row - min_row + 1) > MAX_SUM_CELLS:
        raise SourceUnreadable("来源区域太大，请缩小到 2 万个单元格以内")
    total, counted = Decimal(0), 0
    # A sum of exact decimals must not round at the default 28 digits.
    with localcontext() as context:
        context.prec = 200
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                coordinate = f"{get_column_letter(col)}{row}"
                if coordinate in cells.errors:
                    raise SourceUnreadable(f"来源区域里 {coordinate} 是错误值")
                if coordinate in cells.uncached:
                    raise SourceUnreadable(f"来源区域里 {coordinate} 是公式，但文件里没有保存计算结果")
                value = worksheet.cell(row=row, column=col).value
                if _is_date(value):
                    raise SourceUnreadable(f"来源区域里 {coordinate} 是日期；合计只加数字")
                if isinstance(value, bool):
                    raise SourceUnreadable(f"来源区域里 {coordinate} 是 TRUE/FALSE；合计只加数字")
                if coordinate in cells.numbers:
                    total += cells.numbers[coordinate]
                    counted += 1
                    continue
                if value is None or (isinstance(value, str) and not value.strip()):
                    continue
                # 「1,200」 typed as text is not a blank: skipping it would hand over a short total.
                shown = str(value).strip().replace("\n", " ")[:20]
                raise SourceUnreadable(f"来源区域里 {coordinate} 是文字「{shown}」；合计只加数字，请修正资料或缩小区域")
    if not counted:
        raise SourceUnreadable("来源区域里没有数字")
    return SourceValue(number=total)


def read_run_source(run: dict[str, Any], row: int, column: str) -> SourceValue:
    """One cell of a calculation result, typed by its column rather than by how it looks."""
    if run.get("execution_status") != "succeeded":
        raise SourceUnreadable("这次计算没有成功")
    result = run.get("result") or {}
    rows = result.get("rows") or []
    if not 0 <= row < len(rows):
        raise SourceUnreadable(f"计算结果没有第 {row + 1} 行")
    if not isinstance(rows[row], dict) or column not in rows[row]:
        raise SourceUnreadable(f"计算结果没有「{column}」这一列")
    value = rows[row][column]
    if value is None:
        raise SourceUnreadable("计算结果这一格是空的")
    if isinstance(value, bool):
        raise SourceUnreadable("计算结果这一格是 TRUE/FALSE；这里只搬数字和文字")
    declared = next(
        (str(item.get("type") or "") for item in result.get("columns") or [] if isinstance(item, dict) and item.get("name") == column),
        "",
    )
    if NUMERIC_RESULT_TYPE.match(declared):
        try:
            number = Decimal(str(value))
        except (InvalidOperation, ValueError):
            raise SourceUnreadable("计算结果这一格不是有效数字") from None
        if not number.is_finite():
            raise SourceUnreadable("计算结果这一格不是有效数字")
        return SourceValue(number=number)
    return SourceValue(text=_cell_text(str(value), "计算结果这一格"))


def read_manager_value(text: str) -> SourceValue:
    if not text.strip():
        raise SourceUnreadable("手填值是空的")
    if PLAIN_DECIMAL.fullmatch(text.strip()):
        return SourceValue(number=Decimal(text.strip()))
    return SourceValue(text=_cell_text(text, "手填值"))


def _cell_text(text: str, what: str) -> str:
    if len(text) > 32767:
        raise SourceUnreadable(f"{what}超过 Excel 单元格能放下的 32767 个字符")
    if ILLEGAL_CHARACTERS_RE.search(text):
        raise SourceUnreadable(f"{what}里有 Excel 不接受的控制字符")
    return text


def rendered(value: SourceValue) -> tuple[str, str, list[str]]:
    """(new value as text, how the cell will hold it, notes the manager should see)."""
    if value.number is not None:
        exact = format(value.number, "f")
        if survives_a_cell(value.number):
            return exact, "number", []
        # A figure Excel would round is written as its exact digits, as text.
        return exact, "text", ["decimal_written_as_text"]
    text = value.text or ""
    return text, "text", ["formula_like_text"] if FORMULA_LIKE.match(text) else []


# ----------------------------------------------------------- the target side


def _shown(cell: Any, numbers: dict[str, Decimal]) -> Any:
    value = cell.value
    if isinstance(value, ArrayFormula):
        return value.text
    if isinstance(value, DataTableFormula):
        return "（模拟运算表公式）"
    if cell.data_type == "f":
        return str(value)
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    # A date is a serial number in the XML; show the date the manager sees.
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if cell.coordinate in numbers:
        return format(numbers[cell.coordinate], "f")
    if isinstance(value, (int, float)):
        return str(value)
    return value


def _array_members(worksheet: Any) -> list[tuple[str, tuple[int, int, int, int]]]:
    members = []
    for anchor, ref in worksheet.array_formulae.items():
        try:
            min_col, min_row, max_col, max_row = range_boundaries(str(ref).replace("$", ""))
        except ValueError:
            continue
        members.append((anchor, (min_col, min_row, max_col, max_row)))
    return members


def _blocked_target(worksheet: Any, cell: Any, members: list[tuple[str, tuple[int, int, int, int]]]) -> str | None:
    if isinstance(cell, MergedCell):
        return "merged_non_anchor"
    if cell.data_type == "f" or isinstance(cell.value, (ArrayFormula, DataTableFormula)):
        return "formula_cell"
    for anchor, (min_col, min_row, max_col, max_row) in members:
        if anchor != cell.coordinate and min_col <= cell.column <= max_col and min_row <= cell.row <= max_row:
            return "array_formula_member"
    if worksheet.protection.sheet and cell.protection.locked:
        return "protected_cell"
    if cell.data_type == "e":
        return "error_cell"
    return None


def _unchanged(cell: Any, numbers: dict[str, Decimal], new_value: str, cell_type: str) -> bool:
    if cell_type == "number":
        return cell.data_type == "n" and numbers.get(cell.coordinate) == Decimal(new_value)
    return cell.data_type == "s" and isinstance(cell.value, str) and cell.value == new_value


def plan_rows(
    book: Any,
    data: bytes,
    numeric_parser: str,
    mappings: list[dict[str, Any]],
    resolve: Any,
) -> list[dict[str, Any]]:
    """Classify every mapping against the workbook before anything is written.

    ``resolve(source)`` returns a SourceValue and a public description of the
    source, or raises SourceUnreadable. A blocked target still shows the figure
    its source holds, so the control list says what to type by hand.
    """
    parsed = []
    for mapping in mappings:
        target = mapping["target"]
        coordinate = cell_address(target["cell"]) if target["sheet"] else None
        parsed.append((target["sheet"], coordinate))
    seen: dict[tuple[str, str], int] = {}
    for sheet, coordinate in parsed:
        if coordinate and sheet in book.sheetnames:
            key = (sheet.casefold(), coordinate)
            seen[key] = seen.get(key, 0) + 1
    numbers_by_sheet: dict[str, dict[str, Decimal]] = {}
    members_by_sheet: dict[str, list] = {}
    rows = []
    for index, (mapping, (sheet, coordinate)) in enumerate(zip(mappings, parsed)):
        target = mapping["target"]
        row: dict[str, Any] = {
            "index": index,
            "target": f"{target['sheet']}!{coordinate or target['cell']}",
            "sheet": target["sheet"],
            "cell": coordinate or target["cell"],
            "status": BLOCKED,
            "old_value": None,
            "new_value": None,
            "cell_type": None,
            "number_format": None,
            "notes": [],
        }
        try:
            value, described = resolve(mapping["source"])
            row["source"] = described
            row["new_value"], row["cell_type"], row["notes"] = rendered(value)
            unreadable = ""
        except SourceUnreadable as exc:
            row["source"] = getattr(exc, "described", None) or dict(mapping["source"])
            unreadable = str(exc)
        code = None
        if coordinate is None:
            code = "bad_address"
        elif sheet not in book.sheetnames:
            code = "missing_sheet"
        elif seen.get((sheet.casefold(), coordinate), 0) > 1:
            code = "duplicate_target"
        cell = None
        if code is None or code == "duplicate_target":
            worksheet = book[sheet]
            if sheet not in numbers_by_sheet:
                numbers_by_sheet[sheet] = sheet_cells(data, sheet, numeric_parser).numbers
                members_by_sheet[sheet] = _array_members(worksheet)
            cell = worksheet[coordinate]
            row["old_value"] = _shown(cell, numbers_by_sheet[sheet])
            row["number_format"] = None if isinstance(cell, MergedCell) else cell.number_format
            code = code or _blocked_target(worksheet, cell, members_by_sheet[sheet])
        if code is None and unreadable:
            code = "source_unreadable"
        if code is not None:
            row["reason_code"] = code
            row["reason"] = unreadable if code == "source_unreadable" else REASONS[code]
        elif _unchanged(cell, numbers_by_sheet[sheet], row["new_value"], row["cell_type"]):
            row["status"] = UNCHANGED
        else:
            row["status"] = WRITE
        rows.append(row)
    return rows


def counts(rows: list[dict[str, Any]]) -> dict[str, int]:
    return {status: sum(1 for row in rows if row["status"] == status) for status in (WRITE, BLOCKED, UNCHANGED)}


def title_hints(book: Any, rows: list[dict[str, Any]], period: str) -> list[dict[str, Any]]:
    """Cells near the top of each sheet that still name some other month.

    The copy keeps last month's title text. A sheet whose title still says
    「2026年08月」 would make the September copy read as August to readiness and
    to whoever opens it, so these are offered for mapping, never changed here.
    """
    from foundry.finance_readiness import _declares_span, _period_match

    mapped = {(row["sheet"].casefold(), row["cell"]) for row in rows if row["status"] in {WRITE, UNCHANGED}}
    hints: list[dict[str, Any]] = []
    for worksheet in book.worksheets:
        top = min(worksheet.max_row or 1, 6)
        width = min(worksheet.max_column or 1, 12)
        for line in worksheet.iter_rows(min_row=1, max_row=top, max_col=width):
            for cell in line:
                if isinstance(cell, MergedCell) or cell.data_type == "f" or not isinstance(cell.value, str):
                    continue
                text = cell.value.strip()
                found, _words = _period_match(text)
                if not (found or _declares_span(text)) or found == period:
                    continue
                hints.append({
                    "sheet": worksheet.title,
                    "cell": cell.coordinate,
                    "target": f"{worksheet.title}!{cell.coordinate}",
                    "text": text[:200],
                    "period": found,
                    "mapped": (worksheet.title.casefold(), cell.coordinate) in mapped,
                })
                if len(hints) >= MAX_TITLE_HINTS:
                    return hints
    return hints


# ------------------------------------------------------------------ writing


def write_copy(data: bytes, destination: Path, rows: list[dict[str, Any]]) -> dict[str, Any]:
    """Write the planned cells into a new copy of the archived bytes.

    Each cell keeps its number format. A number goes in as a number only when
    Excel would hold it exactly; text always goes in as text, so 「=A1」 typed
    by the manager is a label, not a formula.
    """
    if destination.exists():
        raise FoundryError("rollforward output already exists")
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_bytes(data)
    book = _open(destination)
    try:
        # openpyxl drops every formula's cached result on save, written cell or not.
        formulas = any(
            cell.data_type == "f" for worksheet in book.worksheets for cell in worksheet._cells.values()
        )
        for row in rows:
            if row["status"] != WRITE:
                continue
            if row["sheet"] not in book.sheetnames:
                raise FoundryError("rollforward target changed since preview")
            cell = book[row["sheet"]][row["cell"]]
            if isinstance(cell, MergedCell) or cell.data_type in {"f", "e"}:
                raise FoundryError("rollforward target changed since preview")
            number_format = cell.number_format
            if row["cell_type"] == "number":
                exact = Decimal(row["new_value"])
                cell.value = int(exact) if exact == exact.to_integral_value() else float(exact)
            else:
                cell.value = row["new_value"]
                cell.data_type = "s"
            cell.number_format = number_format
        book.calculation.fullCalcOnLoad = True
        book.save(destination)
    finally:
        book.close()
    return {"formula_cache_cleared": formulas}


def output_filename(source_name: str, period: str) -> str:
    stem = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "", Path(str(source_name or "")).stem).strip()[:80] or "workbook"
    return f"{stem}（{period} 滚动）.xlsx"


def output_name(name: str, period: str) -> str:
    return f"{str(name or '工作簿')[:160]}（{period} 滚动）"


# ---------------------------------------------------------------- change log


def _one_line(text: Any) -> str:
    return re.sub(r"[\r\n  \t]+", " ", "" if text is None else str(text)).strip()


def _tsv_cell(value: Any) -> str:
    """A pasted cell that starts like a formula stays text; a plain number stays a number."""
    text = _one_line(value)
    if PLAIN_DECIMAL.fullmatch(text):
        return text
    return "'" + text if text[:1] in {"=", "+", "-", "@"} else text


def _md_cell(value: Any) -> str:
    return _one_line(value).replace("|", "\\|") or "—"


def source_label(source: dict[str, Any]) -> str:
    kind = source.get("kind")
    if kind == "dataset":
        where = f"{source.get('sheet', '')}!{source.get('range', '')}"
        total = "（合计）" if ":" in str(source.get("range", "")) else ""
        return f"{source.get('name') or '资料'} · {where}{total}"
    if kind == "run":
        return f"计算结果 第 {int(source.get('row', 0)) + 1} 行 · {source.get('column', '')}"
    note = source.get("note") or ""
    return "手填" + (f"：{note}" if note else "")


def _state(row: dict[str, Any], applied: bool) -> str:
    if row["status"] == WRITE:
        return "已写入" if applied else "将写入"
    if row["status"] == UNCHANGED:
        return "无变化"
    return f"未写入：{row.get('reason') or REASONS.get(row.get('reason_code', ''), '')}"


def change_log(
    rows: list[dict[str, Any]],
    *,
    period: str,
    workbook: str,
    copy: str | None = None,
) -> dict[str, str]:
    """Every mapped cell with its old and new value, to paste into Excel or a mail."""
    applied = copy is not None
    tally = counts(rows)
    header = ["单元格", "原值", "新值", "来源", "状态/原因"]
    table = [[row["target"], row.get("old_value"), row.get("new_value"), source_label(row.get("source") or {}), _state(row, applied)] for row in rows]
    lines = [
        "# 下月套表变更记录" if applied else "# 下月套表滚动清单",
        "",
        f"- 新月份：{period}",
        f"- 上月套表：{_one_line(workbook)}（原工作簿未改动）",
        *( [f"- 下月副本：{_one_line(copy)}"] if applied else [] ),
        f"- {'写入' if applied else '将写入'} {tally[WRITE]} 格 · 未写入 {tally[BLOCKED]} 格 · 无变化 {tally[UNCHANGED]} 格",
        "",
        "| " + " | ".join(header) + " |",
        "|" + "---|" * len(header),
        *("| " + " | ".join(_md_cell(cell) for cell in line) + " |" for line in table),
    ]
    tsv = "\n".join("\t".join(_tsv_cell(cell) for cell in line) for line in [header, *table])
    return {"markdown": "\n".join(lines) + "\n", "tsv": tsv + "\n"}


# ---------------------------------------------------------------- templates


def sheet_pattern(sheet_names: list[str]) -> str:
    """A workbook's sheets with the digits masked, so August's pack matches September's."""
    from foundry.structure_templates import sheet_pattern as masked

    return "|".join(masked(name) for name in sheet_names)


def template_mappings(rows: list[dict[str, Any]], mappings: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """The mapping without what belongs to one month: no snapshot, hash, run or typed value.

    A dataset source keeps where to read and the file's name with its digits
    masked, so next month the manager picks that month's file herself. A typed
    value keeps last month's text only as a reminder of what to type.
    """
    from foundry.structure_templates import sheet_pattern as masked

    names = {row["index"]: (row.get("source") or {}).get("name") for row in rows}
    kept = []
    for index, mapping in enumerate(mappings):
        source = mapping["source"]
        if source["kind"] == "dataset":
            stripped = {"kind": "dataset", "sheet": source["sheet"], "range": source["range"], "hint": masked(names.get(index) or "")[:200]}
        elif source["kind"] == "run":
            stripped = {"kind": "run", "row": source["row"], "column": source["column"]}
        else:
            stripped = {"kind": "value", "note": source["note"], "last_value": source["manager_value"][:200]}
        kept.append({"target": dict(mapping["target"]), "source": stripped})
    return kept


TEMPLATE_DIRNAME = "rollforward_templates"
TEMPLATE_LOCK = "rollforward_templates.lock"
TEMPLATE_DOMAIN = "foundry.finance.rollforward_templates.v1"


class RollforwardTemplates:
    """Saved roll-forward mappings for one var root, beside ``tasks/``.

    A task is one month, so a mapping kept inside it could never reach the next
    month's task. Like the structure templates, the shelf sits at the var root;
    unlike them, a saved mapping is never rewritten: a changed mapping is saved
    again under a new name.
    """

    def __init__(self, var_root: Any) -> None:
        self.var_root = Path(var_root)
        self.root = self.var_root / TEMPLATE_DIRNAME
        self.lock_path = self.var_root / TEMPLATE_LOCK

    def _inside(self, path: Path) -> Path:
        root = self.root.resolve()
        if not path.resolve().is_relative_to(root):
            raise FoundryError("rollforward template path escapes the registry")
        current = path
        while current.resolve() != root:
            if current.is_symlink():
                raise FoundryError("rollforward template paths cannot use symbolic links")
            if current == current.parent:
                raise FoundryError("rollforward template path has no registry ancestor")
            current = current.parent
        if current.is_symlink():
            raise FoundryError("rollforward template paths cannot use symbolic links")
        return path

    def _load(self, path: Path) -> dict[str, Any]:
        import json

        from foundry.analysis import _hash

        try:
            record = json.loads(path.read_text(encoding="utf8"))
        except (OSError, ValueError) as exc:
            raise FoundryError(f"rollforward template is unavailable: {path.name}") from exc
        check = dict(record)
        digest = check.pop("_integrity_hash", None)
        if not digest or digest != _hash(check):
            raise FoundryError("rollforward template integrity check failed")
        return record

    def list(self) -> list[dict[str, Any]]:
        if not self.root.is_dir():
            return []
        records = [self._load(self._inside(path)) for path in self.root.glob("*.json") if ID.fullmatch(path.stem)]
        return sorted(records, key=lambda row: (row.get("seq", 0), row["id"]))

    def create(self, **fields: Any) -> dict[str, Any]:
        import json

        from foundry.analysis import _hash
        from foundry.analysis_data import json_value
        from foundry.fsutil import write_json_atomic
        from foundry.locking import exclusive
        from foundry.task import now_iso

        record = {"id": uuid.uuid4().hex, "schema_version": 1, "domain": TEMPLATE_DOMAIN, "created_at": now_iso(), **json_value(fields)}
        self.root.mkdir(parents=True, exist_ok=True)
        with exclusive(self.lock_path, timeout=10):
            sequence = self._inside(self.root / "sequence.json")
            try:
                value = json.loads(sequence.read_text())["seq"] if sequence.exists() else 0
                if isinstance(value, bool) or not isinstance(value, int) or value < 0:
                    raise ValueError("invalid sequence")
            except (OSError, ValueError, KeyError, TypeError) as exc:
                raise FoundryError("rollforward template sequence is invalid") from exc
            write_json_atomic(sequence, {"seq": value + 1})
            record["seq"] = value + 1
            record["_integrity_hash"] = _hash(record)
            path = self._inside(self.root / f"{record['id']}.json")
            if path.exists():
                raise FoundryError("rollforward template already exists")
            write_json_atomic(path, record)
        return record


def public_template(record: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": record["id"],
        "name": record.get("name", ""),
        "sheet_pattern": record.get("sheet_pattern", ""),
        "created_at": record.get("created_at"),
        "mapping_count": len(record.get("mappings") or []),
        "mappings": record.get("mappings") or [],
    }
