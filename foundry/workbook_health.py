"""Formula health hints for an uploaded workbook (工作簿公式体检).

Run once when an XLSX file is added, and on demand for snapshots filed before
this check existed. The result is hints only: nothing here refuses an import,
changes a value or decides anything. The manager reads the hints and decides
what, if anything, to change in her own workbook.

Pure by contract: bytes in, a small dict out. It reads no task directory and
writes no record. It never raises: a workbook it cannot read comes back as
``{"schema_version": 1, "status": "unavailable", "reason": <code>}``.

What it reports (code -> severity):

- ``external_links`` (warning): the package carries links to other workbooks.
  Only their count and each target's file name are kept, never a path.
- ``iterative_calculation`` (warning): ``<calcPr iterate="1">``; circular
  references then settle silently instead of being reported by Excel.
- ``hidden_sheets`` (info): sheets whose state is hidden or veryHidden.
- ``external_reference`` (warning): a formula reads another workbook
  (``[1]Sheet!A1``, ``'dir\\[book.xlsx]Sheet'!A1``, ``[2]!Name``).
- ``volatile_function`` (info): OFFSET, INDIRECT, NOW, TODAY, RAND,
  RANDBETWEEN, CELL or INFO (an ``_xlfn.`` prefix is ignored).
- ``hardcoded_number`` (warning): a number typed into a formula; see the
  policy below.
- ``error_in_formula`` (warning): an error literal in the formula itself
  (``#REF!``, or a ``Sheet1!#REF!`` reference). Quoted text such as
  ``"#REF!"`` is text, not an error, and is not reported.
- ``cached_error`` (warning): a cell whose saved value is an error
  (``t="e"``), formula or not.
- ``formula_without_cache`` (info): a formula with no saved result. A cell is
  cached when its ``<v>`` holds text, or when it is ``t="str"`` with a
  ``<v>`` at all (Excel's blank-string result), the rule catalog.py uses.
  Such a sheet cannot be read as a table until the file is saved in Excel.
- ``inconsistent_row_formula`` / ``inconsistent_column_formula`` (warning):
  a formula that breaks the pattern its row (column) was copied with. Each
  A1 reference is rewritten relative to its own cell (``R[dr]C[dc]``, with
  ``$``-absolute parts kept absolute; external references are kept as
  written), so copies of one formula share a signature. When at least
  ``MIN_PATTERN_CELLS`` formulas of the row share one signature and they are
  at least ``PATTERN_SHARE_PERCENT`` percent of the row's formulas, a formula
  that differs is reported if a pattern cell sits on both sides of it. A cell
  at either end of the run is not reported: that is where a total, a ratio
  or an opening balance normally sits, and the hint is about a copy that was
  overwritten in the middle.
- ``sheet_too_large`` (info): a sheet beyond ``MAX_ROWS``/``MAX_COLUMNS`` was
  skipped rather than scanned.

Hard-coded number policy. A NUMBER operand is reported unless it is

- a common constant: 0, 1, 2, 10, 12, 100, 1000, 10000, 365 or 366, or any
  power of ten (1E+n), which is how amounts are converted between 元, 万元
  and 亿元; a leading minus sign is an operator, so -1 counts as 1;
- inside an array constant (``{1,2;3,4}``);
- directly in a structural argument position, where it says where or how to
  read rather than how much: the digits of ROUND/ROUNDUP/ROUNDDOWN/TRUNC,
  the row/column/area of INDEX, MATCH's match type, the column/row index and
  range-lookup flag of VLOOKUP/HLOOKUP, OFFSET's offsets and size, CHOOSE's
  index, LEFT/RIGHT/MID character positions, TEXT's format, plus the
  equivalents in XLOOKUP/XMATCH (modes), SUBTOTAL/AGGREGATE (function
  number), LARGE/SMALL (k), EOMONTH/EDATE (months), FIND/SEARCH (start)
  and WEEKDAY (return type).

Numbers inside text operands are text and are never seen as numbers.
Everything else is reported, for example ``=B5*1.13`` or
``=SUM(B2:B9)+2793.28``: such a number is usually a plugged amount or an
assumption that belongs in an input cell with its source written down.

``count`` is the number of occurrences (cells, sheets or links) for that code;
``issue_count`` is the number of codes with at least one occurrence. At most
``MAX_SAMPLES`` samples are kept per code, each naming a sheet and a cell and,
for a cell, a bounded ``range`` around it that a read-only preview can open.
At most ``MAX_FORMULA_CELLS`` formula cells are scanned; beyond that the
result says ``truncated: true``. The output holds ints, booleans and strings
only, in a fixed order.
"""

from __future__ import annotations

import posixpath
import re
import zlib
from collections import Counter
from decimal import Decimal, InvalidOperation
from typing import Any
from urllib.parse import unquote
from xml.etree import ElementTree as ET
from zipfile import BadZipFile

from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.utils import column_index_from_string, get_column_letter

from foundry.analysis_data import MAX_COLUMNS, MAX_ROWS, _open_xlsx_archive
from foundry.errors import FoundryError

SCHEMA_VERSION = 1
MAX_FORMULA_CELLS = 200_000
MAX_SAMPLES = 5
MAX_SHEET_NAMES = 50
MAX_LINK_TARGETS = 20
MAX_FORMULA_TEXT = 200
MAX_DETAIL_TEXT = 80
MIN_PATTERN_CELLS = 3
PATTERN_SHARE_PERCENT = 65

# Report order, which is also the order a reader meets them in.
CODES: tuple[tuple[str, str], ...] = (
    ("external_links", "warning"),
    ("iterative_calculation", "warning"),
    ("hidden_sheets", "info"),
    ("external_reference", "warning"),
    ("volatile_function", "info"),
    ("hardcoded_number", "warning"),
    ("error_in_formula", "warning"),
    ("cached_error", "warning"),
    ("formula_without_cache", "info"),
    ("inconsistent_row_formula", "warning"),
    ("inconsistent_column_formula", "warning"),
    ("sheet_too_large", "info"),
)
SEVERITY = dict(CODES)

VOLATILE = frozenset({"OFFSET", "INDIRECT", "NOW", "TODAY", "RAND", "RANDBETWEEN", "CELL", "INFO"})
COMMON_NUMBERS = frozenset(Decimal(text) for text in ("0", "1", "2", "10", "12", "100", "1000", "10000", "365", "366"))
# Zero-based argument positions whose number says where or how to read, not how much.
STRUCTURAL_ARGUMENTS: dict[str, frozenset[int]] = {
    name: frozenset(positions)
    for name, positions in {
        "ROUND": (1,), "ROUNDUP": (1,), "ROUNDDOWN": (1,), "TRUNC": (1,),
        "INDEX": (1, 2, 3), "MATCH": (2,), "VLOOKUP": (2, 3), "HLOOKUP": (2, 3),
        "OFFSET": (1, 2, 3, 4), "CHOOSE": (0,),
        "LEFT": (1,), "RIGHT": (1,), "MID": (1, 2), "TEXT": (1,),
        "XLOOKUP": (4, 5), "XMATCH": (2, 3), "SUBTOTAL": (0,), "AGGREGATE": (0, 1),
        "LARGE": (1,), "SMALL": (1,), "EOMONTH": (1,), "EDATE": (1,),
        "FIND": (2,), "SEARCH": (2,), "WEEKDAY": (1,),
    }.items()
}
FUNCTION_PREFIXES = ("_xlfn.", "_xlws.")

# The bracketed book comes before the sheet's "!": [1]Sheet1!A1, '..\[b.xlsx]S'!A1, [2]!Name.
# A structured reference (Table1[Amount]) has no "!" after its bracket.
_EXTERNAL = re.compile(r"^'?[^'!\[]*\[[^\]]+\][^!]*!")
_REFERENCE = re.compile(
    r"^(?P<prefix>.*!)?(?P<ref>"
    r"\$?[A-Za-z]{1,3}\$?\d+(?::\$?[A-Za-z]{1,3}\$?\d+)?"
    r"|\$?[A-Za-z]{1,3}:\$?[A-Za-z]{1,3}"
    r"|\$?\d+:\$?\d+)$"
)
_AXIS = re.compile(r"(\$?)([A-Za-z]{1,3})?(\$?)(\d+)?")
_CELL = re.compile(r"^([A-Z]{1,3})(\d+)$")
# A directory in front of an external book: '<dir>\[book.xlsx]Sheet'!A1.
_BOOK_DIRECTORY = re.compile(r"'[^'\[]*[\\/]\[")
_PATH_LIKE = re.compile(r"(?i)(/Users/|/home/|/tmp/|[A-Z]:\\|\\\\|file:/|https?://|wss?://)")

_PKG_REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"
_TRANSITIONAL = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def check_workbook(data: bytes) -> dict[str, Any]:
    """Formula health hints for one XLSX file. Never raises."""
    try:
        return _check(data)
    except FoundryError as exc:
        reason = "workbook_too_large" if "expands beyond" in str(exc) else "unreadable_workbook"
    except (BadZipFile, ET.ParseError, KeyError, ValueError, UnicodeError, EOFError, zlib.error,
            NotImplementedError, OSError):
        reason = "unreadable_workbook"
    except Exception:  # noqa: BLE001 - a hint must never cost the manager her import
        reason = "check_failed"
    return {"schema_version": SCHEMA_VERSION, "status": "unavailable", "reason": reason}


def _check(data: bytes) -> dict[str, Any]:
    with _open_xlsx_archive(data) as archive:
        book_part = _workbook_part(archive)
        book = ET.fromstring(archive.read(book_part))
        ns = book.tag[: book.tag.index("}") + 1] if book.tag.startswith("{") else _TRANSITIONAL
        folder = posixpath.dirname(book_part)
        relationships = _relationships(archive, book_part)
        report = _Report()

        sheets: list[tuple[str, str | None]] = []
        for sheet in book.iter(f"{ns}sheet"):
            name = str(sheet.get("name") or "")
            state = str(sheet.get("state") or "visible")
            rid = next((value for key, value in sheet.attrib.items() if key.endswith("}id")), None)
            kind, target, mode = relationships.get(str(rid), ("", "", ""))
            part = _part_name(folder, target) if kind.endswith("/worksheet") and mode != "External" else None
            sheets.append((name, part))
            if state in ("hidden", "veryHidden"):
                report.add("hidden_sheets", name, sample={"sheet": name, "detail": state})
        report.sheet_count = len(sheets)

        calculation = book.find(f"{ns}calcPr")
        if calculation is not None and str(calculation.get("iterate") or "").lower() in ("1", "true"):
            report.add("iterative_calculation", None)

        members = set(archive.namelist())
        links, targets = _external_links(archive, relationships, folder, members)
        if links:
            report.add("external_links", None, count=links)
        report.targets = targets

        formulas = _FormulaCache()
        for name, part in sheets:
            if part is None or part not in members:
                continue
            if report.truncated:
                break
            _scan_sheet(archive, part, name, ns, formulas, report)
        return report.result()


def _workbook_part(archive) -> str:
    try:
        package = ET.fromstring(archive.read("_rels/.rels"))
    except KeyError:
        return "xl/workbook.xml"
    for rel in package.iter(f"{_PKG_REL}Relationship"):
        if str(rel.get("Type") or "").endswith("/officeDocument"):
            return _part_name("", str(rel.get("Target") or ""))
    return "xl/workbook.xml"


def _part_name(base_dir: str, target: str) -> str:
    if target.startswith("/"):
        return target[1:]
    return posixpath.normpath(posixpath.join(base_dir, target) if base_dir else target)


def _relationships(archive, part: str) -> dict[str, tuple[str, str, str]]:
    folder, name = posixpath.split(part)
    try:
        root = ET.fromstring(archive.read(posixpath.join(folder, "_rels", f"{name}.rels")))
    except KeyError:
        return {}
    return {
        str(rel.get("Id")): (str(rel.get("Type") or ""), str(rel.get("Target") or ""), str(rel.get("TargetMode") or ""))
        for rel in root.iter(f"{_PKG_REL}Relationship")
    }


def _external_links(archive, relationships, folder: str, members: set[str]) -> tuple[int, list[str]]:
    """How many external link parts the package carries, and their targets' file names."""
    parts = {
        _part_name(folder, target)
        for kind, target, mode in relationships.values()
        if kind.endswith("/externalLink") and mode != "External"
    }
    parts.update(name for name in members if re.fullmatch(r"xl/externalLinks/externalLink\d+\.xml", name))
    parts &= members
    targets: list[str] = []
    for part in sorted(parts):
        for _, target, mode in _relationships(archive, part).values():
            name = _file_name(target) if mode == "External" else None
            if name and name not in targets and len(targets) < MAX_LINK_TARGETS:
                targets.append(name)
    return len(parts), targets


def _file_name(target: str) -> str | None:
    """The last segment of a link target. The directory never leaves this function."""
    text = unquote(target.split("?", 1)[0].split("#", 1)[0]).strip().rstrip("\\/")
    name = re.split(r"[\\/]", text)[-1].strip()
    if not name or ":" in name or _PATH_LIKE.search(name):
        return None
    return name[:120]


class _Report:
    def __init__(self) -> None:
        self.sheet_count = 0
        self.checked_formula_cells = 0
        self.truncated = False
        self.targets: list[str] = []
        self.counts: Counter[str] = Counter()
        self.sheets: dict[str, list[str]] = {}
        self.samples: dict[str, list[dict[str, Any]]] = {}

    def add(self, code: str, sheet: str | None, count: int = 1, sample: dict[str, Any] | None = None) -> None:
        self.counts[code] += count
        names = self.sheets.setdefault(code, [])
        if sheet is not None and sheet not in names and len(names) < MAX_SHEET_NAMES:
            names.append(sheet)
        samples = self.samples.setdefault(code, [])
        if sample is not None and len(samples) < MAX_SAMPLES:
            samples.append(sample)

    def result(self) -> dict[str, Any]:
        issues = []
        for code, severity in CODES:
            if not self.counts[code]:
                continue
            issue: dict[str, Any] = {
                "code": code,
                "severity": severity,
                "count": self.counts[code],
                "sheets": self.sheets.get(code, []),
                "samples": self.samples.get(code, []),
            }
            if code == "external_links":
                issue["targets"] = self.targets
            issues.append(issue)
        return {
            "schema_version": SCHEMA_VERSION,
            "status": "checked",
            "checked_formula_cells": self.checked_formula_cells,
            "truncated": self.truncated,
            "sheet_count": self.sheet_count,
            "issue_count": len(issues),
            "issues": issues,
        }


class _Formula:
    """What one formula text says, independent of the cell it sits in."""

    __slots__ = ("volatile", "external", "numbers", "errors", "template")

    def __init__(self, volatile=(), external=False, numbers=(), errors=(), template=None):
        self.volatile: tuple[str, ...] = volatile
        self.external: bool = external
        self.numbers: tuple[str, ...] = numbers
        self.errors: tuple[str, ...] = errors
        self.template: tuple | None = template


class _FormulaCache:
    """Tokenise each formula shape once.

    A formula copied down a column has a different text in every cell
    (``=B5*C5``, ``=B6*C6``) but one shape relative to its own cell. The cache
    key is that shape, found with a cheap pass over the text that rewrites
    each A1 cell reference outside quotes and brackets relative to the cell.
    Two cells with the same key read the same once tokenised, so they share
    one analysis and one signature, and the tokenizer runs once per shape.
    """

    def __init__(self) -> None:
        self.shapes: dict[str, tuple[_Formula, int | None]] = {}
        self.signatures: dict[tuple, int] = {}

    def at(self, text: str, row: int, column: int) -> tuple[_Formula, int | None]:
        key = _shape(text, row, column)
        found = self.shapes.get(key)
        if found is None:
            formula = _analyse(text)
            found = self.shapes[key] = (formula, self.signature(formula, row, column))
        return found

    def signature(self, formula: _Formula, row: int, column: int) -> int | None:
        if formula.template is None:
            return None
        rendered = tuple(
            piece if piece.__class__ is str else _render_reference(piece, row, column)
            for piece in formula.template
        )
        return self.signatures.setdefault(rendered, len(self.signatures))


# Quoted text, quoted sheet names and bracketed parts are kept as written; any
# other A1 cell reference becomes relative to its cell. Function names such as
# LOG10( and number exponents such as 1E3 are not references.
_SHAPE_PARTS = re.compile(
    r'"(?:[^"]|"")*"?'
    r"|'(?:[^']|'')*'?"
    r"|\[[^\]]*\]?"
    r"|(?<![A-Za-z0-9_.])(\$?)([A-Za-z]{1,3})(\$?)(\d+)(?![A-Za-z0-9_(.])"
)


def _shape(text: str, row: int, column: int) -> str:
    def relative(match: re.Match) -> str:
        if match.group(2) is None:
            return match.group(0)
        column_absolute, letters, row_absolute, digits = match.groups()
        index = _column_index(letters.upper())
        return (
            (f"R{digits}" if row_absolute else f"R[{int(digits) - row}]")
            + (f"C{index}" if column_absolute else f"C[{index - column}]")
        )

    return _SHAPE_PARTS.sub(relative, text)


def _analyse(text: str) -> _Formula:
    try:
        items = Tokenizer(text if text.startswith("=") else "=" + text).items
    except Exception:  # noqa: BLE001 - an unreadable formula gets no hint, not a failure
        return _Formula()
    volatile: list[str] = []
    numbers: list[str] = []
    errors: list[str] = []
    external = False
    template: list[Any] = []
    # One frame per open function, parenthesis or array constant: [function name or None or "{", argument index].
    frames: list[list[Any]] = []
    for index, token in enumerate(items):
        kind, subtype, value = token.type, token.subtype, token.value
        if kind == Token.WSPACE:
            continue
        if subtype == Token.OPEN:
            if kind == Token.FUNC:
                name = _function_name(value)
                if name in VOLATILE and name not in volatile:
                    volatile.append(name)
                frames.append([name, 0])
                template.append(name + "(")
            else:
                frames.append(["{" if kind == Token.ARRAY else None, 0])
                template.append(value)
            continue
        if subtype == Token.CLOSE:
            if frames:
                frames.pop()
            template.append(value)
            continue
        if kind == Token.SEP:
            if subtype == Token.ARG and frames:
                frames[-1][1] += 1
            template.append(value)
            continue
        if kind != Token.OPERAND:
            template.append(value)
            continue
        if subtype == Token.NUMBER:
            template.append(value)
            if not _in_array(frames) and not _structural(frames) and not _common(value):
                following = _next_token(items, index)
                percent = following is not None and following.type == Token.OP_POST and following.value == "%"
                numbers.append(value + ("%" if percent else ""))
        elif subtype == Token.ERROR:
            template.append(value.upper())
            errors.append(value.upper())
        elif subtype == Token.RANGE:
            if _EXTERNAL.match(value):
                external = True
                template.append(value)
                continue
            if "#REF!" in value.upper():
                errors.append("#REF!")
            template.append(_reference_template(value))
        else:
            template.append(value)
    return _Formula(
        volatile=tuple(volatile),
        external=external,
        numbers=tuple(dict.fromkeys(numbers)),
        errors=tuple(dict.fromkeys(errors)),
        template=tuple(template),
    )


def _function_name(value: str) -> str:
    name = value[:-1] if value.endswith("(") else value
    for prefix in FUNCTION_PREFIXES:
        if name.lower().startswith(prefix):
            name = name[len(prefix):]
    return name.upper()


def _next_token(items, index: int):
    for token in items[index + 1:]:
        if token.type != Token.WSPACE:
            return token
    return None


def _in_array(frames) -> bool:
    return any(frame[0] == "{" for frame in frames)


def _structural(frames) -> bool:
    """Whether a number here sits directly in a structural argument of the nearest function."""
    for name, argument in reversed(frames):
        if name is None:
            continue  # a parenthesised part of the same argument
        positions = STRUCTURAL_ARGUMENTS.get(name)
        return positions is not None and argument in positions
    return False


def _common(value: str) -> bool:
    try:
        number = Decimal(value)
    except InvalidOperation:
        return False
    if not number.is_finite():
        return False
    if number in COMMON_NUMBERS:
        return True
    # 1E+n: 元 to 万元 to 亿元 and back.
    return number >= 1 and number == number.to_integral_value() and number.normalize().as_tuple().digits == (1,)


def _reference_template(value: str) -> Any:
    match = _REFERENCE.match(value)
    if match is None:
        return value  # a defined name or a structured reference reads the same from any cell
    prefix = match.group("prefix") or ""
    parts: list[Any] = [prefix]
    for index, endpoint in enumerate(match.group("ref").split(":")):
        if index:
            parts.append(":")
        axis = _AXIS.fullmatch(endpoint)
        if axis is None:
            return value
        column_absolute, letters, row_absolute, digits = axis.groups()
        if not letters:
            row_absolute = column_absolute or row_absolute  # "$5" in a whole-row reference
        if letters:
            parts.append(("C", bool(column_absolute), _column_index(letters.upper())))
        if digits:
            parts.append(("R", bool(row_absolute), int(digits)))
    return tuple(parts)


def _render_reference(parts: tuple, row: int, column: int) -> str:
    rendered = []
    for part in parts:
        if part.__class__ is str:
            rendered.append(part)
            continue
        axis, absolute, number = part
        if absolute:
            rendered.append(f"{axis}{number}")
        else:
            rendered.append(f"{axis}[{number - (row if axis == 'R' else column)}]")
    return "".join(rendered)


_COLUMNS: dict[str, int] = {}


def _column_index(letters: str) -> int:
    index = _COLUMNS.get(letters)
    if index is None:
        index = _COLUMNS[letters] = column_index_from_string(letters)
    return index


class _SheetScan:
    """One sheet's findings, kept apart until the sheet is known to be within limits."""

    def __init__(self, name: str) -> None:
        self.name = name
        self.formula_cells = 0
        self.max_row = 0
        self.max_column = 0
        self.counts: Counter[str] = Counter()
        self.samples: dict[str, list[tuple[int, int, Any, str | None]]] = {}
        # (row, column, signature, formula text or shared master) per scanned formula.
        self.patterns: list[tuple[int, int, int, Any]] = []

    def add(self, code: str, row: int, column: int, text: Any, detail: str | None) -> None:
        self.counts[code] += 1
        samples = self.samples.setdefault(code, [])
        if len(samples) < MAX_SAMPLES:
            samples.append((row, column, text, detail))


def _scan_sheet(archive, part: str, name: str, ns: str, formulas: _FormulaCache, report: _Report) -> None:
    row_tag, cell_tag, formula_tag, value_tag, dimension_tag = (
        f"{ns}row", f"{ns}c", f"{ns}f", f"{ns}v", f"{ns}dimension",
    )
    scan = _SheetScan(name)
    shared: dict[str, tuple[_Formula, int | None, tuple[str, int, int]]] = {}
    row_number = 0
    with archive.open(part) as stream:
        for _, element in ET.iterparse(stream, events=("end",)):
            tag = element.tag
            if tag == dimension_tag:
                if _beyond_limits(str(element.get("ref") or "")):
                    report.add("sheet_too_large", name, sample={"sheet": name})
                    return
                continue
            if tag != row_tag:
                continue
            given = element.get("r")
            row_number = int(given) if given and given.isdigit() else row_number + 1
            cells = element.findall(cell_tag)
            last = len(cells) - 1
            for index, cell in enumerate(cells):
                cell_type = cell.get("t")
                formula_element = cell.find(formula_tag)
                # Plain values only matter for the sheet's extent, which its last cell in the row gives.
                if formula_element is None and cell_type != "e" and index != last:
                    continue
                reference = cell.get("r")
                matched = _CELL.match(reference) if reference else None
                if matched:
                    column_number = _column_index(matched.group(1))
                    row = int(matched.group(2))
                else:
                    column_number = index + 1
                    row = row_number
                if row > MAX_ROWS + 1 or column_number > MAX_COLUMNS:
                    # One oversized sheet must not cost the manager the rest of the workbook.
                    report.add("sheet_too_large", name, sample={"sheet": name})
                    return
                if row > scan.max_row:
                    scan.max_row = row
                if column_number > scan.max_column:
                    scan.max_column = column_number
                if formula_element is None and cell_type != "e":
                    continue
                value = cell.find(value_tag)
                text: Any = None
                if formula_element is not None:
                    if report.checked_formula_cells + scan.formula_cells >= MAX_FORMULA_CELLS:
                        report.truncated = True
                        _finish(scan, report)
                        return
                    scan.formula_cells += 1
                    text = _scan_formula(
                        scan, formulas, shared, formula_element, row, column_number,
                        cached=value is not None and (bool(value.text) or cell_type in ("str", "e")),
                    )
                if cell_type == "e":
                    scan.add("cached_error", row, column_number, text,
                             (value.text or "").strip()[:MAX_DETAIL_TEXT] if value is not None else None)
            element.clear()
    _finish(scan, report)


def _scan_formula(scan: _SheetScan, formulas: _FormulaCache, shared, element, row: int, column: int, *, cached: bool) -> Any:
    """Record one formula cell's hints; return what renders its text for a sample."""
    kind = element.get("t")
    body = element.text or ""
    formula: _Formula | None = None
    signature: int | None = None
    text: Any = None
    if kind == "shared" and element.get("si") in shared:
        # Excel stores a copied formula once; every copy reads the same, shifted.
        formula, signature, text = shared[element.get("si")]
    elif kind != "dataTable" and body:
        formula, signature = formulas.at(body, row, column)
        text = body
        if kind == "shared" and element.get("si") is not None:
            shared[element.get("si")] = (formula, signature, (body, row, column))
            text = (body, row, column)
    if not cached:
        scan.add("formula_without_cache", row, column, text, None)
    if formula is None:
        return text
    if formula.external:
        scan.add("external_reference", row, column, text, None)
    if formula.volatile:
        scan.add("volatile_function", row, column, text, ", ".join(formula.volatile))
    if formula.numbers:
        scan.add("hardcoded_number", row, column, text, _detail(formula.numbers))
    if formula.errors:
        scan.add("error_in_formula", row, column, text, _detail(formula.errors))
    if signature is not None:
        scan.patterns.append((row, column, signature, text))
    return text


def _detail(values: tuple[str, ...]) -> str:
    return ", ".join(values[:3])[:MAX_DETAIL_TEXT]


def _beyond_limits(reference: str) -> bool:
    last = reference.split(":")[-1].replace("$", "").upper()
    matched = _CELL.match(last)
    if matched is None:
        return False
    return int(matched.group(2)) > MAX_ROWS + 1 or _column_index(matched.group(1)) > MAX_COLUMNS


def _finish(scan: _SheetScan, report: _Report) -> None:
    report.checked_formula_cells += scan.formula_cells
    for code, flagged in (
        ("inconsistent_row_formula", _breaks(scan.patterns, along_rows=True)),
        ("inconsistent_column_formula", _breaks(scan.patterns, along_rows=False)),
    ):
        for row, column, text, example in flagged:
            scan.add(code, row, column, text, example)
    for code, _ in CODES:
        count = scan.counts[code]
        if not count:
            continue
        report.add(code, scan.name, count=count)
        samples = report.samples[code]
        for row, column, text, detail in scan.samples.get(code, []):
            if len(samples) >= MAX_SAMPLES:
                break
            sample: dict[str, Any] = {"sheet": scan.name, "cell": f"{get_column_letter(column)}{row}",
                                      "range": _window(row, column, scan.max_row, scan.max_column)}
            shown = _formula_text(text, row, column)
            if shown is not None:
                sample["formula"] = shown
            if detail:
                sample["detail"] = detail
            samples.append(sample)


def _breaks(patterns, *, along_rows: bool) -> list[tuple[int, int, Any, str]]:
    """Formulas inside a copied run of their row (or column) that read differently."""
    lines: dict[int, list[tuple[int, int, int, int, Any]]] = {}
    for row, column, signature, text in patterns:
        key, position = (row, column) if along_rows else (column, row)
        lines.setdefault(key, []).append((position, signature, row, column, text))
    flagged = []
    for members in lines.values():
        if len(members) < MIN_PATTERN_CELLS:
            continue
        signature, share = Counter(member[1] for member in members).most_common(1)[0]
        if share == len(members) or share < MIN_PATTERN_CELLS or share * 100 < PATTERN_SHARE_PERCENT * len(members):
            continue
        anchors = sorted(member[0] for member in members if member[1] == signature)
        for position, other, row, column, text in members:
            if other == signature or not anchors[0] < position < anchors[-1]:
                continue
            nearest = min(anchors, key=lambda anchor: (abs(anchor - position), anchor))
            # The detail names the closest cell that follows the pattern.
            example = f"{get_column_letter(nearest)}{row}" if along_rows else f"{get_column_letter(column)}{nearest}"
            flagged.append((row, column, text, example))
    return sorted(flagged, key=lambda item: (item[0], item[1]))


def _window(row: int, column: int, max_row: int, max_column: int) -> str:
    """A bounded range around the cell, inside the cells the sheet actually has."""
    top, bottom = max(1, row - 4), max(row, min(max_row, row + 5))
    left, right = max(1, column - 3), max(column, min(max_column, column + 3))
    return f"{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}"


def _formula_text(text: Any, row: int, column: int) -> str | None:
    if text is None:
        return None
    if isinstance(text, tuple):
        body, origin_row, origin_column = text
        if (origin_row, origin_column) == (row, column):
            formula = "=" + body
        else:
            try:
                formula = Translator("=" + body, origin=f"{get_column_letter(origin_column)}{origin_row}").translate_formula(
                    f"{get_column_letter(column)}{row}"
                )
            except Exception:  # noqa: BLE001 - the sample then names the cell without its text
                return None
    else:
        formula = text if text.startswith("=") else "=" + text
    formula = _BOOK_DIRECTORY.sub("'[", formula)
    if _PATH_LIKE.search(formula):
        return None
    return formula if len(formula) <= MAX_FORMULA_TEXT else formula[: MAX_FORMULA_TEXT - 1] + "…"
