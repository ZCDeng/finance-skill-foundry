"""Task-owned analysis snapshots, immutable revisions, execution and deliverables.

These objects are parallel to legacy closing windows. Creating an analysis never
creates a ledger, imports a chart of accounts, or grants a payment instruction.
"""

from __future__ import annotations

import hashlib
import io
import json
import platform
import re
import sys
import uuid
from decimal import Decimal
from pathlib import Path
from typing import Any, Callable

import duckdb
from openpyxl import Workbook
from openpyxl.comments import Comment

from foundry.analysis_data import MAX_SOURCE_BYTES, SUPPORTED, json_value, parse_source
from foundry.analysis_sql import CONFIG, execute_sql, method_template, validate_relationship_contract
from foundry.analysis_parameters import normalize_parameter_metadata, method_parameters
from foundry.analysis_period import bind_period, normalize_period_binding, resolve_period
from foundry.errors import FoundryError
from foundry.fsutil import write_json_atomic
from foundry.locking import exclusive
from foundry.task import Task, now_iso
from foundry.sheet import survives_a_cell, put_number
from foundry.report_documents import report_document_excel_rows

ID = re.compile(r"^[0-9a-f]{32}$")
ALIAS = re.compile(r"^[A-Za-z_][A-Za-z0-9_]{0,63}$")
# estimate_acceptances holds a manager's acceptance of a specific estimate recipe (phase 2 adds the
# manager-only command); requirement checks only read it and never create one.
KINDS = {"revisions", "runs", "drafts", "methods", "requirements", "evidence_checks", "completions",
         "estimate_acceptances", "dataset_display_names"}
REPORT_EXPORT_FORMATS = {"md", "xlsx", "docx", "pdf", "html"}
RESULT_ROLES = {"business", "activity", "probe", "preanalysis", "calculation", "draft_support", "audit_check"}
# What a script stands in for, when an engine path could not run yet: 回款健康 drawn by a script while the workbook's
# structure waits on the manager. The page shows it on 经营分析 as a script draft, never as the engine's result.
SCRIPT_TOPICS = {"collection_health"}
REVIEW = {"unreviewed", "reviewed", "accepted"}
PUBLICATION = {"draft", "archived"}
SCRIPT_RESULT_PREVIEW_ROWS = 1000
# One answer for a parent draft a session may not read and for one that does not exist, so it cannot tell them apart.
PARENT_UNAVAILABLE = "parent draft is unavailable or outside this session context"
# The same for a draft a call names by id (validate, review, circle-select).
DRAFT_UNAVAILABLE = "draft is unavailable or outside this session context"
# The same for a requirements version a call names by id (a spec's requirements_id): AnalysisStore.named_requirements.
REQUIREMENTS_UNAVAILABLE = "requirements version is unavailable or outside this session context"
# The same for a run an explanation names (its run_ids, a claim's or a recipe's run): FinanceStore._named_run.
RUN_UNAVAILABLE = "run is unavailable or outside this session context"
# The same for a saved analysis a command names (to run it, to save it as a method, or as a revision's parent).
ANALYSIS_UNAVAILABLE = "saved analysis is unavailable or outside this session context"
# What a revision takes over from its parent when the spec leaves it out (save_draft copies it): what the parent's
# author wrote or the server generated from the parent's files. (The runs and requirements version it takes are read
# as its own.) A revision records which of these it took (`inherited_fields`), and one that took any rests on
# everything its parent rests on (AnalysisStore.draft_snapshots).
INHERITED_FIELDS = ("title", "body", "claims", "fact_ids", "materiality", "limitations", "source_read_failures",
                    "evidence_appendix", "evidence_appendix_meta", "evidence_recipes", "consistency_rules",
                    "consistency_sources")
# Frozen legal-workflow identity. Service-owned required_check_authorship stays
# on the record and inside _integrity_hash; it is not part of this hash.
DRAFT_CONTENT_FIELDS = (
    "title", "requirements_id", "materiality", "body", "body_template", "document",
    "run_ids", "fact_ids", "claims", "limitations",
)
DEFAULT_DRAFT_TITLE = "分析报告"


def _draft_required_check_authorship(prepared, runs, *, claims=(), fact_ids=(), template_fact_ids=(), inherited=None):
    """Keep explicit calculation declarations separate from a generated evidence appendix.

    The appendix remains readable and fully cited. A supplied calculation claim
    or fact selection must still acknowledge its mandatory residual checks; repeating other
    claims or inheriting the appendix cannot supply that acknowledgement.
    """
    declared = [claim for claim in claims if isinstance(claim, dict) and claim.get("basis") == "calculation"]
    if not declared and not fact_ids and not template_fact_ids and inherited is None:
        return None
    facts = {fact["id"]: fact for fact in prepared.get("facts", [])}
    required = sorted(fact_id for fact_id in prepared.get("required_fact_ids", [])
                      if fact_id in facts and facts[fact_id].get("column") == "residual")
    if not required:
        return None
    hashes = {run["id"]: run["output_hash"] for run in runs}
    authored = set(inherited.get("authored_fact_ids", [])) if inherited and inherited.get("run_output_hashes") == hashes else set()
    authored.update(fact_id for fact_id in [*fact_ids, *template_fact_ids] if isinstance(fact_id, str) and fact_id in facts)
    locations = {(fact["run_id"], fact["row"], fact["column"]): fact for fact in facts.values()}
    for claim in declared:
        fact = facts.get(claim.get("fact_id")) if isinstance(claim.get("fact_id"), str) else None
        if not claim.get("fact_id"):
            fact = next((fact for location, fact in locations.items()
                         if location == (claim.get("run_id"), claim.get("row"), claim.get("column"))), None)
        if fact and all(claim.get(key) == fact[key] for key in ("run_id", "row", "column")):
            try:
                matches = Decimal(str(claim.get("value"))) == Decimal(str(fact["value"]))
            except ArithmeticError:
                matches = False
            if matches:
                authored.add(fact["id"])
    return {"schema_version": 1, "run_output_hashes": hashes,
            "required_fact_ids": required, "authored_fact_ids": sorted(authored & facts.keys())}
# What a session reads of a completion receipt whose messages or fixups name a run outside it: the verdict.
COMPLETION_VERDICT = (
    "id", "task_id", "schema_version", "created_at", "seq", "draft_id", "draft_content_hash",
    "completion_policy_version", "requirements_id", "status", "repair_attempts", "max_repairs", "response_status",
    "original_requirements_completed", "repair_allowed", "stop_reason", "manager_reasons", "checks",
)
# What a receipt decides from the runs it judged against the current requirements: a session reads them again when
# it may read one of those runs otherwise than whoever filed it (AnalysisStore._receipt_verdict_in_view).
RECEIPT_RUN_VERDICT = (
    "reused_run_ids", "issue_fingerprint", "status", "repair_attempts", "response_status",
    "original_requirements_completed", "repair_allowed", "stop_reason", "manager_reasons", "actions",
)
BASIS = {
    "source_fact",
    "calculation",
    "external_reference",
    "user_assumption",
    "professional_judgment",
}


def _bytes(value: Any) -> bytes:
    return json.dumps(
        json_value(value),
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        allow_nan=False,
    ).encode("utf8")


def _hash(value: Any) -> str:
    return hashlib.sha256(_bytes(value)).hexdigest()


def _draft_content_hash(**fields) -> str:
    """Hash the frozen ten-field draft identity, excluding optional authorship."""
    return _hash({key: fields[key] for key in DRAFT_CONTENT_FIELDS})


def _hash_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _export_reader(mode: str | None, allowed: set[str] | None, runs: list[dict] = ()) -> str:
    """What keeps a session's audit export apart from other readers' in the archive.

    An audit projection carries the evidence labels, which each reader reads by the files it may read
    (AnalysisStore.draft_in_view). Were a session's audit export filed under the manager's name whenever the
    two agree, getting the manager's file back would tell the session whether a newer file it may not read
    exists. So do the run lines of an export citing a SQL or structured_aggregate run, in either mode: each
    reader reads those runs' reproducibility by the runs it may read (AnalysisStore.run_in_view, ``runs`` as this
    reader reads them). Any other management projection reads the same for everyone, and export_draft renders it
    on the server, so its file is the same whoever asks.
    """
    if allowed is None:
        return ""
    shown = [[run["id"], run.get("reproducibility_status")] for run in runs if _snapshot_run(run)]
    if mode != "audit" and not shown:
        return ""
    # The run lines are no part of the projection hash: their view names the file too, so a file filed for another
    # reading of them (before sessions read runs so) is never handed back for this one.
    return _registration_reader(allowed) + (f"-r{_hash(shown)[:8]}" if shown else "")


def _registration_reader(allowed: set[str] | None) -> str:
    """What keeps a session's registered export apart from every other reader's, in any mode.

    A registration files bytes the caller rendered, not ones the server did: were they filed under a name
    another reader also files under, whoever registered first would decide what the other gets back.
    """
    return "" if allowed is None else f"-{_hash(sorted(allowed))[:12]}"






def _quality_status_dimensions(
    *,
    execution_status: str,
    review_status: str,
    quality_issues: list[dict],
    evidence_report: dict,
    consistency_report: dict,
    numeric_configured: bool = True,
) -> dict[str, dict]:
    evidence_issues = [item for item in quality_issues if str(item.get("code", "")).startswith("evidence_recipe_")]
    consistency_issues = [item for item in quality_issues if str(item.get("code", "")).startswith("financial_consistency_")]
    revision_issues = [item for item in quality_issues if str(item.get("code", "")).startswith("revision_")]
    numeric_issues = [
        item
        for item in quality_issues
        if item not in evidence_issues and item not in consistency_issues and item not in revision_issues
    ]
    evidence_status = (evidence_report.get("summary") or {}).get("status") or "not_configured"
    consistency_status = (consistency_report.get("summary") or {}).get("status") or "pass"
    source_status = "needs_attention" if consistency_status == "needs_attention" else evidence_status
    source_issues = [*evidence_issues, *consistency_issues]
    return {
        "execution": {
            "label": "执行状态",
            "status": "pass" if execution_status == "succeeded" else execution_status,
            "issues": [],
        },
        "numeric_citation": {
            "label": "数值引用",
            "status": (
                "needs_attention"
                if numeric_issues or revision_issues
                else "pass"
                if numeric_configured
                else "not_configured"
            ),
            "issues": [*numeric_issues, *revision_issues],
        },
        "source_validation": {
            "label": "来源核验",
            "status": source_status,
            "issues": source_issues,
        },
        "professional_review": {
            "label": "专业复核",
            "status": "pass" if review_status in {"reviewed", "accepted"} else "unreviewed",
            "issues": [],
        },
    }

def _fact_ids_in_runs(analysis: Any, fact_ids: list, runs: dict[str, dict], materiality: Any) -> list:
    from foundry.finance_explanations import prepare_explanation

    succeeded = [run for run in runs.values() if run.get("execution_status") == "succeeded"]
    available = {fact["id"] for fact in prepare_explanation(analysis, succeeded, materiality=materiality)["facts"]}
    # Non-string ids are kept so validate_explanation still reports them as unknown_fact.
    return [fact_id for fact_id in fact_ids if not isinstance(fact_id, str) or fact_id in available]


def _revision_quality_issues(comparison: dict) -> list[dict]:
    issues = []
    for item in comparison.get("issues", []):
        code = item.get("code")
        if code == "revision_numeric_regression":
            issues.append({
                **item,
                "severity": item.get("severity", "error"),
                "message": item.get("message") or "同源同口径业务数值发生变化，需要解释或修正。",
            })
        elif code in {"revision_basis_changed", "revision_comparison_ambiguous"}:
            issues.append({
                **item,
                "severity": item.get("severity", "warning"),
                "message": item.get("message") or "修订比较口径发生变化，需要复核。",
            })
    return issues


def _snapshot_reproducibility(succeeded: bool, deterministic: bool, previous: list[dict], output_hash: Any) -> str:
    """A SQL or structured_aggregate run's reproducibility against ``previous``, the earlier succeeded runs of its
    fingerprint: unverified with none, reproducible when every one gave the same output, else non_reproducible (as
    is a nondeterministic run)."""
    if succeeded and not deterministic:
        return "non_reproducible"
    if succeeded and previous:
        return "reproducible" if all(r["output_hash"] == output_hash for r in previous) else "non_reproducible"
    return "unverified"


def _run_order(run: dict) -> tuple:
    """The order runs were committed in (AnalysisStore._list)."""
    return run.get("seq", 0), run["created_at"], run["id"]


def _snapshot_run(run: dict) -> bool:
    """A run whose reproducibility each reader decides again (AnalysisStore.run_in_view): a SQL or
    structured_aggregate run, whose fingerprint keys its inputs by their rows and not their snapshots. A finance
    fingerprint hashes its context, so every run sharing it rests on the same files; a script run is always
    unverified."""
    return run.get("engine") in (None, "sql", "structured_aggregate") and bool(run.get("fingerprint"))


def _body_before_marks(draft: dict) -> Any:
    """The body save_draft compared with its parent's to decide ``changed``: a document's rendering before
    mark_visual_estimates added its notices (save_draft renders the body again after deciding). materialize_document
    strips every bar or line notice, so each one a stored document carries was added there."""
    document = draft.get("document")
    if not isinstance(document, dict):
        return draft.get("body")
    import copy

    from foundry.report_documents import report_document_markdown
    from foundry.report_visuals import document_visuals

    bare = copy.deepcopy(document)
    marked = False
    for visual in document_visuals(bare):
        if isinstance(visual, dict) and visual.get("kind") in {"bar", "line"} and visual.pop("estimate_notice", None):
            marked = True
    return report_document_markdown(bare) if marked else draft.get("body")


def _without_revision(check: Any) -> Any:
    """A stored check as save_draft would have composed it with no parent: the same check less its revision issues.

    Only issues set ``needs_attention``, on the check and on its numeric_citation dimension, and a revision
    issue needs cited runs, so a check left with no issue reads ``pass`` (_quality_status_dimensions).
    """
    if not isinstance(check, dict):
        return check

    def kept(issues: Any) -> list:
        return [issue for issue in issues or []
                if not (isinstance(issue, dict) and str(issue.get("code", "")).startswith("revision_"))]

    issues = kept(check.get("issues"))
    masked = {**check, "issues": issues, **({"status": "pass"} if check.get("status") == "needs_attention" and not issues else {})}
    dimensions = check.get("status_dimensions")
    numeric = dimensions.get("numeric_citation") if isinstance(dimensions, dict) else None
    if isinstance(numeric, dict):
        left = kept(numeric.get("issues"))
        masked["status_dimensions"] = {**dimensions, "numeric_citation": {
            **numeric, "issues": left, **({"status": "pass"} if numeric.get("status") == "needs_attention" and not left else {})}}
    return masked


def _receipt_runs(receipt: dict) -> set[str]:
    """Every run a completion receipt names: its hashes and reused runs, and its actions' runs and fixups."""
    runs = set(receipt.get("run_hashes") or {}) | set(receipt.get("reused_run_ids") or [])
    for action in receipt.get("actions") or []:
        if not isinstance(action, dict):
            continue
        runs |= {run_id for run_id in action.get("run_ids") or [] if isinstance(run_id, str)}
        for fixup in action.get("fixups") or []:
            if not isinstance(fixup, dict):
                continue
            location = fixup.get("location") if isinstance(fixup.get("location"), dict) else {}
            runs |= {run_id for run_id in (fixup.get("run_id"), location.get("run_id")) if isinstance(run_id, str)}
            runs |= {row["run_id"] for row in fixup.get("runs") or []
                     if isinstance(row, dict) and isinstance(row.get("run_id"), str)}
    return runs


def _rests_on_parent(draft: dict) -> bool:
    """Whether a revision took anything of its parent (INHERITED_FIELDS). One saved before revisions recorded what they
    took cannot tell, so it is read as having taken it."""
    inherited = draft.get("inherited_fields")
    return bool(draft.get("parent_id")) and (inherited is None or bool(inherited))


def _default_materiality(value: Any) -> bool:
    """Whether a draft's materiality is the one a draft given none is saved with (finance_explanations._materiality)."""
    from foundry.finance_explanations import _materiality

    try:
        return value in (None, {}) or _materiality(value) == _materiality({})
    except FoundryError:
        return False


def _inherited_fields(spec: dict, parent: dict) -> set[str]:
    """The INHERITED_FIELDS a revision saved from ``spec`` takes from ``parent``: what the spec leaves out and the
    parent carries. A parent's default title, or an empty value, is what a draft saved with no parent carries too."""
    if not parent:
        return set()
    inherited = set()
    for key in INHERITED_FIELDS:
        if key == "body":
            # The body goes with the document it is projected from.
            taken = spec.get("document") is None and "body" not in spec and (
                parent.get("document") is not None or bool(parent.get("body")))
        elif key == "materiality":
            taken = key not in spec and not _default_materiality(parent.get(key))
        else:
            taken = key not in spec and parent.get(key) not in (None, "", [], {}) and not (
                key == "title" and parent[key] == DEFAULT_DRAFT_TITLE)
        if taken:
            inherited.add(key)
    return inherited


def _engine_runs(draft: dict) -> set[str]:
    """The engine runs a draft's requirement items name (finance_requirements.requirements_in_view)."""
    runs: set[str] = set()
    for key in ("explanation_check", "quality"):
        requirements = draft.get(key, {}).get("requirements") if isinstance(draft.get(key), dict) else None
        for item in (requirements.get("items") or []) if isinstance(requirements, dict) else []:
            for run in (item.get("engine_runs") or []) if isinstance(item, dict) else []:
                if isinstance(run, dict) and isinstance(run.get("run_id"), str):
                    runs.add(run["run_id"])
    return runs


def _repair_budget(receipt: dict, history: list[dict]) -> dict:
    """What a receipt says of its repair budget, counted over ``history`` (the earlier receipts of its chain a reader
    may read) by record_completion's rules: what that reader's own save of the draft would have spent."""
    from foundry.finance_completion import MAX_REPAIRS

    maximum = receipt["max_repairs"] if isinstance(receipt.get("max_repairs"), int) else MAX_REPAIRS
    status = receipt.get("status")
    stalled = bool(history) and history[-1].get("issue_fingerprint") == receipt.get("issue_fingerprint") and status != "complete"
    exhausted = len(history) >= maximum
    if status in {"needs_review", "repair_available"}:
        status = "needs_review" if stalled or exhausted else "repair_available"
    fields = {"repair_attempts": min(len(history), maximum), "status": status,
              "repair_allowed": status in {"repair_available", "needs_requirements"} and not exhausted and not stalled}
    if status not in {"complete", "awaiting_input", "awaiting_manager"}:
        fields["stop_reason"] = "no_progress" if stalled else "repair_budget_exhausted" if exhausted else None
    return fields


def _new(task: Task, **fields) -> dict:
    return {
        "id": uuid.uuid4().hex,
        "task_id": task.id,
        "schema_version": 1,
        "created_at": now_iso(),
        **json_value(fields),
    }


def _saved_by(session_id: str | None, context_revision: int | None) -> dict:
    """The fields of an analysis revision naming the session that saved it, as a run names its session
    (AnalysisStore._author_ceiling): none for the manager's, saved as before."""
    return {"session_id": session_id, "context_revision": context_revision} if session_id else {}


def _assert_id(value: str) -> str:
    if not isinstance(value, str) or not ID.fullmatch(value):
        raise FoundryError(f"invalid analysis object id: {value}")
    return value


def _safe_text(value: Any, label: str, *, required=False, maximum=1_000_000) -> str:
    if (
        not isinstance(value, str)
        or (required and not value.strip())
        or len(value) > maximum
    ):
        raise FoundryError(
            f"{label} must be {'nonempty ' if required else ''}text (at most {maximum} characters)"
        )
    return value




def _result_role(value: Any, *, default: str = "business") -> str:
    if value in (None, ""):
        return default
    if not isinstance(value, str) or value not in RESULT_ROLES:
        raise FoundryError("result_role is unsupported")
    return value


def _script_topic(value: Any) -> str | None:
    if value in (None, ""):
        return None
    if not isinstance(value, str) or value not in SCRIPT_TOPICS:
        raise FoundryError("script topic must be collection_health")
    return value


def _scope_metadata(value: Any) -> dict:
    if value in (None, ""):
        return {}
    if not isinstance(value, dict):
        raise FoundryError("scope must be an object")
    return json_value(value)


def _requirements_id(value: Any) -> str | None:
    if value in (None, ""):
        return None
    return _assert_id(value)

def _append_excel(sheet, row):
    parts = [
        (
            [value[i : i + 32000] for i in range(0, len(value), 32000)]
            if isinstance(value, str) and len(value) > 32000
            else [value]
        )
        for value in row
    ]
    for index in range(max((len(part) for part in parts), default=1)):
        sheet.append([part[index] if index < len(part) else None for part in parts])


def _script_result_path(value: Any) -> str | None:
    if value is None:
        return None
    path = _safe_text(value, "script result_file", required=True, maximum=1000)
    relative = Path(path)
    if relative.is_absolute() or ".." in relative.parts:
        raise FoundryError("script result_file must be a relative output path")
    if relative.suffix.casefold() not in {".csv", ".json"}:
        raise FoundryError("script result_file must be a CSV or JSON table")
    return relative.as_posix()


def _empty_script_result() -> dict:
    return {"columns": [], "rows": [], "row_count": 0, "truncated": False}


class AnalysisStore:
    def __init__(self, task: Task):
        self.task = task
        self.root = task.root / "analysis"
        if self.root.is_symlink():
            raise FoundryError("analysis root cannot use symbolic links")
        if task.root.is_symlink() or not task.root.resolve().is_relative_to(
            task.store.tasks_dir.resolve()
        ):
            raise FoundryError("analysis task path escapes task store")
        self._inside(self.root)

    def _inside(self, path: Path) -> Path:
        if not path.resolve().is_relative_to(self.task.root.resolve()):
            raise FoundryError("analysis path escapes this task")
        current = path
        while True:
            if current.is_symlink():
                raise FoundryError("analysis paths cannot use symbolic links")
            if current.resolve() == self.task.root.resolve():
                break
            if current == current.parent:
                raise FoundryError("analysis path has no task ancestor")
            current = current.parent
        return path

    def _object_path(self, kind: str, object_id: str) -> Path:
        _assert_id(object_id)
        if kind not in KINDS:
            raise FoundryError("unknown analysis object kind")
        return self._inside(self.root / kind / (object_id + ".json"))

    def _read(self, path: Path) -> dict:
        self._inside(path)
        try:
            record = json.loads(path.read_text(encoding="utf8"))
        except (OSError, ValueError) as exc:
            raise FoundryError(f"analysis object is unavailable: {path.name}") from exc
        check = dict(record)
        digest = check.pop("_integrity_hash", None)
        if (
            not digest
            or digest != _hash(check)
            or record.get("task_id") != self.task.id
        ):
            raise FoundryError("analysis object integrity check failed")
        return record

    def _write(self, path: Path, record: dict) -> dict:
        self._inside(path)
        record = json_value(record)
        with exclusive(self.task.lock_path, timeout=10):
            return self._write_locked(path, record)

    def _write_locked(self, path: Path, record: dict) -> dict:
        """Commit an immutable object while the caller holds the task lock."""
        self._inside(path)
        if path.exists():
            raise FoundryError("immutable analysis object already exists")
        record = json_value(record)
        record["seq"] = self._next_seq()
        record["_integrity_hash"] = _hash(record)
        write_json_atomic(path, record)
        return record

    def _next_seq(self) -> int:
        # Called only while holding the task lock; gaps after interruption are
        # acceptable, duplicate/reordered committed versions are not.
        path = self._inside(self.root / "sequence.json")
        if path.exists():
            try:
                value = json.loads(path.read_text())["seq"]
                if isinstance(value, bool) or not isinstance(value, int) or value < 0:
                    raise ValueError("invalid sequence")
            except (OSError, ValueError, KeyError, TypeError) as exc:
                raise FoundryError("analysis sequence record is invalid") from exc
        else:
            value = 0
        value += 1
        write_json_atomic(path, {"task_id": self.task.id, "seq": value})
        return value

    def _get(self, kind: str, object_id: str) -> dict:
        return self._read(self._object_path(kind, object_id))

    def _list(self, kind: str, *, read: dict[Path, dict] | None = None) -> list[dict]:
        """Every record of ``kind`` in commit order. ``read``, when given, maps each path already read and checked to
        its record and takes the ones read now: records are immutable, so listing again with it reads only those
        committed since."""
        root = self._inside(self.root / kind)
        paths = root.glob("*.json") if root.exists() else []
        read = {} if read is None else read
        for path in paths:
            if path not in read:
                read[path] = self._read(path)
        return sorted(read.values(), key=lambda r: (r.get("seq", 0), r["created_at"], r["id"]))

    def _get_method(self, method_id: str) -> dict:
        _assert_id(method_id)
        for task in [
            self.task,
            *[task for task in self.task.store.list("all") if task.id != self.task.id],
        ]:
            source = AnalysisStore(task)
            path = source._object_path("methods", method_id)
            if path.exists():
                return source._read(path)
        raise FoundryError("saved analysis method does not exist")

    def _methods(self) -> list[dict]:
        # Only explicitly saved, sanitized methods cross task boundaries. This
        # lookup never reads another task's snapshots, revisions, or run outputs.
        methods = [
            method
            for task in self.task.store.list("all")
            for method in AnalysisStore(task)._list("methods")
        ]
        return sorted(
            methods,
            key=lambda row: (row["created_at"], row["task_id"], row.get("seq", 0)),
        )

    def published_datasets(self) -> list[dict]:
        """Every published dataset record of the task, oldest first, without reading its rows."""
        root = self._inside(self.root / "datasets")
        datasets = (
            [record for path in root.glob("*/meta.json")
             if self._dataset_published(record := self._read(path))]
            if root.exists()
            else []
        )
        return sorted(datasets, key=lambda r: (r.get("seq", 0), r["created_at"], r["id"]))

    def state(self, allowed: set[str] | None = None, *, cache: dict | None = None) -> dict:
        """Every record of the task. ``allowed``, a session's snapshot ceiling, keeps the requirements versions it
        may read, each as it reads it (finance_requirements.requirements_version_view), reads each saved analysis and
        run on files it may read as it reads them (analysis_in_view, run_in_view), and runs are read against the
        latest of those versions as it reads it (as analysis_requirements_list names its current one); the caller
        filters the rest. ``cache`` as for run_in_view, so the drafts and receipts of the same read share it."""
        from foundry.result_semantics import declared_runs

        requirements = self._list("requirements")
        shown = requirements
        analyses, runs = self._list("revisions"), self._list("runs")
        if allowed is not None:
            from foundry.finance_requirements import requirements_readable, requirements_version_view

            cache = {} if cache is None else cache
            cache.setdefault("requirements", requirements)
            cache.setdefault("runs", {run["id"]: run for run in runs})
            cache.setdefault("analyses", {row["id"]: row for row in analyses})
            versions = {row.get("id"): row for row in requirements}
            requirements = [row for row in requirements if requirements_readable(row, allowed)]
            drafts: dict = {}
            shown = [requirements_version_view(self, row, allowed, versions=versions, drafts=drafts)
                     for row in requirements]

            def readable(row: dict) -> bool:
                return set((row.get("bindings") or {}).values()) <= set(allowed)

            # Only the records on files it may read: the caller leaves out the rest, which no view need read.
            analyses = [self.analysis_in_view(row, allowed, cache=cache) if readable(row) else row for row in analyses]
            runs = [self.run_in_view(run, allowed, cache=cache) if readable(run) else run for run in runs]
        from foundry.finance_controls import superseded_snapshots

        datasets = self.published_datasets()
        return {
            "schema_version": 1,
            "task_id": self.task.id,
            "datasets": [self.dataset_display_view(row) for row in datasets],
            # Snapshots a newer file or a published source refresh has replaced: a result that read one is out of
            # date until it is calculated again. The rule the controls and the close status use.
            "superseded_datasets": sorted(superseded_snapshots(
                self.task, datasets, allowed=allowed,
                # A reader's batches come from the listing its drafts and runs share (_batches_in_view).
                batches=None if allowed is None else self._batches_in_view(set(allowed), cache))),
            "analyses": analyses,
            # Row declarations are read as the workbench shows them: recomputed when stored under older
            # rules, and with the latest requirements' estimates added. Stored records are never rewritten.
            "runs": declared_runs(
                self, runs, shown[-1] if shown else None
            ),
            "drafts": self._list("drafts"),
            "requirements": shown,
            "evidence_checks": self._list("evidence_checks"),
            "completions": self._list("completions"),
            "methods": self._methods(),
        }

    def add_dataset(
        self,
        path: Path,
        *,
        name: str,
        metadata: dict | None = None,
        allow_duplicate: bool = False,
        allowed: set[str] | None = None,
    ) -> dict:
        """File these bytes as a dataset. ``allowed``, a session's snapshot ceiling, keeps the earlier uploads it
        may read for the same-bytes answer (_same_source); None is the manager, who is answered from every upload."""
        path = Path(path)
        _safe_text(name, "dataset name", required=True, maximum=200)
        if metadata is not None and not isinstance(metadata, dict):
            raise FoundryError("dataset metadata must be an object")
        metadata = json_value(metadata or {})
        suffix = path.suffix.lower()
        if suffix not in SUPPORTED:
            raise FoundryError(
                "unsupported format .xls (the Excel 97-2003 format): save it as .xlsx first" if suffix == ".xls"
                else f"unsupported format {suffix}; supported: CSV, XLSX, JSON, Markdown and TXT"
            )
        try:
            if not path.exists():
                raise FoundryError(f"dataset source not found: {path.name}")
            if not path.is_file() or path.stat().st_size > MAX_SOURCE_BYTES:
                raise FoundryError("dataset source must be a file of at most 32 MiB")
            with path.open("rb") as handle:
                data = handle.read(MAX_SOURCE_BYTES + 1)
        except OSError as exc:
            raise FoundryError(f"cannot read dataset source: {exc}") from exc
        source_hash = _hash_bytes(data)
        if not allow_duplicate:
            existing = self._same_source(source_hash, allowed)
            if existing is not None:
                # The manager already has these bytes. Hand back the copy she has,
                # marked as the same upload, rather than filing a second one she
                # then has to tell apart.
                return {**existing, "duplicate_of": existing["id"]}
        parsed = parse_source(data, suffix, metadata)
        for paragraph in parsed.get("paragraphs", []):
            paragraph["source_hash"] = source_hash
        encoded = _bytes(parsed)
        extra: dict[str, Any] = {}
        if suffix == ".xlsx":
            # Hints only, read before the lock: check_workbook never raises, so
            # nothing it finds can refuse or delay the import beyond its own scan.
            from foundry.workbook_health import check_workbook

            extra["workbook_health"] = check_workbook(data)
        record = _new(
            self.task,
            name=name,
            kind=parsed["kind"],
            source_name=path.name,
            source_hash=source_hash,
            source_bytes=len(data),
            data_hash=_hash_bytes(encoded),
            schema=parsed["schema"],
            row_count=parsed["row_count"],
            metadata=metadata,
            transformation=parsed["transformation"],
            source_file="source" + suffix,
            **extra,
        )
        folder = self._inside(self.root / "datasets" / record["id"])
        with exclusive(self.task.lock_path, timeout=10):
            self._inside(folder)
            folder.mkdir(parents=True, exist_ok=False)
            (folder / record["source_file"]).write_bytes(data)
            (folder / "data.json").write_bytes(encoded)
            # Commit metadata last: readers never observe a partially imported dataset.
            record["seq"] = self._next_seq()
            record["_integrity_hash"] = _hash(record)
            write_json_atomic(folder / "meta.json", record)
        return record

    def _same_source(self, source_hash: str, allowed: set[str] | None = None) -> dict | None:
        """An earlier upload of exactly these bytes, if the task has one the caller may read.

        A table derived from a confirmed structure carries its original's bytes
        too, so it is skipped here: it is the same file seen through a structure,
        not a second upload of it. An upload outside ``allowed`` (a session's
        snapshot ceiling) is no part of what the session reads: its bytes are
        filed as the session's own, exactly as when no one holds them.
        """
        for row in self.state()["datasets"]:
            if allowed is not None and row["id"] not in allowed:
                continue
            if row.get("source_hash") != source_hash:
                continue
            if (row.get("transformation") or {}).get("structure"):
                continue
            return row
        return None

    def derive_table_snapshot(self, source_dataset_id: str, declaration: dict) -> dict:
        """Read the table a declaration describes without filing it, so a confirmation of many
        sheets can read every one before it files any. Pass the result to add_table_snapshot."""
        from foundry.structure_snapshot import derive_table

        if not isinstance(declaration, dict):
            raise FoundryError("structure declaration must be an object")
        shell = self.read_dataset(source_dataset_id)
        folder = self._inside(self.root / "datasets" / source_dataset_id)
        data = (folder / shell["source_file"]).read_bytes()
        if _hash_bytes(data) != shell["source_hash"]:
            raise FoundryError("dataset integrity check failed")
        declaration = json_value({**declaration, "derived_from": source_dataset_id})
        return {"source_dataset_id": source_dataset_id, "shell": shell, "data": data,
                "declaration": declaration, "parsed": derive_table(data, declaration)}

    def add_table_snapshot(
        self, source_dataset_id: str, declaration: dict, *, name: str, derived: dict | None = None
    ) -> dict:
        """File the calculable table a confirmed structure declaration describes.

        The original workbook bytes are stored again, so the snapshot answers for
        its own source and a located figure can still be opened in the file the
        manager uploaded. ``derived`` is derive_table_snapshot's reading of the same declaration.
        """
        _safe_text(name, "dataset name", required=True, maximum=200)
        if derived is None or derived["source_dataset_id"] != source_dataset_id:
            derived = self.derive_table_snapshot(source_dataset_id, declaration)
        shell, data, declaration, parsed = derived["shell"], derived["data"], derived["declaration"], derived["parsed"]
        encoded = _bytes(
            {
                "kind": parsed["kind"],
                "schema": parsed["schema"],
                "rows": parsed["rows"],
                "row_count": parsed["row_count"],
                "transformation": parsed["transformation"],
            }
        )
        record = _new(
            self.task,
            name=name,
            kind="table",
            source_name=shell["source_name"],
            source_hash=shell["source_hash"],
            source_bytes=len(data),
            data_hash=_hash_bytes(encoded),
            schema=parsed["schema"],
            row_count=parsed["row_count"],
            metadata={"derived_from": source_dataset_id, "sheet": declaration["sheet"]},
            transformation=parsed["transformation"],
            source_file=shell["source_file"],
        )
        destination = self._inside(self.root / "datasets" / record["id"])
        with exclusive(self.task.lock_path, timeout=10):
            self._inside(destination)
            destination.mkdir(parents=True, exist_ok=False)
            (destination / record["source_file"]).write_bytes(data)
            (destination / "data.json").write_bytes(encoded)
            record["seq"] = self._next_seq()
            record["_integrity_hash"] = _hash(record)
            write_json_atomic(destination / "meta.json", record)
        return record

    def read_dataset(self, dataset_id: str) -> dict:
        _assert_id(dataset_id)
        folder = self._inside(self.root / "datasets" / dataset_id)
        metadata = self._read(folder / "meta.json")
        if not self._dataset_published(metadata):
            raise FoundryError("source batch is not published; snapshot is unavailable")
        data_path = self._inside(folder / "data.json")
        source_path = self._inside(folder / metadata["source_file"])
        try:
            data = data_path.read_bytes()
            source = source_path.read_bytes()
            if (
                _hash_bytes(data) != metadata["data_hash"]
                or _hash_bytes(source) != metadata["source_hash"]
            ):
                raise FoundryError("dataset integrity check failed")
            return self.dataset_display_view({**metadata, **json.loads(data)})
        except (OSError, ValueError) as exc:
            raise FoundryError(
                "dataset integrity check failed: missing or invalid snapshot"
            ) from exc

    def dataset_display_view(self, dataset: dict) -> dict:
        """Display labels never replace the source name used by financial lineage."""
        records = self._list("dataset_display_names")
        latest = next((row for row in reversed(records) if row["dataset_id"] == dataset["id"]), None)
        if latest is None:
            return dataset
        return {**dataset, "display_name": latest["display_name"], "display_name_revision": latest["id"]}

    def rename_dataset_display(self, dataset_id: str, name: str) -> dict:
        name = _safe_text(name, "display name", required=True, maximum=200)
        dataset = self.read_dataset(dataset_id)
        with exclusive(self.task.lock_path, timeout=10):
            previous = self.dataset_display_view(dataset)
            record = _new(self.task, dataset_id=dataset_id, display_name=name,
                          parent_id=previous.get("display_name_revision"),
                          source_hash=dataset["source_hash"], data_hash=dataset["data_hash"])
            self._write_locked(self._object_path("dataset_display_names", record["id"]), record)
        return self.read_dataset(dataset_id)

    def _dataset_published(self, metadata: dict) -> bool:
        batch_id = metadata.get("metadata", {}).get("source_refresh", {}).get("source_batch_id")
        if not batch_id:
            return True
        batch_path = self._inside(self.root / "source_batches" / f"{_assert_id(batch_id)}.json")
        if not batch_path.exists():
            return False
        batch = self._read(batch_path)
        if batch.get("publication_status") != "published" or batch.get("dataset_id") != metadata["id"]:
            raise FoundryError("source batch publication does not match snapshot")
        return True

    def save_analysis(self, spec: dict, *, session_id: str | None = None,
                      context_revision: int | None = None) -> dict:
        """Save an analysis revision. A session's carries its ``session_id`` and ``context_revision``, as its runs do
        (_author_ceiling); the manager's is saved as before."""
        if not isinstance(spec, dict):
            raise FoundryError("analysis spec must be an object")
        question = _safe_text(
            spec.get("question"), "question", required=True, maximum=20000
        )
        engine = spec.get("engine", "sql")
        if not isinstance(engine, str) or engine not in {"sql", "structured_aggregate"}:
            raise FoundryError("analysis engine must be sql or structured_aggregate")
        structured = engine == "structured_aggregate"
        sql = _safe_text(spec.get("sql", ""), "SQL", required=not structured, maximum=100000)
        bindings = spec.get("bindings") or {}
        if not isinstance(bindings, dict) or not bindings or len(bindings) > 32:
            raise FoundryError("analysis needs 1–32 explicit snapshot bindings")
        if len({str(a).casefold() for a in bindings}) != len(bindings):
            raise FoundryError("binding aliases must be unique (case insensitive)")
        for alias, dataset_id in bindings.items():
            if not isinstance(alias, str) or not ALIAS.fullmatch(alias):
                raise FoundryError(f"invalid dataset alias: {alias}")
            self.read_dataset(dataset_id)
        aggregate_plan = None
        if structured:
            from foundry.analysis_aggregate import validate_aggregate_plan
            plan = spec.get("aggregate_plan")
            if (not isinstance(plan, dict) or len(bindings) != 1
                    or not isinstance(plan.get("dataset_alias"), str)
                    or plan["dataset_alias"] not in bindings):
                raise FoundryError("structured aggregate needs exactly its single snapshot binding")
            if sql or spec.get("parameters") or spec.get("parameter_metadata") or spec.get("relationships"):
                raise FoundryError("structured aggregate uses its plan, not SQL, parameters or relationships")
            aggregate_plan = validate_aggregate_plan(plan, self.read_dataset(bindings[plan['dataset_alias']]))
        elif spec.get("aggregate_plan") is not None:
            raise FoundryError("aggregate_plan requires structured_aggregate engine")
        parent_id = spec.get("parent_id")
        if parent_id:
            self.get_analysis(parent_id)
        requirements_id = _requirements_id(spec.get("requirements_id"))
        if requirements_id:
            self._get("requirements", requirements_id)
        result_role = _result_role(spec.get("result_role"))
        scope = _scope_metadata(spec.get("scope"))
        parameters = spec.get("parameters") or []
        if not isinstance(parameters, (list, dict)):
            raise FoundryError("SQL parameters must be an array or object")
        parameter_metadata = [] if structured else normalize_parameter_metadata(
            sql, parameters, spec.get("parameter_metadata")
        )
        period_binding = normalize_period_binding(spec.get("period_binding"), engine=engine, sql=sql,
                                                  parameters=parameters, aggregate_plan=aggregate_plan)
        timeout = spec.get("timeout_seconds", 15)
        max_rows = spec.get("max_rows", 1000)
        if (
            isinstance(timeout, bool)
            or not isinstance(timeout, (int, float, Decimal))
            or not 0.01 <= timeout <= 60
        ):
            raise FoundryError("timeout_seconds must be between 0.01 and 60")
        if (
            isinstance(max_rows, bool)
            or not isinstance(max_rows, int)
            or not 1 <= max_rows <= 10000
        ):
            raise FoundryError("max_rows must be between 1 and 10000")
        timeout = float(timeout)
        fields = {}
        for key in [
            "assumptions",
            "relationships",
            "metrics",
            "grain",
            "missing_inputs",
            "conversion_rules",
        ]:
            value = spec.get(key, [])
            if not isinstance(value, list):
                raise FoundryError(f"{key} must be an array")
            fields[key] = value
        fields["relationships"] = validate_relationship_contract(
            fields["relationships"], set(bindings)
        )
        record = _new(
            self.task,
            **({"engine": engine, "aggregate_plan": aggregate_plan} if structured else {}),
            question=question,
            sql=sql,
            bindings=bindings,
            parameters=parameters,
            parameter_metadata=parameter_metadata,
            **({"period_binding": period_binding} if period_binding else {}),
            parent_id=parent_id,
            requirements_id=requirements_id,
            result_role=result_role,
            scope=scope,
            max_rows=max_rows,
            timeout_seconds=timeout,
            method_id=spec.get("method_id"),
            **_saved_by(session_id, context_revision),
            **fields,
        )
        if record["method_id"]:
            self._get_method(record["method_id"])
        return self._write(self._object_path("revisions", record["id"]), record)

    def get_analysis(self, analysis_id: str) -> dict:
        return self._get("revisions", analysis_id)

    def get_run(self, run_id: str) -> dict:
        run = self._get("runs", run_id)
        for output in run.get("outputs", []):
            path = self._inside(self.root / output["artifact_path"])
            if not path.is_file() or _hash_bytes(path.read_bytes()) != output["sha256"]:
                raise FoundryError("run output integrity check failed")
        return run

    def verified_run(self, run_id: str, allowed: set[str] | None, cache: dict) -> dict:
        """get_run, the run's outputs verified once a session read (``cache``, as for run_in_view), however many drafts
        and receipts citing it the read judges again as its twin would: a session's completion save does so under
        completion.lock. The manager verifies them on every call, as it always did."""
        if allowed is None:
            return self.get_run(run_id)
        verified = cache.setdefault("verified_runs", {})
        if run_id not in verified:
            verified[run_id] = self.get_run(run_id)
        return verified[run_id]

    def run_in_view(self, run: dict, allowed: set[str] | None, *, cache: dict | None = None) -> dict:
        """A stored run as a reader whose snapshot ceiling is ``allowed`` reads it; None (the manager) reads it as stored.

        A session reads a run as it would have saved it from the files it may read: pinned to the requirements version
        _pin_in_view names, a script run's row declarations read again against that version
        (result_semantics.declare_against) or, declared with a fingerprint of that version other than its own, with
        its own (_captured_otherwise), and its reproducibility as _reproducibility_in_view reads it. No
        run a session reads carries the integrity hash: it covers the stored record, which a view may not be, and its
        absence must not tell a view from a stored record. ``cache``, shared by the runs, drafts and receipts of one
        read, holds the task's requirements versions, their fingerprints, its runs as listed and as verified
        (verified_run), saved analyses, sessions and refresh batches, the row declarations read again and the
        requirement support of the receipts judged again.
        """
        if allowed is None:
            return run
        cache = {} if cache is None else cache
        view = {key: value for key, value in run.items() if key != "_integrity_hash"}
        fields, substitute = self._pin_in_view(run, allowed, cache) or ({}, None)
        view.update(fields)
        declaration = run.get("row_declarations")
        if "requirements_id" in fields and isinstance(declaration, dict):
            from foundry.result_semantics import declare_against

            # Runs and versions are immutable: a read declares each run's result again once, however many drafts and
            # receipts cite the run, and hands each its own copy.
            declarations = cache.setdefault("declarations", {})
            key = (run["id"], (substitute or {}).get("id"))
            if key not in declarations:
                # Keyed as a stored declaration reads back (write_json_atomic sorts keys), so the view reads as one.
                declarations[key] = json.dumps(json_value(declare_against(
                    self, run, substitute, self.fingerprint_in_view(substitute, allowed, cache=cache))), sort_keys=True)
            view["row_declarations"] = json.loads(declarations[key])
        elif isinstance(declaration, dict):
            declared = self._requirements_version(declaration.get("requirements_id"), cache)
            if self._captured_otherwise(declared, declaration.get("requirements_fingerprint"), allowed, cache):
                view["row_declarations"] = {**declaration, "requirements_fingerprint": self.fingerprint_in_view(
                    declared, allowed, cache=cache)}
        if "reproducibility_status" in run:
            view["reproducibility_status"] = self._reproducibility_in_view(run, allowed, cache)
        return view

    def analysis_in_view(self, analysis: dict, allowed: set[str] | None, *, cache: dict | None = None) -> dict:
        """A saved analysis as a reader whose snapshot ceiling is ``allowed`` reads it; None (the manager) reads it as
        stored. A session reads one naming a requirements version it may not read as naming the version the runs made
        from it read as pinned to (_pin_in_view), one the session that saved it may read too, so a listing never shows
        both, and without the integrity hash, as it reads runs."""
        if allowed is None:
            return analysis
        view = {key: value for key, value in analysis.items() if key != "_integrity_hash"}
        pin = self._pin_in_view(analysis, allowed, {} if cache is None else cache)
        if pin is not None and "requirements_id" in pin[0]:
            view["requirements_id"] = pin[0]["requirements_id"]
        return view

    def _pin_in_view(self, record: dict, allowed: set[str], cache: dict) -> tuple[dict, dict | None] | None:
        """The pin a run or saved analysis reads with to a reader whose snapshot ceiling is ``allowed``: the fields that
        read otherwise than stored, and the version it reads as pinned to; None when it reads as stored.

        A run is pinned to the version current when it started (capture_requirements), or to the one its saved analysis
        names, with the verified fingerprint of the version current when it started (_versions_at_start). Of a run a
        session saved, only the versions that session may read count (_author_ceiling): it captured the newest of those
        and named no other. A reader that may not read the pinned version reads the record as saved where that version
        was never written: pinned to the newest version older than the pinned one it may read, with the fingerprint of
        the newest it may read of those saved when the run started, which capture_requirements would have taken then;
        a run with no older one is pinned to that one, as a run of an analysis read as naming none is. So does one that
        may read the pin but not the version current when the run started (an analysis naming an older version). One
        that may read that version reads the fingerprint captured, as stored or as the run's session read the version,
        as the one it reads that version with (_captured_otherwise). Each fingerprint is the one the reader reads
        (fingerprint_in_view), and nothing of a version it may not read shows or decides anything. Of an analysis a
        session saved, too, only the versions that session may read count: it named no other. So, of a run pinned to the
        version its analysis names, the version it reads as pinned to in that one's place is one the session that saved
        the analysis may read too (_namer_ceiling), while its fingerprint is still the run's own.
        """
        from foundry.finance_requirements import requirements_readable

        pinned = record.get("requirements_id")
        if not isinstance(pinned, str) or not pinned:
            return None
        versions = self._requirements_versions(cache)
        version = next((row for row in versions if row.get("id") == pinned), None)
        if version is None:
            return None
        author = self._author_ceiling(record, cache)
        if author is not None:
            versions = [row for row in versions if requirements_readable(row, author)]
        if requirements_readable(version, allowed):
            if "requirements_semantic_fingerprint" not in record:
                return None
            before = self._versions_at_start(record, versions, author, cache)
            if not before:
                return None
            if before[-1].get("id") == pinned or requirements_readable(before[-1], allowed):
                then = before[-1]
                if not self._captured_otherwise(then, record["requirements_semantic_fingerprint"], allowed, cache):
                    return None
            else:
                then = next((row for row in reversed(before) if requirements_readable(row, allowed)), None)
            fingerprint = self.fingerprint_in_view(then, allowed, cache=cache)
            if fingerprint == record.get("requirements_semantic_fingerprint"):
                return None
            return {"requirements_semantic_fingerprint": fingerprint}, version
        # Where the pin was never saved, a run of an analysis naming it ran what that analysis named instead.
        namer = self._namer_ceiling(record, pinned, cache)
        named = versions if namer is None else [row for row in versions if requirements_readable(row, namer)]
        substitute = next((row for row in reversed(named)
                           if row.get("seq", 0) < version.get("seq", 0) and requirements_readable(row, allowed)), None)
        then = next((row for row in reversed(self._versions_at_start(record, versions, author, cache))
                     if requirements_readable(row, allowed)), None)
        if substitute is None and "requirements_semantic_fingerprint" in record:
            # Where the pin was never saved, the run's analysis named none to whoever ran it (analysis_in_view), and the
            # run took the version current then.
            substitute = then
        return ({"requirements_id": substitute["id"] if substitute else None,
                 "requirements_semantic_fingerprint": self.fingerprint_in_view(then, allowed, cache=cache)}, substitute)

    def _versions_at_start(self, record: dict, versions: list[dict], author: set[str] | None,
                           cache: dict) -> list[dict]:
        """Of ``versions``, oldest first, those saved before run ``record`` started. A run captures its version when it
        starts (capture_requirements) but takes its seq when it commits, and versions may be saved while it runs: the
        one it started under is the newest saved before it whose fingerprint, as stored or as the session that saved
        the run (``author``, its snapshot ceiling) reads it, is the one it captured. Without one, every one saved
        before it."""
        before = [row for row in versions if row.get("seq", 0) < record.get("seq", 0)]
        captured = record.get("requirements_semantic_fingerprint")
        if captured:
            # The session reads fingerprints of its own: only the versions are shared.
            read = None if author is None else cache.setdefault("authors", {}).setdefault(frozenset(author), {
                "requirements": self._requirements_versions(cache), "versions": self._requirements_by_id(cache)})
            for index in range(len(before) - 1, -1, -1):
                if before[index].get("semantic_fingerprint") == captured or (
                        read is not None and self.fingerprint_in_view(before[index], author, cache=read) == captured):
                    return before[:index + 1]
        return before

    def _namer_ceiling(self, record: dict, pinned: str, cache: dict) -> set[str] | None:
        """The snapshot ceiling of the session that saved the analysis run ``record`` ran (_author_ceiling) where that
        analysis names ``pinned``, the run's pin: that session named only a version it may read. None for any other
        record, a run pinned otherwise, or an analysis the manager saved or a session saved before analyses named their
        session. Each analysis is read once per ``cache``."""
        analysis_id = record.get("analysis_id")
        if not isinstance(analysis_id, str):
            return None
        analyses = cache.setdefault("analyses", {})
        if analysis_id not in analyses:
            try:
                analyses[analysis_id] = self.get_analysis(analysis_id)
            except FoundryError:
                analyses[analysis_id] = None
        analysis = analyses[analysis_id]
        if analysis is None or analysis.get("requirements_id") != pinned:
            return None
        return self._author_ceiling(analysis, cache)

    def _author_ceiling(self, record: dict, cache: dict) -> set[str] | None:
        """The snapshot ceiling of the session that saved run or analysis ``record``, at the context revision it saved
        it under (Task.update_session_context keeps each earlier one), or None: the manager's, a session's the task does
        not know, or one whose recorded ceiling does not hold the record's files. A run filed under an MCP grant reads
        by its session's ceiling, which holds the grant's. The task's sessions are read once per ``cache``."""
        session_id = record.get("session_id")
        if not isinstance(session_id, str) or not session_id:
            return None
        sessions = cache.setdefault("sessions", {})
        if session_id not in sessions:
            sessions.update((row.get("session_id"), row) for row in self.task.sessions())
            sessions.setdefault(session_id, None)
        row = sessions[session_id]
        if row is None:
            return None
        state = next((past for past in row.get("context_history") or []
                      if past.get("context_revision") == record.get("context_revision")), row)
        ceiling = set(state.get("snapshot_ids") or [])
        # A session runs only on files it may read: one recorded before sessions had a ceiling ran under none.
        return ceiling if set((record.get("bindings") or {}).values()) <= ceiling else None

    def fingerprint_in_view(self, record: dict | None, allowed: set[str] | None, *,
                            cache: dict | None = None) -> str | None:
        """Requirements version ``record``'s verified semantic fingerprint as a reader whose snapshot ceiling is
        ``allowed`` reads it (finance_requirements.semantics_in_view); None (the manager) reads the stored one, as does
        every reader for no version. ``cache`` as for run_in_view: a session reads each version's once."""
        from foundry.finance_requirement_semantics import verified_requirement_fingerprint

        if allowed is None or not record:
            return verified_requirement_fingerprint(self, record)
        cache = {} if cache is None else cache
        fingerprints = cache.setdefault("fingerprints", {})
        if record["id"] not in fingerprints:
            fingerprints[record["id"]] = verified_requirement_fingerprint(
                self, self._semantics_in_view(record, allowed, cache))
        return fingerprints[record["id"]]

    def _semantics_in_view(self, record: dict, allowed: set[str], cache: dict) -> dict:
        """finance_requirements.semantics_in_view of requirements version ``record``, read once per ``cache``."""
        from foundry.finance_requirements import semantics_in_view

        semantics = cache.setdefault("semantics", {})
        if record["id"] not in semantics:
            view = semantics_in_view(self, record, allowed, versions=self._requirements_by_id(cache))
            # None where it reads as stored, so whichever copy of the version a caller holds comes back as itself.
            semantics[record["id"]] = None if view is record else view
        return semantics[record["id"]] or record

    def _captured_otherwise(self, version: dict | None, captured: Any, allowed: set[str], cache: dict) -> bool:
        """Whether ``captured``, the fingerprint a run, declaration or receipt captured for requirements ``version`` as
        whoever filed it read that version, reads to a reader whose snapshot ceiling is ``allowed`` as the one it reads
        that version with (fingerprint_in_view): when it may read the version and that one is another. A session files
        the one it reads, which may rest on versions only it may read, so the reader reads any other than its own as its
        own, whoever filed it, and nothing another session read shows. A record with none captured stands. Telling
        reads no file where the record captured the stored one."""
        from foundry.finance_requirements import requirements_readable

        if version is None or not isinstance(captured, str) or not requirements_readable(version, allowed):
            return False
        if captured == version.get("semantic_fingerprint"):
            return self._semantics_in_view(version, allowed, cache) is not version
        return captured != self.fingerprint_in_view(version, allowed, cache=cache)

    def _reproducibility_in_view(self, run: dict, allowed: set[str], cache: dict) -> Any:
        """A run's reproducibility as a session whose snapshot ceiling is ``allowed`` reads it.

        A SQL or structured_aggregate run (_snapshot_run) was compared, when saved, with every earlier run of its
        fingerprint its saver could read, and the fingerprint keys inputs by their rows, so those may rest on files the
        reader may not read (the same rows filed again). A session reads its reproducibility as run_analysis would have
        saved it for the session: against the earlier succeeded runs of the fingerprint it may read, whoever saved the
        run. Any other run reads as stored."""
        if not _snapshot_run(run):
            return run.get("reproducibility_status")
        if "runs_by_fingerprint" not in cache:
            index: dict[str, list[dict]] = {}
            for row in cache["runs"].values() if "runs" in cache else self._list("runs"):
                if row.get("fingerprint") and row.get("execution_status") == "succeeded":
                    index.setdefault(row["fingerprint"], []).append(row)
            cache["runs_by_fingerprint"] = index
        order = _run_order(run)
        earlier = [row for row in cache["runs_by_fingerprint"].get(run["fingerprint"], [])
                   if _run_order(row) < order and set((row.get("bindings") or {}).values()) <= set(allowed)]
        return _snapshot_reproducibility(run.get("execution_status") == "succeeded", run.get("deterministic", True),
                                         earlier, run.get("output_hash"))

    def _requirements_versions(self, cache: dict) -> list[dict]:
        """Every requirements version of the task, oldest first, read once per ``cache``."""
        if "requirements" not in cache:
            cache["requirements"] = self._list("requirements")
        return cache["requirements"]

    def _requirements_by_id(self, cache: dict) -> dict:
        """The task's requirements versions by id (finance_requirements._version_record), read once per ``cache``."""
        if "versions" not in cache:
            cache["versions"] = {row.get("id"): row for row in self._requirements_versions(cache)}
        return cache["versions"]

    def _listed_versions(self, cache: dict) -> dict:
        """_requirements_by_id for a lineage walk (finance_requirements._version_record). Where a version does not
        read, an empty map: the walk then reads each version it needs by _get and passes over one that does not read,
        so a damaged version outside a draft's lineage fails no read of it."""
        try:
            return self._requirements_by_id(cache)
        except FoundryError:
            return {}

    def _requirements_version(self, version_id: Any, cache: dict) -> dict | None:
        """The requirements version ``version_id`` names (_requirements_by_id), or None."""
        return self._requirements_by_id(cache).get(version_id) if isinstance(version_id, str) else None

    def _batches_in_view(self, allowed: set[str], cache: dict) -> list[dict]:
        """The task's refresh batches as a reader whose snapshot ceiling is ``allowed`` reads them
        (SourceRefreshService.batches_in_view), read once per ``cache``."""
        if "batches_in_view" not in cache:
            from foundry.source_refresh import SourceRefreshService

            cache["batches_in_view"] = SourceRefreshService(self.task, allowed=set(allowed)).batches_in_view()
        return cache["batches_in_view"]

    def _stored_run(self, run_id: Any, cache: dict) -> dict:
        """A run record as stored, from ``cache``'s listing when it holds one, without reading its outputs (get_run)."""
        runs = cache.get("runs")
        if runs is not None and isinstance(run_id, str) and run_id in runs:
            return runs[run_id]
        return self._get("runs", run_id)

    def aggregate_sources_verified(self, run: dict) -> bool:
        """Validate the persisted executor boundary, never caller source claims."""
        if run.get("engine") != "structured_aggregate" or run.get("execution_status") != "succeeded":
            return False
        definition = self.get_analysis(run["analysis_id"])
        result = run.get("result", {})
        provenance = result.get("provenance", {})
        if (definition.get("engine") != "structured_aggregate"
                or definition.get("aggregate_plan") != run.get("aggregate_plan")
                or definition.get("bindings") != run.get("bindings")
                or _hash(result) != run.get("output_hash")
                or provenance.get("mode") != "executor_verified_source_map"
                or provenance.get("status") != "verified"):
            return False
        for alias, dataset_id in run["bindings"].items():
            data = self.read_dataset(dataset_id)
            if (run.get("input_hashes", {}).get(alias) != data["data_hash"]
                    or run.get("source_hashes", {}).get(alias) != data["source_hash"]):
                return False
        return True

    def run_analysis(
        self,
        analysis_id: str,
        *,
        session_id: str | None = None,
        context_revision: int | None = None,
        allowed: set[str] | None = None,
    ) -> dict:
        """Execute a saved revision. ``allowed`` is the executing session's snapshot ceiling (None: the manager): the
        run is saved as the session would have saved it from those files alone, against the newest requirements
        version it may read and compared only with earlier runs it may read."""
        from foundry.finance_requirement_semantics import capture_requirements
        requirements_snapshot = capture_requirements(self, allowed=allowed)
        # A session runs a saved analysis as it reads it: one naming a version it may not read names the version
        # analysis_in_view reads it as naming, or none.
        analysis = self.analysis_in_view(self.get_analysis(analysis_id), allowed)
        requirements_snapshot = {
            **requirements_snapshot,
            "requirements_id": analysis.get("requirements_id") or requirements_snapshot.get("requirements_id"),
        }
        if analysis.get("engine") == "finance":
            from foundry.finance_service import FinanceStore

            return FinanceStore(
                self.task, snapshot_ids=allowed, session_id=session_id, context_revision=context_revision
            ).run_revision(analysis)
        if analysis.get("engine") in {"python", "shell"}:
            return self.run_script(
                {
                    "code": analysis["code"],
                    "language": analysis["engine"],
                    "bindings": analysis["bindings"],
                    "question": analysis["question"],
                    "timeout_seconds": analysis["timeout_seconds"],
                    "assumptions": analysis["assumptions"],
                    "result_file": analysis.get("result_file"),
                    "requirements_id": analysis.get("requirements_id"),
                    "result_role": analysis.get("result_role"),
                    "scope": analysis.get("scope", {}),
                    "period_binding": analysis.get("period_binding"),
                },
                session_id=session_id,
                context_revision=context_revision,
                analysis_id=analysis_id,
                allowed=allowed,
            )
        datasets = {
            alias: self.read_dataset(dataset_id)
            for alias, dataset_id in analysis["bindings"].items()
        }
        resolved_period = None
        if analysis.get("period_binding"):
            resolved_period = resolve_period(self, datasets, analysis.get("scope", {}), allowed=allowed)
            analysis = bind_period(analysis, resolved_period)
            if analysis.get("engine") != "structured_aggregate":
                normalize_parameter_metadata(analysis["sql"], analysis["parameters"], analysis.get("parameter_metadata"))
        structured = analysis.get("engine") == "structured_aggregate"
        environment = {
            "python": platform.python_version(),
            "duckdb": duckdb.__version__,
            "platform": sys.platform,
            "runner": "snapshot-aggregate-v1" if structured else "snapshot-sql-v1",
            "settings": CONFIG,
            "configuration_locked": True,
        }
        code_hash = _hash(
            {"aggregate_plan": analysis["aggregate_plan"]} if structured
            else {"sql": analysis["sql"], "parameters": analysis["parameters"]}
        )
        input_hashes = {alias: data["data_hash"] for alias, data in datasets.items()}
        fingerprint = _hash(
            {
                "code_hash": code_hash,
                "inputs": input_hashes,
                "environment": environment,
                "relationships": analysis["relationships"],
                "max_rows": analysis["max_rows"],
                **({"resolved_period": resolved_period} if resolved_period else {}),
            }
        )
        if structured:
            from foundry.analysis_aggregate import execute_aggregate
            import time
            started = time.monotonic()
            try:
                aggregate = execute_aggregate(
                    datasets[analysis["aggregate_plan"]["dataset_alias"]],
                    analysis["aggregate_plan"], max_rows=analysis["max_rows"],
                    timeout_seconds=analysis["timeout_seconds"],
                )
                result = {"result": aggregate, "deterministic": True,
                          "duration_seconds": time.monotonic() - started}
            except FoundryError as exc:
                result = {"error": str(exc), "deterministic": True,
                          "duration_seconds": time.monotonic() - started}
        else:
            result = execute_sql(datasets, analysis)
        succeeded = "error" not in result
        output_hash = _hash(result["result"]) if succeeded else None
        # The earlier runs are read under the lock the run is committed under: none lands between the two, so a reader
        # comparing by commit order (run_in_view) reads what was saved. Those committed already are read and checked
        # first, outside it, so the lock waits only on the runs committed since, not the task's whole run history.
        read: dict[Path, dict] = {}
        self._list("runs", read=read)
        with exclusive(self.task.lock_path, timeout=10):
            previous = [
                r
                for r in self._list("runs", read=read)
                if r.get("fingerprint") == fingerprint
                and r["execution_status"] == "succeeded"
                # The fingerprint keys inputs by their data, not their snapshot: a run on a file outside the session
                # that holds the same rows decides nothing the session saves.
                and (allowed is None or set((r.get("bindings") or {}).values()) <= allowed)
            ]
            # A failed SQL result has no "deterministic" key; it is unverified whatever it says.
            reproducibility = _snapshot_reproducibility(succeeded, result.get("deterministic", True), previous,
                                                        output_hash)
            record = _new(
                self.task,
                analysis_id=analysis_id,
                **({"engine": "structured_aggregate", "aggregate_plan": analysis["aggregate_plan"],
                    "finance_scope": analysis["aggregate_plan"].get("scope", {})}
                   if structured else {}),
                **requirements_snapshot,
                result_role=analysis.get("result_role", "business"),
                scope=analysis.get("scope", {}),
                session_id=session_id,
                context_revision=context_revision,
                sql=analysis["sql"],
                parameters=analysis["parameters"],
                parameter_metadata=analysis.get("parameter_metadata", []),
                **({"period_binding": analysis["period_binding"], "resolved_period": resolved_period,
                    "execution_input_hash": _hash({"datasets": input_hashes, "period": resolved_period})}
                   if resolved_period else {}),
                bindings=analysis["bindings"],
                input_hashes=input_hashes,
                source_hashes={
                    alias: data["source_hash"] for alias, data in datasets.items()
                },
                code_hash=code_hash,
                fingerprint=fingerprint,
                environment=environment,
                execution_status="succeeded" if succeeded else "failed",
                reproducibility_status=reproducibility,
                review_status="unreviewed",
                publication_status="draft",
                output_hash=output_hash,
                missing_inputs=analysis["missing_inputs"],
                assumptions=analysis["assumptions"],
                **result,
            )
            return self._write_locked(self._object_path("runs", record["id"]), record)

    def get_draft(self, draft_id: str) -> dict:
        return self._get("drafts", draft_id)

    def draft_snapshots(self, draft: dict, *, cache: dict[str, frozenset[str]] | None = None) -> set[str]:
        """Every snapshot a saved draft rests on: its runs and claims, its evidence recipes and consistency
        sources, the requirements versions it was pinned to and checked against, and, for a revision that took
        anything of its parent (_rests_on_parent), everything the parent rests on: that text was written or
        generated from the parent's files. A session may read the draft only if it may read all of them; a run or
        record that no longer reads raises. ``cache`` (draft id to its snapshots), when given, is shared by the
        drafts of one listing, so each draft of a revision chain is read once."""
        from foundry.financial_evidence import consistency_dataset_ids, recipe_dataset_ids

        read: dict[Any, frozenset[str]] = {}

        def run(run_id: Any) -> frozenset[str]:
            if not (isinstance(run_id, str) and run_id in read):
                read[run_id] = frozenset((self.get_run(run_id).get("bindings") or {}).values())
            return read[run_id]

        def own(record: dict) -> set[str]:
            snapshots: set[str] = set()
            for run_id in record.get("run_ids") or []:
                snapshots |= run(run_id)
            for claim in record.get("claims") or []:
                if isinstance(claim, dict):
                    if claim.get("dataset_id"):
                        snapshots.add(claim["dataset_id"])
                    if claim.get("run_id"):
                        snapshots |= run(claim["run_id"])
            recipes = record.get("evidence_recipes") or []
            snapshots |= recipe_dataset_ids(recipes) | consistency_dataset_ids(record.get("consistency_sources") or {})
            for recipe in recipes:
                claim = recipe.get("claim") if isinstance(recipe, dict) else None
                if isinstance(claim, dict) and claim.get("run_id"):
                    snapshots |= run(claim["run_id"])
            checked = [check["requirements"].get("requirements_id")
                       for check in (record.get("explanation_check"), record.get("quality"))
                       if isinstance(check, dict) and isinstance(check.get("requirements"), dict)]
            for requirements_id in dict.fromkeys([record.get("requirements_id"), *checked]):
                if requirements_id:
                    bindings = self._get("requirements", requirements_id).get("bindings")
                    bindings = {} if bindings is None else bindings
                    # Bindings that are no map of snapshot ids no session reads (requirements_readable): the version
                    # raises as one that no longer reads.
                    if not isinstance(bindings, dict) or any(not isinstance(value, str) for value in bindings.values()):
                        raise FoundryError(REQUIREMENTS_UNAVAILABLE)
                    snapshots.update(bindings.values())
            return snapshots

        # Up the chain while each revision rests on its parent, then back down, each draft's own snapshots added once.
        chain: list[dict] = []
        seen: set[Any] = {draft.get("id")}
        total: set[str] = set()
        current = draft
        while True:
            if cache is not None and current.get("id") in cache:
                total = set(cache[current["id"]])
                break
            chain.append(current)
            if not _rests_on_parent(current):
                break
            if current["parent_id"] in seen:
                raise FoundryError("draft revision chain is invalid")
            seen.add(current["parent_id"])
            current = self.get_draft(current["parent_id"])
        for record in reversed(chain):
            total |= own(record)
            if cache is not None and isinstance(record.get("id"), str):
                cache[record["id"]] = frozenset(total)
        return total

    def draft_readable(self, draft: dict, allowed: set[str] | None, *,
                       cache: dict[str, frozenset[str]] | None = None) -> bool:
        """Whether a reader whose snapshot ceiling is ``allowed`` (None: the manager) may read a saved draft.
        ``cache`` as for draft_snapshots."""
        if allowed is None:
            return True
        try:
            return self.draft_snapshots(draft, cache=cache) <= set(allowed)
        except FoundryError:
            return False

    def _parent_readable(self, parent_id: str, allowed: set[str] | None) -> bool:
        try:
            return self.draft_readable(self.get_draft(parent_id), allowed)
        except FoundryError:
            return False

    def named_requirements(self, requirements_id: Any, allowed: set[str] | None) -> dict:
        """A requirements version a call names by id (a spec's requirements_id, a saved analysis's pin). A reader whose
        snapshot ceiling is ``allowed`` is refused one it may not read (finance_requirements.requirements_readable)
        exactly as an unknown or malformed id (REQUIREMENTS_UNAVAILABLE). The manager (None) reads it, or is refused,
        as stored; its callers check the bindings as before."""
        if allowed is None:
            return self._get("requirements", requirements_id)
        from foundry.finance_requirements import requirements_readable

        try:
            record = self._get("requirements", requirements_id)
        except FoundryError:
            raise FoundryError(REQUIREMENTS_UNAVAILABLE) from None
        if not requirements_readable(record, allowed):
            raise FoundryError(REQUIREMENTS_UNAVAILABLE)
        return record

    def draft_in_view(self, draft: dict, allowed: set[str] | None, *, readable_ids: set[str] | None = None,
                      versions: dict | None = None, cache: dict | None = None) -> dict:
        """A stored draft as a reader whose snapshot ceiling is ``allowed`` reads it; None is the full view.

        Whoever saved the draft read its evidence over the files and runs they could read. Every reader reads
        the labels, posture and freshness as it would have saved them (evidence_labels.evidence_in_view): a
        newer file outside a session decides nothing the session reads, save that a refresh it may read filed against
        one replaces what that one replaced (finance_controls.refreshed_snapshots), and the manager reads a session's
        draft with the newer files the session could not read. The basis the labels were read from names
        every such file, so it never reaches a session; the integrity hash covers the stored record, so it
        does not go with a view that differs from it (save that the manager reads a revision with its record's hash,
        which also covers what the revision took from its parent, as it always read every draft). A draft saved before drafts kept a basis cannot tell
        what its labels rested on: a session reads it without labels, posture or freshness, the full view as
        saved. A session that may not read a revision's parent reads it as saved with no parent; one reading a
        check saved before completion policy 13 reads no engine run it may not read
        (finance_requirements.requirements_in_view); and every session reads what a check took from its requirements
        version's lineage as its own check would have (finance_requirements.check_lineage_in_view). A session reads
        the row estimates of a cited run as it reads the run (_row_estimates_in_view), and the draft's reproducibility
        by its runs' as it reads them (_draft_reproducibility_in_view). ``readable_ids``, when given, are the drafts
        this reader may read (draft_readable), so a listing need not read every parent again; ``versions``, the
        requirements versions a listing has read (finance_requirements._version_record); ``cache`` as for run_in_view.
        """
        from foundry.evidence_labels import EVIDENCE_FIELDS, evidence_in_view
        from foundry.finance_professional import draft_professional_view

        cache = {} if cache is None else cache
        draft = draft_professional_view(self, draft, allowed, cache=cache)

        views: dict[str, dict] = {}
        loaded: dict[str, dict] = {}
        # A revision is compared with its parent's runs: a reader that may not read the parent reads the draft as
        # it would have saved it, with no parent, and the labels read again from that check.
        orphaned = (allowed is not None and bool(draft.get("parent_id"))
                    and not (draft["parent_id"] in readable_ids if readable_ids is not None
                             else self._parent_readable(draft["parent_id"], allowed)))
        # Its reproducibility too, as saved with no parent: a body changed from that parent decides nothing.
        reproducibility = self._draft_reproducibility_in_view(draft, allowed, cache, orphaned=orphaned)
        if orphaned:
            from foundry.finance_revision import compare_draft_runs

            # Review stays as stored: a change from the parent that nobody reviewed never reads as reviewed.
            draft = {**draft, "parent_id": None, "revision_comparison": compare_draft_runs(self, {}, {}),
                     **{key: _without_revision(draft[key]) for key in ("explanation_check", "quality") if key in draft}}
        from foundry.finance_requirements import OWN_FILES_POLICY, requirements_in_view

        if allowed is not None and (draft.get("check_policy_version") or 0) < OWN_FILES_POLICY:
            # Saved when a requirement still read runs on other files too: a session reads none it may not read.
            def readable_run(run_id: Any) -> bool:
                try:
                    return isinstance(run_id, str) and set(
                        (self.verified_run(run_id, allowed, cache).get("bindings") or {}).values()) <= set(allowed)
                except FoundryError:
                    return False

            draft = {**draft, **{key: requirements_in_view(draft[key], readable_run)
                                 for key in ("explanation_check", "quality") if isinstance(draft.get(key), dict)}}
        if allowed is not None:
            from foundry.finance_requirements import check_lineage_in_view

            # The versions this read has listed (_listed_versions), so each lineage walk reads none again.
            versions = self._listed_versions(cache) if versions is None else versions
            draft = {**draft, **{key: check_lineage_in_view(self, draft[key], allowed, versions=versions)
                                 for key in ("explanation_check", "quality") if isinstance(draft.get(key), dict)}}

        def runs() -> dict[str, dict]:
            if not loaded:
                loaded.update({run_id: self.verified_run(run_id, allowed, cache)
                               for run_id in dict.fromkeys(draft.get("run_ids") or [])})
            return loaded

        replaced: list[dict[str, set[str]]] = []

        def refreshed() -> dict[str, set[str]]:
            # The refreshes filed before the draft, each against the prior this reader reads it with, from batches
            # listed once for every draft of the read.
            if not replaced:
                from foundry.finance_controls import refreshed_snapshots

                ceiling = set(allowed)
                replaced.append(refreshed_snapshots(self.task, ceiling, before=draft.get("seq"),
                                                    batches=self._batches_in_view(ceiling, cache)))
            return replaced[0]

        estimates: Any = False  # not read yet
        for key in ("explanation_check", "quality"):
            check = draft.get(key)
            if not isinstance(check, dict) or not any(field in check for field in EVIDENCE_FIELDS):
                continue
            if not isinstance(draft.get("evidence_basis"), dict):
                if allowed is not None:
                    views[key] = {name: value for name, value in check.items() if name not in EVIDENCE_FIELDS}
                continue
            try:
                if estimates is False:
                    estimates = self._row_estimates_in_view(draft, allowed, runs, cache)
                fields = evidence_in_view(draft, check, None if allowed is None else set(allowed), runs=runs,
                                          force=orphaned, row_estimates=estimates,
                                          refreshed=None if allowed is None else refreshed)
            except FoundryError:
                # A run it cites no longer reads, so they cannot be read again: a session reads the draft without
                # them, the full view as saved.
                if allowed is not None:
                    views[key] = {name: value for name, value in check.items() if name not in EVIDENCE_FIELDS}
                continue
            if fields is not None:
                views[key] = {**check, **fields}
        # What a revision took from its parent decides who may read it, and is no part of what they read. The manager
        # reads a revision with the integrity hash of its stored record, as every other draft.
        shown = {key: value for key, value in draft.items() if key != "inherited_fields"} if "inherited_fields" in draft else draft
        if allowed is None and not views:
            return shown
        hidden = {"_integrity_hash"} if allowed is None else {"evidence_basis", "_integrity_hash"}
        if allowed is not None and "reproducibility_status" in draft:
            views["reproducibility_status"] = reproducibility
        return {**{key: value for key, value in shown.items() if key not in hidden}, **views}

    def _draft_reproducibility_in_view(self, draft: dict, allowed: set[str] | None, cache: dict, *,
                                       orphaned: bool = False) -> Any:
        """A draft's reproducibility as a session whose ceiling is ``allowed`` would have saved it (save_draft): by its
        own conditions, and its runs' reproducibility as the session reads them (_reproducibility_in_view). The manager
        reads it as stored.

        save_draft stores reproducible only when the draft's own conditions hold (cited runs, calculation claims only,
        the same body and claims as its parent) and every cited run read reproducible to its saver. A stored
        reproducible draft therefore keeps its conditions. A stored unverified one may have failed only on a run as its
        saver read it, so its conditions are read again from the record and its parent: the body before its estimate
        notices (_body_before_marks), as save_draft compared it. A revision whose parent the session may not read
        (``orphaned``, draft_in_view) it would have saved with no parent: its conditions are read again from the record
        alone, however its runs read, and the parent is never read.
        """
        stored = draft.get("reproducibility_status")
        run_ids = list(dict.fromkeys(draft.get("run_ids") or []))
        if allowed is None or not run_ids:
            return stored
        try:
            runs = [self._stored_run(run_id, cache) for run_id in run_ids]
        except FoundryError:
            return stored
        # Every reader reads these runs as stored, so the draft too, unless a parent it may not read decided it.
        if not orphaned and not any(_snapshot_run(run) for run in runs):
            return stored
        if not all(self._reproducibility_in_view(run, allowed, cache) == "reproducible" for run in runs):
            return "unverified"
        if stored == "reproducible":
            return stored
        claims = draft.get("claims") or []
        if not claims or not all(isinstance(claim, dict) and claim.get("basis") == "calculation" for claim in claims):
            return stored
        if draft.get("parent_id") and not orphaned:
            # A revision is reproducible only with its parent's body and claims; a parent that no longer reads cannot
            # show that it kept them.
            try:
                parent = self.get_draft(draft["parent_id"])
            except FoundryError:
                return stored
            if _body_before_marks(draft) != parent.get("body") or claims != parent.get("claims"):
                return stored
        return "reproducible"

    def _judged_otherwise(self, draft: dict, allowed: set[str], cache: dict, *, rows: bool = False) -> bool:
        """Whether a reader whose snapshot ceiling is ``allowed`` may judge ``draft``'s cited runs otherwise than its row
        estimates (``rows``) or its receipts were judged, so that it judges them again.

        A session saves a draft and files its receipts on the runs as it reads them (run_in_view), and whoever saved
        them may read every file the draft rests on (draft_snapshots): at least those its cited runs and the versions
        it was pinned to and checked against rest on. A run reads otherwise where a version a reader may not read
        decides it, so a reader that may read only those files reads a cited run otherwise wherever any reader of the
        draft does, and where a session saved it with the fingerprint it read the version it started under with, to
        any reader that reads that version with another (_captured_otherwise). Where it and this reader read every
        cited run as stored, the draft was judged as this reader's own save would have judged it, whoever saved it;
        otherwise this reader judges it again on its own reading. A run reads otherwise pinned to another version
        (_pin_in_view) or, but for ``rows``, with another fingerprint, or pinned to or checked against a version read
        with another (_semantics_in_view). Telling reads no run's outputs, and no pin where every version rests on
        those files and every cited run captured the stored fingerprint of the version it started under, if any.
        ``cache`` as for run_in_view."""
        from foundry.finance_requirements import requirements_readable

        judged = cache.setdefault("judged", {})
        key = (draft.get("id"), rows)
        if key not in judged:
            runs = [self._stored_run(run_id, cache) for run_id in dict.fromkeys(draft.get("run_ids") or [])]
            checked = [check["requirements"].get("requirements_id")
                       for check in (draft.get("explanation_check"), draft.get("quality"))
                       if isinstance(check, dict) and isinstance(check.get("requirements"), dict)]
            versions = [self._requirements_version(version_id, cache)
                        for version_id in dict.fromkeys([draft.get("requirements_id"), *checked])]
            files = {file for record in [*runs, *versions] if record
                     for file in (record.get("bindings") or {}).values()}

            def otherwise(ceiling: set[str], read: dict) -> bool:
                for run in runs:
                    pin = self._pin_in_view(run, ceiling, read)
                    if pin is not None and (not rows or "requirements_id" in pin[0]):
                        return True
                if rows:
                    return False
                deciding = [*versions, *(self._requirements_version(run.get("requirements_id"), read) for run in runs)]
                return any(version is not None and requirements_readable(version, ceiling)
                           and self._semantics_in_view(version, ceiling, read) is not version for version in deciding)

            def captured_as_stored(run: dict) -> bool:
                # Every reader of the draft reads every version as stored, and so a run as stored where it captured
                # none or the stored fingerprint of the newest version before it that its saver may read
                # (_versions_at_start).
                captured = run.get("requirements_semantic_fingerprint")
                if not isinstance(captured, str):
                    return True
                author = self._author_ceiling(run, cache)
                before = [row for row in self._requirements_versions(cache) if row.get("seq", 0) < run.get("seq", 0)
                          and (author is None or requirements_readable(row, author))]
                return not before or before[-1].get("semantic_fingerprint") == captured

            if all(requirements_readable(row, files) for row in self._requirements_versions(cache)) and (
                    rows or all(captured_as_stored(run) for run in runs)):
                judged[key] = False
            else:
                # The narrowest reader reads fingerprints of its own: only versions, sessions and analyses are shared.
                narrowest = cache.setdefault("narrowest", {}).setdefault(frozenset(files), {
                    "requirements": self._requirements_versions(cache), "versions": self._requirements_by_id(cache),
                    "sessions": cache.setdefault("sessions", {}), "analyses": cache.setdefault("analyses", {})})
                judged[key] = otherwise(set(allowed), cache) or otherwise(files, narrowest)
        return judged[key]

    def _row_estimates_in_view(self, draft: dict, allowed: set[str] | None, runs: Callable[[], dict[str, dict]],
                               cache: dict) -> list | None:
        """The basis's row estimates as a reader whose snapshot ceiling is ``allowed`` would have saved them, or None
        when they stand as saved (always, for the manager).

        draft_evidence_basis read each cited run's row declarations, as its saver read the run, against the version it
        was pinned to, with the estimates of the version the draft was checked against added. A reader may read that
        version of every draft it reads (draft_snapshots holds it), so only a cited run it or its saver may read pinned
        to another version reads otherwise (_judged_otherwise): the rows are declared again as it reads the runs, with
        that version's estimates (for a check that files no requirements, those of the newest version it may read that
        is older than the draft).
        Whether a run reads so is told from its stored record, so a read where no version rests on files other than
        the draft's, or none pins a cited run otherwise, loads no run (``runs`` loads them, as evidence_in_view does).
        """
        if allowed is None:
            return None
        from foundry.evidence_labels import declared_row_estimates
        from foundry.finance_requirements import requirements_readable

        if not self._judged_otherwise(draft, allowed, cache, rows=True):
            return None
        versions = self._requirements_versions(cache)
        # The version the draft was checked against, as draft_evidence_basis read it, not the newest one before the
        # draft's seq: a version may be saved while the draft is saved, after its check.
        check = draft.get("explanation_check")
        requirements = check.get("requirements") if isinstance(check, dict) else None
        if isinstance(requirements, dict) and "requirements_id" in requirements:
            checked = self._requirements_version(requirements["requirements_id"], cache)
            limit = checked.get("seq", 0) + 1 if checked else 0
        else:
            limit = draft.get("seq", 0)
        latest = next((row for row in reversed(versions)
                       if row.get("seq", 0) < limit and requirements_readable(row, allowed)), None)
        estimates = declared_row_estimates(self, runs(), latest, set(allowed), cache=cache)
        listed = [[run_id, row, sources] for (run_id, row), sources in sorted(estimates.items())]
        return None if listed == (draft["evidence_basis"].get("row_estimates") or []) else listed

    def completion_reader(self, allowed: set[str], *, readable_ids: set[str] | None = None) -> Callable[[dict], bool]:
        """Whether a session whose snapshot ceiling is ``allowed`` may read a completion receipt at all, as a test that
        reads each draft and requirements version once however many receipts it is asked about.

        A receipt is the verdict on one draft, so it reads the way its draft does: not for a draft the session may
        not read, nor for a revision whose parent it may not read (the draft reads to it with no parent; the receipt
        was judged with one), nor for a verdict judged against a requirements version it may not read. A record that
        no longer reads counts as one it may not read. ``readable_ids`` as for draft_in_view.
        """
        from foundry.finance_requirements import requirements_readable

        allowed = set(allowed)
        loaded: dict[str, dict | None] = {}
        drafts: dict[str, bool] = {}
        versions: dict[str, bool] = {}
        receipts: dict[Any, bool] = {}

        def load(draft_id: Any) -> dict | None:
            if not isinstance(draft_id, str):
                return None
            if draft_id not in loaded:
                try:
                    loaded[draft_id] = self.get_draft(draft_id)
                except FoundryError:
                    loaded[draft_id] = None
            return loaded[draft_id]

        def draft(draft_id: str) -> bool:
            if readable_ids is not None:
                return draft_id in readable_ids
            if draft_id not in drafts:
                stored = load(draft_id)
                drafts[draft_id] = stored is not None and self.draft_readable(stored, allowed)
            return drafts[draft_id]

        def version(requirements_id: str) -> bool:
            if requirements_id not in versions:
                try:
                    versions[requirements_id] = requirements_readable(
                        self._get("requirements", requirements_id), allowed)
                except FoundryError:
                    versions[requirements_id] = False
            return versions[requirements_id]

        def readable(receipt: dict) -> bool:
            key = receipt.get("id")
            if key not in receipts:
                stored = load(receipt.get("draft_id"))
                receipts[key] = (stored is not None and draft(stored["id"])
                                 and (not stored.get("parent_id") or draft(stored["parent_id"]))
                                 and (not receipt.get("requirements_id") or version(receipt["requirements_id"])))
            return receipts[key]

        return readable

    def completion_in_view(self, receipt: dict, allowed: set[str] | None, *,
                           readable_ids: set[str] | None = None, cache: dict | None = None) -> dict | None:
        """A completion receipt as a reader whose snapshot ceiling is ``allowed`` reads it; None when it may not.

        The manager reads it as stored. A session reads none it may not read (completion_reader). A receipt counts
        its repair budget over every earlier receipt of its chain (record_completion); a session reads it counted over
        those it may read (_repair_budget), so a draft it may not read neither spends nor shows through its budget,
        and its completion operation settles on that reading. A receipt's messages and fixups name runs and carry
        their values, as do the engine runs of a requirement checked before completion policy 13: when one is outside
        the session, they and the digests over them are withheld, and the session reads the verdict, judged without
        the engine runs it may not read. ``readable_ids`` as for draft_in_view, ``cache`` as for run_in_view.
        """
        return self.completions_in_view([receipt], allowed, readable_ids=readable_ids, cache=cache)[0]

    def completions_in_view(self, receipts: list[dict], allowed: set[str] | None, *, readable_ids: set[str] | None = None,
                            versions: dict | None = None, cache: dict | None = None) -> list[dict | None]:
        """completion_in_view of each receipt, reading every draft and earlier receipt once. ``versions`` as for
        draft_in_view, ``cache`` as for run_in_view.

        A receipt judged the draft's runs as whoever filed it read them: one this reader may judge otherwise
        (_judged_otherwise), or that judged a check it reads without runs it may not read, reads with the verdict its
        own save would have filed (_receipt_verdict_in_view). A receipt filed under a fingerprint other than the one
        this reader reads its version with, whoever filed it, reads as filed under that one, and so does its chain
        (_chain_in_view). A stall is judged on the receipt before, as the reader reads it, and the budget counts the
        earlier receipts of the chain the reader may read."""
        from foundry.finance_professional import receipt_professional_view
        if allowed is None:
            return [receipt_professional_view(self, receipt, allowed, cache=cache) for receipt in receipts]
        from foundry.finance_completion import receipt_lineage_in_view
        from foundry.finance_requirements import OWN_FILES_POLICY

        allowed = set(allowed)
        cache = {} if cache is None else cache
        versions = self._listed_versions(cache) if versions is None else versions
        readable = self.completion_reader(allowed, readable_ids=readable_ids)
        chain: list[dict] | None = None
        # Each receipt as this reader reads it, whole (None: it may not read it), and whether every run it names is
        # one the reader may read.
        read: dict[Any, tuple[dict, bool] | None] = {}
        # The chain each receipt reads filed under where it reads otherwise than stored (_chain_in_view).
        chains: dict[Any, dict] = {}

        def order(row: dict) -> tuple:
            return (row.get("seq", 0), row.get("created_at", ""), row.get("id", ""))

        def rekeyed(row: dict) -> dict:
            if row.get("id") not in chains:
                chains[row.get("id")] = self._chain_in_view(row, allowed, cache)
            return chains[row.get("id")]

        def last_read(row: dict) -> dict:
            # The receipt before, with the issue_fingerprint this reader reads it with: a stall is judged on it alone,
            # and its verdict on its own draft, whatever came before it, so none of those is read again.
            if read.get(row.get("id")) is not None:
                return read[row.get("id")][0]
            try:
                verdict = self._receipt_verdict_in_view(row, self.get_draft(row["draft_id"]), allowed, [],
                                                        versions=versions, cache=cache)
            except FoundryError:
                return row
            return row if verdict is None else {**row, "issue_fingerprint": verdict["issue_fingerprint"]}

        def view(receipt: dict) -> tuple[dict, bool] | None:
            nonlocal chain
            if receipt.get("id") in read:
                return read[receipt.get("id")]
            read[receipt.get("id")] = None
            if not readable(receipt):
                return None
            try:
                draft = self.get_draft(receipt["draft_id"])
                named = _receipt_runs(receipt)
                if (draft.get("check_policy_version") or 0) < OWN_FILES_POLICY:
                    named |= _engine_runs(draft)
                whole = all(set((self.verified_run(run_id, allowed, cache).get("bindings") or {}).values()) <= allowed
                            for run_id in sorted(named))
            except FoundryError:
                return None
            filed = receipt
            # The verdict reads the requirements version's lineage as the draft's check reads it to the session, and
            # the receipt the chain it reads it filed under.
            receipt = receipt_lineage_in_view(self, receipt, draft, allowed, versions=versions, chain=rekeyed(receipt))
            if chain is None:
                chain = self._list("completions")
            earlier = [row for row in chain if receipt.get("chain_key") is not None
                       and rekeyed(row).get("chain_key", row.get("chain_key")) == receipt["chain_key"]
                       and order(row) < order(receipt)]
            kept = [row for row in earlier if readable(row)]
            # A count reads how many receipts came before and the last of them, as the reader reads it.
            seen = [*kept[:-1], last_read(kept[-1])] if kept else kept
            # Stored over every earlier receipt of the chain it was filed under: one the session may not read, one it
            # reads in another chain, or the one before reading otherwise than stored, is the only reason to count
            # again.
            counted = [row for row in chain if filed.get("chain_key") is not None
                       and row.get("chain_key") == filed["chain_key"] and order(row) < order(receipt)]
            moved = bool(kept) and seen[-1].get("issue_fingerprint") != kept[-1].get("issue_fingerprint")
            budget = {} if kept == counted and not moved else _repair_budget(receipt, seen)
            verdict = self._receipt_verdict_in_view(filed, draft, allowed, seen, versions=versions, cache=cache)
            if verdict is not None:
                # Keyed as a stored receipt reads back (write_json_atomic sorts keys), so the view reads as one.
                receipt, budget = dict(sorted({**{key: value for key, value in receipt.items()
                                                   if key not in RECEIPT_RUN_VERDICT}, **verdict}.items())), {}
            read[receipt.get("id")] = ({**{key: value for key, value in receipt.items() if key != "_integrity_hash"},
                                        **budget}, whole)
            return read[receipt.get("id")]

        views: list[dict | None] = []
        for receipt in receipts:
            shown = view(receipt)
            if shown is None or shown[1]:
                views.append(shown and shown[0])
                continue
            views.append({**{key: shown[0][key] for key in COMPLETION_VERDICT if key in shown[0]},
                          "actions": [{key: action[key] for key in ("code", "requirement_item_id") if key in action}
                                      for action in shown[0].get("actions") or [] if isinstance(action, dict)]})
        return [receipt_professional_view(self, receipt, allowed, cache=cache) for receipt in views]

    def _receipt_verdict_in_view(self, receipt: dict, draft: dict, allowed: set[str], kept: list[dict], *,
                                 versions: dict, cache: dict) -> dict | None:
        """The RECEIPT_RUN_VERDICT a session reads of a receipt on runs it may judge otherwise than whoever filed it
        (_judged_otherwise), or against a fingerprint of its version other than the one it reads (_captured_otherwise),
        or on a check it reads without runs it may not read, or None when it reads the stored one. Whether a run is
        bound to the current requirements was judged on the runs as the filer read them: the session reads the verdict
        its own save of the draft would have filed (finance_completion.assess_completion with its ceiling), over the
        check as it reads it (draft_in_view) and counted over ``kept``, the earlier receipts of the chain as it reads
        them. A check read without such runs reads its issues, and so their digest, without them. Whether a run or
        version reads so is told from its stored record, so a receipt whose runs, version and check all read as stored
        is assessed no further."""
        from foundry.finance_completion import assess_completion, receipt_fields
        from foundry.finance_requirements import OWN_FILES_POLICY, check_lineage_in_view, requirements_in_view

        quality = draft.get("explanation_check")
        if receipt.get("draft_content_hash") != draft.get("content_hash") or not isinstance(quality, dict):
            return None
        shown = quality
        if (draft.get("check_policy_version") or 0) < OWN_FILES_POLICY:
            # As draft_in_view reads a check saved when a requirement still read runs on other files too.
            def readable_run(run_id: Any) -> bool:
                try:
                    return isinstance(run_id, str) and set(
                        (self.verified_run(run_id, allowed, cache).get("bindings") or {}).values()) <= allowed
                except FoundryError:
                    return False

            shown = requirements_in_view(quality, readable_run)
        try:
            if (shown is quality and not self._judged_otherwise(draft, allowed, cache)
                    and not self._captured_otherwise(self._requirements_version(receipt.get("requirements_id"), cache),
                                                     receipt.get("requirements_semantic_fingerprint"), allowed, cache)):
                return None
        except FoundryError:
            return None
        # A read assesses a version's support once for every receipt on it that it judges again (with_input_gaps).
        assessed = assess_completion(self, draft, check_lineage_in_view(self, shown, allowed, versions=versions),
                                     allowed=allowed, cache=cache, supports=cache.setdefault("supports", {}))
        stale = next((action.get("run_ids") or [] for action in receipt.get("actions") or []
                      if isinstance(action, dict) and action.get("code") == "requirements_run_outdated"), [])
        if (sorted(assessed["reused_runs"]) == sorted(receipt.get("reused_run_ids") or [])
                and sorted(assessed["stale_runs"]) == sorted(stale)
                and (shown is quality or assessed["issue_fingerprint"] == receipt.get("issue_fingerprint"))):
            return None
        fields = receipt_fields(self, assessed, kept)
        return {key: fields[key] for key in RECEIPT_RUN_VERDICT if key in fields}

    def _chain_in_view(self, receipt: dict, allowed: set[str], cache: dict) -> dict:
        """The requirements fingerprint and chain key a reader whose snapshot ceiling is ``allowed`` reads a completion
        receipt filed under, where they read otherwise than stored; nothing where they stand. One filed under a
        fingerprint other than the one the reader reads the version with (_captured_otherwise), whoever filed it, reads
        as filed under that one, as the reader's own save would have filed it (finance_completion.assess_completion)."""
        from foundry.finance_completion import completion_chain_key

        version = self._requirements_version(receipt.get("requirements_id"), cache)
        captured = receipt.get("requirements_semantic_fingerprint")
        if (not self._captured_otherwise(version, captured, allowed, cache) or receipt.get("chain_key")
                != completion_chain_key(captured, version["id"], receipt.get("run_hashes") or {})):
            return {}
        fingerprint = self.fingerprint_in_view(version, allowed, cache=cache)
        return {"requirements_semantic_fingerprint": fingerprint,
                "chain_key": completion_chain_key(fingerprint, version["id"], receipt.get("run_hashes") or {})}

    def _registered_report_source_ids(self) -> set[str]:
        return set(self._registered_report_sources())

    def _registered_report_sources(self) -> dict[str, dict]:
        root = self._inside(self.root / "report_sources")
        sources = {}
        for path in root.glob("*.json"):
            if ID.fullmatch(path.stem):
                try:
                    record = self._read(path)
                except FoundryError:
                    continue
                sources[record["id"]] = self._report_source_public(record)
        return sources

    def save_draft(self, spec: dict, *, _recipe_report: dict | None = None, _allowed: set[str] | None = None,
                   _inherited: set[str] | frozenset[str] = frozenset(), _cache: dict | None = None,
                   _required_check_authorship: dict | None = None) -> dict:
        """Save an immutable draft version.

        `_recipe_report` is for internal callers that already evaluated this draft's evidence_recipes
        with evaluate_evidence_recipes in the same call; a report inside `spec` is never read.
        `_allowed` is the calling session's snapshot ceiling, so evidence freshness and the requirement check read only
        what it may.
        `_inherited` is the INHERITED_FIELDS an internal caller took from the parent itself before composing `spec`.
        `_cache` is the read cache (run_in_view) of the caller's request, which reads the saved draft with it: a save
        writes no run, requirements version or refresh batch, so what it holds stays true.
        """
        if not isinstance(spec, dict):
            raise FoundryError("draft spec must be an object")
        if {"required_check_authorship", "_required_check_authorship"}.intersection(spec):
            raise FoundryError("required check authorship is computed by the service, not draft spec")
        if _allowed is not None:
            from foundry.finance_requirements import latest_requirements, requirements_readable

            # A draft is checked against the current requirements: a session may not check one against a version it
            # may not read (requirements_readable: on files it may not read, or with malformed bindings), as the
            # finance saves refuse it.
            current = latest_requirements(self)
            if current and not requirements_readable(current, _allowed):
                raise FoundryError("snapshots are outside this session context; use analysis-refresh to include task datasets")
        parent = self.get_draft(spec["parent_id"]) if spec.get("parent_id") else {}
        inherited = _inherited_fields(spec, parent) | set(_inherited) if parent else None
        title = _safe_text(
            spec.get("title", parent.get("title", DEFAULT_DRAFT_TITLE)),
            "draft title",
            required=True,
            maximum=1000,
        )
        document = None
        visual_run_ids, visual_claims = [], []
        if spec.get("document") is not None:
            from foundry.report_documents import report_document_markdown
            from foundry.report_visuals import materialize_document

            report_sources = self._registered_report_sources()
            document, visual_run_ids, visual_claims = materialize_document(self, spec["document"])
            if document.get("sources"):
                document["sources"] = [report_sources[source["id"]] for source in document["sources"]]
            projected = report_document_markdown(document)
            body = _safe_text(projected, "draft body")
        elif "body" not in spec and parent.get("document") is not None:
            from foundry.report_documents import report_document_markdown
            from foundry.report_visuals import materialize_document

            document, visual_run_ids, visual_claims = materialize_document(self, parent["document"])
            body = _safe_text(report_document_markdown(document), "draft body")
        else:
            body = _safe_text(spec.get("body", parent.get("body", "")), "draft body")
        body_template = spec.get("body_template")
        if body_template is not None:
            body_template = _safe_text(body_template, "draft body_template")
        run_ids = spec.get("run_ids", parent.get("run_ids", []))
        if not isinstance(run_ids, list):
            raise FoundryError("run_ids must be an array")
        run_ids = list(dict.fromkeys([*run_ids, *visual_run_ids]))
        runs = {run_id: self.get_run(run_id) for run_id in dict.fromkeys(run_ids)}
        fact_ids = spec.get("fact_ids", parent.get("fact_ids", []))
        if not isinstance(fact_ids, list):
            raise FoundryError("fact_ids must be an array")
        claims = json_value(spec.get("claims", parent.get("claims", [])))
        if not isinstance(claims, list):
            raise FoundryError("claims must be an array")
        for visual_claim in visual_claims:
            if not any(isinstance(claim, dict) and all(claim.get(key) == visual_claim[key]
                       for key in ("basis", "run_id", "row", "column")) for claim in claims):
                claims.append(visual_claim)
        for claim in claims:
            if not isinstance(claim, dict) or claim.get("basis") not in BASIS:
                raise FoundryError("claim requires an explicit basis")
            if claim["basis"] == "calculation":
                run = runs.get(claim.get("run_id"))
                if not run or run["execution_status"] != "succeeded":
                    raise FoundryError(
                        "calculation claim requires a referenced successful run"
                    )
                try:
                    row = claim["row"]
                    if isinstance(row, bool) or not isinstance(row, int) or row < 0:
                        raise KeyError("row")
                    actual = run["result"]["rows"][row][claim["column"]]
                    expected = claim["value"]
                    matches = actual == expected
                    if not matches and actual is not None and expected is not None:
                        from decimal import Decimal, InvalidOperation

                        try:
                            matches = Decimal(str(actual)) == Decimal(str(expected))
                        except InvalidOperation:
                            pass
                    if not matches:
                        error = FoundryError(
                            "claim value differs from execution result "
                            f"(run_id={run['id']}, row={row}, column={claim['column']}, "
                            f"expected={expected}, actual={actual})"
                        )
                        error.details = {
                            "code": "claim_value_mismatch",
                            "run_id": run["id"],
                            "row": row,
                            "column": claim["column"],
                            "expected": str(expected),
                            "actual": str(actual),
                        }
                        raise error
                except (KeyError, IndexError, TypeError) as exc:
                    raise FoundryError("claim has an invalid output location") from exc
                claim["verification_status"] = "matches_run"
                claim["dependencies"] = {
                    "snapshot_ids": list(run["bindings"].values()),
                    "assumptions": run["assumptions"],
                }
            else:
                claim["verification_status"] = "unverified"
                if claim.get("dataset_id"):
                    data = self.read_dataset(claim["dataset_id"])
                    if claim.get("paragraph"):
                        if not any(
                            p["paragraph"] == claim["paragraph"]
                            for p in data.get("paragraphs", [])
                        ):
                            raise FoundryError(
                                "claim background paragraph does not exist"
                            )
                    claim["source_hash"] = data["source_hash"]
        publication = spec.get("publication_status", "draft")
        review = spec.get("review_status", "unreviewed")
        limitations = spec.get("limitations", parent.get("limitations", []))
        if (
            publication not in PUBLICATION
            or review not in REVIEW
            or not isinstance(limitations, list)
        ):
            raise FoundryError("invalid draft status or limitations")
        if (
            publication == "archived"
            and any(r["execution_status"] != "succeeded" for r in runs.values())
            and not limitations
        ):
            raise FoundryError(
                "archiving failed execution references requires explicit limitations"
            )
        changed = bool(parent) and (
            body != parent.get("body") or claims != parent.get("claims")
        )
        if changed:
            review = "unreviewed"
        # A session reads each run's reproducibility by the runs it may read (run_in_view); the manager as stored.
        cache: dict = {} if _cache is None else _cache
        reproducibility = (
            "reproducible"
            if claims
            and runs
            and all(
                self.run_in_view(r, _allowed, cache=cache)["reproducibility_status"] == "reproducible"
                for r in runs.values()
            )
            and not changed
            and all(c["basis"] == "calculation" for c in claims)
            else "unverified"
        )
        from foundry.finance_explanations import validate_explanation

        materiality = spec.get("materiality", parent.get("materiality", {}))
        source_read_failures = spec.get("source_read_failures", parent.get("source_read_failures", []))
        evidence_appendix = spec.get("evidence_appendix", parent.get("evidence_appendix", ""))
        evidence_appendix_meta = spec.get(
            "evidence_appendix_meta", parent.get("evidence_appendix_meta", {})
        )
        from foundry.finance_requirements import augment_quality, latest_requirements

        current_requirements = latest_requirements(self)
        requirements_id = spec.get("requirements_id", parent.get("requirements_id")) or (current_requirements or {}).get("id")
        if "fact_ids" not in spec and fact_ids and set(runs) != set(parent.get("run_ids", [])):
            # A revision that swaps its runs must not carry the parent's facts from runs it no longer cites.
            fact_ids = _fact_ids_in_runs(self, fact_ids, runs, materiality)
        from foundry.finance_completion import COMPLETION_POLICY_VERSION
        from foundry.finance_revision import compare_draft_runs
        from foundry.financial_consistency import consistency_quality_issues, evaluate_consistency_rules
        from foundry.financial_evidence import evaluate_evidence_recipes, prepare_consistency_inputs

        # Evaluate this draft's recipes once, before requirement coverage reads them: coverage may
        # only use recipe results the server computed in this call.
        evidence_recipes = spec.get("evidence_recipes", parent.get("evidence_recipes", []))
        evidence_report = (
            _recipe_report if _recipe_report is not None else evaluate_evidence_recipes(self, evidence_recipes)
        )
        required_check_authorship = _required_check_authorship
        if required_check_authorship is None and parent.get("required_check_authorship"):
            from foundry.finance_explanations import FACT_TOKEN, prepare_explanation

            explicit_claims = spec.get("claims", [])
            template_fact_ids = FACT_TOKEN.findall(body_template or "")
            required_check_authorship = _draft_required_check_authorship(
                prepare_explanation(self, list(runs.values()), materiality=materiality, _allowed=_allowed),
                list(runs.values()), claims=explicit_claims,
                fact_ids=fact_ids if "fact_ids" in spec else [],
                template_fact_ids=template_fact_ids,
                inherited=(None if (any(isinstance(claim, dict) and claim.get("basis") == "calculation"
                                        for claim in explicit_claims)
                                    or spec.get("fact_ids") or template_fact_ids)
                           else parent["required_check_authorship"]),
            )
        explanation_check = validate_explanation(
            self,
            {
                "requirements_id": requirements_id,
                "body": body,
                "body_template": body_template,
                "run_ids": list(runs),
                "fact_ids": fact_ids,
                "claims": claims,
                "materiality": materiality,
                "evidence_appendix": evidence_appendix,
                "evidence_appendix_meta": evidence_appendix_meta,
            },
            # Recipe issues are added below from the same report; validate only needs it to rebuild
            # the generated appendix exactly as the requirement check sees it.
            _requirement_recipe_report=evidence_report,
            _allowed=_allowed,
            _required_check_authorship=required_check_authorship,
        )
        explanation_check = augment_quality(self, {
            "body": body, "run_ids": list(runs), "claims": claims,
            "requirements_id": requirements_id,
            "source_read_failures": source_read_failures,
            "evidence_recipes": evidence_recipes,
        }, explanation_check, _recipe_report=evidence_report, _allowed=_allowed)
        consistency_rules = spec.get("consistency_rules", parent.get("consistency_rules", []))
        consistency_sources = spec.get("consistency_sources", parent.get("consistency_sources", {}))

        revision_comparison = compare_draft_runs(self, parent, {"run_ids": list(runs)})
        resolved_rules, consistency_source_auth = prepare_consistency_inputs(self, consistency_rules, consistency_sources)
        consistency_report = evaluate_consistency_rules(
            resolved_rules,
            recipe_report=evidence_report,
            source_auth=consistency_source_auth,
        )
        quality_issues = list(explanation_check.get("issues", []))
        quality_issues.extend(_revision_quality_issues(revision_comparison))
        quality_issues.extend(evidence_report.get("issues", []))
        quality_issues.extend(consistency_quality_issues(consistency_report))
        execution_status = (
            "succeeded"
            if runs and all(r["execution_status"] == "succeeded" for r in runs.values())
            else "failed"
            if runs
            else "not_run"
        )
        status_dimensions = _quality_status_dimensions(
            execution_status=execution_status,
            review_status=review,
            quality_issues=quality_issues,
            evidence_report=evidence_report,
            consistency_report=consistency_report,
            numeric_configured=bool(runs or claims),
        )
        explanation_check = {
            **explanation_check,
            "issues": quality_issues,
            "status": "needs_attention" if quality_issues else explanation_check.get("status", "pass"),
            "status_dimensions": status_dimensions,
        }
        from foundry.evidence_labels import draft_evidence_basis

        # Beside the status, never in it: labels and posture classify the evidence and add no issue.
        evidence, evidence_basis = draft_evidence_basis(
            self, explanation_check, claims=claims, fact_ids=fact_ids, runs=runs,
            evidence_report=evidence_report, consistency_report=consistency_report,
            source_read_failures=source_read_failures, allowed=_allowed, cache=cache,
        )
        explanation_check = {**explanation_check, **evidence}
        if document is not None:
            from foundry.report_estimates import estimated_cells
            from foundry.report_visuals import mark_visual_estimates
            from foundry.report_documents import report_document_markdown

            if mark_visual_estimates(document, estimated_cells({"explanation_check": explanation_check}, runs)):
                body = report_document_markdown(document)
        record = _new(
            self.task,
            requirements_id=requirements_id,
            materiality=materiality,
            completion_request_key=spec.get("completion_request_key"),
            completion_operation_id=spec.get("completion_operation_id"),
            completion_operation_revision=spec.get("completion_operation_revision"),
            source_read_failures=source_read_failures,
            evidence_appendix=evidence_appendix,
            evidence_appendix_meta=evidence_appendix_meta,
            **({"required_check_authorship": required_check_authorship} if required_check_authorship is not None else {}),
            evidence_recipes=evidence_recipes,
            evidence_report=evidence_report,
            consistency_rules=consistency_rules,
            consistency_sources=consistency_sources,
            consistency_report=consistency_report,
            title=title,
            body=body,
            **({"document": document} if document is not None else {}),
            **({"body_template": body_template} if body_template is not None else {}),
            run_ids=list(runs),
            fact_ids=fact_ids,
            claims=claims,
            revision_comparison=revision_comparison,
            explanation_check=explanation_check,
            quality=explanation_check,
            # What the labels were read from, so every reader reads them by the files it may read (draft_in_view).
            evidence_basis=evidence_basis,
            # The completion policy this check was computed under; a replay never reuses an older one.
            check_policy_version=COMPLETION_POLICY_VERSION,
            parent_id=parent.get("id"),
            # What this revision took from its parent, so it is read as resting on the parent's files (draft_snapshots).
            # The parent's facts it no longer cites are not taken.
            **({"inherited_fields": sorted(inherited - (set() if fact_ids else {"fact_ids"}))}
               if inherited is not None else {}),
            execution_status=execution_status,
            reproducibility_status=reproducibility,
            review_status=review,
            publication_status=publication,
            limitations=limitations,
            content_hash=_draft_content_hash(
                title=title,
                requirements_id=requirements_id,
                materiality=materiality,
                body=body,
                body_template=body_template,
                document=document,
                run_ids=list(runs),
                fact_ids=fact_ids,
                claims=claims,
                limitations=limitations,
            ),
        )
        return self._write(self._object_path("drafts", record["id"]), record)

    def add_report_source(self, path: Path, *, name: str | None = None) -> dict:
        from foundry.report_documents import check_report_source_bytes, report_format_for_path

        source = Path(path).expanduser()
        if not source.is_file() or source.is_symlink():
            raise FoundryError("report source must be a regular file")
        fmt = report_format_for_path(source)
        data = source.read_bytes()
        check_report_source_bytes(data, fmt)
        digest = _hash_bytes(data)
        record = _new(
            self.task,
            # A caller that built the bytes itself passes the name the manager
            # would recognise; otherwise the file on disk is the name.
            name=str(name).strip() if str(name or "").strip() else source.name,
            format=fmt,
            sha256=digest,
            bytes=len(data),
            source_file="source" + source.suffix.casefold(),
            imported_review_trust="not_trusted",
        )
        with exclusive(self.task.lock_path, timeout=10):
            root = self._inside(self.root / "report_sources" / record["id"])
            root.mkdir(parents=True, exist_ok=False)
            archived = self._inside(root / record["source_file"])
            archived.write_bytes(data)
            written = self._write_locked(self._inside(self.root / "report_sources" / (record["id"] + ".json")), record)
        return {"source": self._report_source_public(written), "path": str(archived.resolve())}

    def _report_source_path(self, source_id: str) -> Path:
        return self._inside(self.root / "report_sources" / (_assert_id(source_id) + ".json"))

    def _report_source_public(self, record: dict) -> dict:
        return {
            "id": record["id"],
            "name": record["name"],
            "sha256": record["sha256"],
            "format": record["format"],
        }

    def read_report_source(self, source_id: str) -> dict:
        record = self._read(self._report_source_path(source_id))
        path = self._inside(self.root / "report_sources" / record["id"] / record["source_file"])
        data = path.read_bytes()
        if _hash_bytes(data) != record["sha256"]:
            raise FoundryError("report source integrity check failed")
        return {"source": self._report_source_public(record), "path": str(path.resolve())}

    def read_report_xlsx(self, source_id: str, *, sheet: str | None = None, range_ref: str | None = None) -> dict:
        from foundry.report_documents import report_document_from_xlsx

        resolved = self.read_report_source(source_id)
        if resolved["source"]["format"] != "xlsx":
            raise FoundryError("report xlsx read requires an XLSX source")
        return report_document_from_xlsx(
            Path(resolved["path"]).read_bytes(),
            resolved["source"],
            sheet=sheet,
            range_ref=range_ref,
        )

    def read_report_pptx(self, source_id: str) -> dict:
        from foundry.report_pptx import report_document_from_pptx

        resolved = self.read_report_source(source_id)
        if resolved["source"]["format"] != "pptx":
            raise FoundryError("report pptx read requires a PPTX source")
        return report_document_from_pptx(Path(resolved["path"]).read_bytes(), resolved["source"])

    def register_report_export(
        self,
        draft_id: str,
        format: str,
        path: Path,
        *,
        content_mode: str | None = None,
        projection_hash: str | None = None,
        allowed: set[str] | None = None,
    ) -> dict:
        if format not in REPORT_EXPORT_FORMATS:
            raise FoundryError("report export format must be md, xlsx, docx, pdf or html")
        draft = self.get_draft(draft_id)
        if projection_hash is not None and not re.fullmatch(r"[0-9a-f]{64}", projection_hash):
            raise FoundryError("report export projection_hash must be a sha256")
        if projection_hash is None:
            content_mode = None
        elif content_mode not in {"management", "audit"}:
            raise FoundryError("report export mode must be management or audit")
        template_version = (draft.get("document") or {}).get("template_version")
        projection = None
        if projection_hash is not None and allowed is not None:
            # A session's receipt names a projection the server reads for it, never one it only asserts: an asserted
            # hash could take the archive name of another export and break it. (The CLI files only the desktop's.)
            projection = self.projection_for_draft(draft_id, content_mode, allowed=allowed)
            if projection["projection_hash"] != projection_hash:
                raise FoundryError("report export projection_hash is not this draft's projection; read the projection again")
        source = Path(path).expanduser()
        if not source.is_file() or source.is_symlink():
            raise FoundryError("report export source must be a regular file")
        data = source.read_bytes()
        limit = MAX_SOURCE_BYTES
        if len(data) > limit and projection_hash is not None:
            # The estimate declaration and labels bring their own room (issue #32): a file that fits the limit
            # without them still registers.
            projection = projection or self.projection_for_draft(draft_id, content_mode, allowed=allowed)
            if projection["projection_hash"] == projection_hash:
                limit += int((projection.get("estimate_headroom") or {}).get("bytes") or 0)
        if len(data) > limit:
            raise FoundryError("report export exceeds 32 MiB limit")
        digest = _hash_bytes(data)
        filename = (
            f"{draft_id}-{content_mode}-{projection_hash[:12]}{_registration_reader(allowed)}.{format}"
            if content_mode and projection_hash
            else f"{draft_id}{_registration_reader(allowed)}.{format}"
        )
        target = self._inside(self.root / "exports" / "registered" / filename)
        manifest_path = self._inside(target.with_suffix(target.suffix + ".meta.json"))
        with exclusive(self.task.lock_path, timeout=10):
            target.parent.mkdir(parents=True, exist_ok=True)
            if target.exists():
                manifest = self._read(manifest_path)
                if (
                    manifest.get("sha256") != _hash_bytes(target.read_bytes())
                    or manifest.get("content_hash") != draft["content_hash"]
                    or manifest.get("content_mode") != content_mode
                    or manifest.get("projection_hash") != projection_hash
                ):
                    raise FoundryError("report export integrity check failed; archived bytes differ from their receipt")
                if manifest.get("format") != format or manifest.get("template_version") != template_version:
                    raise FoundryError("report export receipt does not match the requested format/template")
                digest = manifest["sha256"]
            else:
                temporary = self._inside(target.with_name("." + uuid.uuid4().hex + ".tmp"))
                try:
                    temporary.write_bytes(data)
                    temporary.replace(target)
                    manifest = _new(
                        self.task,
                        draft_id=draft_id,
                        format=format,
                        sha256=digest,
                        bytes=len(data),
                        content_hash=draft["content_hash"],
                        content_mode=content_mode,
                        projection_hash=projection_hash,
                        template_version=template_version,
                        source_sha256=digest,
                    )
                    manifest["_integrity_hash"] = _hash(manifest)
                    write_json_atomic(manifest_path, manifest)
                finally:
                    temporary.unlink(missing_ok=True)
        return {
            "task_id": self.task.id,
            "draft_id": draft_id,
            "format": format,
            "path": str(target.resolve()),
            "sha256": digest,
            "content_hash": draft["content_hash"],
            **({"content_mode": content_mode, "projection_hash": projection_hash} if content_mode and projection_hash else {}),
            "publication_status": draft["publication_status"],
        }

    def projection_for_draft(self, draft_id: str, mode: str = "management", *, allowed: set[str] | None = None,
                             cache: dict | None = None) -> dict:
        """The draft as a report projects it, for a reader whose snapshot ceiling is ``allowed`` (draft_in_view;
        ``cache`` as for run_in_view)."""
        from foundry.financial_evidence import consistency_dataset_ids, recipe_dataset_ids
        from foundry.report_projection import build_report_projection

        draft = self.draft_in_view(self.get_draft(draft_id), allowed, cache=cache)
        dataset_ids: set[str] = set()
        runs: dict[str, dict] = {}
        for run_id in draft.get("run_ids", []) or []:
            try:
                run = self.get_run(run_id)
            except FoundryError:
                continue
            runs[run_id] = run
            bindings = run.get("bindings", {})
            if isinstance(bindings, dict):
                dataset_ids.update(value for value in bindings.values() if isinstance(value, str))
        for claim in draft.get("claims", []) or []:
            if not isinstance(claim, dict):
                continue
            if isinstance(claim.get("dataset_id"), str):
                dataset_ids.add(claim["dataset_id"])
            if isinstance(claim.get("run_id"), str):
                try:
                    run = self.get_run(claim["run_id"])
                except FoundryError:
                    continue
                runs[claim["run_id"]] = run
                bindings = run.get("bindings", {})
                if isinstance(bindings, dict):
                    dataset_ids.update(value for value in bindings.values() if isinstance(value, str))
        dataset_ids.update(recipe_dataset_ids(draft.get("evidence_recipes", [])))
        dataset_ids.update(consistency_dataset_ids(draft.get("consistency_sources", {})))
        resolved = []
        for dataset_id in sorted(dataset_ids):
            try:
                dataset = self.read_dataset(dataset_id)
            except FoundryError:
                continue
            resolved.append({
                "id": dataset["id"],
                "name": dataset.get("name") or dataset.get("source_name") or dataset["id"],
                "kind": dataset.get("kind") or dataset.get("metadata", {}).get("mode") or "dataset",
                "source_hash": dataset.get("source_hash"),
            })
        return build_report_projection(draft, mode=mode, resolved_datasets=resolved, runs=runs)

    def report_projection(self, draft_id: str, mode: str = "management", *, allowed: set[str] | None = None,
                          cache: dict | None = None) -> dict:
        return self.projection_for_draft(draft_id, mode, allowed=allowed, cache=cache)

    def export_draft(self, draft_id: str, format: str = "md", *, mode: str = "management",
                     allowed: set[str] | None = None) -> dict:
        # The draft, its projection and its run lines are one read, of the versions and runs listed once (run_in_view).
        cache: dict = {}
        draft = self.draft_in_view(self.get_draft(draft_id), allowed, cache=cache)
        if format not in {"md", "xlsx"}:
            raise FoundryError("draft export format must be md or xlsx")
        projection = self.report_projection(draft_id, mode=mode, allowed=allowed, cache=cache)
        # Each run line reads the run as this reader reads it (run_in_view).
        runs = [self.run_in_view(self.get_run(run_id), allowed, cache=cache) for run_id in draft["run_ids"]]
        path = self._inside(self.root / "exports" / (
            f"{draft_id}-{mode}-{projection['projection_hash'][:12]}{_export_reader(mode, allowed, runs)}.{format}"))
        with exclusive(self.task.lock_path, timeout=10):
            self._inside(path)
            manifest_path = self._inside(path.with_suffix(path.suffix + ".meta.json"))
            if path.exists():
                manifest = self._read(manifest_path)
                if (
                    manifest.get("sha256") != _hash_bytes(path.read_bytes())
                    or manifest.get("content_hash") != draft["content_hash"]
                    or manifest.get("projection_hash") != projection["projection_hash"]
                ):
                    raise FoundryError(
                        "export integrity check failed; archived bytes differ from their receipt"
                    )
            else:
                if format == "md":
                    lines = [
                        projection["body"].rstrip(),
                        "",
                        "---",
                        "",
                        f"Version: {draft_id}",
                        f"Projection: {projection['content_mode']}; SHA-256: {projection['projection_hash']}",
                        f"Status: {projection['publication_status']}; review: {projection['review_status']}",
                        f"Content SHA-256: {projection['content_hash']}",
                    ]
                    if projection.get("document"):
                        import html
                        from foundry.report_documents import _escape_markdown_text
                        from foundry.report_visuals import document_visuals

                        for visual in document_visuals(projection["document"]):
                            if visual["kind"] == "flow":
                                continue
                            title = _escape_markdown_text(html.escape(visual["title"], quote=False)).replace("\n", " ")
                            lines.append(f"Figure: {title}; chart_id: {visual['chart_id']}; run_id: {visual['run_id']}; snapshot_hash: {visual['snapshot_hash']}")
                    if draft["limitations"]:
                        lines += [
                            "",
                            "局限：",
                            *["- " + str(v) for v in draft["limitations"]],
                        ]
                    for run in runs:
                        lines += [
                            "",
                            f"Run: {run['id']}; execution: {run['execution_status']}; reproducibility: {run['reproducibility_status']}",
                            f"Output SHA-256: {run.get('output_hash')}",
                            "```" + (run.get("engine") or "sql"),
                            run.get("sql") or run.get("code") or (
                                json.dumps(run["aggregate_plan"], ensure_ascii=False, indent=2)
                                if run.get("aggregate_plan") else ""
                            ),
                            "```",
                            "Inputs: "
                            + json.dumps(run["input_hashes"], ensure_ascii=False),
                            "Sources: " + json.dumps(run.get("source_hashes", {}), ensure_ascii=False),
                        ]
                        if run.get("engine") == "structured_aggregate":
                            result = run.get("result", {})
                            lines += ["来源核验：" + str(result.get("provenance", {}).get("status", "unverified")),
                                      "```json", json.dumps({key: result.get(key) for key in (
                                          "provenance", "detail_rows", "excluded_detail_rows")},
                                          ensure_ascii=False, indent=2), "```"]
                    for run in runs:
                        if run.get("engine") in {"python", "shell"}:
                            lines += [
                                "",
                                "Script stdout:",
                                "```text",
                                run.get("stdout", ""),
                                "```",
                                "Script stderr:",
                                "```text",
                                run.get("stderr", ""),
                                "```",
                            ]
                            for artifact in run.get("outputs", []):
                                lines.append(
                                    "Artifact: "
                                    + artifact["artifact_path"]
                                    + "; SHA-256: "
                                    + artifact["sha256"]
                                )
                    data = ("\n".join(lines) + "\n").encode("utf8")
                else:
                    workbook = Workbook()
                    sheet = workbook.active
                    sheet.title = "Report"
                    for row in [
                        [projection["title"]],
                        [projection["marker"]],
                        [projection["body"]],
                        ["Version", draft_id],
                        ["Projection", projection["content_mode"]],
                        ["Projection SHA-256", projection["projection_hash"]],
                        ["Content SHA-256", projection["content_hash"]],
                        ["Publication", projection["publication_status"]],
                        ["Review", projection["review_status"]],
                        ["Reproducibility", projection["reproducibility_status"]],
                    ]:
                        _append_excel(sheet, row)
                    if projection.get("document"):
                        # The projection's estimate declaration and labels get the room it gave them.
                        headroom = projection.get("estimate_headroom") or {}
                        for sheet_name, rows in report_document_excel_rows(
                            projection["document"], projection,
                            text_headroom=headroom.get("text", 0), node_headroom=headroom.get("nodes", 0),
                        ).items():
                            structured_sheet = workbook.create_sheet(sheet_name)
                            for row in rows:
                                _append_excel(structured_sheet, row)
                    from foundry.report_estimates import ESTIMATE_LABEL, estimated_cells

                    # Cells the server's requirement check stored on this draft found estimated or declared.
                    estimates = {entry["cell"] for entry in estimated_cells(draft, {run["id"]: run for run in runs})
                                 if entry["cell"] is not None}
                    for number, run in enumerate(runs, start=1):
                        target = workbook.create_sheet("Result " + str(number))
                        result = run.get("result") or {"columns": [], "rows": []}
                        columns = [c["name"] for c in result["columns"]]
                        _append_excel(target, columns)
                        for row_index, row in enumerate(result["rows"]):
                            row_start = target.max_row + 1
                            _append_excel(
                                target,
                                [
                                    str(row.get(c)) if row.get(c) is not None else None
                                    for c in columns
                                ],
                            )
                            for column_index, column in enumerate(result["columns"], start=1):
                                value = row.get(column["name"])
                                numeric_type = str(column.get("type", "")).upper()
                                if value is None or not numeric_type.startswith(("DECIMAL", "NUMERIC", "INTEGER", "BIGINT", "SMALLINT", "HUGEINT", "DOUBLE", "FLOAT")):
                                    continue
                                from decimal import InvalidOperation

                                try:
                                    amount_value = Decimal(str(value))
                                except InvalidOperation:
                                    continue
                                cell = target.cell(row_start, column_index)
                                if amount_value.is_finite() and survives_a_cell(amount_value):
                                    put_number(cell, value)
                                else:
                                    cell.comment = Comment("超过 Excel 可安全保留的数值精度，按精确文本保存。", "Foundry")
                            for column_index, column in enumerate(columns, start=1):
                                cell = target.cell(row_start, column_index)
                                if (run["id"], row_index, column) in estimates and cell.value not in (None, ""):
                                    # The value stays what the run computed; its display carries the label.
                                    shown = cell.number_format if cell.data_type == "n" else "@"
                                    cell.number_format = f'{shown}"{ESTIMATE_LABEL}"'
                        evidence = workbook.create_sheet("Evidence " + str(number))
                        for key in [
                            "id",
                            "engine",
                            "sql",
                            "code",
                            "stdout",
                            "stderr",
                            "outputs",
                            "parameters",
                            "bindings",
                            "input_hashes",
                            "source_hashes",
                            "aggregate_plan",
                            "output_hash",
                            "execution_status",
                            "reproducibility_status",
                            "assumptions",
                        ]:
                            _append_excel(
                                evidence,
                                [
                                    key,
                                    json.dumps(run.get(key), ensure_ascii=False)
                                    if isinstance(run.get(key), (list, dict))
                                    else str(run.get(key)),
                                ],
                            )
                        if run.get("engine") == "structured_aggregate":
                            _append_excel(evidence, ["provenance", json.dumps(result.get("provenance", {}), ensure_ascii=False)])
                            for section, label in (("detail_rows", "Sources "), ("excluded_detail_rows", "Excluded ")):
                                source_sheet = workbook.create_sheet(label + str(number))
                                fields = ["result_key", "included", "exclusion_reason", "source", "measure_values"]
                                _append_excel(source_sheet, fields)
                                for detail in result.get(section, []):
                                    _append_excel(source_sheet, [
                                        json.dumps(detail.get(key), ensure_ascii=False)
                                        if isinstance(detail.get(key), (dict, list)) else detail.get(key)
                                        for key in fields
                                    ])
                    # Keep text literal, including identifiers and amounts that
                    # exceed Excel precision. Safe typed amounts stay numeric.
                    from foundry.report_visuals import add_visual_worksheets

                    add_visual_worksheets(workbook, projection["document"])
                    for sheet in workbook:
                        for row in sheet:
                            for cell in row:
                                if isinstance(cell.value, str):
                                    cell.data_type = "s"
                    output = io.BytesIO()
                    workbook.save(output)
                    data = output.getvalue()
                path.parent.mkdir(parents=True, exist_ok=True)
                temporary = self._inside(
                    path.with_name("." + uuid.uuid4().hex + ".tmp")
                )
                try:
                    temporary.write_bytes(data)
                    temporary.replace(path)
                    manifest = _new(
                        self.task,
                        draft_id=draft_id,
                        format=format,
                        sha256=_hash_bytes(data),
                        content_hash=draft["content_hash"],
                        content_mode=projection["content_mode"],
                        projection_hash=projection["projection_hash"],
                    )
                    manifest["_integrity_hash"] = _hash(manifest)
                    write_json_atomic(manifest_path, manifest)
                finally:
                    temporary.unlink(missing_ok=True)
            digest = _hash_bytes(path.read_bytes())
        return {
            "task_id": self.task.id,
            "draft_id": draft_id,
            "format": format,
            "path": str(path.resolve()),
            "sha256": digest,
            "content_hash": draft["content_hash"],
            "content_mode": projection["content_mode"],
            "projection_hash": projection["projection_hash"],
            "publication_status": draft["publication_status"],
        }

    def save_method(self, analysis_id: str, name: str) -> dict:
        analysis = self.get_analysis(analysis_id)
        if analysis.get("engine") in {"python", "shell", "structured_aggregate"}:
            raise FoundryError(
                "non-SQL revisions can be rerun; saving a reusable method currently requires a SQL analysis"
            )
        _safe_text(name, "method name", required=True, maximum=200)
        # Values in WHERE clauses are data too. Replace SQL string literals with
        # caller-supplied parameters and omit original question/values/assumptions.
        sql, parameter_contract = method_template(
            analysis["sql"], analysis.get("parameter_metadata", [])
        )
        contracts = {
            alias: {
                "schema": self.read_dataset(dataset_id)["schema"],
                "grain": self.read_dataset(dataset_id)["metadata"].get("grain", []),
            }
            for alias, dataset_id in analysis["bindings"].items()
        }
        record = _new(
            self.task,
            name=name,
            source_task_id=self.task.id,
            sql=sql,
            input_contracts=contracts,
            relationships=analysis["relationships"],
            metrics=analysis["metrics"],
            parameter_contract=parameter_contract,
            parameter_metadata=[
                entry["metadata"] for entry in parameter_contract if "metadata" in entry
            ],
            parameter_policy_version=1,
            parameters_required=any(
                entry["role"] != "formula_constant" for entry in parameter_contract
            ),
            parameter_note="Supply new values by canonical parameter key. Only declared formula constants have defaults; original business values, business defaults and choice lists are not retained.",
            method_kind="financial-analysis",
            max_rows=analysis["max_rows"],
            timeout_seconds=analysis["timeout_seconds"],
        )
        return self._write(self._object_path("methods", record["id"]), record)

    def apply_method(
        self,
        method_id: str,
        bindings: dict,
        *,
        question: str | None = None,
        parameters: Any = None,
        session_id: str | None = None,
        context_revision: int | None = None,
    ) -> dict:
        method = self._get_method(method_id)
        if not isinstance(bindings, dict) or set(bindings) != set(
            method["input_contracts"]
        ):
            raise FoundryError("method bindings must match its input aliases")
        for alias, dataset_id in bindings.items():
            data = self.read_dataset(dataset_id)
            columns = {c["name"]: c["type"] for c in data["schema"]}
            for expected in method["input_contracts"][alias]["schema"]:
                actual = columns.get(expected["name"])
                if actual is None or (
                    expected["type"].startswith("DECIMAL")
                    != actual.startswith("DECIMAL")
                ):
                    raise FoundryError(
                        f"method input {alias} is missing compatible column {expected['name']}"
                    )
        return self.save_analysis(
            {
                "question": question or method["name"],
                "sql": method["sql"],
                "bindings": bindings,
                "parameters": method_parameters(method, parameters),
                "parameter_metadata": method.get("parameter_metadata", []),
                "relationships": method["relationships"],
                "metrics": method["metrics"],
                "method_id": method_id,
                "max_rows": method["max_rows"],
                "timeout_seconds": method["timeout_seconds"],
            },
            session_id=session_id,
            context_revision=context_revision,
        )

    def run_script(
        self,
        spec: dict,
        *,
        session_id: str | None = None,
        context_revision: int | None = None,
        cancel_check=None,
        analysis_id: str | None = None,
        allowed: set[str] | None = None,
        frozen_period: dict | None = None,
    ) -> dict:
        """Broker explicit snapshots to one isolated workspace and archive outputs.

        A job may pass a cancellation callback and frozen call context. The
        caller validates authorization before dispatch; this service validates
        ownership/hashes again and never silently rebinds inputs during a run.
        """
        from foundry.runner import run_script
        from foundry.finance_requirement_semantics import capture_requirements
        requirements_snapshot = capture_requirements(self, allowed=allowed)

        if not isinstance(spec, dict):
            raise FoundryError("script spec must be an object")
        code = _safe_text(
            spec.get("code"), "script code", required=True, maximum=256 * 1024
        )
        language = spec.get("language", "python")
        if language not in {"python", "shell"}:
            raise FoundryError("script language must be python or shell")
        bindings = spec.get("bindings") or {}
        if not isinstance(bindings, dict) or len(bindings) > 32:
            raise FoundryError("script bindings must be a map of at most 32 snapshots")
        if any(
            not isinstance(alias, str) or not ALIAS.fullmatch(alias)
            for alias in bindings
        ):
            raise FoundryError("script binding aliases must be simple identifiers")
        if len({alias.casefold() for alias in bindings}) != len(bindings):
            raise FoundryError("script aliases must be unique")
        datasets = {
            alias: self.read_dataset(dataset_id)
            for alias, dataset_id in bindings.items()
        }
        assumptions = spec.get("assumptions", [])
        if not isinstance(assumptions, list):
            raise FoundryError("script assumptions must be an array")
        result_file = _script_result_path(spec.get("result_file"))
        timeout = spec.get("timeout_seconds", 30)
        if (
            isinstance(timeout, bool)
            or not isinstance(timeout, (int, float, Decimal))
            or not 0 < timeout <= 300
        ):
            raise FoundryError("script timeout_seconds must be between 0 and 300")
        timeout = float(timeout)
        requirements_id = _requirements_id(spec.get("requirements_id"))
        if requirements_id:
            self._get("requirements", requirements_id)
        default_result_role = "business" if result_file else "activity"
        result_role = _result_role(spec.get("result_role"), default=default_result_role)
        scope = _scope_metadata(spec.get("scope"))
        period_binding = normalize_period_binding(spec.get("period_binding"), engine=language)
        topic = _script_topic(spec.get("topic"))
        revision = _new(
            self.task,
            question=spec.get("question", "脚本分析"),
            engine=language,
            **({"topic": topic} if topic else {}),
            code=code,
            sql=None,
            bindings=bindings,
            parameters=[],
            assumptions=assumptions,
            result_file=result_file,
            parent_id=None,
            requirements_id=requirements_id,
            result_role=result_role,
            scope=scope,
            **({"period_binding": period_binding} if period_binding else {}),
            timeout_seconds=timeout,
            relationships=[],
            metrics=[],
            grain=[],
            missing_inputs=[],
            **_saved_by(session_id, context_revision),
        )
        if analysis_id:
            # Pinned as run_analysis runs it: to the version the session reads the saved analysis as naming.
            revision = self.analysis_in_view(self.get_analysis(analysis_id), allowed)
            requirements_id = revision.get("requirements_id")
            result_role = revision.get("result_role", "business")
            scope = revision.get("scope", {})
            period_binding = revision.get("period_binding")
        else:
            self._write(self._object_path("revisions", revision["id"]), revision)
        run_id = uuid.uuid4().hex
        workspace = self._inside(self.root / "run_workspaces" / run_id)
        inputs = self._inside(workspace / "inputs")
        inputs.mkdir(parents=True, exist_ok=False)
        for alias, dataset in datasets.items():
            # Numeric JSON fields are strings intentionally; scripts can convert
            # them with Decimal and have schema/metadata alongside the rows.
            (inputs / (alias + ".json")).write_bytes(_bytes(dataset))
            canonical = self._inside(
                self.root / "datasets" / dataset["id"] / dataset["source_file"]
            )
            (
                inputs / (alias + ".source" + Path(dataset["source_file"]).suffix)
            ).write_bytes(canonical.read_bytes())
        resolved_period = resolve_period(self, datasets, scope, allowed=allowed, frozen=frozen_period) if period_binding else None
        if resolved_period:
            scope = {**scope, "period": resolved_period["period"]}
        options = {"language": language, "timeout": timeout}
        if resolved_period:
            options["context"] = resolved_period
        if cancel_check is not None:
            options["cancel_check"] = cancel_check
        execution = run_script(workspace, code, **options)
        succeeded = execution["status"] == "succeeded"
        outputs = []
        output_root = self._inside(self.root / "run_outputs" / run_id)
        for item in execution.get("outputs", []):
            relative = Path(item["path"])
            if relative.is_absolute() or ".." in relative.parts:
                raise FoundryError("runner output escapes its run workspace")
            source = self._inside(Path(execution["workspace"]) / relative)
            if (
                source.is_symlink()
                or not source.is_file()
                or source.stat().st_nlink != 1
            ):
                raise FoundryError("runner output must be a regular nonlinked file")
            data = source.read_bytes()
            if _hash_bytes(data) != item["sha256"]:
                raise FoundryError("runner output integrity check failed")
            destination = self._inside(output_root / relative)
            destination.parent.mkdir(parents=True, exist_ok=True)
            with destination.open("xb") as handle:
                handle.write(data)
            outputs.append(
                {**item, "artifact_path": str(destination.relative_to(self.root))}
            )
        result = _empty_script_result()
        result_error = None
        if succeeded and result_file:
            try:
                item = next(
                    (item for item in outputs if item["path"] == result_file),
                    None,
                )
                if item is None:
                    raise FoundryError(
                        "script result_file is not in this run output manifest"
                    )
                archived = self._inside(output_root / result_file)
                data = archived.read_bytes()
                if _hash_bytes(data) != item["sha256"]:
                    raise FoundryError("script result_file integrity check failed")
                parsed = parse_source(data, Path(result_file).suffix.casefold(), {})
                if parsed.get("kind") != "table":
                    raise FoundryError("script result_file must be a CSV or JSON table")
                row_count = parsed["row_count"]
                rows = parsed["rows"][:SCRIPT_RESULT_PREVIEW_ROWS]
                result = {
                    "columns": parsed["schema"],
                    "rows": rows,
                    "row_count": row_count,
                    "truncated": row_count > len(rows),
                    "result_source": {
                        "path": result_file,
                        "sha256": item["sha256"],
                        "format": Path(result_file).suffix.lstrip(".").casefold(),
                        "basis": "script_output",
                        "reproducibility_status": "unverified",
                    },
                }
            except FoundryError as exc:
                result_error = str(exc)
                succeeded = False
        output_hash = (
            _hash(
                {
                    "stdout": execution.get("stdout", ""),
                    "outputs": [
                        {k: v for k, v in item.items() if k != "artifact_path"}
                        for item in outputs
                    ],
                    "result": result,
                }
            )
            if succeeded
            else None
        )
        requirements_snapshot = {
            **requirements_snapshot,
            "requirements_id": requirements_id or requirements_snapshot.get("requirements_id"),
        }
        # Declared from the full parsed table (the stored result keeps only a preview) and kept beside the
        # result, never in it: output_hash covers the result, so the same output keeps the same hash.
        declaration = {}
        if succeeded and result_file:
            from foundry.result_semantics import declare_script_result

            declaration["row_declarations"] = declare_script_result(
                self, parsed["schema"], parsed["rows"], requirements_snapshot["requirements_id"], allowed=allowed
            )
        record = _new(
            self.task,
            analysis_id=revision["id"],
            **requirements_snapshot,
            **declaration,
            result_role=result_role,
            scope=scope,
            session_id=session_id,
            context_revision=context_revision,
            engine=language,
            code=code,
            sql=None,
            parameters=[],
            bindings=bindings,
            input_hashes={
                alias: dataset["data_hash"] for alias, dataset in datasets.items()
            },
            **({"period_binding": period_binding, "resolved_period": resolved_period,
                "execution_input_hash": _hash({"datasets": {alias: data["data_hash"] for alias, data in datasets.items()},
                                                "period": resolved_period})}
               if resolved_period else {}),
            source_hashes={
                alias: dataset["source_hash"] for alias, dataset in datasets.items()
            },
            code_hash=_hash_bytes(code.encode("utf8")),
            environment=execution.get("environment", {}),
            execution_status="succeeded" if succeeded else "failed",
            runner_status=execution["status"],
            reproducibility_status="unverified",
            review_status="unreviewed",
            publication_status="draft",
            output_hash=output_hash,
            stdout=execution.get("stdout", ""),
            stderr=execution.get("stderr", ""),
            exit_code=execution.get("exit_code"),
            # Beside the result, never in it: output_hash covers the result, and a slower run of the same output
            # must keep its hash.
            duration_seconds=execution.get("duration_seconds"),
            outputs=outputs,
            assumptions=assumptions,
            missing_inputs=[],
            result=result,
            warnings=[
                "脚本输出已随代码和输入一并保存，尚未核验能否复现。"
            ],
        )
        record["id"] = run_id
        if not succeeded:
            record["error"] = (
                result_error or execution.get("stderr") or execution["status"]
            )
            record["error_code"] = (
                "result_file_invalid" if result_error else execution["status"]
            )
        return self._write(self._object_path("runs", run_id), record)

    def output_path(self, run_id: str, output: str) -> dict:
        """Resolve a recorded artifact name; never accept a renderer host path."""
        run = self.get_run(run_id)
        if not isinstance(output, str):
            raise FoundryError("output must be a recorded artifact name")
        item = next(
            (item for item in run.get("outputs", []) if item["path"] == output), None
        )
        if item is None:
            raise FoundryError("output is not in this run manifest")
        path = self._inside(self.root / item["artifact_path"])
        return {
            "task_id": self.task.id,
            "run_id": run_id,
            "output": output,
            "path": str(path.resolve()),
            "sha256": item["sha256"],
            "size": item["size"],
        }
