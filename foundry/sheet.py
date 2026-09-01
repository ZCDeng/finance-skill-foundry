from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from foundry.questions import load_questions
from foundry.references import load_references
from foundry.tags import AGREED, CLAIMED, EXTERNAL, NOT_IN_WINDOW, OUTSIDE, RECOMPUTED
from foundry.window import Window
from foundry.workpaper import display_origin, display_tag, hash_pairs, load_schedules

"""
The close as a workbook, written from the lock.

A finance manager reads a spreadsheet. Everything this product produces is
markdown, which is right for the audit trail and wrong for the person: the one
document they most want to sort, filter and pivot arrives as 394 rows of pipes.

What makes this safe is where the bytes come from. The workbook is rendered
from ``schedules.json`` and ``brief_facts.json`` -- the same two files
``workpaper.md`` and ``brief.md`` are rendered from -- by this process, not by
a model with a spreadsheet tool. Nothing is computed here. If a figure is in
this file it was already in the lock, and the mark it carries is the mark the
brief gave it.

The three columns that make it a working paper rather than a table dump are
`条目`, `标签` and `依据`: which schedule a line belongs to, what the figure
claims about itself, and whether the lock produced it. Repeated on every row on
purpose -- a pivot needs them on the row, not in a heading above it.
"""

# xlsx numbers are binary64. Beyond that the format cannot hold what the
# catalog can, and `catalog.py` already refuses to IMPORT such a value. The
# symmetric rule on the way out is not to refuse -- the figure is real and the
# operator needs to see it -- but never to write one as a number that a
# spreadsheet would then round. Those go in as text, exact.
MONEY = "#,##0.00"
INK = "1F1F1F"
RULE = Side(style="thin", color="D9D9D9")


def _decimal(value: Any) -> Decimal | None:
    """The exact value behind a schedule figure, or nothing if it is not one.

    `schedules.json` carries an integer as an int and everything else as a
    decimal STRING, so that JSON never turns 282670.7525 into a float. Both
    shapes arrive here.
    """
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, Decimal):
        return value
    if isinstance(value, str):
        try:
            return Decimal(value.strip())
        except (InvalidOperation, ValueError):
            return None
    return None


def survives_a_cell(value: Decimal) -> bool:
    """Whether binary64 can hold this exactly.

    `repr` of a float is the shortest string that round-trips, so this asks the
    only question that matters: read back out of the file, is it the same
    number the lock holds.
    """
    try:
        return Decimal(repr(float(value))) == value
    except (OverflowError, ValueError, InvalidOperation):
        return False


def put_number(cell: Any, value: Any) -> None:
    """One figure into one cell, as a number when that is honest.

    A value the format cannot hold is written as text rather than as a number
    that is nearly right. It reads the same to a person and refuses to lie to
    a formula, which is the correct trade for a working paper: a figure a
    spreadsheet silently rounded is the exact failure this product exists to
    make impossible.
    """
    exact = _decimal(value)
    if exact is None:
        cell.value = "" if value is None else str(value)
        return
    if survives_a_cell(exact):
        cell.value = int(exact) if exact == exact.to_integral_value() else float(exact)
        cell.number_format = MONEY
        return
    cell.value = format(exact.normalize(), "f")
    cell.alignment = Alignment(horizontal="right")


def _head(sheet: Any, row: int, labels: list[str]) -> None:
    fill = PatternFill("solid", fgColor="F2F2F2")
    for column, label in enumerate(labels, start=1):
        cell = sheet.cell(row=row, column=column, value=label)
        cell.font = Font(bold=True, color=INK)
        cell.fill = fill
        cell.border = Border(bottom=RULE)
    sheet.freeze_panes = sheet.cell(row=row + 1, column=1)


def _widths(sheet: Any, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


LEGEND = [
    (CLAIMED, "报表自己声称的合计，原样抄过来，没有动过"),
    (RECOMPUTED, "本窗口的锁库重新算出来的"),
    (AGREED, "上面两个一致"),
    (OUTSIDE, "结论稿上没有的追问，重算出来的"),
    (EXTERNAL, "不出自本窗口的锁库。外部材料，带出处和查到日期"),
    (NOT_IN_WINDOW, "这个问题要用的表不在本窗口，拒绝作答"),
]


def _cover(book: Workbook, window: Window, task: dict[str, Any] | None) -> None:
    sheet = book.active
    sheet.title = "本期"
    _widths(sheet, [16, 72])

    sheet["A1"] = "本期结账"
    sheet["A1"].font = Font(bold=True, size=16, color=INK)

    row = 3
    facts: list[tuple[str, str]] = [("窗口", window.id), ("状态", window.state)]
    if task:
        facts += [
            ("任务", str(task.get("title") or task.get("id") or "")),
            ("期间", str(task.get("period") or "")),
        ]
    for label, value in facts:
        sheet.cell(row=row, column=1, value=label).font = Font(bold=True, color=INK)
        sheet.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="这份表怎么来的").font = Font(bold=True, color=INK)
    sheet.cell(
        row=row + 1,
        column=2,
        value="由 foundry 从锁后的窗口目录直接写出，不是模型起草，也不是归档 Excel 的副本。"
        "每个数字的标签见下表；底稿页每一行都带着它属于哪一条。",
    ).alignment = Alignment(wrap_text=True, vertical="top")
    row += 3

    sheet.cell(row=row, column=1, value="标签").font = Font(bold=True, color=INK)
    sheet.cell(row=row, column=2, value="意思").font = Font(bold=True, color=INK)
    row += 1
    for tag, meaning in LEGEND:
        sheet.cell(row=row, column=1, value=tag)
        sheet.cell(row=row, column=2, value=meaning)
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="锁哈希").font = Font(bold=True, color=INK)
    row += 1
    for table, digest in hash_pairs(window.hashes()):
        sheet.cell(row=row, column=1, value=table)
        sheet.cell(row=row, column=2, value=digest).font = Font(name="Menlo", size=9)
        row += 1


def _conclusions(book: Workbook, window: Window, facts: list[dict[str, Any]]) -> None:
    sheet = book.create_sheet("结论")
    _widths(sheet, [14, 44, 12, 10, 18, 14, 30, 12, 34])
    _head(sheet, 1, ["条目", "问题", "标签", "依据", "数值", "来源表", "出处", "查到于", "说明"])

    questions = {str(q.get("id")): q for q in load_questions(window)}
    row = 2
    for fact in facts:
        item_id = str(fact.get("id") or "")
        sheet.cell(row=row, column=1, value=item_id)
        sheet.cell(row=row, column=2, value=str(fact.get("text") or "")).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
        sheet.cell(row=row, column=3, value=str(fact.get("tag") or ""))
        sheet.cell(row=row, column=4, value=str(fact.get("basis") or "lock"))
        if fact.get("tag") == NOT_IN_WINDOW:
            # No figure, and the 数值 column stays empty rather than holding a
            # table name. A reader scanning that column for numbers should find
            # numbers or nothing.
            sheet.cell(
                row=row,
                column=9,
                value=f"缺表 {'；'.join(fact.get('missing') or [])}，拒绝作答。",
            )
        elif fact.get("basis") == "external":
            sheet.cell(row=row, column=9, value=f"{EXTERNAL}，不出自本窗口的锁库。")
        else:
            put_number(sheet.cell(row=row, column=5), fact.get("value"))
        required = fact.get("required_table_ids") or questions.get(item_id, {}).get(
            "required_table_ids"
        )
        sheet.cell(
            row=row,
            column=6,
            value=str(fact.get("table") or "") or "，".join(str(t) for t in (required or [])),
        )
        sheet.cell(row=row, column=7, value=str(fact.get("source") or "")).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
        sheet.cell(row=row, column=8, value=str(fact.get("at") or ""))
        sheet.cell(row=row, column=9).alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    sheet.auto_filter.ref = f"A1:I{max(row - 1, 1)}"


def _paper(book: Workbook, schedules: list[dict[str, Any]]) -> None:
    sheet = book.create_sheet("底稿")
    _widths(sheet, [14, 34, 12, 12, 12, 14, 22, 18, 16, 12, 40])
    # `说明` exists for the rows that have no arithmetic. Without it a
    # reference's source would have to borrow the 科目名称 column and a refusal
    # would borrow 对照, and a sheet that reuses a column for two meanings is a
    # sheet nobody can filter.
    _head(
        sheet,
        1,
        ["条目", "问题", "标签", "来源", "表", "科目", "科目名称", "辅助", "金额", "对照", "说明"],
    )

    row = 2
    for item in schedules:
        item_id = str(item.get("id") or "")
        text = str(item.get("text") or "")
        tag = display_tag(item)
        origin = display_origin(item)
        table = str(item.get("table") or "")

        def stamp(target: int) -> None:
            sheet.cell(row=target, column=1, value=item_id)
            sheet.cell(row=target, column=2, value=text).alignment = Alignment(
                wrap_text=True, vertical="top"
            )
            sheet.cell(row=target, column=3, value=tag)
            sheet.cell(row=target, column=4, value=origin)
            sheet.cell(row=target, column=5, value=table)

        lines = item.get("lines") or []
        if not lines:
            # A reference has no arithmetic and a refusal has no table. Both
            # still get a row, or the sheet would quietly hold fewer schedules
            # than the workpaper does.
            stamp(row)
            reference = item.get("reference") or {}
            if reference:
                sheet.cell(
                    row=row,
                    column=11,
                    value=f"{EXTERNAL}，不出自本窗口的锁库，无附表行。"
                    f"出处 {reference.get('source') or '—'}，查到于 {reference.get('at') or '—'}。",
                )
            elif item.get("missing"):
                sheet.cell(
                    row=row,
                    column=11,
                    value=f"缺表 {'；'.join(item.get('missing') or [])}，拒绝作答。",
                )
                sheet.cell(row=row, column=10, value=NOT_IN_WINDOW)
            sheet.cell(row=row, column=11).alignment = Alignment(wrap_text=True, vertical="top")
            row += 1
            continue

        for line in lines:
            stamp(row)
            # Text, always. An account code is an identifier that happens to be
            # digits, and a spreadsheet that read it as a number would drop the
            # leading zero that tells 0501 from 501.
            code = sheet.cell(row=row, column=6, value=str(line.get("account") or ""))
            code.alignment = Alignment(horizontal="left")
            code.number_format = "@"
            sheet.cell(row=row, column=7, value=str(line.get("account_name") or ""))
            sheet.cell(row=row, column=8, value=str(line.get("aux") or ""))
            put_number(sheet.cell(row=row, column=9), line.get("amount"))
            sheet.cell(row=row, column=10, value=str(line.get("match") or ""))
            row += 1
    sheet.auto_filter.ref = f"A1:K{max(row - 1, 1)}"


def _references(book: Workbook, window: Window) -> None:
    entries = load_references(window)
    if not entries:
        return
    sheet = book.create_sheet(EXTERNAL)
    _widths(sheet, [14, 52, 40, 14, 30])
    _head(sheet, 1, ["条目", "内容", "来源出处", "查到于", "备注"])
    for row, entry in enumerate(entries, start=2):
        sheet.cell(row=row, column=1, value=str(entry.get("id") or ""))
        sheet.cell(row=row, column=2, value=str(entry.get("text") or "")).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
        sheet.cell(row=row, column=3, value=str(entry.get("source") or "")).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
        sheet.cell(row=row, column=4, value=str(entry.get("at") or ""))
        sheet.cell(row=row, column=5, value=str(entry.get("note") or ""))


# openpyxl writes `dcterms:created` and `dcterms:modified` into
# `docProps/core.xml` from the clock, which makes two renders of one unchanged
# window different files. That is not cosmetic here: a deliverable version is
# keyed on its bytes, so every export past the first would file a new version
# of an identical workbook and the timeline would fill with versions that are
# not versions -- the exact thing `register_bytes` dedupes to prevent.
#
# So the properties describe the window instead of the moment. The epoch is a
# sentinel and is meant to read as one; the honest timestamp is on the
# deliverable row, where the store stamps it, and the honest provenance is the
# window id and the lock hashes on the cover.
STAMP = datetime(1970, 1, 1, tzinfo=timezone.utc).replace(tzinfo=None)


def _stamp(book: Workbook, window: Window) -> None:
    """Say what this file is, in the file's own properties."""
    book.properties.creator = "foundry"
    book.properties.lastModifiedBy = "foundry"
    book.properties.title = f"本期结账 {window.id}"
    book.properties.created = STAMP
    book.properties.modified = STAMP


STAMP_TEXT = "1970-01-01T00:00:00Z"
MODIFIED = re.compile(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)")
# Any fixed pair works; this one is the DOS epoch a zip can actually store.
ZIP_EPOCH = (1980, 1, 1, 0, 0, 0)


def _settle(data: bytes) -> bytes:
    """Take the clock back out of a saved workbook.

    Two writers put it in. `openpyxl.writer.excel.save_workbook` assigns
    `properties.modified = now` immediately before writing, so setting that
    property beforehand is overwritten every time; and `ZipFile` stamps each
    member from the clock, at two-second resolution, which is why a pair of
    renders sometimes matched and sometimes did not.

    Rebuilt rather than patched in place because both live in different parts
    of the archive. Entry order and compression are preserved -- this is the
    same workbook, with the two fields that describe *when it was written*
    replaced by the sentinel that says the question is answered elsewhere.
    """
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            payload = source.read(entry.filename)
            if entry.filename == "docProps/core.xml":
                payload = MODIFIED.sub(rb"\g<1>" + STAMP_TEXT.encode() + rb"\g<2>", payload)
            settled = zipfile.ZipInfo(entry.filename, date_time=ZIP_EPOCH)
            settled.compress_type = entry.compress_type
            settled.external_attr = entry.external_attr
            target.writestr(settled, payload)
    return out.getvalue()


def workbook_bytes(window: Window, *, task: dict[str, Any] | None = None) -> bytes:
    """The whole close as one file, rendered from the window's own two records.

    Read rather than recomputed, deliberately. `brief.md` and `workpaper.md`
    come from `brief_facts.json` and `schedules.json`; so does this. Three
    renderings of one set of facts cannot disagree; three computations of them
    could.
    """
    import json

    facts_path = window.root / "brief_facts.json"
    facts = json.loads(facts_path.read_text()) if facts_path.exists() else []
    # `load_schedules` hands back the whole payload -- window id, hashes,
    # items -- because that is what `schedules.json` holds.
    schedules = load_schedules(window).get("items") or []

    book = Workbook()
    _stamp(book, window)
    _cover(book, window, task)
    _conclusions(book, window, facts)
    _paper(book, schedules)
    _references(book, window)

    buffer = io.BytesIO()
    book.save(buffer)
    return _settle(buffer.getvalue())
