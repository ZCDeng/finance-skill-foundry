from __future__ import annotations

import csv
import hashlib
import io
import json
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import duckdb
from openpyxl import load_workbook

from foundry.errors import LockRefused

MAX_SAFE_XLSX_INTEGER = 9_007_199_254_740_991


def file_hash(path: Path) -> str:
    try:
        payload = path.read_bytes()
    except OSError as exc:
        raise LockRefused(f"source file is unreadable: {path}") from exc
    return hashlib.sha256(payload).hexdigest()


def table_hash(headers: list[str], rows: list[tuple[Any, ...]]) -> str:
    payload = json.dumps({"headers": headers, "rows": rows}, default=str, sort_keys=True)
    return hashlib.sha256(payload.encode()).hexdigest()


def _normalize(name: str) -> str:
    return sql_identifier(name)


def sql_identifier(name: str) -> str:
    cleaned = "".join(ch if ch.isalnum() or ch == "_" else "_" for ch in str(name).strip().lower())
    while "__" in cleaned:
        cleaned = cleaned.replace("__", "_")
    cleaned = cleaned.strip("_")
    if not cleaned:
        raise LockRefused("empty identifier")
    if cleaned[0].isdigit():
        cleaned = f"t_{cleaned}"
    return cleaned


def _normalize_headers(raw_headers: list[Any], path: Path) -> list[str]:
    headers = [_normalize(str(h or "")) for h in raw_headers]
    _assert_no_identifier_collisions(f"{path} column", headers)
    return headers


def _normalize_rows(path: Path, headers: list[str], rows: list[tuple[Any, ...]]) -> list[tuple[Any, ...]]:
    width = len(headers)
    normalized: list[tuple[Any, ...]] = []
    for row_number, row in enumerate(rows, start=2):
        if len(row) != width:
            raise LockRefused(
                f"ragged row in {path}: row {row_number} has {len(row)} cells, expected {width}"
            )
        normalized.append(tuple(row))
    return normalized


def _refuse_unsafe_xlsx_numbers(path: Path, headers: list[str], rows: list[tuple[Any, ...]]) -> None:
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        return
    decimal_indexes = [index for index, header in enumerate(headers) if _is_decimal_column(header)]
    for row_number, row in enumerate(rows, start=2):
        for index in decimal_indexes:
            value = row[index]
            if isinstance(value, bool) or not isinstance(value, (int, float)):
                continue
            if float(value).is_integer() and abs(int(value)) > MAX_SAFE_XLSX_INTEGER:
                raise LockRefused(
                    f"unsafe xlsx numeric precision in {path}: row {row_number} column {headers[index]}"
                )


def _load_json(path: Path) -> tuple[list[str], list[tuple[Any, ...]]]:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise LockRefused(f"invalid JSON table: {path}") from exc
    if isinstance(payload, dict) and "headers" in payload and "rows" in payload:
        if not isinstance(payload["headers"], list) or not isinstance(payload["rows"], list):
            raise LockRefused(f"invalid JSON table: headers and rows must be lists in {path}")
        raw_headers = [str(item) for item in payload["headers"]]
        if not raw_headers:
            raise LockRefused(f"invalid JSON table: no columns in {path}")
        headers = _normalize_headers(raw_headers, path)
        rows = []
        for row in payload["rows"]:
            if isinstance(row, dict):
                rows.append(tuple(row.get(raw, "") for raw in raw_headers))
            elif isinstance(row, (list, tuple)):
                rows.append(tuple(row))
            else:
                raise LockRefused(f"invalid JSON table: row must be an object or list in {path}")
        return headers, _normalize_rows(path, headers, rows)
    if isinstance(payload, list) and all(isinstance(item, dict) for item in payload):
        raw_headers: list[str] = []
        headers: list[str] = []
        for item in payload:
            for key in item:
                raw = str(key)
                normalized = _normalize(raw)
                if normalized in headers:
                    prior = raw_headers[headers.index(normalized)]
                    if prior != raw:
                        raise LockRefused(f"{path} column identifier collision: {prior} and {raw}")
                    continue
                raw_headers.append(raw)
                headers.append(normalized)
        if not headers:
            raise LockRefused(f"invalid JSON table: no columns in {path}")
        rows = [tuple(item.get(raw, "") for raw in raw_headers) for item in payload]
        return headers, _normalize_rows(path, headers, rows)
    raise LockRefused(f"unsupported JSON table shape: {path}")


def load_tabular(path: Path) -> tuple[list[str], list[tuple[Any, ...]]]:
    try:
        return _load_tabular(path)
    except LockRefused:
        raise
    except Exception as exc:
        raise LockRefused(f"failed to read table {path}: {exc}") from exc


def _load_tabular(path: Path) -> tuple[list[str], list[tuple[Any, ...]]]:
    if not path.is_file():
        raise LockRefused(f"named file is unreadable: {path}")
    suffix = path.suffix.lower()
    if suffix == ".csv":
        text = path.read_text(encoding="utf-8-sig")
        reader = csv.reader(io.StringIO(text), strict=True)
        try:
            raw_headers = next(reader)
        except StopIteration as exc:
            raise LockRefused(f"empty named file: {path}") from exc
        headers = _normalize_headers(raw_headers, path)
        rows = [tuple(row) for row in reader]
        return headers, _normalize_rows(path, headers, rows)
    if suffix == ".json":
        return _load_json(path)
    if suffix in {".xlsx", ".xlsm"}:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.active
            if ws is None:
                raise LockRefused(f"empty named file: {path}")
            iterator = ws.iter_rows(values_only=True)
            try:
                raw_headers = next(iterator)
            except StopIteration as exc:
                raise LockRefused(f"empty named file: {path}") from exc
            headers = _normalize_headers(list(raw_headers), path)
            rows = [tuple(cell if cell is not None else "" for cell in row) for row in iterator]
        finally:
            wb.close()
        _refuse_unsafe_xlsx_numbers(path, headers, rows)
        return headers, _normalize_rows(path, headers, rows)
    raise LockRefused(f"unsupported import format: {path.suffix}")


def require_columns(headers: list[str], required: list[str], label: str) -> None:
    have = set(headers)
    missing = [col for col in required if col not in have]
    if missing:
        raise LockRefused(f"{label} missing columns {missing}")


def quote_ident(name: str) -> str:
    return f'"{sql_identifier(name)}"'


def _decimal_from(value: Any) -> Decimal | None:
    if isinstance(value, bool) or value in (None, ""):
        return None
    if isinstance(value, Decimal):
        return value if value.is_finite() else None
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        converted = Decimal(str(value))
        return converted if converted.is_finite() else None
    if isinstance(value, str):
        try:
            text = value.strip().replace(",", "")
            if text.startswith("(") and text.endswith(")"):
                text = "-" + text[1:-1]
            converted = Decimal(text)
            return converted if converted.is_finite() else None
        except InvalidOperation:
            return None
    return None


def _is_decimal_column(header: str) -> bool:
    normalized = _normalize(header)
    return normalized in {"amount", "debit", "credit", "balance"} or normalized.endswith(
        ("_amount", "_debit", "_credit", "_balance")
    )


def _decimal_precision(value: Decimal) -> tuple[int, int]:
    _sign, digits, exponent = value.as_tuple()
    if value.is_zero():
        return 1, 0
    scale = abs(exponent) if exponent < 0 else 0
    integer_digits = len(digits) + exponent if exponent > 0 else len(digits) - scale
    if integer_digits < 1:
        integer_digits = 1
    return integer_digits, scale


def _column_spec(headers: list[str], rows: list[tuple[Any, ...]], index: int) -> tuple[str, list[Any]]:
    values = [row[index] for row in rows if row[index] not in (None, "")]
    if values and all(isinstance(v, bool) for v in values):
        return "BOOLEAN", [bool(row[index]) if row[index] not in (None, "") else None for row in rows]
    decimals = [_decimal_from(v) for v in values]
    if values and _is_decimal_column(headers[index]) and not all(item is not None for item in decimals):
        raise LockRefused(f"column {headers[index]} contains a non-numeric amount")
    if values and _is_decimal_column(headers[index]):
        converted = [_decimal_from(row[index]) for row in rows]
        precisions = [_decimal_precision(item) for item in converted if item is not None]
        integer_digits = max((item[0] for item in precisions), default=1)
        scale = max((item[1] for item in precisions), default=0)
        width = integer_digits + scale
        if width > 38:
            raise LockRefused(f"column {headers[index]} exceeds DECIMAL(38) precision")
        return f"DECIMAL({max(1, width)},{scale})", converted
    return "VARCHAR", [None if row[index] is None else str(row[index]) for row in rows]


def _assert_no_identifier_collisions(label: str, names: list[str]) -> None:
    seen: dict[str, str] = {}
    for name in names:
        ident = sql_identifier(name)
        prior = seen.get(ident)
        if prior is not None:
            raise LockRefused(f"{label} identifier collision: {prior} and {name}")
        seen[ident] = name


def write_table(con: duckdb.DuckDBPyConnection, table: str, headers: list[str], rows: list[tuple[Any, ...]]) -> None:
    _assert_no_identifier_collisions(f"{table} column", headers)
    rows = _normalize_rows(Path(table), headers, rows)
    if rows:
        specs = [_column_spec(headers, rows, i) for i in range(len(headers))]
        cols = ", ".join(f"{quote_ident(h)} {sql_type}" for h, (sql_type, _values) in zip(headers, specs))
    else:
        cols = ", ".join(f"{quote_ident(h)} VARCHAR" for h in headers)
    con.execute(f"CREATE OR REPLACE TABLE {quote_ident(table)} ({cols})")
    if not rows:
        return
    placeholders = ", ".join(["?"] * len(headers))
    converted_rows = [tuple(values[row_index] for _sql_type, values in specs) for row_index in range(len(rows))]
    con.executemany(f"INSERT INTO {quote_ident(table)} VALUES ({placeholders})", converted_rows)


def checkpoint_and_close(con: duckdb.DuckDBPyConnection) -> None:
    con.execute("CHECKPOINT")
    con.close()
