"""看板快照: the manager's accepted cards as one page (PDF, PPT) or a workbook, filed in 成果库.

The desktop lays the page out from what the board shows; this module decides what may go on it and files what
comes back. A snapshot names the board version it was taken from and, for each card, the run that card reads.
Before anything is filed the snapshot is held to that board:

- only accepted cards, each still on the run it was accepted on (a 待认可 card is not a figure anyone signed,
  and a 回款情景 card is a what-if, never accepted, never in a snapshot);
- every figure is the very cell its run holds, every bar a row of its run in the card's order, every step of a
  profit bridge the variance line it stands for, and what the page prints of each is that number, rounded to the
  places it shows;
- every accepted card is either on the page or named as left off it, and the cards not taken are the board's.

A page laid out from a board that has changed since is refused, and the manager exports again.

Formats:

- ``pdf``: the desktop prints the page (A4 landscape) and hands the bytes here to be filed.
- ``pptx``: the desktop renders the same page as a PNG; it fills one A4 slide, its figures and sources in the
  picture's description. The page is a picture on purpose: a figure retyped in a deck no longer matches its source.
- ``xlsx``: written here, from the same snapshot, every figure as the exact number its run holds (text where a
  spreadsheet cell would round it), with a sheet of where each one came from.
"""

from __future__ import annotations

import hashlib
import io
import re
import struct
import uuid
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape, quoteattr

from foundry.errors import FoundryError

FORMATS = ("pdf", "pptx", "xlsx")
MAX_FILE_BYTES = 32 * 1024 * 1024
MAX_TEXT = 200
MAX_LINE = 600
# A run's decimals can run long (DuckDB keeps every digit of a division); the cell is held as the run holds it.
MAX_CELL = 1000
MAX_CARDS = 40
MAX_ITEMS = 50
# The screening lights a collection-health figure carries (foundry/collection_metrics.py).
LEVELS = {"red", "orange", "yellow", "green", "unknown"}
CONTROL = re.compile(r"[\u0000-\u0008\u000b\u000c\u000e-\u001f\u007f]")
ID = re.compile(r"^[0-9a-f]{32}$")
# A4 landscape, in EMU (36,000 to the millimetre): the slide is the page.
SLIDE = (10_692_000, 7_560_000)


# --- the snapshot as the desktop laid it out ----------------------------------------------------


def _object(value: Any, label: str) -> dict:
    if not isinstance(value, dict):
        raise FoundryError(f"dashboard snapshot {label} must be an object")
    return value


def _list(value: Any, label: str, maximum: int) -> list:
    if not isinstance(value, list) or len(value) > maximum:
        raise FoundryError(f"dashboard snapshot {label} must be a list of at most {maximum}")
    return value


def _text(value: Any, label: str, maximum: int = MAX_TEXT, *, blank: bool = True) -> str:
    if not isinstance(value, str) or CONTROL.search(value) or len(value) > maximum or (not blank and not value.strip()):
        raise FoundryError(f"dashboard snapshot {label} must be text of at most {maximum} characters")
    return value


def _number(value: Any, label: str) -> float:
    # The CLI reads JSON decimals as Decimal (analysis_data.load_json); a direct call hands floats.
    number = float(value) if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool) else None
    if number is None or number != number or number in (float("inf"), float("-inf")):
        raise FoundryError(f"dashboard snapshot {label} must be a finite number")
    return number


def _id(value: Any, label: str) -> str:
    if not isinstance(value, str) or not ID.match(value):
        raise FoundryError(f"dashboard snapshot {label} must be an id")
    return value


def _shown(value: Any, label: str) -> dict:
    shown = _object(value, label)
    if not isinstance(shown.get("negative"), bool):
        raise FoundryError(f"dashboard snapshot {label} must say whether it is negative")
    return {"text": _text(shown.get("text"), label), "negative": shown["negative"]}


def _named(value: Any, label: str) -> list[dict]:
    return [{"card_id": _id(_object(item, label).get("card_id"), f"{label} card"), "title": _text(item.get("title"), f"{label} title")}
            for item in _list(value, label, MAX_CARDS)]


def normalize(raw: Any) -> dict:
    """The snapshot's shape, every text bounded and every figure finite. What it says is checked against the board in
    ``check``."""
    model = _object(raw, "")
    if model.get("version") != 1:
        raise FoundryError("dashboard snapshot version must be 1")
    figures = []
    for item in _list(model.get("figures"), "figures", MAX_CARDS):
        item = _object(item, "figure")
        cell = _object(item.get("cell"), "figure cell")
        level = item.get("level")
        if level is not None and level not in LEVELS:
            raise FoundryError("dashboard snapshot figure level is not one the board shows")
        figures.append({
            "card_id": _id(item.get("card_id"), "figure card"), "title": _text(item.get("title"), "figure title", blank=False),
            "subject": _text(item.get("subject"), "figure subject"), "part": _text(item.get("part"), "figure part"),
            "value": _shown(item.get("value"), "figure value"), "level": level,
            "level_label": _text(item.get("level_label"), "figure level label"),
            "beside": [{"label": _text(_object(pair, "beside").get("label"), "beside label"), "text": _text(pair.get("text"), "beside text")}
                       for pair in _list(item.get("beside"), "beside", 6)],
            "note": _text(item.get("note"), "figure note"), "source": int(_number(item.get("source"), "figure source")),
            "cell": {"column": _text(cell.get("column"), "figure cell column", blank=False), "value": _text(cell.get("value"), "figure cell value", MAX_CELL, blank=False)},
        })
    charts = []
    for item in _list(model.get("charts"), "charts", MAX_CARDS):
        item = _object(item, "chart")
        common = {"card_id": _id(item.get("card_id"), "chart card"), "title": _text(item.get("title"), "chart title", blank=False),
                  "unit": _text(item.get("unit"), "chart unit"), "note": _text(item.get("note"), "chart note"),
                  "source": int(_number(item.get("source"), "chart source"))}
        if item.get("kind") == "rank":
            reference = item.get("reference")
            charts.append({**common, "kind": "rank",
                           "items": [{"label": _text(_object(row, "bar").get("label"), "bar label", blank=False), "value": _number(row.get("value"), "bar value"),
                                      "shown": _shown(row.get("shown"), "bar figure"), "note": _text(row.get("note"), "bar note")}
                                     for row in _list(item.get("items"), "bars", MAX_ITEMS)],
                           "reference": None if reference is None else {"value": _number(_object(reference, "reference").get("value"), "reference value"),
                                                                         "label": _text(reference.get("label"), "reference label")}})
        elif item.get("kind") == "bridge":
            steps = []
            for step in _list(item.get("steps"), "bridge steps", 8):
                step = _object(step, "bridge step")
                if step.get("kind") not in {"total", "delta"}:
                    raise FoundryError("dashboard snapshot bridge step is a total or a delta")
                steps.append({"label": _text(step.get("label"), "bridge step label", blank=False), "kind": step["kind"],
                              "value": _number(step.get("value"), "bridge step value"), "shown": _shown(step.get("shown"), "bridge step figure")})
            charts.append({**common, "kind": "bridge", "steps": steps})
        else:
            raise FoundryError("dashboard snapshot chart is a rank or a bridge")
    sources = [{"n": int(_number(_object(item, "source").get("n"), "source number")), "card_id": _id(item.get("card_id"), "source card"),
                "title": _text(item.get("title"), "source title"), "run": _text(item.get("run"), "source run", 8, blank=False),
                "lines": [_text(line, "source line", MAX_LINE) for line in _list(item.get("lines"), "source lines", 12)]}
               for item in _list(model.get("sources"), "sources", MAX_CARDS)]
    runs = [{"key": _text(_object(item, "run").get("key"), "run key", 8, blank=False), "run_id": _id(item.get("run_id"), "run id"),
             "label": _text(item.get("label"), "run label"), "computed": _text(item.get("computed"), "run computed"),
             "files": [_text(name, "run file", MAX_LINE) for name in _list(item.get("files"), "run files", 20)]}
            for item in _list(model.get("runs"), "runs", MAX_CARDS)]
    left = _object(model.get("left_out"), "left_out")
    return {
        "version": 1, "board_id": _id(model.get("board_id"), "board"),
        "task": _text(model.get("task"), "task"), "period": _text(model.get("period"), "period"),
        "figures": figures, "charts": charts, "sources": sources, "runs": runs,
        "beyond": _named(model.get("beyond"), "beyond"),
        "left_out": {"pinned": _named(left.get("pinned"), "pinned"), "scenario": _named(left.get("scenario"), "scenario")},
        # The 署名 it was exported under (设置 › 使用者): a name on the copy, not a permission.
        **({"signed_by": _text(model["signed_by"], "signed by", 80, blank=False)} if model.get("signed_by") else {}),
    }


# --- held to the board ---------------------------------------------------------------------------


def _same(stored: Any, shown: str) -> bool:
    """Whether the cell a run holds is the text the snapshot shows for it: the same figure, however written."""
    if stored is None or isinstance(stored, bool):
        return False
    text = stored if isinstance(stored, str) else repr(stored) if isinstance(stored, float) else str(stored)
    if text.strip() == shown.strip():
        return True
    try:
        return Decimal(text.strip()) == Decimal(shown.strip())
    except InvalidOperation:
        return False


# A figure as the page writes it (app/src/lib/figures.ts): accounting parentheses or a sign, a currency mark, grouped
# digits, a %, then words without a digit (万元, 个百分点, （估算）).
WRITTEN = re.compile(r"^([(（])?\s*([-+−])?\s*(?:¥|￥|\$)?\s*(\d[\d,]*(?:\.\d+)?)\s*(%)?\s*[)）]?\s*(\D*)$")


def _decimal(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        number = Decimal(value.strip().replace(",", "") if isinstance(value, str) else repr(value) if isinstance(value, float) else str(value))
    except InvalidOperation:
        return None
    return number if number.is_finite() else None


def _written(text: str) -> tuple[Decimal, int, bool] | None:
    """The figure a text writes: its number, the places it shows, and whether it reads as a percentage."""
    found = WRITTEN.match(text.strip())
    if found is None:
        return None
    paren, sign, digits, mark, words = found.groups()
    number = Decimal(digits.replace(",", ""))
    return (-number if paren or sign in ("-", "−") else number), len(digits.partition(".")[2]), bool(mark) or "%" in words or "百分点" in words


def _reads(stored: Any, text: str) -> bool:
    """Whether a figure the page writes is the run's own: its number rounded to the places the page shows, as it is or
    ×100 where the page writes a percentage of a fraction (0.8732 as 87.32%). Text that writes no figure states no
    amount; a cell that is no number is shown as it is."""
    written = _written(text)
    if written is None:
        return True
    number = _decimal(stored)
    if number is None:
        return stored is not None and not isinstance(stored, bool) and str(stored).strip() == text.strip()
    figure, places, percent = written
    half = Decimal(5).scaleb(-places - 1)
    return any(abs(number * scale - figure) <= half for scale in ((1, 100) if percent else (1,)))


def _health(run: dict) -> bool:
    result = run.get("result") or {}
    return run.get("engine") == "finance" and result.get("scenario_kind") == "collection_health" and isinstance(result.get("rows"), list)


def _says(row: dict, text: str) -> bool:
    """Whether the page writes a collection-health line's figure as the run itself wrote it (display_value)."""
    written, own = _written(text), _written(str(row.get("display_value") or ""))
    return written is None or (own is not None and written[0] == own[0])


def _float(value: Any) -> float | None:
    if value is None or isinstance(value, bool) or (isinstance(value, str) and not value.strip()):
        return None
    try:
        number = float(value.strip() if isinstance(value, str) else value)
    except (TypeError, ValueError):
        return None
    return number if number == number and number not in (float("inf"), float("-inf")) else None


def accepted(card: dict) -> bool:
    """A card the manager accepted on the run it reads now: the only kind a snapshot shows."""
    from foundry.dashboards import WHAT_IF

    return (card.get("state") == "accepted" and card.get("kind") not in WHAT_IF
            and (card.get("accepted") or {}).get("run_id") == card.get("run_id"))


def _rows(run: dict, match: dict) -> list[dict]:
    from foundry.dashboards import matching_rows

    result = run.get("result") or {}
    return [row for _, row in matching_rows(result.get("rows") if isinstance(result.get("rows"), list) else [], match)]


def _check_figure(card: dict, run: dict, shown: dict) -> None:
    rows = _rows(run, card["select"].get("match") or {})
    column = card["select"].get("column")
    if len(rows) != 1 or shown["cell"]["column"] != column or not _same(rows[0].get(column), shown["cell"]["value"]):
        raise FoundryError(f"「{card['title']}」 is not the figure its run holds; export again")
    # What the page prints is held too: the figure, its light, and the figures beside it.
    row = rows[0]
    if _health(run):
        # The same measure's other lines, as the board lists them beside it (dashboardViews.figureView), at most six.
        beside = [(str(item.get("scenario")), str(item.get("display_value"))) for item in run["result"]["rows"]
                  if item is not row and item.get("entity") == row.get("entity") and item.get("metric_key") == row.get("metric_key")][:6]
        honest = (_says(row, shown["value"]["text"]) and shown["level"] == row.get("level")
                  and shown["level_label"] == (row.get("level_label") or "")
                  and [(pair["label"], pair["text"]) for pair in shown["beside"]] == beside)
    else:
        honest = (_reads(row.get(column), shown["value"]["text"]) and shown["level"] is None and not shown["level_label"]
                  and all(any(_reads(value, pair["text"]) for value in row.values()) for pair in shown["beside"]))
    if not honest:
        raise FoundryError(f"「{card['title']}」 does not print the figure its run holds; export again")


def ranked(card: dict, run: dict) -> list[tuple[dict, float]]:
    """A rank card's rows as the board draws them: labelled, a finite figure, in the card's order, cut at its limit."""
    select = card["select"]
    label, column = select.get("label"), select.get("column")
    rows = [(row, _float(row.get(column))) for row in _rows(run, select.get("match") or {})
            if row.get(label) not in (None, "")]
    rows = [(row, value) for row, value in rows if value is not None]
    # Python's reverse sort keeps equal rows in their order, as the board's does.
    rows.sort(key=lambda pair: pair[1], reverse=select.get("order", "desc") == "desc")
    return rows[: select.get("limit", 10)]


def _check_rank(card: dict, run: dict, shown: dict) -> None:
    label, column = card["select"].get("label"), card["select"].get("column")
    rows = ranked(card, run)
    if [(item["label"], item["value"]) for item in shown["items"]] != [(str(row.get(label)), value) for row, value in rows]:
        raise FoundryError(f"「{card['title']}」 is not the ranking its run holds; export again")
    says = (lambda row, text: _says(row, text)) if _health(run) else (lambda row, text: _reads(row.get(column), text))
    if not all(says(row, item["shown"]["text"]) for (row, _), item in zip(rows, shown["items"])):
        raise FoundryError(f"「{card['title']}」 does not print the figures its run holds; export again")
    reference = card["select"].get("reference")
    if (reference is None) != (shown["reference"] is None) or (
            reference is not None and (_float(reference.get("value")) != shown["reference"]["value"] or reference.get("label") != shown["reference"]["label"])):
        raise FoundryError(f"「{card['title']}」 does not draw its own reference line; export again")


BRIDGE_DELTAS = ("revenue", "cost", "expense")


def bridge_lines(card: dict, run: dict) -> list[tuple[str, Any]] | None:
    """The variance lines a profit bridge stands on, in its order: (line, the run's own figure)."""
    rows = _rows(run, card["select"].get("match") or {})
    profit = [row for row in rows if row.get("metric") == "profit"]
    deltas = [[row for row in rows if row.get("metric") == metric] for metric in BRIDGE_DELTAS]
    if len(profit) != 1 or any(len(found) != 1 for found in deltas):
        return None
    return [("baseline", profit[0].get("baseline")),
            *[(metric, found[0].get("contribution")) for metric, found in zip(BRIDGE_DELTAS, deltas)],
            ("actual", profit[0].get("actual"))]


def _check_bridge(card: dict, run: dict, shown: dict) -> None:
    lines = bridge_lines(card, run)
    if lines is None or [step["value"] for step in shown["steps"]] != [_float(value) for _, value in lines] \
            or [step["kind"] for step in shown["steps"]] != ["total", "delta", "delta", "delta", "total"]:
        raise FoundryError(f"「{card['title']}」 is not the profit bridge its run holds; export again")
    if not all(_reads(value, step["shown"]["text"]) for (_, value), step in zip(lines, shown["steps"])):
        raise FoundryError(f"「{card['title']}」 does not print the profit bridge its run holds; export again")


def check(model: dict, board: dict | None, run_of, *, whole: bool) -> dict[str, dict]:
    """Hold the snapshot to the board it says it was taken from. ``whole``: every accepted card is on it (a workbook);
    otherwise a card not on the page is named as left off. Returns the board's cards by id."""
    if board is None or model["board_id"] != board["id"]:
        raise FoundryError("the dashboard changed since this snapshot was laid out; export again")
    cards = [card for card in board.get("cards", []) if card.get("state") != "explore"]
    by_id = {card["id"]: card for card in cards}
    taken = [item["card_id"] for item in model["figures"]] + [item["card_id"] for item in model["charts"]]
    beyond = [item["card_id"] for item in model["beyond"]]
    wanted = [card["id"] for card in cards if accepted(card)]
    if not wanted:
        raise FoundryError("no card on the dashboard is accepted yet; accept the ones to export first")
    if len(set(taken + beyond)) != len(taken + beyond) or sorted(taken + beyond) != sorted(wanted) or (whole and beyond):
        raise FoundryError("a snapshot shows the dashboard's accepted cards and no others; export again")
    from foundry.dashboards import WHAT_IF

    # What-ifs (情景、倒推、敏感性) are listed together under 情景.
    left = {"pinned": [card for card in cards if card["kind"] not in WHAT_IF and not accepted(card)],
            "scenario": [card for card in cards if card["kind"] in WHAT_IF]}
    for key, shown in model["left_out"].items():
        if [item["card_id"] for item in shown] != [card["id"] for card in left[key]]:
            raise FoundryError("a snapshot names the cards it leaves out as the dashboard has them; export again")
    sources = {item["card_id"]: item for item in model["sources"]}
    runs = {item["key"]: item for item in model["runs"]}
    for kind, items in (("figure", model["figures"]), ("chart", model["charts"])):
        for item in items:
            card = by_id[item["card_id"]]
            if (kind == "figure") != (card["kind"] == "figure") or (kind == "chart" and item["kind"] != card["kind"]):
                raise FoundryError(f"「{card['title']}」 is drawn as another kind of card; export again")
            source = sources.get(card["id"])
            if source is None or source["n"] != item["source"] or runs.get(source["run"], {}).get("run_id") != card["run_id"]:
                raise FoundryError(f"「{card['title']}」 does not name the run it reads; export again")
            run = run_of(card["run_id"])
            if kind == "figure":
                _check_figure(card, run, item)
            elif item["kind"] == "rank":
                _check_rank(card, run, item)
            else:
                _check_bridge(card, run, item)
    return by_id


# --- the files ------------------------------------------------------------------------------------


def _file(spec: dict, magic: bytes, label: str) -> bytes:
    raw = spec.get("file")
    if not isinstance(raw, str) or not raw.strip():
        raise FoundryError(f"a {label} snapshot needs the file the desktop rendered")
    path = Path(raw)
    if not path.is_absolute() or path.is_symlink() or not path.is_file():
        raise FoundryError(f"the rendered {label} is not a regular file")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise FoundryError(f"the rendered {label} is larger than 32 MiB")
    data = path.read_bytes()
    if not data.startswith(magic):
        raise FoundryError(f"the rendered file is not a {label}")
    return data


def _png_size(data: bytes) -> tuple[int, int]:
    if len(data) < 24 or data[12:16] != b"IHDR":
        raise FoundryError("the rendered page is not a PNG")
    width, height = struct.unpack(">II", data[16:24])
    if not (0 < width <= 8000 and 0 < height <= 8000):
        raise FoundryError("the rendered page has an unusable size")
    return width, height


def describe(model: dict) -> str:
    """The page in words: the picture's description in the deck, for a reader that cannot see it."""
    lines = [f"经营看板 · {model['task']}{' · ' + model['period'] if model['period'] else ''}"]
    for item in model["figures"]:
        lines.append(f"{item['title']}：{item['value']['text']}{'（' + item['level_label'] + '）' if item['level_label'] else ''}")
    for item in model["charts"]:
        if item["kind"] == "rank":
            lines.append(f"{item['title']}：" + "，".join(f"{row['label']} {row['shown']['text']}" for row in item["items"]))
        else:
            lines.append(f"{item['title']}：" + "，".join(f"{step['label']} {step['shown']['text']}" for step in item["steps"]))
    return "\n".join(lines)[:4000]


_NS = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
       'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
       'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"')
_HEAD = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_TREE = ('<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr><a:xfrm>'
         '<a:off x="0" y="0"/><a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>')


def _rels(*rows: tuple[str, str, str]) -> str:
    body = "".join(f'<Relationship Id="{rid}" Type="{kind}" Target="{target}"/>' for rid, kind, target in rows)
    return f'{_HEAD}<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{body}</Relationships>'


def _theme() -> str:
    colours = [("dk1", "1C2232"), ("lt1", "FFFFFF"), ("dk2", "1C2232"), ("lt2", "F5F6F8"), ("accent1", "1C2232"),
               ("accent2", "8A6729"), ("accent3", "426786"), ("accent4", "4F7553"), ("accent5", "A35F1F"),
               ("accent6", "94463D"), ("hlink", "426786"), ("folHlink", "5B6176")]
    scheme = "".join(f'<a:{name}><a:srgbClr val="{value}"/></a:{name}>' for name, value in colours)
    fonts = '<a:latin typeface="Arial"/><a:ea typeface=""/><a:cs typeface=""/>'
    fill = '<a:solidFill><a:schemeClr val="phClr"/></a:solidFill>'
    line = f'<a:ln w="6350">{fill}</a:ln>'
    effect = "<a:effectStyle><a:effectLst/></a:effectStyle>"
    return (f'{_HEAD}<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="FA Foundry"><a:themeElements>'
            f'<a:clrScheme name="FA Foundry">{scheme}</a:clrScheme>'
            f'<a:fontScheme name="FA Foundry"><a:majorFont>{fonts}</a:majorFont><a:minorFont>{fonts}</a:minorFont></a:fontScheme>'
            f'<a:fmtScheme name="FA Foundry"><a:fillStyleLst>{fill * 3}</a:fillStyleLst><a:lnStyleLst>{line * 3}</a:lnStyleLst>'
            f'<a:effectStyleLst>{effect * 3}</a:effectStyleLst><a:bgFillStyleLst>{fill * 3}</a:bgFillStyleLst></a:fmtScheme>'
            '</a:themeElements><a:objectDefaults/><a:extraClrSchemeLst/></a:theme>')


def one_slide(png: bytes, *, title: str, description: str) -> bytes:
    """One A4 landscape slide filled by the page, written with the standard library: no deck template, no macros."""
    width, height = _png_size(png)
    cx, cy = SLIDE
    # The page keeps its own proportions, centred; an A4 page fills the slide.
    if width * cy > height * cx:
        ext = (cx, round(cx * height / width))
    else:
        ext = (round(cy * width / height), cy)
    off = ((cx - ext[0]) // 2, (cy - ext[1]) // 2)
    layout_type = f"{_REL}/slideLayout"
    parts = {
        "[Content_Types].xml": (
            f'{_HEAD}<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/><Default Extension="png" ContentType="image/png"/>'
            '<Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>'
            '<Override PartName="/ppt/slideMasters/slideMaster1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slideMaster+xml"/>'
            '<Override PartName="/ppt/slideLayouts/slideLayout1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slideLayout+xml"/>'
            '<Override PartName="/ppt/slides/slide1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>'
            '<Override PartName="/ppt/theme/theme1.xml" ContentType="application/vnd.openxmlformats-officedocument.theme+xml"/>'
            '<Override PartName="/ppt/presProps.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presProps+xml"/>'
            '<Override PartName="/ppt/viewProps.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.viewProps+xml"/>'
            '<Override PartName="/ppt/tableStyles.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.tableStyles+xml"/>'
            '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>'
            '</Types>'),
        "_rels/.rels": _rels(("rId1", f"{_REL}/officeDocument", "ppt/presentation.xml"),
                             ("rId2", "http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties", "docProps/core.xml")),
        "docProps/core.xml": (
            f'{_HEAD}<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
            'xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>' + escape(title) + '</dc:title>'
            '<dc:creator>FA Foundry</dc:creator></cp:coreProperties>'),
        "ppt/presentation.xml": (
            f'{_HEAD}<p:presentation {_NS} saveSubsetFonts="1"><p:sldMasterIdLst><p:sldMasterId id="2147483648" r:id="rId1"/></p:sldMasterIdLst>'
            f'<p:sldIdLst><p:sldId id="256" r:id="rId2"/></p:sldIdLst><p:sldSz cx="{cx}" cy="{cy}"/><p:notesSz cx="6858000" cy="9144000"/>'
            '</p:presentation>'),
        "ppt/_rels/presentation.xml.rels": _rels(
            ("rId1", f"{_REL}/slideMaster", "slideMasters/slideMaster1.xml"), ("rId2", f"{_REL}/slide", "slides/slide1.xml"),
            ("rId3", f"{_REL}/theme", "theme/theme1.xml"), ("rId4", f"{_REL}/presProps", "presProps.xml"),
            ("rId5", f"{_REL}/viewProps", "viewProps.xml"), ("rId6", f"{_REL}/tableStyles", "tableStyles.xml")),
        "ppt/presProps.xml": f"{_HEAD}<p:presentationPr {_NS}/>",
        "ppt/viewProps.xml": f'{_HEAD}<p:viewPr {_NS} lastView="sldView"/>',
        "ppt/tableStyles.xml": f'{_HEAD}<a:tblStyleLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" def="{{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}}"/>',
        "ppt/theme/theme1.xml": _theme(),
        "ppt/slideMasters/slideMaster1.xml": (
            f'{_HEAD}<p:sldMaster {_NS}><p:cSld><p:bg><p:bgRef idx="1001"><a:schemeClr val="bg1"/></p:bgRef></p:bg><p:spTree>{_TREE}</p:spTree></p:cSld>'
            '<p:clrMap bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2" accent1="accent1" accent2="accent2" accent3="accent3" accent4="accent4" '
            'accent5="accent5" accent6="accent6" hlink="hlink" folHlink="folHlink"/>'
            '<p:sldLayoutIdLst><p:sldLayoutId id="2147483649" r:id="rId1"/></p:sldLayoutIdLst></p:sldMaster>'),
        "ppt/slideMasters/_rels/slideMaster1.xml.rels": _rels(("rId1", layout_type, "../slideLayouts/slideLayout1.xml"),
                                                             ("rId2", f"{_REL}/theme", "../theme/theme1.xml")),
        "ppt/slideLayouts/slideLayout1.xml": (
            f'{_HEAD}<p:sldLayout {_NS} type="blank" preserve="1"><p:cSld name="Blank"><p:spTree>{_TREE}</p:spTree></p:cSld>'
            '<p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sldLayout>'),
        "ppt/slideLayouts/_rels/slideLayout1.xml.rels": _rels(("rId1", f"{_REL}/slideMaster", "../slideMasters/slideMaster1.xml")),
        "ppt/slides/slide1.xml": (
            f'{_HEAD}<p:sld {_NS}><p:cSld><p:spTree>{_TREE}'
            f'<p:pic><p:nvPicPr><p:cNvPr id="2" name={quoteattr(title)} descr={quoteattr(description)}/>'
            '<p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr><p:nvPr/></p:nvPicPr>'
            '<p:blipFill><a:blip r:embed="rId2"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
            f'<p:spPr><a:xfrm><a:off x="{off[0]}" y="{off[1]}"/><a:ext cx="{ext[0]}" cy="{ext[1]}"/></a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>'
            '</p:spTree></p:cSld><p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sld>'),
        "ppt/slides/_rels/slide1.xml.rels": _rels(("rId1", layout_type, "../slideLayouts/slideLayout1.xml"),
                                                 ("rId2", f"{_REL}/image", "../media/image1.png")),
    }
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        for name, text in parts.items():
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, text.encode("utf-8"))
        info = zipfile.ZipInfo("ppt/media/image1.png", date_time=(1980, 1, 1, 0, 0, 0))
        info.compress_type = zipfile.ZIP_STORED
        archive.writestr(info, png)
    return out.getvalue()


def workbook(model: dict, cards: dict[str, dict], run_of) -> bytes:
    """The snapshot as a workbook: each figure the exact number its run holds, each chart its rows, and the sources."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill

    from foundry.sheet import INK, RULE, _widths, put_number

    def _figure(cell, value) -> None:
        # Kept as text it reads left to right, on into the empty columns beside it.
        put_number(cell, value)
        if isinstance(cell.value, str):
            cell.alignment = Alignment(horizontal="left")

    def _head(sheet, labels: list[str]) -> None:
        sheet.append(labels)
        for cell in sheet[sheet.max_row]:
            cell.font, cell.fill, cell.border = Font(bold=True, color=INK), PatternFill("solid", fgColor="F2F2F2"), Border(bottom=RULE)

    book = Workbook()
    sheet = book.active
    sheet.title = "看板快照"
    sheet.append([f"经营看板 · {model['task']}{' · ' + model['period'] if model['period'] else ''}"])
    sheet["A1"].font = Font(bold=True, size=14, color=INK)
    sheet.append(["只放经理认可的卡。「计算里的数」是它读的那次计算里原样的数（位数太多、表格存不下的写成文字），来源见「来源」页。"])
    if model.get("signed_by"):
        sheet.append([f"导出人：{model['signed_by']}"])
    left = model["left_out"]
    if left["pinned"] or left["scenario"]:
        sheet.append(["没有放进来：" + "；".join(filter(None, [
            f"待认可 {len(left['pinned'])} 张（{'、'.join(item['title'] for item in left['pinned'])}）" if left["pinned"] else "",
            f"情景 {len(left['scenario'])} 张（情景不是实际数）" if left["scenario"] else ""]))])
    sheet.append([])
    # The run's own figure goes last: a long one runs on into the empty cells beside it, never over another column.
    if model["figures"]:
        sheet.append(["关键数字"])
        sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True, color=INK)
        _head(sheet, ["卡片", "主体或期间", "取的是", "页面上写作", "级别", "来源", "计算里的数"])
        for item in model["figures"]:
            card = cards[item["card_id"]]
            row = _rows(run_of(card["run_id"]), card["select"].get("match") or {})[0]
            sheet.append([item["title"], item["subject"], item["part"], item["value"]["text"], item["level_label"], item["source"]])
            _figure(sheet.cell(row=sheet.max_row, column=7), row.get(card["select"]["column"]))
        sheet.append([])
    for item in model["charts"]:
        card = cards[item["card_id"]]
        run = run_of(card["run_id"])
        sheet.append([f"{item['title']}{'（单位：' + item['unit'] + '）' if item['unit'] else ''}", None, None, None, None, item["source"]])
        sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True, color=INK)
        if item["kind"] == "rank":
            _head(sheet, ["名次", "主体", "", "页面上写作", "说明", "", "计算里的数"])
            for place, ((row, _), bar) in enumerate(zip(ranked(card, run), item["items"]), start=1):
                sheet.append([place, bar["label"], None, bar["shown"]["text"], bar["note"]])
                sheet.cell(row=sheet.max_row, column=1).alignment = Alignment(horizontal="left")
                _figure(sheet.cell(row=sheet.max_row, column=7), row.get(card["select"]["column"]))
            if item["reference"]:
                sheet.append(["参考线", item["reference"]["label"]])
                _figure(sheet.cell(row=sheet.max_row, column=7), card["select"]["reference"]["value"])
        else:
            _head(sheet, ["步骤", "项目", "", "页面上写作", "", "", "计算里的数"])
            for (line, value), step in zip(bridge_lines(card, run) or [], item["steps"]):
                sheet.append(["起点" if line == "baseline" else "终点" if line == "actual" else "变动", step["label"], None, step["shown"]["text"]])
                _figure(sheet.cell(row=sheet.max_row, column=7), value)
        sheet.append([])
    if model["beyond"]:
        sheet.append(["一页纸上放不下、只在这里的卡：" + "、".join(item["title"] for item in model["beyond"])])
    _widths(sheet, [34, 18, 10, 20, 14, 8, 22])
    for row in sheet.iter_rows(min_row=1, max_row=3):
        row[0].alignment = Alignment(wrap_text=False)
    sources = book.create_sheet("来源")
    _head(sources, ["编号", "卡片", "取数和口径", "计算", "计算编号", "读的资料"])
    runs = {item["key"]: item for item in model["runs"]}
    for item in model["sources"]:
        run = runs[item["run"]]
        sources.append([item["n"], item["title"], "\n".join(item["lines"]), f"{run['label']} · {run['computed']}", run["run_id"], "\n".join(run["files"])])
        for column in (3, 6):
            sources.cell(row=sources.max_row, column=column).alignment = Alignment(wrap_text=True, vertical="top")
    _widths(sources, [6, 30, 60, 36, 34, 48])
    # Printed, each sheet is one page across, landscape like the page it snapshots.
    for page in (sheet, sources):
        page.page_setup.orientation = "landscape"
        page.page_setup.paperSize = page.PAPERSIZE_A4
        page.sheet_properties.pageSetUpPr.fitToPage = True
        page.page_setup.fitToWidth, page.page_setup.fitToHeight = 1, 0
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


# --- filed in 成果库 -------------------------------------------------------------------------------


def _folder(analysis) -> Path:
    return analysis._inside(analysis.root / "finance" / "dashboard_snapshots")


UNSAFE_NAME = re.compile(r'[\\/:*?"<>|]+')
DOT = re.compile(r"\s*·\s*")
SPACE = re.compile(r"\s+")


def _filename(model: dict, fmt: str) -> str:
    """经营看板_示例_华宸物业8月月结_2026年8月.pdf: the page's own words, with nothing a file system refuses."""
    words = [SPACE.sub("", DOT.sub("_", word)) for word in ("经营看板", model["task"], model["period"]) if word]
    stem = UNSAFE_NAME.sub("_", "_".join(words))
    return f"{stem[:120]}.{fmt}"


def save(task, spec: Any, *, actor: str | None) -> dict:
    """File one snapshot of the board in the format asked; the desktop copies the filed bytes where the manager chose."""
    from foundry.analysis import AnalysisStore, _hash, _new
    from foundry.finance_service import FinanceStore
    from foundry.fsutil import write_json_atomic
    from foundry.locking import exclusive

    spec = _object(spec, "spec")
    fmt = spec.get("format")
    if fmt not in FORMATS:
        raise FoundryError("a dashboard snapshot is a pdf, a pptx or an xlsx")
    model = normalize(spec.get("model"))
    store = FinanceStore(task, actor=actor)
    store._manager_only("export the dashboard")
    run_of = store._dashboard_runs()
    cards = check(model, store._dashboard(), run_of, whole=fmt == "xlsx")
    title = f"经营看板 · {model['task']}{' · ' + model['period'] if model['period'] else ''}"
    if fmt == "xlsx":
        data = workbook(model, cards, run_of)
    elif fmt == "pdf":
        data = _file(spec, b"%PDF-", "PDF")
    else:
        data = one_slide(_file(spec, b"\x89PNG\r\n\x1a\n", "PNG"), title=title, description=describe(model))
    analysis = AnalysisStore(task)
    snapshot_id = uuid.uuid4().hex
    relative = Path("finance") / "dashboard_snapshots" / "files" / f"{snapshot_id}.{fmt}"
    target = analysis._inside(analysis.root / relative)
    record_path = analysis._inside(_folder(analysis) / f"{snapshot_id}.json")
    taken = [*model["figures"], *model["charts"]]
    # Filed only while the board is still the version the page was laid out from: a save in the meantime waits on the
    # same lock, and task.lock is taken inside it as dashboard_save takes it.
    with store._dashboard_lock():
        if (store._dashboard() or {}).get("id") != model["board_id"]:
            raise FoundryError("the dashboard changed since this snapshot was laid out; export again")
        with exclusive(task.lock_path, timeout=10):
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(data)
            record = _new(
                task, domain="foundry.dashboard_snapshot.v1", format=fmt, filename=_filename(model, fmt), artifact_path=str(relative),
                sha256=hashlib.sha256(data).hexdigest(), board_id=model["board_id"], title=title, period=model["period"],
                cards=[{"id": item["card_id"], "title": item["title"], "kind": cards[item["card_id"]]["kind"],
                        "run_id": cards[item["card_id"]]["run_id"]} for item in taken],
                beyond=[item["title"] for item in model["beyond"]],
                left_out={key: [item["title"] for item in items] for key, items in model["left_out"].items()},
            )
            record["id"] = snapshot_id
            record["seq"] = analysis._next_seq()
            record["_integrity_hash"] = _hash(record)
            write_json_atomic(record_path, record)
    return {"snapshot": record, "path": str(target.resolve())}


def snapshots(task) -> list[dict]:
    """Every snapshot this task filed, oldest first."""
    from foundry.analysis import AnalysisStore

    analysis = AnalysisStore(task)
    folder = _folder(analysis)
    if not folder.is_dir():
        return []
    records = [analysis._read(path) for path in sorted(folder.glob("*.json"))]
    return sorted(records, key=lambda record: (record.get("seq", 0), record["created_at"]))


def snapshot_path(task, snapshot_id: str) -> Path:
    """A snapshot's file, from its id alone, refused when its bytes no longer match the record."""
    from foundry.analysis import AnalysisStore

    if not isinstance(snapshot_id, str) or not ID.match(snapshot_id):
        raise FoundryError(f"not a dashboard snapshot id: {snapshot_id}")
    analysis = AnalysisStore(task)
    record_path = analysis._inside(_folder(analysis) / f"{snapshot_id}.json")
    if not record_path.is_file():
        raise FoundryError(f"no dashboard snapshot {snapshot_id} on task {task.id}")
    record = analysis._read(record_path)
    target = analysis._inside(analysis.root / str(record.get("artifact_path") or ""))
    if target.suffix.lstrip(".") not in FORMATS or target.suffix.lstrip(".") != record.get("format") or not target.is_file() or target.is_symlink():
        raise FoundryError(f"dashboard snapshot {snapshot_id} is missing from task {task.id}")
    if hashlib.sha256(target.read_bytes()).hexdigest() != record.get("sha256"):
        raise FoundryError("dashboard snapshot integrity check failed; its bytes differ from its record")
    return target.resolve()
