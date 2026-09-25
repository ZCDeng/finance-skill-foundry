"""Deterministic source evidence recipes for finance drafts.

Recipes are manager-visible source checks. They execute only against archived,
task-bound snapshots and produce issues; they never promote a model assertion into
source truth.
"""

from __future__ import annotations

import io
import math
from decimal import Decimal, InvalidOperation
from fractions import Fraction
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

from foundry.analysis import _hash, _hash_bytes
from foundry.analysis_data import archived_numeric_parser, read_exact_numerics
from foundry.errors import FoundryError

RECIPE_KINDS = {"direct", "sum", "diff", "ratio"}
MAX_RECIPE_CELLS = 8000
MAX_RECIPE_COUNT = 256
MAX_CONSISTENCY_SOURCE_COUNT = 256
MAX_CONSISTENCY_RULE_COUNT = 256
MAX_RECIPE_ANCHORS = 4
EXCEL_MAX_ROW = 1048576
EXCEL_MAX_COLUMN = 16384
# What a passing recipe proves about its claim cell. Only the first two are source evidence; an
# estimate needs a factor, and a free-form scale proves only a multiplication (never a source value).
SOURCE_VALUE = "source_value"
SOURCE_CALCULATION = "source_calculation"
ESTIMATE = "estimate"
COEFFICIENT_UNBOUND = "coefficient_unbound"


def evaluate_evidence_recipes(analysis: Any, recipes: list[Any] | None) -> dict[str, Any]:
    clean = _clean_recipes(recipes)
    recipe_by_id = {recipe["id"]: recipe for recipe in clean}
    factor_cache: dict[str, tuple[Decimal, Fraction | None]] = {}
    results = [
        {
            **_evaluate_recipe(analysis, recipe, recipe_by_id, factor_cache),
            **_recipe_identity(recipe, recipe_by_id),
            **_anchor_details(analysis, recipe),
        }
        for recipe in clean
    ]
    for result in results:
        result["evidence_class"] = recipe_evidence_class(recipe_by_id[result["recipe_id"]], result["status"])
    issues = [result["issue"] for result in results if result.get("issue")]
    status = "not_configured" if not clean else "needs_attention" if issues else "pass"
    return {
        "schema_version": 1,
        "summary": {
            "status": status,
            "recipe_count": len(clean),
            "issue_count": len(issues),
        },
        "results": results,
        "issues": issues,
    }


def recipe_evidence_class(recipe: dict[str, Any], status: str) -> str | None:
    """What a recipe result proves about its claim cell; None unless it passed.

    A factor-bound recipe is an estimate. A scale that is neither 1 nor an exact unit power
    (千/万/百万/亿) is a hand-written coefficient that only proves a multiplication. A ratio
    recipe is a source calculation of a ratio; callers must not let it prove an amount.
    """
    if status != "pass":
        return None
    if recipe.get("factor_recipe_id"):
        return ESTIMATE
    if not _is_unit_scale(recipe["scale"]):
        return COEFFICIENT_UNBOUND
    return SOURCE_VALUE if recipe["kind"] == "direct" else SOURCE_CALCULATION


def _is_unit_scale(scale: Decimal) -> bool:
    if scale == 1:
        return True
    if not scale.is_finite() or scale <= 0:
        return False
    sign, digits, exponent = _significand(scale)
    return digits == (1,) and abs(exponent) in UNIT_SCALE_EXPONENTS


def _recipe_identity(recipe: dict[str, Any], recipe_by_id: dict[str, dict[str, Any]]) -> dict[str, Any]:
    """Server-side facts about a recipe that requirement checks read instead of trusting a caller."""
    factor = recipe_by_id.get(recipe.get("factor_recipe_id") or "")
    return {
        "kind": recipe["kind"],
        "claim": dict(recipe["claim"]),
        "scale": _format_decimal(recipe["scale"]),
        "recipe_hash": recipe_hash(recipe),
        "source_refs": _source_refs(recipe),
        **(
            {
                "factor": {
                    "recipe_id": factor["id"],
                    "recipe_hash": recipe_hash(factor),
                    "source_refs": _source_refs(factor),
                    "location": _recipe_location(factor),
                }
            }
            if factor
            else {}
        ),
        "location": _recipe_location(recipe),
    }


def recipe_hash(recipe: dict[str, Any]) -> str:
    """Hash of the cleaned recipe that decides the result; display-only anchors are excluded."""
    return _hash({key: value for key, value in recipe.items() if key != "anchors"})


def _source_refs(recipe: dict[str, Any]) -> list[dict[str, str]]:
    sources = [recipe] if recipe["kind"] in {"direct", "sum"} else [recipe["left"], recipe["right"]]
    return [{key: source[key] for key in ("dataset_id", "source_hash", "sheet", "range")} for source in sources]


def _anchor_details(analysis: Any, recipe: dict[str, Any]) -> dict[str, Any]:
    """Archived label/header text next to the recipe range, for a manager to read; never a pass/fail input."""
    anchors = recipe.get("anchors") or {}
    if not anchors:
        return {}
    source = recipe if recipe["kind"] in {"direct", "sum"} else recipe["left"]
    details = {}
    for name, cell in anchors.items():
        reference = {"dataset_id": source["dataset_id"], "source_hash": source["source_hash"], "sheet": source["sheet"], "range": cell}
        try:
            dataset = _read_authorized_dataset(analysis, reference)
            details[name] = _archived_workbook_text_cell(analysis, dataset, reference)
        except (FoundryError, OSError, ValueError, KeyError) as exc:
            details[name] = {"sheet": source["sheet"], "range": cell, "error": str(exc)}
    return {"anchors": details}


def _clean_anchors(value: Any) -> dict[str, str]:
    if not isinstance(value, dict) or len(value) > MAX_RECIPE_ANCHORS:
        raise FoundryError(f"evidence recipe anchors must be an object with at most {MAX_RECIPE_ANCHORS} cells")
    anchors = {}
    for name, cell in value.items():
        label = _text(name, "evidence recipe anchor name")
        reference = _text(cell, "evidence recipe anchor cell").replace("$", "")
        try:
            min_col, min_row, max_col, max_row = range_boundaries(reference)
        except ValueError as exc:
            raise FoundryError("evidence recipe anchor must be a single A1 cell") from exc
        if None in {min_col, min_row} or min_col != max_col or min_row != max_row:
            raise FoundryError("evidence recipe anchor must be a single A1 cell")
        anchors[label[:40]] = reference
    return anchors


def recipe_dataset_ids(recipes: list[Any] | None) -> set[str]:
    ids: set[str] = set()
    for recipe in _clean_recipes(recipes):
        if recipe["kind"] in {"direct", "sum"}:
            ids.add(recipe["dataset_id"])
        else:
            ids.add(recipe["left"]["dataset_id"])
            ids.add(recipe["right"]["dataset_id"])
    return ids


def evidence_quality_issues(report: dict[str, Any]) -> list[dict[str, Any]]:
    return list(report.get("issues") or [])


def consistency_dataset_ids(sources: Any) -> set[str]:
    ids: set[str] = set()
    for source in _clean_consistency_sources(sources).values():
        ids.add(source["dataset_id"])
    return ids


def prepare_consistency_inputs(analysis: Any, rules: list[Any] | None, sources: Any) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    clean_sources = _clean_consistency_sources(sources)
    values: dict[str, Decimal] = {}
    source_auth: dict[str, Any] = {}
    for source_id, source in clean_sources.items():
        dataset = _read_authorized_dataset(analysis, source)
        reference = {key: source[key] for key in ("dataset_id", "source_hash")}
        if source["kind"] == "number":
            reference.update({key: source[key] for key in ("sheet", "range")})
            values[source_id] = _source_value(analysis, source) * source["scale"]
            source_auth[source_id] = {"kind": "number", "reference": reference}
        elif source["kind"] == "text":
            paragraph = _archived_paragraph(dataset, source, analysis)
            source_auth[source_id] = {
                "kind": "text",
                "text": paragraph["text"],
                "reference": {**reference, **{key: paragraph[key] for key in paragraph if key != "text"}},
            }
        else:  # pragma: no cover - guarded by cleaner
            raise FoundryError("unsupported consistency source kind")
    return _resolve_consistency_rules(rules, values), source_auth


def _clean_recipes(recipes: list[Any] | None) -> list[dict[str, Any]]:
    if recipes in (None, []):
        return []
    if not isinstance(recipes, list):
        raise FoundryError("evidence_recipes must be an array")
    if len(recipes) > MAX_RECIPE_COUNT:
        raise FoundryError("evidence_recipes exceeds the recipe count limit")
    cleaned = [_clean_recipe(recipe, index) for index, recipe in enumerate(recipes, start=1)]
    ids = [recipe["id"] for recipe in cleaned]
    if len(ids) != len(set(ids)):
        raise FoundryError("evidence recipe ids must be unique")
    by_id = {recipe["id"]: recipe for recipe in cleaned}
    for recipe in cleaned:
        factor_id = recipe.get("factor_recipe_id")
        if not factor_id:
            continue
        factor = by_id.get(factor_id)
        if factor is None:
            raise FoundryError("factor_recipe_id must reference a same-batch evidence recipe")
        if factor["kind"] != "ratio":
            raise FoundryError("factor_recipe_id must reference a ratio evidence recipe")
        if factor["scale"] != Decimal(1):
            raise FoundryError("factor ratio recipe scale must be 1")
        if recipe["id"] == factor_id:
            raise FoundryError("factor_recipe_id cannot reference itself")
        _assert_matching_factor_business_key(recipe, factor)
    return cleaned


def _clean_recipe(recipe: Any, index: int) -> dict[str, Any]:
    if not isinstance(recipe, dict):
        raise FoundryError("evidence recipe must be an object")
    kind = recipe.get("kind")
    if kind not in RECIPE_KINDS:
        raise FoundryError("evidence recipe kind must be direct, sum, diff or ratio")
    recipe_id = _text(recipe.get("id") or f"recipe-{index}", "evidence recipe id")
    cleaned = {
        "schema_version": 1,
        "id": recipe_id,
        "label": _text(recipe.get("label") or recipe_id, "evidence recipe label"),
        "kind": kind,
        "unit": _text(recipe.get("unit", ""), "evidence recipe unit", required=False),
        "scale": _decimal(recipe.get("scale", "1"), "evidence recipe scale"),
        "claim": _claim(recipe.get("claim")),
        "business_key": recipe.get("business_key") if isinstance(recipe.get("business_key"), dict) else {},
    }
    if recipe.get("anchors") is not None:
        cleaned["anchors"] = _clean_anchors(recipe.get("anchors"))
    if kind in {"direct", "sum"}:
        cleaned.update(_source(recipe))
        if recipe.get("factor_recipe_id") is not None:
            # A unit conversion (元 → 万元) may sit next to the factor; a coefficient may not, it belongs in
            # the referenced ratio recipe, where it is recomputed from the archive.
            if "scale" in recipe and not _is_unit_scale(cleaned["scale"]):
                raise FoundryError(
                    "factor_recipe_id cannot be combined with an explicit scale other than 1 or a unit power "
                    "(千/万/百万/亿); put the coefficient in the referenced ratio recipe"
                )
            cleaned["factor_recipe_id"] = _text(recipe.get("factor_recipe_id"), "evidence recipe factor_recipe_id")
    elif recipe.get("factor_recipe_id") is not None:
        raise FoundryError("factor_recipe_id is only supported on direct and sum evidence recipes")
    else:
        left = recipe.get("left")
        right = recipe.get("right")
        if not isinstance(left, dict) or not isinstance(right, dict):
            raise FoundryError("diff and ratio recipes require left and right sources")
        cleaned["left"] = _source(left)
        cleaned["right"] = _source(right)
    return cleaned


def _source(recipe: dict[str, Any]) -> dict[str, str]:
    dataset_id = _text(recipe.get("dataset_id"), "evidence recipe dataset_id")
    source_hash = _text(recipe.get("source_hash"), "evidence recipe source_hash")
    sheet = _text(recipe.get("sheet"), "evidence recipe sheet")
    range_ref = _text(recipe.get("range"), "evidence recipe range")
    try:
        min_col, min_row, max_col, max_row = range_boundaries(range_ref.replace("$", ""))
    except ValueError as exc:
        raise FoundryError("evidence recipe range must be a bounded A1 reference") from exc
    if None in {min_col, min_row, max_col, max_row}:
        raise FoundryError("evidence recipe range must be a bounded A1 reference")
    if min_row < 1 or min_col < 1 or max_row < min_row or max_col < min_col:
        raise FoundryError("evidence recipe range must be a bounded A1 reference")
    if max_row > EXCEL_MAX_ROW or max_col > EXCEL_MAX_COLUMN:
        raise FoundryError("evidence recipe range exceeds Excel worksheet bounds")
    height = max_row - min_row + 1
    width = max_col - min_col + 1
    if height > 200 or width > 40 or height * width > MAX_RECIPE_CELLS:
        raise FoundryError("evidence recipe range is too large")
    return {"dataset_id": dataset_id, "source_hash": source_hash, "sheet": sheet, "range": range_ref.replace("$", "")}


def _evaluate_recipe(
    analysis: Any,
    recipe: dict[str, Any],
    recipe_by_id: dict[str, dict[str, Any]] | None = None,
    factor_cache: dict[str, tuple[Decimal, Fraction | None]] | None = None,
) -> dict[str, Any]:
    factor_value: Decimal | None = None
    try:
        if recipe.get("factor_recipe_id"):
            _assert_claim_business_key(analysis, recipe)
        # `actual` is the server's Decimal rendering (shown, and read for hints); `exact` decides.
        actual, exact = _recipe_actual(analysis, recipe)
        if recipe.get("factor_recipe_id"):
            if recipe_by_id is None or factor_cache is None:  # pragma: no cover - compatibility guard
                raise FoundryError("factor recipe context is unavailable")
            factor_value, factor_exact = _factor_value(analysis, recipe, recipe_by_id, factor_cache)
            actual = actual * factor_value
            exact = _exact_product(exact, factor_exact)
        actual = actual * recipe["scale"]
        exact = _exact_product(exact, _exact(recipe["scale"]))
        expected = _claim_expected(analysis, recipe["claim"])
    except FoundryError as exc:
        # A factor that failed only on the precision of its run value makes the dependent fail for the
        # same reason; say so, so the fix (full precision) is named where the dependent is reported.
        rounded = bool(getattr(exc, "precision_hint", False))
        return {
            "recipe_id": recipe["id"],
            "status": "unreadable",
            **({"precision_hint": FULL_PRECISION_REQUIRED} if rounded else {}),
            "issue": _issue(recipe, "evidence_recipe_unreadable", str(exc), actual=None, expected=None,
                            precision_hint=rounded),
        }
    if not _matches_exact(exact, expected, actual):
        rounded = _rounding_hint(actual, expected)
        return {
            "recipe_id": recipe["id"],
            "status": "mismatch",
            "actual": _format_decimal(actual),
            "expected": _format_decimal(expected),
            **({"precision_hint": FULL_PRECISION_REQUIRED} if rounded else {}),
            **_factor_result_details(recipe, factor_value),
            "issue": _issue(
                recipe,
                "evidence_recipe_mismatch",
                None,
                actual=actual,
                expected=expected,
                scale_hint=_accepts_scale(recipe, recipe_by_id),
                precision_hint=rounded,
            ),
        }
    return {
        "recipe_id": recipe["id"],
        "status": "pass",
        "actual": _format_decimal(actual),
        "expected": _format_decimal(expected),
        **_factor_result_details(recipe, factor_value),
    }


def _recipe_actual(analysis: Any, recipe: dict[str, Any]) -> tuple[Decimal, Fraction | None]:
    """The recipe's value as the server's Decimal (default context) and as an exact rational (see _exact)."""
    if recipe["kind"] == "direct":
        value = _single_value(analysis, recipe)
        return value, _exact(value)
    if recipe["kind"] == "sum":
        value = _range_sum(analysis, recipe)
        return value, _exact(value)
    if recipe["kind"] == "diff":
        left, right = _source_value(analysis, recipe["left"]), _source_value(analysis, recipe["right"])
        exact_left, exact_right = _exact(left), _exact(right)
        return left - right, (exact_left - exact_right if exact_left is not None and exact_right is not None else None)
    if recipe["kind"] == "ratio":
        return _ratio_value(analysis, recipe)
    raise FoundryError("unsupported evidence recipe kind")  # pragma: no cover - guarded by clean step


def _ratio_value(analysis: Any, recipe: dict[str, Any]) -> tuple[Decimal, Fraction | None]:
    denominator = _source_value(analysis, recipe["right"])
    if denominator == 0:
        raise FoundryError("ratio evidence recipe denominator is zero")
    numerator = _source_value(analysis, recipe["left"])
    exact_numerator, exact_denominator = _exact(numerator), _exact(denominator)
    exact = (exact_numerator / exact_denominator
             if exact_numerator is not None and exact_denominator is not None else None)
    return numerator / denominator, exact


def _factor_value(
    analysis: Any,
    dependent: dict[str, Any],
    recipe_by_id: dict[str, dict[str, Any]],
    factor_cache: dict[str, tuple[Decimal, Fraction | None]],
) -> tuple[Decimal, Fraction | None]:
    """The factor's Decimal rendering and its exact quotient. The dependent is computed on the exact one, so
    how the run rounded its own copy of the factor never decides the estimate."""
    factor_id = dependent["factor_recipe_id"]
    if factor_id in factor_cache:
        return factor_cache[factor_id]
    factor = recipe_by_id.get(factor_id)
    if factor is None or factor["kind"] != "ratio":  # pragma: no cover - guarded by clean step
        raise FoundryError("factor_recipe_id must reference a ratio evidence recipe")
    _assert_claim_business_key(analysis, factor)
    actual, exact = _ratio_value(analysis, factor)
    expected = _claim_expected(analysis, factor["claim"])
    if not _matches_exact(exact, expected, actual):
        error = FoundryError("系数规则未通过归档来源的精确复算")
        error.precision_hint = _rounding_hint(actual, expected)
        raise error
    factor_cache[factor_id] = (actual, exact)
    return actual, exact


def _factor_result_details(recipe: dict[str, Any], factor_value: Decimal | None) -> dict[str, str]:
    if not recipe.get("factor_recipe_id") or factor_value is None:
        return {}
    return {"factor_recipe_id": recipe["factor_recipe_id"], "factor_value": _format_decimal(factor_value)}


def _source_value(analysis: Any, source: dict[str, str]) -> Decimal:
    min_col, min_row, max_col, max_row = range_boundaries(source["range"])
    if min_col == max_col and min_row == max_row:
        return _single_value(analysis, source)
    return _range_sum(analysis, source)


def _single_value(analysis: Any, source: dict[str, str]) -> Decimal:
    min_col, min_row, max_col, max_row = range_boundaries(source["range"])
    if min_col != max_col or min_row != max_row:
        raise FoundryError("direct evidence recipe requires a single cell")
    values = _exact_values(analysis, source)
    from openpyxl.utils import get_column_letter

    coordinate = f"{get_column_letter(min_col)}{min_row}"
    if coordinate not in values:
        raise FoundryError(f"source cell {source['sheet']}!{coordinate} is not numeric")
    return values[coordinate]


def _range_sum(analysis: Any, source: dict[str, str]) -> Decimal:
    min_col, min_row, max_col, max_row = range_boundaries(source["range"])
    values = _exact_values(analysis, source)
    from openpyxl.utils import get_column_letter

    total = Decimal("0")
    seen = False
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            coordinate = f"{get_column_letter(col)}{row}"
            if coordinate not in values:
                raise FoundryError(f"source cell {source['sheet']}!{coordinate} is not numeric")
            total += values[coordinate]
            seen = True
    if not seen:
        raise FoundryError(f"source range {source['sheet']}!{source['range']} has no numeric cells")
    return total


def _exact_values(analysis: Any, source: dict[str, str]) -> dict[str, Decimal]:
    dataset = analysis.read_dataset(source["dataset_id"])
    if dataset.get("source_hash") != source["source_hash"]:
        raise FoundryError("evidence recipe source_hash does not match archived snapshot")
    data = _dataset_bytes(analysis, dataset, source)
    # The rule the snapshot was filed under, so a run computed from its stored
    # rows is checked against the same reading of the same bytes.
    return read_exact_numerics(data, source["sheet"], archived_numeric_parser(dataset, data))


def _issue(
    recipe: dict[str, Any],
    code: str,
    message: str | None,
    *,
    actual: Decimal | None,
    expected: Decimal | None,
    scale_hint: bool = False,
    precision_hint: bool = False,
) -> dict[str, Any]:
    expected_text = _format_decimal(expected) if expected is not None else ""
    actual_text = _format_decimal(actual) if actual is not None else None
    exponent = _unit_scale_hint(actual, expected) if scale_hint else None
    location = _recipe_location(recipe)
    if recipe["kind"] in {"direct", "sum"}:
        source_label = f"{location['sheet']}!{location['range']}"
    else:
        operator = "÷" if recipe["kind"] == "ratio" else "−"
        source_label = (
            f"{location['left']['sheet']}!{location['left']['range']} "
            f"{operator} {location['right']['sheet']}!{location['right']['range']}"
        )
    if message is None:
        message = _mismatch_message(source_label, actual_text, expected_text, recipe["scale"], exponent)
        if precision_hint:
            message += (
                "运行引用值与独立计算值只差舍入（按引用值的小数位舍入，或脚本用二进制浮点数计算），"
                "差异来自运行结果的精度损失；请用精确小数（如 Decimal）计算，并在运行结果中保存完整精度的数值"
                f"（展示时再舍入）。{PRECISION_RULE}"
            )
    elif precision_hint:
        message = message.rstrip("。") + "。" + (
            "所引用的系数规则只因运行值的精度损失（舍入或二进制浮点数）而未通过；请用精确小数计算系数和估算，"
            f"并在运行结果中保存完整精度的数值。{PRECISION_RULE}"
        )
    return {
        "code": code,
        "severity": "error",
        "recipe_id": recipe["id"],
        "label": recipe["label"],
        "business_key": recipe.get("business_key", {}),
        "expected": expected_text,
        **({"actual": actual_text} if actual_text is not None else {}),
        "location": location,
        "claim": recipe.get("claim", {}),
        **({"suggested_scale": _format_decimal(_shifted(recipe["scale"], exponent))} if exponent is not None else {}),
        **({"precision_hint": FULL_PRECISION_REQUIRED} if precision_hint else {}),
        "message": message,
    }


UNIT_SCALE_EXPONENTS = {3, 4, 6, 8}  # 千元、万元、百万元、亿元
FULL_PRECISION_REQUIRED = "full_precision_required"
# A quotient that does not terminate (344487.0552 ÷ 527999.9999) has no exact decimal: every run value of it
# is rounded, the server's own 28-digit Decimal included. Such a value matches when it keeps at least a
# Decimal context's full precision (28 significant digits) and agrees with the exact quotient to 10^-25 of
# its size, so neither a higher precision nor another order of exact-decimal operations decides the result.
# Every value that terminates (each direct/sum/diff recipe, 300 ÷ 480 = 0.625) is still compared exactly.
FULL_PRECISION_DIGITS = 28
QUOTIENT_TOLERANCE = Fraction(1, 10**25)
PRECISION_RULE = (
    "能除尽的值须与独立计算值完全相等；除不尽的比值及由它算出的估算至少保留 28 位有效数字"
    "（Python Decimal 的默认精度），更高精度或不同运算顺序都可以，精确比对不放宽。"
)


# Decimals with an exponent or a digit count beyond this are never expanded into an exact rational (1E-1000000
# would be a million-digit integer); a recipe touching one is compared as Decimals, as before round 7.
EXACT_SIZE_LIMIT = 400


def _exact(value: Decimal) -> Fraction | None:
    sign, digits, exponent = value.as_tuple()
    if not value.is_finite() or abs(exponent) > EXACT_SIZE_LIMIT or len(digits) > EXACT_SIZE_LIMIT:
        return None
    return Fraction(value)


def _exact_product(left: Fraction | None, right: Fraction | None) -> Fraction | None:
    return left * right if left is not None and right is not None else None


def _matches_exact(exact: Fraction | None, claimed: Decimal, actual: Decimal) -> bool:
    """Whether the run value is the recipe's value: exactly, or for a quotient that does not terminate, to
    full precision (see FULL_PRECISION_DIGITS). Without an exact rational (extreme sizes) the server's
    Decimal `actual` is compared as it was."""
    if exact is None:
        return actual == claimed
    value = _exact(claimed)
    if value is None:
        return False
    if value == exact:
        return True
    if _terminates(exact) or len(claimed.as_tuple().digits) < FULL_PRECISION_DIGITS:
        return False
    return abs(value - exact) <= abs(exact) * QUOTIENT_TOLERANCE


def _terminates(value: Fraction) -> bool:
    denominator = value.denominator
    for prime in (2, 5):
        while denominator % prime == 0:
            denominator //= prime
    return denominator == 1


def _rounding_hint(actual: Decimal | None, expected: Decimal | None) -> bool:
    """Whether the claimed value is the independent value with its precision lost: rounded to the claim's
    own decimals, or computed as a binary float by the script that wrote the run.

    Like _unit_scale_hint this only explains a mismatch; the comparison stays exact. Both half-up
    and half-even rounding are recognised, since scripts and SQL engines differ.
    """
    if actual is None or expected is None or not actual.is_finite() or not expected.is_finite():
        return False
    if actual == expected:
        return False
    return _decimal_rounding(actual, expected) or _float_rounding(actual, expected)


# A float result drifts from the exact value by a few units in the last place per operation; 16 ulps covers a
# short chain (a factor, then a product) and is far below any difference a wrong cell or value produces.
FLOAT_ROUNDING_ULPS = 16


def _float_rounding(actual: Decimal, expected: Decimal) -> bool:
    """Whether the claimed value is what a script computing in binary floats would write (its repr).

    Such a value is not a decimal rounding of the exact one (repr(0.1 + 0.2) is 0.30000000000000004), so the
    decimal check misses it. The claim must carry a double's worth of digits (15 to 17: a repr never has
    more) and sit within a few ulps; a 28- or 30-digit value is a Decimal result, never a float repr.
    """
    try:
        claimed = float(expected)
    except (OverflowError, ValueError):  # pragma: no cover - finite Decimal always converts
        return False
    if not math.isfinite(claimed) or claimed == 0 or not 15 <= len(expected.as_tuple().digits) <= 17:
        return False
    tolerance = Decimal(math.ulp(claimed)) * FLOAT_ROUNDING_ULPS
    return abs(actual - expected) <= tolerance


def _decimal_rounding(actual: Decimal, expected: Decimal) -> bool:
    sign, digits, actual_exponent = actual.as_tuple()
    exponent = expected.as_tuple().exponent
    if actual_exponent >= exponent:
        return False
    # Rounded on the digit tuple (like _unit_scale_hint) so extreme exponents cannot overflow.
    dropped_count = exponent - actual_exponent
    if dropped_count > len(digits):
        kept, dropped = 0, (0,)
    else:
        kept = int("".join(map(str, digits[: len(digits) - dropped_count])) or "0")
        dropped = digits[len(digits) - dropped_count :]
    first, rest = dropped[0], any(dropped[1:])
    half_up = first >= 5
    half_even = first > 5 or (first == 5 and (rest or kept % 2 == 1))
    for round_up in {half_up, half_even}:
        value = kept + (1 if round_up else 0)
        if Decimal((sign, tuple(int(char) for char in str(value)), exponent)) == expected:
            return True
    return False


def _accepts_scale(recipe: dict[str, Any], recipe_by_id: dict[str, dict[str, Any]] | None) -> bool:
    """A factor-bound recipe and the ratio it references must keep scale 1 (see _clean_recipes)."""
    if recipe.get("factor_recipe_id"):
        return False
    return not any(other.get("factor_recipe_id") == recipe["id"] for other in (recipe_by_id or {}).values())


def _unit_scale_hint(actual: Decimal | None, expected: Decimal | None) -> int | None:
    """Power k with expected == actual × 10^k exactly, when k is a unit step (千/万/百万/亿).

    Signs and significant digits are compared instead of dividing, so rounding at the
    context precision cannot fake an exact power and extreme exponents cannot overflow.
    """
    if actual is None or expected is None or not actual.is_finite() or not expected.is_finite():
        return None
    if actual == 0 or expected == 0:
        return None
    actual_sign, actual_digits, actual_exponent = _significand(actual)
    expected_sign, expected_digits, expected_exponent = _significand(expected)
    if actual_sign != expected_sign or actual_digits != expected_digits:
        return None
    exponent = expected_exponent - actual_exponent
    return exponent if abs(exponent) in UNIT_SCALE_EXPONENTS else None


def _significand(value: Decimal) -> tuple[int, tuple[int, ...], int]:
    """Sign, digits without trailing zeros and the matching exponent; exact, unlike normalize()."""
    sign, digits, exponent = value.as_tuple()
    kept = len(digits)
    while kept > 1 and digits[kept - 1] == 0:
        kept -= 1
    return sign, digits[:kept], exponent + len(digits) - kept


def _shifted(value: Decimal, exponent: int) -> Decimal:
    """value × 10^exponent, built from the digit tuple so no context rounding or overflow applies."""
    sign, digits, value_exponent = _significand(value)
    return Decimal((sign, digits, value_exponent + exponent))


def _mismatch_message(
    source_label: str,
    actual_text: str | None,
    expected_text: str,
    scale: Decimal,
    exponent: int | None,
) -> str:
    message = f"来源区域 {source_label} 的独立计算为 {actual_text}，运行引用值为 {expected_text}。"
    if exponent is None:
        return message
    cause = "缺少单位换算" if scale == 1 else "单位换算有误"
    return (
        f"{message}两者恰好相差 {10 ** abs(exponent)} 倍，疑似{cause}（如元与万元）；"
        f"若确属换算，请在该核验规则中设置 scale = {_format_decimal(_shifted(scale, exponent))}。"
    )


def _claim(value: Any) -> dict[str, Any]:
    if not isinstance(value, dict):
        raise FoundryError("evidence recipe requires a claim run_id, row, column binding")
    run_id = _text(value.get("run_id"), "evidence recipe claim run_id")
    row = value.get("row")
    if isinstance(row, bool) or not isinstance(row, int) or row < 0:
        raise FoundryError("evidence recipe claim row must be a non-negative integer")
    column = _text(value.get("column"), "evidence recipe claim column")
    return {"run_id": run_id, "row": row, "column": column}


def _claim_expected(analysis: Any, claim: dict[str, Any]) -> Decimal:
    try:
        value = _claim_row(analysis, claim)[claim["column"]]
        return _decimal(value, "evidence recipe claim value")
    except FoundryError:
        raise
    except Exception as exc:
        raise FoundryError("evidence recipe claim does not point to a numeric run value") from exc


def _claim_row(analysis: Any, claim: dict[str, Any]) -> dict[str, Any]:
    try:
        run = analysis.get_run(claim["run_id"])
        row = run["result"]["rows"][claim["row"]]
    except FoundryError:
        raise
    except Exception as exc:
        raise FoundryError("evidence recipe claim does not point to a run row") from exc
    if not isinstance(row, dict):
        raise FoundryError("evidence recipe claim does not point to a run row")
    return row


def _recipe_location(recipe: dict[str, Any]) -> dict[str, Any]:
    if recipe["kind"] in {"direct", "sum"}:
        location = {key: recipe[key] for key in ("dataset_id", "source_hash", "sheet", "range")}
        if recipe.get("factor_recipe_id"):
            location["factor_recipe_id"] = recipe["factor_recipe_id"]
        return location
    return {
        "left": {key: recipe["left"][key] for key in ("dataset_id", "source_hash", "sheet", "range")},
        "right": {key: recipe["right"][key] for key in ("dataset_id", "source_hash", "sheet", "range")},
    }


def _assert_matching_factor_business_key(dependent: dict[str, Any], factor: dict[str, Any]) -> None:
    dependent_key = dependent.get("business_key") or {}
    factor_key = factor.get("business_key") or {}
    if (
        not dependent_key.get("metric_key")
        or not dependent_key.get("period")
        or dependent_key.get("metric_key") != factor_key.get("metric_key")
        or dependent_key.get("period") != factor_key.get("period")
    ):
        raise FoundryError("factor_recipe_id requires matching business_key metric_key and period")


def _assert_claim_business_key(analysis: Any, recipe: dict[str, Any]) -> None:
    business_key = recipe.get("business_key") or {}
    metric_key = business_key.get("metric_key")
    period = business_key.get("period")
    if not metric_key or not period:
        raise FoundryError("factor_recipe_id requires business_key metric_key and period")
    row = _claim_row(analysis, recipe["claim"])
    if row.get("metric_key") != metric_key or row.get("period") != period:
        raise FoundryError("evidence recipe business_key metric_key and period must match claim row")


def _text(value: Any, label: str, *, required: bool = True) -> str:
    if not isinstance(value, str):
        raise FoundryError(f"{label} must be text")
    text = value.strip()
    if required and not text:
        raise FoundryError(f"{label} is required")
    return text


def _decimal(value: Any, label: str) -> Decimal:
    try:
        decimal = Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError) as exc:
        raise FoundryError(f"{label} must be decimal text") from exc
    if not decimal.is_finite():
        raise FoundryError(f"{label} must be finite")
    return decimal


def _format_decimal(value: Decimal) -> str:
    return format(value, "f")


def _clean_consistency_sources(sources: Any) -> dict[str, dict[str, Any]]:
    if sources in (None, {}, []):
        return {}
    items: list[Any]
    if isinstance(sources, dict):
        items = [{**value, "id": key} if isinstance(value, dict) else value for key, value in sources.items()]
    elif isinstance(sources, list):
        items = sources
    else:
        raise FoundryError("consistency_sources must be an object or array")
    if len(items) > MAX_CONSISTENCY_SOURCE_COUNT:
        raise FoundryError("consistency_sources exceeds the source count limit")
    cleaned: dict[str, dict[str, Any]] = {}
    for item in items:
        if not isinstance(item, dict):
            raise FoundryError("consistency source must be an object")
        source_id = _text(item.get("id"), "consistency source id")
        kind = item.get("kind")
        if kind not in {"number", "text"}:
            raise FoundryError("consistency source kind must be number or text")
        source = {
            "id": source_id,
            "kind": kind,
            "dataset_id": _text(item.get("dataset_id"), "consistency source dataset_id"),
            "source_hash": _text(item.get("source_hash"), "consistency source source_hash"),
            "label": _text(item.get("label", source_id), "consistency source label", required=False),
        }
        if kind == "number":
            source.update(_source(item))
            source["kind"] = kind
            source["id"] = source_id
            source["label"] = _text(item.get("label", source_id), "consistency source label", required=False)
            source["scale"] = _decimal(item.get("scale", "1"), "consistency source scale")
        else:
            paragraph = item.get("paragraph")
            has_paragraph = paragraph is not None
            has_workbook_cell = any(item.get(key) is not None for key in ("sheet", "range"))
            if has_paragraph and has_workbook_cell:
                raise FoundryError("text consistency source must use either paragraph or workbook sheet/range")
            if has_paragraph:
                if isinstance(paragraph, bool) or not isinstance(paragraph, int) or paragraph < 1:
                    raise FoundryError("consistency source paragraph must be a positive integer")
                source["paragraph"] = paragraph
            elif has_workbook_cell:
                workbook_source = _source(item)
                min_col, min_row, max_col, max_row = range_boundaries(workbook_source["range"])
                if min_col != max_col or min_row != max_row:
                    raise FoundryError("text consistency workbook source range must be a single cell")
                source.update(workbook_source)
                source["kind"] = kind
                source["id"] = source_id
                source["label"] = _text(item.get("label", source_id), "consistency source label", required=False)
            else:
                raise FoundryError("text consistency source requires paragraph or workbook sheet/range")
        if source_id in cleaned:
            raise FoundryError("consistency source ids must be unique")
        cleaned[source_id] = source
    return cleaned


def _read_authorized_dataset(analysis: Any, source: dict[str, Any]) -> dict[str, Any]:
    dataset = analysis.read_dataset(source["dataset_id"])
    if dataset.get("source_hash") != source["source_hash"]:
        raise FoundryError("consistency source_hash does not match archived snapshot")
    return dataset


def _dataset_bytes(analysis: Any, dataset: dict[str, Any], source: dict[str, Any]) -> bytes:
    folder = analysis._inside(analysis.root / "datasets" / source["dataset_id"])
    path = analysis._inside(folder / dataset["source_file"])
    data = path.read_bytes()
    if _hash_bytes(data) != source["source_hash"]:
        raise FoundryError("archived source integrity check failed")
    return data


def _archived_paragraph(dataset: dict[str, Any], source: dict[str, Any], analysis: Any | None = None) -> dict[str, Any]:
    if source.get("sheet") or source.get("range"):
        if analysis is None:
            raise FoundryError("workbook text source requires archived workbook access")
        return _archived_workbook_text_cell(analysis, dataset, source)
    if dataset.get("kind") != "background":
        raise FoundryError("text consistency paragraph source requires an archived text dataset")
    paragraphs = dataset.get("paragraphs") or []
    if source.get("paragraph"):
        paragraph = next((item for item in paragraphs if item.get("paragraph") == source["paragraph"]), None)
        if not paragraph:
            raise FoundryError("consistency source paragraph was not found")
        return {"paragraph": paragraph.get("paragraph"), "line": paragraph.get("line"), "text": str(paragraph.get("text") or "")}
    text = str(dataset.get("text") or "")
    if not text.strip():
        raise FoundryError("consistency source text is empty")
    return {"text": text}


def _archived_workbook_text_cell(analysis: Any, dataset: dict[str, Any], source: dict[str, Any]) -> dict[str, Any]:
    data = _dataset_bytes(analysis, dataset, source)
    min_col, min_row, max_col, max_row = range_boundaries(source["range"])
    if min_col != max_col or min_row != max_row:
        raise FoundryError("text consistency workbook source range must be a single cell")
    book = load_workbook(filename=io.BytesIO(data), data_only=False, read_only=True)
    try:
        if source["sheet"] not in book.sheetnames:
            raise FoundryError(f"no worksheet: {source['sheet']}")
        sheet = book[source["sheet"]]
        coordinate = f"{get_column_letter(min_col)}{min_row}"
        value = sheet[coordinate].value
    finally:
        book.close()
    if value is None or (isinstance(value, str) and not value.strip()):
        raise FoundryError(f"text consistency source cell {source['sheet']}!{source['range']} is empty")
    if isinstance(value, str) and value.startswith("="):
        raise FoundryError("text consistency source cell formula text is not accepted")
    return {"sheet": source["sheet"], "range": source["range"], "text": str(value)}


def _resolve_consistency_rules(rules: list[Any] | None, source_values: dict[str, Decimal]) -> list[dict[str, Any]]:
    if rules in (None, []):
        return []
    if not isinstance(rules, list):
        raise FoundryError("consistency_rules must be an array")
    if len(rules) > MAX_CONSISTENCY_RULE_COUNT:
        raise FoundryError("consistency_rules exceeds the rule count limit")
    resolved: list[dict[str, Any]] = []
    for rule in rules:
        if not isinstance(rule, dict):
            raise FoundryError("consistency rule must be an object")
        copy = dict(rule)
        operands = copy.get("values", copy.get("declared_values"))
        if isinstance(operands, list):
            new_operands = []
            for operand in operands:
                if not isinstance(operand, dict):
                    new_operands.append(operand)
                    continue
                item = dict(operand)
                if "value" in item or "declared_value" in item:
                    raise FoundryError("consistency direct values must be read from an authorized source")
                source_id = item.get("source_id")
                if source_id and not item.get("recipe_result_id"):
                    if source_id not in source_values:
                        raise FoundryError("consistency numeric source_id is not authorized")
                    item["value"] = _format_decimal(source_values[source_id])
                new_operands.append(item)
            if "values" in copy:
                copy["values"] = new_operands
            else:
                copy["declared_values"] = new_operands
        resolved.append(copy)
    return resolved
