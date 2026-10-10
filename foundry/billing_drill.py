"""往下拆: a collection rate of a 收费明细 run, taken apart 城市公司 → 项目 → 房产 → 账单.

Each level is the bills of the one above split by one more field, with the same 口径, the same column mapping and
the same statistics date as the run (方案 notes/intelligent-ui-borrowing-plan-20261010.md, 三 4). Every level says
whether its rows add up to the row it was opened from, and the whole to the figure the run filed; a bill names the
lines of the export it was read from, so a manager can go from 82.12% to the file and line of one unpaid bill.

The bills are worked out once per run, as the run worked them (billing_detail._prepare), and kept beside the run:
the run, its detail tables and its mapping never change, so neither does what they give. A run filed by another
engine version is not taken apart: the same rows would no longer give its figures.
"""

from __future__ import annotations

import os
import uuid
from decimal import Decimal
from pathlib import Path
from typing import Any

import duckdb

from foundry import duckdb_offline
from foundry.billing_detail import COLLECTION_ENGINE_VERSION, BillingDetailStore, _money, _prepare, _rate, _window
from foundry.errors import FoundryError

# Raise when the cached tables change shape or meaning; a cache of another version is built again.
CACHE_VERSION = 1
# The rows a level shows, those with the most still owed first; the export holds them all.
MAX_ROWS = 200
# The lines of the export a bill names on screen, its own and the receipts against it; an export names them all.
MAX_LINES = 5
MAX_VALUE = 200

# What a bill owes and has paid toward each rate, as the run's FIGURE_SQL counts it.
_PAID = "least(paid_before + paid_within, greatest(net, 0))"
_OPENING = "greatest(net - paid_before, 0)"
_RECOVERED = f"least(paid_within, {_OPENING})"
FIGURES: dict[str, dict[str, Any]] = {
    "collection_rate_current": {"label": "当期收缴率", "base_label": "本年应收", "done_label": "当期已收",
                                "where": "place = 'current'", "base": "net", "done": _PAID, "run": (("receivable",), ("collected",))},
    "collection_rate_month": {"label": "本月收缴率", "base_label": "本月应收", "done_label": "本月已收",
                              "where": "place = 'current' AND period = getvariable('as_of_month')", "base": "net", "done": _PAID,
                              "run": (("month_receivable",), ("month_collected",))},
    "arrears_recovery_rate": {"label": "清欠率", "base_label": "年初往年欠费", "done_label": "本年收回",
                              "where": "place = 'arrears'", "base": _OPENING, "done": _RECOVERED,
                              "run": (("arrears_opening",), ("arrears_collected",))},
    "collection_rate_total": {"label": "综合收缴率", "base_label": "本年应收 + 年初欠费", "done_label": "当期已收 + 本年清欠",
                              "where": "place IN ('current', 'arrears')",
                              "base": f"CASE place WHEN 'current' THEN net ELSE {_OPENING} END",
                              "done": f"CASE place WHEN 'current' THEN {_PAID} ELSE {_RECOVERED} END",
                              "run": (("receivable", "arrears_opening"), ("collected", "arrears_collected"))},
}
# The levels a rate is taken apart by, outermost first; a level whose column the run did not map is skipped.
LEVELS = (("city", "城市公司"), ("project", "项目"), ("property_id", "房产"))
BILL_LEVEL = {"key": "bill", "label": "账单"}
EMPTY = "（空）"


def _levels(run: dict) -> list[dict]:
    inputs = run["inputs"]
    mapped = {"city": "city" in (inputs.get("property") or {}).get("mapping", {}),
              "project": "project" in inputs["receivable"]["mapping"], "property_id": True}
    return [{"key": key, "label": label} for key, label in LEVELS if mapped[key]]


def _figure(run: dict, figure: Any) -> dict:
    if figure not in FIGURES:
        raise FoundryError("figure is one of: " + ", ".join(FIGURES))
    if figure == "collection_rate_total" and run["params"]["total_method"] != "bounded":
        raise FoundryError("综合收缴率按本年全部实收算，含没对上账单的预收，拆不到账单；拆当期收缴率或清欠率")
    return FIGURES[figure]


def _path(levels: list[dict], path: Any) -> list[str | None]:
    if path is None:
        return []
    if not isinstance(path, list) or len(path) > len(levels):
        raise FoundryError(f"path names up to {len(levels)} values, one per level: " + "、".join(level["label"] for level in levels))
    for value in path:
        if value is not None and (not isinstance(value, str) or not value.strip() or len(value) > MAX_VALUE):
            raise FoundryError(f"each value of path is a name of at most {MAX_VALUE} characters, or null for {EMPTY}")
    return list(path)


class Drill:
    """往下拆 over one collection run of a task."""

    def __init__(self, store: BillingDetailStore, run_id: Any):
        self.store = store
        self.run = store.read_run(run_id)
        if self.run.get("engine_version") != COLLECTION_ENGINE_VERSION:
            raise FoundryError("this run was worked out by an earlier version of the calculation; run it again to take it apart")
        self.levels = _levels(self.run)

    # The bills of the run, worked out once.

    def _cache(self) -> Path:
        path = self.store.root / "runs" / f"{self.run['id']}.drill-v{CACHE_VERSION}.duckdb"
        if path.is_file():
            return path
        metas = {}
        for kind, used in self.run["inputs"].items():
            meta = self.store.get(used["id"])
            if meta["database_sha256"] != used["database_sha256"]:
                raise FoundryError(f"{meta['name']} is not the table this run read; run it again")
            # The columns as the run read them, whatever has been confirmed since.
            metas[kind] = {**meta, "confirmation": {"mapping": used["mapping"]}}
        building = path.with_name(f".{path.stem}-{uuid.uuid4().hex}.duckdb")
        connection = duckdb_offline.connect(":memory:")
        try:
            for kind, meta in metas.items():
                source = str(self.store._folder(meta["id"]) / "detail.duckdb").replace("'", "''")
                connection.execute(f"ATTACH '{source}' AS {kind} (READ_ONLY)")
            _window(connection, self.run["params"])
            _prepare(connection, metas, self.run["params"])
            connection.execute(f"ATTACH '{str(building).replace(chr(39), chr(39) * 2)}' AS cache")
            connection.execute("CREATE TABLE cache.ledger AS SELECT bill_id, property_id, fee_item, project, city, period, "
                               "net, paid_before, paid_within, related, place FROM ledger")
            connection.execute("CREATE TABLE cache.bill_source AS SELECT bill_id, file, line FROM bill_lines "
                               "WHERE NOT void AND bill_id IN (SELECT bill_id FROM ledger)")
            connection.execute("CREATE TABLE cache.receipt_source AS SELECT bill_id, file, line FROM placed "
                               "WHERE bill_id IN (SELECT bill_id FROM ledger)")
            connection.execute("DETACH cache")
        except duckdb.Error as exc:
            building.unlink(missing_ok=True)
            raise FoundryError(f"cannot take this run apart: {exc}") from exc
        except BaseException:
            building.unlink(missing_ok=True)
            raise
        finally:
            connection.close()
        # Two drills at once each build a whole copy; whichever lands last stands, and both are the same.
        os.replace(building, path)
        return path

    # One level.

    def level(self, figure: Any, path: Any = None, *, limit: int | None = MAX_ROWS) -> dict:
        spec = _figure(self.run, figure)
        chosen = _path(self.levels, path)
        connection = duckdb_offline.connect(str(self._cache()), read_only=True)
        try:
            _window(connection, self.run["params"])
            return self._level(connection, figure, spec, chosen, limit)
        except duckdb.Error as exc:
            raise FoundryError(f"cannot take this run apart: {exc}") from exc
        finally:
            connection.close()

    def _level(self, connection, figure: str, spec: dict, path: list, limit: int | None) -> dict:
        main = "NOT related" if self.run["params"]["exclude_related"] else "true"
        rows_of = f"(SELECT *, {spec['base']} AS base, {spec['done']} AS done FROM ledger WHERE {main} AND {spec['where']})"

        def where(depth: int) -> tuple[str, list]:
            terms = [f"{self.levels[index]['key']} IS NOT DISTINCT FROM ?" for index in range(depth)]
            return " AND ".join(["true", *terms]), path[:depth]

        def grouped(depth: int) -> tuple[str, list]:
            key = self.levels[depth]["key"]
            condition, values = where(depth)
            return (f"(SELECT {key} AS value, coalesce(sum(base), 0) AS base, coalesce(sum(done), 0) AS done, "
                    f"count(DISTINCT property_id) FILTER (WHERE base <> 0) AS households "
                    f"FROM {rows_of} WHERE {condition} GROUP BY {key})"), values

        def group(value, base, done, households) -> dict:
            return {"value": value, "label": EMPTY if value is None else value, **_sums(base, done), "properties": households}

        condition, values = where(0)
        base, done = connection.execute(f"SELECT coalesce(sum(base), 0), coalesce(sum(done), 0) FROM {rows_of} WHERE {condition}",
                                        values).fetchone()
        whole = {"label": "全部", **_sums(base, done)}
        figures = self.run["result"]["figures"]
        filed = [sum(Decimal(figures[name]) for name in names) for names in spec["run"]]
        run_check = {"base": _money(filed[0]), "done": _money(filed[1]),
                     "ties": (Decimal(whole["base"]), Decimal(whole["done"])) == tuple(filed)}
        trail = [whole]
        for depth, value in enumerate(path):
            groups, values = grouped(depth)
            found = connection.execute(f"SELECT * FROM {groups} WHERE value IS NOT DISTINCT FROM ?", [*values, value]).fetchone()
            row = group(*found) if found else None
            if row is None:
                raise FoundryError(f"{trail[-1]['label']} has no {self.levels[depth]['label']} {EMPTY if value is None else value}")
            trail.append({"level": self.levels[depth]["key"], "value": value, "label": row["label"],
                          **{name: row[name] for name in ("base", "done", "outstanding", "rate")}})
        # Only the rows shown are read out: a level of every property can run to hundreds of thousands. Its count and
        # totals are summed where they are.
        page = "" if limit is None else f" LIMIT {int(limit)}"
        if len(path) < len(self.levels):
            level = self.levels[len(path)]
            groups, values = grouped(len(path))
            rows = [group(*found) for found in connection.execute(
                f"SELECT * FROM {groups} ORDER BY base - done DESC, value IS NULL, CAST(value AS VARCHAR){page}", values).fetchall()]
            count, base, done = connection.execute(f"SELECT count(*), coalesce(sum(base), 0), coalesce(sum(done), 0) FROM {groups}",
                                                   values).fetchone()
        else:
            level = BILL_LEVEL
            condition, values = where(len(path))
            rows = self._bills(connection, rows_of, (condition, values), page, every_line=limit is None)
            count, base, done = connection.execute(
                f"SELECT count(*), coalesce(sum(base), 0), coalesce(sum(done), 0) FROM {rows_of} WHERE {condition} AND (base <> 0 OR done <> 0)",
                values).fetchone()
        total = {"base": _money(base), "done": _money(done)}
        parent = trail[-1]
        return {
            "run_id": self.run["id"], "figure": figure, "label": spec["label"], "base_label": spec["base_label"],
            "done_label": spec["done_label"], "as_of": self.run["params"]["as_of"], "caliber": self.run["caliber"],
            "levels": [*self.levels, BILL_LEVEL], "path": trail[1:], "level": level,
            "run_check": run_check, "whole": whole,
            "total": {**total, "outstanding": _money(Decimal(total["base"]) - Decimal(total["done"])),
                      "rate": _rate(total["done"], total["base"])},
            "ties": (total["base"], total["done"]) == (parent["base"], parent["done"]),
            "gap": {name: _money(Decimal(total[name]) - Decimal(parent[name])) for name in ("base", "done")},
            "row_count": count, "rows": rows,
        }

    def _bills(self, connection, rows_of: str, condition: tuple[str, list], page: str, *, every_line: bool) -> list[dict]:
        """The bills shown, each with the export lines it was read from: the first few on screen, every one in an export."""
        found = connection.execute(
            f"SELECT bill_id, fee_item, strftime(period, '%Y-%m'), place, base, done FROM {rows_of} "
            f"WHERE {condition[0]} AND (base <> 0 OR done <> 0) ORDER BY base - done DESC, period, bill_id{page}", condition[1]).fetchall()
        bills = [row[0] for row in found]
        files = {kind: [entry["name"] for entry in self.store.get(used["id"])["files"]] for kind, used in self.run["inputs"].items()}
        lines: dict[tuple[str, str], list[str]] = {}
        for kind, table in (("receivable", "bill_source"), ("receipt", "receipt_source")):
            if not bills:
                break
            for bill_id, file, line in connection.execute(
                    f"SELECT bill_id, file, line FROM {table} WHERE bill_id IN (SELECT unnest(?)) ORDER BY bill_id, file, line", [bills]).fetchall():
                lines.setdefault((kind, bill_id), []).append(f"{files[kind][file]} 第 {line} 行")
        shown = slice(None) if every_line else slice(MAX_LINES)
        return [{"value": bill_id, "label": bill_id, "fee_item": fee_item, "period": period, "place": place, **_sums(base, done),
                 "lines": lines.get(("receivable", bill_id), [])[shown],
                 "line_count": len(lines.get(("receivable", bill_id), [])),
                 "receipts": lines.get(("receipt", bill_id), [])[shown],
                 "receipt_count": len(lines.get(("receipt", bill_id), []))}
                for bill_id, fee_item, period, place, base, done in found]


def _sums(base: Any, done: Any) -> dict:
    base, done = Decimal(base or 0), Decimal(done or 0)
    return {"base": _money(base), "done": _money(done), "outstanding": _money(base - done), "rate": _rate(done, base)}


def export(drill: Drill, figure: Any, path: Any, destination: Any) -> dict:
    """One level, every row of it, as an Excel workbook at ``destination`` (a path the manager chose)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    target = Path(str(destination or "")).expanduser()
    if not target.is_absolute() or target.suffix.lower() != ".xlsx":
        raise FoundryError("destination is an absolute path ending in .xlsx")
    level = drill.level(figure, path, limit=None)
    book = Workbook()
    sheet = book.active
    sheet.title = "往下拆"
    trail = " › ".join(["全部", *(step["label"] for step in level["path"])])
    caliber = level["caliber"]
    words = {"confirmed": f"已确认口径 · 第 {caliber.get('version')} 版", "default": "默认口径（未确认）", "trial": "试算"}
    head = [
        [f"{level['label']} · 往下拆到{level['level']['label']}"],
        [f"截至 {level['as_of']} · {words.get(caliber.get('status'), '')} · 计算编号 {level['run_id']}"],
        [f"位置：{trail}"],
        [f"本层合计 {level['base_label']} {level['total']['base']}、{level['done_label']} {level['total']['done']}；"
         + ("等于上一层的数。" if level["ties"] else "和上一层的数对不上。")],
    ]
    for line in head:
        sheet.append(line)
    sheet["A1"].font = Font(bold=True)
    sheet.append([])
    bill = level["level"]["key"] == "bill"
    columns = [level["level"]["label"], *(("收费项目", "账期") if bill else ("户数",)), level["base_label"], level["done_label"], "未收",
               level["label"], *(("应收明细行", "收款行") if bill else ())]
    sheet.append(columns)
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    for row in level["rows"]:
        rate = None if row["rate"] is None else float(Decimal(row["rate"]))
        sheet.append([row["label"], *((row["fee_item"], row["period"]) if bill else (row["properties"],)),
                      float(Decimal(row["base"])), float(Decimal(row["done"])), float(Decimal(row["outstanding"])), rate,
                      *(("；".join(row["lines"]), "；".join(row["receipts"])) if bill else ())])
        for index, name in enumerate(columns):
            cell = sheet.cell(sheet.max_row, index + 1)
            if name in (level["base_label"], level["done_label"], "未收"):
                cell.number_format = "#,##0.00;(#,##0.00)"
            elif name == level["label"]:
                cell.number_format = "0.00%"
    target.parent.mkdir(parents=True, exist_ok=True)
    scratch = target.with_name(f".{target.stem}-{uuid.uuid4().hex}.xlsx")
    try:
        book.save(scratch)
        os.replace(scratch, target)
    finally:
        scratch.unlink(missing_ok=True)
    return {"path": str(target), "rows": level["row_count"], "level": level["level"]}

