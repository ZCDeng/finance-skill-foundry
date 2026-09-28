"""Bounded report visuals backed by immutable task runs, never caller-authored amounts."""

from __future__ import annotations

import copy
import html
import json
import math
import re
from decimal import Decimal, InvalidOperation
from typing import Any, Callable
from types import SimpleNamespace

from foundry.errors import FoundryError

IDENTITY = re.compile(r"^[0-9a-f]{32}$")
DECIMAL = re.compile(r"^[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?$")
CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
KINDS = {"bar", "line", "flow"}
PROOF_KEYS = {"series", "run_id", "snapshot_hash", "source_label", "estimate_notice"}
ESTIMATE_NOTICE = "本图含估算值，经理确认前不作为批复预算。"
COLORS = ("8A6729", "426786", "626262")


def _exact_label(value: Any) -> str:
    if value is None:
        return "缺失"
    raw = str(value).strip()
    match = re.fullmatch(r"([+-]?)(\d*\.?\d*)[eE]([+-]?\d+)", raw)
    if not match:
        return raw
    sign, coefficient, exponent = match.groups()
    digits = coefficient.replace(".", "")
    if not re.search(r"[1-9]", digits):
        return "-0" if sign == "-" else "0"
    point = (coefficient.index(".") if "." in coefficient else len(coefficient)) + int(exponent)
    expanded = ("0." + "0" * -point + digits if point <= 0 else
                digits + "0" * (point - len(digits)) if point >= len(digits) else
                digits[:point] + "." + digits[point:])
    return sign + expanded


def _wrapped_lines(text: str, width: int, size: int) -> int:
    """Match the deterministic glyph budget in the offline TypeScript renderer."""
    lines = 0
    for paragraph in re.split(r"\r?\n", text):
        used = 0.0
        line = ""
        for char in paragraph:
            space = char in "\t\n\r \u00a0\u1680\u2000\u2001\u2002\u2003\u2004\u2005\u2006\u2007\u2008\u2009\u200a\u2028\u2029\u202f\u205f\u3000\ufeff"
            char_width = size * (0.34 if space else 0.62 if ord(char) <= 255 else 1.02)
            if used + char_width > width and line:
                lines += 1
                line, used = "", 0.0
            line += char
            used += char_width
        lines += 1
    return lines


def visual_height(visual: dict) -> int:
    """Keep this arithmetic aligned with app/electron/reports/visuals.ts::visualHeight."""
    n = _wrapped_lines
    kind = "作者解释流程" if visual["kind"] == "flow" else ("类别比较" if visual["kind"] == "bar" else "趋势变化") + " · 单位：" + visual["unit"]
    y = 50 + n(visual["title"], 880, 24) * 36 + n(kind, 880, 14) * 21 + 12
    if visual["takeaway"]:
        y += n("作者结论：" + visual["takeaway"], 880, 17) * 25.5 + 24
    if visual["kind"] != "flow" and visual.get("estimate_notice"):
        y += n(visual["estimate_notice"], 880, 15) * 22.5 + 24
    y += 24
    if visual["kind"] == "flow":
        y += 8
        for index, step in enumerate(visual["steps"]):
            y += 32 + n(step["label"], 684, 19) * 28.5 + (12 + n(step["detail"], 684, 16) * 24 if step["detail"] else 0)
            if index < len(visual["steps"]) - 1:
                y += max(56, n(step["connector"] or "下一步", 648, 14) * 21 + 24)
        y += 24
    elif visual["kind"] == "bar":
        y += max(36, n(visual["category"], 220, 14) * 21 + 12)
        rows = visual["series"]["rows"]
        grouped = len(visual["value_columns"]) > 1
        for index, row in enumerate(rows):
            header = n(str(row[visual["category"]]), 220, 18) * 27 + (8 if grouped else 0)
            if grouped:
                y += header
            for column in visual["value_columns"]:
                y += max(44, n(_exact_label(row[column]), 216, 18) * 27 + 12,
                         n(column, 220, 16) * 24 + 12 if grouped else 12, 0 if grouped else header + 12)
            y += 20 if index == len(rows) - 1 else 28
    else:
        rows = visual["series"]["rows"]
        chart_top = y + 24
        chart_bottom = chart_top + 288
        label_y = chart_top + 12
        for column in visual["value_columns"]:
            label_y += max(76, n(column + " · 末项 " + _exact_label(rows[-1][column]), 232, 18) * 27 + 24)
        y = max(chart_bottom + 68, label_y + 12) + 32
        for start in range(0, len(rows), 4):
            y += max(n(f"{start + offset + 1}. {row[visual['category']]}", 204, 16) * 24 + 16
                     for offset, row in enumerate(rows[start:start + 4]))
        y += 8
    source = "来源：" + visual["source_label"]
    if visual["kind"] == "flow":
        source += " · 作者解释，不构成计算证据"
    y += 20 + n(source, 880, 16) * 24
    if visual["kind"] != "flow":
        y += n("数值保持原精度；几何位置为近似比例。单位按声明展示，不自动换算。", 880, 14) * 21 + 4
        if any(row[column] is None for row in visual["series"]["rows"] for column in visual["value_columns"]):
            y += n("缺失值显示为“缺失”，不补零。", 880, 14) * 21 + 4
    return math.ceil((y + 28) / 4) * 4


def _check_density(visual: dict) -> None:
    if visual_height(visual) > 1200:
        raise FoundryError("report visual is too dense (height exceeds 1200); split the figure or shorten its labels")


def _bounded_series(value: Any, depth: int = 0) -> None:
    if depth > 12:
        raise FoundryError("report visual series is too deeply nested")
    if isinstance(value, str):
        if len(value) > 16 * 1024 or CONTROL.search(value):
            raise FoundryError("report visual series text is too long")
    elif isinstance(value, float):
        if not math.isfinite(value):
            raise FoundryError("report visual series must contain finite JSON values")
    elif isinstance(value, dict):
        for key, item in value.items():
            if not isinstance(key, str) or len(key) > 1000 or key in {"__proto__", "prototype", "constructor"}:
                raise FoundryError("report visual series keys must be text")
            _bounded_series(key, depth + 1)
            _bounded_series(item, depth + 1)
    elif isinstance(value, list):
        for item in value:
            _bounded_series(item, depth + 1)
    elif value is not None and not isinstance(value, (bool, int)):
        raise FoundryError("report visual series must contain JSON values")


def _text(value: Any, name: str, maximum: int, *, required: bool = False) -> str:
    if not isinstance(value, str) or CONTROL.search(value) or len(value) > maximum or (required and not value.strip()):
        raise FoundryError(f"report visual {name} must be {'nonempty ' if required else ''}text (at most {maximum} characters)")
    return value


def decimal_value(value: Any) -> Decimal:
    if isinstance(value, bool) or not isinstance(value, (str, int, float, Decimal)):
        raise FoundryError("report visual values must be finite decimal numbers or null")
    text = str(value)
    if not DECIMAL.fullmatch(text):
        raise FoundryError("report visual values must be finite decimal numbers or null")
    try:
        number = Decimal(text)
        plotted = float(number)
    except (InvalidOperation, OverflowError, ValueError) as exc:
        raise FoundryError("report visual value is outside the plotting range") from exc
    if not number.is_finite() or not math.isfinite(plotted) or (number != 0 and plotted == 0):
        raise FoundryError("report visual value is outside the plotting range")
    return number


def validate_visual(value: Any, *, allow_unresolved: bool = True) -> dict[str, Any]:
    if not isinstance(value, dict) or value.get("version") != 1 or isinstance(value.get("version"), bool):
        raise FoundryError("report visual version must be 1")
    kind = value.get("kind")
    if kind not in KINDS:
        raise FoundryError("report visual kind must be bar, line or flow")
    clean = {"version": 1, "kind": kind,
             "title": _text(value.get("title"), "title", 160, required=True),
             "takeaway": _text(value.get("takeaway", ""), "takeaway", 300)}
    if kind == "flow":
        steps = value.get("steps")
        if not isinstance(steps, list) or not 2 <= len(steps) <= 6:
            raise FoundryError("report flow requires 2-6 steps")
        clean["steps"] = []
        for step in steps:
            if not isinstance(step, dict):
                raise FoundryError("report flow steps must be objects")
            clean["steps"].append({
                "label": _text(step.get("label"), "step label", 60, required=True),
                "detail": _text(step.get("detail", ""), "step detail", 140),
                "connector": _text(step.get("connector", ""), "step connector", 40),
            })
        if clean["steps"][-1]["connector"]:
            raise FoundryError("report flow last connector must be empty")
        clean["source_label"] = _text(value.get("source_label", "作者分析"), "source label", 200) or "作者分析"
        _check_density(clean)
        return clean
    chart_id = value.get("chart_id")
    if not isinstance(chart_id, str) or not IDENTITY.fullmatch(chart_id):
        raise FoundryError("report visual chart_id must name a task-owned chart")
    selected = value.get("value_columns")
    if not isinstance(selected, list) or not 1 <= len(selected) <= 3:
        raise FoundryError("report visual requires 1-3 value columns")
    selected = [_text(item, "value column", 200, required=True) for item in selected]
    if len(set(selected)) != len(selected):
        raise FoundryError("report visual value columns must be unique")
    clean.update(chart_id=chart_id, category=_text(value.get("category"), "category", 200, required=True),
                 value_columns=selected, unit=_text(value.get("unit"), "unit", 40, required=True))
    if clean["category"] in selected:
        raise FoundryError("report visual category cannot also be a value column")
    if "series" not in value:
        if allow_unresolved:
            return clean
        raise FoundryError("report visual requires a frozen chart snapshot")
    series = value["series"]
    if not isinstance(series, dict) or not isinstance(series.get("columns"), list) or not isinstance(series.get("rows"), list):
        raise FoundryError("report visual series requires columns and rows")
    _bounded_series(series)
    if len(series["columns"]) > 256 or len(json.dumps(series, ensure_ascii=False, separators=(",", ":")).encode("utf8")) > 512 * 1024:
        raise FoundryError("report visual series exceeds the size limit")
    names = [column.get("name") if isinstance(column, dict) else None for column in series["columns"]]
    if any(not isinstance(name, str) or not name for name in names) or len(set(names)) != len(names):
        raise FoundryError("report visual series columns must have unique names")
    for column in series["columns"]:
        _text(column["name"], "column name", 200, required=True)
        if "type" in column and not isinstance(column["type"], str):
            raise FoundryError("report visual column type must be text")
    if any(column not in names for column in [clean["category"], *selected]):
        raise FoundryError("report visual selected column is absent from its run")
    rows = series["rows"]
    if not 1 <= len(rows) <= (12 if kind == "bar" else 24):
        raise FoundryError(f"report {kind} requires 1-{12 if kind == 'bar' else 24} rows; narrow the calculation first")
    if kind == "bar" and len(rows) * len(selected) > 12:
        raise FoundryError("report bar can contain at most 12 marks; split the figure or narrow the calculation")
    categories: set[str] = set()
    present = {column: False for column in selected}
    for row in rows:
        if not isinstance(row, dict):
            raise FoundryError("report visual rows must be objects")
        category = row.get(clean["category"])
        if isinstance(category, bool) or not isinstance(category, (str, int, float)) or not str(category).strip():
            raise FoundryError("report visual categories must be nonempty and unique")
        if isinstance(category, float) and not math.isfinite(category):
            raise FoundryError("report visual category must be finite")
        if str(category).strip() in categories:
            raise FoundryError("report visual categories must be nonempty and unique")
        categories.add(str(category).strip())
        for column in selected:
            if column not in row:
                raise FoundryError("report visual selected value is absent; missing values must be null")
            if row[column] is not None:
                decimal_value(row[column])
                present[column] = True
    if not all(present.values()):
        raise FoundryError("report visual value series cannot be entirely null")
    run_id, snapshot_hash = value.get("run_id"), value.get("snapshot_hash")
    if not isinstance(run_id, str) or not IDENTITY.fullmatch(run_id):
        raise FoundryError("report visual requires its frozen run_id")
    if not isinstance(snapshot_hash, str) or not re.fullmatch(r"[0-9a-f]{64}", snapshot_hash):
        raise FoundryError("report visual requires its frozen snapshot hash")
    clean.update(series=copy.deepcopy(series), run_id=run_id, snapshot_hash=snapshot_hash,
                 source_label=_text(value.get("source_label", "归档图"), "source label", 200))
    if value.get("estimate_notice") == ESTIMATE_NOTICE:
        clean["estimate_notice"] = ESTIMATE_NOTICE
    elif value.get("estimate_notice") not in (None, ""):
        raise FoundryError("report visual estimate notice is invalid")
    _check_density(clean)
    return clean


def document_visuals(document: dict) -> list[dict]:
    result = []
    def visit(nodes: list) -> None:
        for node in nodes:
            if not isinstance(node, dict):
                continue
            if node.get("type") == "reportVisual":
                attrs = node.get("attrs")
                result.append(attrs.get("visual") if isinstance(attrs, dict) else None)
            children = node.get("content")
            if isinstance(children, list):
                visit(children)
    if isinstance(document, dict) and isinstance(document.get("content"), list):
        visit(document["content"])
    return result


def series_for_report(series: dict, *, selected_columns: list[str] | None = None) -> dict:
    """Transport displayed scalars losslessly while retaining the archive's original series."""
    result = copy.deepcopy(series)
    for row in result["rows"]:
        if not isinstance(row, dict):
            continue
        for column in row if selected_columns is None else selected_columns:
            number = row.get(column)
            if isinstance(number, (int, float, Decimal)) and not isinstance(number, bool):
                if ((isinstance(number, float) and not math.isfinite(number)) or
                        (isinstance(number, Decimal) and not number.is_finite())):
                    raise FoundryError("report visual displayed numbers must be finite")
                row[column] = str(number)
    return result


def materialize_document(store, document: dict, *, authorize: Callable[[dict], None] | None = None) -> tuple[dict, list[str], list[dict]]:
    """Resolve every input reference anew; a caller's alleged frozen values never confer authority."""
    from foundry.analysis import _hash
    from foundry.chart_deliverables import normalize_chart, verify_snapshots
    from foundry.report_documents import validate_report_document

    resolved = copy.deepcopy(document)
    # Strip proof before document validation, so even malformed client proof is overwritten.
    for visual in document_visuals(resolved):
        if isinstance(visual, dict) and visual.get("kind") in {"bar", "line"}:
            for key in PROOF_KEYS:
                visual.pop(key, None)
    resolved = validate_report_document(resolved, registered_source_ids=store._registered_report_source_ids())
    run_ids, claims = [], []
    for visual in document_visuals(resolved):
        if visual["kind"] == "flow":
            continue
        chart = store._read(store._inside(store.root / "charts" / f"{visual['chart_id']}.json"))
        run = store.get_run(chart["run_id"])
        if authorize:
            authorize(chart.get("bindings", {}))
            authorize(run.get("bindings", {}))
        if chart.get("bindings") != run.get("bindings") or chart.get("snapshot_hashes") != run.get("input_hashes"):
            raise FoundryError("report chart no longer matches its bound run")
        # The existing archive is verified; the immutable run remains the authority for series.
        verify_snapshots(SimpleNamespace(analysis=store), chart)
        frozen = normalize_chart({"label": chart.get("label")}, run=run)
        # Freeze displayed scalar numbers as text before JSON reaches JavaScript's Number.
        # Unselected evidence stays untouched, and geometry alone may use approximate floats.
        series = series_for_report(frozen["series"], selected_columns=[visual["category"], *visual["value_columns"]])
        visual.update(series=series, run_id=run["id"], snapshot_hash=_hash(frozen["snapshot_hashes"]),
                      source_label=frozen["label"])
        visual.update(validate_visual(visual, allow_unresolved=False))
        if run["id"] not in run_ids:
            run_ids.append(run["id"])
        for index, row in enumerate(visual["series"]["rows"]):
            for column in visual["value_columns"]:
                if row[column] is not None:
                    claims.append({"basis": "calculation", "run_id": run["id"], "row": index,
                                   "column": column, "value": row[column]})
    return resolved, run_ids, claims


def visual_markdown(value: dict) -> str:
    from foundry.report_documents import _escape_markdown_text, _table_cell_markdown
    visual = validate_visual(value)
    def escape(text: str) -> str:
        return _escape_markdown_text(html.escape(text, quote=False))
    def cell(text: str) -> str:
        return _table_cell_markdown(escape(text))
    lines = ["### " + escape(visual["title"])]
    if visual["takeaway"]:
        lines.extend(["", escape(visual["takeaway"])])
    if visual.get("estimate_notice"):
        lines.extend(["", visual["estimate_notice"]])
    if visual["kind"] == "flow":
        lines.extend(["", "解释流程（作者分析，不是计算证据）"])
        for index, step in enumerate(visual["steps"], 1):
            lines.extend(["", f"{index}. {escape(step['label'])}" + (f"：{escape(step['detail'])}" if step["detail"] else "")])
            if step["connector"]:
                lines.append("   → " + escape(step["connector"]))
    else:
        lines.extend(["", "单位：" + escape(visual["unit"])])
        if "series" not in visual:
            lines.extend(["", "图表待绑定计算快照。"])
        else:
            columns = [visual["category"], *visual["value_columns"]]
            lines.extend(["", "| " + " | ".join(cell(c) for c in columns) + " |",
                          "| " + " | ".join("---" for _ in columns) + " |"])
            for row in visual["series"]["rows"]:
                lines.append("| " + " | ".join(cell("缺失" if row[c] is None else str(row[c])) for c in columns) + " |")
    lines.extend(["", "来源：" + escape(visual.get("source_label", "待绑定计算快照"))])
    return "\n".join(lines)


def add_visual_worksheets(workbook, document: dict) -> None:
    """Exact text data accompany native Excel charts; plotted numbers are labelled approximations."""
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.comments import Comment

    def literal_text(sheet) -> None:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"

    for number, raw in enumerate(document_visuals(document), 1):
        visual = validate_visual(raw, allow_unresolved=False)
        sheet = workbook.create_sheet(f"图表 {number}")
        sheet.append([visual["title"]])
        sheet.append([visual["takeaway"], visual.get("estimate_notice", "")])
        sheet.append(["来源", visual["source_label"]])
        if visual["kind"] == "flow":
            sheet.append(["解释流程（作者分析，不是计算证据）"])
            sheet.append(["步骤", "标题", "说明", "至下一步关系"])
            for index, step in enumerate(visual["steps"], 1):
                sheet.append([index, step["label"], step["detail"], step["connector"]])
            literal_text(sheet)
            continue
        sheet.append(["单位", visual["unit"]])
        sheet.append(["运行", visual["run_id"], "快照哈希", visual["snapshot_hash"]])
        sheet.append(["精确数据（文本原值；缺失值留空）"])
        columns = [visual["category"], *visual["value_columns"]]
        sheet.append(columns)
        rows = visual["series"]["rows"]
        for row in rows:
            sheet.append([None if row[column] is None else str(row[column]) for column in columns])
        plot_start = len(rows) + 10
        sheet.cell(plot_start, 1, "图形坐标（Excel 数值近似；核对使用上方精确数据）")
        for column, name in enumerate(columns, 1):
            sheet.cell(plot_start + 1, column, name)
        for index, row in enumerate(rows, plot_start + 2):
            sheet.cell(index, 1, str(row[visual["category"]]))
            for column, name in enumerate(visual["value_columns"], 2):
                if row[name] is not None:
                    cell = sheet.cell(index, column, float(decimal_value(row[name])))
                    cell.comment = Comment("原值：" + str(row[name]), "Foundry")
        chart = BarChart() if visual["kind"] == "bar" else LineChart()
        if visual["kind"] == "bar":
            chart.type = "bar"
            chart.grouping = "clustered"
            chart.overlap = 0
        chart.title = visual["title"] + ("（含估算）" if visual.get("estimate_notice") else "")
        chart.y_axis.title = visual["unit"] if visual["kind"] == "line" else visual["category"]
        chart.x_axis.title = visual["category"] if visual["kind"] == "line" else visual["unit"]
        chart.display_blanks = "gap"
        chart.height, chart.width = 12, 24
        chart.add_data(Reference(sheet, min_col=2, max_col=len(columns), min_row=plot_start + 1, max_row=plot_start + len(rows) + 1), titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=1, min_row=plot_start + 2, max_row=plot_start + len(rows) + 1))
        values = [decimal_value(row[column]) for row in rows for column in visual["value_columns"] if row[column] is not None]
        if min(values) >= 0:
            chart.y_axis.scaling.min = 0
        if max(values) <= 0:
            chart.y_axis.scaling.max = 0
        for index, series in enumerate(chart.series):
            series.graphicalProperties.solidFill = COLORS[index]
            series.graphicalProperties.line.solidFill = COLORS[index]
        sheet.add_chart(chart, "F7")
        literal_text(sheet)


def mark_visual_estimates(document: dict, cells: list[dict]) -> int:
    """Add a visible notice only for selected cells the stored server check identifies as estimated."""
    estimated = {entry["cell"] for entry in cells if entry.get("cell") is not None}
    added = 0
    for visual in document_visuals(document):
        if visual["kind"] == "flow" or "series" not in visual:
            continue
        if any((visual["run_id"], row, column) in estimated
               for row in range(len(visual["series"]["rows"])) for column in visual["value_columns"]):
            if visual.get("estimate_notice") != ESTIMATE_NOTICE:
                visual["estimate_notice"] = ESTIMATE_NOTICE
                added += 1
    return added
