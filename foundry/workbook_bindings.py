from __future__ import annotations

import hashlib
import io
import json
import re
import shutil
import zipfile
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.utils.cell import absolute_coordinate

from foundry.errors import FoundryError
from foundry.sheet import put_number, survives_a_cell

MAX_ROWS = 100_000
MAX_COLUMNS = 256
EXCEL_MAX_ROW = 1_048_576
EXCEL_MAX_COLUMN = 16_384
SOURCE_KINDS = {"range", "table", "named_range"}
ROLES = {"input", "output"}
TEXT_POLICIES = {"text", "identifier", "long_identifier"}
DECIMAL_POLICIES = {"amount", "decimal", "normal_amount"}
UNSUPPORTED_ARCHIVE_PREFIXES = ("xl/externalLinks/", "xl/drawings/", "xl/charts/")
UNSUPPORTED_ARCHIVE_NAMES = {"xl/vbaProject.bin"}
FORMULA_LIKE = re.compile(r"^[=+\-@]")


def _sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _json_hash(payload: Any) -> str:
    data = json.dumps(
        payload,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    ).encode("utf8")
    return _sha256_bytes(data)


def workbook_hash(path: str | Path) -> str:
    return _sha256_bytes(Path(path).read_bytes())


def _unsupported_features(path: Path) -> list[str]:
    suffix = path.suffix.lower()
    features: set[str] = set()
    if suffix in {".xlsm", ".xltm"}:
        features.add("macro-enabled extension")
    try:
        with zipfile.ZipFile(path) as archive:
            for name in archive.namelist():
                if name in UNSUPPORTED_ARCHIVE_NAMES:
                    features.add("macro project")
                for prefix in UNSUPPORTED_ARCHIVE_PREFIXES:
                    if name.startswith(prefix):
                        features.add(prefix.removeprefix("xl/").removesuffix("/").replace("/", " "))
    except zipfile.BadZipFile as exc:
        raise FoundryError(f"cannot inspect workbook: {exc}") from exc
    return sorted(features)


@dataclass(frozen=True)
class ResolvedRange:
    sheet: str
    min_col: int
    min_row: int
    max_col: int
    max_row: int

    @property
    def ref(self) -> str:
        return f"{get_column_letter(self.min_col)}{self.min_row}:{get_column_letter(self.max_col)}{self.max_row}"

    @property
    def height(self) -> int:
        return self.max_row - self.min_row + 1

    @property
    def width(self) -> int:
        return self.max_col - self.min_col + 1

    def contains(self, coordinate: str) -> bool:
        min_col, min_row, max_col, max_row = range_boundaries(coordinate)
        return (
            min_col == max_col
            and min_row == max_row
            and self.min_col <= min_col <= self.max_col
            and self.min_row <= min_row <= self.max_row
        )


def _check_bounds(resolved: ResolvedRange) -> None:
    if (
        resolved.min_col < 1
        or resolved.min_row < 1
        or resolved.max_col > EXCEL_MAX_COLUMN
        or resolved.max_row > EXCEL_MAX_ROW
    ):
        raise FoundryError("binding range exceeds Excel worksheet bounds")
    if resolved.width > MAX_COLUMNS or resolved.height > MAX_ROWS + 1:
        raise FoundryError("binding range exceeds foundry workbook limits")


def _resolve_source(book: Workbook, source: dict[str, Any]) -> ResolvedRange:
    kind = source.get("kind")
    if kind not in SOURCE_KINDS:
        raise FoundryError("binding source kind must be range, table, or named_range")
    if kind == "range":
        sheet_name = str(source.get("sheet") or "")
        if sheet_name not in book.sheetnames:
            raise FoundryError(f"no worksheet: {sheet_name}")
        try:
            min_col, min_row, max_col, max_row = range_boundaries(str(source.get("range") or ""))
        except ValueError as exc:
            raise FoundryError("invalid binding range") from exc
        resolved = ResolvedRange(sheet_name, min_col, min_row, max_col, max_row)
        _check_bounds(resolved)
        return resolved
    if kind == "table":
        table_name = str(source.get("name") or "")
        for sheet in book.worksheets:
            if table_name in sheet.tables:
                table = sheet.tables[table_name]
                min_col, min_row, max_col, max_row = range_boundaries(table.ref)
                resolved = ResolvedRange(sheet.title, min_col, min_row, max_col, max_row)
                _check_bounds(resolved)
                return resolved
        raise FoundryError(f"no table: {table_name}")
    name = str(source.get("name") or "")
    defined = book.defined_names.get(name)
    if defined is None:
        raise FoundryError(f"no named range: {name}")
    destinations = list(defined.destinations)
    if len(destinations) != 1:
        raise FoundryError(f"named range {name} must resolve to exactly one range")
    sheet_name, ref = destinations[0]
    if sheet_name not in book.sheetnames:
        raise FoundryError(f"no worksheet: {sheet_name}")
    try:
        min_col, min_row, max_col, max_row = range_boundaries(ref.replace("$", ""))
    except ValueError as exc:
        raise FoundryError(f"named range {name} must point to a rectangular cell range") from exc
    resolved = ResolvedRange(sheet_name, min_col, min_row, max_col, max_row)
    _check_bounds(resolved)
    return resolved


def _normalize_field(field: dict[str, Any]) -> dict[str, Any]:
    name = str(field.get("name") or "").strip()
    if not name:
        raise FoundryError("field name must be nonempty")
    kind = str(field.get("type") or field.get("policy") or "text")
    policy = str(field.get("policy") or kind)
    if policy in DECIMAL_POLICIES:
        ftype = "decimal"
    elif policy in TEXT_POLICIES or kind in {"string", "text"}:
        ftype = "string"
    elif kind in {"date", "boolean"}:
        ftype = kind
    else:
        raise FoundryError(f"unsupported field policy: {policy}")
    return {"name": name, "type": ftype, "policy": policy, "unit": str(field.get("unit") or "")}


def _fields(contract: dict[str, Any]) -> list[dict[str, Any]]:
    fields = [_normalize_field(field) for field in contract.get("fields") or []]
    if not fields:
        raise FoundryError("binding contract requires fields")
    if len({field["name"].casefold() for field in fields}) != len(fields):
        raise FoundryError("binding field names must be unique")
    return fields


def _cell_payload(cell: Cell) -> dict[str, Any]:
    value = cell.value
    if isinstance(value, Decimal):
        value = format(value, "f")
    elif hasattr(value, "isoformat"):
        value = value.isoformat()
    return {
        "coordinate": cell.coordinate,
        "value": value,
        "data_type": cell.data_type,
        "number_format": cell.number_format,
        "style_id": cell.style_id,
    }


def cell_hash(cell: Cell) -> str:
    return _json_hash(_cell_payload(cell))


def _range_cell_hashes(sheet: Any, resolved: ResolvedRange) -> dict[str, str]:
    return {
        sheet.cell(row=row, column=col).coordinate: cell_hash(sheet.cell(row=row, column=col))
        for row in range(resolved.min_row, resolved.max_row + 1)
        for col in range(resolved.min_col, resolved.max_col + 1)
    }


def _open(path: str | Path, *, data_only: bool = False) -> Workbook:
    try:
        if Path(path).stat().st_size > 32 * 1024 * 1024:
            raise FoundryError("workbook exceeds 32 MiB limit")
        with zipfile.ZipFile(path) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 128 * 1024 * 1024:
                raise FoundryError("expanded workbook exceeds 128 MiB limit")
        return load_workbook(path, data_only=data_only)
    except Exception as exc:
        raise FoundryError(f"cannot open workbook: {exc}") from exc


def open_workbook_bytes(data: bytes, *, data_only: bool = False) -> Workbook:
    """Open workbook bytes under the same size guards ``_open`` applies to a path.

    A caller holding archived snapshot bytes has no file to point at, and the
    non-read-only workbook is what carries merged-cell ranges and cell
    coordinates.
    """
    try:
        if len(data) > 32 * 1024 * 1024:
            raise FoundryError("workbook exceeds 32 MiB limit")
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 128 * 1024 * 1024:
                raise FoundryError("expanded workbook exceeds 128 MiB limit")
        return load_workbook(io.BytesIO(data), data_only=data_only)
    except FoundryError:
        raise
    except Exception as exc:
        raise FoundryError(f"cannot open workbook: {exc}") from exc


def inspect_workbook(path: str | Path, bindings: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    source = Path(path)
    features = _unsupported_features(source)
    book = _open(source)
    try:
        resolved_bindings = []
        for binding in bindings or []:
            resolved = _resolve_source(book, binding["source"])
            sheet = book[resolved.sheet]
            resolved_bindings.append(
                {
                    "id": binding.get("id") or binding.get("binding_id"),
                    "role": binding.get("role", "input"),
                    "source": binding["source"],
                    "resolved": {
                        "sheet": resolved.sheet,
                        "range": resolved.ref,
                        "rows": resolved.height,
                        "columns": resolved.width,
                    },
                    "baseline_hashes": _range_cell_hashes(sheet, resolved),
                }
            )
        return {
            # No path. The caller passed one in and every reader of this result
            # reaches the workbook by dataset id; returning it only carried the
            # host's directory layout out through the tool surface.
            "sha256": workbook_hash(source),
            "sheet_names": list(book.sheetnames),
            "tables": sorted(table for sheet in book.worksheets for table in sheet.tables),
            "named_ranges": sorted(name for name in book.defined_names),
            "unsupported_features": features,
            "copy_rewrite_supported": not features,
            "bindings": resolved_bindings,
        }
    finally:
        book.close()


def create_binding_contract(
    path: str | Path,
    source: dict[str, Any],
    fields: list[dict[str, Any]],
    *,
    role: str = "input",
    writable: bool = False,
    header: bool = True,
    binding_id: str | None = None,
) -> dict[str, Any]:
    if role not in ROLES:
        raise FoundryError("binding role must be input or output")
    normalized_fields = [_normalize_field(field) for field in fields]
    contract = {
        "id": binding_id or _json_hash({"source": source, "fields": normalized_fields, "role": role})[:16],
        "schema_version": 1,
        "role": role,
        "writable": bool(writable),
        "header": bool(header),
        "source": dict(source),
        "fields": normalized_fields,
    }
    inspection = inspect_workbook(path, [contract])["bindings"][0]
    contract["resolved"] = inspection["resolved"]
    contract["workbook_sha256"] = workbook_hash(path)
    contract["baseline_hashes"] = inspection["baseline_hashes"]
    return contract


def _headers(sheet: Any, resolved: ResolvedRange, fields: list[dict[str, Any]], header: bool) -> dict[str, int]:
    if not header:
        if len(fields) > resolved.width:
            raise FoundryError("binding fields exceed range width")
        return {field["name"]: resolved.min_col + offset for offset, field in enumerate(fields)}
    # 「本期发生额\n借方」 is one header that wraps in the cell. Matching the raw
    # text made the manager type the line break to bind a column she can see.
    from foundry.structure_proposal import normalize_header

    header_values: dict[str, int] = {}
    collisions: set[str] = set()
    for col in range(resolved.min_col, resolved.max_col + 1):
        raw = sheet.cell(row=resolved.min_row, column=col).value
        for spelling in (str(raw or "").strip(), normalize_header(raw)):
            key = spelling.casefold()
            if not key:
                continue
            if key in header_values and header_values[key] != col:
                collisions.add(key)
            header_values.setdefault(key, col)
    wanted = [(field["name"], normalize_header(field["name"]).casefold()) for field in fields]
    missing = [name for name, key in wanted if key not in header_values]
    if missing:
        raise FoundryError("binding header missing fields: " + ", ".join(missing))
    ambiguous = sorted({name for name, key in wanted if key in collisions})
    if ambiguous:
        # Two columns that read the same once the line breaks are gone cannot be
        # told apart by name, and guessing one would bind the wrong cells.
        raise FoundryError("binding header is ambiguous for: " + ", ".join(ambiguous))
    return {name: header_values[key] for name, key in wanted}


def _decimal_text(value: Any) -> Decimal:
    if value is None or value == "":
        raise FoundryError("decimal cell is empty")
    if isinstance(value, bool):
        raise FoundryError("boolean is not a decimal value")
    try:
        number = Decimal(str(value).strip())
        if not number.is_finite():
            raise InvalidOperation
        return number
    except (InvalidOperation, ValueError) as exc:
        raise FoundryError(f"invalid decimal value: {value}") from exc


def _coerce_input(cell: Cell, field: dict[str, Any], diagnostics: list[dict[str, Any]]) -> Any:
    value = cell.value
    if value is None:
        return None
    if field["type"] == "decimal":
        amount = _decimal_text(value)
        if isinstance(value, str):
            return format(amount, "f")
        if not survives_a_cell(amount):
            diagnostics.append(
                {
                    "cell": cell.coordinate,
                    "field": field["name"],
                    "code": "excel_numeric_precision",
                    "message": "numeric Excel cell exceeds safe precision; use text for exact import",
                }
            )
        return format(amount, "f")
    if field["policy"] in {"identifier", "long_identifier"}:
        if not isinstance(value, str):
            diagnostics.append(
                {
                    "cell": cell.coordinate,
                    "field": field["name"],
                    "code": "identifier_not_text",
                    "message": "identifier cells should be stored as text to preserve leading zeros and long IDs",
                }
            )
        return str(value)
    if isinstance(value, str):
        return value
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def read_input(path: str | Path, contract: dict[str, Any]) -> dict[str, Any]:
    fields = _fields(contract)
    book = _open(path, data_only=True)
    try:
        resolved = _resolve_source(book, contract["source"])
        sheet = book[resolved.sheet]
        columns = _headers(sheet, resolved, fields, bool(contract.get("header", True)))
        start_row = resolved.min_row + (1 if contract.get("header", True) else 0)
        rows: list[dict[str, Any]] = []
        diagnostics: list[dict[str, Any]] = []
        for row_number in range(start_row, resolved.max_row + 1):
            row: dict[str, Any] = {}
            if not any(sheet.cell(row=row_number, column=col).value is not None for col in columns.values()):
                continue
            for field in fields:
                cell = sheet.cell(row=row_number, column=columns[field["name"]])
                row[field["name"]] = _coerce_input(cell, field, diagnostics)
            rows.append(row)
        return {"binding_id": contract.get("id"), "rows": rows, "diagnostics": diagnostics, "unit_policy": fields}
    finally:
        book.close()


def _write_coordinate(sheet: Any, resolved: ResolvedRange, columns: dict[str, int], contract: dict[str, Any], write: dict[str, Any]) -> str:
    if write.get("cell"):
        coordinate = absolute_coordinate(str(write["cell"]).replace("$", "")).replace("$", "")
    else:
        row_index = int(write.get("row", 1))
        if row_index < 1:
            raise FoundryError("write row must be one-based")
        field_name = str(write.get("field") or "")
        if field_name not in columns:
            raise FoundryError(f"unknown write field: {field_name}")
        row = resolved.min_row + (1 if contract.get("header", True) else 0) + row_index - 1
        coordinate = f"{get_column_letter(columns[field_name])}{row}"
    if not resolved.contains(coordinate):
        raise FoundryError(f"write target outside declared writable area: {coordinate}")
    return coordinate


def _diagnose_output_value(value: Any, field: dict[str, Any]) -> tuple[Any, str, list[dict[str, Any]]]:
    diagnostics: list[dict[str, Any]] = []
    if value is None:
        return None, "empty", diagnostics
    if field["type"] == "decimal":
        amount = _decimal_text(value)
        if survives_a_cell(amount):
            numeric = int(amount) if amount == amount.to_integral_value() else float(amount)
            return numeric, "number", diagnostics
        diagnostics.append(
            {
                "code": "decimal_written_as_text",
                "message": "decimal exceeds safe Excel numeric precision and is written as exact text",
            }
        )
        return format(amount, "f"), "text", diagnostics
    text = str(value)
    if FORMULA_LIKE.match(text):
        diagnostics.append({"code": "formula_like_text", "message": "text value is stored as a string, not an Excel formula"})
    return text, "text", diagnostics


def preview_output_patch(
    path: str | Path,
    contract: dict[str, Any],
    writes: list[dict[str, Any]],
    baseline_hashes: dict[str, str] | None = None,
) -> dict[str, Any]:
    if contract.get("role") != "output" or not contract.get("writable"):
        raise FoundryError("output patch requires a writable output binding")
    fields = _fields(contract)
    field_by_name = {field["name"]: field for field in fields}
    book = _open(path)
    try:
        resolved = _resolve_source(book, contract["source"])
        sheet = book[resolved.sheet]
        columns = _headers(sheet, resolved, fields, bool(contract.get("header", True)))
        baseline = baseline_hashes or contract.get("baseline_hashes") or {}
        cells = []
        diagnostics: list[dict[str, Any]] = []
        for write in writes:
            coordinate = _write_coordinate(sheet, resolved, columns, contract, write)
            field_name = str(write.get("field") or "")
            field = field_by_name.get(field_name) or fields[0]
            rendered, cell_type, value_diags = _diagnose_output_value(write.get("value"), field)
            for diag in value_diags:
                diagnostics.append({"cell": coordinate, "field": field["name"], **diag})
            current_hash = cell_hash(sheet[coordinate])
            baseline_hash = baseline.get(coordinate)
            conflict = bool(baseline_hash and baseline_hash != current_hash)
            cells.append(
                {
                    "cell": coordinate,
                    "field": field["name"],
                    "old_value": sheet[coordinate].value,
                    "new_value": rendered,
                    "cell_type": cell_type,
                    "baseline_hash": baseline_hash,
                    "current_hash": current_hash,
                    "conflict": conflict,
                }
            )
        features = _unsupported_features(Path(path))
        return {
            "binding_id": contract.get("id"),
            "workbook_sha256": workbook_hash(path),
            "target": {"sheet": resolved.sheet, "range": resolved.ref},
            "cells": cells,
            "conflicts": [cell for cell in cells if cell["conflict"]],
            "diagnostics": diagnostics,
            "unsupported_features": features,
            "copy_rewrite_supported": not features,
            "requires_independent_result_workbook": bool(features),
        }
    finally:
        book.close()


def _apply_cell(cell: Cell, value: Any, cell_type: str) -> None:
    if cell_type == "number":
        put_number(cell, str(value))
        return
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"
        if FORMULA_LIKE.match(value):
            cell.number_format = "@"


def _independent_result(destination: Path, patch: dict[str, Any], resolutions: dict[str, str]) -> dict[str, Any]:
    book = Workbook()
    sheet = book.active
    sheet.title = "Results"
    sheet.append(["cell", "field", "value", "resolution"])
    applied = []
    kept = []
    for entry in patch["cells"]:
        choice = resolutions.get(entry["cell"], "apply")
        if entry.get("conflict") and entry["cell"] not in resolutions:
            raise FoundryError(f"conflict at {entry['cell']} requires explicit resolution")
        if choice == "keep":
            kept.append(entry)
            continue
        sheet.append([entry["cell"], entry["field"], entry["new_value"], choice])
        _apply_cell(sheet.cell(sheet.max_row, 3), entry["new_value"], entry["cell_type"])
        for column in (1, 2, 4):
            if isinstance(sheet.cell(sheet.max_row, column).value, str):
                sheet.cell(sheet.max_row, column).data_type = "s"
        applied.append(entry)
    destination.parent.mkdir(parents=True, exist_ok=True)
    book.save(destination)
    book.close()
    return {
        "path": str(destination),
        "sha256": workbook_hash(destination),
        "mode": "independent_result_workbook",
        "applied": applied,
        "kept": kept,
        "diagnostics": [
            {
                "code": "independent_result_workbook",
                "message": "source workbook has unsupported relationships for automatic rewrite",
            }
        ],
    }


def apply_output_patch(
    path: str | Path,
    destination: str | Path,
    patch: dict[str, Any],
    resolutions: dict[str, str] | None = None,
) -> dict[str, Any]:
    source = Path(path)
    target = Path(destination)
    if source.resolve() == target.resolve():
        raise FoundryError("apply_output_patch writes to a copy, not the source workbook")
    if target.exists():
        raise FoundryError("output workbook already exists; choose a new copy")
    if workbook_hash(source) != patch.get("workbook_sha256"):
        raise FoundryError("workbook changed since preview; preview the current cells again")
    min_col, min_row, max_col, max_row = range_boundaries(patch["target"]["range"])
    allowed = ResolvedRange(patch["target"]["sheet"], min_col, min_row, max_col, max_row)
    _check_bounds(allowed)
    for entry in patch["cells"]:
        if not allowed.contains(entry["cell"]):
            raise FoundryError("patch cell is outside declared output range")
    resolutions = dict(resolutions or {})
    for value in resolutions.values():
        if value not in {"apply", "keep"}:
            raise FoundryError("cell resolution must be apply or keep")
    if patch.get("unsupported_features"):
        return _independent_result(target, patch, resolutions)
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, target)
    book = _open(target)
    try:
        sheet_name = patch["target"]["sheet"]
        if sheet_name not in book.sheetnames:
            raise FoundryError(f"no worksheet: {sheet_name}")
        sheet = book[sheet_name]
        applied = []
        kept = []
        conflicts = []
        for entry in patch["cells"]:
            coordinate = entry["cell"]
            current_hash = cell_hash(sheet[coordinate])
            changed_since_baseline = bool(
                entry.get("baseline_hash") and current_hash != entry["baseline_hash"]
            )
            if changed_since_baseline:
                conflicts.append({**entry, "copy_hash": current_hash})
            choice = resolutions.get(coordinate, "apply")
            if choice == "keep":
                kept.append(entry)
                continue
            if changed_since_baseline and coordinate not in resolutions:
                raise FoundryError(f"conflict at {coordinate} requires explicit resolution")
            _apply_cell(sheet[coordinate], entry["new_value"], entry["cell_type"])
            applied.append(entry)
        book.save(target)
        return {
            "path": str(target),
            "sha256": workbook_hash(target),
            "mode": "copy_patch",
            "applied": applied,
            "kept": kept,
            "conflicts": conflicts,
            "diagnostics": patch.get("diagnostics") or [],
        }
    finally:
        book.close()
