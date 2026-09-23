"""Bounded read-only previews for archived workbook snapshots."""

from __future__ import annotations

import re
import io
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from openpyxl.utils import get_column_letter, range_boundaries

from foundry.analysis_data import _open_xlsx_archive, read_exact_numerics
from foundry.errors import FoundryError
from foundry.workbook_bindings import open_workbook_bytes

DEFAULT_ROWS = 40
DEFAULT_COLUMNS = 12
MAX_ROWS = 200
MAX_COLUMNS = 40

_CELL_REF = re.compile(r"^\$?[A-Za-z]{1,3}\$?\d+$")
_CELL_RANGE = re.compile(r"^\$?[A-Za-z]{1,3}\$?\d+:\$?[A-Za-z]{1,3}\$?\d+$")


def preview_workbook_range(
    source: Path,
    *,
    dataset_id: str,
    source_hash: str,
    source_name: str,
    sheet_name: str,
    range_ref: str | None = None,
) -> dict[str, Any]:
    data = source.read_bytes()
    formula_book = open_workbook_bytes(data, data_only=False)
    value_book = open_workbook_bytes(data, data_only=True)
    try:
        if sheet_name not in formula_book.sheetnames:
            raise FoundryError(f"no worksheet: {sheet_name}")
        if sheet_name not in value_book.sheetnames:
            raise FoundryError(f"no worksheet: {sheet_name}")
        sheet = formula_book[sheet_name]
        value_sheet = value_book[sheet_name]
        bounds = _resolve_range(sheet, range_ref)
        exact = _preview_exact_numerics(data, sheet_name)
        merged = {
            cell.coordinate
            for merged_range in sheet.merged_cells.ranges
            for row in sheet[merged_range.coord]
            for cell in row
        }
        cells = []
        for row in range(bounds["min_row"], bounds["max_row"] + 1):
            for col in range(bounds["min_col"], bounds["max_col"] + 1):
                cell = sheet.cell(row=row, column=col)
                cached = value_sheet.cell(row=row, column=col).value
                cells.append(_cell_preview(cell, cached, exact, cell.coordinate in merged))
        return {
            "dataset_id": dataset_id,
            "source_hash": source_hash,
            "source_name": source_name,
            "sheet": sheet_name,
            "range": _range_text(bounds),
            "bounds": bounds,
            "header_context": _header_context(sheet),
            "limits": {
                "default_rows": DEFAULT_ROWS,
                "default_columns": DEFAULT_COLUMNS,
                "max_rows": MAX_ROWS,
                "max_columns": MAX_COLUMNS,
            },
            "cells": cells,
            "warnings": [],
        }
    finally:
        formula_book.close()
        value_book.close()


def _header_context(sheet: Any) -> list[dict[str, str]]:
    """Show a few literal title cells beside a detached range, without inferring scope."""
    context: list[dict[str, str]] = []
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 1, 3), max_col=min(sheet.max_column or 1, 8)):
        for cell in row:
            if cell.data_type == "f" or not isinstance(cell.value, str):
                continue
            value = cell.value.strip()
            if not value:
                continue
            context.append({"address": cell.coordinate, "text": value[:200]})
            if len(context) == 4:
                return context
    return context


def _resolve_range(sheet: Any, range_ref: str | None) -> dict[str, int]:
    if range_ref is None or not str(range_ref).strip():
        max_row = min(max(sheet.max_row or 1, 1), DEFAULT_ROWS)
        max_col = min(max(sheet.max_column or 1, 1), DEFAULT_COLUMNS)
        return {"min_row": 1, "min_col": 1, "max_row": max_row, "max_col": max_col}
    text = str(range_ref).strip().replace("$", "")
    if "!" in text:
        raise FoundryError("range preview takes sheet and range separately; external references are not allowed")
    if "," in text or " " in text:
        raise FoundryError("range preview requires a single bounded range")
    if _CELL_REF.match(text):
        text = f"{text}:{text}"
    elif not _CELL_RANGE.match(text):
        raise FoundryError("range preview requires a bounded A1 rectangle")
    try:
        min_col, min_row, max_col, max_row = range_boundaries(text)
    except ValueError as exc:
        raise FoundryError("invalid range preview range") from exc
    if min_col < 1 or min_row < 1 or min_col > max_col or min_row > max_row:
        raise FoundryError("invalid range preview range")
    height = max_row - min_row + 1
    width = max_col - min_col + 1
    if height > MAX_ROWS or width > MAX_COLUMNS:
        raise FoundryError("range preview is too large")
    if max_row > (sheet.max_row or 1) or max_col > (sheet.max_column or 1):
        raise FoundryError("range preview is outside worksheet bounds")
    return {"min_row": min_row, "min_col": min_col, "max_row": max_row, "max_col": max_col}


def _preview_exact_numerics(data: bytes, sheet_name: str) -> dict[str, Decimal]:
    try:
        return read_exact_numerics(data, sheet_name)
    except FoundryError as exc:
        if "formula has no cached value" not in str(exc):
            raise
        return _read_exact_numerics_without_formula_cache(data, sheet_name)


def _read_exact_numerics_without_formula_cache(data: bytes, sheet_name: str) -> dict[str, Decimal]:
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with _open_xlsx_archive(data) as archive:
        workbook_xml = ET.fromstring(archive.read("xl/workbook.xml"))
        sheets = workbook_xml.find("s:sheets", ns)
        try:
            rid = next(
                s.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
                for s in sheets
                if s.attrib["name"] == sheet_name
            )
        except StopIteration:
            raise FoundryError(f"no worksheet: {sheet_name}") from None
        rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        target = next(r.attrib["Target"] for r in rels if r.attrib["Id"] == rid)
        member = target.lstrip("/") if target.startswith("/") else "xl/" + target
        raw_cells: dict[str, Decimal] = {}
        for _, cell in ET.iterparse(io.BytesIO(archive.read(member)), events=("end",)):
            if cell.tag != "{" + ns["s"] + "}c":
                continue
            value = cell.find("s:v", ns)
            formula = cell.find("s:f", ns)
            if formula is None and cell.attrib.get("t", "n") == "n" and value is not None and value.text:
                raw_cells[cell.attrib["r"]] = Decimal(value.text)
            cell.clear()
        return raw_cells


def _range_text(bounds: dict[str, int]) -> str:
    start = f"{get_column_letter(bounds['min_col'])}{bounds['min_row']}"
    end = f"{get_column_letter(bounds['max_col'])}{bounds['max_row']}"
    return f"{start}:{end}"


def _cell_preview(cell: Any, cached: Any, exact: dict[str, Decimal], merged: bool) -> dict[str, Any]:
    formula = str(cell.value) if cell.data_type == "f" else None
    exact_value = exact.get(cell.coordinate)
    value = None if formula else (exact_value if exact_value is not None else cell.value)
    display = _display(value, cell.number_format)
    cached_available = formula is not None and cached not in (None, "")
    return {
        "address": cell.coordinate,
        "row": cell.row,
        "column": cell.column,
        "column_letter": get_column_letter(cell.column),
        "kind": _kind(cell, value),
        "display": display,
        "exact": format(exact_value, "f") if exact_value is not None else None,
        "formula": formula,
        "cached_value_available": cached_available,
        "cached_value": _display(cached, cell.number_format) if cached_available else None,
        "number_format": cell.number_format,
        "merged": merged,
    }


def _kind(cell: Any, value: Any) -> str:
    if cell.data_type == "f":
        return "formula"
    if value is None:
        return "blank"
    if cell.data_type == "b" or isinstance(value, bool):
        return "boolean"
    if cell.data_type == "e":
        return "error"
    if getattr(cell, "is_date", False):
        return "date"
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return "number"
    return "text"


def _display(value: Any, number_format: str = "General") -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return _display_decimal(value, number_format)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return _display_decimal(Decimal(str(value)), number_format)
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _display_decimal(value: Decimal, number_format: str) -> str:
    """Human display for common Excel numeric formats; exact XML stays separate."""
    fmt = (number_format or "General").split(";")[0]
    if fmt == "General" or "General" in fmt:
        return format(value, "f")
    percent = "%" in fmt
    shown = value * Decimal("100") if percent else value
    decimal_places = _decimal_places(fmt)
    if decimal_places is not None:
        quant = Decimal(1).scaleb(-decimal_places)
        shown = shown.quantize(quant, rounding=ROUND_HALF_UP)
    use_grouping = "," in fmt.split(".")[0]
    spec = f"{',' if use_grouping else ''}.{decimal_places}f" if decimal_places is not None else f"{',' if use_grouping else ''}f"
    text = format(shown, spec)
    return text + ("%" if percent else "")


def _decimal_places(fmt: str) -> int | None:
    if "." not in fmt:
        if any(token in fmt for token in ("0", "#")):
            return 0
        return None
    tail = fmt.split(".", 1)[1]
    count = 0
    for char in tail:
        if char in {"0", "#"}:
            count += 1
            continue
        break
    return count
