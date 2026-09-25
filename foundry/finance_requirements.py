"""Versioned task obligations, independent of the draft's chosen executions."""
from __future__ import annotations

import re
from decimal import Decimal, InvalidOperation

from foundry.analysis import _new, _safe_text, _hash
from foundry.errors import FoundryError
from foundry.locking import exclusive

ITEM_ID = re.compile(r'^[A-Za-z0-9_\-]{1,80}$')
COMPARISON_COMMON_SCOPE_KEYS = {
    'tenant_id', 'ledger', 'ledger_id', 'organization', 'organization_id', 'currency', 'unit'
}
SCOPE_CONTROL_KEYS = {
    'actual', 'actual_scope', 'baseline', 'baseline_scope', 'budget', 'budget_scope',
    'source', 'source_scope', 'comparison_scope',
    'baseline_period', 'comparison_period', 'period_baseline', 'baseline_scenario',
}
METRIC_LABELS = {
    'profit': (
        'profit', 'management_profit', 'operating_profit', 'gross_profit', 'net_profit',
        'operating_profit_ex_other_income', '利润', '管理利润', '管理利润合计', '营业利润', '毛利', '净利润',
    ),
    'revenue': ('revenue', 'operating_revenue', '收入', '营业收入'),
    'cost': ('cost', 'operating_cost', '营业成本', '成本'),
    'expense': ('expense', 'period_expense', 'total_expense', '费用', '期间费用', '总费用'),
    'ratio': ('ratio', 'rate', 'achievement_ratio', 'attainment_rate', '达成率', '比率'),
    'other_income': ('other_income', 'other_gains', 'other_income_and_gains_4items', '其他收益', '其他损益'),
}
RATIO_SCENARIOS = {
    'achievement_ratio', 'achievement_ratio_estimate',
    'attainment_rate', 'attainment_ratio',
    'completion_rate',
    'actual_budget_ratio', 'actual_to_budget',
    'budget_phase_ratio',
    'ratio',
}
PRICE_QUANTITY_METHODS = {'quantity_then_price', 'base_weights_with_interaction'}
PRICE_QUANTITY_REQUIRED_COLUMNS = (
    'total_variance', 'quantity_contribution', 'price_contribution', 'interaction'
)
RESULT_SET_METRICS = {'control_totals'}
INTENT_DIMENSION_KEYS = {'dimension', 'scenario_compare'}
# The P&L lines whose amounts need engine closure (or the engine's own account total, or a source recipe);
# the finance engine's plain output of these lines does not count (run_metric management_pnl rows never close).
CLOSED_EVIDENCE_METRICS = {'profit', 'revenue', 'cost', 'expense'}
# Amounts that only engine closure or a source recipe answer: the four P&L lines and the price-quantity
# bridge (price_quantity, price_volume …), which only the finance engine's explain_variance bridge closes.
# Engine or aggregate output of these names is not their answer.
CLOSURE_OR_RECIPE_AMOUNTS = CLOSED_EVIDENCE_METRICS | {'price_quantity'}
# The amount gate is inverted: every item is an amount unless it is a known non-amount (below). The item
# metric is free text, so these stems and tokens only pick the P&L class used in guidance (利润总额 → profit,
# NetProfit → profit); a name they miss is still gated, as the plain class 'amount'.
AMOUNT_LABEL_STEMS = (
    ('profit', ('利润', '毛利', '净利', '净收益', '盈利', '亏损')),
    ('revenue', ('收入', '营收', '营业额', '销售额')),
    ('cost', ('成本',)),
    # 管理费/销售费/财务费/研发费 are the same expense lines written without 用.
    ('expense', ('费用', '管理费', '销售费', '财务费', '研发费')),
)
AMOUNT_LABEL_TOKENS = {
    'profit': 'profit', 'profits': 'profit', 'ebit': 'profit', 'ebitda': 'profit', 'earnings': 'profit',
    'revenue': 'revenue', 'revenues': 'revenue', 'sales': 'revenue',
    'cost': 'cost', 'costs': 'cost', 'cogs': 'cost',
    'expense': 'expense', 'expenses': 'expense', 'opex': 'expense',
}
AMOUNT_LABEL_TOKEN_PAIRS = {
    ('net', 'income'): 'profit', ('operating', 'income'): 'profit',
    ('gross', 'income'): 'profit', ('pretax', 'income'): 'profit',
}
# Known non-amounts, the only items the gate exempts: ratios (decision 4 keeps their phase-1 script rule),
# counts and physical quantities. They are recognised by the head of the name (its last word), so 汇率损益
# or business_units_revenue stay amounts; an unrecognised name is never exempt.
NON_AMOUNT_METRICS = frozenset({'ratio'})
RATIO_LABEL_SUFFIXES = ('率', '占比', '比例', '比重', '比率', '倍数', '%', '％')
QUANTITY_LABEL_SUFFIXES = ('数量', '销量', '销售量', '件数', '台数', '人数', '户数', '次数', '个数', '天数', '面积')
RATIO_HEAD_TOKENS = frozenset({'ratio', 'ratios', 'rate', 'rates', 'pct', 'percent', 'percentage'})
QUANTITY_HEAD_TOKENS = frozenset({
    'volume', 'volumes', 'quantity', 'quantities', 'qty', 'units', 'count', 'counts', 'headcount',
    'days', 'area', 'sqm',
})
# A count, quantity or area head exempts a name only when nothing else in it names an amount: price_quantity,
# sales_price_volume and 价量 are price-volume bridges, revenue_by_area and 收入按面积 split an amount by area.
# Only `sales` right before the head is a qualifier of the quantity (sales_volume, sales_units).
AMOUNT_MARKER_WORDS = frozenset({'price', 'prices', 'pricing', 'amount', 'amounts', 'value', 'values'})
AMOUNT_MARKER_STEMS = ('价', '金额')
QUANTITY_QUALIFIERS = frozenset({'sales'})
# Amounts outside the four P&L lines that a finance-engine run states itself (see NATIVE_METRICS).
ENGINE_OUTPUT_AMOUNTS = {
    'open_amount': {'tool': 'finance_scenario_run', 'scenario_kind': 'receivables_aging', 'column': 'open_amount'},
    'closing_cash': {'tool': 'finance_scenario_run', 'scenario_kind': 'cash_13_weeks', 'column': 'closing_cash'},
    **{column: {'tool': 'explain_variance', 'column': column} for column in (
        'amount_change', 'total_variance', 'quantity_contribution', 'price_contribution', 'interaction', 'residual')},
}
PROFIT_BRIDGE_REQUIRED_LINES = ('revenue', 'cost', 'expense', 'profit')
CONTROL_TOTAL_COMPARISON_GRAIN_KEYS = (
    'ledger', 'ledger_id', 'organization', 'organization_id', 'period', 'currency', 'unit',
    'management_account', 'account_id', 'account_code',
)
# Hierarchy levels say how far to drill, not which result row answers; a row cannot carry them.
HIERARCHY_DIMENSION_KEYS = frozenset({
    'drilldown', 'drill_down', 'breakdown', 'granularity', 'level', 'group_by', 'hierarchy',
})
# Book identity a run without finance_scope proves by its bound snapshot hashes, not by row text.
LEDGER_IDENTITY_KEYS = frozenset({'ledger', 'ledger_id', 'tenant_id'})
# A ratio has no money unit or currency; the shared scope's unit/currency only constrain amounts.
AMOUNT_ONLY_SCOPE_KEYS = frozenset({'unit', 'currency'})
GRAIN_KEYS = ('organization', 'period', 'scenario')
GRAIN_LISTS = (('organization', 'organizations'), ('period', 'periods'), ('scenario', 'scenarios'))
GRAIN_BASIS_KINDS = ('estimate', 'unavailable')
# Evidence classes that answer an amount; an estimate only ever reaches `estimated`.
SOURCE_EVIDENCE_BASES = frozenset({'engine_closed', 'source_value', 'source_calculation'})
# Blockers only the manager can clear; everything else stays a repairable gap.
MANAGER_BLOCKERS = frozenset({'structure_pending_confirmation'})


def _canonical_metric_name(metric):
    from foundry.finance_requirement_support import canonical_metric
    return canonical_metric(metric)


def _canonical_method_name(method):
    from foundry.finance_requirement_support import canonical_method
    return canonical_method(method)


def _assess_requirement_support(analysis, spec):
    from foundry.finance_requirement_support import assess_requirement_support
    return assess_requirement_support(analysis, spec)


def latest_requirements(analysis):
    records = analysis._list('requirements')
    return records[-1] if records else None


def save_requirements(analysis, spec):
    prepared = prepare_requirements(analysis, spec)
    question, scope, bindings = prepared['question'], prepared['scope'], prepared['bindings']
    normalized, selection = prepared['items'], prepared['selection']
    # Read before taking the task lock: scanning workbook text is slow and only ever warns.
    warnings = [*prepared['warnings'], *_organization_warnings(analysis, normalized, bindings)]
    with exclusive(analysis.task.lock_path, timeout=10):
        parent, removed, reason = _version_change(analysis, spec, normalized)
        assessed_spec = {**spec, 'question': question, 'scope': scope, 'items': normalized, 'bindings': bindings}
        from foundry.finance_requirement_semantics import requirement_semantics
        semantics = requirement_semantics(analysis, {**assessed_spec, 'removed_item_ids': removed})
        record = _new(analysis.task, **semantics, question=question, scope=scope, bindings=bindings, context_id=spec.get("context_id"),
                      items=normalized, parent_id=(parent or {}).get('id'),
                      removed_item_ids=removed, change_reason=reason, selection=selection,
                      support=_assess_requirement_support(analysis, assessed_spec),
                      **({'authoring_warnings': warnings} if warnings else {}),
                      content_hash=_hash({'question': question, 'scope': scope, 'items': normalized, 'bindings': bindings}))
        return analysis._write_locked(analysis._object_path('requirements', record['id']), record)


def prepare_requirements(analysis, spec):
    """Validate and normalize a requirements spec without writing anything; raises what save would refuse.

    circle-select calls it before writing its child context, so a refused save leaves nothing behind.
    The version checks run again under the task lock when the record is written.
    """
    if not isinstance(spec, dict):
        raise FoundryError('analysis requirements must be an object')
    question = _safe_text(spec.get('question'), 'requirements question', required=True, maximum=10000)
    scope = spec.get('scope', {})
    items = spec.get('items')
    bindings = spec.get('bindings', {})
    selection = _selection(spec.get('selection'))
    if not isinstance(scope, dict) or not isinstance(bindings, dict):
        raise FoundryError('requirements scope and bindings must be objects')
    if not isinstance(items, list) or len(items) > 100:
        raise FoundryError('requirements items must be an array of at most 100 items')
    normalized = []
    seen = set()
    for item in items:
        if not isinstance(item, dict) or not isinstance(item.get('id'), str) or not ITEM_ID.fullmatch(item['id']):
            raise FoundryError('requirement item requires a stable id')
        if item['id'] in seen:
            raise FoundryError('requirement item ids must be unique')
        seen.add(item['id'])
        metric = _safe_text(item.get('metric'), 'requirement metric', required=True, maximum=200)
        dimensions = item.get('dimensions', {})
        periods = item.get('periods', [])
        organizations = item.get('organizations', [])
        scenarios = item.get('scenarios', [])
        if not isinstance(dimensions, dict) or not isinstance(item.get('requires_drilldown', False), bool):
            raise FoundryError('invalid requirement dimensions or drilldown flag')
        if isinstance(dimensions.get('period'), (list, tuple, set, frozenset, dict)):
            raise FoundryError('requirement dimensions.period must be a scalar; use periods for a set')
        if periods in (None, ''):
            periods = []
        if (
            not isinstance(periods, list)
            or len(periods) > 120
            or any(not isinstance(period, str) or not period.strip() for period in periods)
        ):
            raise FoundryError('requirement periods must be an array of non-empty strings')
        periods = list(dict.fromkeys(period.strip() for period in periods))
        if periods and scope.get('period') not in (None, '') and periods != [scope['period']]:
            raise FoundryError('requirement periods conflict with scope.period; set common period only when every item shares it')
        if organizations in (None, ''):
            organizations = []
        if (
            not isinstance(organizations, list)
            or len(organizations) > 100
            or any(not isinstance(organization, str) or not organization.strip() for organization in organizations)
        ):
            raise FoundryError('requirement organizations must be an array of non-empty strings')
        organizations = list(dict.fromkeys(organization.strip() for organization in organizations))
        if organizations and scope.get('organization') not in (None, ''):
            if organizations != [scope['organization']]:
                raise FoundryError('requirement organizations conflict with scope.organization; set common organization only when every item shares it')
        if scenarios in (None, ''):
            scenarios = []
        if (
            not isinstance(scenarios, list)
            or len(scenarios) > 100
            or any(not isinstance(scenario, str) or not scenario.strip() for scenario in scenarios)
        ):
            raise FoundryError('requirement scenarios must be an array of non-empty strings')
        scenarios = list(dict.fromkeys(scenario.strip() for scenario in scenarios))
        if scenarios and scope.get('scenario') not in (None, ''):
            if scenarios != [scope['scenario']]:
                raise FoundryError('requirement scenarios conflict with scope.scenario; set common scenario only when every item shares it')
        from foundry.finance_requirement_support import canonical_dimensions
        normalized_item = {'id': item['id'], 'metric': metric,
                           'label': _safe_text(item.get('label', metric), 'requirement label', maximum=500),
                           'dimensions': canonical_dimensions(dimensions),
                           'requires_drilldown': item.get('requires_drilldown', False)}
        if periods:
            normalized_item['periods'] = periods
        if organizations:
            normalized_item['organizations'] = organizations
        if scenarios:
            normalized_item['scenarios'] = scenarios
        grain_basis = _grain_basis(item.get('grain_basis'), normalized_item)
        if grain_basis:
            normalized_item['grain_basis'] = grain_basis
        normalized.append(normalized_item)
    normalized, carried_over = _drop_carried_over_hierarchy(analysis, spec.get('parent_id'), normalized)
    hierarchy = {item['id']: _hierarchy_dimensions(item) for item in normalized if _hierarchy_dimensions(item)}
    if hierarchy:
        _reject_hierarchy_dimensions(normalized, hierarchy)
    scope, scope_dropped = _scope_hierarchy(analysis, spec.get('parent_id'), scope)
    _version_change(analysis, spec, normalized)
    return {'question': question, 'scope': scope, 'items': normalized, 'bindings': bindings,
            'selection': selection, 'warnings': [*carried_over, *scope_dropped]}


def _version_change(analysis, spec, items):
    """The parent this save must name and the item ids it removes; save calls it again under the task lock."""
    parent = latest_requirements(analysis)
    if parent and spec.get('parent_id') != parent['id']:
        raise FoundryError('requirements update must reference the latest version')
    if not parent and spec.get('parent_id'):
        raise FoundryError('requirements parent does not exist')
    newly_removed = {item['id'] for item in (parent or {}).get('items', [])} - {item['id'] for item in items}
    removed = sorted(set((parent or {}).get('removed_item_ids', [])) | newly_removed)
    reason = _safe_text(spec.get('change_reason', ''), 'requirements change reason', maximum=2000)
    if newly_removed and not reason.strip():
        raise FoundryError('removing requirements requires a change reason')
    return parent, removed, reason


def _hierarchy_dimensions(item):
    dimensions = item.get('dimensions') if isinstance(item.get('dimensions'), dict) else {}
    return {key: value for key, value in dimensions.items() if str(key).strip().lower() in HIERARCHY_DIMENSION_KEYS}


def _without_hierarchy(item):
    keys = _hierarchy_dimensions(item)
    return {**item, 'dimensions': {key: value for key, value in item.get('dimensions', {}).items() if key not in keys}}


def _parent_record(analysis, parent_id):
    if not parent_id:
        return None
    try:
        return analysis._get('requirements', parent_id)
    except FoundryError:
        return None  # the version check under the task lock reports a missing parent


def _drop_carried_over_hierarchy(analysis, parent_id, items):
    """Drop a hierarchy level an item carries over unchanged from its parent record (saved before phase 1),
    and record every legacy level the new version no longer has.

    The scope card carries items over unchanged and cannot edit them; circle-select and a resubmitted
    refusal's suggested_items arrive already without the level. Either way the dropped level is recorded
    as an authoring warning on the new version. A key the writer adds or changes is still refused.
    """
    parent = _parent_record(analysis, parent_id)
    previous = {item.get('id'): item for item in (parent or {}).get('items') or [] if isinstance(item, dict)}
    if not any(_hierarchy_dimensions(item) for item in previous.values()):
        return items, []
    kept, warnings = [], []
    for item in items:
        keys = _hierarchy_dimensions(item)
        legacy = _hierarchy_dimensions(previous.get(item['id']) or {})
        if legacy and keys in ({}, legacy):
            named = '、'.join(f'{key}={value}' for key, value in legacy.items())
            warnings.append({'code': 'hierarchy_dimension_dropped', 'requirement_item_id': item['id'], 'keys': legacy,
                             'message': f'旧版要求的维度 {named} 是下钻层级，不是结果行字段，新版已去掉；'
                                        '如需按下钻对象逐一作答，请把它们列入 organizations 或另立要求项。'})
            item = _without_hierarchy(item)
        kept.append(item)
    return kept, warnings


# The scope card's 粒度 (scope.granularity) describes the analysis and is kept; it never filters rows
# (_item_scope, _row_scope). The other hierarchy keys in scope are a drilldown moved out of dimensions.
SCOPE_REFUSED_HIERARCHY_KEYS = HIERARCHY_DIMENSION_KEYS - {'granularity'}


def _scope_levels(values):
    return {key: value for key, value in (values or {}).items()
            if str(key).strip().lower() in SCOPE_REFUSED_HIERARCHY_KEYS} if isinstance(values, dict) else {}


def _earlier_scope_levels(analysis, parent_id):
    """Every (key, value) hierarchy level held by the parent version or any version before it.

    A client that loaded a legacy record (the scope card loads once) may resend a level that an earlier save
    already dropped; that is the same legacy level, not a newly written one.
    """
    import json

    seen, held = set(), set()
    record = _parent_record(analysis, parent_id)
    while record and record.get('id') not in seen:
        seen.add(record.get('id'))
        held.update((key, json.dumps(value, sort_keys=True, ensure_ascii=False, default=str))
                    for key, value in _scope_levels(record.get('scope')).items())
        record = _parent_record(analysis, record.get('parent_id'))
    return held


def _scope_hierarchy(analysis, parent_id, scope):
    """Refuse a hierarchy level newly written into scope, as for dimensions; drop and record one carried
    over from the parent version, or resent after an earlier version dropped it."""
    import json

    written = _scope_levels(scope)
    legacy = _scope_levels((_parent_record(analysis, parent_id) or {}).get('scope'))
    earlier = _earlier_scope_levels(analysis, parent_id) if written else set()
    new = {key: value for key, value in written.items()
           if (key, json.dumps(value, sort_keys=True, ensure_ascii=False, default=str)) not in earlier}
    if new:
        suggested = {key: value for key, value in scope.items() if key not in written}
        named = ', '.join(f'{key}={value}' for key, value in new.items())
        error = FoundryError(
            f'requirement scope holds hierarchy levels, not result-row filters ({named}). List each drilldown '
            'target in organizations or save it as its own item; use requires_drilldown for per-cell source '
            'tracing. suggested_scope=' + json.dumps(suggested, ensure_ascii=False)
        )
        error.details = {'code': 'requirement_scope_not_row_filter', 'keys': sorted(new), 'suggested_scope': suggested}
        raise error
    dropped = {**legacy, **written}
    if not dropped:
        return scope, []
    named = '、'.join(f'{key}={value}' for key, value in dropped.items())
    return {key: value for key, value in scope.items() if key not in written}, [{
        'code': 'hierarchy_scope_dropped', 'keys': dropped,
        'message': f'旧版要求范围里的 {named} 是下钻层级，不是结果行字段，新版已去掉；'
                   '如需按下钻对象逐一作答，请把它们列入 organizations 或另立要求项。'}]


def _reject_hierarchy_dimensions(items, hierarchy):
    """A hierarchy level is not a result-row filter: stamping it on rows proved nothing (PR #17).

    The refusal carries the same items without those keys, which can be resubmitted unchanged; the
    drilldown targets themselves belong in organizations or in their own item.
    """
    import json

    suggested = [_without_hierarchy(item) for item in items]
    named = '; '.join(f'{item_id}: ' + ', '.join(f'{key}={value}' for key, value in keys.items())
                      for item_id, keys in hierarchy.items())
    error = FoundryError(
        'requirement dimensions hold hierarchy levels, not result-row filters (' + named + '). '
        'List each drilldown target in organizations or save it as its own item; use requires_drilldown '
        'for per-cell source tracing. suggested_items=' + json.dumps(suggested, ensure_ascii=False)
    )
    error.details = {'code': 'requirement_dimension_not_row_filter',
                     'items': [{'requirement_item_id': item_id, 'keys': sorted(keys)} for item_id, keys in hierarchy.items()],
                     'suggested_items': suggested}
    raise error


def _grain_basis(value, item):
    """Declared non-source grains: `estimate` (with method) or `unavailable` (with reason).

    Only a declared estimate grain can reach `estimated`; nothing declared here answers a grain.
    """
    if value in (None, '', []):
        return []
    if not isinstance(value, list) or len(value) > 100:
        raise FoundryError('requirement grain_basis must be an array of at most 100 declarations')
    declarations = []
    for entry in value:
        if not isinstance(entry, dict):
            raise FoundryError('requirement grain_basis entries must be objects')
        basis = entry.get('basis')
        if basis not in GRAIN_BASIS_KINDS:
            raise FoundryError('requirement grain_basis basis must be estimate or unavailable')
        declaration = {'basis': basis}
        for key, field in GRAIN_LISTS:
            members = entry.get(field)
            if members in (None, []):
                continue
            if (not isinstance(members, list) or any(not isinstance(member, str) or not member.strip() for member in members)):
                raise FoundryError(f'requirement grain_basis {field} must be an array of non-empty strings')
            members = list(dict.fromkeys(member.strip() for member in members))
            if not item.get(field) or not set(members) <= set(item[field]):
                raise FoundryError(f'requirement grain_basis {field} must be a subset of the item {field}')
            declaration[field] = members
        if basis == 'estimate':
            declaration['method'] = _safe_text(entry.get('method'), 'grain_basis estimate method', required=True, maximum=1000)
            reproducible = entry.get('factor_reproducible', True)
            if not isinstance(reproducible, bool):
                raise FoundryError('requirement grain_basis factor_reproducible must be a boolean')
            if not reproducible:
                declaration['factor_reproducible'] = False
                declaration['reason'] = _safe_text(entry.get('reason'), 'grain_basis reason', required=True, maximum=1000)
        else:
            declaration['reason'] = _safe_text(entry.get('reason'), 'grain_basis reason', required=True, maximum=1000)
        declarations.append(declaration)
    grains = _requested_grains(item) or {(None, None, None): []}
    for grain in grains:
        kinds = {(entry['basis'], entry.get('factor_reproducible', True)) for entry in declarations if _declaration_covers(entry, grain)}
        if len(kinds) > 1:
            raise FoundryError('requirement grain_basis declarations conflict for the same grain')
    return declarations


def _declaration_covers(entry, grain):
    return all(not entry.get(field) or value in entry[field] for (_, field), value in zip(GRAIN_LISTS, grain))


def _declared_basis(item, grain):
    return next((entry for entry in item.get('grain_basis') or [] if _declaration_covers(entry, grain)), None)


def _organization_warnings(analysis, items, bindings):
    """Organizations absent from every bound workbook's text cells: a warning, never a refusal."""
    organizations = list(dict.fromkeys(org for item in items for org in item.get('organizations', [])))
    if not organizations or not isinstance(bindings, dict):
        return []
    texts = []
    for dataset_id in dict.fromkeys(bindings.values()):
        try:
            texts.extend(_workbook_texts(analysis, dataset_id))
        except (FoundryError, OSError, ValueError, KeyError, TypeError):
            return []
    if not texts:
        return []
    return [{'code': 'organization_not_found_in_sources', 'organization': org,
             'message': f'绑定资料的文本单元格中没有找到“{org}”；请核对组织名称是否与资料一致。'}
            for org in organizations if not any(org in text for text in texts)]


def _workbook_texts(analysis, dataset_id):
    import io
    from openpyxl import load_workbook

    dataset = analysis.read_dataset(dataset_id)
    if dataset.get('kind') != 'workbook':
        return []
    path = analysis._inside(analysis.root / 'datasets' / dataset_id / dataset['source_file'])
    book = load_workbook(filename=io.BytesIO(path.read_bytes()), read_only=True, data_only=True)
    try:
        return [value.strip() for sheet in book.worksheets for row in sheet.iter_rows(values_only=True)
                for value in row if isinstance(value, str) and value.strip()]
    finally:
        book.close()


SELECTION_KEYS = ('kind', 'draft_id', 'binding_id', 'block_id', 'range')


def _selection(value):
    """The circled table or workbook range a revision came from; never a path."""
    if value is None:
        return None
    if not isinstance(value, dict) or not value:
        raise FoundryError('requirements selection must be an object')
    extra = set(value) - set(SELECTION_KEYS)
    if extra:
        raise FoundryError('requirements selection only records kind, draft_id, binding_id, block_id, range')
    return {key: _safe_text(value[key], f'selection {key}', required=True, maximum=200) for key in SELECTION_KEYS if key in value}


def _same_number(left, right):
    try:
        return Decimal(str(left)).is_finite() and Decimal(str(left)) == Decimal(str(right))
    except (InvalidOperation, ValueError):
        return left == right and left is not None


def _bindings_match(required, run_bindings):
    if not required:
        return True
    if not isinstance(required, dict) or not isinstance(run_bindings, dict):
        return False
    if len(required) == 1:
        key, value = next(iter(required.items()))
        return run_bindings.get(key) == value or value in set(run_bindings.values())
    return all(run_bindings.get(key) == value for key, value in required.items())


def _has_finance_scope(run):
    scopes = run.get('finance_scope')
    return isinstance(scopes, dict) and bool(scopes)


def _item_scope(item, scope):
    """The requirement scope that applies to this item. A hierarchy level (scope.granularity from the scope
    card's 粒度) describes the analysis and never filters rows; a known non-amount (a ratio, a count, a
    quantity) carries no money unit or currency."""
    scope = _scope_without(scope if isinstance(scope, dict) else {}, HIERARCHY_DIMENSION_KEYS)
    if not _known_non_amount(item):
        return scope
    return _scope_without(scope, AMOUNT_ONLY_SCOPE_KEYS)


def _row_scope(scope):
    """The requirement scope rows and run scopes are compared with: hierarchy levels are not row filters."""
    return _scope_without(scope if isinstance(scope, dict) else {}, HIERARCHY_DIMENSION_KEYS)


def _scope_without(scope, keys):
    return {
        key: (_scope_without(value, keys) if isinstance(value, dict) else value)
        for key, value in scope.items() if key not in keys
    }


def _requirement_sources(latest):
    basis = latest.get('semantic_basis') if isinstance(latest, dict) else None
    sources = basis.get('sources') if isinstance(basis, dict) else None
    return sources if isinstance(sources, dict) and sources else {}


def _ledger_identity_proven(run, latest):
    """A run without finance_scope proves its books only by running on the requirement's own snapshots.

    Without a recorded semantic basis there is nothing to compare hashes with, so the identity is
    unproven; row text (a copied ledger name) never stands in for it.
    """
    return bool(
        _requirement_sources(latest)
        and _bindings_match(latest.get('bindings', {}), run.get('bindings', {}))
        and _run_inputs_match_requirement(run, latest)
    )


def _scope_identity_keys(scope):
    expected = _comparison_scopes(scope)
    return {key for side in expected.values() for key in side if key in LEDGER_IDENTITY_KEYS}


def _requirement_scope_matches(run, row, item, latest):
    scope = _item_scope(item, latest.get('scope') or {})
    if _has_finance_scope(run) or not _scope_identity_keys(scope):
        return _scope_matches(run, row, scope)
    if not _ledger_identity_proven(run, latest):
        return False
    return _scope_matches(run, row, _scope_without(scope, LEDGER_IDENTITY_KEYS))


def _scope_matches(run, row, scope):
    if not scope:
        return True
    expected = _comparison_scopes(scope)
    scopes = run.get('finance_scope') or {}
    if isinstance(scopes, dict):
        actual_scope = row.get('actual_scope') if isinstance(row.get('actual_scope'), dict) else scopes.get('actual')
        baseline_scope = (
            row.get('baseline_scope') if isinstance(row.get('baseline_scope'), dict)
            else scopes.get('baseline') if isinstance(scopes.get('baseline'), dict)
            else scopes.get('budget')
        )
        source_scope = (
            row.get('source_scope') if isinstance(row.get('source_scope'), dict)
            else scopes.get('source')
        )
        if isinstance(actual_scope, dict) or isinstance(baseline_scope, dict):
            actual_candidate = _merge_scope_with_row(actual_scope, row, expected['actual'])
            baseline_candidate = _merge_scope_with_row(baseline_scope, row, expected['baseline'])
            baseline_matches = (
                _scope_dict_matches(baseline_candidate, expected['baseline'])
                if expected['baseline'] or isinstance(baseline_scope, dict)
                else True
            )
            return (
                _scope_dict_matches(actual_candidate, expected['actual'])
                and baseline_matches
            )
        if isinstance(source_scope, dict):
            return _scope_dict_matches(_merge_scope_with_row(source_scope, row, expected['source']), expected['source'])
        bound_scopes = [scope for scope in scopes.values() if isinstance(scope, dict)]
        if bound_scopes:
            return _bound_scopes_match(bound_scopes, expected, row)
    return _scope_dict_matches(row, expected['actual']) or (
        isinstance(scopes, dict) and _scope_dict_matches(scopes, expected['source'])
    )


def _comparison_scopes(scope):
    flat = {
        key: value for key, value in scope.items()
        if key not in SCOPE_CONTROL_KEYS and not isinstance(value, dict)
    }
    actual = _clean_scope(scope.get('actual') or scope.get('actual_scope') or {
        key: value for key, value in flat.items() if key != 'scenario'
    })
    baseline = _clean_scope(scope.get('baseline') or scope.get('baseline_scope') or scope.get('budget') or scope.get('budget_scope') or scope.get('comparison_scope') or {
        key: value for key, value in flat.items()
        if key in COMPARISON_COMMON_SCOPE_KEYS
    })
    baseline_period = scope.get('baseline_period', scope.get('comparison_period', scope.get('period_baseline')))
    if baseline_period is not None and 'period' not in baseline:
        baseline = {**baseline, 'period': baseline_period}
    baseline_scenario = scope.get('baseline_scenario')
    if baseline_scenario is not None and 'scenario' not in baseline:
        baseline = {**baseline, 'scenario': baseline_scenario}
    source = _clean_scope(scope.get('source') or scope.get('source_scope') or {
        key: value for key, value in flat.items() if key != 'scenario'
    })
    return {'actual': actual, 'baseline': baseline, 'source': source}


def _clean_scope(scope):
    if not isinstance(scope, dict):
        return {}
    return {
        key: value for key, value in scope.items()
        if value not in (None, '')
    }


def _scope_dict_matches(candidate, scope):
    return isinstance(candidate, dict) and all(candidate.get(k) == value for k, value in scope.items())


def _merge_scope_with_row(scope, row, expected):
    if not isinstance(scope, dict):
        return scope
    merged = dict(scope)
    for key in expected:
        row_value = row.get(key)
        if row_value in (None, ''):
            continue
        if key in merged and merged[key] != row_value:
            return {}
        merged.setdefault(key, row_value)
    return merged


def _bound_scopes_match(candidates, expected, row=None):
    if len(candidates) != 1:
        return False
    candidate = _merge_scope_with_row(candidates[0], row or {}, expected['actual'])
    if not _scope_dict_matches(candidate, expected['actual']):
        return False
    baseline_scope = expected['baseline']
    if baseline_scope and not _scope_dict_matches(candidates[0], baseline_scope):
        return False
    source_scope = expected['source']
    return not source_scope or _scope_dict_matches(candidate, source_scope)


def _row_period(run, row):
    if row.get('period') not in (None, ''):
        return row['period']
    scopes = (run.get('finance_scope') or {})
    if not isinstance(scopes, dict):
        return None
    for side in ('actual', 'source'):
        scope = scopes.get(side)
        if isinstance(scope, dict) and scope.get('period') not in (None, ''):
            return scope['period']
    return None


def _period_matches(item, run, row):
    periods = item.get('periods', [])
    return not periods or _row_period(run, row) in periods


def _row_organization(run, row):
    if row.get('organization') not in (None, ''):
        return row['organization']
    scopes = (run.get('finance_scope') or {})
    if not isinstance(scopes, dict):
        return None
    for side in ('actual', 'source'):
        scope = scopes.get(side)
        if isinstance(scope, dict) and scope.get('organization') not in (None, ''):
            return scope['organization']
    return None


def _organization_matches(item, run, row):
    organizations = item.get('organizations', [])
    return not organizations or _row_organization(run, row) in organizations


def _row_scenario(_run, row):
    return row.get('scenario') if row.get('scenario') not in (None, '') else None


def _scenario_matches(item, run, row):
    scenarios = item.get('scenarios', [])
    return not scenarios or _row_scenario(run, row) in scenarios


def _requested_grain_matches(item, run, row):
    return (
        _period_matches(item, run, row)
        and _organization_matches(item, run, row)
        and _scenario_matches(item, run, row)
    )


def _requested_grains(item):
    if not item.get('periods') and not item.get('organizations') and not item.get('scenarios'):
        return None
    return {
        (organization, period, scenario): []
        for organization in (item.get('organizations') or [None])
        for period in (item.get('periods') or [None])
        for scenario in (item.get('scenarios') or [None])
    }


def _row_grain(item, run, row):
    return (
        _row_organization(run, row) if item.get('organizations') else None,
        _row_period(run, row) if item.get('periods') else None,
        _row_scenario(run, row) if item.get('scenarios') else None,
    )


def _matches(item, row, column, run=None):
    if not _dimensions_match(item, row, run or {}):
        return False
    return _metric_matches(item, row, column)


def _metric_matches(item, row, column):
    """The metric half of _matches: does this cell state the item's metric, whatever its other fields."""
    metric = _canonical_metric_name(item['metric'])
    aliases = METRIC_LABELS.get(metric, (metric,))
    if column == metric or column in {f'{alias}_{suffix}' for alias in aliases for suffix in ('variance','difference','contribution','change','impact')}:
        return True
    amount_columns = {
        'contribution','profit_impact','amount_change','variance','difference','value','amount','open_amount',
    }
    if column not in amount_columns:
        return False
    if _metric_key_ratio_pin_matches(item, row, metric):
        return True
    if _metric_key_matches(row.get('metric_key'), metric, aliases):
        return True
    labels = [str(row.get(key, '')).lower() for key in ('metric', 'line', 'contribution_line', 'management_account')]
    return any(alias.lower() in labels for alias in aliases)


def _metric_key_ratio_pin_matches(item, row, metric):
    if metric != 'ratio':
        return False
    dimensions = item.get('dimensions', {})
    if not isinstance(dimensions, dict) or dimensions.get('metric_key') in (None, ''):
        return False
    scenario = row.get('scenario')
    return isinstance(scenario, str) and scenario.strip().lower() in RATIO_SCENARIOS


def _metric_key_matches(value, metric, aliases):
    if not isinstance(value, str):
        return False
    raw = value.strip()
    if not raw:
        return False
    canonical = _canonical_metric_name(raw)
    candidates = {metric}
    candidates.update(str(alias).strip() for alias in aliases if str(alias).strip())
    candidates.update(_canonical_metric_name(alias) for alias in list(candidates))
    return canonical in candidates or raw.lower() in {str(candidate).lower() for candidate in candidates}


def _run_method(run):
    return _canonical_method_name(run.get('result', {}).get('method') or run.get('method'))


def _dimensions_match(item, row, run=None):
    dimensions = item.get('dimensions', {})
    if not isinstance(dimensions, dict):
        return False
    # Older records may still hold a hierarchy level; a row stamped with it proves nothing.
    if _hierarchy_dimensions(item):
        return False
    method = _canonical_method_name(dimensions.get('method'))
    if method is not None and _run_method(run or {}) != method:
        return False
    for key, value in dimensions.items():
        if key in {'method', *INTENT_DIMENSION_KEYS}:
            continue
        if row.get(key) != value:
            return False
    if 'scenario_compare' in dimensions and not _scenario_compare_matches(dimensions['scenario_compare'], row, run or {}):
        return False
    if 'dimension' in dimensions and not _dimension_intent_matches(dimensions['dimension'], row, run or {}):
        return False
    return True


def _scenario_compare_matches(value, row, run):
    if value != 'actual_vs_budget':
        return False
    result = run.get('result', {}) if isinstance(run, dict) else {}
    if result.get('metric_id') == 'control_totals':
        return _control_totals_actual_budget_pairs_complete(
            [(index, item) for index, item in enumerate(result.get('rows', [])) if isinstance(item, dict)]
        )
    scopes = (run.get('finance_scope') or {}) if isinstance(run, dict) else {}
    actual_scope = row.get('actual_scope') if isinstance(row.get('actual_scope'), dict) else scopes.get('actual')
    baseline_scope = row.get('baseline_scope') if isinstance(row.get('baseline_scope'), dict) else scopes.get('baseline') or scopes.get('budget')
    actual_scenario = actual_scope.get('scenario') if isinstance(actual_scope, dict) else row.get('actual_scenario')
    baseline_scenario = baseline_scope.get('scenario') if isinstance(baseline_scope, dict) else row.get('baseline_scenario')
    return actual_scenario == 'actual' and baseline_scenario == 'budget'


def _dimension_intent_matches(value, row, run):
    if value != 'management_account':
        return False
    if row.get('management_account') is not None:
        return True
    if _run_method(run) == 'revenue_minus_cost_minus_expense' and row.get('metric') in set(METRIC_LABELS):
        return True
    return False


def _run_finance_metric(analysis, run):
    if run.get('engine') != 'finance':
        return None
    result = run.get('result', {})
    if not isinstance(result, dict) or result.get('metric_id') not in {'account_total', 'control_totals'}:
        return None
    analysis_id = run.get('analysis_id')
    if not analysis_id:
        return None
    try:
        revision = analysis.get_analysis(analysis_id)
    except FoundryError:
        return None
    if (
        revision.get('engine') != 'finance'
        or revision.get('operation') != 'metric'
        or revision.get('finance_context_id') != run.get('finance_context_id')
        or revision.get('bindings') != run.get('bindings')
    ):
        return None
    spec = revision.get('finance_spec', {})
    if not isinstance(spec, dict):
        return None
    metric = spec.get('metric')
    if not isinstance(metric, dict) or metric.get('kind') != result.get('metric_id'):
        return None
    return metric


def _prepared_evidence(analysis, runs):
    from foundry.finance_explanations import prepare_explanation

    prepared = prepare_explanation(analysis, list(runs.values())) if runs else {'evidence_groups': []}
    member_map = {}
    group_map = {}
    for group in prepared.get('evidence_groups', []):
        if not isinstance(group, dict):
            continue
        group_key = (group.get('run_id'), group.get('row'))
        if group_key[0] is not None and group_key[1] is not None:
            group_map[group_key] = group
        engine = (runs.get(group.get('run_id')) or {}).get('engine')
        closure_status = (group.get('closure') or {}).get('status')
        for member in group.get('members', []):
            if not isinstance(member, dict):
                continue
            fact_id = member.get('fact_id')
            if not fact_id:
                continue
            member_map[(group.get('run_id'), member.get('row'), member.get('column'))] = {
                'group_id': group.get('id'),
                'group_type': group.get('group_type'),
                'group_status': group.get('status'),
                'closure_status': closure_status,
                # Only the finance engine's own closure is evidence; a script that builds a
                # bridge-shaped table which adds up (999/-1/-1/997) is shown, never counted.
                'engine_closed': engine == 'finance' and closure_status == 'passed',
                'engine_output': engine == 'finance',
                # A server-side declarative aggregate of a bound snapshot (no user code): not a script, and
                # drilldown_fact verifies its rows. It answers amounts outside the four P&L lines.
                'aggregate_output': engine == 'structured_aggregate',
                'member_state': member.get('state'),
                'fact_id': fact_id,
            }
    return {'members': member_map, 'groups': group_map}


def _recipe_support(analysis, spec, latest, report=None):
    """Recipe results the server computed in this call, keyed by the run cell each recipe claims.

    The report comes only from evaluate_evidence_recipes on spec['evidence_recipes'] here, or from an
    internal caller that just ran it (the private `_recipe_report`); an `evidence_report` supplied by a
    caller or stored on an old draft is never read. A result counts only if it passed and every source
    (including a factor's) is a snapshot bound to the requirement; whether the claim cell is cited is
    decided by the caller, which only looks up cited cells.
    """
    if report is None:
        from foundry.financial_evidence import evaluate_evidence_recipes
        try:
            report = evaluate_evidence_recipes(analysis, spec.get('evidence_recipes') or [])
        except FoundryError:
            # A malformed batch is reported by the draft's own recipe check; here it proves nothing.
            report = {'results': []}
    bound = _bound_source_hashes(analysis, latest)
    support = {}
    for result in report.get('results') or []:
        claim = result.get('claim') if isinstance(result, dict) else None
        if not isinstance(claim, dict):
            continue
        entry = _recipe_entry(result, bound)
        key = (claim.get('run_id'), claim.get('row'), claim.get('column'))
        previous = support.get(key)
        if previous is None or _recipe_rank(entry) > _recipe_rank(previous):
            support[key] = {**entry, 'recipe_ids': [*(previous or {}).get('recipe_ids', []), *entry['recipe_ids']]}
        else:
            previous['recipe_ids'] = [*previous['recipe_ids'], *entry['recipe_ids']]
    return support


def _bound_source_hashes(analysis, latest):
    """dataset_id → source_hash for the requirement's bound snapshots (the recorded basis when present)."""
    bound = {}
    sources = _requirement_sources(latest)
    for alias, dataset_id in (latest.get('bindings') or {}).items():
        if not isinstance(dataset_id, str):
            continue
        recorded = (sources.get(alias) or {}).get('source_hash') if isinstance(sources.get(alias), dict) else None
        if recorded is None:
            try:
                recorded = analysis.read_dataset(dataset_id).get('source_hash')
            except (FoundryError, OSError, ValueError, KeyError, TypeError):
                recorded = None
        bound[dataset_id] = recorded
    return bound


def _recipe_entry(result, bound):
    evidence_class = result.get('evidence_class')
    refs = [*(result.get('source_refs') or []), *((result.get('factor') or {}).get('source_refs') or [])]
    inside = bool(refs) and all(
        isinstance(ref, dict) and ref.get('dataset_id') in bound
        and bound[ref['dataset_id']] is not None and ref.get('source_hash') == bound[ref['dataset_id']]
        for ref in refs
    )
    rejected = None
    if result.get('status') != 'pass':
        rejected = 'recipe_failed_precision' if result.get('precision_hint') else 'recipe_failed'
    elif not inside:
        rejected = 'recipe_source_outside_bindings'
    elif evidence_class == 'coefficient_unbound':
        rejected = 'recipe_coefficient_unbound'
    return {
        'recipe_ids': [result.get('recipe_id')],
        'recipe_hash': result.get('recipe_hash'),
        'kind': result.get('kind'),
        'scale': result.get('scale'),
        'evidence_class': evidence_class,
        'counted': rejected is None and evidence_class is not None,
        'rejected_reason': rejected,
        'location': result.get('location'),
        **({'factor': {key: result['factor'][key] for key in ('recipe_id', 'recipe_hash', 'location')}}
           if isinstance(result.get('factor'), dict) else {}),
        **({'precision_hint': result['precision_hint']} if result.get('precision_hint') else {}),
    }


def _recipe_rank(entry):
    if not entry.get('counted'):
        return 0
    if entry.get('kind') == 'ratio':
        return 1
    if entry.get('evidence_class') == 'estimate':
        return 2
    return 3


def _price_quantity_requirement_matches(item, run, row, group):
    method = _canonical_method_name(item.get('dimensions', {}).get('method'))
    run_method = _run_method(run)
    metric = _canonical_metric_name(item.get('metric'))
    if metric == 'amount':
        if method not in PRICE_QUANTITY_METHODS:
            return False
    elif metric in {'revenue', 'price_quantity'}:
        if method is not None and method not in PRICE_QUANTITY_METHODS:
            return False
    else:
        return False
    if run_method not in PRICE_QUANTITY_METHODS:
        return False
    if method is not None and run_method != method:
        return False
    if run.get('engine') != 'finance':
        return False
    if not isinstance(group, dict) or group.get('group_type') != 'price_quantity_bridge':
        return False
    if group.get('status') != 'complete' or (group.get('closure') or {}).get('status') != 'passed':
        return False
    if any(row.get(k) != value for k, value in item.get('dimensions', {}).items() if k != 'method'):
        return False
    members = _price_quantity_required_members(group)
    return {member.get('column') for member in members} == set(PRICE_QUANTITY_REQUIRED_COLUMNS)


def _price_quantity_required_members(group):
    if not isinstance(group, dict):
        return []
    members = [
        member for member in group.get('members', [])
        if isinstance(member, dict) and member.get('column') in PRICE_QUANTITY_REQUIRED_COLUMNS
    ]
    complete = []
    for column in PRICE_QUANTITY_REQUIRED_COLUMNS:
        matches = [member for member in members if member.get('column') == column]
        if len(matches) != 1 or matches[0].get('fact_id') is None:
            return []
        complete.append(matches[0])
    return complete


def _account_total_requirement_matches(analysis, item, run, row, column):
    if column != 'amount':
        return False
    metric = _run_finance_metric(analysis, run)
    if not metric or metric.get('kind') != 'account_total':
        return False
    accounts = metric.get('accounts')
    if not isinstance(accounts, list) or len(accounts) != 1:
        return False
    if str(accounts[0]) != item.get('metric'):
        return False
    return not any(row.get(k) != value for k, value in item.get('dimensions', {}).items())


def _account_amount_claim_requires_native_total(item, row, column):
    return (
        column == 'amount'
        and _canonical_metric_name(item.get('metric')) not in METRIC_LABELS
        and row.get('management_account') == item.get('metric')
    )


def _result_set_requirement_status(analysis, item, latest, runs, claimed, checks, evidence):
    kind = _result_set_kind(item)
    if not kind:
        return None
    bases = []
    for run in runs.values():
        if not _bindings_match(latest.get('bindings', {}), run.get('bindings', {})):
            continue
        metric = _run_finance_metric(analysis, run)
        if not metric or metric.get('kind') != kind:
            continue
        if kind == 'account_total' and not _account_total_metric_matches_item(metric, item):
            continue
        if not _metric_filters_cover_requirement(metric, item, _row_scope(latest['scope'])):
            continue
        if _run_scope_is_narrower_than_requirement(run, item, _row_scope(latest['scope'])):
            continue
        rows = _result_set_required_rows(item, run, _row_scope(latest['scope']))
        if rows and _result_set_comparison_complete(kind, item, latest, run, rows):
            bases.append((run, rows))
    if not bases:
        return {'locations': [], 'status': 'pending'}
    max_rows = max(len(rows) for _, rows in bases)
    for run, rows in bases:
        if len(rows) != max_rows:
            continue
        locations = [
            {'run_id': run['id'], 'row': index, 'column': 'amount', 'output_hash': run.get('output_hash')}
            for index, _ in rows
        ]
        if not all((run['id'], location['row'], 'amount') in claimed for location in locations):
            continue
        if not all(_source_verified(location, checks) for location in locations):
            return {
                'locations': [
                    _result_set_location(location, evidence, 'pending')
                    for location in locations
                ],
                'status': 'pending',
            }
        return {
            'locations': [
                _result_set_location(location, evidence, 'answered')
                for location in locations
            ],
            'status': 'answered',
        }
    return {'locations': [], 'status': 'pending'}


def _result_set_kind(item):
    metric = _canonical_metric_name(item.get('metric'))
    if metric in RESULT_SET_METRICS:
        return metric
    if (item.get('periods') or item.get('organizations')) and metric not in RESULT_SET_METRICS and metric not in METRIC_LABELS:
        return 'account_total'
    return None


def _account_total_metric_matches_item(metric, item):
    accounts = metric.get('accounts')
    return (
        isinstance(accounts, list)
        and len(accounts) == 1
        and str(accounts[0]) == item.get('metric')
    )


def _metric_filters_cover_requirement(metric, item, scope):
    filters = metric.get('filters') or {}
    if not isinstance(filters, dict):
        return False
    expected = _comparison_scopes(scope)
    dimensions = item.get('dimensions', {})
    for key, filtered in filters.items():
        if key == 'organization' and item.get('organizations'):
            required = item['organizations']
        elif key == 'period' and item.get('periods'):
            required = item['periods']
        elif key == 'scenario' and item.get('scenarios'):
            required = item['scenarios']
        elif key == 'management_account' and metric.get('kind') == 'account_total':
            required = item['metric']
        elif key in dimensions and key not in {'method', *INTENT_DIMENSION_KEYS}:
            required = dimensions[key]
        else:
            required = expected['actual'].get(key, expected['source'].get(key))
        if required is None:
            return False
        filtered_values = filtered if isinstance(filtered, (list, tuple, set, frozenset)) else [filtered]
        required_values = required if isinstance(required, (list, tuple, set, frozenset)) else [required]
        try:
            if set(filtered_values) != set(required_values):
                return False
        except TypeError:
            return False
    return True


def _run_scope_is_narrower_than_requirement(run, item, scope):
    scopes = (run.get('finance_scope') or {})
    if not isinstance(scopes, dict):
        return False
    expected = _comparison_scopes(scope)
    periods = set(item.get('periods', []))
    scenarios = set(item.get('scenarios', []))
    for candidate in scopes.values():
        if not isinstance(candidate, dict):
            continue
        for key, value in candidate.items():
            if value in (None, ''):
                continue
            if key in expected['actual'] or key in expected['source']:
                expected_value = expected['actual'].get(key, expected['source'].get(key))
                if expected_value != value:
                    return True
                continue
            if key == 'organization' and item.get('organizations'):
                if set(item['organizations']) == {value}:
                    continue
                return True
            if key == 'scenario' and scenarios:
                if value in scenarios:
                    continue
                return True
            if (
                key == 'scenario'
                and _result_set_kind(item) == 'control_totals'
                and not expected['actual'].get('scenario')
                and not expected['source'].get('scenario')
                and value in {'actual', 'budget'}
            ):
                continue
            if key in item.get('dimensions', {}):
                if item['dimensions'][key] != value:
                    return True
                continue
            if key == 'period' and periods:
                if periods == {value}:
                    continue
                return True
            return True
    return False


def _result_set_required_rows(item, run, scope):
    rows = []
    for index, row in enumerate(run.get('result', {}).get('rows', [])):
        if not _scope_matches(run, row, scope):
            continue
        if not _item_dimensions_match(item, row, run):
            continue
        if not _requested_grain_matches(item, run, row):
            continue
        rows.append((index, row))
    requested_grains = _requested_grains(item)
    if requested_grains is not None and {_row_grain(item, run, row) for _, row in rows} != set(requested_grains):
        return []
    return rows


def _item_dimensions_match(item, row, run=None):
    return _dimensions_match(item, row, run or {})


def _result_set_comparison_complete(kind, item, latest, run, rows):
    if kind != 'control_totals':
        return True
    bindings = latest.get('bindings', {})
    dimensions = item.get('dimensions', {}) if isinstance(item.get('dimensions'), dict) else {}
    needs_actual_budget = (
        {'actual', 'budget'}.issubset(set(bindings))
        or dimensions.get('scenario_compare') == 'actual_vs_budget'
    )
    if not needs_actual_budget:
        return True
    return _control_totals_actual_budget_pairs_complete(rows)


def _control_totals_actual_budget_pairs_complete(rows):
    scenarios_by_grain = {}
    for _, row in rows:
        if not isinstance(row, dict):
            return False
        scenario = row.get('scenario')
        if scenario not in {'actual', 'budget'}:
            return False
        grain = tuple((key, row.get(key)) for key in CONTROL_TOTAL_COMPARISON_GRAIN_KEYS if row.get(key) not in (None, ''))
        if not grain:
            return False
        scenarios_by_grain.setdefault(grain, set()).add(scenario)
    return bool(scenarios_by_grain) and all(scenarios == {'actual', 'budget'} for scenarios in scenarios_by_grain.values())


def _profit_bridge_group_required(item):
    dimensions = item.get('dimensions', {}) if isinstance(item.get('dimensions'), dict) else {}
    return (
        _canonical_metric_name(item.get('metric')) == 'profit'
        and _canonical_method_name(dimensions.get('method')) == 'revenue_minus_cost_minus_expense'
        and dimensions.get('dimension') == 'management_account'
        and dimensions.get('scenario_compare') == 'actual_vs_budget'
    )


def _run_finance_variance_method(analysis, run):
    if run.get('engine') != 'finance':
        return None
    analysis_id = run.get('analysis_id')
    if not analysis_id:
        return None
    try:
        revision = analysis.get_analysis(analysis_id)
    except FoundryError:
        return None
    if (
        revision.get('engine') != 'finance'
        or revision.get('operation') != 'variance'
        or revision.get('finance_context_id') != run.get('finance_context_id')
        or revision.get('bindings') != run.get('bindings')
    ):
        return None
    result = run.get('result', {})
    if not isinstance(result, dict):
        return None
    return _canonical_method_name(result.get('method'))


def _profit_bridge_candidate_locations(analysis, item, run, evidence, scope):
    by_line = _profit_bridge_native_group(analysis, item, run, evidence, scope)
    if by_line is None:
        return []
    return [by_line[line] for line in PROFIT_BRIDGE_REQUIRED_LINES]


def _profit_bridge_locations(analysis, item, run, claimed_by_row, evidence, scope):
    locations = _profit_bridge_candidate_locations(analysis, item, run, evidence, scope)
    if not locations:
        return []
    if any('contribution' not in claimed_by_row.get((run['id'], location['row']), set()) for location in locations):
        return []
    return locations


def _profit_bridge_native_group(analysis, item, run, evidence, scope):
    if not _profit_bridge_group_required(item):
        return None
    method = 'revenue_minus_cost_minus_expense'
    if _run_finance_variance_method(analysis, run) != method:
        return None
    rows = run.get('result', {}).get('rows', [])
    by_line = {}
    group_ids = set()
    for index, row in enumerate(rows):
        if not _scope_matches(run, row, scope) or not _dimensions_match(item, row, run):
            continue
        metric = row.get('metric')
        if metric not in PROFIT_BRIDGE_REQUIRED_LINES:
            continue
        member = evidence['members'].get((run['id'], index, 'contribution'))
        if not member or member.get('group_type') != 'profit_bridge':
            return None
        if member.get('group_status') != 'complete' or member.get('closure_status') != 'passed':
            return None
        group_ids.add(member.get('group_id'))
        by_line[metric] = {'run_id': run['id'], 'row': index, 'column': 'contribution', 'output_hash': run.get('output_hash')}
    if set(by_line) != set(PROFIT_BRIDGE_REQUIRED_LINES) or len(group_ids) != 1:
        return None
    return by_line


def _result_set_location(location, evidence, status):
    member = evidence['members'].get((location['run_id'], location['row'], location['column']))
    metadata = {}
    if member:
        metadata = {
            'fact_id': member.get('fact_id'),
            'group_id': member.get('group_id'),
            'group_type': member.get('group_type'),
            'group_status': member.get('group_status'),
            'closure_status': member.get('closure_status'),
            'member_state': member.get('member_state'),
        }
    return {**location, 'evidence': metadata, 'status': status}


def _closed_metric_key(item, scenario=None):
    """The amount class an item is gated as, or None for a known non-amount.

    Every item is an amount unless its metric is a ratio, a count or a physical quantity, or it asks only
    for ratio scenarios (achievement_ratio …); `scenario`, the cited row's scenario, exempts that one grain
    of an item that lists it. A count or quantity whose name also names an amount (price_volume,
    revenue_by_area) is an amount. A recognised P&L name gives its class (营业收入 → revenue, NetProfit →
    profit, price_volume → price_quantity); any other name, including one nothing here recognises, is the
    class 'amount' and is gated all the same.
    """
    if _known_non_amount(item, scenario):
        return None
    return _pnl_class(item.get('metric')) or 'amount'


def _known_non_amount(item, scenario=None):
    metric = item.get('metric')
    if _canonical_metric_name(metric) in NON_AMOUNT_METRICS or str(metric or '').strip().lower() in {
            str(alias).lower() for alias in METRIC_LABELS['ratio']}:
        return True
    scenarios = [str(value).strip().lower() for value in item.get('scenarios') or []]
    if scenarios and all(value in RATIO_SCENARIOS for value in scenarios):
        return True
    if scenarios and isinstance(scenario, str) and scenario.strip().lower() in RATIO_SCENARIOS:
        return True
    return _non_amount_label(str(metric or ''))


def _label_words(text):
    """Lowercase English words of a label, camelCase split (NetProfit → net, profit)."""
    return re.findall(r'[a-z]+', re.sub(r'([a-z0-9])([A-Z])', r'\1_\2', text).lower())


def _non_amount_label(text):
    head = text.strip()
    while True:
        # A trailing （%） or (次) names the unit: a percentage is a ratio; otherwise read the name before it.
        match = re.search(r'\s*[（(]([^（）()]*)[)）]\s*$', head)
        if not match:
            break
        if '%' in match.group(1) or '％' in match.group(1):
            return True
        head = head[:match.start()]
    words = _label_words(head)
    english_head = words[-1] if words and not re.search(r'[\u4e00-\u9fff]\s*$', head) else None
    # Any ratio is exempt, whatever it is a ratio of (decision 4): 毛利率, revenue_growth_rate.
    if head.endswith(RATIO_LABEL_SUFFIXES) or english_head in RATIO_HEAD_TOKENS:
        return True
    if head.endswith(QUANTITY_LABEL_SUFFIXES) or english_head in QUANTITY_HEAD_TOKENS:
        return not _names_an_amount(head, words)
    return False


def _price_volume_name(text):
    """price_quantity, price_volume_mix, 价量分析 …: a price and a volume named together is the price-quantity bridge."""
    words = set(_label_words(text))
    if words & {'price', 'prices', 'pricing'} and words & {'volume', 'volumes', 'quantity', 'quantities', 'qty'}:
        return True
    return '价量' in text or '量价' in text


def _names_an_amount(text, words):
    """Whether a name with a count/quantity/area head also names an amount (price_volume, revenue_by_area)."""
    if len(words) >= 2 and words[-2] in QUANTITY_QUALIFIERS:
        words = [*words[:-2], words[-1]]  # sales_volume: `sales` qualifies the quantity
    if any(word in AMOUNT_LABEL_TOKENS or word in AMOUNT_MARKER_WORDS for word in words):
        return True
    if any(bigram in AMOUNT_LABEL_TOKEN_PAIRS for bigram in zip(words, words[1:])):
        return True
    return any(stem in text for _key, stems in AMOUNT_LABEL_STEMS for stem in stems) or any(
        stem in text for stem in AMOUNT_MARKER_STEMS)


def _pnl_class(metric):
    """revenue/profit/cost/expense when the name is recognisably one of those lines, price_quantity for a
    price-quantity bridge, else None."""
    canonical = _canonical_metric_name(metric)
    if canonical in CLOSED_EVIDENCE_METRICS:
        return canonical
    text = str(metric or '').strip()
    if canonical == 'price_quantity' or _price_volume_name(text):
        return 'price_quantity'
    lowered = text.lower()
    for key in sorted(CLOSED_EVIDENCE_METRICS):
        if lowered in {str(alias).lower() for alias in METRIC_LABELS.get(key, ())}:
            return key
    if any(lowered in {str(alias).lower() for alias in aliases} or lowered == key
           for key, aliases in METRIC_LABELS.items() if key not in CLOSED_EVIDENCE_METRICS):
        return None  # other_income and its aliases are amounts of their own class
    for key, stems in AMOUNT_LABEL_STEMS:
        if any(stem in text for stem in stems):
            return key
    words = _label_words(text)
    pair = next((AMOUNT_LABEL_TOKEN_PAIRS[bigram] for bigram in zip(words, words[1:])
                 if bigram in AMOUNT_LABEL_TOKEN_PAIRS), None)
    return pair or next((AMOUNT_LABEL_TOKENS[word] for word in words if word in AMOUNT_LABEL_TOKENS), None)


def _location_evidence_status(item, location, member, recipe=None, *, native_total=False, scenario=None):
    """Status and server-derived evidence_basis of one cited cell.

    Every item except a known non-amount (see _closed_metric_key) is an amount. An amount counts only on
    engine closure (engine == 'finance'), on the finance engine's own run_metric account_total for exactly
    this account (`native_total`), on a passing, cited, in-binding direct/sum/diff recipe with a unit
    scale, or, for an amount outside revenue/profit/cost/expense and the price-quantity bridge, on the
    finance engine's own output of it or a server-side structured_aggregate of the bound snapshot.
    A factor recipe reaches `estimated` for every metric, never answered; script output, script-built
    closures, ratio recipes and free scales are not counted for amounts.
    """
    if not member and not recipe:
        return 'pending', {'evidence_basis': None, 'not_counted_reason': 'no_evidence_record'}
    metadata = {
        'fact_id': (member or {}).get('fact_id'),
        'group_id': (member or {}).get('group_id'),
        'group_type': (member or {}).get('group_type'),
        'group_status': (member or {}).get('group_status'),
        'closure_status': (member or {}).get('closure_status'),
        'member_state': (member or {}).get('member_state'),
        **_recipe_metadata(recipe),
    }
    member = member or {}
    if member.get('member_state') == 'missing':
        return 'missing_data', {**metadata, 'evidence_basis': None, 'not_counted_reason': 'evidence_incomplete'}
    if member.get('member_state') == 'not_applicable':
        return 'not_applicable', {**metadata, 'evidence_basis': None, 'not_counted_reason': 'method_not_applicable'}
    amount = _closed_metric_key(item, scenario)
    source_recipe = bool(
        recipe and recipe.get('counted') and recipe['evidence_class'] in SOURCE_EVIDENCE_BASES
        and recipe.get('kind') != 'ratio' and amount is not None
    )
    # A plain result row has no closure rule: an empty sibling column says nothing about a cell that a
    # source recipe recomputes, so only rule-bearing groups (bridges, ratios) can void it.
    if member.get('group_status') in {'missing_data', 'scope_mismatch', 'unit_mismatch', 'closure_mismatch'} and not (
            source_recipe and member.get('group_type') == 'calculation_row'):
        return 'missing_data', {**metadata, 'evidence_basis': None, 'not_counted_reason': 'evidence_incomplete'}
    engine_closed = bool(member.get('engine_closed')) and member.get('group_status') == 'complete'
    estimate = bool(recipe and recipe.get('counted') and recipe['evidence_class'] == 'estimate')
    if amount is None:
        # Known non-amounts keep their current rule (script ratio values are a recorded phase-1 gap),
        # except that a cell the writer backed with a factor recipe is an estimate, never an answer.
        counted = recipe if (recipe or {}).get('counted') and recipe['evidence_class'] in SOURCE_EVIDENCE_BASES else None
        if estimate and not engine_closed:
            return 'estimated', {**metadata, 'evidence_basis': 'estimate'}
        if not member and not counted:
            return 'pending', {**metadata, 'evidence_basis': None, 'not_counted_reason': 'no_evidence_record'}
        basis = 'engine_closed' if engine_closed else counted['evidence_class'] if counted else _output_basis(member)
        return 'answered', {**metadata, 'evidence_basis': basis}
    if source_recipe:
        return 'answered', {**metadata, 'evidence_basis': recipe['evidence_class']}
    if engine_closed:
        return 'answered', {**metadata, 'evidence_basis': 'engine_closed'}
    if native_total and member.get('engine_output'):
        # The finance engine's own total of exactly this account (not a script's copy of it).
        return 'answered', {**metadata, 'evidence_basis': 'engine_output'}
    if amount not in CLOSURE_OR_RECIPE_AMOUNTS and _server_output(member) and not estimate:
        # Outside the four P&L lines the engine's own output of the amount (finance_scenario_run,
        # run_metric columns) or a server-side structured aggregate of the bound snapshot answers it;
        # only revenue/profit/cost/expense need closure or a recipe (decision 1).
        return 'answered', {**metadata, 'evidence_basis': _output_basis(member)}
    if estimate:
        return 'estimated', {**metadata, 'evidence_basis': 'estimate'}
    return 'not_applicable', {**metadata, 'evidence_basis': _uncounted_basis(member),
                              'not_counted_reason': _amount_not_counted_reason(member, recipe)}


def _server_output(member):
    return bool(member.get('engine_output') or member.get('aggregate_output'))


def _output_basis(member):
    if member.get('engine_output'):
        return 'engine_output'
    return 'aggregate_output' if member.get('aggregate_output') else 'script_output'


def _uncounted_basis(member):
    if member.get('closure_status') == 'passed' and not member.get('engine_output'):
        return 'internal_only'
    return _output_basis(member) if member else None


def _amount_not_counted_reason(member, recipe):
    if recipe and recipe.get('rejected_reason'):
        return recipe['rejected_reason']
    if recipe and recipe.get('counted') and recipe.get('kind') == 'ratio':
        return 'ratio_recipe_cannot_prove_amount'
    if member.get('engine_output'):
        return 'engine_output_without_closure'
    if member.get('aggregate_output'):
        return 'aggregate_output_without_closure'
    if member.get('closure_status') == 'passed':
        return 'script_closure_not_counted'
    return 'script_output_without_recipe'


def _recipe_metadata(recipe):
    if not recipe:
        return {}
    return {
        'recipe_ids': list(recipe['recipe_ids']),
        'recipe_hash': recipe['recipe_hash'],
        'recipe_kind': recipe.get('kind'),
        'recipe_scale': recipe.get('scale'),
        'recipe_location': recipe.get('location'),
        **({'factor': recipe['factor']} if recipe.get('factor') else {}),
        **({'precision_hint': recipe['precision_hint']} if recipe.get('precision_hint') else {}),
        # Which cells the assistant chose is not yet reviewed; the draft watermark stays.
        'mapping_review': 'unreviewed',
    }


def _cell_evidence_status(item, location, evidence, *, native_total=False, row=None):
    key = (location['run_id'], location['row'], location['column'])
    return _location_evidence_status(item, location, evidence['members'].get(key), (evidence.get('recipes') or {}).get(key),
                                     native_total=native_total, scenario=(row or {}).get('scenario'))


def _claim_fixup(item, run, row_index, row, column, evidence, reason='valid_candidate_not_claimed', *, native_total=False):
    location = {'run_id': run['id'], 'row': row_index, 'column': column, 'output_hash': run.get('output_hash')}
    status, metadata = _cell_evidence_status(item, location, evidence, native_total=native_total, row=row)
    fact_id = metadata.get('fact_id')
    if status != 'answered' or not fact_id:
        return None
    return {
        'code': 'add_calculation_claim',
        'requirement_item_id': item['id'],
        **location,
        'value': row.get(column),
        'fact_id': fact_id,
        'reason': reason,
    }


def _run_inputs_match_requirement(run, requirement):
    """The run read the requirement's recorded snapshots (same dataset, data and source hashes).

    Like _bindings_match, a requirement with a single binding accepts that snapshot under any run alias;
    with several bindings each snapshot must sit under the requirement's own alias.
    """
    basis = requirement.get('semantic_basis') if isinstance(requirement, dict) else None
    sources = basis.get('sources') if isinstance(basis, dict) else None
    if not sources:
        return True
    run_bindings = run.get('bindings') or {}
    single = len(requirement.get('bindings') or {}) == 1
    for alias, source in sources.items():
        aliases = [alias, *sorted(name for name, dataset_id in run_bindings.items()
                                  if single and name != alias and dataset_id == source.get('dataset_id'))]
        if not any(_run_alias_reads(run, name, source) for name in aliases):
            return False
    return True


def _run_alias_reads(run, alias, source):
    return ((run.get('bindings') or {}).get(alias) == source.get('dataset_id')
            and (run.get('input_hashes') or {}).get(alias) == source.get('data_hash')
            and (run.get('source_hashes') or {}).get(alias) == source.get('source_hash'))


# Fixups that trace a cell to its source: drilldown_fact for engine rows, a source recipe for script rows.
SOURCE_FIXUP_CODES = frozenset({'verify_source', 'add_source_recipe'})
# Runs whose rows drilldown_fact can record as verified (finance_service.drilldown_fact trusted_run).
DRILLDOWN_VERIFIABLE_ENGINES = frozenset({'finance', 'structured_aggregate'})


def _source_fixup(item, location, checks, run=None, *, required=False):
    """How to trace a cited cell to its source, when the item (or its result set) needs that.

    drilldown_fact only verifies finance-engine and verified-aggregate rows; on a script cell it returns
    'missing', so script cells are pointed at a passing source recipe instead, and only when the cell
    already counts (an uncounted cell's own fixups say what would count, which also traces it).
    """
    if not (item.get('requires_drilldown') or required) or _location_provenance(location, checks):
        return None
    fact_id = (location.get('evidence') or {}).get('fact_id')
    if not fact_id:
        return None
    base = {key: location[key] for key in ('run_id', 'row', 'column', 'output_hash')}
    if not required and location.get('status') not in {'answered', 'estimated'}:
        # Tracing a cell that does not count changes nothing (an aggregated revenue stays uncounted after
        # drilldown_fact); the item's own fixups say what would count. Result-set rows are always traced.
        return None
    if run is not None and run.get('engine') not in DRILLDOWN_VERIFIABLE_ENGINES:
        if location.get('status') != 'answered':
            return None
        return {
            'code': 'add_source_recipe',
            'requirement_item_id': item['id'],
            'location': base,
            'fact_id': fact_id,
            'recipe_kinds': ['direct', 'sum', 'diff', 'ratio'],
            'reason': 'script_cell_traced_by_recipe',
        }
    return {
        'code': 'verify_source',
        'requirement_item_id': item['id'],
        'location': base,
        'fact_id': fact_id,
        'reason': 'drilldown_required' if item.get('requires_drilldown') else 'result_set_requires_source_check',
    }


def _dedupe_fixups(fixups):
    seen = set()
    deduped = []
    for fixup in fixups:
        if not isinstance(fixup, dict):
            continue
        key = _hash(fixup)
        if key in seen:
            continue
        seen.add(key)
        deduped.append(fixup)
    return deduped


def _result_set_candidate_fixups(analysis, item, latest, runs, checks, evidence, claimed):
    kind = _result_set_kind(item)
    if not kind:
        return []
    fixups = []
    for run in runs.values():
        if not _bindings_match(latest.get('bindings', {}), run.get('bindings', {})):
            continue
        if not _run_inputs_match_requirement(run, latest):
            continue
        metric = _run_finance_metric(analysis, run)
        if not metric or metric.get('kind') != kind:
            continue
        if kind == 'account_total' and not _account_total_metric_matches_item(metric, item):
            continue
        if not _metric_filters_cover_requirement(metric, item, _row_scope(latest['scope'])):
            continue
        if _run_scope_is_narrower_than_requirement(run, item, _row_scope(latest['scope'])):
            continue
        rows = _result_set_required_rows(item, run, _row_scope(latest['scope']))
        if not rows:
            continue
        if run.get('result', {}).get('truncated'):
            fixups.append({
                'code': 'result_set_incomplete',
                'requirement_item_id': item['id'],
                'run_id': run['id'],
                'row_count': run.get('result', {}).get('row_count', len(run.get('result', {}).get('rows', []))),
                'reason': 'run_result_truncated',
            })
            continue
        row_fixups = []
        source_fixups = []
        for index, row in rows:
            if (run['id'], index, 'amount') in claimed:
                continue
            # The finance engine's own run_metric rows are what the result set counts (see
            # _result_set_requirement_status), whatever amount class the item's name has.
            claim = _claim_fixup(item, run, index, row, 'amount', evidence, native_total=True)
            if not claim:
                row_fixups = []
                break
            row_fixups.append(claim)
            location = _result_set_location(
                {'run_id': run['id'], 'row': index, 'column': 'amount', 'output_hash': run.get('output_hash')},
                evidence, 'answered'
            )
            # The result set only answers once every row is source-checked, requires_drilldown or not.
            source = _source_fixup(item, location, checks, run, required=True)
            if source:
                source_fixups.append(source)
        fixups.extend(row_fixups)
        fixups.extend(source_fixups)
    return fixups


def _price_quantity_candidate_fixups(item, latest, runs, checks, evidence, claimed):
    fixups = []
    for run in runs.values():
        if not _bindings_match(latest.get('bindings', {}), run.get('bindings', {})):
            continue
        if not _run_inputs_match_requirement(run, latest):
            continue
        for index, row in enumerate(run.get('result', {}).get('rows', [])):
            group = evidence['groups'].get((run['id'], index))
            if not _requirement_scope_matches(run, row, item, latest):
                continue
            if not _price_quantity_requirement_matches(item, run, row, group):
                continue
            for member in _price_quantity_required_members(group):
                if (run['id'], index, member['column']) in claimed:
                    continue
                claim = _claim_fixup(item, run, index, row, member['column'], evidence)
                if claim:
                    fixups.append(claim)
                    location = _result_set_location(
                        {'run_id': run['id'], 'row': index, 'column': member['column'], 'output_hash': run.get('output_hash')},
                        evidence, 'answered'
                    )
                    source = _source_fixup(item, location, checks, run)
                    if source:
                        fixups.append(source)
    return fixups


def _generic_candidate_fixups(item, latest, runs, checks, evidence, claimed):
    if _profit_bridge_group_required(item):
        return []
    fixups = []
    for run in runs.values():
        if not _bindings_match(latest.get('bindings', {}), run.get('bindings', {})):
            continue
        if not _run_inputs_match_requirement(run, latest):
            continue
        for index, row in enumerate(run.get('result', {}).get('rows', [])):
            if not _requirement_scope_matches(run, row, item, latest) or not _requested_grain_matches(item, run, row):
                continue
            for column in row:
                if _account_amount_claim_requires_native_total(item, row, column):
                    continue
                if not _matches(item, row, column, run):
                    continue
                if (run['id'], index, column) in claimed:
                    continue
                claim = _claim_fixup(item, run, index, row, column, evidence)
                if claim:
                    fixups.append(claim)
                    location = {'run_id': run['id'], 'row': index, 'column': column, 'output_hash': run.get('output_hash')}
                    status, metadata = _cell_evidence_status(item, location, evidence, row=row)
                    source = _source_fixup(item, {**location, 'evidence': metadata, 'status': status}, checks, run)
                    if source:
                        fixups.append(source)
    return fixups


def _profit_bridge_candidate_fixups(analysis, item, latest, runs, checks, evidence, claimed):
    if not _profit_bridge_group_required(item):
        return []
    fixups = []
    for run in runs.values():
        if not _bindings_match(latest.get('bindings', {}), run.get('bindings', {})):
            continue
        if not _run_inputs_match_requirement(run, latest):
            continue
        if _profit_bridge_native_group(analysis, item, run, evidence, _row_scope(latest['scope'])) is None:
            continue
        for location in _profit_bridge_candidate_locations(analysis, item, run, evidence, _row_scope(latest['scope'])):
            key = (location['run_id'], location['row'], location['column'])
            if key not in claimed:
                row = run['result']['rows'][location['row']]
                claim = _claim_fixup(item, run, location['row'], row, location['column'], evidence)
                if claim:
                    fixups.append(claim)
            source = _source_fixup(item, _result_set_location(location, evidence, 'answered'), checks, run)
            if source:
                fixups.append(source)
    return fixups


def _method_conflict_fixups(item, latest, runs, evidence):
    expected_method = _canonical_method_name(item.get('dimensions', {}).get('method'))
    if expected_method is None:
        return []
    fixups = []
    for run in runs.values():
        actual_method = _run_method(run)
        if actual_method in (None, expected_method):
            continue
        if not _bindings_match(latest.get('bindings', {}), run.get('bindings', {})):
            continue
        if not _run_inputs_match_requirement(run, latest):
            continue
        for index, row in enumerate(run.get('result', {}).get('rows', [])):
            if not _requirement_scope_matches(run, row, item, latest):
                continue
            relaxed = {**item, 'dimensions': {k: v for k, v in item.get('dimensions', {}).items() if k != 'method'}}
            group = evidence['groups'].get((run['id'], index))
            if not (
                _price_quantity_requirement_matches(relaxed, run, row, group)
                or any(_matches(relaxed, row, column, run) for column in row)
            ):
                continue
            fixups.append({
                'code': 'method_conflict',
                'requirement_item_id': item['id'],
                'run_id': run['id'],
                'row': index,
                'expected_method': expected_method,
                'actual_method': actual_method,
                'reason': 'requirement_method_does_not_match_run',
            })
            break
    return fixups


def _requirement_fixups(analysis, item, latest, runs, checks, evidence, claimed):
    return _dedupe_fixups([
        *_result_set_candidate_fixups(analysis, item, latest, runs, checks, evidence, claimed),
        *_price_quantity_candidate_fixups(item, latest, runs, checks, evidence, claimed),
        *_generic_candidate_fixups(item, latest, runs, checks, evidence, claimed),
        *_profit_bridge_candidate_fixups(analysis, item, latest, runs, checks, evidence, claimed),
        *_method_conflict_fixups(item, latest, runs, evidence),
    ])


def check_requirements(analysis, spec, *, _requirements_record=None, _recipe_report=None):
    latest = _requirements_record or latest_requirements(analysis)
    if not latest:
        return {'requirements_id': None, 'items': [], 'coverage': {'status': 'not_assessed'},
                'provenance': {'status': 'not_assessed'}, 'scope_changed': False,
                'original_requirements_completed': False}
    selected_id = spec.get('requirements_id') or latest['id']
    selected = latest if _requirements_record and latest.get('id') == selected_id else analysis._get('requirements', selected_id)
    # Public checks use current obligations; only internal appendix reconstruction
    # supplies an immutable historical record to recover its original text.
    runs = {r['id']: r for r in analysis._list('runs') if r.get('execution_status') == 'succeeded'}
    evidence = _prepared_evidence(analysis, runs)
    # Recipe results are computed by the server in this call (or handed over by an internal caller
    # that just computed them); spec['evidence_report'] is never read.
    evidence['recipes'] = _recipe_support(analysis, spec, latest, _recipe_report)
    claimed = set()
    for claim in spec.get('claims', []):
        if not isinstance(claim, dict) or claim.get('basis') != 'calculation':
            continue
        run = runs.get(claim.get('run_id'))
        index = claim.get('row')
        if not run or run['id'] not in spec.get('run_ids', []) or isinstance(index, bool) or not isinstance(index, int) or index < 0:
            continue
        rows = run.get('result', {}).get('rows', [])
        if index < len(rows) and claim.get('column') in rows[index] and _same_number(rows[index][claim['column']], claim.get('value')):
            claimed.add((run['id'], index, claim['column']))
    # Sorted so locations, messages and completion fingerprints do not depend on hash seeds.
    claimed_cells = sorted(claimed, key=lambda cell: (str(cell[0]), cell[1], str(cell[2])))
    claimed_by_row = {}
    for run_id, index, column in claimed_cells:
        claimed_by_row.setdefault((run_id, index), set()).add(column)
    fixup_runs = {run_id: runs[run_id] for run_id in spec.get('run_ids', []) if run_id in runs}
    checks = analysis._list('evidence_checks')
    source_failures = _source_read_failures(spec.get('source_read_failures', []))
    environment = _evidence_environment(analysis, latest, runs)
    acceptances = _estimate_acceptances(analysis, latest)
    items = []
    for item in latest['items']:
        locations = []
        statuses = []
        processed_groups = set()
        result_set = _result_set_requirement_status(analysis, item, latest, runs, claimed, checks, evidence)
        if result_set is not None:
            locations = result_set['locations']
            statuses = [result_set['status']]
        for run_id, index, column in claimed_cells:
            if result_set is not None:
                break
            run = runs[run_id]
            row = run['result']['rows'][index]
            if not _requested_grain_matches(item, run, row):
                continue
            if not _bindings_match(latest.get('bindings', {}), run.get('bindings', {})):
                # Cited, about this item, but on other bindings: say so (ruling (a)), never "no cited row".
                if _matches(item, row, column, run):
                    locations.append(_bindings_mismatch_location(run, index, column))
                continue
            group = evidence['groups'].get((run_id, index))
            profit_bridge_locations = _profit_bridge_locations(analysis, item, run, claimed_by_row, evidence, _row_scope(latest['scope']))
            if profit_bridge_locations and ('profit_bridge', run_id) not in processed_groups:
                processed_groups.add(('profit_bridge', run_id))
                for location in profit_bridge_locations:
                    status, metadata = _cell_evidence_status(item, location, evidence)
                    locations.append({**location, 'evidence': metadata, 'status': status})
                continue
            if (
                _price_quantity_requirement_matches(item, run, row, group)
                and (run_id, index) not in processed_groups
                and set(PRICE_QUANTITY_REQUIRED_COLUMNS).issubset(claimed_by_row.get((run_id, index), set()))
                and _requirement_scope_matches(run, row, item, latest)
            ):
                processed_groups.add((run_id, index))
                for member in _price_quantity_required_members(group):
                    location = {'run_id': run_id, 'row': index, 'column': member['column'], 'output_hash': run.get('output_hash')}
                    status, metadata = _cell_evidence_status(item, location, evidence)
                    locations.append({**location, 'evidence': metadata, 'status': status})
                continue
            if _profit_bridge_group_required(item):
                continue
            account_total_match = (_account_total_requirement_matches(analysis, item, run, row, column)
                                   and _requirement_scope_matches(run, row, item, latest))
            generic_match_allowed = not _account_amount_claim_requires_native_total(item, row, column)
            if account_total_match or (
                generic_match_allowed and _matches(item, row, column, run) and _requirement_scope_matches(run, row, item, latest)
            ):
                location = {'run_id': run_id, 'row': index, 'column': column, 'output_hash': run.get('output_hash')}
                status, metadata = _cell_evidence_status(item, location, evidence, native_total=account_total_match, row=row)
                locations.append({**location, 'evidence': metadata, 'status': status})
        if result_set is None:
            # An estimate only counts where the item declared that grain an estimate (grain_basis).
            locations = [_with_declared_basis(item, location, runs) for location in locations]
            statuses = [location['status'] for location in locations]
        item_status = _item_status(statuses)
        grain_locations = _requested_grains(item) if result_set is None else None
        grain_states = {}
        if result_set is None:
            grains = grain_locations if grain_locations is not None else {(None, None, None): []}
            for location in locations:
                run = runs[location['run_id']]
                row = run['result']['rows'][location['row']]
                key = _row_grain(item, run, row) if grain_locations is not None else (None, None, None)
                grains.get(key, []).append(location)
            grain_states = {grain: _grain_state(item, grain, members) for grain, members in grains.items()}
            if grain_locations is not None:
                states = set(grain_states.values())
                item_status = ('answered' if states == {'answered'}
                               else 'estimated' if states <= {'answered', 'estimated'} else 'pending')
        failed_source = item['id'] in source_failures
        provenance = [_location_provenance(location, checks) for location in locations]
        if _profit_bridge_group_required(item):
            source_complete = bool(provenance) and all(value == 'verified' for value in provenance)
        elif grain_locations is not None:
            source_complete = all(
                any(location['status'] == 'answered' and _location_provenance(location, checks) == 'verified'
                    for location in members)
                for members in grain_locations.values()
            )
        else:
            source_complete = any(value == 'verified' for value in provenance)
        estimate_inputs = item_status == 'estimated' and all(
            any(_location_provenance(location, checks) in {'verified', 'inputs_verified'}
                and location['status'] in {'answered', 'estimated'} for location in members)
            for members in (grain_locations or {(None, None, None): locations}).values()
        )
        provenance_status = 'not_requested'
        if item['requires_drilldown']:
            provenance_status = (
                'failed' if failed_source
                else 'verified' if item_status == 'answered' and source_complete
                else 'inputs_verified' if estimate_inputs
                else 'pending'
            )
        fixups = _requirement_fixups(analysis, item, latest, fixup_runs, checks, evidence, claimed)
        if item_status == 'answered':
            fixups = [fixup for fixup in fixups if fixup.get('code') in SOURCE_FIXUP_CODES]
        item_record = {**item, 'status': item_status, 'locations': locations,
                       'source_read_failure': source_failures.get(item['id']),
                       'provenance_status': provenance_status}
        paths = _evidence_paths(item, environment)
        if paths['accepted'] or paths['blocked']:
            item_record['evidence_paths'] = paths
        hierarchy = _hierarchy_dimensions(item)
        if hierarchy:
            item_record['hierarchy_dimensions'] = hierarchy
            fixups = [*fixups, _revise_dimension_fixup(item)]
        if item_status != 'answered':
            item_record, fixups = _unanswered_details(
                item, item_record, fixups, latest, runs, claimed_cells, evidence, grain_states, paths, acceptances,
                result_set is not None,
            )
        source_fixups = [fixup for location in locations
                         for fixup in [_source_fixup(item, location, checks, runs.get(location['run_id']),
                                                     required=result_set is not None)] if fixup]
        fixups = _dedupe_fixups([*fixups, *source_fixups])
        if fixups:
            item_record['fixups'] = fixups
        items.append(item_record)
    scope_changed = False if _requirements_record else selected['id'] != latest['id']
    complete = all(item['status'] == 'answered' for item in items) and not scope_changed
    accepted_estimates = (
        not complete and not scope_changed
        and all(item['status'] in {'answered', 'estimated'} for item in items)
        and all(grain.get('accepted') for item in items for grain in item.get('estimated_grains', []))
        and any(item['status'] == 'estimated' for item in items)
    )
    trace_complete = all(item['provenance_status'] in {'verified', 'not_requested'} for item in items)
    trace_with_estimates = accepted_estimates and all(
        item['provenance_status'] in {'verified', 'not_requested', 'inputs_verified'} for item in items)
    return {'requirements_id': latest['id'], 'draft_requirements_id': selected['id'],
            'items': items, 'scope_changed': scope_changed,
            'coverage': {'status': 'pass' if complete else 'pass_with_accepted_estimates' if accepted_estimates
                         else 'needs_attention'},
            'provenance': {'status': 'pass' if trace_complete else 'pass_with_accepted_estimates' if trace_with_estimates
                           else 'needs_attention'},
            # Accepted estimates never complete the original requirements.
            'original_requirements_completed': complete and not latest.get('removed_item_ids')}


def _bindings_mismatch_location(run, index, column):
    """A cited cell about the item on a run whose bindings differ from the requirement's: never counted
    (the run is not proven to have read the requirement's snapshots under their names), always reported."""
    return {'run_id': run['id'], 'row': index, 'column': column, 'output_hash': run.get('output_hash'),
            'evidence': {'evidence_basis': None, 'not_counted_reason': 'run_bindings_differ'}, 'status': 'pending'}


def _with_declared_basis(item, location, runs):
    if location['status'] != 'estimated':
        return location
    run = runs[location['run_id']]
    grain = _row_grain(item, run, run['result']['rows'][location['row']])
    declared = _declared_basis(item, grain)
    if declared and declared['basis'] == 'estimate' and declared.get('factor_reproducible', True):
        return location
    # A grain declared non-reproducible waits for the manager's basis; a reproducible factor cited anyway
    # is another coefficient standing in for it (decision 3), not a source grain cited with an estimate.
    reason = ('estimate_factor_declared_not_reproducible'
              if declared and declared['basis'] == 'estimate' else 'estimate_cited_for_source_grain')
    # PRD R04: where the source has the number, an estimate does not answer it.
    return {**location, 'status': 'pending', 'evidence': {**location['evidence'], 'not_counted_reason': reason}}


def _grain_state(item, grain, members):
    if any(location['status'] == 'answered' for location in members):
        return 'answered'
    if any(location['status'] == 'estimated' for location in members):
        return 'estimated'
    declared = _declared_basis(item, grain)
    if declared and declared['basis'] == 'unavailable':
        return 'unavailable_declared'
    if declared and declared['basis'] == 'estimate' and not declared.get('factor_reproducible', True):
        return 'estimate_factor_not_reproducible'
    return 'uncounted' if members else 'missing'


MANAGER_GRAIN_STATES = {
    'estimated': 'estimate_requires_acceptance',
    'unavailable_declared': 'unavailable_declared',
    'estimate_factor_not_reproducible': 'estimate_factor_not_reproducible',
}


def _grain_record(item, grain):
    return {key: value for (key, field), value in zip(GRAIN_LISTS, grain) if item.get(field)}


def _unanswered_details(item, record, fixups, latest, runs, claimed_cells, evidence, grain_states, paths,
                        acceptances, result_set):
    """Why an unanswered item is unanswered, split by grain, and what can still answer it."""
    missing = [_grain_record(item, grain) for grain, state in grain_states.items() if state == 'missing']
    uncounted = []
    for grain, state in grain_states.items():
        if state != 'uncounted':
            continue
        members = [location for location in record['locations'] if _location_grain(item, location, runs) == grain]
        reasons = [location['evidence'].get('not_counted_reason') for location in members]
        uncounted.append({**_grain_record(item, grain),
                          'statuses': list(dict.fromkeys(location['status'] for location in members)),
                          'reasons': list(dict.fromkeys(reason for reason in reasons if reason))})
    estimated = []
    for grain, state in grain_states.items():
        if state != 'estimated':
            continue
        members = [location for location in record['locations']
                   if location['status'] == 'estimated' and _location_grain(item, location, runs) == grain]
        estimated.append({**_grain_record(item, grain),
                          'accepted': all(_estimate_accepted(acceptances, item, grain, location) for location in members),
                          'locations': [{key: location[key] for key in ('run_id', 'row', 'column')} for location in members]})
    declared = [{**_grain_record(item, grain), 'state': state,
                 'reason': (_declared_basis(item, grain) or {}).get('reason')}
                for grain, state in grain_states.items()
                if state in {'unavailable_declared', 'estimate_factor_not_reproducible'}]
    if missing:
        record['missing_grains'] = missing
    if uncounted:
        record['uncounted_grains'] = uncounted
    if estimated:
        record['estimated_grains'] = estimated
    if declared:
        record['declared_grains'] = declared
    route = _method_guidance(item)
    if route:
        record['required_evidence'] = route
    if missing or uncounted:
        required_fields = _required_row_fields(item, latest)
        field_gaps = _cited_row_field_gaps(item, latest, runs, claimed_cells, required_fields)
        if field_gaps:
            record['row_field_gaps'] = field_gaps
        if (missing and not _has_finance_scope_runs(runs, claimed_cells)
                and _scope_identity_keys(_item_scope(item, latest.get('scope') or {})) and not _requirement_sources(latest)):
            record['ledger_identity'] = 'unproven_without_bound_snapshots'
        if missing and _requested_grains(item) is not None:
            fixups = [*fixups, {**_missing_rows_fixup(item, missing, required_fields, field_gaps, paths, route),
                                **_grain_basis_guidance(item, missing, paths)}]
        if uncounted:
            uncounted_keys = {grain for grain, state in grain_states.items() if state == 'uncounted'}
            fixups = [*fixups, {**_uncounted_rows_fixup(item, record['locations'], uncounted, paths, route, runs, uncounted_keys),
                                **_grain_basis_guidance(item, uncounted, paths)}]
    rebind = list(dict.fromkeys(location['run_id'] for location in record['locations']
                                if location['evidence'].get('not_counted_reason') == 'run_bindings_differ'))
    if rebind:
        record['required_bindings'] = dict(latest.get('bindings') or {})
        fixups = [*fixups, {
            'code': 'rebind_run',
            'requirement_item_id': item['id'],
            # Bind every requirement snapshot under the requirement's own alias (one binding: any alias).
            'required_bindings': dict(latest.get('bindings') or {}),
            'runs': [{'run_id': run_id, 'bindings': dict(runs[run_id].get('bindings') or {})} for run_id in rebind[:12]],
            'reason': 'run_bindings_differ',
        }]
    if result_set and paths['blocked'] and not paths['accepted']:
        fixups = [*fixups, _blocked_path_fixup(item, paths)]
    cited_for_source = [grain for grain in uncounted if 'estimate_cited_for_source_grain' in grain['reasons']]
    if cited_for_source:
        fixups = [*fixups, {
            'code': 'estimate_cited_for_source_grain',
            'requirement_item_id': item['id'],
            'grains': cited_for_source[:12],
            'reason': 'estimate_cited_for_source_grain',
        }]
    for entry in declared:
        fixups = [*fixups, {
            'code': 'estimate_factor_not_reproducible' if entry['state'] == 'estimate_factor_not_reproducible'
            else 'confirm_unavailable',
            'requirement_item_id': item['id'],
            'grain': {key: value for key, value in entry.items() if key in GRAIN_KEYS},
            'declared_reason': entry.get('reason'),
            'needs_manager': True,
            'reason': entry['state'],
        }]
    # A hierarchy dimension on an old record is the assistant's to revise first (revise_requirement_dimension);
    # until then the item cannot be answered whatever the manager confirms, so it is not handed over.
    reasons = [] if _hierarchy_dimensions(item) else _manager_reasons(grain_states, paths, result_set, record['status'])
    if reasons:
        record['awaiting_manager'] = reasons
    return record, fixups


def _location_grain(item, location, runs):
    run = runs.get(location['run_id']) or {}
    rows = (run.get('result') or {}).get('rows') or []
    if not item.get('periods') and not item.get('organizations') and not item.get('scenarios'):
        return (None, None, None)
    return _row_grain(item, run, rows[location['row']]) if location['row'] < len(rows) else None


def _has_finance_scope_runs(runs, claimed_cells):
    return any(_has_finance_scope(runs[run_id]) for run_id, _, _ in claimed_cells if run_id in runs)


def _manager_reasons(grain_states, paths, result_set, status):
    """Reasons only the manager can clear; set only when nothing else is left for the assistant."""
    if status == 'answered':
        return []
    manager_blocked = not paths['accepted'] and bool(paths['blocked']) and all(
        set(path['blockers']) & MANAGER_BLOCKERS for path in paths['blocked'])
    blocker_reasons = sorted({blocker for path in paths['blocked'] for blocker in path['blockers']
                              if blocker in MANAGER_BLOCKERS}) if manager_blocked else []
    if result_set:
        return blocker_reasons
    unresolved = [state for state in grain_states.values() if state != 'answered']
    if not unresolved:
        return []
    if all(state in MANAGER_GRAIN_STATES for state in unresolved):
        return sorted({MANAGER_GRAIN_STATES[state] for state in unresolved})
    if manager_blocked:
        return sorted({*blocker_reasons, *(MANAGER_GRAIN_STATES[state] for state in unresolved if state in MANAGER_GRAIN_STATES)})
    return []


def _estimate_acceptances(analysis, latest):
    """Manager acceptances of specific estimate recipes (written by the phase-2 manager command only)."""
    try:
        records = analysis._list('estimate_acceptances')
    except FoundryError:
        return []
    return [record for record in records
            if record.get('requirements_id') == latest.get('id') and record.get('accepted_by') == 'manager']


def _estimate_accepted(acceptances, item, grain, location):
    evidence = location.get('evidence') or {}
    expected = {
        'requirement_item_id': item['id'],
        'grain': {key: value for key, value in zip(GRAIN_KEYS, grain) if value is not None},
        'recipe_hash': evidence.get('recipe_hash'),
        'factor_recipe_hash': (evidence.get('factor') or {}).get('recipe_hash'),
    }
    return bool(expected['recipe_hash']) and any(
        all(record.get(key) == value for key, value in expected.items()) for record in acceptances)


def _evidence_environment(analysis, latest, runs=None):
    """What the requirement's bound data allows right now: archived workbooks for recipes, and whether the
    finance engine can run on it (a saved finance context or finance runs on these snapshots, and
    structures the manager has confirmed)."""
    dataset_ids = [value for value in (latest.get('bindings') or {}).values() if isinstance(value, str)]
    if runs is None:
        try:
            runs = {run['id']: run for run in analysis._list('runs') if run.get('execution_status') == 'succeeded'}
        except FoundryError:
            runs = {}
    engine_runs = any(
        run.get('engine') == 'finance'
        and (not dataset_ids or set((run.get('bindings') or {}).values()) & set(dataset_ids))
        for run in runs.values()
    )
    workbooks, tables, accounts = [], [], set()
    for dataset_id in dataset_ids:
        try:
            dataset = analysis.read_dataset(dataset_id)
        except (FoundryError, OSError, ValueError, KeyError, TypeError):
            continue
        # evaluate_evidence_recipes reads the archived XLSX bytes whatever the snapshot was filed as
        # (a one-sheet workbook is filed as a table by default).
        if str(dataset.get('source_file') or '').lower().endswith('.xlsx'):
            workbooks.append(dataset_id)
        if isinstance(dataset.get('rows'), list) and dataset['rows']:
            tables.append(dataset_id)
        # Management accounts the finance engine can total on its own (run_metric account_total).
        accounts.update(str(row['management_account']) for row in dataset.get('rows') or []
                        if isinstance(row, dict) and row.get('management_account') not in (None, ''))
    finance_context, pending = None, []
    try:
        from foundry.finance_service import FinanceStore
        from foundry.structure_records import PENDING_CONFIRMATION

        store = FinanceStore(analysis.task)
        contexts = store._list('contexts')
        finance_context = bool(latest.get('context_id')) or engine_runs or any(
            not dataset_ids or set((context.get('bindings') or {}).values()) & set(dataset_ids) for context in contexts)
        pending = [dataset_id for dataset_id in dataset_ids
                   if store.dataset_structure_state(dataset_id) == PENDING_CONFIRMATION]
    except (FoundryError, OSError, ValueError, KeyError, TypeError):
        finance_context = True if engine_runs else None
    return {'workbook_sources': workbooks, 'finance_context': finance_context, 'structure_pending': pending,
            'accounts': sorted(accounts),
            # structured_aggregate reads exactly one snapshot, so it can meet only a single-binding requirement.
            'aggregate_sources': tables if len(dataset_ids) == 1 else []}


def _engine_blockers(environment):
    blockers = []
    if environment.get('structure_pending'):
        blockers.append('structure_pending_confirmation')
    if environment.get('finance_context') is False:
        blockers.append('no_finance_context')
    return blockers


def _evidence_paths(item, environment):
    """Evidence paths that can answer this item in the current code, and the ones blocked right now.

    Only paths the matcher can actually count are named: run_metric management_pnl rows never close
    for revenue/profit/cost/expense, and profit-bridge rows carry no scenario, so neither is offered
    where it cannot answer.
    """
    accepted, blocked = [], []
    engine_blockers = _engine_blockers(environment)

    def engine(path):
        if engine_blockers:
            blocked.append({**path, 'blockers': list(engine_blockers)})
        else:
            accepted.append(path)

    kind = _result_set_kind(item)
    route = _method_guidance(item)
    if kind:
        engine({'path': 'run_metric_result_set', 'tool': 'run_metric', 'metric_kind': kind})
    elif route.get('cite_row_metrics'):
        engine({'path': 'profit_bridge', 'tool': 'explain_variance', 'method': route['required_method']})
    elif route.get('cite_columns'):
        engine({'path': 'price_quantity_bridge', 'tool': 'explain_variance', 'method': route['required_method']})
    elif route.get('required_method') is not None:
        accepted.append({'path': 'method_run', 'method': route['required_method']})
    else:
        if _ledger_account_item(item, environment):
            # The engine's own total of this one account answers it (see _account_total_requirement_matches).
            engine({'path': 'run_metric_account_total', 'tool': 'run_metric', 'metric_kind': 'account_total',
                    'account': item['metric']})
        if _closed_metric_key(item) is not None:
            native = ENGINE_OUTPUT_AMOUNTS.get(_canonical_metric_name(item.get('metric')))
            if native:
                # An amount the finance engine states itself (not one of the four P&L lines).
                engine({'path': 'engine_output', **native})
            if _canonical_metric_name(item.get('metric')) == 'price_quantity':
                # The price-quantity bridge closes only on the finance engine (_price_quantity_requirement_matches).
                engine({'path': 'price_quantity_bridge', 'tool': 'explain_variance',
                        'methods': sorted(PRICE_QUANTITY_METHODS)})
            if _closed_metric_key(item) not in CLOSURE_OR_RECIPE_AMOUNTS and environment.get('aggregate_sources'):
                # A server-side aggregate of the bound snapshot answers amounts outside the four P&L lines
                # and the price-quantity bridge.
                accepted.append({'path': 'aggregate_output', 'tool': 'analysis-save', 'engine': 'structured_aggregate',
                                 'measure': item['metric']})
            recipe = {'path': 'evidence_recipe', 'tool': 'save_explanation_draft.evidence_recipes',
                      'recipe_kinds': ['direct', 'sum', 'diff'], 'scale': 'one_or_unit_power'}
            if environment.get('workbook_sources'):
                accepted.append(recipe)
            else:
                blocked.append({**recipe, 'blockers': ['no_workbook_source']})
            # explain_variance rows name the canonical lines (revenue/cost/expense/profit) and carry no
            # scenario, so closure is offered only where _metric_matches can meet those rows.
            if not item.get('scenarios') and _canonical_metric_name(item.get('metric')) in CLOSED_EVIDENCE_METRICS:
                engine({'path': 'engine_closure', 'tool': 'explain_variance'})
        else:
            accepted.append({'path': 'cited_result_row'})
    paths = {'accepted': accepted, 'blocked': blocked}
    if not kind and not route and any(entry.get('basis') == 'estimate' and entry.get('factor_reproducible', True)
                                      for entry in item.get('grain_basis') or []):
        # Grains declared estimates reach `estimated` (never answered) by a factor recipe on the archive.
        estimate = {'path': 'estimate_recipe', 'tool': 'save_explanation_draft.evidence_recipes',
                    'recipe_kinds': ['direct', 'sum'], 'scale': 'one_or_unit_power',
                    'factor_recipe_kind': 'ratio', 'reaches': 'estimated'}
        paths['estimate'] = estimate if environment.get('workbook_sources') else {**estimate, 'blockers': ['no_workbook_source']}
    return paths


def _ledger_account_item(item, environment):
    metric = item.get('metric')
    return (
        isinstance(metric, str)
        and _canonical_metric_name(metric) not in METRIC_LABELS
        and _canonical_metric_name(metric) not in RESULT_SET_METRICS
        and metric in (environment.get('accounts') or ())
    )


def _revise_dimension_fixup(item):
    keys = _hierarchy_dimensions(item)
    suggested = {key: value for key, value in _without_hierarchy(item).items()
                 if key in {'id', 'metric', 'label', 'dimensions', 'periods', 'organizations', 'scenarios',
                            'requires_drilldown', 'grain_basis'}}
    return {
        'code': 'revise_requirement_dimension',
        'requirement_item_id': item['id'],
        'keys': sorted(keys),
        'suggested_item': suggested,
        'reason': 'hierarchy_dimension_not_row_filter',
    }


def _required_row_fields(item, latest):
    """Fixed fields a cited row must carry besides metric/grain/value: non-intent dimensions and, when the
    run has no finance_scope, the requirement scope (see _requirement_scope_matches' row fallback).

    Grain keys are dropped only when the item lists them (the row states them as its grain); the book
    identity (ledger) is proven by the bound snapshot hashes and is never copied onto rows.
    """
    fields = {}
    dimensions = item.get('dimensions') if isinstance(item.get('dimensions'), dict) else {}
    hierarchy = _hierarchy_dimensions(item)
    for key, value in dimensions.items():
        if key not in {'method', *INTENT_DIMENSION_KEYS} and key not in hierarchy:
            fields[key] = value
    listed = {key for key, field in GRAIN_LISTS if item.get(field)}
    scope = _item_scope(item, latest.get('scope') or {})
    for key, value in (_comparison_scopes(scope)['actual'] if scope else {}).items():
        if key in listed or key in LEDGER_IDENTITY_KEYS:
            continue
        fields.setdefault(key, value)
    return fields


def _cited_row_field_gaps(item, latest, runs, claimed_cells, required_fields):
    """Fields that already-cited rows about this metric and grain miss or state differently.

    Uses the real metric matcher (a ratio metric_key pin included) and reports gaps in required_fields
    order, so the message and completion fingerprint are reproducible.
    """
    missing = set()
    for run_id, index, column in claimed_cells:
        run = runs.get(run_id)
        if not run or not _bindings_match(latest.get('bindings', {}), run.get('bindings', {})):
            continue
        row = run['result']['rows'][index]
        if not _requested_grain_matches(item, run, row) or not _metric_matches(item, row, column):
            continue
        dimensions_ok = _dimensions_match(_without_hierarchy(item), row, run)
        scope_ok = _requirement_scope_matches(run, row, item, latest)
        if dimensions_ok and scope_ok:
            continue
        for key, value in required_fields.items():
            is_dimension = key in (item.get('dimensions') or {})
            if (is_dimension and dimensions_ok) or (not is_dimension and scope_ok):
                continue
            if row.get(key) != value:
                missing.add(key)
    return {key: value for key, value in required_fields.items() if key in missing}


def _row_keys(item, grain, dimension_fields):
    return list(dict.fromkeys(['metric_key', *[key for key in GRAIN_KEYS if key in grain], *dimension_fields, 'value']))


def _missing_rows_fixup(item, missing_grains, required_fields, field_gaps, paths, route):
    """Tell the writer which cited results would answer the item; wide tables carry none of these keys."""
    dimensions = item.get('dimensions') if isinstance(item.get('dimensions'), dict) else {}
    dimension_fields = {key: value for key, value in required_fields.items() if key in dimensions}
    scope_fields = {key: value for key, value in required_fields.items() if key not in dimensions}
    uses_rows = any(path['path'] in {'evidence_recipe', 'cited_result_row', 'method_run'} for path in paths['accepted'])
    return {
        'code': 'add_requirement_rows',
        'requirement_item_id': item['id'],
        'metric': item.get('metric'),
        **({'row_keys': _row_keys(item, missing_grains[0], dimension_fields),
            # The row's metric_key names the requirement's own metric unless a dimension pins it.
            **({'metric_key': _canonical_metric_name(item.get('metric'))} if 'metric_key' not in dimension_fields else {}),
            **({'row_fields': dimension_fields} if dimension_fields else {}),
            # On every row: only a finance-engine run's finance_scope stands in for row fields, and a script
            # or SQL run has none (its run.scope metadata is not read by requirement matching).
            **({'scope_fields': scope_fields, 'scope_fields_on': 'rows'} if scope_fields else {})} if uses_rows else {}),
        **({'cited_row_field_gaps': field_gaps} if field_gaps else {}),
        **_path_fields(item, paths, route),
        'missing_count': len(missing_grains),
        'missing': missing_grains[:12],
        'reason': 'cited_rows_miss_required_fields' if field_gaps else 'no_cited_row_for_requested_grain',
    }


def _uncounted_rows_fixup(item, locations, uncounted, paths, route, runs, uncounted_keys):
    """Grains whose cited results matched the item but did not count, each with its statuses and reasons.

    evidence_recipe_cells lists only the cells of those grains that did not count: an estimated cell, or a
    cell of a grain another cell already answers, is not something to repair.
    """
    recipe_cells = []
    if any(path['path'] == 'evidence_recipe' for path in paths['accepted']):
        for location in locations:
            reason = location['evidence'].get('not_counted_reason')
            if (location['status'] in {'answered', 'estimated'}
                    or reason in {'estimate_cited_for_source_grain', 'run_bindings_differ'}
                    or _location_grain(item, location, runs) not in uncounted_keys):
                continue
            row = runs[location['run_id']]['result']['rows'][location['row']]
            recipe_cells.append({'run_id': location['run_id'], 'row': location['row'], 'column': location['column'],
                                 'value': row.get(location['column']), 'not_counted_reason': reason})
    return {
        'code': 'add_requirement_rows',
        'requirement_item_id': item['id'],
        'metric': item.get('metric'),
        **_path_fields(item, paths, route),
        **({'evidence_recipe_cells': recipe_cells[:12]} if recipe_cells else {}),
        'not_counted_count': len(uncounted),
        'not_counted': uncounted[:12],
        'reason': 'cited_result_not_counted',
    }


# Uncounted reasons where the source may not state the grain's value, so declaring it in grain_basis can be
# the way on: the writer is estimating (a coefficient, an undeclared estimate, a ratio standing in for the
# amount), or cited a script value that no recipe ties to an archived cell (Huachen's 区域全年预算 × 进度率).
# A missing grain always qualifies. A rounded value, a recipe on an unbound snapshot or a run on other
# bindings shows the source does hold the number, so those never get the hint.
ESTIMATING_REASONS = frozenset({
    'recipe_coefficient_unbound', 'estimate_cited_for_source_grain', 'ratio_recipe_cannot_prove_amount',
    'script_output_without_recipe', 'recipe_failed',
})
DECLARE_GRAIN_BASIS = {
    'tool': 'analysis_requirements_save', 'field': 'grain_basis', 'with': 'parent_id',
    'options': [{'basis': 'estimate', 'requires': ['method'], 'then': 'estimate_recipe'},
                {'basis': 'estimate', 'factor_reproducible': False, 'requires': ['method', 'reason'], 'then': 'manager'},
                {'basis': 'unavailable', 'requires': ['reason'], 'then': 'manager'}],
    'never_answers': True,
}


def _grain_key(grain):
    return tuple(grain.get(key) for key in GRAIN_KEYS)


def _reproducible_estimate_declared(item, grain):
    declared = _declared_basis(item, _grain_key(grain))
    return bool(declared and declared['basis'] == 'estimate' and declared.get('factor_reproducible', True))


def _needs_grain_basis_hint(item, grain):
    """A closed-metric grain with no declaration that nothing counted for, where the source may lack it."""
    return (
        _closed_metric_key(item) is not None
        and _declared_basis(item, _grain_key(grain)) is None
        and ('reasons' not in grain or bool(set(grain['reasons']) & ESTIMATING_REASONS))
    )


def _grain_basis_guidance(item, grains, paths):
    """How the grains the source does not state are handled: a declared estimate grain needs a factor
    recipe (it reaches `estimated`, never answered); an undeclared one is declared in grain_basis first."""
    guidance = {}
    estimates = [{key: grain[key] for key in GRAIN_KEYS if key in grain}
                 for grain in grains if _reproducible_estimate_declared(item, grain)]
    if estimates and paths.get('estimate'):
        guidance.update(estimate_grains=estimates[:12], estimate_evidence=paths['estimate'])
    if any(_needs_grain_basis_hint(item, grain) for grain in grains):
        guidance['if_source_lacks_value'] = DECLARE_GRAIN_BASIS
    return guidance


def _blocked_path_fixup(item, paths):
    blockers = sorted({blocker for path in paths['blocked'] for blocker in path['blockers']})
    return {
        'code': 'evidence_path_blocked',
        'requirement_item_id': item['id'],
        'metric': item.get('metric'),
        'blocked_evidence': paths['blocked'],
        'needs_manager': bool(set(blockers) & MANAGER_BLOCKERS),
        'reason': blockers[0] if len(blockers) == 1 else 'evidence_paths_blocked',
    }


def _path_fields(item, paths, route):
    return {
        **({'requires_closed_evidence': True} if _closed_metric_key(item) is not None else {}),
        'accepted_evidence': [path['path'] for path in paths['accepted']],
        **({'blocked_evidence': paths['blocked']} if paths['blocked'] else {}),
        **route,
    }


def _method_guidance(item):
    """What a counted result must come from when the item pins a method: _dimensions_match needs the run
    method; profit-bridge and price-quantity items only match their native bridge cells."""
    if _profit_bridge_group_required(item):
        return {'required_method': 'revenue_minus_cost_minus_expense', 'cite_columns': ['contribution'],
                'cite_row_metrics': list(PROFIT_BRIDGE_REQUIRED_LINES)}
    dimensions = item.get('dimensions') if isinstance(item.get('dimensions'), dict) else {}
    method = _canonical_method_name(dimensions.get('method'))
    if method is None:
        return {}
    if (
        isinstance(method, str) and method in PRICE_QUANTITY_METHODS
        and _canonical_metric_name(item.get('metric')) in {'amount', 'revenue', 'price_quantity'}
    ):
        return {'required_method': method, 'cite_columns': list(PRICE_QUANTITY_REQUIRED_COLUMNS)}
    return {'required_method': method}


def _item_status(statuses):
    if not statuses:
        return 'pending'
    if 'answered' in statuses:
        return 'answered'
    if 'estimated' in statuses:
        return 'estimated'
    if 'missing_data' in statuses:
        return 'missing_data'
    if 'not_applicable' in statuses:
        return 'not_applicable'
    return 'pending'


def _source_read_failures(failures):
    if not isinstance(failures, list):
        return {}
    indexed = {}
    for failure in failures:
        if not isinstance(failure, dict):
            continue
        item_id = failure.get('requirement_item_id') or failure.get('item_id')
        if isinstance(item_id, str):
            indexed[item_id] = failure
    return indexed


def _location_provenance(location, checks):
    """'verified' for a cell traced to source (drilldown check, or a counted source recipe that recomputes
    it from the archived snapshot); 'inputs_verified' for an estimate whose inputs were recomputed."""
    basis = (location.get('evidence') or {}).get('evidence_basis')
    if location.get('status') == 'answered' and basis in {'source_value', 'source_calculation'}:
        return 'verified'
    if location.get('status') == 'estimated' and basis == 'estimate':
        return 'inputs_verified'
    return 'verified' if _source_verified(location, checks) else None


def _source_verified(location, checks):
    base = {key: location[key] for key in ('run_id', 'row', 'column', 'output_hash')}
    matching = [
        check for check in checks
        if check.get('status') == 'verified'
        and _location_matches(check.get('location'), base)
    ]
    if not matching and location.get('column') is not None:
        row_base = {**base, 'column': None}
        matching = [
            check for check in checks
            if check.get('status') == 'verified'
            and _location_matches(check.get('location'), row_base)
        ]
    if not matching:
        return False
    if any(not _has_paging_metadata(check) for check in matching):
        return True
    totals = {check.get('total') for check in matching if isinstance(check.get('total'), int)}
    if len(totals) != 1:
        return False
    total = totals.pop()
    if total < 0:
        return False
    intervals = sorted(
        (check.get('offset'), check.get('offset') + check.get('read_count'))
        for check in matching
        if (
            isinstance(check.get('offset'), int)
            and isinstance(check.get('read_count'), int)
            and check.get('offset') >= 0
            and check.get('read_count') >= 0
        )
    )
    cursor = 0
    for start, end in intervals:
        if start > cursor:
            return False
        cursor = max(cursor, end)
        if cursor >= total:
            break
    return cursor >= total and any(check.get('has_more') is False for check in matching)


def _has_paging_metadata(check):
    return all(key in check for key in ('total', 'offset', 'read_count', 'has_more'))


def _location_matches(candidate, expected):
    return isinstance(candidate, dict) and all(candidate.get(key) == value for key, value in expected.items())


NOT_COUNTED_TEXT = {
    'script_output_without_recipe': '只有脚本输出，没有通过的来源复算规则',
    'script_closure_not_counted': '闭合由脚本自行算出，只作展示，不计为引擎闭合',
    'engine_output_without_closure': '财务引擎的这项输出没有闭合规则，也没有通过的来源复算规则',
    'aggregate_output_without_closure': ('服务端汇总（structured_aggregate）的这项输出不是闭合结果，也没有通过的来源复算规则；'
                                         '收入、成本、费用、利润和量价拆分要引擎闭合或来源复算'),
    'recipe_failed': '复算规则未通过',
    'recipe_failed_precision': ('复算规则因运行值的精度损失（被舍入，或脚本用二进制浮点数计算）而未通过，'
                                '需用精确小数计算并在运行结果中保存完整精度'),
    'recipe_coefficient_unbound': ('复算规则的 scale 不是 1 或单位换算，只证明了乘法'
                                   '（系数要写成一条 ratio 规则，由 factor_recipe_id 引用，且只对 grain_basis 声明为估算的组合计为估算）'),
    'ratio_recipe_cannot_prove_amount': 'ratio 规则不能证明金额',
    'recipe_source_outside_bindings': '复算规则的来源不在本要求绑定的资料快照内',
    'estimate_cited_for_source_grain': '该组合未在 grain_basis 声明为估算，却引用了估算值',
    'estimate_factor_declared_not_reproducible': '该组合已声明系数无法复算，引用的估算不计入，等经理决定口径',
    'run_bindings_differ': '运行绑定的资料与本要求的绑定不一致，这个运行的结果不计入',
    'evidence_incomplete': '证据缺项、口径或单位不一致，或复算不闭合',
    'method_not_applicable': '计算方法未定义该项',
    'no_evidence_record': '该数值没有对应的证据记录',
}
BLOCKER_TEXT = {
    'structure_pending_confirmation': '工作簿结构待经理确认',
    'no_finance_context': '尚无覆盖这些资料的财务口径',
    'no_workbook_source': '绑定资料中没有可复算的归档工作簿',
}


def _grain_example(grain):
    return '、'.join(str(grain[key]) for key in GRAIN_KEYS if grain.get(key) not in (None, ''))


def _grain_count(grains, verb):
    example = _grain_example(grains[0])
    if not example:
        return f'{verb}（本项不分组织、期间和情景）'
    return f'{len(grains)} 个要求组合{verb}（如 {example}）'


def _path_text(path, keys):
    name = path['path']
    if name == 'evidence_recipe':
        return (f'在文稿中引用含 {keys} 的结果行，并为每个引用数值附一条通过的 evidence_recipe'
                '（direct、sum 或 diff，scale 只能是 1 或千、万、百万、亿的单位换算，来源须是本要求绑定的资料快照）')
    if name == 'engine_closure':
        return '引用 explain_variance 在财务口径上生成、闭合通过的差异结果（contribution 等列）'
    if name == 'profit_bridge':
        return ('引用 explain_variance 按 revenue_minus_cost_minus_expense 生成的利润桥，'
                '收入、成本、费用、利润四行的 contribution 都要引用')
    if name == 'price_quantity_bridge':
        method = path.get('method') or ' 或 '.join(path.get('methods') or [])
        return (f'引用 explain_variance 按 {method} 生成的同一行 total_variance、'
                'quantity_contribution、price_contribution 和 interaction')
    if name == 'method_run':
        return f'引用 method 为 {path["method"]} 的运行产出的含 {keys} 的结果行'
    if name == 'run_metric_result_set':
        return (f'在财务口径上运行 run_metric（{path["metric_kind"]}），引用范围内全部结果行的 amount，'
                '并逐行用 drilldown_fact 核查来源')
    if name == 'run_metric_account_total':
        return (f'在财务口径上运行 run_metric（kind 为 account_total，accounts 只列“{path["account"]}”），'
                '并引用要求期间那一行的 amount')
    if name == 'aggregate_output':
        return (f'用 analysis-save 保存 engine 为 structured_aggregate 的汇总计划（只绑定本要求的那份资料快照，'
                f'度量列命名为 {path["measure"]}），用 analysis-run 运行后引用该列的汇总结果，再用 drilldown_fact 核查来源')
    if name == 'engine_output':
        kind = f'（{path["scenario_kind"]}）' if path.get('scenario_kind') else ''
        return f'引用 {path["tool"]}{kind} 在财务口径上产出的 {path["column"]} 结果'
    return f'需要含 {keys} 的结果行，并在文稿中引用该数值'


def _unanswered_message(item):
    """What is missing and which evidence paths can close it; names only paths the checker counts."""
    status = item.get('status')
    label = item.get('label') or item.get('metric') or item.get('id')
    if status == 'estimated':
        message = f'“{label}”含估算组合，经理确认前不计为已回答。'
    else:
        message = f'尚未提供“{label}”的计算依据。'
    hierarchy = item.get('hierarchy_dimensions') or {}
    if hierarchy:
        named = '、'.join(f'{key}={value}' for key, value in hierarchy.items())
        message += (f'要求把 {named} 写成了维度；下钻层级不是结果行字段，在行上补这个字段不计入。'
                    '请用 parent_id 重新保存要求：去掉该维度，把下钻对象逐一列入 organizations 或另立要求项，'
                    '逐格来源核查用 requires_drilldown。')
    missing = item.get('missing_grains') or []
    uncounted = item.get('uncounted_grains') or []
    if missing:
        message += _grain_count(missing, '没有已引用的结果行') + '。'
    if uncounted:
        reasons = dict.fromkeys(reason for grain in uncounted for reason in grain.get('reasons', []))
        statuses = dict.fromkeys(status for grain in uncounted for status in grain.get('statuses', []))
        why = '、'.join(NOT_COUNTED_TEXT.get(reason, reason) for reason in reasons) or '、'.join(statuses)
        message += _grain_count(uncounted, '已引用了结果但未计入') + f'，状态为 {"、".join(statuses)}：{why}。'
    estimated = item.get('estimated_grains') or []
    if estimated:
        message += _grain_count(estimated, '为估算') + '，按声明的估算方法由系数规则复算，始终标注“估算”。'
    for grain in item.get('declared_grains') or []:
        example = _grain_example(grain) or '本项'
        if grain['state'] == 'estimate_factor_not_reproducible':
            message += f'{example} 声明为估算，但系数无法由归档单元格复算（{grain.get("reason")}），需经理决定口径；不得换用其他系数凑数。'
        else:
            message += f'{example} 声明为资料缺失（{grain.get("reason")}），需经理确认后才记为缺数。'
    paths = item.get('evidence_paths') or {'accepted': [], 'blocked': []}
    if missing or uncounted:
        grain = (missing or uncounted)[0]
        fields = [key for key in GRAIN_KEYS if key in grain]
        dimension_keys = [key for key in (item.get('dimensions') or {})
                          if key not in {'method', *INTENT_DIMENSION_KEYS} and key not in hierarchy]
        pinned = (item.get('dimensions') or {}).get('metric_key')
        metric_key = f'metric_key（{pinned if pinned not in (None, "") else _canonical_metric_name(item.get("metric"))}）'
        named = list(dict.fromkeys([metric_key, *fields, *[key for key in dimension_keys if key != 'metric_key']]))
        keys = '、'.join(named) + ' 和 value'
        if paths['accepted']:
            message += '可计入的依据：' + '；或'.join(_path_text(path, keys) for path in paths['accepted']) + '。'
        if _closed_metric_key(item) is not None and any(path['path'] == 'evidence_recipe' for path in paths['accepted']):
            message += '金额类要求（比率、数量、人数等以外的一切指标）的脚本输出和脚本自行算出的闭合本身都不计为已回答。'
        if item.get('required_bindings'):
            named = '、'.join(f'{alias}={dataset_id}' for alias, dataset_id in item['required_bindings'].items())
            message += (f'要让这些结果计入，运行须按本要求的绑定重新运行（{named}）：别名与快照都要一致，'
                        '要求只有一个绑定时别名不限。')
        estimates = [grain for grain in [*missing, *uncounted] if _reproducible_estimate_declared(item, grain)]
        if estimates and paths.get('estimate'):
            message += (_grain_count(estimates, '已在 grain_basis 声明为估算')
                        + '：引用该格的结果行，附一条 direct 或 sum 规则（scale 只能是 1 或单位换算），并用 factor_recipe_id '
                        '引用本批次一条通过的 ratio 规则作系数，系数不写进 scale；这些组合只计为估算，经理确认前不计为已回答。')
        if any(_needs_grain_basis_hint(item, grain) for grain in [*missing, *uncounted]):
            message += ('若绑定资料里确实没有某个组合的数：用 parent_id 重新保存要求，在该项 grain_basis 里声明这个组合。'
                        '系数能由归档单元格复算时声明 estimate 并写 method，再按估算附规则；系数写不出时声明 estimate、'
                        'factor_reproducible:false 并写 reason；资料缺失时声明 unavailable 并写 reason。'
                        '声明只让该组合计为估算或交经理确认，不计为已回答。')
    elif paths['accepted'] and status != 'answered' and not estimated and not item.get('declared_grains'):
        # A result set has no per-grain detail; still name what answers it.
        message += '可计入的依据：' + '；或'.join(_path_text(path, 'value') for path in paths['accepted']) + '。'
    blocked = paths.get('blocked') or []
    if blocked and status != 'answered':
        texts = [f'{path.get("tool") or path["path"]}（{"、".join(BLOCKER_TEXT.get(b, b) for b in path["blockers"])}）'
                 for path in blocked]
        message += '当前走不通的路径：' + '；'.join(texts) + '。'
        blockers = {blocker for path in blocked for blocker in path['blockers']}
        if not paths['accepted'] and item.get('awaiting_manager'):
            # Only when nothing is left for the assistant (see _manager_reasons); completion stops too.
            message += '需经理先处理上述阻塞，助手不要重复计算。'
        elif not paths['accepted'] and 'no_finance_context' in blockers:
            message += '可先用 finance_context_save 为绑定资料保存财务口径，再按被阻塞的引擎路径计算。'
    gaps = item.get('row_field_gaps') or {}
    if gaps:
        fields = '、'.join(f'{key}={value}' for key, value in gaps.items())
        message += (f'已引用的行还缺少或不一致的字段：{fields}。这些字段写在每一条结果行上；'
                    '脚本和 SQL 运行没有 finance_scope，运行的 scope 设置不参与要求核对。')
    if item.get('ledger_identity') == 'unproven_without_bound_snapshots':
        message += '没有财务口径的运行只能用本要求绑定快照的哈希证明账套，本要求没有记录绑定快照，行上的账套文字不计入。'
    return message


def _estimated_issues(item):
    """One requirement_estimated issue per estimated cell, with its source and factor ranges."""
    issues = []
    for location in item.get('locations', []):
        if location.get('status') != 'estimated':
            continue
        evidence = location.get('evidence') or {}
        grain = next((grain for grain in item.get('estimated_grains', [])
                      if {key: location[key] for key in ('run_id', 'row', 'column')} in grain.get('locations', [])), {})
        if grain.get('accepted'):
            continue
        factor = evidence.get('factor') or {}
        formula = _estimate_formula(evidence)
        issues.append({
            'code': 'requirement_estimated',
            'severity': 'warning',
            'requirement_item_id': item['id'],
            'business_key': {'metric': item.get('metric'),
                             **{key: grain[key] for key in GRAIN_KEYS if key in grain}},
            'message': (f'“{item.get("label") or item.get("metric")}”'
                        + (f'的 {_grain_example(grain)}' if _grain_example(grain) else '')
                        + f'为估算：{formula}。经理确认前不计为已回答，报告中始终标注“估算”。'),
            'location': {'reference': _plain_location(evidence.get('recipe_location')),
                         **({'factor': factor.get('location')} if factor.get('location') else {})},
            'recipe_ids': evidence.get('recipe_ids', []),
            'recipe_hash': evidence.get('recipe_hash'),
            **({'factor_recipe_hash': factor.get('recipe_hash')} if factor.get('recipe_hash') else {}),
            'accepted': bool(grain.get('accepted')),
            'run_id': location['run_id'], 'row': location['row'], 'column': location['column'],
        })
    return issues


def _plain_location(location):
    if not isinstance(location, dict):
        return None
    return {key: location[key] for key in ('dataset_id', 'source_hash', 'sheet', 'range') if key in location}


def _range_text(location):
    if not isinstance(location, dict):
        return ''
    if 'left' in location and 'right' in location:
        return f'{_range_text(location["left"])} ÷ {_range_text(location["right"])}'
    return f'{location.get("sheet")}!{location.get("range")}'


def _estimate_formula(evidence):
    dependent = _range_text(evidence.get('recipe_location'))
    factor = _range_text((evidence.get('factor') or {}).get('location'))
    return f'{dependent} × ({factor})' if factor else dependent


def augment_quality(analysis, spec, quality, *, _recipe_report=None):
    requirements = check_requirements(analysis, spec, _recipe_report=_recipe_report)
    from foundry.finance_requirement_support import with_input_gaps
    requirements = with_input_gaps(analysis, requirements)
    issues = list(quality.get('issues', []))
    if requirements['scope_changed']:
        issues.append({'code': 'requirements_changed', 'severity': 'warning', 'message': '分析范围已更新，请按当前要求复核文稿。'})
    for item in requirements['items']:
        if item['status'] == 'estimated':
            issues.extend(_estimated_issues(item))
        elif item['status'] != 'answered':
            issues.append({'code': 'requirement_unanswered', 'severity': 'warning', 'message': _unanswered_message(item), 'requirement_item_id': item['id'],
                           **({'fixups': item['fixups']} if item.get('fixups') else {})})
            issues.extend(_estimated_issues(item))
        if item['provenance_status'] in {'pending', 'failed'}:
            source_fixups = [fixup for fixup in item.get('fixups', []) if fixup.get('code') in SOURCE_FIXUP_CODES]
            issues.append({'code': 'requested_source_missing', 'severity': 'warning', 'message': f'尚未完成“{item["label"]}”的指定来源核查。', 'requirement_item_id': item['id'],
                           **({'fixups': source_fixups} if source_fixups else {})})
    return {**quality, 'issues': issues, 'status': 'needs_attention' if issues else quality['status'],
            'requirements': requirements,
            'checks': {'numbers': {'status': quality['status']}, 'coverage': requirements['coverage'], 'provenance': requirements['provenance']}}
